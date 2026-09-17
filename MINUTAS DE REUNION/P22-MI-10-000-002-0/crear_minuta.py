#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera la minuta P22-MI-10-000-002-0
(Minutes of Meeting - Follow-up Meeting of 17 September 2026)
usando el template-adasa. Idioma ingles. Destinatario: BW Water.

Codigo:  P22-MI-10-000-002-0  (correlativo siguiente a la minuta del 16-Sep)
Fecha:   17-Sep-2026 (jueves)
Rev:     0

FUENTE: transcripcion Plaud "2026-09-17 MINUTA BW WATERS"
(file of_6f0d85a3d29794045a7e1ee98891f5d2, 30 min, 12:31 UTC = 09:31 Chile).
El PDF `REUNION 17-09-26.pdf` es el resumen automatico de esa grabacion; su
extraccion esta en `md/REUNION 17-09-26_extracted.md`.

REGLAS DE CONTENIDO (heredadas de la minuta P22-MI-10-000-001-0):
  - Solo lo que se dijo en la reunion; donde el resumen y la transcripcion
    difieren, manda la transcripcion:
      * "jerk" = yoke; "Chilean PA" = Chilean Professional Engineer.
      * "Buroveritas" = Bureau Veritas; "Antofagasta plans three D" = AutoCAD
        Plant 3D.
      * El resumen afirma que la causa es un error de fabricacion de dos
        spools. En la transcripcion BW Water lo plantea como probable ("maybe
        the length of the pipe") y ADASA dice que la causa esta por
        establecerse: se registra asi.
      * El resumen no registra que la explicacion del 16-Sep (turbo distinto a
        los planos Fedco) quedo descartada por el propio BW Water: se registra.
  - Sin nombres internos de BW Water (Adni, Lockman, Iqbal, Nick, LG): se
    usan "the Penang team" y "the engineering team".
  - Sin seccion DOCUMENT HISTORY, sin TOC (igual que la minuta del 16-Sep).

Path al skill: absoluto via ~/.claude/skills/template-adasa segun CLAUDE.md Seccion 3.
NOTA: el template numera H1/H2 automaticamente. NO escribir numeros en add_heading.
"""

import os
import sys

SKILL_PATH = os.path.expanduser("~/.claude/skills/template-adasa")
sys.path.insert(0, SKILL_PATH)
# docx_metadata vive en la carpeta compartida de skills, no en template-adasa
sys.path.insert(0, os.path.expanduser("~/.claude/skills/_shared"))

from ejemplo_documento import (  # noqa: E402
    crear_documento_adasa,
    add_simple_table,
)
import docx_metadata  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(
    SCRIPT_DIR,
    "P22-MI-10-000-002-0_Minutes-Follow-up-Meeting-17Sep2026_ADASA.docx",
)


def add_para(doc, text, size=11):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    return para


def ajustar_tabla(table, anchos=None, mantener_junta=False):
    """Celdas alineadas a la izquierda (el justificado abre huecos en columnas
    angostas) y filas que no se parten entre paginas. Si se dan anchos (pulgadas),
    se escriben en el tblGrid y en cada celda: Word obedece al tblGrid."""
    if anchos:
        for col, ancho in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), anchos):
            col.set(qn("w:w"), str(int(ancho * 1440)))
        for row in table.rows:
            for cell, ancho in zip(row.cells, anchos):
                cell.width = Inches(ancho)
    for i, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                # tabla corta: todas las filas salvo la ultima quedan unidas
                # a la siguiente para que no se parta entre paginas
                if mantener_junta and i < len(table.rows) - 1:
                    para.paragraph_format.keep_with_next = True
                # el encabezado de tabla no queda solo al pie de pagina
                if i == 0:
                    para.paragraph_format.keep_with_next = True


def crear_documento():
    crear_documento_adasa(
        titulo="MINUTES OF MEETING — FOLLOW-UP MEETING OF 17 SEPTEMBER 2026",
        codigo="P22-MI-10-000-002-0",
        preparado_por="Luis Rivera",
        revisado_por="Luis Rivera",
        aprobado_por="Victor Gutierrez",
        nombre_planta="TALTAL",
        cliente="ADASA",
        output_filename=OUTPUT,
        incluir_toc=False,
    )

    doc = Document(OUTPUT)

    # ---- limpiar placeholder ------------------------------------------------
    elementos_a_eliminar = []
    encontrado = False
    for para in doc.paragraphs:
        if para.style and para.style.name == "Heading 1" and not encontrado:
            encontrado = True
        if encontrado:
            elementos_a_eliminar.append(para)
    for para in elementos_a_eliminar:
        p = para._element
        p.getparent().remove(p)

    # =========================================================================
    # MEETING INFORMATION
    # =========================================================================
    doc.add_heading("MEETING INFORMATION", level=1)
    ajustar_tabla(add_simple_table(doc, [
        ("Item", "Detail"),
        ("Project", "Second Stage RO Brine Module, Taltal (Contract C-4300)"),
        ("Date and time", "Thursday 17 September 2026, 09:30 Chile time"),
        ("Format", "Video call, approximately 30 minutes"),
        ("ADASA", "Luis Rivera, Victor Gutierrez"),
        ("BW Water", "Eduardo Yamauchi and the Penang team"),
        ("Remarks", "Follow-up to the weekly coordination meeting of "
         "16 September 2026 (P22-MI-10-000-001-0)"),
        ("Next meeting", "Weekly coordination meeting, Wednesday "
         "23 September 2026"),
    ]))

    # =========================================================================
    # DISCUSSION AND AGREEMENTS
    # =========================================================================
    doc.add_heading("DISCUSSION AND AGREEMENTS", level=1)

    # ---- 2.1 ----------------------------------------------------------------
    doc.add_heading("ADASA Point of Contact", level=2)
    add_para(
        doc,
        "Luis Rivera will be out of the office for the next two weeks. During "
        "that period Victor Gutierrez will be in charge of ADASA's communication "
        "with BW Water and the Penang team and will take the decisions, in "
        "particular those concerning the shipping procedure. Luis Rivera will "
        "remain in copy. Victor Gutierrez asked BW Water to send him all project "
        "information and to keep the scheduled meetings.",
    )

    # ---- 2.2 ----------------------------------------------------------------
    doc.add_heading("Factory Acceptance Test Procedure", level=2)
    add_para(
        doc,
        "BW Water stated that the individual procedures exist but have not yet "
        "been compiled into the single document requested by ADASA. ADASA noted "
        "that the Factory Acceptance Test procedure will also serve as the "
        "baseline for the Site Installation Test procedure that BW Water must "
        "submit for the electrical reconnection on site, and that ADASA needs to "
        "share it with Bureau Veritas for review, even as a preliminary issue.",
    )
    add_para(
        doc,
        "On 17 September the Penang team prepared a Factory Acceptance Test "
        "procedure based on Section 7 of the Inspection and Test Plan, which is "
        "now under BW Water internal review. Once that review is complete, "
        "BW Water expects to submit the procedure to ADASA on 18 September 2026.",
    )

    # ---- 2.3 ----------------------------------------------------------------
    doc.add_heading("Review of the Lifting Package by the Chilean Professional "
                    "Engineer", level=2)
    add_para(
        doc,
        "BW Water has issued the purchase order for the Chilean Professional "
        "Engineer's review of the lifting procedure and the yoke drawings. "
        "Confirmation is still pending, and BW Water will send a reminder. ADASA "
        "will also make contact to expedite the reply, noting that 18 and "
        "19 September are public holidays in Chile.",
    )
    add_para(
        doc,
        "The lifting calculations and the yoke drawings are being prepared by "
        "BW Water's structural engineering team, which aims to complete them on "
        "18 September 2026. The Chilean Professional Engineer is engaged only to "
        "review them.",
    )

    # ---- 2.4 ----------------------------------------------------------------
    doc.add_heading("Turbocharger Piping Misalignment", level=2)
    add_para(
        doc,
        "On 17 September the Penang team assembled the piping and aligned it "
        "with the turbochargers. On one turbocharger the piping was aligned by "
        "slightly adjusting the base plate thickness, within the adjustment "
        "allowed by the flexible coupling. On the other, two spools at the top "
        "connection clash with the turbocharger. BW Water showed photographs of "
        "the clash during the meeting.",
    )
    add_para(
        doc,
        "BW Water stated that the turbocharger matches the 3D model, that its "
        "main dimensions conform to the drawing, and that the turbochargers are "
        "the units considered in the design. At the meeting of 16 September the "
        "issue had been attributed to a difference between the delivered "
        "turbocharger and the Fedco drawings; that explanation was not "
        "confirmed. The dimensional difference reported for a valve concerns "
        "only its actuator and does not change the flange-to-flange distance. "
        "According to BW Water, the resulting interference on the top pipe can be "
        "avoided without hot work.",
    )
    add_para(
        doc,
        "BW Water stated that the spools were fabricated according to the "
        "drawing and attributed the misalignment, without confirming it, to the "
        "length and inclination of the two spools. ADASA stated that the root "
        "cause is still to be established.",
    )
    add_para(
        doc,
        "One of the two spools has not yet been hydrostatically tested and can "
        "still be adjusted before full welding. The other has already been cut at "
        "the connection beyond the tie-in point and coupling assembly, and has "
        "not been rewelded. According to the Bureau Veritas report received by "
        "ADASA, the hydrostatic test scheduled for 16 September was cancelled, "
        "and BW Water confirmed that the cancellation was due to this spool.",
    )
    add_para(
        doc,
        "ADASA set two priorities: confirmation that the turbocharger complies "
        "with the specification approved in ADASA's review, which Bureau Veritas "
        "will also check, and the rescheduling of the high pressure test for the "
        "whole line once the adjustment is made. ADASA also asked BW Water to "
        "confirm how many lines are affected. Two lines were mentioned on "
        "16 September, while at this meeting the Penang team referred to a single "
        "line, the same one whose pressure test is pending.",
    )

    # ---- 2.5 ----------------------------------------------------------------
    doc.add_heading("Rework Plan and Amended Drawings", level=2)
    add_para(
        doc,
        "The engineering team will amend the drawings of the two spools using "
        "the dimensions measured on 17 September, and BW Water will issue them on "
        "18 September 2026. BW Water presented the following plan:",
    ).paragraph_format.keep_with_next = True
    ajustar_tabla(add_simple_table(doc, [
        ("Date", "Activity"),
        ("Friday 18-Sep-2026", "Issue of the amended drawings; cutting and "
         "preparation of the spools"),
        ("Saturday 19-Sep-2026", "Fit-up and root weld"),
        ("Monday 21-Sep-2026", "Penetrant test by a third party, followed by "
         "full welding"),
        ("Tuesday 22-Sep-2026", "Hydrostatic test"),
        ("Wednesday 23-Sep-2026", "Reassembly and fit-up"),
    ]), anchos=[2.0, 4.5], mantener_junta=True)
    add_para(
        doc,
        "BW Water asked whether ADASA must approve the amended drawings before "
        "root welding, since waiting for that approval, with the public holidays "
        "in Chile, would delay the work by two to three days. It was agreed that "
        "the rework will proceed in parallel with ADASA's review, and that "
        "BW Water will send the amended drawings as soon as they are issued. "
        "ADASA needs them to identify the error that caused the misalignment and "
        "to share them immediately with Bureau Veritas, so that the inspection "
        "is carried out against the drawings actually issued.",
    )
    add_para(
        doc,
        "ADASA asked whether the modification will be made in the AutoCAD Plant "
        "3D model or directly in the CAD drawings. The question was not answered "
        "during the meeting.",
    )

    # ---- 2.6 ----------------------------------------------------------------
    doc.add_heading("Bureau Veritas Inspection", level=2)
    add_para(
        doc,
        "Bureau Veritas will visit the Penang workshop on 18 September 2026 to "
        "inspect the misalignment and report to ADASA. ADASA asked BW Water to "
        "give the inspectors access to all information and to any measurements "
        "they require. Victor Gutierrez asked BW Water to keep supporting Bureau "
        "Veritas with the information on these modifications over the coming "
        "weeks, so that it can issue a complete report.",
    )

    # ---- 2.7 ----------------------------------------------------------------
    doc.add_heading("Schedule and Shipment", level=2)
    add_para(
        doc,
        "ADASA stated that the Factory Acceptance Test must be rescheduled and "
        "restated that it is waiting for the pending engineering documentation, "
        "asking BW Water to close all pending engineering milestones. BW Water "
        "will prepare the updated schedule once the piping issue is resolved, "
        "which it expects on Wednesday 23 or Thursday 24 September 2026. The "
        "latest date estimated by BW Water was 10 October 2026, and BW Water is "
        "working to recover time.",
    )
    add_para(
        doc,
        "The shipment of the plant is on hold until ADASA receives the updated "
        "schedule. The freight forwarder requires about one week's notice of the "
        "date on which the plant will be ready to ship, so that coordination has "
        "to run in parallel with the remaining work. BW Water noted that the "
        "results of the tests in this final phase may affect the shipment date.",
    )

    # ---- 2.8 ----------------------------------------------------------------
    doc.add_heading("Shipment of the Membranes", level=2)
    add_para(
        doc,
        "During the week of 21 September 2026, Victor Gutierrez will send "
        "BW Water the information for coordinating the shipment of the "
        "membranes. ADASA is ready to start that coordination. BW Water will "
        "forward the notification internally and prepare the shipment.",
    )

    # =========================================================================
    # ACTION ITEMS
    # =========================================================================
    doc.add_heading("ACTION ITEMS", level=1)
    ajustar_tabla(add_simple_table(doc, [
        ("No.", "Action", "Responsible", "Due date"),
        ("1", "Submit the Factory Acceptance Test procedure to ADASA for review, "
         "once BW Water internal review is complete", "BW Water", "18-Sep-2026"),
        ("2", "Share the Factory Acceptance Test procedure with Bureau Veritas",
         "ADASA", "Upon receipt"),
        ("3", "Issue the amended drawings of the two affected spools, based on "
         "the dimensions measured on 17 September", "BW Water", "18-Sep-2026"),
        ("4", "Share the amended drawings with Bureau Veritas", "ADASA",
         "Upon receipt"),
        ("5", "State whether the modification is made in the AutoCAD Plant 3D "
         "model or directly in the CAD drawings", "BW Water", "Not stated"),
        ("6", "Confirm the number of lines affected", "BW Water", "Not stated"),
        ("7", "Give Bureau Veritas access to all information and measurements, "
         "starting with the visit of 18 September", "BW Water",
         "From 18-Sep-2026"),
        ("8", "Complete the rework per Section 2.5 - Rework Plan and Amended "
         "Drawings, including the high pressure test of the whole affected line",
         "BW Water", "22-Sep-2026 (test), 23-Sep-2026 (reassembly)"),
        ("9", "Obtain the Chilean Professional Engineer's confirmation of the "
         "purchase order for the review of the lifting procedure and the yoke "
         "drawings", "BW Water", "Not stated"),
        ("10", "Complete the lifting calculations and the yoke drawings",
         "BW Water", "18-Sep-2026"),
        ("11", "Issue the updated schedule, including the new Factory Acceptance "
         "Test date", "BW Water", "Once the piping issue is resolved "
         "(23 or 24-Sep-2026)"),
        ("12", "Send the information for coordinating the shipment of the "
         "membranes", "ADASA (Victor Gutierrez)", "Week of 21-Sep-2026"),
        ("13", "Send all project information to Victor Gutierrez, with Luis "
         "Rivera in copy", "BW Water", "During Luis Rivera's absence"),
    ]), anchos=[0.45, 3.2, 1.25, 1.6])

    # los titulos del template no traen "mantener con el siguiente" efectivo:
    # se fija explicito para que ningun titulo quede solo al pie de pagina
    for para in doc.paragraphs:
        if para.style is not None and para.style.name.startswith("Heading"):
            para.paragraph_format.keep_with_next = True

    # add_simple_table deja un parrafo vacio tras la tabla; al final del documento
    # sobra y puede abrir una pagina en blanco
    ultimo = doc.paragraphs[-1]
    if not ultimo.text.strip():
        ultimo._element.getparent().remove(ultimo._element)

    # ---- metadatos limpios (global sec. 2.3) --------------------------------
    docx_metadata.apply_core_properties(
        doc,
        title="Minutes of Meeting - Follow-up Meeting of 17 September 2026",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="Second Stage RO Brine Module - Taltal",
        comments="Aguas de Antofagasta S.A.",
        language="en-US",
    )
    doc.save(OUTPUT)
    docx_metadata.fix_app_xml(OUTPUT, company="Aguas de Antofagasta S.A.")
    print(f"Documento generado: {OUTPUT}")


if __name__ == "__main__":
    crear_documento()
