#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extrae el tracker semanal de procurement BW Water (.xlsx) a tabla Markdown,
preservando todas las filas y columnas relevantes para auditoria. Reutilizable
para futuras semanas: pasa la ruta del .xlsx como argumento.

Uso:
    python extraer_tracker_xlsx.py "<ruta_xlsx>" "<ruta_md_salida>"

Asume la hoja "Weekly Dashboard (2)" con las columnas estandar BW Water.
"""

from __future__ import annotations

import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl


def fmt(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    return text.replace("\n", " / ").replace("|", "/")


def extract_to_md(xlsx_path: Path, md_path: Path) -> None:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "Weekly Dashboard (2)" not in wb.sheetnames:
        raise SystemExit(f"Hoja 'Weekly Dashboard (2)' no encontrada en {xlsx_path}")
    ws = wb["Weekly Dashboard (2)"]
    rows = list(ws.iter_rows(values_only=True))

    title = fmt(rows[0][0]) if rows else ""
    subtitle = fmt(rows[1][0]) if len(rows) > 1 else ""
    header = [fmt(c) for c in rows[2]] if len(rows) > 2 else []
    body = []
    for row in rows[3:]:
        cells = [fmt(c) for c in row]
        if any(cells):
            body.append(cells)

    out = []
    out.append(f"# Tracker Procurement BW Water — extraccion automatica")
    out.append("")
    out.append(f"**Archivo origen:** `{xlsx_path.name}`")
    out.append(f"**Hoja:** Weekly Dashboard (2)")
    out.append(f"**Titulo del tracker:** {title}")
    out.append(f"**Cabecera:** {subtitle}")
    out.append("")
    out.append("## Tabla principal")
    out.append("")
    out.append("| " + " | ".join(header) + " |")
    out.append("|" + "|".join(["---"] * len(header)) + "|")
    for cells in body:
        while len(cells) < len(header):
            cells.append("")
        cells = cells[: len(header)]
        out.append("| " + " | ".join(cells) + " |")
    out.append("")
    out.append("## Notas de extraccion")
    out.append("")
    out.append(
        "Extraccion automatica via openpyxl preservando fechas en formato ISO "
        "(AAAA-MM-DD). Filas vacias omitidas. Caracteres `|` y saltos de linea "
        "internos a celdas reemplazados por `/`. Status legend (filas finales del "
        "Excel) se preserva como parte de la tabla."
    )

    md_path.write_text("\n".join(out), encoding="utf-8")
    print(f"OK — generado: {md_path}  ({len(body)} filas con datos)")


def main(argv: list[str]) -> None:
    if len(argv) != 3:
        raise SystemExit(
            "Uso: python extraer_tracker_xlsx.py <ruta_xlsx> <ruta_md_salida>"
        )
    xlsx_path = Path(argv[1])
    md_path = Path(argv[2])
    if not xlsx_path.exists():
        raise SystemExit(f"Archivo no encontrado: {xlsx_path}")
    extract_to_md(xlsx_path, md_path)


if __name__ == "__main__":
    main(sys.argv)
