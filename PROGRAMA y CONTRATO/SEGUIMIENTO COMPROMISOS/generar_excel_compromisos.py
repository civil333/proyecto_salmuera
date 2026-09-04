#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generar_excel_compromisos.py — Genera el Registro de Compromisos P22-IT-06-000-006-0.

FUENTE UNICA: `compromisos.yaml` + `hitos.yaml`. El .xlsx es DERIVADO y no se
edita a mano: cualquier cambio manual se pierde en la siguiente corrida.

Principio de diseno que gobierna todo el archivo:

    lo que cambia por un EVENTO va como VALOR
    lo que cambia por el PASO DEL TIEMPO va como FORMULA

Por eso el semaforo y los dias de mora son formulas (`TODAY()`) y el color lo
pone el formato condicional. Un relleno fijado en generacion es correcto el dia
que se corre el script y falso al siguiente: es el modo de falla silencioso de
estos registros.

Reglas duras de openpyxl que este script respeta:
  - Ningun string de celda empieza con "=" (se escribiria como formula y
    LibreOffice renderiza Err:501). El helper `texto()` fuerza data_type='s'.
  - `number_format` en convencion INGLESA ('dd-mmm-yyyy', '#,##0'). Excel y
    LibreOffice traducen al locale chileno al abrir.
  - Formulas en ingles: COUNTIFS, TODAY, IF.
  - Fechas como `datetime.date`, nunca como string ya formateado.
  - En `FormulaRule` la formula va SIN "=" inicial y anclada a `$G2`
    (columna absoluta, fila relativa) o se pinta toda la tabla del color
    de la fila 2.

Uso:
    python generar_excel_compromisos.py
    python generar_excel_compromisos.py --sin-backup
"""
import argparse
import datetime as dt
import os
import shutil
import sys

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BASE = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(BASE, "P22-IT-06-000-006-0_Registro-de-Compromisos.xlsx")
BACKUPS = os.path.join(BASE, "_backups")

# --- vocabularios controlados: si un valor no pertenece, se ABORTA ---
VOC = {
    "frente": {"BUREAU VERITAS", "PROGRAMA", "COMERCIAL", "INTERNO ADASA"},
    "obligado": {"BW WATER", "ADASA", "BUREAU VERITAS", "FEDCO"},
    "beneficiario": {"BW WATER", "ADASA", "BUREAU VERITAS", "FEDCO"},
    "estado": {"ABIERTO", "CERRADO", "CERRADO PARCIAL", "EN DISPUTA"},
    "criticidad": {"CRITICA", "ALTA", "MEDIA", "BAJA"},
    "origen_fecha": {"CONTRACTUAL", "COMPROMISO ESCRITO BW",
                     "COMPROMISO VERBAL-MINUTA", "FIJADA ADASA"},
}
VOC_HITOS = {
    "firmeza": {"CONTRACTUAL FIRME", "REFERENCIA FIJA ADASA",
                "DECLARADO PROVEEDOR", "INDICATIVO"},
}

OBLIGATORIOS = ["id", "frente", "obligado", "beneficiario", "compromiso",
                "estado", "criticidad", "criterio_cierre", "fuente_contractual",
                "origen_fecha", "consecuencia", "accion_adasa",
                "evidencia_origen", "reprogramaciones", "historial_fechas"]

COLUMNAS = [
    ("ID", 9, "id"),
    ("Frente", 15, "frente"),
    ("Obligado", 15, "obligado"),
    ("Beneficiario", 14, "beneficiario"),
    ("Compromiso", 58, "compromiso"),
    ("Fecha comprometida", 15, "fecha_comprometida"),
    ("Semaforo", 13, None),                 # G - formula
    ("Estado", 16, "estado"),
    ("Dias vs. compromiso", 12, None),      # I - formula
    ("Criticidad", 11, "criticidad"),
    ("Criterio de cierre", 45, "criterio_cierre"),
    ("Fuente contractual", 45, "fuente_contractual"),
    ("Origen de la fecha", 22, "origen_fecha"),
    ("Fecha origen", 13, "fecha_origen"),
    ("Fecha de cierre", 13, "fecha_cierre"),
    ("Consecuencia", 50, "consecuencia"),
    ("Accion siguiente ADASA", 45, "accion_adasa"),
    ("Evidencia de origen", 45, "evidencia_origen"),
    ("Evidencia de cierre", 45, "evidencia_cierre"),
    ("Ref. cruzada", 16, "ref_cruzada"),
    ("Reprogramaciones", 11, "reprogramaciones"),
    ("Historial de fechas", 45, "historial_fechas"),
    ("Ultima actualizacion", 15, None),     # W - valor de meta.corte
]
COL_FECHA = {"fecha_comprometida", "fecha_origen", "fecha_cierre"}

# --- paleta ---
F_HEADER = PatternFill("solid", start_color="4472C4")
F_VENCIDO = PatternFill("solid", start_color="FFC7CE")
F_POR_VENCER = PatternFill("solid", start_color="FCD5B4")
F_EN_PLAZO = PatternFill("solid", start_color="C6EFCE")
F_SIN_FECHA = PatternFill("solid", start_color="FFE4B5")
F_CERRADO = PatternFill("solid", start_color="D9D9D9")
F_TITULO = PatternFill("solid", start_color="D9E1F2")

# Rellenos para FORMATO CONDICIONAL. No son los mismos objetos que los de arriba:
# en un estilo diferencial (dxf) Excel toma el color de `bgColor`, no de
# `start_color`/fgColor. Un PatternFill construido con start_color se guarda sin
# bgColor y la regla se aplica sin pintar nada — la celda queda blanca y el
# defecto solo se ve al renderizar, no al inspeccionar el archivo.
D_VENCIDO = PatternFill(bgColor="FFC7CE", fill_type="solid")
D_POR_VENCER = PatternFill(bgColor="FCD5B4", fill_type="solid")
D_EN_PLAZO = PatternFill(bgColor="C6EFCE", fill_type="solid")
D_SIN_FECHA = PatternFill(bgColor="FFE4B5", fill_type="solid")
D_CERRADO = PatternFill(bgColor="D9D9D9", fill_type="solid")

FT_HEADER = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FT_NORMAL = Font(name="Arial", size=9)
FT_BOLD = Font(name="Arial", size=9, bold=True)
FT_TITULO = Font(name="Arial", size=14, bold=True, color="1F3864")
FT_SECCION = Font(name="Arial", size=11, bold=True, color="1F3864")

BORDE = Border(left=Side(style="thin"), right=Side(style="thin"),
               top=Side(style="thin"), bottom=Side(style="thin"))

FMT_FECHA = "dd-mmm-yyyy"   # convencion INGLESA. Excel traduce al abrir.
FMT_ENTERO = "#,##0"


def texto(celda, valor):
    """Escribe un string forzando tipo texto.

    Un string que empieza con '=' se guardaria como formula y LibreOffice
    renderiza Err:501. Ocurre justo en la Leyenda, que es donde se explica
    la formula del semaforo.
    """
    celda.value = valor
    if isinstance(valor, str) and valor.startswith("="):
        celda.data_type = "s"
    return celda


def validar_fuente(datos, hitos):
    """Aborta la generacion ante cualquier defecto. No advierte: aborta."""
    errores = []

    if "meta" not in datos or "corte" not in datos.get("meta", {}):
        errores.append("falta meta.corte en compromisos.yaml")

    vistos = set()
    for i, c in enumerate(datos.get("compromisos", []), 1):
        cid = c.get("id", "<sin id>")
        for campo in OBLIGATORIOS:
            if campo not in c or c[campo] is None or c[campo] == "":
                if campo == "reprogramaciones" and c.get(campo) == 0:
                    continue
                errores.append("%s: falta el campo obligatorio '%s'" % (cid, campo))
        if cid in vistos:
            errores.append("%s: id duplicado" % cid)
        vistos.add(cid)
        for campo, permitidos in VOC.items():
            v = c.get(campo)
            if v is not None and v not in permitidos:
                errores.append("%s: %s='%s' fuera del vocabulario %s"
                               % (cid, campo, v, sorted(permitidos)))
        for campo in COL_FECHA:
            v = c.get(campo)
            if v is not None and not isinstance(v, dt.date):
                errores.append("%s: %s='%s' no es una fecha (usar ISO sin comillas)"
                               % (cid, campo, v))
        # Coherencia estado <-> fecha de cierre
        if c.get("estado") in ("CERRADO", "CERRADO PARCIAL") and not c.get("fecha_cierre"):
            errores.append("%s: estado '%s' sin fecha_cierre" % (cid, c["estado"]))
        if c.get("estado") == "ABIERTO" and c.get("fecha_cierre"):
            errores.append("%s: estado ABIERTO con fecha_cierre" % cid)
        if c.get("estado") in ("CERRADO", "CERRADO PARCIAL") and not c.get("evidencia_cierre"):
            errores.append("%s: cerrado sin evidencia_cierre" % cid)
        if not isinstance(c.get("reprogramaciones"), int):
            errores.append("%s: reprogramaciones debe ser entero" % cid)

    # Referencias cruzadas: cada id citado debe existir
    for c in datos.get("compromisos", []):
        for ref in [r.strip() for r in str(c.get("ref_cruzada") or "").split(",") if r.strip()]:
            if ref not in vistos:
                errores.append("%s: ref_cruzada '%s' no existe" % (c.get("id"), ref))

    hid = set()
    for h in hitos.get("hitos", []):
        i = h.get("id", "<sin id>")
        if i in hid:
            errores.append("hito %s: id duplicado" % i)
        hid.add(i)
        if h.get("firmeza") not in VOC_HITOS["firmeza"]:
            errores.append("hito %s: firmeza='%s' fuera del vocabulario"
                           % (i, h.get("firmeza")))
        for campo in ("ventana_inicio", "ventana_fin"):
            if not isinstance(h.get(campo), dt.date):
                errores.append("hito %s: %s no es una fecha" % (i, campo))
        if isinstance(h.get("ventana_inicio"), dt.date) and \
           isinstance(h.get("ventana_fin"), dt.date) and \
           h["ventana_fin"] < h["ventana_inicio"]:
            errores.append("hito %s: ventana_fin anterior a ventana_inicio" % i)
        ref = h.get("ref_compromiso") or ""
        if ref and ref not in vistos:
            errores.append("hito %s: ref_compromiso '%s' no existe" % (i, ref))

    if errores:
        print("ABORTADO — la fuente tiene %d defecto(s):\n" % len(errores))
        for e in errores:
            print("  -", e)
        sys.exit(1)


def hoja_compromisos(wb, comps, corte):
    ws = wb.create_sheet("Compromisos")
    for j, (titulo, ancho, _) in enumerate(COLUMNAS, 1):
        c = ws.cell(row=1, column=j, value=titulo)
        c.font = FT_HEADER
        c.fill = F_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.row_dimensions[1].height = 30

    for i, comp in enumerate(comps, start=2):
        for j, (_, _, campo) in enumerate(COLUMNAS, 1):
            c = ws.cell(row=i, column=j)
            c.font = FT_NORMAL
            c.border = BORDE
            c.alignment = Alignment(vertical="top", wrap_text=True)

            letra = get_column_letter(j)
            if letra == "G":
                # Semaforo: CERRADO por evento; el resto lo decide el reloj.
                c.value = ('=IF($H{r}="CERRADO","CERRADO",'
                           'IF($F{r}="","SIN FECHA",'
                           'IF($F{r}<TODAY(),"VENCIDO",'
                           'IF($F{r}-TODAY()<={u},"POR VENCER","EN PLAZO"))))'
                           ).format(r=i, u=UMBRAL)
                c.font = FT_BOLD
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif letra == "I":
                # Cerrado: dias de desviacion al cierre. Abierto: mora a hoy.
                c.value = ('=IF($F{r}="","",'
                           'IF($H{r}="CERRADO",IF($O{r}="","",$O{r}-$F{r}),'
                           'TODAY()-$F{r}))').format(r=i)
                c.number_format = FMT_ENTERO
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif letra == "W":
                c.value = corte
                c.number_format = FMT_FECHA
                c.alignment = Alignment(horizontal="center", vertical="top")
            else:
                v = comp.get(campo)
                if campo in COL_FECHA:
                    if isinstance(v, dt.date):
                        c.value = v
                        c.number_format = FMT_FECHA
                    c.alignment = Alignment(horizontal="center", vertical="top")
                elif campo == "reprogramaciones":
                    c.value = int(v or 0)
                    c.number_format = FMT_ENTERO
                    c.alignment = Alignment(horizontal="center", vertical="top")
                    if int(v or 0) >= 2:
                        c.font = FT_BOLD
                else:
                    texto(c, v if v is not None else "")
                    if campo == "id":
                        c.font = FT_BOLD
                        c.alignment = Alignment(horizontal="center", vertical="top")

    n = len(comps) + 1
    rango = "A2:I%d" % n
    for expr, fill in (('$G2="VENCIDO"', D_VENCIDO),
                       ('$G2="POR VENCER"', D_POR_VENCER),
                       ('$G2="EN PLAZO"', D_EN_PLAZO),
                       ('$G2="SIN FECHA"', D_SIN_FECHA),
                       ('$G2="CERRADO"', D_CERRADO)):
        ws.conditional_formatting.add(rango, FormulaRule(formula=[expr], fill=fill))

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:W%d" % n
    return ws


def hoja_hitos(wb, hitos):
    ws = wb.create_sheet("Hitos")
    cols = [("ID", 8), ("Hito", 52), ("Ventana inicio", 15), ("Ventana fin", 15),
            ("Firmeza", 24), ("Fuente", 60), ("Estado real", 60),
            ("Consecuencia", 50), ("Ref. compromiso", 15)]
    for j, (t, a) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=t)
        c.font = FT_HEADER
        c.fill = F_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
        ws.column_dimensions[get_column_letter(j)].width = a
    ws.row_dimensions[1].height = 30

    claves = ["id", "hito", "ventana_inicio", "ventana_fin", "firmeza",
              "fuente", "estado_real", "consecuencia", "ref_compromiso"]
    for i, h in enumerate(hitos, start=2):
        for j, k in enumerate(claves, 1):
            c = ws.cell(row=i, column=j)
            c.font = FT_NORMAL
            c.border = BORDE
            c.alignment = Alignment(vertical="top", wrap_text=True)
            v = h.get(k)
            if k in ("ventana_inicio", "ventana_fin"):
                c.value = v
                c.number_format = FMT_FECHA
                c.alignment = Alignment(horizontal="center", vertical="top")
            else:
                texto(c, v if v is not None else "")
                if k == "id":
                    c.font = FT_BOLD
                    c.alignment = Alignment(horizontal="center", vertical="top")
        # El color del hito lo fija su firmeza, que no depende del reloj.
        firm = h.get("firmeza")
        if firm == "CONTRACTUAL FIRME":
            ws.cell(row=i, column=5).fill = F_VENCIDO
        elif firm == "REFERENCIA FIJA ADASA":
            ws.cell(row=i, column=5).fill = F_POR_VENCER
        elif firm == "DECLARADO PROVEEDOR":
            ws.cell(row=i, column=5).fill = F_SIN_FECHA
    ws.freeze_panes = "B2"
    return ws


def hoja_movimientos(wb, comps):
    """El historial de desplazamientos. Es lo que convierte el registro en prueba."""
    ws = wb.create_sheet("Movimientos")
    cols = [("ID", 9), ("Obligado", 15), ("Compromiso", 58),
            ("Reprogramaciones", 12), ("Historial de fechas", 70),
            ("Fecha vigente", 15), ("Estado", 16), ("Origen de la fecha", 22)]
    for j, (t, a) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=t)
        c.font = FT_HEADER
        c.fill = F_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
        ws.column_dimensions[get_column_letter(j)].width = a
    ws.row_dimensions[1].height = 30

    movidos = sorted([c for c in comps if int(c.get("reprogramaciones") or 0) > 0],
                     key=lambda c: -int(c["reprogramaciones"]))
    for i, comp in enumerate(movidos, start=2):
        vals = [comp["id"], comp["obligado"], comp["compromiso"],
                int(comp["reprogramaciones"]), comp["historial_fechas"],
                comp.get("fecha_comprometida"), comp["estado"], comp["origen_fecha"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j)
            c.font = FT_NORMAL
            c.border = BORDE
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if j == 4:
                c.value = v
                c.number_format = FMT_ENTERO
                c.font = FT_BOLD
                c.alignment = Alignment(horizontal="center", vertical="top")
                c.fill = F_VENCIDO if v >= 3 else (F_POR_VENCER if v == 2 else F_SIN_FECHA)
            elif j == 6:
                if isinstance(v, dt.date):
                    c.value = v
                    c.number_format = FMT_FECHA
                c.alignment = Alignment(horizontal="center", vertical="top")
            else:
                texto(c, v if v is not None else "")
                if j == 1:
                    c.font = FT_BOLD
    ws.freeze_panes = "B2"
    return ws, len(movidos)


def hoja_panel(wb, comps, meta, n_hitos, n_mov):
    ws = wb.create_sheet("Panel", 0)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 34
    ws.column_dimensions["E"].width = 12

    n = len(comps) + 1
    G = "Compromisos!$G$2:$G$%d" % n
    C = "Compromisos!$C$2:$C$%d" % n
    B = "Compromisos!$B$2:$B$%d" % n
    H = "Compromisos!$H$2:$H$%d" % n
    J = "Compromisos!$J$2:$J$%d" % n

    r = 1
    texto(ws.cell(row=r, column=1), "REGISTRO DE COMPROMISOS").font = FT_TITULO
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.cell(row=r, column=1).fill = F_TITULO
    ws.row_dimensions[r].height = 24
    r += 1
    texto(ws.cell(row=r, column=1), "%s  |  %s  |  Contrato %s"
          % (meta["codigo"], meta["titulo"], meta["contrato"])).font = FT_NORMAL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 2

    texto(ws.cell(row=r, column=1), "Corte de la fuente").font = FT_BOLD
    c = ws.cell(row=r, column=2, value=meta["corte"])
    c.number_format = FMT_FECHA
    c.font = FT_NORMAL
    r += 1
    texto(ws.cell(row=r, column=1), "Vista al").font = FT_BOLD
    c = ws.cell(row=r, column=2, value="=TODAY()")
    c.number_format = FMT_FECHA
    c.font = FT_BOLD
    r += 1
    texto(ws.cell(row=r, column=1), "Umbral de 'por vencer' (dias)").font = FT_BOLD
    ws.cell(row=r, column=2, value=meta["umbral_por_vencer_dias"]).font = FT_NORMAL
    r += 2

    def bloque(titulo, filas, rango, ancla_col=1):
        nonlocal r
        inicio = r
        texto(ws.cell(row=r, column=ancla_col), titulo).font = FT_SECCION
        r += 1
        for etiqueta, fill in filas:
            c1 = texto(ws.cell(row=r, column=ancla_col), etiqueta)
            c1.font = FT_NORMAL
            c1.border = BORDE
            if fill:
                c1.fill = fill
            c2 = ws.cell(row=r, column=ancla_col + 1,
                         value='=COUNTIFS(%s,"%s")' % (rango, etiqueta))
            c2.font = FT_BOLD
            c2.border = BORDE
            c2.number_format = FMT_ENTERO
            c2.alignment = Alignment(horizontal="center")
            r += 1
        r += 1
        return inicio

    bloque("POR SEMAFORO", [("VENCIDO", F_VENCIDO), ("POR VENCER", F_POR_VENCER),
                            ("EN PLAZO", F_EN_PLAZO), ("SIN FECHA", F_SIN_FECHA),
                            ("CERRADO", F_CERRADO)], G)
    bloque("POR OBLIGADO", [("BW WATER", None), ("ADASA", None),
                            ("BUREAU VERITAS", None)], C)
    bloque("POR CRITICIDAD", [("CRITICA", F_VENCIDO), ("ALTA", F_POR_VENCER),
                              ("MEDIA", F_SIN_FECHA), ("BAJA", None)], J)

    # Segunda columna del panel
    r2 = 9
    for titulo, filas, rango in (
            ("POR FRENTE", [("BUREAU VERITAS", None), ("PROGRAMA", None),
                            ("COMERCIAL", None), ("INTERNO ADASA", None)], B),
            ("POR ESTADO", [("ABIERTO", None), ("EN DISPUTA", F_POR_VENCER),
                            ("CERRADO PARCIAL", None), ("CERRADO", F_CERRADO)], H)):
        texto(ws.cell(row=r2, column=4), titulo).font = FT_SECCION
        r2 += 1
        for etiqueta, fill in filas:
            c1 = texto(ws.cell(row=r2, column=4), etiqueta)
            c1.font = FT_NORMAL
            c1.border = BORDE
            if fill:
                c1.fill = fill
            c2 = ws.cell(row=r2, column=5,
                         value='=COUNTIFS(%s,"%s")' % (rango, etiqueta))
            c2.font = FT_BOLD
            c2.border = BORDE
            c2.number_format = FMT_ENTERO
            c2.alignment = Alignment(horizontal="center")
            r2 += 1
        r2 += 1

    # Vencidos por obligado: el dato que se cita en el Estado Vigente
    texto(ws.cell(row=r2, column=4), "VENCIDOS POR OBLIGADO").font = FT_SECCION
    r2 += 1
    for etiqueta in ("BW WATER", "ADASA", "BUREAU VERITAS"):
        c1 = texto(ws.cell(row=r2, column=4), etiqueta)
        c1.font = FT_NORMAL
        c1.border = BORDE
        c2 = ws.cell(row=r2, column=5,
                     value='=COUNTIFS(%s,"VENCIDO",%s,"%s")' % (G, C, etiqueta))
        c2.font = FT_BOLD
        c2.border = BORDE
        c2.number_format = FMT_ENTERO
        c2.alignment = Alignment(horizontal="center")
        r2 += 1

    r = max(r, r2) + 1
    texto(ws.cell(row=r, column=1), "TOTALES").font = FT_SECCION
    r += 1
    for etiqueta, valor in (("Compromisos registrados", len(comps)),
                            ("Hitos registrados", n_hitos),
                            ("Compromisos con reprogramacion", n_mov)):
        texto(ws.cell(row=r, column=1), etiqueta).font = FT_NORMAL
        ws.cell(row=r, column=2, value=valor).font = FT_BOLD
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
        r += 1
    r += 1
    texto(ws.cell(row=r, column=1),
          "Regla anti-contradiccion: si el README y este registro discrepan en una "
          "fecha vigente, manda el registro. Una fecha nunca se escribe dos veces "
          "como estado vivo.").font = FT_BOLD
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 28
    return ws


def hoja_leyenda(wb, meta):
    ws = wb.create_sheet("Leyenda")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 105

    r = 1
    texto(ws.cell(row=r, column=1), "LEYENDA DEL REGISTRO DE COMPROMISOS").font = FT_TITULO
    r += 2

    def seccion(titulo, filas):
        nonlocal r
        texto(ws.cell(row=r, column=1), titulo).font = FT_SECCION
        r += 1
        for a, b, fill in filas:
            c1 = texto(ws.cell(row=r, column=1), a)
            c1.font = FT_BOLD
            if fill:
                c1.fill = fill
            c2 = texto(ws.cell(row=r, column=2), b)
            c2.font = FT_NORMAL
            c2.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
        r += 1

    seccion("SEMAFORO (columna G, calculado)", [
        ("VENCIDO", "Compromiso abierto cuya fecha ya paso.", F_VENCIDO),
        ("POR VENCER", "Abierto y vence dentro de los proximos %d dias, una weekly "
                       "del martes." % meta["umbral_por_vencer_dias"], F_POR_VENCER),
        ("EN PLAZO", "Abierto y con mas de %d dias de margen."
         % meta["umbral_por_vencer_dias"], F_EN_PLAZO),
        ("SIN FECHA", "Abierto y sin fecha comprometida. NO es un estado neutro: un "
                      "compromiso sin vencimiento es un defecto de gestion, y por eso "
                      "se colorea.", F_SIN_FECHA),
        ("CERRADO", "Cumplido. El color lo fija el Estado, no el reloj.", F_CERRADO),
    ])

    seccion("COMO SE CALCULA", [
        ("Semaforo",
         "Formula viva, no un relleno fijo. Si el Estado es CERRADO devuelve CERRADO; "
         "si no hay fecha comprometida devuelve SIN FECHA; si la fecha es anterior a "
         "TODAY() devuelve VENCIDO; si faltan %d dias o menos devuelve POR VENCER; en "
         "otro caso EN PLAZO. Un relleno fijado al generar seria correcto el dia de la "
         "corrida y falso al siguiente." % meta["umbral_por_vencer_dias"], None),
        ("Dias vs. compromiso",
         "Para un compromiso cerrado, dias de desviacion entre la fecha de cierre y la "
         "comprometida (positivo = se cerro tarde). Para uno abierto, dias transcurridos "
         "desde la fecha comprometida (positivo = mora, negativo = margen).", None),
        ("Color de la fila",
         "Formato condicional anclado a la columna G del semaforo, con columna absoluta "
         "y fila relativa, de modo que cada fila se pinta con su propio estado.", None),
    ])

    seccion("ORIGEN DE LA FECHA (columna M)", [
        ("CONTRACTUAL", "Fijada por la BAE, la ET, el PIE o el Contrato. Es la unica "
                        "categoria que se puede citar como obligacion exigible con "
                        "consecuencia contractual.", None),
        ("COMPROMISO ESCRITO BW", "Comprometida por escrito por el proveedor, en correo o "
                                  "documento. Exigible como compromiso propio, no como "
                                  "clausula.", None),
        ("COMPROMISO VERBAL-MINUTA", "Consta solo en minuta de reunion. Se recuerda; no "
                                     "sostiene un reclamo por si sola.", None),
        ("FIJADA ADASA", "Plazo fijado unilateralmente por ADASA. Su fuerza depende de la "
                         "obligacion contractual que lo respalde, no de la fecha misma.", None),
    ])

    seccion("ESTADO (columna H, cambia por evento)", [
        ("ABIERTO", "Pendiente de cumplimiento.", None),
        ("CERRADO", "Cumplido, con evidencia de cierre registrada.", None),
        ("CERRADO PARCIAL", "Cumplido en forma pero deficiente en fondo. Sigue "
                            "generando accion.", None),
        ("EN DISPUTA", "Las partes no coinciden en la fecha, el alcance o la "
                       "exigibilidad.", None),
    ])

    seccion("REPROGRAMACIONES (columna U)", [
        ("Que cuenta", "Los desplazamientos efectivos de la fecha comprometida, no las "
                       "veces que el compromiso se menciona. Convierte en un numero lo "
                       "que de otro modo queda disuelto en prosa: tres reprogramaciones "
                       "son la prueba objetiva de un patron, no una impresion.", None),
    ])

    seccion("GOBERNANZA", [
        ("Fuente unica", "compromisos.yaml y hitos.yaml. Este .xlsx es DERIVADO: se "
                         "regenera con generar_excel_compromisos.py y cualquier edicion "
                         "manual se pierde en la siguiente corrida.", None),
        ("Anti-contradiccion", "Si el README del proyecto y este registro discrepan en una "
                               "fecha vigente, manda el registro.", None),
        ("Fechas", "Ninguna fecha se infiere. Una fuente que dice 'proxima semana' u "
                   "'once ready' produce SIN FECHA, nunca una fecha estimada.", None),
    ])

    texto(ws.cell(row=r, column=1), "Documento").font = FT_BOLD
    texto(ws.cell(row=r, column=2),
          "%s | Preparado por: %s | Corte: %s"
          % (meta["codigo"], meta["preparado_por"], meta["corte"].strftime("%d-%b-%Y"))
          ).font = FT_NORMAL
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sin-backup", action="store_true")
    args = ap.parse_args()

    with open(os.path.join(BASE, "compromisos.yaml"), encoding="utf-8") as f:
        datos = yaml.safe_load(f)
    with open(os.path.join(BASE, "hitos.yaml"), encoding="utf-8") as f:
        hitos_doc = yaml.safe_load(f)

    validar_fuente(datos, hitos_doc)

    meta = datos["meta"]
    comps = datos["compromisos"]
    hitos = hitos_doc["hitos"]

    global UMBRAL
    UMBRAL = int(meta["umbral_por_vencer_dias"])

    if os.path.exists(OUTPUT) and not args.sin_backup:
        os.makedirs(BACKUPS, exist_ok=True)
        dst = os.path.join(BACKUPS, "P22-IT-06-000-006-0_Registro-de-Compromisos_pre-%s.xlsx"
                           % dt.date.today().strftime("%Y%m%d"))
        shutil.copy2(OUTPUT, dst)
        print("Backup:", os.path.relpath(dst, BASE))

    wb = Workbook()
    wb.remove(wb.active)
    hoja_compromisos(wb, comps, meta["corte"])
    hoja_hitos(wb, hitos)
    _, n_mov = hoja_movimientos(wb, comps)
    hoja_panel(wb, comps, meta, len(hitos), n_mov)
    hoja_leyenda(wb, meta)
    wb.move_sheet("Panel", offset=-4)
    wb.save(OUTPUT)

    # --- resumen para pegar en el Estado Vigente del README ---
    hoy = dt.date.today()
    def sem(c):
        if c["estado"] == "CERRADO":
            return "CERRADO"
        f = c.get("fecha_comprometida")
        if not f:
            return "SIN FECHA"
        if f < hoy:
            return "VENCIDO"
        return "POR VENCER" if (f - hoy).days <= UMBRAL else "EN PLAZO"

    est = {}
    for c in comps:
        est[sem(c)] = est.get(sem(c), 0) + 1
    vivos = sum(v for k, v in est.items() if k != "CERRADO")
    venc = [c for c in comps if sem(c) == "VENCIDO"]
    v_bw = sum(1 for c in venc if c["obligado"] == "BW WATER")
    v_ad = sum(1 for c in venc if c["obligado"] == "ADASA")
    v_ot = len(venc) - v_bw - v_ad

    print("\nExcel generado:", os.path.basename(OUTPUT))
    print("  compromisos %d | hitos %d | con reprogramacion %d" % (len(comps), len(hitos), n_mov))
    print("  semaforo:", " · ".join("%s %d" % (k, est[k]) for k in
                                    ("VENCIDO", "POR VENCER", "EN PLAZO", "SIN FECHA", "CERRADO")
                                    if k in est))
    print("\n--- linea para el Estado Vigente ---")
    partes = "%d BW Water / %d ADASA" % (v_bw, v_ad)
    if v_ot:
        partes += " / %d otros" % v_ot
    print("Compromisos: %d vivos · %d vencidos (%s) · %d vencen esta semana — corte %s"
          % (vivos, len(venc), partes, est.get("POR VENCER", 0),
             hoy.strftime("%d-%b-%Y")))
    return 0


UMBRAL = 7

if __name__ == "__main__":
    sys.exit(main())
