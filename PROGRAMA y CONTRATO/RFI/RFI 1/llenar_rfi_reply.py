# -*- coding: utf-8 -*-
"""Llena el bloque 'Replied Information' del form RFI-001 (round-trip).
Patron patcher (CLAUDE.md 3.10): preserva el .docx original recibido y emite
RFI-001 ..._ADASA_REPLY.docx. Metadatos Word limpios (global 2.3)."""
import os, sys, shutil, copy
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
from docx_metadata import apply_core_properties, fix_app_xml

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "RFI-001 Clarification AI Module with HART.docx")
OUT = os.path.join(HERE, "RFI-001 Clarification AI Module with HART_ADASA_REPLY.docx")
BACKUP = os.path.join(HERE, "RFI-001 Clarification AI Module with HART_ORIGINAL.docx.bak")

REPLIED_DATE = "24 Jun. 2026"
REPLY_PARAS = [
    "ADASA confirms that compliance with the Technical Specification, Section 5.5 - "
    "Instrumentation Specification, is achieved under the configuration described in the "
    "Contractor Understanding, namely: all field instruments supplied as 4-20 mA with HART "
    "communication capability; HART communication available locally at the field instrument "
    "through a handheld HART communicator; and the PLC analog input module (Allen-Bradley "
    "5069-IF8) acquiring the 4-20 mA process variable, with the 250 ohms resistor noted on "
    "the module datasheet allowing an external HART device (e.g. a handheld communicator) on "
    "the loop, preserving field HART access.",

    "Section 5.5 sets the instrumentation signal protocol as 4-20 mA + HART at the "
    "field-instrument level. It does not require HART pass-through to the PLC or SCADA, "
    "HART-capable I/O modules, or a dedicated HART asset-management solution. The "
    "control-system communication required by the Specification is Modbus TCP/IP over "
    "Ethernet, which is unaffected. The alternative HART-capable I/O hardware referenced in "
    "the RFI is therefore not required, with no associated cost, engineering, procurement or "
    "schedule impact.",

    "Transmittal N24 already confirmed the 4-20 mA + HART instrumentation requirement as met "
    "by the offered transmitters (all 4-20 mA + HART) and closed the point; this disposition "
    "is consistent with that.",

    "This confirmation is conditioned on every field instrument being supplied as 4-20 mA + "
    "HART-capable. Any instrument supplied as 4-20 mA only, without HART capability, would "
    "not comply with Section 5.5; the observations previously issued on instruments lacking "
    "HART capability remain in force.",
]

def template_rpr(cell):
    """Devuelve el rPr (formato de run) del primer run del cell, para clonarlo."""
    for p in cell.paragraphs:
        for r in p.runs:
            if r._element.rPr is not None:
                return r._element.rPr
    return None

def set_field(cell, label, value):
    """Reemplaza el texto del cell por 'label value' preservando el formato del run."""
    p = cell.paragraphs[0]
    rpr = template_rpr(cell)
    # limpiar runs existentes
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(f"{label} {value}")
    if rpr is not None:
        run._element.insert(0, copy.deepcopy(rpr))
    else:
        run.font.name = "Arial"

def main():
    if not os.path.exists(BACKUP):
        shutil.copy2(SRC, BACKUP)
    doc = docx.Document(SRC)
    t = doc.tables[1]

    # row 5: To / From  | row 6: Company / Company
    set_field(t.rows[5].cells[0], "To (Name) :", "Billy Tan")
    set_field(t.rows[5].cells[1], "From (Name) :", "Luis Rivera Gonzalez")
    set_field(t.rows[6].cells[0], "Company / Department :", "BW Water / Engineering")
    set_field(t.rows[6].cells[1], "Company / Department :", "Aguas Antofagasta / Projects")

    # row 7: celda de cuerpo (vacia, con parrafos en blanco del template) ->
    # limpiar los parrafos vacios y cargar Replied Date + parrafos de respuesta
    body = t.rows[7].cells[0]
    # tomar formato de un run existente de la tabla para clonar Arial
    rpr = template_rpr(t.rows[1].cells[0])
    # eliminar todos los parrafos salvo el primero (una celda requiere >=1 parrafo)
    for p in list(body.paragraphs[1:]):
        p._element.getparent().remove(p._element)
    # vaciar runs del parrafo remanente
    for r in list(body.paragraphs[0].runs):
        r._element.getparent().remove(r._element)

    def add_para(text, bold=False, space_after=8):
        p = body.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        if rpr is not None:
            run._element.insert(0, copy.deepcopy(rpr))
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.bold = bold
        return p

    # reutilizar el parrafo vacio inicial de la celda si existe
    first = body.paragraphs[0]
    if not first.runs and not first.text.strip():
        r = first.add_run(f"Replied Date: {REPLIED_DATE}")
        if rpr is not None:
            r._element.insert(0, copy.deepcopy(rpr))
        r.font.name = "Arial"; r.font.size = Pt(10); r.font.bold = True
        first.paragraph_format.space_after = Pt(8)
    else:
        add_para(f"Replied Date: {REPLIED_DATE}", bold=True)

    for para in REPLY_PARAS:
        add_para(para)

    # --- metadatos limpios (global 2.3) ---
    apply_core_properties(
        doc,
        title="RFI-001 Reply - 4-20 mA + HART (AI Module)",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="RFI 25007-RO-RFI-0001 - Reply",
        comments="",
        language="en-US",
    )
    doc.save(OUT)
    fix_app_xml(OUT, application="Microsoft Office Word", company="Aguas Antofagasta")
    print("[OK] generado:", OUT)

if __name__ == "__main__":
    main()
