# -*- coding: utf-8 -*-
"""Llena el bloque 'Replied Information' del form RFI-002 (round-trip).
Patron patcher (CLAUDE.md 3.10/3.12): preserva el .docx original recibido y
emite RFI-002 ..._ADASA_REPLY.docx. Metadatos Word limpios (global 2.3).
Clon de RFI 1/llenar_rfi_reply.py (misma estructura de form: tables[1],
rows 5=To/From, 6=Company, 7=cuerpo; gridSpan -> escribir solo cells[0])."""
import os, sys, shutil, copy
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
from docx_metadata import apply_core_properties, fix_app_xml

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "RFI-002 Panel Enclosure Nema 4_IP66.docx")
OUT = os.path.join(HERE, "RFI-002 Panel Enclosure Nema 4_IP66_ADASA_REPLY.docx")
BACKUP = os.path.join(HERE, "RFI-002 Panel Enclosure Nema 4_IP66_ORIGINAL.docx.bak")

REPLIED_DATE = "07 Jul. 2026"
REPLY_PARAS = [
    "ADASA confirms that the external enclosure body, external door, roof, rear "
    "panel, plinth and the gasketed gland plates on the upper and lower bases "
    "shall be Stainless Steel 316L, unpainted, in accordance with the approved "
    "Datasheet of Local Control Panel (LCP), P22-ET-09-007-005 Rev 1 (nVent "
    "Hoffman Type FS FS66S, SS316L, NEMA 4X / IP66). The Technical Specification "
    "(P22-ET-09-000-001-0), Section 5.4.1 - Constructive Characteristics, "
    "requires the "
    "upper and lower bases to carry a bolted metal plate with gaskets for the "
    "cable glands and a protection class of NEMA 4X or its IP equivalent; ADASA "
    "notes and accepts BW Water's confirmation that the gland plates change to "
    "SS316L, which maintains that requirement.",

    "The internal mounting components that are not exposed to the external "
    "environment - mounting plate, internal swing door and internal frame - may "
    "be supplied in galvanized or cold-rolled steel in accordance with the "
    "manufacturer's standard construction. The Technical Specification, "
    "Section 5.4.1 - Constructive Characteristics, permits measuring, "
    "switching and control equipment to be mounted inside the cabinet provided "
    "the protection grade is not lower than NEMA 4X; these internal materials do "
    "not reduce the enclosure protection, which remains NEMA 4X / IP66. The Gray "
    "RAL 7035 finishing is accepted only on internal components and surfaces; "
    "the external SS316L enclosure remains unpainted as specified in the "
    "approved datasheet.",

    "This is a documentation clarification, not a material or design deviation. "
    "The Technical Specification, Section 5.4.3 - Finishing, "
    "allows a stainless steel enclosure, and the approved LCP Datasheet fixes "
    "that stainless steel as SS316L; the Outline Panel Drawing must align to it. "
    "BW Water shall re-issue the PLC-LCP Outline Panel Drawing "
    "(P22-CD-09-008-001) at the next revision, distinguishing the external "
    "SS316L enclosure components from the internal mounting components, "
    "restricting the Gray RAL 7035 finishing to internal surfaces, and stating "
    "the protection class in full as \"NEMA 4X / IP66\". This closes the "
    "observation on the Outline Panel Drawing tracked since Transmittal N20 and "
    "re-issued at Transmittal N25.",

    "This confirmation is conditional on the external enclosure and gasketed "
    "gland plates being Stainless Steel 316L and on the enclosure maintaining a "
    "protection class not lower than NEMA 4X per the Technical Specification, "
    "Section 5.4.1 - Constructive Characteristics. The "
    "observations on the Outline Panel Drawing remain in force until the "
    "corrected revision is submitted and approved.",
]


def template_rpr(cell):
    for p in cell.paragraphs:
        for r in p.runs:
            if r._element.rPr is not None:
                return r._element.rPr
    return None


def set_field(cell, label, value):
    p = cell.paragraphs[0]
    rpr = template_rpr(cell)
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(f"{label} {value}")
    if rpr is not None:
        run._element.insert(0, copy.deepcopy(rpr))
    else:
        run.font.name = "Arial"


def main():
    if not os.path.exists(BACKUP):
        shutil.copy2(SRC, BACKUP)
    doc = docx.Document(SRC)
    t = doc.tables[1]

    set_field(t.rows[5].cells[0], "To (Name) :", "Billy Tan")
    set_field(t.rows[5].cells[1], "From (Name) :", "Luis Rivera Gonzalez")
    set_field(t.rows[6].cells[0], "Company / Department :", "BW Water / Engineering")
    set_field(t.rows[6].cells[1], "Company / Department :", "Aguas Antofagasta / Projects")

    body = t.rows[7].cells[0]
    rpr = template_rpr(t.rows[1].cells[0])
    for p in list(body.paragraphs[1:]):
        p._element.getparent().remove(p._element)
    for r in list(body.paragraphs[0].runs):
        r._element.getparent().remove(r._element)

    def add_para(text, bold=False, space_after=8):
        p = body.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        if rpr is not None:
            run._element.insert(0, copy.deepcopy(rpr))
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.bold = bold
        return p

    first = body.paragraphs[0]
    if not first.runs and not first.text.strip():
        r = first.add_run(f"Replied Date: {REPLIED_DATE}")
        if rpr is not None:
            r._element.insert(0, copy.deepcopy(rpr))
        r.font.name = "Arial"; r.font.size = Pt(10); r.font.bold = True
        first.paragraph_format.space_after = Pt(8)
    else:
        add_para(f"Replied Date: {REPLIED_DATE}", bold=True)

    for para in REPLY_PARAS:
        add_para(para)

    apply_core_properties(
        doc,
        title="RFI-002 Reply - LCP Enclosure Material (NEMA 4X/IP66)",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="RFI 25007-RO-RFI-0002 - Reply",
        comments="",
        language="en-US",
    )
    doc.save(OUT)
    fix_app_xml(OUT, application="Microsoft Office Word", company="Aguas Antofagasta")
    print("[OK] generado:", OUT)


if __name__ == "__main__":
    main()
