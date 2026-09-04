"""
Correo: RE: Modulo RO Segunda Etapa para Salmuera // Listado de Entregables
        Observaciones ADASA al listado entregado por L&A Ingenieria
Fecha: 30-Abr-2026
Destinatario: Pablo Castillo (L&A Ingenieria y Proyectos)
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "2026-04-30_RE-Listado-Entregables.docx")


def set_font(run, name="Arial", size=11, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text="", bold=False, size=11, space_before=0, space_after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_font(run, bold=bold, size=size)
    return p


def add_separator(doc):
    sep = doc.add_paragraph()
    sep.paragraph_format.space_before = Pt(0)
    sep.paragraph_format.space_after = Pt(10)
    pPr = sep._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_table(doc, headers, rows, col_widths=None, font_size=9):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"

    # Header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for run in hdr_cells[i].paragraphs[0].runs:
            set_font(run, bold=True, size=font_size)
        tc = hdr_cells[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "D9D9D9")
        tcPr.append(shd)

    # Data rows
    for ri, row_data in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row_data):
            cells[ci].text = val
            for run in cells[ci].paragraphs[0].runs:
                set_font(run, size=font_size)

    if col_widths:
        for row in table.rows:
            for ci, width in enumerate(col_widths):
                row.cells[ci].width = Cm(width)

    return table


def add_heading(doc, text, level=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    size = 13 if level == 1 else 12
    r = p.add_run(text)
    set_font(r, bold=True, size=size)
    return p


def main():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # --- HEADER ---
    add_para(doc, "ADASA — Aguas de Antofagasta S.A.", bold=True, size=12, space_after=2)
    add_para(
        doc,
        "Departamento de Ingeniería y Optimización (DIO) | "
        "Contrato 067 — Ingeniería Detalle OOCC Módulo RO Salmuera Taltal",
        size=10, space_after=8,
    )

    # --- SUBJECT BLOCK ---
    subj = add_para(doc, space_before=0, space_after=4)
    r = subj.add_run("Asunto: ")
    set_font(r, bold=True, size=11)
    r2 = subj.add_run(
        "RE: Módulo RO Segunda Etapa para Salmuera // Listado de Entregables — "
        "Observaciones ADASA"
    )
    set_font(r2, size=11)

    for label, value in [
        ("Fecha: ", "30 de abril de 2026"),
        ("Para: ", "Pablo Castillo — L&A Ingeniería y Proyectos (pcastillo@lyaingenieria.cl)"),
        (
            "CC: ",
            "Yohana Rodríguez Flores, Cristhian Sánchez, "
            "Luciano Méndez Huidobro, Dio Documentos, Víctor Gutiérrez",
        ),
        ("De: ", "Luis Rivera — Departamento de Ingeniería y Optimización (DIO), ADASA"),
        ("Ref: ", "TdR P22-TR-00-010-01-1 Rev 1"),
    ]:
        p = add_para(doc, space_before=0, space_after=2)
        r = p.add_run(label)
        set_font(r, bold=True, size=11)
        r2 = p.add_run(value)
        set_font(r2, size=11)

    add_separator(doc)

    # --- SALUTATION ---
    add_para(doc, "Estimado Pablo,", size=11, space_after=10)

    # --- OPENING ---
    add_para(
        doc,
        "Recibida la cotización por el servicio. Revisé su listado contra el TdR "
        "P22-TR-00-010-01-1 (Rev 1) y comparto las observaciones para poder cerrar "
        "la oferta.",
        size=11, space_after=10,
    )

    # --- 1. CLASE 2 AACE ---
    add_heading(doc, "1. Cubicaciones y Presupuesto, Clase 2 AACE")

    add_para(
        doc,
        "Acepto el cambio a Clase 2 AACE en lugar de Clase 1 dado que el plazo de "
        "cuatro semanas es incompatible con la solicitud de cotizaciones a firme con "
        "constructora. La planilla Excel debe mantener el desglose mínimo indicado en "
        "la Sección 6 del TdR (Cubicaciones y Presupuesto): movimiento de tierras y "
        "rellenos, hormigón armado por elemento, acero de refuerzo, acero estructural, "
        "soldadura, protección superficial por m², techado por m² e impermeabilización "
        "por m². Incluir además vínculos a planos y memorias correspondientes, y dejar "
        "el nivel Clase 2 explícito en la portada.",
        size=11, space_after=10,
    )

    # --- 2. ENTREGABLES POR ACLARAR ---
    add_heading(doc, "2. Entregables por aclarar o adicionar")

    add_para(
        doc,
        "Cruzando su listado con la Sección 6 del TdR identifiqué siete ítems que "
        "no aparecen explícitamente en la cotización:",
        size=11, space_after=8,
    )

    headers = ["Nº", "Entregable solicitado", "Código TdR", "Sección TdR", "Comentario"]
    rows = [
        (
            "1",
            "Plano de Excavaciones y Movimiento de Tierras "
            "(planta general y cortes)",
            "P22-DWG-00-001-001",
            "Sección 6 (Planos de Construcción)",
            "No identificable en las 18 láminas ofertadas; ninguna línea "
            "menciona excavaciones o movimiento de tierras.",
        ),
        (
            "2",
            "Especificación Técnica de Movimiento de Tierras y Excavaciones",
            "P22-ET-00-010-101-0",
            "Sección 6 (Especificaciones Técnicas)",
            "La oferta tiene “ET para Construcción de Obras Civiles” "
            "(2 unidades); el título no incluye movimiento de tierras. El TdR "
            "solicita tres ETs distintas: Movimiento de Tierras, Hormigón Armado "
            "y Estructura Metálica.",
        ),
        (
            "3",
            "Plano de Cubierta Metálica sistema CIP "
            "(planta, elevaciones, cortes, detalles de conexión)",
            "P22-DWG-00-003-001",
            "Sección 6 (Planos de Construcción)",
            "Las 18 láminas ofertadas se identifican como “Obras Civiles”; "
            "no hay plano de estructura metálica, siendo la cubierta CIP la "
            "única estructura metálica del alcance.",
        ),
        (
            "4",
            "Plano de Canalizaciones y Red de Drenajes, tres láminas: "
            "L1 planta general y trazado; L2 cámaras CD-06-00N "
            "(detalle típico, armaduras, tapa y sellos); "
            "L3 perfil longitudinal con cotas y pendientes",
            "P22-DWG-00-002-006",
            "Sección 6 (Planos de Construcción)",
            "Posible solapamiento con “Disposición General Obras Civiles "
            "Secciones” (4 unidades) pero no se identifican las tres "
            "láminas específicas.",
        ),
        (
            "5",
            "Programa detallado de trabajo (Hito H0)",
            "—",
            "Sección 8 (Plazos de Ejecución)",
            "Tres días hábiles desde la adjudicación; confirmar si está "
            "cubierto en las HH de Control de Proyectos.",
        ),
        (
            "6",
            "Informes semanales de avance",
            "—",
            "Sección 10 (Comunicación y Estados de Avance)",
            "Confirmar si las HH de Control de Proyectos (17 hr) y Adm. Contrato "
            "(19 hr) cubren este entregable.",
        ),
        (
            "7",
            "Reuniones semanales de coordinación por videoconferencia",
            "—",
            "Sección 10 (Comunicación y Estados de Avance)",
            "Modalidad remota, sin presencialidad en sitio.",
        ),
    ]
    add_table(doc, headers, rows, col_widths=[0.8, 5.0, 2.5, 3.7, 5.0], font_size=9)

    add_para(doc, space_after=8)

    add_para(
        doc,
        "La propuesta revisada debe identificar cada ítem con su código del TdR y "
        "las HH asignadas. Confirme además por escrito que las cuatro memorias de "
        "cálculo de Obras Civiles ofertadas corresponden a P22-MC-00-002-001 "
        "(Fundación TK-06-001), P22-MC-00-002-002 (Fundación BH-06-001), "
        "P22-MC-00-002-003 (Sistema de Drenajes) y P22-MC-00-002-004 (Fundación "
        "Compartida Contenedor + CIP).",
        size=11, space_after=8,
    )

    # Nota — planos georreferenciados
    p_geo = doc.add_paragraph()
    p_geo.paragraph_format.space_before = Pt(0)
    p_geo.paragraph_format.space_after = Pt(8)
    r1 = p_geo.add_run("Nota sobre planos de planta: ")
    set_font(r1, bold=True, size=11)
    r2 = p_geo.add_run(
        "todos los planos de planta entregados (Implantación general, fundaciones, "
        "fosa, canalizaciones y drenajes) deberán ser georreferenciados en "
        "coordenadas UTM WGS84 Huso 19 Sur, tomando como base el Levantamiento DIO "
        "Abr-2026 incluido en el paquete de antecedentes."
    )
    set_font(r2, size=11)

    # Nota — sistema de drenajes
    p_dren = doc.add_paragraph()
    p_dren.paragraph_format.space_before = Pt(0)
    p_dren.paragraph_format.space_after = Pt(10)
    r1 = p_dren.add_run("Nota sobre el sistema de drenajes: ")
    set_font(r1, bold=True, size=11)
    r2 = p_dren.add_run(
        "No queda claro en la oferta si se considera (a) el diseño de las cámaras "
        "de drenajes CD-06-001, CD-06-002 y CD-06-003, ni (b) el perfil longitudinal "
        "del sistema con cotas y pendientes hacia la fosa TK-06-002. Ambos son "
        "requeridos en la Sección 6 del TdR: el detalle típico prefabricado de las "
        "cámaras (dimensiones, armaduras de tapa y cuerpo, sellos de ingreso y "
        "salida de tuberías) corresponde a la lámina L2 del plano P22-DWG-00-002-006 "
        "y el perfil longitudinal a la lámina L3. El diseño estructural respectivo "
        "va en la memoria P22-MC-00-002-003. Como referencia, el plano de conjunto "
        "P22-DWG-06-005-104 (PL. DRENAJES) de Van Doorn, disponible en el paquete "
        "de antecedentes, identifica las tres cámaras intermedias y sus cotas, base "
        "sobre la cual se desarrollan el detalle típico y el perfil longitudinal."
    )
    set_font(r2, size=11)

    # --- 3. NAVISWORKS ---
    add_heading(doc, "3. Maqueta 3D Navisworks")

    add_para(
        doc,
        "Verifiqué el archivo Maqueta Gral.nwd disponible en el enlace de "
        "antecedentes indicado en la Sección 5 del TdR (Antecedentes) y abre "
        "correctamente. Le solicito volver a descargar el archivo desde ese "
        "enlace; la copia local pudo quedar incompleta durante la descarga anterior. "
        "Como referencia, los planos 2D entregados como antecedentes son la fuente "
        "vinculante; la maqueta tiene carácter de apoyo a la verificación "
        "tridimensional de interferencias, conforme a la Sección 4 del TdR "
        "(Desarrollo de Implantación del Proyecto, Modelo 3D del Proyecto).",
        size=11, space_after=10,
    )

    # --- 4. CIERRE ---
    add_heading(doc, "4. Cierre")

    add_para(
        doc,
        "Quedo atento a la oferta revisada con las aclaraciones señaladas y el "
        "ajuste a Clase 2 AACE acordado. Remitir la propuesta a "
        "dio_documentos@aguasantofagasta.cl con copia a los destinatarios de este "
        "correo.",
        size=11, space_after=14,
    )

    # --- FIRMA ---
    add_para(doc, "Saludos cordiales,", size=11, space_after=10)
    add_para(doc, "Luis Rivera", bold=True, size=11, space_after=2)
    add_para(doc, "Departamento de Ingeniería y Optimización (DIO)", size=11, space_after=2)
    add_para(doc, "ADASA — Aguas de Antofagasta S.A.", size=11, space_after=2)
    add_para(
        doc,
        "Contrato 067 / Servicio Ingeniería Detalle OOCC y Estructuras Metálicas",
        size=11, space_after=0,
    )

    doc.save(OUTPUT)
    print(f"Generado: {OUTPUT}")


if __name__ == "__main__":
    main()
