# -*- coding: utf-8 -*-
"""
Extrae las 5 Memorias de Calculo de ENTREGA 3 (.docx editable) a Markdown,
recorriendo el cuerpo del documento para intercalar parrafos y tablas en orden.
El cronograma .xlsx se extrae con openpyxl. Toda la revision razona contra
estos .md (cumple la prohibicion de Read sobre PDF, CLAUDE.md 1.1).

Uso:  python extraer_mc_md.py
Salida:  ...\ENTREGAS\ENTREGA 3\md\*.md
"""
import os

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ENTREGA3 = os.path.normpath(
    os.path.join(
        SCRIPT_DIR, "..", "..", "..",
        "P22-TR-00-010-01-0", "ENTREGAS", "ENTREGA 3",
    )
)
OUT_DIR = os.path.join(ENTREGA3, "md")

DOCX_MC = [
    "P22-MC-00-002-001_B.docx",
    "P22-MC-00-002-002_B.docx",
    "P22-MC-00-002-004_B.docx",
    "P22-MC-00-002-005_B.docx",
    "P22-MC-00-003-001_B.docx",
]
XLSX_CRONO = "067-032-032-COR-TT-003_0.xlsx"


def _heading_level(style_name: str):
    if not style_name:
        return None
    s = style_name.lower()
    if s.startswith("heading") or s.startswith("titulo") or s.startswith("título"):
        for tok in s.replace("heading", "").replace("titulo", "").replace("título", "").split():
            if tok.isdigit():
                return min(int(tok), 6)
        return 2
    if s in ("title", "titulo del documento"):
        return 1
    return None


def _para_md(p: Paragraph) -> str:
    text = p.text.strip()
    if not text:
        return ""
    lvl = _heading_level(p.style.name if p.style else "")
    if lvl:
        return f"\n{'#' * lvl} {text}\n"
    return text


def _table_md(tbl: Table) -> str:
    rows = []
    for r in tbl.rows:
        cells = [" ".join(c.text.split()).replace("|", "\\|") for c in r.cells]
        rows.append(cells)
    if not rows:
        return ""
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |",
           "| " + " | ".join(["---"] * ncol) + " |"]
    for r in rows[1:]:
        out.append("| " + " | ".join(r) + " |")
    return "\n".join(out) + "\n"


def docx_to_md(path: str) -> str:
    doc = Document(path)
    body = doc.element.body
    tbl_iter = iter(doc.tables)
    par_iter = iter(doc.paragraphs)
    chunks = []
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            try:
                p = next(par_iter)
            except StopIteration:
                continue
            md = _para_md(p)
            if md:
                chunks.append(md)
        elif tag == "tbl":
            try:
                t = next(tbl_iter)
            except StopIteration:
                continue
            chunks.append(_table_md(t))
    return "\n\n".join(c for c in chunks if c.strip())


def xlsx_to_md(path: str) -> str:
    wb = load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"\n## Hoja: {ws.title}\n")
        grid = []
        for row in ws.iter_rows(values_only=True):
            if all(c is None for c in row):
                continue
            grid.append(["" if c is None else str(c).strip() for c in row])
        if not grid:
            continue
        ncol = max(len(r) for r in grid)
        grid = [r + [""] * (ncol - len(r)) for r in grid]
        out.append("| " + " | ".join(grid[0]) + " |")
        out.append("| " + " | ".join(["---"] * ncol) + " |")
        for r in grid[1:]:
            out.append("| " + " | ".join(c.replace("|", "\\|") for c in r) + " |")
    return "\n".join(out)


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    for fn in DOCX_MC:
        src = os.path.join(ENTREGA3, fn)
        dst = os.path.join(OUT_DIR, fn.replace(".docx", ".md"))
        md = docx_to_md(src)
        with open(dst, "w", encoding="utf-8") as fh:
            fh.write(f"# {fn.replace('.docx', '')}\n\n")
            fh.write(md)
        print(f"OK  {fn}  ->  {os.path.basename(dst)}  ({len(md):,} chars)")

    src = os.path.join(ENTREGA3, XLSX_CRONO)
    dst = os.path.join(OUT_DIR, XLSX_CRONO.replace(".xlsx", ".md"))
    md = xlsx_to_md(src)
    with open(dst, "w", encoding="utf-8") as fh:
        fh.write(f"# {XLSX_CRONO.replace('.xlsx', '')}\n")
        fh.write(md)
    print(f"OK  {XLSX_CRONO}  ->  {os.path.basename(dst)}  ({len(md):,} chars)")


if __name__ == "__main__":
    main()
