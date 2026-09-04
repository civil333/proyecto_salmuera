#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera TRANSMITTAL N2 ADASA-LYA (stream OOCC, espanol, para L&A).
Codigo P22-TM-00-010-002-0. ENTREGA 4 (cover L&A 067-032-032-COR-TT-005, 28-May-2026).
Fecha: 01-Jun-2026.

v2 tras revision del usuario (11 comentarios). Cambios principales:
- Leyenda de codigos 1/2/3/4 en el Resumen Ejecutivo (L&A no la tiene incorporada).
- NPT: el escalon entre zonas es correcto; estanque/bomba/fosa coinciden con el
  montaje (se deja pasar en sus memorias); contenedor (-003) y CIP (-007)
  reconcilian con el plano de montaje. Se quita el NPT como observacion de las
  memorias.
- MC-004 peso ~17,3 t aceptado como conservador (NOTA), no rechazo. MC-005
  anclaje CIP diferido aceptado (datos del proveedor). -> MC-004 y MC-005 a Codigo 2.
- ET genericas no se rechazan por titulo (Codigo 2). ET-102 recubrimientos 50/70
  por TdR Seccion 3.3.4.
- Itemizados: declarar Clase 2 AACE (acordada en la reunion de arranque); el
  rotulo "referencial" es aceptable. Se elimina el parrafo interno alcance vs montaje.
- Sin tabla de estado del Transmittal N1; estado actual con notas inline. Un solo
  listado de comentarios para los planos. Solicitud explicita de los entregables
  no recibidos (Excavaciones, Detalles de Anclaje, MC de Drenajes).

Recuento: 7 Codigo 2 + 11 Codigo 3 sobre 18 documentos. Veredicto: 3 - POR REVISAR.
Comentarios como directivas (Defecto / Corregir: / Requisito:).
"""

import sys
import os

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))

from ejemplo_documento import (
    crear_documento_adasa,
    aplicar_arial_12,
    add_simple_table,
    add_bullet as _add_bullet_native,
)
from docx import Document


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "TRANSMITTAL N2 ADASA-LYA.docx")


def add_para(doc, runs):
    para = doc.add_paragraph()
    for item in runs:
        text = item[0]
        attrs = item[1] if len(item) > 1 else {}
        run = para.add_run(text)
        if attrs.get("bold"):
            run.bold = True
        if attrs.get("italic"):
            run.italic = True
    aplicar_arial_12(para)
    return para


def obs(doc, idtit, defecto, corregir, requisito=None):
    runs = [(idtit + " ", {"bold": True}), (defecto + " ",),
            ("Corregir: ", {"bold": True}), (corregir,)]
    if requisito:
        runs += [(" ",), ("Requisito: ", {"bold": True}), (requisito,)]
    add_para(doc, runs)


def add_bullet(doc, text):
    _add_bullet_native(doc, text, size=11, space_after_pt=12)


def main() -> None:
    crear_documento_adasa(
        titulo=("TRANSMITTAL DE REVISIÓN TÉCNICA N2 — INGENIERÍA DE "
                "DETALLE OOCC MÓDULO RO 2DA ETAPA SALMUERA TALTAL"),
        codigo="P22-TM-00-010-002-0",
        output_filename=OUTPUT,
        incluir_toc=True,
    )

    doc = Document(OUTPUT)

    # =========================================================
    # 1. RESUMEN EJECUTIVO
    # =========================================================
    doc.add_heading("RESUMEN EJECUTIVO", level=1)

    add_para(doc, [
        ("VEREDICTO DEL TRANSMITTAL: 3 — POR REVISAR.", {"bold": True}),
        (" Revisión de la ENTREGA 4: cinco Memorias de Cálculo, tres "
         "Especificaciones Técnicas, tres Itemizados de presupuesto y siete "
         "planos de obras civiles (18 documentos). Recuento: 7 Código 2 y 11 "
         "Código 3.",),
    ])

    add_para(doc, [
        ("Códigos de respuesta (según los Términos de Referencia):",
         {"bold": True}),
    ])
    add_simple_table(doc, [
        ("Código", "Significado"),
        ("1 — Aprobado",
         "Sin observaciones; el documento avanza a Rev 0."),
        ("2 — Aprobado con comentarios",
         "Notas a incorporar directamente en Rev 0, sin nueva revisión "
         "intermedia."),
        ("3 — Por revisar",
         "Requiere nueva Rev B incorporando las observaciones antes de "
         "Rev 0."),
        ("4 — Rechazado",
         "Requiere rehacer el documento con revisión completa."),
    ])

    add_para(doc, [
        ("Los documentos de la entrega que no aparezcan en el listado de "
         "respuesta se entienden Aprobados (Código 1).",),
    ])

    add_para(doc, [
        ("Nivel de Piso Terminado (NPT): ", {"bold": True}),
        ("es dato fijo de los planos de montaje del módulo (Términos de "
         "Referencia, Criterios de Diseño). Los planos del estanque, la bomba "
         "y la fosa ya lo declaran y coinciden con el montaje (N.T.N. ≈ "
         "+5,75). Los planos del contenedor RO (cotas en datum relativo) y "
         "del Sistema CIP (+5,610) deben reconciliar el N.T.N. absoluto con el "
         "plano de montaje P22-DWG-06-005-103, verificado contra el "
         "Levantamiento DIO Abr-2026. No se requiere re-declararlo en las "
         "memorias cuyo plano ya lo refleja.",),
    ])

    add_para(doc, [
        ("Entregables comprometidos no recibidos: ", {"bold": True}),
        ("el plano de Excavaciones y Movimiento de Tierras "
         "(P22-DWG-00-001-001), el plano de Detalles de Anclaje y Conexiones "
         "(P22-DWG-00-002-005) y la Memoria de Cálculo del Sistema de "
         "Drenajes (P22-MC-00-002-003), exigidos por los Términos de "
         "Referencia y comprometidos en el Listado de Entregables "
         "067-032-032-COR-LI-001. ADASA los solicita en la próxima entrega.",),
    ])

    # =========================================================
    # 2. OBSERVACIONES POR DOCUMENTO
    # =========================================================
    doc.add_heading("OBSERVACIONES POR DOCUMENTO", level=1)

    # ---------- 2.1 Memorias ----------
    doc.add_heading("Memorias de Cálculo", level=2)

    doc.add_heading(
        "Fundación Estanque TK-06-001 — P22-MC-00-002-001 (Código 3)",
        level=3)
    add_para(doc, [
        ("La losa de fundación y las presiones de contacto son conformes; "
         "el veredicto lo fija el diseño del anclaje. Comentarios anotados: "
         "P22-MC-00-002-001_0_CC_ADASA.pdf.",)])
    obs(doc, "OBS-01 — Anclaje del estanque.",
        "La memoria mantiene el anclaje como postinstalado adhesivo, cuando "
        "debe ser preinstalado colado en sitio — como ya se observó en el "
        "Transmittal N1; además se densificó la armadura de la losa, lo que "
        "agrava la interferencia con un anclaje perforado.",
        "rediseñar a pernos preinstalados con cabeza (ASTM F1554) verificados "
        "por ACI 318-19 Capítulo 17 y Sección 17.10, reconciliando "
        "cantidad/diámetro/círculo con el plano Exfibro EX-26005-F01 Rev C.",
        "ACI 318-19 Capítulo 17 y Sección 17.10; plano Exfibro EX-26005-F01 "
        "Rev C.")
    obs(doc, "NOTA-01 — Trazabilidad de las reacciones basales.",
        "La memoria adopta las reacciones corregidas por ADASA pero cita como "
        "fuente la \"Memoria AFTA\".",
        "citar P22-IT-06-000-005-0 como fuente de las reacciones (Ez = "
        "±3.357 kgf, momento volcante 431.846 kgf·cm, cortante basal 4.617 "
        "kgf).",
        "P22-IT-06-000-005-0.")
    obs(doc, "NOTA-02 — Designación del perno y armadura.",
        "La memoria escribe \"ASTM F1553\" (debe ser F1554) y \"HIR-RE 500\" "
        "(debe ser HIT-RE 500); el cambio de armadura de losa no está "
        "justificado.",
        "corregir las designaciones y declarar el motivo del cambio de "
        "armadura; declarar recubrimientos 50 mm expuesto / 70 mm en contacto "
        "con terreno.",
        "consistencia interna; Términos de Referencia, Sección 3.3.4.")

    doc.add_heading(
        "Fundación Dinámica Bomba BH-06-001 — P22-MC-00-002-002 (Código 2)",
        level=3)
    add_para(doc, [
        ("El análisis dinámico ACI 351.3R es conforme (peso del equipo "
         "trazado a la ficha del fabricante; relación de masa y separación "
         "de frecuencia cumplen). Para incorporar en Rev 0:",)])
    obs(doc, "NOTA-01 — Recubrimientos y trazabilidad normativa.",
        "No se declaran los recubrimientos mínimos; persiste la referencia "
        "residual a NCh 2369:2003.",
        "declarar 50 mm en elementos expuestos y 70 mm en contacto con "
        "terreno, y limpiar la referencia normativa residual.",
        "Términos de Referencia, Sección 3.3.4 (ambiente marino).")

    doc.add_heading(
        "Fundación Contenedor RO — P22-MC-00-002-004 (Código 2)",
        level=3)
    add_para(doc, [
        ("Las cargas del Sistema CIP fueron retiradas y el alcance "
         "reconciliado al contenedor RO. Para incorporar en Rev 0:",)])
    obs(doc, "NOTA-01 — Peso del contenedor (adopción conservadora).",
        "La memoria adopta un peso de contenedor de ~17,3 t, conservador "
        "frente al valor vinculante de 14.934 kg indicado en el Transmittal "
        "N1; ADASA lo acepta como conservador.",
        "corregir la unidad y la notación del valor y citar su fuente; si se "
        "mantiene el ~17,3 t, declararlo explícitamente como adopción "
        "conservadora.",
        "Términos de Referencia, Sección 2.1.2.")
    obs(doc, "NOTA-02 — Versión de ACI 318 y recubrimientos.",
        "El listado de códigos cita ACI 318-19 pero el diseño de pedestales "
        "aún cita ACI 318-10.",
        "unificar la versión a ACI 318-19 en todo el documento y declarar "
        "los recubrimientos 50/70 mm.",
        "Términos de Referencia (ACI 318-19; Sección 3.3.4).")

    doc.add_heading(
        "Fundación Sistema CIP — P22-MC-00-002-005 (Código 2)",
        level=3)
    add_para(doc, [
        ("El split de código y los pesos de dosificación quedaron "
         "reconciliados; la base sísmica adopta NCh 2369:2025.",)])
    add_para(doc, [
        ("Dependencia aceptada — anclajes del Sistema CIP: ", {"bold": True}),
        ("la verificación de los anclajes de los equipos CIP queda diferida "
         "por falta de los planos y fichas finales del proveedor de equipos "
         "(BW Water), según lo acordado en el Transmittal N1. Se completará "
         "con anclajes postinstalados que cumplan las solicitaciones (ACI "
         "318-19; precalificación sísmica ACI 355.2/355.4) cuando se "
         "disponga de los datos. No condiciona la aprobación del diseño.",),
    ])
    obs(doc, "NOTA-01 — Recubrimientos.",
        "No se declaran los recubrimientos mínimos.",
        "declarar 50 mm expuesto / 70 mm en contacto con terreno.",
        "Términos de Referencia, Sección 3.3.4.")

    doc.add_heading(
        "Cubierta Metálica Sistema CIP Exterior — P22-MC-00-003-001 "
        "(Código 3)", level=3)
    add_para(doc, [
        ("La clasificación de suelo quedó reconciliada a Tipo E (consistente "
         "con las fundaciones) y la versión de NCh 427/1 unificada. "
         "Pendiente:",)])
    obs(doc, "OBS-01 — Protección superficial C5-M.",
        "No declara el sistema de protección anticorrosiva C5-M (ISO 12944) "
        "para ambiente marino.",
        "declarar el sistema C5-M o referir explícitamente a la "
        "Especificación de Estructura Metálica P22-ET-00-010-103-0.",
        "Términos de Referencia (protección C5-M, ISO 12944).")
    obs(doc, "OBS-02 — Anexos en revisión anterior y consistencia.",
        "Los anexos de cálculo conservan el cajetín \"Revisión B\" mientras "
        "el cuerpo es Rev 0; el coeficiente sísmico vertical figura como 0,74 "
        "en tabla y 0,67 en el texto; las cargas de pedestal/zapata usan un "
        "factor 1,4E no declarado en las combinaciones.",
        "reemitir los anexos a Rev 0, unificar el coeficiente vertical y "
        "declarar la combinación de cargas efectivamente usada; declarar los "
        "recubrimientos 50/70 mm.",
        "consistencia interna; Términos de Referencia, Sección 3.3.4.")

    # ---------- 2.2 Especificaciones ----------
    doc.add_heading("Especificaciones Técnicas", level=2)
    add_para(doc, [
        ("Las tres Especificaciones son Código 2 — Aprobado con comentarios. "
         "El carácter genérico (emplazamiento \"Antofagasta\" y código de "
         "portada) se corrige al particularizar las Especificaciones a Taltal, "
         "pero no es causal de rechazo. Comentarios a incorporar en Rev 0:",)])
    obs(doc, "OBS-01 (ET Movimiento de Tierra) — Mejoramiento de suelo.",
        "El mejoramiento/tratamiento de suelo queda diferido a obra, sin "
        "especificarse como condición de diseño.",
        "especificar las condiciones de mejoramiento/tratamiento de suelo "
        "(material, espesor, compactación) y particularizar la Especificación "
        "al emplazamiento de Taltal.",
        "Términos de Referencia, Secciones 1 y 3.3.2.")
    obs(doc, "OBS-02 (ET Obras Civiles) — Recubrimientos e impermeabilización.",
        "Los Términos de Referencia (Sección 3.3.4, Recubrimientos) exigen "
        "50 mm en elementos expuestos y 70 mm en elementos en contacto con el "
        "terreno; la Especificación declara solo 70 mm y no especifica la "
        "impermeabilización de elementos enterrados.",
        "declarar ambos recubrimientos (50 mm expuesto / 70 mm en contacto "
        "con terreno) y el sistema de impermeabilización de elementos "
        "enterrados.",
        "Términos de Referencia, Secciones 3.3.4 y 2.1.3.")
    obs(doc, "OBS-03 (ET Estructura Metálica) — Esquema C5-M y soldadura.",
        "Declara la clasificación C5-M pero sin el esquema por capas "
        "(preparación de superficie SSPC-SP10 + zinc + epóxico + poliuretano) "
        "y cita AWS A5.1/A5.17 (electrodos) en vez del código estructural "
        "AWS D1.1.",
        "completar el esquema de protección por capas con su preparación de "
        "superficie y citar AWS D1.1; exigir pernería galvanizada en caliente "
        "o inoxidable para ambiente expuesto.",
        "Términos de Referencia, Sección 3.4.2; ISO 12944; AWS D1.1.")

    # ---------- 2.3 Itemizados ----------
    doc.add_heading("Itemizados de Presupuesto", level=2)
    add_para(doc, [
        ("Los tres Itemizados son Código 3. El rótulo \"presupuesto "
         "referencial\" es aceptable; los comentarios apuntan a la clase de "
         "estimación, la trazabilidad de los tags y la completitud de las "
         "partidas.",)])
    obs(doc, "OBS-01 — Clase de estimación.",
        "Los Itemizados no declaran la clase de estimación.",
        "declarar la clase de estimación acordada en la reunión de arranque "
        "— Clase 2 AACE (Minuta 067-032-032-COR-MI-001), no la Clase 1 del "
        "Términos de Referencia.",
        "Minuta de reunión de arranque 067-032-032-COR-MI-001.")
    obs(doc, "OBS-02 — Tag de la fosa y pestaña ajena.",
        "La fosa de drenajes se rotula TK-06-002 (que es el estanque CIP de "
        "BW Water) y las planillas traen una pestaña \"Interconexión La "
        "Chimba\", ajena al proyecto.",
        "corregir el tag de la fosa a TK-06-004 y eliminar la pestaña "
        "\"La Chimba\".",
        "codificación del proyecto.")
    obs(doc, "OBS-03 — Partidas faltantes.",
        "El Itemizado de Obras Civiles no incluye la impermeabilización y el "
        "de Estructura Metálica no incluye la protección anticorrosiva ni la "
        "soldadura.",
        "incorporar las partidas de impermeabilización, protección "
        "superficial y soldadura.",
        "Términos de Referencia (partidas mínimas del presupuesto).")

    # ---------- 2.4 Planos ----------
    doc.add_heading("Planos de Obras Civiles", level=2)
    add_para(doc, [
        ("Comentarios a los planos (incorporar en la próxima revisión). Los "
         "identificadores OBS/NOTA coinciden con los del PDF anotado adjunto "
         "de cada lámina (CC_ADASA).",)])
    add_simple_table(doc, [
        ("Plano / Lámina", "ID", "Comentario"),
        ("P22-DWG-00-002-001 LAM1\nDisposición general", "NOTA-01",
         "Orientación del estanque sin fijar: agregar las coordenadas "
         "adicionales que definan el azimut del estanque."),
        ("P22-DWG-00-002-001 LAM1", "OBS-01",
         "Acotar el N.T.N. por zona, coincidente con los planos de montaje "
         "P22-DWG-06-005-101 / -103 / -105 y el Levantamiento DIO."),
        ("P22-DWG-00-002-002 LAM1\nEquipos exteriores", "OBS-01",
         "Agregar nota con las condiciones de mejoramiento / tratamiento de "
         "suelo (Especificación de Movimiento de Tierra)."),
        ("P22-DWG-00-002-002 LAM2", "OBS-01",
         "Rediseñar el anclaje del estanque a perno preinstalado colado en "
         "sitio (ASTM F1554), reconciliando con el plano Exfibro "
         "EX-26005-F01 Rev C (ACI 318-19 Capítulo 17 y Sección 17.10)."),
        ("P22-DWG-00-002-002 LAM4", "OBS-01",
         "Agregar nota con las condiciones de mejoramiento / tratamiento de "
         "suelo."),
        ("P22-DWG-00-002-002 LAM4", "OBS-02",
         "Anclaje de la bomba BH-06-001 figura postinstalado adhesivo: "
         "evaluar anclaje preinstalado o justificar el postinstalado según la "
         "categoría sísmica del proyecto (ACI 318-19 Capítulo 17)."),
        ("P22-DWG-00-002-003 LAM1\nFundación contenedor RO", "OBS-01",
         "Expresar las cotas en N.T.N. absoluto coincidente con el plano de "
         "montaje P22-DWG-06-005-103 (hoy en datum relativo)."),
        ("P22-DWG-00-002-003 LAM1", "OBS-02",
         "Agregar el callout que remite al detalle del inserto INS-1 "
         "(ubicado en LAM2)."),
        ("P22-DWG-00-002-003 LAM1", "NOTA-01",
         "Agregar nota con las condiciones de mejoramiento de suelo."),
        ("P22-DWG-00-002-004 LAM1\nFosa de drenajes", "OBS-01",
         "Agregar notas de mejoramiento de suelo e impermeabilización de los "
         "elementos enterrados."),
        ("P22-DWG-00-002-004 LAM1", "NOTA-01",
         "Corregir el tag de la fosa a TK-06-004 (TK-06-002 es el estanque "
         "CIP de BW Water)."),
        ("P22-DWG-00-002-004 LAM2", "OBS-01",
         "Especificar la parrilla en pultruida de PRFV, solo para tránsito "
         "liviano de personas."),
        ("P22-DWG-00-002-004 LAM2", "OBS-02",
         "Agregar y acotar las aperturas de rebalse del estanque."),
        ("P22-DWG-00-002-006 LAM1\nCanalizaciones y red de drenajes", "OBS-01",
         "Adoptar la nomenclatura CD-06-00N para las cámaras."),
        ("P22-DWG-00-002-006 LAM1", "NOTA-01",
         "Incorporar la lámina de detalle típico de cámara prefabricada y "
         "acotar las pendientes longitudinales."),
        ("P22-DWG-00-002-007 LAM1\nFundaciones equipos Sistema CIP", "OBS-01",
         "Agregar notas de mejoramiento de suelo e impermeabilización."),
        ("P22-DWG-00-002-007 LAM1", "OBS-02",
         "Reconciliar el N.T.N. +5,610 con el plano de montaje "
         "P22-DWG-06-005-103."),
        ("P22-DWG-00-002-007 LAM3", "OBS-01",
         "Agregar notas de mejoramiento de suelo e impermeabilización."),
        ("P22-DWG-00-003-001 LAM1\nCubierta metálica Sistema CIP", "OBS-01",
         "Declarar el sistema de protección C5-M referido a la Especificación "
         "de Estructura Metálica P22-ET-00-010-103-0 (ISO 12944)."),
    ])

    # =========================================================
    # 3. ADJUNTOS
    # =========================================================
    doc.add_heading("ADJUNTOS", level=1)
    add_para(doc, [
        ("Las memorias y planos con observaciones llevan PDF anotado "
         "(CC_ADASA): las cinco Memorias de Cálculo, las tres "
         "Especificaciones Técnicas y los siete planos de obras civiles con "
         "comentario.",)])

    # =========================================================
    # 4. RESUMEN DE RESPUESTA
    # =========================================================
    doc.add_heading("RESUMEN DE RESPUESTA", level=1)
    add_simple_table(doc, [
        ("Código Documento", "Título", "Rev", "Código de Respuesta"),
        ("P22-MC-00-002-001", "Fundación Estanque TK-06-001", "0",
         "3 — Por revisar"),
        ("P22-MC-00-002-002", "Fundación Dinámica Bomba BH-06-001", "0",
         "2 — Aprobado con comentarios"),
        ("P22-MC-00-002-004", "Fundación Contenedor RO", "0",
         "2 — Aprobado con comentarios"),
        ("P22-MC-00-002-005", "Fundación Sistema CIP", "0",
         "2 — Aprobado con comentarios"),
        ("P22-MC-00-003-001", "Cubierta Metálica Sistema CIP Exterior", "0",
         "3 — Por revisar"),
        ("P22-ET-00-010-101", "Especificación Movimiento de Tierra", "B",
         "2 — Aprobado con comentarios"),
        ("P22-ET-00-010-102", "Especificación Obras Civiles", "B",
         "2 — Aprobado con comentarios"),
        ("P22-ET-00-010-103", "Especificación Estructura Metálica", "B",
         "2 — Aprobado con comentarios"),
        ("P22-IT-00-010-101", "Itemizado Movimiento de Tierra", "B",
         "3 — Por revisar"),
        ("P22-IT-00-010-102", "Itemizado Obras Civiles", "B",
         "3 — Por revisar"),
        ("P22-IT-00-010-103", "Itemizado Estructura Metálica", "B",
         "3 — Por revisar"),
        ("P22-DWG-00-002-001", "Disposición general OOCC", "B",
         "3 — Por revisar"),
        ("P22-DWG-00-002-002", "Fundaciones equipos exteriores", "B",
         "3 — Por revisar"),
        ("P22-DWG-00-002-003", "Fundación contenedor RO", "B",
         "3 — Por revisar"),
        ("P22-DWG-00-002-004", "Fosa de drenajes", "B",
         "3 — Por revisar"),
        ("P22-DWG-00-002-006", "Canalizaciones y red de drenajes", "B",
         "3 — Por revisar"),
        ("P22-DWG-00-002-007", "Fundaciones equipos Sistema CIP", "B",
         "3 — Por revisar"),
        ("P22-DWG-00-003-001", "Cubierta metálica Sistema CIP", "B",
         "2 — Aprobado con comentarios"),
    ])
    add_para(doc, [
        ("Veredicto global del transmittal: 3 — POR REVISAR.",
         {"bold": True}),
        (" La mayoría de los planos y los tres Itemizados requieren revisión; "
         "las Especificaciones y tres de las Memorias se aprueban con "
         "comentarios a incorporar en Rev 0. No se entregaron el plano de "
         "Excavaciones (P22-DWG-00-001-001), el plano de Detalles de Anclaje "
         "(P22-DWG-00-002-005) ni la Memoria de Cálculo de Drenajes "
         "(P22-MC-00-002-003), comprometidos en el Listado de Entregables; "
         "ADASA los solicita. El N.T.N. de cada zona debe declararse "
         "coincidente con los planos de montaje y verificado contra el "
         "Levantamiento DIO Abr-2026.",),
    ])

    doc.save(OUTPUT)
    print(f"Generado: {OUTPUT}")


if __name__ == "__main__":
    main()
