#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera TRANSMITTAL N1 ADASA-LYA (stream OOCC, español, para L&A).
Codigo P22-TM-00-010-001-0. Entrega 3 (cover L&A 067-032-032-COR-TT-003).
Fecha: 19-May-2026

Version EJECUTIVA (comprimida in situ, estructura ADASA de 5 secciones).
Veredicto global: 3 - POR REVISAR.
Recuento: 3 Codigo 2 (MC-002, MC-004, MC-005) + 2 Codigo 3 (MC-001 anclaje
preinstalado ACI 318-19; MC-003-001 suelo D vs E + C5-M).
MC-005: anclaje postinstalado diferido por datos del proveedor (no rechazo)
-> Codigo 2; PEND-03 en Seccion 3. NOTA-07 dosificacion ~550 kg.
Comentarios como directivas de una linea (Defecto / Corregir: / Requisito:).
"""

import sys
import os

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))

from ejemplo_documento import (
    crear_documento_adasa,
    aplicar_arial_12,
    add_simple_table,
    add_bullet as _add_bullet_native,
)
from docx import Document


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "TRANSMITTAL N1 ADASA-LYA.docx")


def add_para(doc, runs):
    """runs = lista de (texto, kwargs) con kwargs en {'bold', 'italic'}."""
    para = doc.add_paragraph()
    for item in runs:
        text = item[0]
        attrs = item[1] if len(item) > 1 else {}
        run = para.add_run(text)
        if attrs.get("bold"):
            run.bold = True
        if attrs.get("italic"):
            run.italic = True
    aplicar_arial_12(para)
    return para


def obs(doc, idtit, defecto, corregir, requisito=None):
    """Observacion ejecutiva de una linea."""
    runs = [(idtit + " ", {"bold": True}), (defecto + " ",),
            ("Corregir: ", {"bold": True}), (corregir,)]
    if requisito:
        runs += [(" ",), ("Requisito: ", {"bold": True}), (requisito,)]
    add_para(doc, runs)


def add_bullet(doc, text):
    _add_bullet_native(doc, text, size=11, space_after_pt=12)


def main() -> None:
    crear_documento_adasa(
        titulo=("TRANSMITTAL DE REVISIÓN TÉCNICA N1 — INGENIERÍA DE "
                "DETALLE OOCC MÓDULO RO 2DA ETAPA SALMUERA TALTAL"),
        codigo="P22-TM-00-010-001-0",
        output_filename=OUTPUT,
        incluir_toc=True,
    )

    doc = Document(OUTPUT)

    # =========================================================
    # 1. RESUMEN EJECUTIVO
    # =========================================================
    doc.add_heading("RESUMEN EJECUTIVO", level=1)

    add_para(doc, [
        ("VEREDICTO DEL TRANSMITTAL: 3 — POR REVISAR.", {"bold": True}),
        (" Cinco Memorias de Cálculo Rev B. Recuento: 3 Código 2, 2 Código "
         "3. Lo fijan la Fundación del Estanque y la Cubierta Metálica; "
         "las tres restantes se aprueban con comentarios a incorporar "
         "en Rev 0.",),
    ])
    add_simple_table(doc, [
        ("Documento", "Código", "Acción"),
        ("P22-MC-00-002-001 Fundación Estanque TK-06-001",
         "3 — Por revisar",
         "Rev B.1: rediseñar anclaje preinstalado (colado en sitio) "
         "ACI 318-19; declarar NPT y verificar contra Levantamiento "
         "DIO"),
        ("P22-MC-00-002-002 Fundación Dinámica Bomba BH-06-001",
         "2 — Aprobado con comentarios",
         "Rev 0: NPT contra DIO; peso BH-06-001 según ficha técnica "
         "KSB + rehacer ACI 351.3R; recubrimientos"),
        ("P22-MC-00-002-004 Fundación Contenedor RO",
         "2 — Aprobado con comentarios",
         "Rev 0: reconciliar código/alcance; retirar cargas CIP (van "
         "en MC-005); peso contenedor 14.934 kg; NPT contra DIO"),
        ("P22-MC-00-002-005 Fundación Sistema CIP",
         "2 — Aprobado con comentarios",
         "Rev 0: anclajes postinstalados diferidos por datos del "
         "proveedor (Sección 3); reconciliar código; NPT contra DIO; "
         "pesos dosificación ≈550 kg"),
        ("P22-MC-00-003-001 Cubierta Metálica Sistema CIP",
         "3 — Por revisar",
         "Rev B.1: reconciliar clasificación de suelo (D vs E); "
         "declarar protección C5-M"),
    ])
    add_para(doc, [
        ("Determinantes del Código 3: ", {"bold": True}),
        ("P22-MC-00-002-001 (anclaje preinstalado ACI 318-19) y "
         "P22-MC-00-003-001 (clasificación de suelo D vs E; protección "
         "C5-M).",),
    ])
    add_para(doc, [
        ("Transversal: ", {"bold": True}),
        ("ninguna memoria declara el NPT (usan \"Modelo Navis "
         "Referencial\", no vinculante); el TR lo hace dato fijo "
         "verificable contra el Levantamiento DIO Abr-2026 antes de "
         "Rev A. Pendientes y dependencias (MC de Drenajes no "
         "entregada; antecedente Exfibro; datos del proveedor de "
         "equipos CIP) en la Sección 3.",),
    ])

    # =========================================================
    # 2. OBSERVACIONES POR DOCUMENTO
    # =========================================================
    doc.add_heading("OBSERVACIONES POR DOCUMENTO", level=1)

    # ---------- 2.1 MC-001 (Código 3) ----------
    doc.add_heading(
        "Fundación Estanque TK-06-001 Rev B — P22-MC-00-002-001",
        level=2)
    add_para(doc, [
        ("Código de Respuesta: 3 — Por revisar.", {"bold": True}),
        (" Losa de fundación individual del estanque PRFV TK-06-001; "
         "presiones de contacto estática 0,40 / dinámica 0,47 kgf/cm² "
         "bajo el admisible, 100 % de apoyo. El veredicto lo fija el "
         "diseño del anclaje (OBS-01). Comentarios detallados: "
         "P22-MC-00-002-001_B_CC_ADASA.pdf.",),
    ])
    add_simple_table(doc, [
        ("ID", "Severidad", "Tema"),
        ("OBS-01", "Mayor",
         "Anclaje definido como postinstalado; rediseñar preinstalado "
         "(colado en sitio) ACI 318-19"),
        ("OBS-02", "Mayor",
         "NPT no declarado ni verificado contra planos de montaje / "
         "Levantamiento DIO"),
        ("NOTA-01", "Mayor",
         "Anclajes y hormigón armado con ACI 318-14; el TR exige ACI "
         "318-19"),
        ("NOTA-02", "Menor",
         "Recubrimientos mínimos (50/70 mm) no declarados"),
        ("NOTA-03", "Menor",
         "Origen de las reacciones basales no declarado"),
        ("NOTA-04", "Menor",
         "Base normativa sísmica: la memoria lista NCh 2369:2003 y "
         ":2025"),
    ])
    obs(doc, "OBS-01 — Anclaje del estanque.",
        "La memoria especifica ASTM F1554 Gr.36 con cabeza hexagonal "
        "pesada y empotramiento 30 cm (= anclaje preinstalado) pero lo "
        "verifica como mecánico postinstalado (PROFIS, ACI 318-14): "
        "contradicción; el postinstalado es inviable sobre losa con "
        "doble malla Ø16@150 y patrón de pernos fijo del estanque.",
        "rediseñar como pernos preinstalados con cabeza (colados en "
        "sitio, ASTM F1554) verificados por ACI 318-19 Capítulo 17 y "
        "Sección 17.10, con plantilla coordinada con la enferradura, "
        "reconciliando cantidad/diámetro/círculo con el plano Exfibro "
        "EX-26005-F01 Rev C.",
        "TR (ACI 318-19); ACI 318-19 Capítulo 17 y Sección 17.10; "
        "plano Exfibro EX-26005-F01 Rev C.")
    obs(doc, "OBS-02 — NPT no declarado ni verificado.",
        "La memoria no declara el NPT ni las cotas; usa solo el "
        "\"Modelo Navis Referencial\" (de consulta; prevalecen los "
        "planos 2D).",
        "declarar el NPT de los planos de montaje P22-DWG-06-005-103 / "
        "P22-DWG-06-005-101, referir a él cotas de "
        "fundación/pedestales/pendientes, reflejarlo en "
        "P22-DWG-00-002-001 y verificarlo contra el Levantamiento DIO "
        "Abr-2026 antes de Rev A; declarar la cota de la zona del "
        "estanque (N.T.N. ≈ +5,75).",
        "TR Secciones 3.3.1, 4.1 y 2.1.3.")
    obs(doc, "NOTA-01 — Versión de ACI 318.",
        "Anclajes y hormigón armado verificados con ACI 318-14.",
        "ejecutar el anclaje rediseñado directamente con ACI 318-19; "
        "para el resto del hormigón armado, demostrar "
        "equivalencia/conservadurismo o re-verificar con ACI 318-19.",
        "TR (ACI 318-19).")
    obs(doc, "NOTA-02 — Recubrimientos.",
        "No se declaran los recubrimientos mínimos.",
        "declarar 50 mm en elementos expuestos y 70 mm en contacto con "
        "el terreno.",
        "TR (ambiente marino corrosivo).")
    obs(doc, "NOTA-03 — Origen de las reacciones basales.",
        "La memoria no declara la fuente.",
        "emplear y declarar las reacciones de la versión corregida por "
        "ADASA del antecedente Exfibro — Ez = ±3.357 kgf, momento "
        "volcante 431.846 kgf·cm, cortante basal 4.617 kgf, peso "
        "propio 585 kgf, líquido 11.986 kgf — citando "
        "P22-IT-06-000-005-0; la ratificación Exfibro Rev B la "
        "gestiona ADASA (Sección 3).",
        "P22-IT-06-000-005-0; antecedente Exfibro EX-26005-F01 Rev C.")
    obs(doc, "NOTA-04 — Base normativa sísmica.",
        "Lista NCh 2369:2003 y :2025 y adopta :2025.",
        "declarar que las fuerzas sísmicas heredadas del antecedente "
        "(base 2003) quedan acotadas o son conservadoras frente a NCh "
        "2369:2025, o reconciliarlas.",
        "TR (NCh 2369:2025).")

    # ---------- 2.2 MC-002 (Código 2) ----------
    doc.add_heading(
        "Fundación Dinámica Bomba BH-06-001 Rev B — "
        "P22-MC-00-002-002", level=2)
    add_para(doc, [
        ("Código de Respuesta: 2 — Aprobado con comentarios.",
         {"bold": True}),
        (" Losa dinámica de la bomba KSB con el marco ACI 351.3R; "
         "presiones 0,17 / 0,29 kgf/cm², apoyo 100 % / 92,38 %, "
         "conformes. Comentarios detallados: "
         "P22-MC-00-002-002_B_CC_ADASA.pdf.",),
    ])
    add_simple_table(doc, [
        ("ID", "Severidad", "Tema"),
        ("OBS-01", "Mayor",
         "NPT no declarado ni verificado contra planos de montaje / "
         "Levantamiento DIO"),
        ("NOTA-01", "Mayor",
         "Verificación con ACI 318-14; el TR exige ACI 318-19"),
        ("NOTA-02", "Menor", "Recubrimientos mínimos no declarados"),
        ("NOTA-05", "Menor",
         "Base normativa sísmica: la memoria lista NCh 2369:2003 y "
         ":2025"),
        ("NOTA-06", "Menor",
         "Resultados de masa/frecuencia presentados solo en figuras"),
        ("NOTA-07", "Mayor", "Peso de BH-06-001 no declarado"),
    ])
    obs(doc, "OBS-01 — NPT no declarado ni verificado.",
        "No declara el NPT; usa solo el \"Modelo Navis Referencial\".",
        "declarar el NPT de P22-DWG-06-005-101 / P22-DWG-06-005-103, "
        "referir cotas de fundación/pedestal, reflejarlo en "
        "P22-DWG-00-002-001 y verificarlo contra el Levantamiento DIO "
        "Abr-2026 antes de Rev A; declarar la cota de la zona de la "
        "bomba (N.T.N. ≈ +5,75).",
        "TR Secciones 3.3.1, 4.1 y 2.1.3.")
    obs(doc, "NOTA-01 — Versión de ACI 318.",
        "Verificación con ACI 318-14.",
        "demostrar equivalencia/conservadurismo o re-verificar con "
        "ACI 318-19.",
        "TR (ACI 318-19).")
    obs(doc, "NOTA-02 — Recubrimientos.",
        "No se declaran los recubrimientos mínimos.",
        "declarar 50 mm expuesto / 70 mm en contacto con terreno.")
    obs(doc, "NOTA-05 — Base normativa sísmica.",
        "Lista NCh 2369:2003 y :2025.",
        "declarar que las fuerzas (base 2003) son "
        "acotadas/conservadoras frente a NCh 2369:2025, o "
        "reconciliarlas.")
    obs(doc, "NOTA-06 — Margen masa/frecuencia.",
        "Criterios (relación ≥3; frecuencia fuera de ±20 %) solo en "
        "figuras.",
        "reportar los valores numéricos efectivos (relación de masa y "
        "margen de frecuencia).")
    obs(doc, "NOTA-07 — Peso de BH-06-001.",
        "La memoria no declara el peso de equipo adoptado.",
        "emplear y declarar el peso operativo (bomba + motor + base) "
        "de la ficha técnica del fabricante KSB (KNCPP 11-050+160M); "
        "ADASA constata divergencia rótulo montaje (250 kg) vs "
        "estimación TR (~400 kg); resolverla con la ficha y rehacer "
        "ACI 351.3R (relación masa ≥3; separación de frecuencia). Si "
        "no satisface ACI 351.3R, escala a Código 3.",
        "ficha técnica KSB KNCPP 11-050+160M; ACI 351.3R.")
    add_para(doc, [
        ("Acción para emitir en Rev 0 — sin nueva Rev B:",
         {"bold": True}),
        (" (1) NPT declarado y verificado contra el Levantamiento DIO "
         "antes de Rev A. (2) Peso BH-06-001 según ficha técnica KSB; "
         "rehacer ACI 351.3R. (3) Reportar relación de masa y margen "
         "de frecuencia. (4) Recubrimientos. (5) Trazabilidad "
         "normativa (ACI 318, base sísmica).",),
    ])

    # ---------- 2.3 MC-004 (Código 2) ----------
    doc.add_heading(
        "Fundación Contenedor RO Rev B — P22-MC-00-002-004",
        level=2)
    add_para(doc, [
        ("Código de Respuesta: 2 — Aprobado con comentarios.",
         {"bold": True}),
        (" Fundación independiente del contenedor RO (losa, vigas, "
         "pedestales, inserto embebido); presiones 0,44 / 0,63 "
         "kgf/cm², apoyo 100 %. Observaciones de codificación/"
         "trazabilidad documental. Comentarios detallados: "
         "P22-MC-00-002-004_B_CC_ADASA.pdf.",),
    ])
    add_simple_table(doc, [
        ("ID", "Severidad", "Tema"),
        ("OBS-01", "Mayor",
         "Codificación/alcance: título cuerpo vs cover vs TR no "
         "coinciden; incluye cargas CIP"),
        ("OBS-02", "Menor",
         "Inconsistencia interna de versión de ACI 318 (texto cita "
         "318-10 y 318-14)"),
        ("OBS-03", "Mayor",
         "NPT no declarado ni verificado contra planos de montaje / "
         "Levantamiento DIO"),
        ("NOTA-01", "Mayor",
         "Verificación con ACI 318-14; el TR exige ACI 318-19"),
        ("NOTA-02", "Menor", "Recubrimientos mínimos no declarados"),
        ("NOTA-08", "Menor",
         "Peso del contenedor RO no declarado"),
    ])
    obs(doc, "OBS-01 — Codificación y alcance.",
        "Cuerpo \"Fundación Contenedor RO\" vs cover \"Fundación "
        "Compartida Contenedor RO + Sistema CIP\" vs TR Sección 6 "
        "\"Fundación compartida\"; además incluye cargas del "
        "estanque/equipo de dosificación (Sistema CIP).",
        "reconciliar título/código/alcance y formalizar el split y el "
        "código P22-MC-00-002-005 vía la lista "
        "067-032-032-COR-LI-001; retirar de esta memoria (solo "
        "contenedor) las cargas de dosificación — se verifican en "
        "MC-005 con los pesos del TR Sección 2.1.2 (TK-09-002 = 490 "
        "kg; BDS-09-001/002 = 57,4 kg). El enfoque de independencia no "
        "requiere cambios.",
        "TR Sección 6; Minuta 067-032-032-COR-MI-001 ítem 1.7.")
    obs(doc, "OBS-02 — Versión interna de ACI 318.",
        "El diseño de pedestales cita ACI 318-10 y el listado de "
        "códigos ACI 318-14.",
        "unificar la versión y reconciliarla con el TR (ACI 318-19, "
        "ver NOTA-01).",
        "TR (ACI 318-19).")
    obs(doc, "OBS-03 — NPT no declarado ni verificado.",
        "No declara el NPT; usa solo el \"Modelo Navis Referencial\".",
        "declarar el NPT de P22-DWG-06-005-103 / P22-DWG-06-005-101, "
        "referir cotas de fundación/pedestales/vigas, reflejarlo en "
        "P22-DWG-00-002-001 y verificarlo contra el Levantamiento DIO "
        "Abr-2026 antes de Rev A; declarar la cota de la zona del "
        "contenedor (N.T.N. ≈ +6,00).",
        "TR Secciones 3.3.1, 4.1 y 2.1.3.")
    obs(doc, "NOTA-01 — Versión de ACI 318.",
        "Verificación con ACI 318-14.",
        "demostrar equivalencia/conservadurismo o re-verificar con "
        "ACI 318-19.")
    obs(doc, "NOTA-02 — Recubrimientos.",
        "No se declaran los recubrimientos mínimos.",
        "declarar 50 mm expuesto / 70 mm en contacto con terreno.")
    obs(doc, "NOTA-08 — Peso del contenedor.",
        "La memoria no explicita el peso del contenedor RO.",
        "emplear y declarar el peso vinculante de 14.934 kg (casco ≥ "
        "5.000 kg + equipos internos 9.934 kg) según el TR Sección "
        "2.1.2 y el plano civil y de cargas P22-DWG-09-005-001 Rev A.",
        "TR Sección 2.1.2; plano P22-DWG-09-005-001 Rev A.")
    add_para(doc, [
        ("Acción para emitir en Rev 0 — sin nueva Rev B:",
         {"bold": True}),
        (" (1) NPT declarado y verificado contra el Levantamiento DIO "
         "antes de Rev A. (2) Reconciliar título/código/alcance; "
         "retirar cargas CIP; formalizar el split en la lista de "
         "entregables. (3) Unificar versión de ACI 318. (4) Declarar "
         "peso del contenedor 14.934 kg. (5) Recubrimientos.",),
    ])

    # ---------- 2.4 MC-005 (Código 2) ----------
    doc.add_heading(
        "Fundación Sistema CIP Rev B — P22-MC-00-002-005", level=2)
    add_para(doc, [
        ("Código de Respuesta: 2 — Aprobado con comentarios.",
         {"bold": True}),
        (" Fundaciones independientes de los equipos CIP y el estanque "
         "CIP; presiones 0,41 / 0,45 kgf/cm², apoyo 100 %, conformes. "
         "Observaciones de codificación, NPT y datos pendientes del "
         "proveedor. Comentarios detallados: "
         "P22-MC-00-002-005_B_CC_ADASA.pdf.",),
    ])
    add_simple_table(doc, [
        ("ID", "Severidad", "Tema"),
        ("OBS-01", "Menor",
         "Verificación de anclajes diferida por falta de datos "
         "finales del proveedor de equipos CIP"),
        ("OBS-02", "Mayor",
         "Código P22-MC-00-002-005 no previsto en los Términos de "
         "Referencia"),
        ("OBS-03", "Mayor",
         "NPT no declarado ni verificado contra planos de montaje / "
         "Levantamiento DIO"),
        ("NOTA-01", "Mayor",
         "Verificación con ACI 318-14; el TR exige ACI 318-19"),
        ("NOTA-02", "Menor", "Recubrimientos mínimos no declarados"),
        ("NOTA-05", "Menor",
         "Base normativa sísmica: la memoria lista NCh 2369:2003 y "
         ":2025"),
        ("NOTA-07", "Menor",
         "Pesos de dosificación asumidos en 500 + 500 kg (≈1000 kg)"),
    ])
    obs(doc, "OBS-01 — Verificación de anclajes diferida.",
        "La verificación de anclajes de los equipos y del estanque "
        "CIP queda diferida porque aún no se dispone de los planos/"
        "fichas finales del proveedor de equipos CIP; los anclajes "
        "serán postinstalados. No es motivo de rechazo: el diseño es "
        "aceptable si los anclajes cumplen las solicitaciones.",
        "completar la verificación con anclajes postinstalados "
        "dimensionados a las cargas de los equipos CIP, según ACI "
        "318-19 Capítulo 17 y Sección 17.10, precalificados para uso "
        "sísmico (ACI 355.2/355.4 Cat. 1, reporte de evaluación "
        "ICC-ES vigente), cuando se disponga de los datos del "
        "proveedor (dependencia trackeada, Sección 3 PEND-03).",
        "TR (diseño de anclajes); ACI 318-19; ACI 355.2/355.4.")
    obs(doc, "OBS-02 — Codificación.",
        "El código P22-MC-00-002-005 no está previsto en el TR "
        "Sección 6 (fundación CIP dentro de P22-MC-00-002-004 "
        "\"Fundación compartida\").",
        "formalizar el split a fundaciones independientes y el nuevo "
        "código vía la lista 067-032-032-COR-LI-001 (ver subsección "
        "2.3, OBS-01).",
        "TR Sección 6.")
    obs(doc, "OBS-03 — NPT no declarado ni verificado.",
        "No declara el NPT; usa solo el \"Modelo Navis Referencial\".",
        "declarar el NPT de P22-DWG-06-005-103 / P22-DWG-06-005-101, "
        "referir cotas de fundación/pedestales, reflejarlo en "
        "P22-DWG-00-002-001 y verificarlo contra el Levantamiento DIO "
        "Abr-2026 antes de Rev A; declarar la cota de la zona del "
        "Sistema CIP (N.T.N. ≈ +6,00).",
        "TR Secciones 3.3.1, 4.1 y 2.1.3.")
    obs(doc, "NOTA-01 — Versión de ACI 318.",
        "Verificación con ACI 318-14.",
        "ejecutar la verificación de anclajes con ACI 318-19; "
        "demostrar equivalencia o re-verificar el resto del hormigón "
        "armado con ACI 318-19.")
    obs(doc, "NOTA-02 — Recubrimientos.",
        "No se declaran los recubrimientos mínimos.",
        "declarar 50 mm expuesto / 70 mm en contacto con terreno.")
    obs(doc, "NOTA-05 — Base normativa sísmica.",
        "Lista NCh 2369:2003 y :2025.",
        "declarar que las fuerzas (base 2003) son "
        "acotadas/conservadoras frente a NCh 2369:2025, o "
        "reconciliarlas.")
    obs(doc, "NOTA-07 — Pesos de dosificación.",
        "La memoria asume 500 kg + 500 kg (≈1000 kg total).",
        "declarar los pesos correctos del TR Sección 2.1.2 — "
        "TK-09-002 = 490 kg y BDS-09-001/002 = 57,4 kg (≈550 kg "
        "total). El diseño actual (≈1000 kg) es conservador y no "
        "requiere recálculo, pero el valor declarado debe corregirse "
        "a ≈550 kg.",
        "TR Sección 2.1.2.")
    add_para(doc, [
        ("Acción para emitir en Rev 0 — sin nueva Rev B:",
         {"bold": True}),
        (" (1) NPT declarado y verificado contra el Levantamiento DIO "
         "antes de Rev A. (2) Formalizar el split y el código en la "
         "lista de entregables. (3) Completar la verificación de "
         "anclajes postinstalados cuando lleguen los datos del "
         "proveedor (Sección 3 PEND-03). (4) Corregir pesos de "
         "dosificación a ≈550 kg. (5) Recubrimientos; trazabilidad "
         "normativa.",),
    ])

    # ---------- 2.5 MC-003-001 (Código 3) ----------
    doc.add_heading(
        "Cubierta Metálica Sistema CIP Exterior Rev B — "
        "P22-MC-00-003-001", level=2)
    add_para(doc, [
        ("Código de Respuesta: 3 — Por revisar.", {"bold": True}),
        (" Estructura metálica de cobertizo de los equipos CIP "
         "exteriores y sus fundaciones. Verificaciones estructurales "
         "conformes (FU máx 57,5 %; deformaciones OK; "
         "zapata/pedestal/anclaje; contacto 0,34 / 0,69 kgf/cm²; 87 % "
         "en condición sísmica). El veredicto lo fija una "
         "inconsistencia en la base de diseño. Comentarios "
         "detallados: P22-MC-00-003-001_B_CC_ADASA.pdf.",),
    ])
    add_simple_table(doc, [
        ("ID", "Severidad", "Tema"),
        ("OBS-01", "Mayor",
         "Clasificación de suelo tipo D, inconsistente con el tipo E "
         "de las memorias de fundaciones"),
        ("OBS-02", "Mayor",
         "Sistema de protección superficial C5-M no abordado"),
        ("OBS-03", "Menor",
         "Inconsistencia interna de versión de NCh 427/1"),
        ("NOTA-01", "Mayor",
         "Verificación con ACI 318-14; el TR exige ACI 318-19"),
        ("NOTA-02", "Menor", "Recubrimientos mínimos no declarados"),
        ("NOTA-09", "Menor",
         "Verificaciones estructurales conformes; observaciones "
         "acotadas"),
    ])
    obs(doc, "OBS-01 — Clasificación de suelo.",
        "Adopta suelo tipo D y coeficiente sísmico vertical 0,74; las "
        "cuatro memorias de fundaciones adoptan suelo tipo E y "
        "vertical 0,45 para el mismo emplazamiento.",
        "reconciliar la clasificación de suelo del sitio (única y "
        "consistente entre todas las memorias, conforme al informe "
        "sísmico del emplazamiento) y re-verificar la demanda sísmica "
        "de la cubierta con el tipo de suelo que corresponda.",
        "informe sísmico del emplazamiento; NCh 2369:2025.")
    obs(doc, "OBS-02 — Protección superficial.",
        "No aborda el sistema de pintura/protección C5-M (ISO 12944) "
        "exigido para estructura metálica en ambiente marino.",
        "declarar el sistema C5-M en la memoria o referir "
        "explícitamente a la Especificación Técnica de Estructura "
        "Metálica P22-ET-00-010-103-0.",
        "TR (protección C5-M, ISO 12944).")
    obs(doc, "OBS-03 — Versión de NCh 427/1.",
        "Normativa cita NCh 427/1 2006 y el análisis NCh 427/1 2016 "
        "(AISC 360-2016).",
        "unificar la versión efectivamente aplicada.",
        "TR (normativa estructural).")
    obs(doc, "NOTA-01 — Versión de ACI 318.",
        "Verificación con ACI 318-14.",
        "demostrar equivalencia o re-verificar las fundaciones de la "
        "cubierta con ACI 318-19.")
    obs(doc, "NOTA-02 — Recubrimientos.",
        "No se declaran los recubrimientos mínimos.",
        "declarar 50 mm expuesto / 70 mm en contacto con terreno para "
        "zapata y pedestal.")
    add_para(doc, [
        ("NOTA-09 — Conformidad estructural. ", {"bold": True}),
        ("Resistencia, deformaciones y estabilidad de las fundaciones "
         "conformes; las observaciones se acotan a clasificación de "
         "suelo, protección superficial y consistencia normativa.",),
    ])

    # =========================================================
    # 3. OBSERVACIONES PENDIENTES Y DEPENDENCIAS
    # =========================================================
    doc.add_heading(
        "OBSERVACIONES PENDIENTES Y DEPENDENCIAS", level=1)
    add_para(doc, [(
        "Primer transmittal del stream OOCC; sin observaciones "
        "heredadas. Elementos abiertos:",)])
    add_simple_table(doc, [
        ("Origen", "Documento", "Materia", "Estado"),
        ("PEND-01",
         "Sistema de Drenajes — P22-MC-00-002-003 y planos de "
         "fosa/canalizaciones",
         "Entregable del TR Sección 6 no incluido en ENTREGA 3",
         "PENDIENTE INFORMATIVO — se espera en entrega posterior; no "
         "condiciona ENTREGA 3"),
        ("PEND-02",
         "Fundación Estanque TK-06-001 — P22-MC-00-002-001",
         "Las reacciones basales dependen de la ratificación del "
         "antecedente Exfibro (Memoria Rev B), gestionada por ADASA "
         "vía consulta técnica P22-CT-06-000-002-0",
         "ABIERTO — ADASA gestiona; L&A ya adoptó los valores "
         "corregidos"),
        ("PEND-03",
         "Fundación Sistema CIP — P22-MC-00-002-005",
         "La verificación de anclajes (postinstalados) depende de "
         "los planos/fichas finales del proveedor de equipos CIP, no "
         "disponibles",
         "ABIERTO — completar con anclajes postinstalados que cumplan "
         "las solicitaciones cuando se disponga de los datos; "
         "en seguimiento hasta Rev 0"),
    ])
    add_para(doc, [
        ("Seguimiento para Rev 0", {"bold": True}),
        (" (memorias Código 2; sin nueva revisión intermedia): "
         "P22-MC-00-002-002, P22-MC-00-002-004 y P22-MC-00-002-005 "
         "según sus bloques \"Acción para emitir en Rev 0\".",),
    ])

    # =========================================================
    # 4. ADJUNTOS
    # =========================================================
    doc.add_heading("ADJUNTOS", level=1)
    add_simple_table(doc, [
        ("Documento", "Archivo Anotado", "Anotaciones"),
        ("Fundación Estanque TK-06-001 Rev B",
         "P22-MC-00-002-001_B_CC_ADASA.pdf",
         "OBS-01, OBS-02, NOTA-01, NOTA-02, NOTA-03, NOTA-04"),
        ("Fundación Dinámica Bomba BH-06-001 Rev B",
         "P22-MC-00-002-002_B_CC_ADASA.pdf",
         "OBS-01, NOTA-01, NOTA-02, NOTA-05, NOTA-06, NOTA-07"),
        ("Fundación Contenedor RO Rev B",
         "P22-MC-00-002-004_B_CC_ADASA.pdf",
         "OBS-01, OBS-02, OBS-03, NOTA-01, NOTA-02, NOTA-08"),
        ("Fundación Sistema CIP Rev B",
         "P22-MC-00-002-005_B_CC_ADASA.pdf",
         "OBS-01, OBS-02, OBS-03, NOTA-01, NOTA-02, NOTA-05, NOTA-07"),
        ("Cubierta Metálica Sistema CIP Rev B",
         "P22-MC-00-003-001_B_CC_ADASA.pdf",
         "OBS-01, OBS-02, OBS-03, NOTA-01, NOTA-02, NOTA-09"),
    ])
    add_para(doc, [(
        "Las cinco memorias llevan PDF anotado con sus observaciones "
        "y notas.",)])

    # =========================================================
    # 5. RESUMEN DE RESPUESTA
    # =========================================================
    doc.add_heading("RESUMEN DE RESPUESTA", level=1)
    add_simple_table(doc, [
        ("Código Documento", "Título", "Rev", "Código de Respuesta"),
        ("P22-MC-00-002-001", "Fundación Estanque TK-06-001", "B",
         "3 — Por revisar"),
        ("P22-MC-00-002-002", "Fundación Dinámica Bomba BH-06-001", "B",
         "2 — Aprobado con comentarios"),
        ("P22-MC-00-002-004", "Fundación Contenedor RO", "B",
         "2 — Aprobado con comentarios"),
        ("P22-MC-00-002-005", "Fundación Sistema CIP", "B",
         "2 — Aprobado con comentarios"),
        ("P22-MC-00-003-001",
         "Cubierta Metálica Sistema CIP Exterior", "B",
         "3 — Por revisar"),
    ])
    add_para(doc, [
        ("Veredicto global del transmittal: 3 — POR REVISAR.",
         {"bold": True}),
        (" Lo fijan P22-MC-00-002-001 (anclaje preinstalado ACI "
         "318-19) y P22-MC-00-003-001 (clasificación de suelo D vs E; "
         "protección C5-M). Las tres restantes son Código 2 — "
         "Aprobado con comentarios, con puntos a incorporar en Rev 0 "
         "sin nueva revisión intermedia; la verificación del NPT "
         "contra el Levantamiento DIO Abr-2026 es condición previa a "
         "la Revisión A.",),
    ])

    doc.save(OUTPUT)
    print(f"Generado: {OUTPUT}")


if __name__ == "__main__":
    main()
