# -*- coding: utf-8 -*-
"""
Script para generar Informe Comparativo de Instrumentacion
Conductividad y Caudal - Proyecto BAE 12803 Taltal UHPRO

Codigo: P22-CD-09-008-001-0
Fecha: 28-Ene-2026

Usa template ADASA v6.0
"""

import sys
import os

# Agregar path al skill template-adasa
skill_path = os.path.join(
    os.path.dirname(__file__), "..", "..", "..", ".claude", "skills", "template-adasa"
)
sys.path.insert(0, os.path.abspath(skill_path))

from ejemplo_documento import crear_documento_adasa
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


def set_table_borders(table):
    """Aplica bordes a todas las celdas de una tabla y la centra en el documento"""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tblBorders = parse_xml(
        r'<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        r'<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r"</w:tblBorders>"
    )
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = parse_xml(
            r'<w:tblPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
        )
        tbl.insert(0, tblPr)
    tblPr.append(tblBorders)


def set_repeat_table_header(row):
    """Configura una fila para repetirse en cada pagina"""
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    tblHeader = parse_xml(
        r'<w:tblHeader xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
    )
    trPr.append(tblHeader)


def add_header_row(table, texts, bold=True):
    """Agrega fila de encabezado con formato"""
    row = table.rows[0]
    for i, text in enumerate(texts):
        cell = row.cells[i]
        cell.text = text
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.runs[0] if para.runs else para.add_run(text)
        run.bold = bold
        run.font.size = Pt(10)
        run.font.name = "Arial"
        # Fondo gris claro
        shading = parse_xml(r'<w:shd {} w:fill="D9D9D9"/>'.format(nsdecls("w")))
        cell._tc.get_or_add_tcPr().append(shading)
    set_repeat_table_header(row)


def add_data_row(table, texts, highlight=False, highlight_color="FFFF00"):
    """Agrega fila de datos"""
    row = table.add_row()
    for i, text in enumerate(texts):
        cell = row.cells[i]
        cell.text = str(text)
        para = cell.paragraphs[0]
        if para.runs:
            run = para.runs[0]
        else:
            run = para.add_run(str(text))
        run.font.size = Pt(9)
        run.font.name = "Arial"
        if highlight:
            shading = parse_xml(
                r'<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), highlight_color)
            )
            cell._tc.get_or_add_tcPr().append(shading)
    return row


def main():
    # Crear documento base usando template ADASA
    output_file = "P22-CD-09-008-001-0_Informe-Conductividad-Caudal.docx"

    crear_documento_adasa(
        titulo="INFORME COMPARATIVO\nINSTRUMENTACION CONDUCTIVIDAD Y CAUDAL",
        codigo="P22-CD-09-008-001-0",
        preparado_por="ADASA - Ingenieria de Proyectos",
        revisado_por="Luis Riquelme",
        output_filename=output_file,
    )

    # Abrir el documento generado para modificarlo
    doc = Document(output_file)

    # Limpiar contenido de ejemplo generado por template
    paragraphs_to_remove = []
    for i, para in enumerate(doc.paragraphs):
        if (
            "RESUMEN EJECUTIVO" in para.text
            or "INTRODUCCION" in para.text
            or "INTRODUCCIÓN" in para.text
            or "1. INTRO" in para.text
        ):
            paragraphs_to_remove.append(para)
        elif (
            "Este documento ha sido generado" in para.text
            or "Agregue aqui el contenido" in para.text
            or "Agregue aquí el contenido" in para.text
        ):
            paragraphs_to_remove.append(para)

    for para in paragraphs_to_remove:
        p = para._element
        p.getparent().remove(p)

    # =========================================================================
    # 1. RESUMEN EJECUTIVO
    # =========================================================================
    h1 = doc.add_heading("1. RESUMEN EJECUTIVO", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    doc.add_heading("1.1 Objetivo", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "Verificar el cumplimiento de la instrumentacion de conductividad y caudal propuesta por BW Water contra los requisitos contractuales establecidos en las Especificaciones Tecnicas (ET) del proyecto BAE 12803 - Planta Modular UHPRO Taltal."
    )

    doc.add_heading("1.2 Hallazgos Principales", level=2)

    # Tabla resumen
    table = doc.add_table(rows=1, cols=5)
    set_table_borders(table)
    add_header_row(
        table, ["Variable", "ET Requiere", "Oferta", "Instrument List", "Estado"]
    )
    add_data_row(
        table,
        ["Conductividad", "5 ubicaciones", "1 instrumento", "5 instrumentos", "CUMPLE"],
    )
    add_data_row(
        table,
        ["Caudal", "5 ubicaciones", "4 instrumentos", "4 TAGs unicos", "TAG DUPLICADO"],
        highlight=True,
        highlight_color="FFCCCC",
    )
    doc.add_paragraph()

    doc.add_heading("1.3 Veredicto", level=2)
    p = doc.add_paragraph()
    run = p.add_run("3 - TO BE REVISED")
    run.bold = True
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.font.size = Pt(14)

    p = doc.add_paragraph()
    p.add_run(
        "La Instrument List cubre todas las ubicaciones requeridas por la ET para conductividad y caudal. Sin embargo, se detecta un "
    )
    run = p.add_run("error critico de codificacion")
    run.bold = True
    p.add_run(": el TAG ")
    run = p.add_run("FIT-09-001 esta DUPLICADO")
    run.bold = True
    run.font.color.rgb = RGBColor(255, 0, 0)
    p.add_run(" (Items 4 y 13), lo que requiere correccion inmediata.")

    # =========================================================================
    # 2. DOCUMENTOS ANALIZADOS
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading("2. DOCUMENTOS ANALIZADOS", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    table = doc.add_table(rows=1, cols=5)
    set_table_borders(table)
    add_header_row(table, ["Documento", "Codigo", "Rev.", "Fecha", "Fuente"])
    add_data_row(
        table,
        [
            "Especificacion Tecnica Modulo OI",
            "P22-ET-09-000-001-0",
            "0",
            "Sep-2025",
            "ADASA",
        ],
    )
    add_data_row(
        table, ["Oferta Tecnica BW Water", "-", "Rev.1", "Oct-2025", "BW Water"]
    )
    add_data_row(
        table,
        ["P&ID Sistema UHPRO", "P22-DWG-09-009-002-A", "A", "Dic-2025", "BW Water"],
    )
    add_data_row(
        table, ["Instrument List", "P22-LI-09-008-003-A", "A", "Ene-2026", "BW Water"]
    )
    doc.add_paragraph()

    # =========================================================================
    # 3. REQUISITOS CONTRACTUALES
    # =========================================================================
    h1 = doc.add_heading("3. REQUISITOS CONTRACTUALES (ET)", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    doc.add_heading("3.1 Conductimetros (Seccion 5.5.5)", level=2)
    p = doc.add_paragraph()
    p.add_run("La ET establece los siguientes puntos de medicion de conductividad:")

    table = doc.add_table(rows=1, cols=3)
    set_table_borders(table)
    add_header_row(table, ["#", "Ubicacion Requerida", "Proposito"])
    add_data_row(table, ["1", "Alimentacion al modulo", "Monitoreo calidad entrada"])
    add_data_row(table, ["2", "Salida de permeado", "Verificacion garantia TDS"])
    add_data_row(table, ["3", "Salida permeado 2da etapa", "Control de proceso"])
    add_data_row(table, ["4", "Salida rechazo 1ra etapa", "Monitoreo interstage"])
    add_data_row(table, ["5", "Salida de rechazo", "Balance de masas"])
    doc.add_paragraph()

    doc.add_heading("3.2 Caudalimetros (Seccion 5.5.1)", level=2)
    p = doc.add_paragraph()
    p.add_run("La ET establece los siguientes puntos de medicion de caudal:")

    table = doc.add_table(rows=1, cols=3)
    set_table_borders(table)
    add_header_row(table, ["#", "Ubicacion Requerida", "Proposito"])
    add_data_row(
        table, ["1", "Alimentacion al modulo", "Balance de masas, control HP pump"]
    )
    add_data_row(
        table, ["2", "Salida permeado rack 2da Etapa", "Verificacion recuperacion"]
    )
    add_data_row(table, ["3", "Salida de permeado", "Verificacion garantia capacidad"])
    add_data_row(table, ["4", "Salida de rechazo", "Balance de masas"])
    add_data_row(table, ["5", "Salida filtro cartucho CIP", "Control sistema CIP"])
    doc.add_paragraph()

    doc.add_heading("3.3 Protocolo de Comunicacion (Seccion 5.5)", level=2)
    p = doc.add_paragraph()
    run = p.add_run("Requisito: ")
    run.bold = True
    p.add_run("Todos los transmisores deben operar con protocolo ")
    run = p.add_run("4-20mA + HART")
    run.bold = True
    p.add_run(".")

    # =========================================================================
    # 4. COMPROMISOS OFERTA TECNICA
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading("4. COMPROMISOS OFERTA TECNICA (Rev.1)", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    doc.add_heading("4.1 Conductividad", level=2)
    p = doc.add_paragraph()
    p.add_run("La Oferta Tecnica Rev.1 lista ")
    run = p.add_run("1 instrumento")
    run.bold = True
    p.add_run(" de conductividad:")

    table = doc.add_table(rows=1, cols=4)
    set_table_borders(table)
    add_header_row(table, ["Item", "Descripcion", "Marca", "Ubicacion"])
    add_data_row(
        table,
        [
            "16",
            "Analytical Transmitter CONDUCTIVITY",
            "EMERSON/ABB",
            "RO TRAIN COMBINED PERMEATE",
        ],
    )
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run("Observacion: ")
    run.bold = True
    p.add_run(
        "La oferta solo menciona 1 conductimetro. Las 4 ubicaciones restantes no se mencionan explicitamente."
    )

    doc.add_heading("4.2 Caudal", level=2)
    p = doc.add_paragraph()
    p.add_run("La Oferta Tecnica Rev.1 lista ")
    run = p.add_run("4 instrumentos")
    run.bold = True
    p.add_run(" de caudal:")

    table = doc.add_table(rows=1, cols=4)
    set_table_borders(table)
    add_header_row(table, ["Item", "Descripcion", "Marca", "Ubicacion"])
    add_data_row(
        table, ["13", "Flow Transmitter", "ROSEMOUNT/ABB", "RO TRAIN COMBINED PERMEATE"]
    )
    add_data_row(
        table, ["14", "Flow Transmitter", "ROSEMOUNT/ABB", "RO TRAIN STAGE 2 PERMEATE"]
    )
    add_data_row(table, ["15", "Flow Transmitter", "ROSEMOUNT/ABB", "RO TRAIN REJECT"])
    add_data_row(table, ["33", "Flow Transmitter", "ROSEMOUNT/ABB", "CIP TO RO TRAIN"])
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run("Observacion: ")
    run.bold = True
    p.add_run("La oferta NO menciona caudalimetro para alimentacion al modulo.")

    # =========================================================================
    # 5. IMPLEMENTACION INSTRUMENT LIST
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading("5. IMPLEMENTACION INSTRUMENT LIST", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    p = doc.add_paragraph()
    p.add_run("Documento: ")
    run = p.add_run("P22-LI-09-008-003-A Rev.A")
    run.bold = True

    doc.add_heading("5.1 Conductimetros Incluidos (5 unidades)", level=2)

    table = doc.add_table(rows=1, cols=6)
    set_table_borders(table)
    add_header_row(
        table,
        ["TAG", "Descripcion", "Marca/Modelo", "Rango Op.", "Rango Max", "Protocolo"],
    )
    add_data_row(
        table,
        [
            "CIT-09-001",
            "RO Cartridge Filter Discharge",
            "Rosemount 400",
            "0-2000 uS/cm",
            "0-20 mS/cm",
            "4-20mA HART",
        ],
    )
    add_data_row(
        table,
        [
            "CIT-09-002",
            "RO Train Permeate",
            "Rosemount 400",
            "0-200 uS/cm",
            "0-200 uS/cm",
            "4-20mA HART",
        ],
    )
    add_data_row(
        table,
        [
            "CIT-09-003",
            "RO Stage 2 Permeate",
            "Rosemount 400",
            "0-200 uS/cm",
            "0-200 uS/cm",
            "4-20mA HART",
        ],
    )
    add_data_row(
        table,
        [
            "CIT-09-004",
            "RO Stage 1 Reject",
            "Rosemount 400",
            "0-12 mS/cm",
            "0-20 mS/cm",
            "4-20mA HART",
        ],
    )
    add_data_row(
        table,
        [
            "CIT-09-005",
            "RO Train Reject",
            "Rosemount 400",
            "0-12 mS/cm",
            "0-20 mS/cm",
            "4-20mA HART",
        ],
    )
    doc.add_paragraph()

    doc.add_heading("5.2 Caudalimetros Incluidos - TAG DUPLICADO", level=2)

    p = doc.add_paragraph()
    run = p.add_run("ALERTA CRITICA: ")
    run.bold = True
    run.font.color.rgb = RGBColor(255, 0, 0)
    p.add_run("El TAG FIT-09-001 aparece DOS VECES con ubicaciones diferentes.")

    table = doc.add_table(rows=1, cols=7)
    set_table_borders(table)
    add_header_row(
        table,
        ["TAG", "Descripcion", "Marca", "Rango Op.", "Rango Max", "DN", "Protocolo"],
    )
    add_data_row(
        table,
        [
            "FIT-09-001",
            "RO Cartridge Filter Discharge",
            "Rosemount 8750W",
            "0-49 m3/h",
            "0-100 m3/h",
            "DN100",
            "4-20mA HART",
        ],
    )
    add_data_row(
        table,
        [
            "FIT-09-001",
            "RO 2nd Stage Permeate",
            "Rosemount 8750W",
            "0-9 m3/h",
            "0-18 m3/h",
            "DN50",
            "4-20mA HART",
        ],
        highlight=True,
        highlight_color="FFCCCC",
    )
    add_data_row(
        table,
        [
            "FIT-09-003",
            "RO Train Permeate",
            "Rosemount 8750W",
            "0-21 m3/h",
            "0-40 m3/h",
            "DN80",
            "4-20mA HART",
        ],
    )
    add_data_row(
        table,
        [
            "FIT-09-004",
            "RO Train Reject",
            "Rosemount 8750W",
            "0-28 m3/h",
            "0-60 m3/h",
            "DN65",
            "4-20mA HART",
        ],
    )
    add_data_row(
        table,
        [
            "FIT-09-005",
            "CIP Pump Discharge",
            "Rosemount 8750W",
            "0-54 m3/h",
            "0-100 m3/h",
            "DN100",
            "4-20mA HART",
        ],
    )
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run("Nota: ")
    run.bold = True
    p.add_run("La segunda fila (resaltada) deberia tener TAG ")
    run = p.add_run("FIT-09-002")
    run.bold = True
    p.add_run(" para permeado de segunda etapa.")

    # =========================================================================
    # 6. MATRIZ COMPARATIVA
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading("6. MATRIZ COMPARATIVA", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    doc.add_heading("6.1 Conductividad", level=2)

    table = doc.add_table(rows=1, cols=6)
    set_table_borders(table)
    add_header_row(
        table, ["#", "Requisito ET", "Oferta", "TAG IL", "Descripcion IL", "CUMPLE"]
    )
    add_data_row(
        table,
        [
            "1",
            "Alimentacion al modulo",
            "NO",
            "CIT-09-001",
            "RO Cartridge Filter Discharge",
            "SI",
        ],
    )
    add_data_row(
        table,
        ["2", "Salida de permeado", "Item 16", "CIT-09-002", "RO Train Permeate", "SI"],
    )
    add_data_row(
        table,
        [
            "3",
            "Salida permeado 2da etapa",
            "NO",
            "CIT-09-003",
            "RO Stage 2 Permeate",
            "SI",
        ],
    )
    add_data_row(
        table,
        [
            "4",
            "Salida rechazo 1ra etapa",
            "NO",
            "CIT-09-004",
            "RO Stage 1 Reject",
            "SI",
        ],
    )
    add_data_row(
        table, ["5", "Salida de rechazo", "NO", "CIT-09-005", "RO Train Reject", "SI"]
    )
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run("Resultado Conductividad: ")
    run.bold = True
    p.add_run("5/5 ubicaciones cubiertas. ")
    run = p.add_run("CUMPLE REQUISITOS ET.")
    run.bold = True
    run.font.color.rgb = RGBColor(0, 128, 0)

    doc.add_heading("6.2 Caudal", level=2)

    table = doc.add_table(rows=1, cols=6)
    set_table_borders(table)
    add_header_row(
        table, ["#", "Requisito ET", "Oferta", "TAG IL", "Descripcion IL", "CUMPLE"]
    )
    add_data_row(
        table,
        [
            "1",
            "Alimentacion al modulo",
            "NO",
            "FIT-09-001",
            "RO Cartridge Filter Discharge",
            "SI",
        ],
    )
    add_data_row(
        table,
        [
            "2",
            "Salida permeado 2da etapa",
            "Item 14",
            "FIT-09-001",
            "RO 2nd Stage Permeate",
            "DUPLICADO",
        ],
        highlight=True,
        highlight_color="FFCCCC",
    )
    add_data_row(
        table,
        ["3", "Salida de permeado", "Item 13", "FIT-09-003", "RO Train Permeate", "SI"],
    )
    add_data_row(
        table,
        ["4", "Salida de rechazo", "Item 15", "FIT-09-004", "RO Train Reject", "SI"],
    )
    add_data_row(
        table,
        ["5", "Salida filtro CIP", "Item 33", "FIT-09-005", "CIP Pump Discharge", "SI"],
    )
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run("Resultado Caudal: ")
    run.bold = True
    p.add_run("5/5 ubicaciones cubiertas, ")
    run = p.add_run("PERO con TAG duplicado que requiere correccion.")
    run.bold = True
    run.font.color.rgb = RGBColor(255, 0, 0)

    # =========================================================================
    # 7. OBSERVACIONES
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading("7. OBSERVACIONES", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    # OBS-01
    doc.add_heading("OBS-01: TAG DUPLICADO FIT-09-001 (CRITICO)", level=2)

    table = doc.add_table(rows=1, cols=2)
    set_table_borders(table)
    add_header_row(table, ["Campo", "Valor"])
    add_data_row(table, ["Documento", "Instrument List P22-LI-09-008-003-A"])
    add_data_row(table, ["Ubicacion", "Items 4 y 13"])
    add_data_row(table, ["Categoria", "Tecnico"])
    add_data_row(
        table, ["Severidad", "CRITICO"], highlight=True, highlight_color="FFCCCC"
    )
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run("Descripcion: ")
    run.bold = True
    p.add_run(
        "El TAG FIT-09-001 aparece asignado a DOS ubicaciones completamente diferentes:"
    )

    table = doc.add_table(rows=1, cols=5)
    set_table_borders(table)
    add_header_row(table, ["Item", "TAG", "Ubicacion", "DN", "Rango"])
    add_data_row(
        table,
        ["4", "FIT-09-001", "RO Cartridge Filter Discharge", "DN100", "0-49 m3/h"],
    )
    add_data_row(
        table,
        ["13", "FIT-09-001", "RO 2nd Stage Permeate", "DN50", "0-9 m3/h"],
        highlight=True,
        highlight_color="FFCCCC",
    )
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run("Impacto: ")
    run.bold = True
    p.add_run(
        "Violacion de estandar de codificacion, imposibilidad de distinguir senales en PLC/SCADA, mapeo incorrecto en IO List."
    )

    p = doc.add_paragraph()
    run = p.add_run("Accion Requerida: ")
    run.bold = True
    p.add_run("Corregir Item 13 asignando TAG ")
    run = p.add_run("FIT-09-002")
    run.bold = True
    p.add_run(" al caudalimetro de permeado de segunda etapa.")

    # OBS-02
    doc.add_heading("OBS-02: Discrepancia Cantidad Conductimetros", level=2)

    table = doc.add_table(rows=1, cols=2)
    set_table_borders(table)
    add_header_row(table, ["Campo", "Valor"])
    add_data_row(table, ["Documento", "Oferta Tecnica Rev.1 vs Instrument List"])
    add_data_row(table, ["Categoria", "Contractual"])
    add_data_row(
        table, ["Severidad", "MAYOR"], highlight=True, highlight_color="FFFFCC"
    )
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run("Descripcion: ")
    run.bold = True
    p.add_run(
        "La Oferta Tecnica Rev.1 menciona solo 1 conductimetro, mientras que la Instrument List incluye 5 conductimetros."
    )

    p = doc.add_paragraph()
    run = p.add_run("Accion Requerida: ")
    run.bold = True
    p.add_run(
        "Solicitar confirmacion escrita de que los 5 conductimetros estan incluidos en el suministro contractual."
    )

    # OBS-06
    doc.add_heading("OBS-06: Rango Conductimetro Permeado - Verificar", level=2)

    table = doc.add_table(rows=1, cols=2)
    set_table_borders(table)
    add_header_row(table, ["Campo", "Valor"])
    add_data_row(table, ["Documento", "Instrument List - TAG CIT-09-002"])
    add_data_row(table, ["Categoria", "Tecnico"])
    add_data_row(
        table, ["Severidad", "REVISAR"], highlight=True, highlight_color="CCE5FF"
    )
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run("Descripcion: ")
    run.bold = True
    p.add_run(
        "CIT-09-002 tiene rango 0-200 uS/cm (~0-140 mg/l TDS). La garantia de TDS es <= 500 mg/l (~700-1000 uS/cm). Si el permeado opera cerca del limite de garantia, el instrumento estaria fuera de rango."
    )

    p = doc.add_paragraph()
    run = p.add_run("Accion Requerida: ")
    run.bold = True
    p.add_run(
        "Solicitar confirmacion de que el rango es adecuado para las condiciones operativas esperadas."
    )

    # =========================================================================
    # 8. VERIFICACION GARANTIAS
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading("8. VERIFICACION GARANTIAS DE DESEMPENO", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    p = doc.add_paragraph()
    p.add_run("Referencia: ET Seccion 10.1")

    table = doc.add_table(rows=1, cols=6)
    set_table_borders(table)
    add_header_row(
        table, ["Garantia", "Valor", "Instrumento", "TAG IL", "Rango", "Estado"]
    )
    add_data_row(
        table,
        [
            "Capacidad Nominal",
            "480 m3/dia (20 m3/h)",
            "FIT permeado",
            "FIT-09-003",
            "0-21 m3/h op",
            "ADECUADO",
        ],
    )
    add_data_row(
        table,
        [
            "TDS Permeado",
            "<= 500 mg/l",
            "CIT permeado",
            "CIT-09-002",
            "0-200 uS/cm",
            "VERIFICAR",
        ],
        highlight=True,
        highlight_color="FFFFCC",
    )
    add_data_row(table, ["Cloruros Permeado", "<= 400 mg/l", "Lab", "N/A", "-", "N/A"])
    add_data_row(
        table, ["SEC", "4.71 kWh/m3 +/-5%", "FIT + EIT", "FIT-09-003", "-", "VERIFICAR"]
    )
    doc.add_paragraph()

    # =========================================================================
    # 9. CONCLUSION
    # =========================================================================
    h1 = doc.add_heading("9. CONCLUSION", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    doc.add_heading("9.1 Cumplimiento General", level=2)

    table = doc.add_table(rows=1, cols=2)
    set_table_borders(table)
    add_header_row(table, ["Aspecto", "Evaluacion"])
    add_data_row(table, ["Conductimetros - Cantidad", "CUMPLE (5/5 ubicaciones)"])
    add_data_row(table, ["Conductimetros - Protocolo", "CUMPLE (4-20mA HART)"])
    add_data_row(table, ["Caudalimetros - Cantidad", "CUMPLE (5/5 ubicaciones)"])
    add_data_row(table, ["Caudalimetros - Protocolo", "CUMPLE (4-20mA HART)"])
    add_data_row(
        table,
        ["Caudalimetros - TAGs", "NO CUMPLE (TAG duplicado)"],
        highlight=True,
        highlight_color="FFCCCC",
    )
    add_data_row(table, ["Verificacion Garantia Capacidad", "CUMPLE"])
    add_data_row(table, ["Verificacion Garantia TDS", "VERIFICAR rango"])
    doc.add_paragraph()

    doc.add_heading("9.2 Veredicto Final", level=2)

    p = doc.add_paragraph()
    run = p.add_run("3 - TO BE REVISED")
    run.bold = True
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.font.size = Pt(14)

    p = doc.add_paragraph()
    p.add_run(
        "La Instrument List P22-LI-09-008-003-A cubre todas las ubicaciones requeridas por la ET. Sin embargo, requiere revision debido a:"
    )

    p = doc.add_paragraph()
    p.add_run("  - ")
    run = p.add_run("Error critico: ")
    run.bold = True
    p.add_run("TAG FIT-09-001 duplicado (debe corregirse a FIT-09-002)")

    p = doc.add_paragraph()
    p.add_run("  - ")
    run = p.add_run("Verificacion pendiente: ")
    run.bold = True
    p.add_run("Rango de CIT-09-002 para condiciones cercanas al limite de garantia")

    # =========================================================================
    # 10. ACCIONES REQUERIDAS
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading("10. ACCIONES REQUERIDAS A BW WATER", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    table = doc.add_table(rows=1, cols=5)
    set_table_borders(table)
    add_header_row(table, ["#", "Accion", "Prioridad", "Documento", "Plazo"])
    add_data_row(
        table,
        [
            "1",
            "Corregir TAG duplicado FIT-09-001 -> FIT-09-002",
            "CRITICA",
            "Instrument List Rev.B",
            "Inmediato",
        ],
        highlight=True,
        highlight_color="FFCCCC",
    )
    add_data_row(
        table,
        [
            "2",
            "Actualizar IO List con TAG corregido",
            "ALTA",
            "IO List Rev.B",
            "Con Rev. IL",
        ],
    )
    add_data_row(
        table,
        ["3", "Verificar P&ID refleje TAGs correctos", "ALTA", "P&ID", "Con Rev. IL"],
    )
    add_data_row(
        table,
        [
            "4",
            "Confirmar 5 conductimetros en suministro",
            "MEDIA",
            "Respuesta tecnica",
            "5 dias",
        ],
    )
    add_data_row(
        table,
        [
            "5",
            "Confirmar rango CIT-09-002 (0-200 uS/cm)",
            "MEDIA",
            "Respuesta tecnica",
            "5 dias",
        ],
    )
    add_data_row(
        table,
        [
            "6",
            "Confirmar FIT-09-001 alimentacion en suministro",
            "BAJA",
            "Respuesta tecnica",
            "5 dias",
        ],
    )

    # =========================================================================
    # GUARDAR DOCUMENTO
    # =========================================================================
    doc.save(output_file)
    print(f"Documento generado: {output_file}")

    return output_file


if __name__ == "__main__":
    main()
