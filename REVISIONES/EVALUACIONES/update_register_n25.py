#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
update_register_n25.py
Actualiza el Master Deliverable Register a TM N25 (submittals 25007-0055 E55 +
25007-0056 E56). Update EXPLICITO modelado en update_register_n24.py.

Particularidad de TM N25: los 6 documentos son re-revisiones de documentos ya
entregados (ya contados como Delivered), por lo que NO se agregan filas nuevas; las
6 filas existentes se ACTUALIZAN (Rev / Delivery / TM / Verdict / Action). Total y
Delivered no cambian; solo la distribucion de veredictos (recomputada desde la hoja:
1-Approved 43->46, 2-AN 23->21, 3 13->12, 4 0).

Nota (re-verdict 29-Jun): el IO List Rev 4 quedo Code 2 - Approved as Noted (no 3).
Las 4 senales de coordinacion con el PLC externo estan cumplidas; el esquema
soft-I/O Ethernet/IP de campo fue aceptado en N20 y NO se reabre (la observacion
'soft-BOOL = no feedback' era over-reach). Driver unico del veredicto 3 del TM =
el Outline Panel Drawing.

Mapeo (matching por codigo de registro):
  #31 IO List (Rev 3->4, N24 2-AN -> N25 2-AN), #81 LCP Datasheet (Rev 0->1,
  N24 2-AN -> N25 1-Approved), #17 DS Static Mixer (Rev C->0, N22 2-AN -> N25
  1-Approved), #98 Outline Panel Drawing (Rev A->B, N20 3 -> N25 3), #107 UHPRO
  Structural (Rev A->B, N23 3 -> N25 1-Approved), #19 Painting Specifications
  (Rev B->C, N11 1 -> N25 1-Approved; OJO: re-codificado por BW Water de
  P22-ET-09-005-002 a P22-ET-09-006-002 - matching por el codigo de registro
  005-002, recode anotado en Action pendiente de confirmacion).
"""
import os
import shutil
import collections
from copy import copy
import openpyxl

EV = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register.xlsx")
BAK = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register_pre-N25.xlsx")

TM = "N25"
TM_DATE = "29-Jun-2026"

# ---- Master Register: updates a filas existentes (por codigo de registro) ----
MR_UPDATES = {
    # IO List Rev 4 -> Code 2 - Approved as Noted
    "P22-LI-09-008-001": dict(
        rev="4", delivery="E56", tm=TM, verdict="2-AN",
        action=(
            "Rev 4 delivered (E56). Code 2 - Approved as Noted (TM N25). The four "
            "module-to-external-PLC interface signals are present and correctly "
            "typed as relay contacts (XA005 enable, YA001 running, YA002 fault, "
            "YA003 local/remote) - the two added at TM N24 are an ADASA extension; "
            "OBS-01 of TM N24 closed, the interface is complete. The soft-I/O field "
            "scheme over Ethernet/IP accepted at TM N20 is not reopened. Two minor "
            "items fold at IFC Rev 0 (no new review cycle): OBS-01 - complete the "
            "signal-type count cell on items 131/136 (typed BOOL) and place the "
            "count in the BOOL column on items 129/134 (the TM N24 count-cell "
            "observation persists); NOTE-02 - confirm the dosing-pump run status "
            "(BOOL over Ethernet/IP, shown FROM the HMI) reflects a verified "
            "pump/VFD feedback, not an operator-faceplate echo. IFC also gated on "
            "the Plant Control Philosophy children (Section 3)."),
    ),
    # LCP Datasheet Rev 1 -> Code 1 - Approved
    "P22-ET-09-007-005": dict(
        rev="1", delivery="E56", tm=TM, verdict="1-Approved",
        action=(
            "Rev 1 delivered (E56). Code 1 - Approved (TM N25). Closes the TM N24 "
            "cycle: sheet 114 now reconciles the panel internal consumption with the "
            "UPS (13 VE-09 valve feeders 1.56 kW + SAI-09-001 0.24 kW = 1.80 kW / "
            "2.00 kW with the factor) and the conflicting DWG code is removed and "
            "unified to P22-ET-09-007-005. Enclosure marine spec (SS316L, NEMA 4X / "
            "IP66) reconfirmed; the Outline Panel Drawing must align to it (still "
            "Code 3 in TM N25). NOTE-01: populate the blank sheet-114 document number "
            "at the IFC Rev 0 issue (housekeeping). Issue directly at IFC Rev 0."),
    ),
    # DS Static Mixer Rev 0 -> Code 1 - Approved
    "P22-ET-09-009-012": dict(
        rev="0", delivery="E56", tm=TM, verdict="1-Approved",
        action=(
            "Rev 0 delivered (E56). Code 1 - Approved (TM N25). Closes the TM N22 "
            "cycle: the datasheet reconciles the design-condition and injection-rate "
            "discrepancy noted at TM N22, adding a conservative-values remark. "
            "Contracted item (Technical Offer Rev1, 1x100% FRP static mixer). Komax "
            "to N-Spindle NS11 substitution permitted under the offer 'or equal' "
            "(CoV 0.05 mixing equivalence); FRP housing, ANSI 150 connections, "
            "design pressure/temperature consistent with the antiscalant service. "
            "NOTE: injection port DN15 vs offer 1 inch - ample for the dosing rate. "
            "Cross-checks (injection rate vs the dosing-pump duty; tag and line "
            "class on the P&ID) tracked in Section 3. Issue directly at IFC Rev 0."),
    ),
    # Outline Panel Drawing Rev B -> Code 3 - To be revised
    "P22-CD-09-008-001": dict(
        rev="B", delivery="E55", tm=TM, verdict="3-To be revised",
        action=(
            "Rev B delivered (E55). Code 3 - To be revised (TM N25). The TM N20 "
            "enclosure contradiction is only half-corrected, so the panel "
            "fabrication hold continues. Rev B adds an 'Exterior SUS316L' colour "
            "line, a sealed NEMA 4X stainless-steel A/C unit (closes the cooling "
            "observation) and the panel weight (814 kg) - but OBS-01 CRITICAL: the "
            "MATERIAL row still specifies sheet steel for the frame/roof/rear "
            "panel/door under an 'interior only' heading that lists weather-exposed "
            "surfaces, and the FINISHING rows paint them RAL 7035, internally "
            "contradicting the sheet's own SUS316L exterior and specifying an "
            "enclosure below the approved documents that govern it - the LCP "
            "Datasheet Rev B (FS66S, unpainted SS316L, NEMA 4X/IP66) and the IFC "
            "Single Line Diagram Rev 0 ('METAL CLAD, NEMA 4X/IP66'). Basis: ET Power "
            "and Control Switchboards Constructive Characteristics requires NEMA 4X "
            "(not lower); the SS316L material is fixed by the approved LCP Datasheet, "
            "NOT by the ET (ET 5.4.3 permits painted steel) - lead with the "
            "approved-document contradiction, not 'the ET requires SS316L'. OBS-02 "
            "MAJOR: bottom gland plate zinc-plated. NOTE-01: NEMA 4X without IP66. "
            "NOTE-02: title-block typo. Re-issue as Rev C aligned to the approved "
            "LCP Datasheet (SS316L on every weather-exposed surface)."),
    ),
    # UHPRO Structural Design Criteria Rev B -> Code 1 - Approved
    "P22-CD-09-005-003": dict(
        rev="B", delivery="E55", tm=TM, verdict="1-Approved",
        action=(
            "Rev B delivered (E55). Code 1 - Approved (TM N25). Closes the TM N23 "
            "cycle: every seismic parameter and the ASD/strength load combinations "
            "are now cited to NCh 2369 Of.2003 (clause 4.5 governing the seismic "
            "ASD combinations), and the lifting load case is included with its "
            "padeye and yoke design criteria (API RP 2A-WSD factors 1.35 / 2.0, "
            "padeye FS 2.0 + 5% lateral, spreader-beam yoke). NOTE: fold the "
            "editorial items at issue (cover date, duplicate table number) and "
            "confirm N/m2 wind units on the issued PDF. The final lifting design "
            "(lifting calc, lifting drawing with weights, yoke design) and the "
            "seismic calculation memo are tracked in Section 3. Issue at IFC Rev 0."),
    ),
    # Painting Specifications Rev C -> Code 1 - Approved (re-coded by BW Water)
    "P22-ET-09-005-002": dict(
        rev="C", delivery="E55", tm=TM, verdict="1-Approved",
        action=(
            "Rev C delivered (E55) as P22-ET-09-006-002 - RE-CODED by BW Water from "
            "the TM N11 approved P22-ET-09-005-002 Rev B (register code retained "
            "pending confirmation, NOTE-03). Code 1 - Approved (TM N25). The "
            "painting SPECIFICATION (distinct from the Painting Procedure "
            "P22-BA-09-000-011, still Code 3 in Section 3) matches the ET marine "
            "system (80/200/75 = 355 um, Sa 2.5); the MAKE column names Sherwin "
            "Williams plus generic equivalents. NOTE-01: the brand list adds 'Jotun "
            "or similar technically equivalent approved by ADASA' - this is not "
            "approval of a Jotun system; the equivalence + C5-M demonstration remain "
            "open at the Painting Procedure (Section 3). NOTE-02: verify the "
            "surface-prep grade for the zinc primer on the safety supports / "
            "platform. NOTE-03: confirm the document-code change. Issue at IFC Rev 0."),
    ),
}

# ---- Revision History: 6 filas nuevas (num, doc, code, rev, delivery, submittal, verdict, keyobs) ----
RH_NEW = [
    (31, "IO List", "P22-LI-09-008-001", "4", "E56", "25007-0056",
     "2-AN",
     "Rev 4. Code 2 - Approved as Noted. The four module-to-external-PLC signals "
     "present and correctly typed as relay contacts (XA005/YA001/YA002/YA003); TM "
     "N24 OBS-01 closed, the interface is complete. The soft-I/O Ethernet/IP field "
     "scheme accepted at TM N20 is not reopened (the prior 'soft-BOOL = no "
     "feedback' framing was over-reach: ADASA itself asked to relabel the dosing "
     "RUNNING to soft I/O at N20 NOTE-02). OBS-01 MINOR: items 131/136 (BOOL) lack "
     "the count cell, 129/134 place it in the DI column - complete at IFC. NOTE-02: "
     "confirm the dosing run-status source. Fold at IFC Rev 0; issue also gated on "
     "the Control Philosophy children (Section 3)."),
    (81, "Datasheet Local Control Panel (LCP)", "P22-ET-09-007-005", "1", "E56",
     "25007-0056", "1-Approved",
     "Rev 1. Code 1 - Approved. Closes the TM N24 cycle: sheet 114 reconciles the "
     "panel internal consumption with the UPS (1.80 / 2.00 kW) and the conflicting "
     "DWG code is unified to P22-ET-09-007-005. Enclosure marine spec (SS316L, "
     "NEMA 4X/IP66) reconfirmed; the Outline must align to it. NOTE-01: populate "
     "the blank sheet-114 document number at issue. Issue at IFC Rev 0."),
    (17, "DS Static Mixer", "P22-ET-09-009-012", "0", "E56", "25007-0056",
     "1-Approved",
     "Rev 0. Code 1 - Approved. Closes the TM N22 cycle: reconciles the "
     "design-condition and injection-rate discrepancy with a conservative-values "
     "remark. Contracted item (Offer Rev1). Komax to N-Spindle NS11 under 'or "
     "equal' (CoV 0.05). FRP, ANSI 150, design pressure/temperature consistent. "
     "Injection DN15 vs offer 1 inch - ample. Cross-checks tracked (Section 3). "
     "Issue at IFC Rev 0."),
    (98, "PLC-LCP Outline Panel Drawing", "P22-CD-09-008-001", "B", "E55",
     "25007-0055", "3-To be revised",
     "Rev B. Code 3. The TM N20 enclosure contradiction is half-corrected: adds "
     "'Exterior SUS316L', a sealed NEMA 4X A/C unit and the 814 kg weight, but "
     "OBS-01 CRITICAL - MATERIAL/FINISHING still build a painted sheet-steel "
     "enclosure (frame/roof/rear panel/door) under an 'interior only' heading that "
     "lists weather-exposed surfaces, contradicting the SUS316L exterior and the "
     "LCP Datasheet. OBS-02: bottom gland plate zinc-plated. NOTE-01: NEMA 4X "
     "without IP66. NOTE-02: title-block typo. Re-issue as Rev C; fabrication "
     "remains gated."),
    (107, "UHPRO Structural Design Criteria", "P22-CD-09-005-003", "B", "E55",
     "25007-0055", "1-Approved",
     "Rev B. Code 1 - Approved. Closes the TM N23 cycle: seismic citations and "
     "ASD/strength combinations now NCh 2369 Of.2003 (clause 4.5 governing), and "
     "the lifting load case with padeye/yoke criteria (API RP 2A-WSD 1.35/2.0, "
     "padeye FS 2.0 + 5% lateral, spreader-beam yoke) is included. NOTE: editorial "
     "items (cover date, table number, N/m2 wind units) fold at issue. Final "
     "lifting design + seismic calc memo tracked (Section 3). Issue at IFC Rev 0."),
    (19, "Painting Specifications", "P22-ET-09-005-002", "C", "E55", "25007-0055",
     "1-Approved",
     "Rev C. Code 1 - Approved. Delivered as P22-ET-09-006-002 (RE-CODED by BW "
     "Water from the TM N11 approved P22-ET-09-005-002 Rev B; NOTE-03 confirm). "
     "Matches the ET marine system (355 um, Sa 2.5). NOTE-01: brand list adds "
     "'Jotun or similar technically equivalent approved by ADASA' - not approval of "
     "Jotun; equivalence + C5-M remain open at the Painting Procedure (Section 3). "
     "NOTE-02: surface-prep grade for the zinc primer. Issue at IFC Rev 0."),
]


def style_row(ws, dst_row, ref_row, ncols):
    for c in range(1, ncols + 1):
        s = ws.cell(ref_row, c)
        d = ws.cell(dst_row, c)
        d._style = copy(s._style)


def main():
    shutil.copyfile(SRC, BAK)
    print(f"Backup: {BAK}")

    wb = openpyxl.load_workbook(SRC)  # data_only=False: preserva formulas/estilos
    mr = wb["Master Register"]
    rh = wb["Revision History"]
    sm = wb["Summary"]

    ncols_mr = mr.max_column
    hdr = {mr.cell(1, c).value: c for c in range(1, ncols_mr + 1)}
    C_CODE = hdr["Code / ET Reference"]
    C_REV = hdr["Rev"]; C_DEL = hdr["Delivery"]; C_TM = hdr["TM"]
    C_VER = hdr["Verdict"]; C_STA = hdr["Status"]; C_ACT = hdr["Action Required"]

    # 1) updates a filas existentes (sin filas nuevas en TM N25)
    seen = set()
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code in MR_UPDATES:
            u = MR_UPDATES[code]
            mr.cell(r, C_REV).value = u["rev"]
            mr.cell(r, C_DEL).value = u["delivery"]
            mr.cell(r, C_TM).value = u["tm"]
            mr.cell(r, C_VER).value = u["verdict"]
            mr.cell(r, C_ACT).value = u["action"]
            seen.add(code)
            print(f"  MR update row {r}: {code} -> {u['verdict']} / {u['tm']} (Rev {u['rev']})")
    missing = set(MR_UPDATES) - seen
    if missing:
        raise SystemExit(f"ERROR: codigos N25 no encontrados en Master Register: {missing}")

    # 2) Revision History
    rcols = rh.max_column
    rhh = {rh.cell(1, c).value: c for c in range(1, rcols + 1)}
    prior = collections.Counter()
    rh_last = 1
    for r in range(2, rh.max_row + 1):
        cd = rh.cell(r, rhh["Code / ET Reference"]).value
        if cd:
            prior[str(cd).strip()] += 1
            rh_last = r
    for i, row in enumerate(RH_NEW, start=1):
        dst = rh_last + i
        num, doc, code, rev, dele, sub, ver, keyobs = row
        cycle = prior[code] + 1
        style_row(rh, dst, rh_last, rcols)
        rh.cell(dst, rhh["#"]).value = num
        rh.cell(dst, rhh["Document"]).value = doc
        rh.cell(dst, rhh["Code / ET Reference"]).value = code
        rh.cell(dst, rhh["Rev"]).value = rev
        rh.cell(dst, rhh["Delivery"]).value = dele
        rh.cell(dst, rhh["Submittal"]).value = sub
        rh.cell(dst, rhh["TM"]).value = TM
        rh.cell(dst, rhh["TM Date"]).value = TM_DATE
        rh.cell(dst, rhh["Verdict"]).value = ver
        rh.cell(dst, rhh["Cycle"]).value = cycle
        rh.cell(dst, rhh["Key Observations"]).value = keyobs
        print(f"  RH new row {dst}: #{num} {code} cycle {cycle} -> {ver}")

    # 3) Summary: recomputar desde Master Register
    rows = []
    for r in range(2, mr.max_row + 1):
        code = mr.cell(r, C_CODE).value
        if code and str(code).strip():
            rows.append((str(mr.cell(r, C_VER).value or "").strip(),
                         str(mr.cell(r, C_STA).value or "")))
    total = len(rows)
    delivered = sum(1 for v, s in rows if "Delivered" in s)
    vd = collections.Counter(v for v, s in rows if "Delivered" in s)
    print(f"  Summary recompute: total={total} delivered={delivered} verdicts={dict(vd)}")

    sm["B4"].value = total
    sm["B5"].value = delivered
    sm["B9"].value = "25 (N1 through N25 - TM N5 issued in Rev 0 and Rev 1)"
    sm["B10"].value = "56 (E1 through E56)"
    sm["B12"].value = "29-Jun-2026 (E56)"
    sm["B13"].value = "29-Jun-2026"

    def pct(n):
        return f"{round(100 * n / delivered)}%" if delivered else "0%"
    sm["B27"].value = vd.get("1-Approved", 0); sm["C27"].value = pct(vd.get("1-Approved", 0))
    sm["B28"].value = vd.get("2-AN", 0); sm["C28"].value = pct(vd.get("2-AN", 0))
    sm["B29"].value = vd.get("3-To be revised", 0); sm["C29"].value = pct(vd.get("3-To be revised", 0))
    sm["B30"].value = vd.get("4-Rejected", 0); sm["C30"].value = pct(vd.get("4-Rejected", 0))
    # ITEMS BY SECTION: sin cambio en TM N25 (re-revisiones de docs ya entregados)

    wb.save(SRC)
    print(f"Guardado: {SRC}")
    print("NOTA: ITEMS BY SECTION y REVIEW CYCLES distribution NO se modificaron "
          "(re-revisiones; recuento manual).")


if __name__ == "__main__":
    main()
