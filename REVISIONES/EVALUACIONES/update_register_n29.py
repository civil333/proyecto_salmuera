#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
update_register_n29.py
Actualiza el Master Deliverable Register a TM N29 (E67, submittal 25007-0067,
4 documentos, TODAS re-revisiones). Modelado en update_register_n28.py.

Idempotencia: el backup _pre-N29.xlsx conserva el estado LIMPIO N28 (48/28/8/0).
Al re-correr, RESTAURA SRC desde ese backup antes de aplicar.

Veredicto N29: 2 - Approved as Noted. Tally 2 Code 2 + 2 Code 1.
  UPDATES (4 re-revisiones de filas existentes; NO hay items nuevos):
    #105 HP and LP Pressure Test Procedure P22-BA-09-000-010 (Rev C->D, 3 -> 2-AN)
    #108 UHPRO Structural Calculation Report P22-CD-09-005-001 (Rev A->B, 3 -> 2-AN)
    #98  PLC/LCP Outline Panel Drawing       P22-CD-09-008-001 (Rev C->0, 2 -> 1-Approved)
    #38  Line List                           P22-LI-09-009-003 (Rev C->0, 1 -> 1-Approved)

Tally esperado tras N29: 49 Code 1 / 29 Code 2 / 6 Code 3 / 0 Code 4;
108 items / 84 delivered; 29 TMs / 67 entregas.
"""
import os
import shutil
import collections
from copy import copy
import openpyxl

EV = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register.xlsx")
BAK = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register_pre-N29.xlsx")

TM = "N29"
TM_DATE = "21-Jul-2026"

MR_UPDATES = {
    # HP and LP Pressure Test Procedure Rev D -> Code 2 (cierra el Code 3 critico de N27)
    "P22-BA-09-000-010": dict(
        rev="D", delivery="E67", tm=TM, verdict="2-AN",
        action=(
            "Rev D delivered (E67). Code 2 - Approved as noted (TM N29). Closes the Code 3 "
            "critical finding of TM N27: hydrostatic test pressures now set per line and "
            "material via the Line List - line DA-PVC-DN65-09-016 (PVC) corrected to 3 bar "
            "hydrotest at 2 bar design, no PVC line above 7.5 bar, 135 bar confined to the "
            "Super Duplex lines; body states 135/7.5 bar and cites the Line List; clause "
            "gap 5.7.2.2 fixed; the four N27 comment-sheet items verified in the body. "
            "NOTE-01: fix the applicable edition/addenda of ASME Section V and B31.3. "
            "Approved for the pressure tests; the HP hydrostatic test remains a Hold Point."),
    ),
    # UHPRO Structural Calculation Report Rev B -> Code 2 (cierra el Code 3 de N26)
    "P22-CD-09-005-001": dict(
        rev="B", delivery="E67", tm=TM, verdict="2-AN",
        action=(
            "Rev B delivered (E67). Code 2 - Approved as noted (TM N29). Resolves the Code 3 "
            "of TM N26: adds the design seismic weight, global base shear and NCh 2369:2003 "
            "base-shear check (operating base shear 41.62 kN, Zone 3, A0=0.40g), the "
            "governing skid utilization ratio (0.454 < 1.0), material grades, Site Class E "
            "basis and container base reactions; EOX combination typo corrected. Seismic "
            "design by IBC 2018/ASCE 7-16 calibrated to the Chilean base shear with the "
            "NCh 2369:2003 check; correctly does NOT adopt the 2025 edition. NOTE-01/02: "
            "the PDF is a duplicated concatenation (about 58 MB) and the comment sheet "
            "omits ADASA's original comment text. The calculation itself is accepted."),
    ),
    # PLC/LCP Outline Panel Drawing Rev 0 IFC -> Code 1 (cierra el gate del enclosure)
    "P22-CD-09-008-001": dict(
        rev="0", delivery="E67", tm=TM, verdict="1-Approved",
        action=(
            "Rev 0 (IFC) delivered (E67). Code 1 - Approved (TM N29). Closes the enclosure "
            "gate: incorporates the RFI 25007-RO-RFI-0002 disposition in full (external "
            "body/door/roof/rear/plinth + gland plates SS316L, internals galvanized/CRS, "
            "NEMA 4X/IP66) and closes the three N27 minors (cable-clamp sizes stated, "
            "gateway labelled Ethernet/IP-Modbus TCP, cable-entry label corrected). SLD "
            "re-issued Rev 1 (SS316L Panel, NEMA 4X/IP66), closing the cross-document "
            "commitment. Two cosmetic labels (IFA status stamp on an IFC drawing; "
            "zinc-plated vs hot-dip galvanized mounting-plate wording) to tidy at next "
            "issue; no change to content. Issue at IFC Rev 0."),
    ),
    # Line List Rev 0 IFC -> Code 1 (consistente con P&ID Rev D)
    "P22-LI-09-009-003": dict(
        rev="0", delivery="E67", tm=TM, verdict="1-Approved",
        action=(
            "Rev 0 (IFC) delivered (E67). Code 1 - Approved (TM N29). Adds a hydrotest "
            "column (1.5x design), consistent internally and with the P&ID Rev D and Valve "
            "List Rev D. The one substantive change from the approved Rev C - design "
            "pressure of DA-PVC-DN65-09-016 reduced 50 -> 2 bar - is correct and verified "
            "against the P&ID Rev D (PVC brine discharge to drain, downstream of the feed "
            "turbocharger; the 50 bar belongs to the Super Duplex turbocharger lines). Not "
            "declared in the comment sheet: record the 09-016 design-pressure correction in "
            "the revision history for traceability. Issue at IFC Rev 0."),
    ),
}

# N29 no incorpora documentos nuevos: los 4 son re-revisiones.
NEW_ITEMS = []

RH_NEW = [
    (None, "HP and LP Pressure Test Procedure", "P22-BA-09-000-010", "D", "E67",
     "25007-0067", "2-AN",
     "Rev D. Code 2 - Approved as noted. Closes the N27 Code 3 (75 bar over PVC): test "
     "pressures per line/material via the Line List, 09-016 corrected to 3 bar hydrotest / "
     "2 bar design, 135 bar only on Super Duplex; clause gap fixed; 4 N27 CCS items "
     "verified. NOTE-01: fix the ASME Section V / B31.3 edition. HP hydrostatic test "
     "remains a Hold Point."),
    (None, "UHPRO Structural Calculation Report", "P22-CD-09-005-001", "B", "E67",
     "25007-0067", "2-AN",
     "Rev B. Code 2 - Approved as noted. Resolves the N26 Code 3: seismic weight, global "
     "base shear, NCh 2369:2003 base-shear check (41.62 kN, Zone 3), utilization 0.454 < "
     "1.0, Site Class E, material grades, base reactions; EOX typo fixed; not the 2025 "
     "edition. NOTE-01/02: duplicated 58 MB PDF; comment sheet without ADASA comment text. "
     "Calculation accepted."),
    (None, "PLC/LCP Outline Panel Drawing", "P22-CD-09-008-001", "0", "E67",
     "25007-0067", "1-Approved",
     "Rev 0, IFC. Code 1 - Approved. Closes the enclosure gate: RFI-002 in full (exterior "
     "+ gland plates SS316L, internals galvanized, NEMA 4X/IP66); 3 N27 minors closed; SLD "
     "re-issued Rev 1. Two cosmetic labels (IFA stamp, plate finish) to tidy at next issue. "
     "Issue at IFC Rev 0."),
    (None, "Line List", "P22-LI-09-009-003", "0", "E67",
     "25007-0067", "1-Approved",
     "Rev 0, IFC. Code 1 - Approved. Hydrotest column added; consistent with P&ID Rev D "
     "and Valve List Rev D. 09-016 design pressure 50 -> 2 bar verified correct vs P&ID Rev "
     "D (PVC brine discharge to drain); record the change in the revision history. Issue at "
     "IFC Rev 0."),
]


def style_row(ws, dst_row, ref_row, ncols):
    for c in range(1, ncols + 1):
        s = ws.cell(ref_row, c)
        d = ws.cell(dst_row, c)
        d._style = copy(s._style)


def main():
    if os.path.exists(BAK):
        shutil.copyfile(BAK, SRC)
        print(f"Restaurado SRC desde baseline limpio: {BAK}")
    else:
        shutil.copyfile(SRC, BAK)
        print(f"Baseline creado: {BAK}")

    wb = openpyxl.load_workbook(SRC)
    mr = wb["Master Register"]
    rh = wb["Revision History"]
    sm = wb["Summary"]

    ncols_mr = mr.max_column
    hdr = {mr.cell(1, c).value: c for c in range(1, ncols_mr + 1)}
    C_NUM = hdr["#"]; C_DOC = hdr["Document"]; C_CODE = hdr["Code / ET Reference"]
    C_REV = hdr["Rev"]; C_DEL = hdr["Delivery"]; C_TM = hdr["TM"]
    C_VER = hdr["Verdict"]; C_STA = hdr["Status"]; C_ACT = hdr["Action Required"]

    # 1) updates (4 re-revisiones)
    seen = set()
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code in MR_UPDATES:
            u = MR_UPDATES[code]
            mr.cell(r, C_REV).value = u["rev"]
            mr.cell(r, C_DEL).value = u["delivery"]
            mr.cell(r, C_TM).value = u["tm"]
            mr.cell(r, C_VER).value = u["verdict"]
            mr.cell(r, C_STA).value = "Delivered"
            mr.cell(r, C_ACT).value = u["action"]
            seen.add(code)
            print(f"  MR update row {r}: {code} -> {u['verdict']} / {u['tm']} (Rev {u['rev']})")
    missing = set(MR_UPDATES) - seen
    if missing:
        raise SystemExit(f"ERROR: codigos N29 no encontrados en Master Register: {missing}")

    # 2) filas NUEVAS -> ninguna en N29
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()

    reg_num = {}
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code:
            try:
                reg_num[code] = int(mr.cell(r, C_NUM).value)
            except (TypeError, ValueError):
                pass

    # 3) Revision History (4 filas nuevas)
    rcols = rh.max_column
    rhh = {rh.cell(1, c).value: c for c in range(1, rcols + 1)}
    prior = collections.Counter(); rh_last = 1
    for r in range(2, rh.max_row + 1):
        cd = rh.cell(r, rhh["Code / ET Reference"]).value
        if cd:
            prior[str(cd).strip()] += 1; rh_last = r
    for i, row in enumerate(RH_NEW, start=1):
        dst = rh_last + i
        _num, doc, code, rev, dele, sub, ver, keyobs = row
        num = reg_num.get(code, "")
        cycle = prior[code] + 1; prior[code] += 1
        style_row(rh, dst, rh_last, rcols)
        rh.cell(dst, rhh["#"]).value = num
        rh.cell(dst, rhh["Document"]).value = doc
        rh.cell(dst, rhh["Code / ET Reference"]).value = code
        rh.cell(dst, rhh["Rev"]).value = rev
        rh.cell(dst, rhh["Delivery"]).value = dele
        rh.cell(dst, rhh["Submittal"]).value = sub
        rh.cell(dst, rhh["TM"]).value = TM
        rh.cell(dst, rhh["TM Date"]).value = TM_DATE
        rh.cell(dst, rhh["Verdict"]).value = ver
        rh.cell(dst, rhh["Cycle"]).value = cycle
        rh.cell(dst, rhh["Key Observations"]).value = keyobs
        print(f"  RH new row {dst}: #{num} {code} cycle {cycle} -> {ver}")

    # 4) Summary: recomputar total/delivered + verdicts desde MR
    stat = collections.Counter(); vd = collections.Counter(); total = 0
    for r in range(2, mr.max_row + 1):
        code = mr.cell(r, C_CODE).value
        if code and str(code).strip():
            total += 1
            su = str(mr.cell(r, C_STA).value or "").strip().upper()
            stat[su] += 1
            if su == "DELIVERED":
                vd[str(mr.cell(r, C_VER).value or "").strip()] += 1
    delivered = stat.get("DELIVERED", 0)
    not_delivered = stat.get("NOT DELIVERED", 0)
    partial = stat.get("PARTIAL", 0)
    print(f"  Summary recompute: total={total} delivered={delivered} "
          f"not_delivered={not_delivered} partial={partial} verdicts={dict(vd)}")

    sm["B4"].value = total
    sm["B5"].value = delivered
    sm["B7"].value = partial
    sm["B8"].value = not_delivered
    sm["B9"].value = "29 (N1 through N29 - TM N5 issued in Rev 0 and Rev 1)"
    sm["B10"].value = "67 (E1 through E67)"
    sm["B12"].value = "21-Jul-2026 (E67)"
    sm["B13"].value = "21-Jul-2026"

    def pct(n):
        return f"{round(100 * n / delivered)}%" if delivered else "0%"
    sm["B27"].value = vd.get("1-Approved", 0); sm["C27"].value = pct(vd.get("1-Approved", 0))
    sm["B28"].value = vd.get("2-AN", 0); sm["C28"].value = pct(vd.get("2-AN", 0))
    sm["B29"].value = vd.get("3-To be revised", 0); sm["C29"].value = pct(vd.get("3-To be revised", 0))
    sm["B30"].value = vd.get("4-Rejected", 0); sm["C30"].value = pct(vd.get("4-Rejected", 0))

    # 5) ITEMS BY SECTION: SIN cambios (N29 no suma items ni delivered nuevos).

    _guardar_robusto(wb, SRC)


def _guardar_robusto(wb, dst):
    import zipfile
    scratch_dir = ("C:/Users/luisr/AppData/Local/Temp/claude/"
                   "C--SynologyDrive-SynologyDrive-DESAROLLO-PROYECTOS-CLAUDE-"
                   "MODULO-DE-SALMUERA-TALTAL/"
                   "1bd6de80-400c-4556-9eaf-13a3464919fc/scratchpad")
    os.makedirs(scratch_dir, exist_ok=True)
    tmp = os.path.join(scratch_dir, "register_n29_build.xlsx")
    wb.save(tmp)
    if zipfile.ZipFile(tmp).testzip() is not None:
        raise SystemExit("ERROR: el build en scratchpad salio corrupto.")
    openpyxl.load_workbook(tmp).close()
    print(f"Build verificado en scratchpad: {tmp}")
    for intento in range(1, 6):
        shutil.copyfile(tmp, dst)
        try:
            if zipfile.ZipFile(dst).testzip() is None:
                openpyxl.load_workbook(dst).close()
                print(f"Guardado y verificado en destino (intento {intento}): {dst}")
                return
        except Exception as e:
            print(f"  intento {intento}: destino aun corrupto ({e}); recopiando...")
    raise SystemExit(f"ERROR: no se pudo dejar copia integra. Build en {tmp}.")


if __name__ == "__main__":
    main()
