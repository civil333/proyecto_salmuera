#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
update_register_n27.py
Actualiza el Master Deliverable Register a TM N27, RE-ESCOPEADO a E63 + E64
(submittals 25007-0063 + 25007-0064, 6 documentos). Modelado en
update_register_n26.py.

Idempotencia: el backup _pre-N27.xlsx conserva el estado LIMPIO N26 (48/21/12/0).
Al re-correr, este script RESTAURA SRC desde ese backup antes de aplicar, de modo
que nunca duplica filas de Revision History ni acumula cambios. (El updater N26
hacia copyfile(SRC,BAK) al inicio; aqui se invierte para poder re-correr.)

Tally acumulado esperado tras N27: 48 Code 1 / 25 Code 2 / 10 Code 3 / 0 Code 4;
107 items / 83 delivered; 27 TMs / 64 entregas.

Cambios de veredicto (desde el baseline N26 48/21/12/0):
  UPDATES (4 re-revisiones de filas existentes):
    #98  Outline Panel   P22-CD-09-008-001 (Rev B->C, 3 -> 2-AN)  gate resuelto
    #104 RO Vessel Hydro P22-BA-09-000-009 (Rev B->C, 3 -> 2-AN)  CRITICAL cerrado
    #105 HP/LP Pressure  P22-BA-09-000-010 (Rev B->C, 3 -> 3)     driver 75 bar PVC
    #106 Painting        P22-BA-09-000-011 (Rev A->B, 3 -> 2-AN)  cierre N23
  NEW_ITEMS (2 documentos nuevos):
    FAT Procedure  P22-PP-09-000-001 (Rev A, 2-AN)  Section 3 Fabrication & FAT
      -> Code 2: defectos intrinsecos clericales/referencia; RTD (OBS-02 previo)
         reclasificado a NOTE-02 = dependencia del Alarm & Interlock List (Seccion 3)
    O&M Manual     P22-BA-09-000-012 (Rev A, 3-To be revised)  Section 2 Handover
"""
import os
import shutil
import collections
from copy import copy
import openpyxl

EV = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register.xlsx")
BAK = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register_pre-N27.xlsx")

TM = "N27"
TM_DATE = "13-Jul-2026"

MR_UPDATES = {
    # PLC-LCP Outline Panel Drawing Rev C -> Code 2 (gate de fabricacion resuelto)
    "P22-CD-09-008-001": dict(
        rev="C", delivery="E64", tm=TM, verdict="2-AN",
        action=(
            "Rev C delivered (E64). Code 2 - Approved as noted (TM N27). Resolves the "
            "enclosure gate open since TM N20 and carried through TM N25. Exterior "
            "body, door, roof, rear panel, plinth and gland plates in SS316L; interior "
            "mounting components in galvanised/CRS painted RAL 7035; protection class "
            "NEMA 4X / IP66 - the RFI-002 disposition (07-Jul-2026), coherent with the "
            "approved LCP Datasheet and IFC SLD. Sheet 5 states it explicitly (Hoffman "
            "Cabinet Material Specifications); 'SECONDE STAGE' typo corrected. Internal "
            "galvanised/CRS accepted and not reopened. OBS-01 MINOR cable-clamp 30-34mm "
            "'TBC'; OBS-02 MINOR MATERIAL row lists exterior parts as sheet steel "
            "interior only (reword); OBS-03 MINOR label typos (CABLE DNRY, TP1 "
            "described as Profibus vs PLX32-EIP-MBTCP gateway). Issue at IFC Rev 0 with "
            "the three corrections; no new revision. BW commits (sheet 14) to re-issue "
            "the SLD to 'SS316L Panel, NEMA 4X/IP66' - tracked in Section 3."),
    ),
    # RO Vessel Hydrostatic Test Procedure Rev C -> Code 2 (CRITICAL N26 cerrado)
    "P22-BA-09-000-009": dict(
        rev="C", delivery="E64", tm=TM, verdict="2-AN",
        action=(
            "Rev C delivered (E64). Code 2 - Approved as noted (TM N27). Closes the "
            "CRITICAL of TM N26: the 45.5 bar Protec form is removed and replaced by "
            "the actual 17-Jun-2026 test report evidencing the vessels tested at their "
            "required pressures - 91.01 bar on the BPV81200SP7 (1,320 psi = 1.1x1,200) "
            "and 136.52 bar on the BPV81800SP7 (1,980 psi = 1.1x1,800), all units O.K., "
            "with gauge calibration certificates. This is the test for which the ASME "
            "code stamp was waived. OBS-01 MINOR: the body still writes only the "
            "generic 1.1x rule with no binding numeric pressure (state 1,320/1,980 psi "
            "on the face). OBS-02 MINOR: neither attached gauge cert falls in the "
            "procedure's own 1.5x-4x window for the 136.5 bar test. NOTE-01: ASME "
            "waiver 02-Jun not reopened; single calibrated gauge accepted. Issue at IFC "
            "Rev 0 with the two form edits; no new revision, no re-test."),
    ),
    # HP and LP Pressure Test Procedure Rev C -> Code 3 (driver del veredicto)
    "P22-BA-09-000-010": dict(
        rev="C", delivery="E63", tm=TM, verdict="3-To be revised",
        action=(
            "Rev C delivered (E63). Code 3 - To be revised (TM N27). Answers the TM "
            "N26 observation by attaching the Line List instead of writing the "
            "pressure in the body, so that table becomes the executable instruction - "
            "and it is not safe. OBS-01 CRITICAL: the attached Line List sets line "
            "DA-PVC-DN65-09-016 (RO Brine Discharge, PVC SCH 80 DN65, Class 150 "
            "flanges, operating at 1 bar) at a 50 bar design pressure, so the "
            "HYDROTEST column orders 75 bar on a plastic line; testing to it would "
            "rupture the line. The 50 bar design value is carried over from Line List "
            "Rev C (approved TM N18); the new hydrotest column is what makes it "
            "executable. OBS-02 MAJOR: the attachment is labelled Rev 0 (never "
            "transmitted) while the approved Rev C carries no hydrotest column. "
            "OBS-03 MINOR: the body still writes no numeric pressure (135 bar HP / "
            "7.5 bar LP per ITP rows 5.2 and 5.1). OBS-04 MINOR: 5.7.2.2 still "
            "missing (5.6.5 corrected). NOTE-01: report form now attached without "
            "pre-printed pressure (closes the N26 note); add a Required Test Pressure "
            "field. Re-issue as Rev D; HP hydrostatic remains a Hold Point."),
    ),
    # Painting Procedure Rev B -> Code 2 - Approved as Noted (cierra el driver N23)
    "P22-BA-09-000-011": dict(
        rev="B", delivery="E63", tm=TM, verdict="2-AN",
        action=(
            "Rev B delivered (E63). Code 2 - Approved as noted (TM N27). Meets the two "
            "conditions ADASA set at TM N25 for a non-Sherwin-Williams system: the "
            "Jotun technical letter of 02-Jul-2026 (TSS-DD-MYPC039-26) plus the three "
            "product data sheets give the coat-by-coat equivalence against the approved "
            "Painting Specification Rev C (zinc-rich epoxy 80 um Barrier 80 / MIO "
            "high-build epoxy 200 um Penguard Midcoat M20 / aliphatic PU 75 um Hardtop "
            "XP = 355 um), and they confirm ISO 12944 C5-M marine durability. Scope "
            "bounded to ASTM A-36 (SS and non-metallic excluded) and QC completed "
            "(adhesion pull-off, ISO 2808, SSPC-SP10). Open on the Appendix A "
            "inspection form only: OBS-01 MAJOR the BLASTING ACTIVITY criterion still "
            "reads 40-75 um against the 50-80 um of the body and the ET 50 um lower "
            "bound; OBS-02 MINOR 'Specified DFT: xxx um' placeholder; OBS-03 MINOR RAL "
            "5012 not stated. Issue at IFC Rev 0 with the three form corrections; no "
            "new revision required."),
    ),
}

# Filas nuevas: (Document, Code, Rev, Delivery, Verdict, Action, Status)
NEW_ITEMS = [
    ("PLC/LCP FAT Procedure - Hardware", "P22-PP-09-000-001", "A", "E64",
     "2-AN",
     ("Rev A delivered (E64), first issue. Code 2 - Approved as Noted (TM N27). Hardware "
      "FAT complete and correct, consistent with the approved IO List Rev 4 (138 "
      "points, 4 relay-contact DCS signals, Pt-100 winding+bearing both motors). Three "
      "corrections to incorporate at issue, none touching the test: OBS-01 MAJOR "
      "governing drawings cited under wrong/not-issued references (Outline as "
      "P22-ET-09-008-001 vs real P22-CD-09-008-001 Rev C; Schematic as P22-ET-09-008-002 "
      "Rev B, but only Rev A exists); OBS-02 MINOR AO tag BDS-09-001->-002, doc number, "
      "KA2/KA3 label; OBS-03 MINOR mounting-plate HDG vs Outline zinc-plated, dosing-"
      "pump running DI vs soft-BOOL IO List (N20, no signal-type change). NOTE-02: the "
      "RTD table is CORRECT (TE-09-001=winding, TE-09-002=bearing = IO List Rev 4); the "
      "swap lives in the open Alarm & Interlock List Rev B (Section 3, cross-document "
      "dependency, does not degrade the FAT) - the RTD protection sign-off is held "
      "until that list is reconciled and the bearing AHH set to 95C. Issue at IFC Rev 0 "
      "with the corrections; no new revision."),
     "Delivered"),
    ("Operating and Maintenance Manual", "P22-BA-09-000-012", "A", "E64",
     "3-To be revised",
     ("Rev A delivered (E64), first issue. Code 3 - To be revised (TM N27). "
      "Installation/membrane-loading/cartridge content serviceable and project data "
      "correct (6x4 array / 70 membranes), but Section 4 is not self-contained. OBS-01 "
      "MAJOR: control sequence (4.4), setpoints and HMI (4.5) reproduce logic the "
      "Control Philosophy assigns to the Operating Sequence Chart (017, not issued), "
      "the Alarm & Control Setpoint List (015, Rev B Code 3) and the HMI Screenshots "
      "(016, Rev A Code 3) - all open; no approved source. OBS-02 MAJOR: bypass "
      "VE-09-002 opened by TDS (Note 5) contradicts the approved Control Philosophy "
      "Rev D ('TDS shall not directly initiate bypass; governed by the pressure "
      "control loop'); + TDS/pressure conflation + 'Interstage' vs 'Feed' tag. OBS-03 "
      "MINOR: setpoints 52/55 vs 51 bar, recovery 42 vs 42.86%, flows. OBS-04 MINOR: "
      "HMI faceplates default Rockwell + simulation data. OBS-05 MINOR: CIP cross-ref, "
      "membrane model, boilerplate. Re-issue as Rev B; cannot go to IFC until the "
      "control documents close."),
     "Delivered"),
]

RH_NEW = [
    (None, "PLC-LCP Outline Panel Drawing", "P22-CD-09-008-001", "C", "E64",
     "25007-0064", "2-AN",
     "Rev C. Code 2 - Approved as noted. Resolves the enclosure gate (TM N20/N25): "
     "exterior body/door/roof/rear panel/plinth/gland plates SS316L, interior "
     "galvanised/CRS RAL 7035, NEMA 4X/IP66 per the RFI-002 disposition; sheet 5 "
     "material spec + 'SECOND STAGE' typo fixed. OBS-01/02/03 MINOR: cable-clamp TBC, "
     "MATERIAL-row wording, label typos. BW commits to re-issue the SLD (Section 3)."),
    (None, "RO Vessel Hydrostatic Test Procedure", "P22-BA-09-000-009", "C", "E64",
     "25007-0064", "2-AN",
     "Rev C. Code 2 - Approved as noted. Closes the CRITICAL of TM N26: the 45.5 bar "
     "form is removed; the 17-Jun test report evidences the vessels at 91.01 bar "
     "(1,320 psi) and 136.52 bar (1,980 psi) = 1.1x design, all O.K., with gauge "
     "certs. OBS-01/02 MINOR: state the binding numeric pressures in the body; "
     "reconcile the gauge-selection window. ASME waiver not reopened."),
    (None, "HP and LP Pressure Test Procedure", "P22-BA-09-000-010", "C", "E63",
     "25007-0063", "3-To be revised",
     "Rev C. Code 3. Attaches the Line List instead of writing the pressure in the "
     "body, making that table the executable instruction. OBS-01 CRITICAL: line "
     "DA-PVC-DN65-09-016 (PVC SCH 80 DN65, Class 150, operating 1 bar) carries a 50 "
     "bar design pressure, so the hydrotest column orders 75 bar on a plastic line. "
     "OBS-02 MAJOR: attachment labelled Rev 0, never transmitted; the approved Rev C "
     "has no hydrotest column. OBS-03/04 MINOR: no numeric pressure in the body; "
     "5.7.2.2 still missing. NOTE-01: report form attached, no pre-printed pressure. "
     "Re-issue as Rev D."),
    (None, "Painting Procedure", "P22-BA-09-000-011", "B", "E63", "25007-0063",
     "2-AN",
     "Rev B. Code 2 - Approved as noted. Closes the two TM N23 major observations: "
     "coat-by-coat equivalence of the Jotun system against the approved Painting "
     "Specification Rev C (Jotun letter TSS-DD-MYPC039-26 + 3 TDS) and ISO 12944 C5-M "
     "marine durability. Scope bounded to ASTM A-36; QC completed (adhesion pull-off, "
     "ISO 2808, SSPC-SP10). Open only on the Appendix A inspection form: BLASTING "
     "ACTIVITY criterion still 40-75 um (OBS-01), 'xxx um' DFT placeholder (OBS-02), "
     "RAL 5012 absent (OBS-03). Issue at IFC Rev 0 with those corrections."),
    (None, "PLC/LCP FAT Procedure - Hardware", "P22-PP-09-000-001", "A", "E64",
     "25007-0064", "2-AN",
     "Rev A, first issue. Code 2 - Approved as Noted. Hardware FAT complete and correct, "
     "consistent with the IO List Rev 4; three corrections to incorporate at issue. "
     "OBS-01 MAJOR: governing drawings cited under wrong/not-issued references (Outline "
     "as ET-001 vs CD-001 Rev C; Schematic as ET-002 Rev B, only Rev A exists). OBS-02/03 "
     "MINOR: tags/typos; material and signal type vs IO List. NOTE-02: RTD table CORRECT "
     "(TE-09-001=winding = IO List Rev 4); swap lives in the open Alarm & Interlock List "
     "Rev B (Section 3, dependency, does not degrade the FAT) - RTD protection sign-off "
     "held until that list is reconciled (bearing AHH 95C). Issue at IFC Rev 0; no new "
     "revision."),
    (None, "Operating and Maintenance Manual", "P22-BA-09-000-012", "A", "E64",
     "25007-0064", "3-To be revised",
     "Rev A, first issue. Code 3 - To be revised. Non-control content serviceable, but "
     "Section 4 is not self-contained. OBS-01 MAJOR: control sequence/setpoints/HMI "
     "reproduce logic from the Operating Sequence Chart (017, not issued), Alarm & "
     "Control Setpoint List (015, Rev B Code 3) and HMI Screenshots (016, Rev A Code "
     "3), all open. OBS-02 MAJOR: bypass VE-09-002 opened by TDS contradicts the "
     "approved Control Philosophy Rev D (pressure control loop). OBS-03/04/05 MINOR: "
     "setpoints, HMI unconfigured, housekeeping. Re-issue as Rev B; no IFC until the "
     "control documents close."),
]


def style_row(ws, dst_row, ref_row, ncols):
    for c in range(1, ncols + 1):
        s = ws.cell(ref_row, c)
        d = ws.cell(dst_row, c)
        d._style = copy(s._style)


def main():
    # Idempotencia: restaurar SRC desde el baseline limpio (N26) si el backup existe.
    if os.path.exists(BAK):
        shutil.copyfile(BAK, SRC)
        print(f"Restaurado SRC desde baseline limpio: {BAK}")
    else:
        shutil.copyfile(SRC, BAK)
        print(f"Baseline creado: {BAK}")

    wb = openpyxl.load_workbook(SRC)  # data_only=False: preserva formulas/estilos
    mr = wb["Master Register"]
    rh = wb["Revision History"]
    sm = wb["Summary"]

    ncols_mr = mr.max_column
    hdr = {mr.cell(1, c).value: c for c in range(1, ncols_mr + 1)}
    C_NUM = hdr["#"]
    C_DOC = hdr["Document"]
    C_CODE = hdr["Code / ET Reference"]
    C_REV = hdr["Rev"]; C_DEL = hdr["Delivery"]; C_TM = hdr["TM"]
    C_VER = hdr["Verdict"]; C_STA = hdr["Status"]; C_ACT = hdr["Action Required"]

    # 1) updates a filas existentes (las 4 re-revisiones)
    seen = set()
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code in MR_UPDATES:
            u = MR_UPDATES[code]
            mr.cell(r, C_REV).value = u["rev"]
            mr.cell(r, C_DEL).value = u["delivery"]
            mr.cell(r, C_TM).value = u["tm"]
            mr.cell(r, C_VER).value = u["verdict"]
            mr.cell(r, C_STA).value = "Delivered"
            mr.cell(r, C_ACT).value = u["action"]
            seen.add(code)
            print(f"  MR update row {r}: {code} -> {u['verdict']} / {u['tm']} "
                  f"(Rev {u['rev']})")
    missing = set(MR_UPDATES) - seen
    if missing:
        raise SystemExit(
            f"ERROR: codigos N27 no encontrados en Master Register: {missing}")

    # 2) filas NUEVAS (FAT + O&M). Guardarrail: no deben existir ya.
    existing = set()
    max_num = 0
    ref_mr_row = 2
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code:
            existing.add(code)
            ref_mr_row = r
            try:
                max_num = max(max_num, int(mr.cell(r, C_NUM).value))
            except (TypeError, ValueError):
                pass
    for i, (doc, code, rev, dele, ver, act, sta) in enumerate(NEW_ITEMS, start=1):
        if code in existing:
            raise SystemExit(f"ERROR: NEW_ITEM {code} ya existe en el register.")
        dst = ref_mr_row + i
        style_row(mr, dst, ref_mr_row, ncols_mr)
        mr.cell(dst, C_NUM).value = max_num + i
        mr.cell(dst, C_DOC).value = doc
        mr.cell(dst, C_CODE).value = code
        mr.cell(dst, C_REV).value = rev
        mr.cell(dst, C_DEL).value = dele
        mr.cell(dst, C_TM).value = TM
        mr.cell(dst, C_VER).value = ver
        mr.cell(dst, C_STA).value = sta
        mr.cell(dst, C_ACT).value = act
        print(f"  MR new row {dst}: #{max_num + i} {code} -> {ver}")

    # numero de registro por codigo (para los RH)
    reg_num = {}
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code:
            try:
                reg_num[code] = int(mr.cell(r, C_NUM).value)
            except (TypeError, ValueError):
                pass

    # 3) Revision History (6 filas nuevas)
    rcols = rh.max_column
    rhh = {rh.cell(1, c).value: c for c in range(1, rcols + 1)}
    prior = collections.Counter()
    rh_last = 1
    for r in range(2, rh.max_row + 1):
        cd = rh.cell(r, rhh["Code / ET Reference"]).value
        if cd:
            prior[str(cd).strip()] += 1
            rh_last = r
    for i, row in enumerate(RH_NEW, start=1):
        dst = rh_last + i
        _num, doc, code, rev, dele, sub, ver, keyobs = row
        num = reg_num.get(code, "")
        cycle = prior[code] + 1
        prior[code] += 1  # por si un mismo code apareciera 2 veces en RH_NEW
        style_row(rh, dst, rh_last, rcols)
        rh.cell(dst, rhh["#"]).value = num
        rh.cell(dst, rhh["Document"]).value = doc
        rh.cell(dst, rhh["Code / ET Reference"]).value = code
        rh.cell(dst, rhh["Rev"]).value = rev
        rh.cell(dst, rhh["Delivery"]).value = dele
        rh.cell(dst, rhh["Submittal"]).value = sub
        rh.cell(dst, rhh["TM"]).value = TM
        rh.cell(dst, rhh["TM Date"]).value = TM_DATE
        rh.cell(dst, rhh["Verdict"]).value = ver
        rh.cell(dst, rhh["Cycle"]).value = cycle
        rh.cell(dst, rhh["Key Observations"]).value = keyobs
        print(f"  RH new row {dst}: #{num} {code} cycle {cycle} -> {ver}")

    # 4) Summary: recomputar desde Master Register
    rows = []
    for r in range(2, mr.max_row + 1):
        code = mr.cell(r, C_CODE).value
        if code and str(code).strip():
            rows.append((str(mr.cell(r, C_VER).value or "").strip(),
                         str(mr.cell(r, C_STA).value or "")))
    total = len(rows)
    delivered = sum(1 for v, s in rows if "Delivered" in s)
    vd = collections.Counter(v for v, s in rows if "Delivered" in s)
    print(f"  Summary recompute: total={total} delivered={delivered} "
          f"verdicts={dict(vd)}")

    sm["B4"].value = total
    sm["B5"].value = delivered
    sm["B9"].value = "27 (N1 through N27 - TM N5 issued in Rev 0 and Rev 1)"
    sm["B10"].value = "64 (E1 through E64)"
    sm["B12"].value = "13-Jul-2026 (E64)"
    sm["B13"].value = "13-Jul-2026"

    def pct(n):
        return f"{round(100 * n / delivered)}%" if delivered else "0%"
    sm["B27"].value = vd.get("1-Approved", 0)
    sm["C27"].value = pct(vd.get("1-Approved", 0))
    sm["B28"].value = vd.get("2-AN", 0)
    sm["C28"].value = pct(vd.get("2-AN", 0))
    sm["B29"].value = vd.get("3-To be revised", 0)
    sm["C29"].value = pct(vd.get("3-To be revised", 0))
    sm["B30"].value = vd.get("4-Rejected", 0)
    sm["C30"].value = pct(vd.get("4-Rejected", 0))

    # 5) ITEMS BY SECTION: +1 a Section 2 Handover (O&M) y +1 a Section 3 Fab&FAT (FAT).
    #    (Se parte del baseline limpio: row19 7/0, row20 10/7.)
    sm.cell(19, 2).value = (sm.cell(19, 2).value or 0) + 1  # Handover total
    sm.cell(19, 3).value = (sm.cell(19, 3).value or 0) + 1  # Handover delivered
    sm.cell(20, 2).value = (sm.cell(20, 2).value or 0) + 1  # Fab & FAT total
    sm.cell(20, 3).value = (sm.cell(20, 3).value or 0) + 1  # Fab & FAT delivered
    print(f"  ITEMS BY SECTION: Handover -> {sm.cell(19,2).value}/{sm.cell(19,3).value}, "
          f"Fab&FAT -> {sm.cell(20,2).value}/{sm.cell(20,3).value}")
    print("  NOTA: REVIEW CYCLES distribution no se auto-modifica (recuento manual).")

    _guardar_robusto(wb, SRC)


def _guardar_robusto(wb, dst):
    """Guarda en scratchpad (fuera de Synology), verifica integridad y copia al
    destino con reintento+verificacion. Evita la corrupcion por sync de Synology
    Drive durante la escritura (worksheet xml con 'Bad magic number')."""
    import zipfile
    scratch_dir = (r"C:\Users\luisr\AppData\Local\Temp\claude"
                   r"\C--SynologyDrive-SynologyDrive-DESAROLLO-PROYECTOS-CLAUDE-"
                   r"MODULO-DE-SALMUERA-TALTAL\32b8961e-5e28-4e6f-a12d-653b30834a1d"
                   r"\scratchpad")
    tmp = os.path.join(scratch_dir, "register_n27_build.xlsx")
    wb.save(tmp)
    if zipfile.ZipFile(tmp).testzip() is not None:
        raise SystemExit("ERROR: el build en scratchpad salio corrupto.")
    openpyxl.load_workbook(tmp).close()  # verifica que openpyxl lo lee
    print(f"Build verificado en scratchpad: {tmp}")

    for intento in range(1, 6):
        shutil.copyfile(tmp, dst)
        try:
            if zipfile.ZipFile(dst).testzip() is None:
                openpyxl.load_workbook(dst).close()
                print(f"Guardado y verificado en destino (intento {intento}): {dst}")
                return
        except Exception as e:
            print(f"  intento {intento}: destino aun corrupto ({e}); recopiando...")
    raise SystemExit(
        "ERROR: no se pudo dejar una copia integra en el destino (Synology sync). "
        f"El build integro esta en {tmp} — copiar manualmente cuando el sync se calme.")


if __name__ == "__main__":
    main()
