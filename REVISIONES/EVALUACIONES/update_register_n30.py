#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
update_register_n30.py
Actualiza el Master Deliverable Register a TM N30 (E68, submittal 25007-0068,
1 documento, re-revision a Rev 0 IFC). Modelado en update_register_n29.py.

Idempotencia: el backup _pre-N30.xlsx conserva el estado LIMPIO N29 (49/29/6/0).
Al re-correr, RESTAURA SRC desde ese backup antes de aplicar.

Veredicto N30: 1 - Approved. Tally 1 Code 1.
  UPDATES (1 re-revision de fila existente; NO hay items nuevos):
    #29 DS PLC & HMI  P22-ET-09-008-001  (Rev C->0, 2-AN -> 1-Approved)

Tally esperado tras N30: 50 Code 1 / 28 Code 2 / 6 Code 3 / 0 Code 4 (el item #29
pasa de 2-AN a 1-Approved); 108 items / 84 delivered; 30 TMs / 68 entregas.
"""
import os
import shutil
import collections
from copy import copy
import openpyxl

EV = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register.xlsx")
BAK = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register_pre-N30.xlsx")

TM = "N30"
TM_DATE = "23-Jul-2026"

MR_UPDATES = {
    # DS PLC & HMI Rev 0 (IFC) -> Code 1 (cierra las 2 NOTE del N26; la
    # reconciliacion del catalogo del terminal cae en OTROS 4 documentos -> Seccion 3)
    "P22-ET-09-008-001": dict(
        rev="0", delivery="E68", tm=TM, verdict="1-Approved",
        action=(
            "Rev 0 (IFC) delivered (E68). Code 1 - Approved (TM N30). Closes the two "
            "TM N26 items: the document code now reads P22-ET-09-008-001 on the cover "
            "and on all seven component headers (only the delivered file name keeps the "
            "former two-digit code), and the RTD-module quantity is confirmed as two "
            "5069-IY4 units on the Control System Architecture Rev D sheet "
            "P22-CD-09-004-001-P2 - eight RTD channels against the four temperature "
            "elements of the I/O List Rev 5. No technical content changed between Rev C "
            "and Rev 0; the datasheet requires no modification to itself. BINDING "
            "DECLARATION: ADASA declares the operator-terminal catalogue number to be "
            "2711P-T10C22D9P (stated here since Rev A; fixed at TM N22 as the hardware "
            "basis for the HMI screen design). The 2711P-T10C21D8S carried by the Outline "
            "Rev 0 IFC bill of material, the Schematic Rev A, the FAT Procedure Rev A and "
            "the Control System Architecture Rev D has a single 10/100Base-T port and 512 "
            "MB RAM, and does not meet the 2 x Ethernet RJ45 and 1 GB of the requirement "
            "rows; those four documents are to be aligned, tracked in Section 3. Not "
            "attributable to BW Water: ADASA coded documents on both sides (N4 validated "
            "the -D8S, N14 approved the Architecture Code 1, N21 accepted the -D9P). Six "
            "documentation items to correct at the next natural issue: the comment sheet "
            "dropped the closure record of the three TM N21 items (including the HART "
            "disposition) and cites a non-existent Rev D; the cover states 49 pages "
            "against 50; the seven headers keep the IFA stamp on an IFC submission; the "
            "analog-output sheet carries a digital-output design intent and leaves its "
            "model as 5069-OF4/OF8; the headers carry no revision index."),
    ),
}

# N30 no incorpora documentos nuevos: es una re-revision.
NEW_ITEMS = []

# Higiene del registro (no es alcance del N30, pero la celda obsoleta es la fuente
# del arrastre que contamino la Seccion 3 de los TM N26 a N29): el HMI Display
# Screenshot NO esta "sin entregar" desde el 15-Jun-2026 - se entrego Rev A en la
# E50 y volvio Code 3 en el TM N22. Solo se corrige "Action Required"; el veredicto
# de la fila ya era correcto.
MR_ACTION_FIXES = {
    "P22-LI-09-008-016": (
        "Rev A delivered (E50, 15-Jun-2026). Code 3 - To be revised (TM N22). First real "
        "screen design; materially closes the HMI commitment opened at TM N4 (former code "
        "P22-BREAD-09-008-001), so the item is no longer 'never submitted'. Open against "
        "the Technical Specification: OBS-01 electrical-variables and energy screen "
        "(kWh/m3), OBS-02 process-variable trending screen, OBS-03 setpoint and "
        "parameterization screen with safe limits, OBS-04 missing process screens and "
        "overview with tags per the approved P&ID and Instrument List, OBS-05 ISA-101 "
        "conformance evidence (full visual verification reserved for the FAT). Re-issue "
        "as Rev B; gates the O&M Manual Rev B."),
}

RH_NEW = [
    (None, "DS PLC & HMI Panel Component", "P22-ET-09-008-001", "0", "E68",
     "25007-0068", "1-Approved",
     "Rev 0, IFC. Code 1 - Approved. Both N26 items closed and verified: document code "
     "aligned to the three-digit P22-ET-09-008-001 on cover and all seven component "
     "headers; 5069-IY4 quantity (2 pcs, 8 RTD channels) confirmed on the Control System "
     "Architecture Rev D sheet P2, covering the four RTD points of the I/O List Rev 5. No "
     "technical content changed from Rev C; no modification required to this document. "
     "ADASA declares the binding operator-terminal catalogue number to be 2711P-T10C22D9P "
     "(2 x Ethernet RJ45, 1 GB RAM per rows 15/18); the 2711P-T10C21D8S in the Outline Rev "
     "0 BOM, Schematic Rev A, FAT Procedure Rev A and Control System Architecture Rev D "
     "(single Ethernet port, 512 MB) is to be corrected in those four documents - tracked "
     "in Section 3. Six documentation items to the next natural issue (comment sheet lost "
     "the TM N21 closure record and cites a non-existent Rev D; page count 49 vs 50; IFA "
     "stamp on an IFC issue; analog-output sheet design intent and OF4/OF8 model; no "
     "revision index on headers)."),
]


def style_row(ws, dst_row, ref_row, ncols):
    for c in range(1, ncols + 1):
        s = ws.cell(ref_row, c)
        d = ws.cell(dst_row, c)
        d._style = copy(s._style)


def main():
    if os.path.exists(BAK):
        shutil.copyfile(BAK, SRC)
        print(f"Restaurado SRC desde baseline limpio: {BAK}")
    else:
        shutil.copyfile(SRC, BAK)
        print(f"Baseline creado: {BAK}")

    wb = openpyxl.load_workbook(SRC)
    mr = wb["Master Register"]
    rh = wb["Revision History"]
    sm = wb["Summary"]

    ncols_mr = mr.max_column
    hdr = {mr.cell(1, c).value: c for c in range(1, ncols_mr + 1)}
    C_NUM = hdr["#"]; C_CODE = hdr["Code / ET Reference"]
    C_REV = hdr["Rev"]; C_DEL = hdr["Delivery"]; C_TM = hdr["TM"]
    C_VER = hdr["Verdict"]; C_STA = hdr["Status"]; C_ACT = hdr["Action Required"]

    # 1) updates (1 re-revision)
    seen = set()
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code in MR_UPDATES:
            u = MR_UPDATES[code]
            mr.cell(r, C_REV).value = u["rev"]
            mr.cell(r, C_DEL).value = u["delivery"]
            mr.cell(r, C_TM).value = u["tm"]
            mr.cell(r, C_VER).value = u["verdict"]
            mr.cell(r, C_STA).value = "Delivered"
            mr.cell(r, C_ACT).value = u["action"]
            seen.add(code)
            print(f"  MR update row {r}: {code} -> {u['verdict']} / {u['tm']} (Rev {u['rev']})")
    missing = set(MR_UPDATES) - seen
    if missing:
        raise SystemExit(f"ERROR: codigos N30 no encontrados en Master Register: {missing}")

    # 1b) higiene: refrescar celdas Action Required obsoletas (sin tocar veredictos)
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code in MR_ACTION_FIXES:
            mr.cell(r, C_ACT).value = MR_ACTION_FIXES[code]
            print(f"  MR action fix row {r}: {code} (celda obsoleta actualizada)")

    reg_num = {}
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code:
            try:
                reg_num[code] = int(mr.cell(r, C_NUM).value)
            except (TypeError, ValueError):
                pass

    # 2) Revision History (1 fila nueva)
    rcols = rh.max_column
    rhh = {rh.cell(1, c).value: c for c in range(1, rcols + 1)}
    prior = collections.Counter(); rh_last = 1
    for r in range(2, rh.max_row + 1):
        cd = rh.cell(r, rhh["Code / ET Reference"]).value
        if cd:
            prior[str(cd).strip()] += 1; rh_last = r
    for i, row in enumerate(RH_NEW, start=1):
        dst = rh_last + i
        _num, doc, code, rev, dele, sub, ver, keyobs = row
        num = reg_num.get(code, "")
        cycle = prior[code] + 1; prior[code] += 1
        style_row(rh, dst, rh_last, rcols)
        rh.cell(dst, rhh["#"]).value = num
        rh.cell(dst, rhh["Document"]).value = doc
        rh.cell(dst, rhh["Code / ET Reference"]).value = code
        rh.cell(dst, rhh["Rev"]).value = rev
        rh.cell(dst, rhh["Delivery"]).value = dele
        rh.cell(dst, rhh["Submittal"]).value = sub
        rh.cell(dst, rhh["TM"]).value = TM
        rh.cell(dst, rhh["TM Date"]).value = TM_DATE
        rh.cell(dst, rhh["Verdict"]).value = ver
        rh.cell(dst, rhh["Cycle"]).value = cycle
        rh.cell(dst, rhh["Key Observations"]).value = keyobs
        print(f"  RH new row {dst}: #{num} {code} cycle {cycle} -> {ver}")

    # 3) Summary: recomputar total/delivered + verdicts desde MR
    stat = collections.Counter(); vd = collections.Counter(); total = 0
    for r in range(2, mr.max_row + 1):
        code = mr.cell(r, C_CODE).value
        if code and str(code).strip():
            total += 1
            su = str(mr.cell(r, C_STA).value or "").strip().upper()
            stat[su] += 1
            if su == "DELIVERED":
                vd[str(mr.cell(r, C_VER).value or "").strip()] += 1
    delivered = stat.get("DELIVERED", 0)
    not_delivered = stat.get("NOT DELIVERED", 0)
    partial = stat.get("PARTIAL", 0)
    print(f"  Summary recompute: total={total} delivered={delivered} "
          f"not_delivered={not_delivered} partial={partial} verdicts={dict(vd)}")

    sm["B4"].value = total
    sm["B5"].value = delivered
    sm["B7"].value = partial
    sm["B8"].value = not_delivered
    sm["B9"].value = "30 (N1 through N30 - TM N5 issued in Rev 0 and Rev 1)"
    sm["B10"].value = "68 (E1 through E68)"
    sm["B12"].value = "23-Jul-2026 (E68)"
    sm["B13"].value = "23-Jul-2026"

    def pct(n):
        return f"{round(100 * n / delivered)}%" if delivered else "0%"
    sm["B27"].value = vd.get("1-Approved", 0); sm["C27"].value = pct(vd.get("1-Approved", 0))
    sm["B28"].value = vd.get("2-AN", 0); sm["C28"].value = pct(vd.get("2-AN", 0))
    sm["B29"].value = vd.get("3-To be revised", 0); sm["C29"].value = pct(vd.get("3-To be revised", 0))
    sm["B30"].value = vd.get("4-Rejected", 0); sm["C30"].value = pct(vd.get("4-Rejected", 0))

    # 4) ITEMS BY SECTION: SIN cambios (N30 no suma items ni delivered nuevos).

    wb.save(SRC)
    print(f"\nGuardado: {SRC}")
    print("Tally esperado: 50 C1 / 28 C2 / 6 C3 / 0 C4 | 108 items / 84 delivered | 30 TMs / 68 entregas")


if __name__ == "__main__":
    main()
