#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
update_register_n26.py
Actualiza el Master Deliverable Register a TM N26 (submittals 25007-0057 E57 a
25007-0062 E62, 9 documentos). Modelado en update_register_n25.py.

Particularidad de TM N26: 7 documentos son re-revisiones (UPDATE de filas
existentes) y 2 son NUEVOS (ADD de filas: UHPRO Structural Calculation Report y
GA of CIP Flushing Tank) -> Total 103->105, Delivered 79->81, Section 1
(Engineering) Total 81->83 / Delivered 73->75.

Tally TM N26: 2 Code 1 + 3 Code 2 + 4 Code 3. Veredicto global 3 - To Be Revised.
Distribucion de veredictos recomputada desde la hoja.

Mapeo (matching por codigo de registro):
  UPDATES (7):
    #92  ITP P22-BA-09-000-004 (Rev C->0, N23 2-AN -> N26 1-Approved; IFC)
    #84  NDE Plan P22-BA-09-000-005 (Rev B->C, N23 3 -> N26 1-Approved)
    #104 RO Vessel Hydrostatic P22-BA-09-000-009 (Rev A->B, N23 3 -> N26 3)
    #105 HP/LP Pressure Test P22-BA-09-000-010 (Rev A->B, N23 3 -> N26 3)
    #29  DS PLC & HMI P22-ET-09-008-001 (Rev B->C, N21 3 -> N26 2-AN; entregado
         re-codificado como P22-ET-09-008-01 - matching por el codigo de registro
         008-001, recode anotado en Action, NOTE-01 pendiente de alinear)
    #53  GA Antiscalant Dosing Tank P22-DWG-09-005-015 (Rev A->B, N10 2-AN -> N26 3)
    #57  GA Antiscalant Dosing Pump P22-DWG-09-005-011 (Rev A->B, N11 2-AN -> N26 2-AN)
  NEW (2):
    UHPRO Structural Calculation Report P22-CD-09-005-001 Rev A (N26 3)
    GA of CIP Flushing Tank P22-DWG-09-005-014 Rev A (N26 2-AN)
"""
import os
import shutil
import collections
from copy import copy
import openpyxl

EV = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register.xlsx")
BAK = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register_pre-N26.xlsx")

TM = "N26"
TM_DATE = "06-Jul-2026"
ET_DEADLINE_ENG = "ET Sec.7: Max. 90 days from NTP"

# ---- Master Register: updates a filas existentes (por codigo de registro) ----
MR_UPDATES = {
    # ITP Rev 0 -> Code 1 - Approved (IFC)
    "P22-BA-09-000-004": dict(
        rev="0", delivery="E57", tm=TM, verdict="1-Approved",
        action=(
            "Rev 0 delivered (E57). Code 1 - Approved (TM N26). Closes the TM N23 "
            "cycle: Rev 0 is the IFC and incorporates both edits required at TM N23 "
            "with no intermediate revision - row 2.2 reads 'manufacture to ASME "
            "Section X, without code stamp' and keeps the binding vessel test "
            "pressures (1800 psi x 1.1 and 1200 psi x 1.1), and the RO vessel "
            "hydrostatic test is raised from Witness to Hold Point on ADASA's column. "
            "ASME Section X correct for FRP vessels, stamp waived. NOTE-01: cite the "
            "02-Jun-2026 waiver as the traceable basis in cell 2.2 at issue "
            "(non-gating). The binding-pressure alignment lives on the RO Vessel "
            "Hydrostatic Test Procedure (#104), still Code 3. Issue at IFC Rev 0."),
    ),
    # NDE Plan Rev C -> Code 1 - Approved
    "P22-BA-09-000-005": dict(
        rev="C", delivery="E61", tm=TM, verdict="1-Approved",
        action=(
            "Rev C delivered (E61). Code 1 - Approved (TM N26). Closes the TM N23 "
            "cycle (3rd review cycle A->B->C): the reference section now fixes the "
            "governing code editions (ASME V-2025, II-2025, B31.3-2024, AWS D1.1-2025, "
            "DVS 2202-1) and the acceptance criteria map each code to its joint family "
            "(B31.3 to the Super Duplex HP circuit, Normal Fluid Service para 341.3.2; "
            "AWS D1.1 to structure; DVS 2202-1 to LP thermoplastic), with the UT "
            "column clarified as a thickness verification. RO vessel scope correctly "
            "sits with the ITP and the RO Vessel Hydrostatic Test Procedure. Coverage "
            "matches the ET. Issue at IFC Rev 0."),
    ),
    # RO Vessel Hydrostatic Test Procedure Rev B -> Code 3
    "P22-BA-09-000-009": dict(
        rev="B", delivery="E61", tm=TM, verdict="3-To be revised",
        action=(
            "Rev B delivered (E61). Code 3 - To be revised (TM N26). Does not close "
            "the TM N23 driver. OBS-01 CRITICAL: the body states no binding test "
            "pressure (only the generic 1.1x ASME / 1.43x CE rule) and the embedded "
            "Protec report form still prints 45.5 bar (~660 psi) against the required "
            "1980 psi (BPV-8-1800-SP-7) / 1320 psi (BPV-8-1200-SP-7) fixed by ITP row "
            "2.2 and the 02-Jun waiver - this is the test for which the ASME stamp was "
            "waived. OBS-02 MAJOR: single gauge, no certificate review (witness "
            "covered by the ITP Hold Point). OBS-03 MINOR: Protec 'Revision 0' left on "
            "one index page. Re-issue as Rev C: state the project test pressure by "
            "model in body and form, remove the 45.5 bar default and the CE branch, "
            "add two-gauge certificate-review instrumentation."),
    ),
    # HP and LP Pressure Test Procedure Rev B -> Code 3
    "P22-BA-09-000-010": dict(
        rev="B", delivery="E61", tm=TM, verdict="3-To be revised",
        action=(
            "Rev B delivered (E61). Code 3 - To be revised (TM N26). Partially closes "
            "the TM N23 observation: Rev B adds the correct ASME B31.3 factor (1.5x "
            "hydrostatic, 1.1x pneumatic) but OBS-01 MAJOR remains - the body still "
            "writes no numeric test pressure (135 bar HP / 7.5 bar LP fixed by the "
            "ITP; both steps defer to the line list) and does not reconcile the HP "
            "design pressure (90 bar ITP vs up to 120 bar ET), leaving the result "
            "indeterminate between 135 and 180 bar on the Super Duplex circuit. OBS-02 "
            "MINOR: step 5.6.5 still missing and 5.7.2 skips 5.7.2.2. Re-issue as Rev "
            "C: state the design pressure per subsystem with a cited source and write "
            "the numeric test pressure in the body."),
    ),
    # DS PLC & HMI Panel Component Rev C -> Code 2 - Approved as Noted
    "P22-ET-09-008-001": dict(
        rev="C", delivery="E58", tm=TM, verdict="2-AN",
        action=(
            "Rev C delivered (E58) as P22-ET-09-008-01 (2-digit code on the cover vs "
            "the register/Rev B 3-digit P22-ET-09-008-001; NOTE-01 confirm/align). "
            "Code 2 - Approved as Noted (TM N26). Closes the three TM N21 items: HART "
            "acquisition provided at the instrument level (external 250-ohm resistor "
            "note on the 5069-IF8 and the added 5069-IY4 for handheld HART) - "
            "satisfies 4-20 mA + HART at the instrument, no central HART decoding "
            "required (the TM N24 over-reach is not reopened); the 5069-IY4 RTD module "
            "for the motor Pt-100 channels is included; the title-block typo is fixed "
            "in the body. Remaining: NOTE-01 align the document code to "
            "P22-ET-09-008-001 on the cover and all headers (the item requiring a "
            "change to the datasheet); NOTE-02 confirm the LCP Datasheet governs the "
            "RTD-module quantity. Fold at IFC Rev 0, no new revision."),
    ),
    # GA Antiscalant Dosing Tank Rev B -> Code 3
    "P22-DWG-09-005-015": dict(
        rev="B", delivery="E60", tm=TM, verdict="3-To be revised",
        action=(
            "Rev B delivered (E60). Code 3 - To be revised (TM N26). Partially closes "
            "TM N10 OBS-05: the total capacity (335 L) and the polyethylene body "
            "material are now declared and compatible with the antiscalant service (2 "
            "of 4 items). OBS-01 MAJOR: the NCh 2369 seismic reaction loads and the "
            "anchor-bolt pattern (Zone 3, operating weight at the C.O.G.) are still "
            "absent - Detail 3 shows only the anchor-lug geometry and BW Water's "
            "comment sheet defers them, leaving the OOCC foundation without input. "
            "OBS-02 MINOR: effective working volume (0.27 m3) not stated. Re-issue as "
            "Rev C: add the anchor pattern and the seismic reaction loads (or "
            "reference the endorsed Module Seismic Calculation Report), state the "
            "effective working volume, and label tag TK-09-002."),
    ),
    # GA Antiscalant Dosing Pump Skid Rev B -> Code 2 - Approved as Noted
    "P22-DWG-09-005-011": dict(
        rev="B", delivery="E62", tm=TM, verdict="2-AN",
        action=(
            "Rev B delivered (E62). Code 2 - Approved as Noted (TM N26). Materially "
            "closes TM N11 NOTE-06: Rev B delivers the anchor-bolt layout (M10, 10 "
            "bolts, 120 mm embedment, edge/spacing distances) with a dimensioned plan, "
            "the NCh 2369 seismic reaction forces, and the equipment mass (100 kg). "
            "OBS-01 MINOR: the per-bolt Bolting Embedment forces (Fz 0.205 kN) are "
            "inconsistent with the total reaction block (total Fz 0.2048 kN) over the "
            "10-bolt count - label each block per-bolt or group-total and reconcile "
            "the per-bolt Fz. The endorsed Module Seismic Calculation Report and the "
            "final bolting details are a separate deliverable (Section 3). Fold at IFC "
            "Rev 0, no new revision."),
    ),
}

# ---- Master Register: filas NUEVAS (documentos entregados por primera vez) ----
# (code, document, rev, delivery, verdict, action)
NEW_ITEMS = [
    ("P22-CD-09-005-001", "UHPRO Structural Calculation Report", "A", "E59",
     "3-To be revised",
     "Rev A delivered (E59). Code 3 - To be revised (TM N26). First issue; the "
     "seismic calculation report tracked as a TM N25 deliverable. Confirms the "
     "frame and lifting basis of the approved Structural Design Criteria Rev B "
     "(NCh 2369 Of.2003 Zone 3, R=3.0, I=1.0, A0=0.40g, clause 4.5 combinations; "
     "lifting API RP 2A-WSD 1.35/2.0). OBS-01 CRITICAL: the Bolt Design at the Base "
     "checks 8 ancillary units but omits the main process equipment (HP pump "
     "BH-09-001, turbos SIP-09-001/002, RO filter FIL-09-001, pressure vessels "
     "BOI-09-001/002 = 4160 kg) the ET requires anchored for NCh 2369 Zone 3. OBS-02 "
     "MINOR: design seismic weight and base shear not consolidated in the body. "
     "OBS-03 MINOR: load-combination matrix duplicates 224/225 and omits the EOX "
     "min-gravity case. OBS-04 MINOR: skid-frame max utilization not stated. OBS-05 "
     "MAJOR: reconcile the frame material (Corten A / S275JR) vs the Design Criteria "
     "and the ET A-36 reference; unify the revision label. Re-issue as Rev B; "
     "provide the foundation base-bolt interface table (NOTE-01)."),
    ("P22-DWG-09-005-014", "GA of CIP Flushing Tank", "A", "E60",
     "2-AN",
     "Rev A delivered (E60). Code 2 - Approved as Noted (TM N26). First issue. Tank "
     "capacity (6800 L), dimensions and HDPE/PVC/EPDM materials consistent with the "
     "approved CIP Tank Datasheet, the Equipment List and the P&ID (TK-09-001), and "
     "compatible with the CIP solution at pH 2-12. OBS-01 MINOR: equipment tag "
     "TK-09-001 not shown on the GA. OBS-02 MINOR: nozzle schedule and top opening "
     "disagree with the datasheet (Manhole 21in vs Handhole DN300; added N42/N97). "
     "Fold at IFC Rev 0 (no new revision): add the tag and reconcile the nozzle "
     "schedule. The NCh 2369 anchor loads are tracked to the Module Seismic "
     "Calculation Report (Section 3)."),
]

# ---- Revision History: 9 filas nuevas (num, doc, code, rev, delivery, submittal, verdict, keyobs) ----
# num se resuelve tras insertar (para los nuevos = # asignado en Master Register).
RH_NEW = [
    (None, "Inspection and Test Plan Offsite", "P22-BA-09-000-004", "0", "E57",
     "25007-0057", "1-Approved",
     "Rev 0 (IFC). Code 1 - Approved. Closes the TM N23 cycle: incorporates both "
     "edits with no intermediate revision - row 2.2 'without code stamp' + binding "
     "pressures (1800/1200 psi x 1.1), vessel hydrostatic test raised Witness->Hold "
     "Point. ASME Section X correct for FRP, stamp waived. NOTE-01 cite the 02-Jun "
     "waiver in cell 2.2 (non-gating). Binding-pressure alignment lives on #104. "
     "Issue at IFC Rev 0."),
    (None, "NDE Plan (Super Duplex)", "P22-BA-09-000-005", "C", "E61", "25007-0061",
     "1-Approved",
     "Rev C. Code 1 - Approved. Closes the TM N23 cycle (3rd cycle): governing code "
     "editions specified (OBS-01 Rev B closed) and joint-to-code mapping + UT column "
     "clarification + B31.3 acceptance paragraph incorporated (OBS-02 Rev B closed). "
     "Coverage matches the ET. Issue at IFC Rev 0."),
    (None, "RO Vessel Hydrostatic Test Procedure", "P22-BA-09-000-009", "B", "E61",
     "25007-0061", "3-To be revised",
     "Rev B. Code 3. Does not close the TM N23 driver. OBS-01 CRITICAL: no binding "
     "test pressure in the body and the Protec form still prints 45.5 bar vs the "
     "required 1980 psi (BPV-8-1800) / 1320 psi (BPV-8-1200) - the test the ASME "
     "stamp was waived for. OBS-02 MAJOR: single gauge, no certificate review. OBS-03 "
     "MINOR: Protec 'Revision 0' on one index page. Re-issue as Rev C."),
    (None, "HP and LP Pressure Test Procedure", "P22-BA-09-000-010", "B", "E61",
     "25007-0061", "3-To be revised",
     "Rev B. Code 3. Partially closes the TM N23 observation: adds the B31.3 factor "
     "(1.5x / 1.1x) but OBS-01 MAJOR - no numeric test pressure (135 bar HP / 7.5 "
     "bar LP) and the HP design pressure 90 vs 120 bar unreconciled, both deferred to "
     "the line list. OBS-02 MINOR: 5.6.5 missing, 5.7.2 skips 5.7.2.2. Re-issue Rev "
     "C."),
    (None, "DS PLC & HMI Panel Component", "P22-ET-09-008-001", "C", "E58",
     "25007-0058", "2-AN",
     "Rev C (delivered as P22-ET-09-008-01, 2-digit code vs the register 3-digit; "
     "NOTE-01 align). Code 2 - Approved as Noted. Closes the three TM N21 items: HART "
     "at the instrument level (250-ohm resistor on the 5069-IF8 and 5069-IY4 for "
     "handheld HART; no central HART decoding required, TM N24 over-reach not "
     "reopened), 5069-IY4 RTD module included, title-block typo fixed. NOTE-01: align "
     "the document code to P22-ET-09-008-001 (touches the datasheet). NOTE-02: "
     "confirm the RTD-module quantity vs the LCP Datasheet. Fold at IFC Rev 0."),
    (None, "UHPRO Structural Calculation Report", "P22-CD-09-005-001", "A", "E59",
     "25007-0059", "3-To be revised",
     "Rev A. Code 3. First issue; the seismic calc report tracked as a TM N25 "
     "deliverable. Confirms the Design Criteria Rev B basis (NCh 2369 Of.2003 Zone 3; "
     "lifting API RP 2A-WSD 1.35/2.0). OBS-01 CRITICAL: anchor design omits the main "
     "process equipment (HP pump, turbos, RO filter, pressure vessels 4160 kg). "
     "OBS-02/03/04 MINOR + OBS-05 MAJOR (frame material / revision label). Re-issue "
     "as Rev B; provide the foundation base-bolt interface table."),
    (None, "GA of CIP Flushing Tank", "P22-DWG-09-005-014", "A", "E60", "25007-0060",
     "2-AN",
     "Rev A. Code 2 - Approved as Noted. First issue. Capacity (6800 L), dimensions "
     "and HDPE/PVC/EPDM materials consistent with the approved CIP Tank Datasheet, "
     "Equipment List and P&ID (TK-09-001), fit for CIP at pH 2-12. OBS-01 MINOR: tag "
     "TK-09-001 not shown. OBS-02 MINOR: nozzle schedule / top-opening disagree with "
     "the datasheet. Fold at IFC Rev 0. NCh 2369 anchor loads tracked (Section 3)."),
    (None, "GA of Antiscalant Dosing Tank", "P22-DWG-09-005-015", "B", "E60",
     "25007-0060", "3-To be revised",
     "Rev B. Code 3. Partially closes TM N10 OBS-05: total capacity (335 L) and PE "
     "body material declared (2 of 4). OBS-01 MAJOR: NCh 2369 seismic reaction loads "
     "and anchor-bolt pattern still absent (BW Water defers them), OOCC foundation "
     "without input. OBS-02 MINOR: effective working volume (0.27 m3) not stated. "
     "Re-issue as Rev C."),
    (None, "GA of Antiscalant Dosing Pump Skid", "P22-DWG-09-005-011", "B", "E62",
     "25007-0062", "2-AN",
     "Rev B. Code 2 - Approved as Noted. Materially closes TM N11 NOTE-06: anchor-"
     "bolt layout (M10, 10 bolts, 120 mm embedment), NCh 2369 seismic reaction "
     "forces, and equipment mass (100 kg) delivered. OBS-01 MINOR: per-bolt Bolting "
     "Embedment forces (Fz 0.205 kN) inconsistent with the total block (Fz 0.2048 "
     "kN) over 10 bolts - label per-bolt vs group-total. Endorsed calc report a "
     "separate deliverable (Section 3). Fold at IFC Rev 0."),
]


def style_row(ws, dst_row, ref_row, ncols):
    for c in range(1, ncols + 1):
        s = ws.cell(ref_row, c)
        d = ws.cell(dst_row, c)
        d._style = copy(s._style)


def main():
    shutil.copyfile(SRC, BAK)
    print(f"Backup: {BAK}")

    wb = openpyxl.load_workbook(SRC)  # data_only=False: preserva formulas/estilos
    mr = wb["Master Register"]
    rh = wb["Revision History"]
    sm = wb["Summary"]

    ncols_mr = mr.max_column
    hdr = {mr.cell(1, c).value: c for c in range(1, ncols_mr + 1)}
    C_NUM = hdr["#"]
    C_DOC = hdr["Document"]
    C_CODE = hdr["Code / ET Reference"]
    C_REV = hdr["Rev"]; C_DEL = hdr["Delivery"]; C_TM = hdr["TM"]
    C_VER = hdr["Verdict"]; C_STA = hdr["Status"]; C_ACT = hdr["Action Required"]
    C_ETD = hdr["ET Deadline"]

    # localizar ultima fila de datos y # maximo
    last_data = 1
    max_num = 0
    existing_codes = set()
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code:
            last_data = r
            existing_codes.add(code)
            try:
                max_num = max(max_num, int(mr.cell(r, C_NUM).value))
            except (TypeError, ValueError):
                pass

    # guardarraíl: los NEW no deben existir ya
    for code, *_ in NEW_ITEMS:
        if code in existing_codes:
            raise SystemExit(f"ERROR: el 'nuevo' {code} ya existe en el Master Register - revisar (deberia ser UPDATE).")

    # 1) updates a filas existentes (re-revisiones)
    seen = set()
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code in MR_UPDATES:
            u = MR_UPDATES[code]
            mr.cell(r, C_REV).value = u["rev"]
            mr.cell(r, C_DEL).value = u["delivery"]
            mr.cell(r, C_TM).value = u["tm"]
            mr.cell(r, C_VER).value = u["verdict"]
            mr.cell(r, C_ACT).value = u["action"]
            seen.add(code)
            print(f"  MR update row {r}: {code} -> {u['verdict']} / {u['tm']} (Rev {u['rev']})")
    missing = set(MR_UPDATES) - seen
    if missing:
        raise SystemExit(f"ERROR: codigos N26 no encontrados en Master Register: {missing}")

    # 2) filas nuevas en Master Register
    new_num_by_code = {}
    for i, (code, doc, rev, dele, verdict, action) in enumerate(NEW_ITEMS, start=1):
        dst = last_data + i
        num = max_num + i
        new_num_by_code[code] = num
        style_row(mr, dst, last_data, ncols_mr)
        mr.cell(dst, C_NUM).value = num
        mr.cell(dst, C_DOC).value = doc
        mr.cell(dst, C_CODE).value = code
        mr.cell(dst, C_REV).value = rev
        mr.cell(dst, C_DEL).value = dele
        mr.cell(dst, C_TM).value = TM
        mr.cell(dst, C_VER).value = verdict
        mr.cell(dst, C_STA).value = "Delivered"
        mr.cell(dst, C_ACT).value = action
        mr.cell(dst, C_ETD).value = ET_DEADLINE_ENG
        print(f"  MR NEW row {dst}: #{num} {code} -> {verdict}")

    # numero de registro para los RH de re-revisiones (por codigo)
    reg_num = {}
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code:
            try:
                reg_num[code] = int(mr.cell(r, C_NUM).value)
            except (TypeError, ValueError):
                pass

    # 3) Revision History (9 filas nuevas)
    rcols = rh.max_column
    rhh = {rh.cell(1, c).value: c for c in range(1, rcols + 1)}
    prior = collections.Counter()
    rh_last = 1
    for r in range(2, rh.max_row + 1):
        cd = rh.cell(r, rhh["Code / ET Reference"]).value
        if cd:
            prior[str(cd).strip()] += 1
            rh_last = r
    for i, row in enumerate(RH_NEW, start=1):
        dst = rh_last + i
        _num, doc, code, rev, dele, sub, ver, keyobs = row
        num = reg_num.get(code, "")
        cycle = prior[code] + 1
        style_row(rh, dst, rh_last, rcols)
        rh.cell(dst, rhh["#"]).value = num
        rh.cell(dst, rhh["Document"]).value = doc
        rh.cell(dst, rhh["Code / ET Reference"]).value = code
        rh.cell(dst, rhh["Rev"]).value = rev
        rh.cell(dst, rhh["Delivery"]).value = dele
        rh.cell(dst, rhh["Submittal"]).value = sub
        rh.cell(dst, rhh["TM"]).value = TM
        rh.cell(dst, rhh["TM Date"]).value = TM_DATE
        rh.cell(dst, rhh["Verdict"]).value = ver
        rh.cell(dst, rhh["Cycle"]).value = cycle
        rh.cell(dst, rhh["Key Observations"]).value = keyobs
        print(f"  RH new row {dst}: #{num} {code} cycle {cycle} -> {ver}")

    # 4) Summary: recomputar desde Master Register
    rows = []
    for r in range(2, mr.max_row + 1):
        code = mr.cell(r, C_CODE).value
        if code and str(code).strip():
            rows.append((str(mr.cell(r, C_VER).value or "").strip(),
                         str(mr.cell(r, C_STA).value or "")))
    total = len(rows)
    delivered = sum(1 for v, s in rows if "Delivered" in s)
    vd = collections.Counter(v for v, s in rows if "Delivered" in s)
    print(f"  Summary recompute: total={total} delivered={delivered} verdicts={dict(vd)}")

    sm["B4"].value = total
    sm["B5"].value = delivered
    sm["B9"].value = "26 (N1 through N26 - TM N5 issued in Rev 0 and Rev 1)"
    sm["B10"].value = "62 (E1 through E62)"
    sm["B12"].value = "06-Jul-2026 (E62)"
    sm["B13"].value = "06-Jul-2026"

    def pct(n):
        return f"{round(100 * n / delivered)}%" if delivered else "0%"
    sm["B27"].value = vd.get("1-Approved", 0); sm["C27"].value = pct(vd.get("1-Approved", 0))
    sm["B28"].value = vd.get("2-AN", 0); sm["C28"].value = pct(vd.get("2-AN", 0))
    sm["B29"].value = vd.get("3-To be revised", 0); sm["C29"].value = pct(vd.get("3-To be revised", 0))
    sm["B30"].value = vd.get("4-Rejected", 0); sm["C30"].value = pct(vd.get("4-Rejected", 0))

    # ITEMS BY SECTION: Section 1 (Engineering) sube 2 (los 2 NEW son de ingenieria)
    # Row 18: B=Total, C=Delivered, D=Progress
    try:
        s1_total = int(sm["B18"].value) + 2
        s1_deliv = int(sm["C18"].value) + 2
        sm["B18"].value = s1_total
        sm["C18"].value = s1_deliv
        sm["D18"].value = f"{round(100 * s1_deliv / s1_total)}%"
        print(f"  Section 1 (Engineering): total->{s1_total} delivered->{s1_deliv}")
    except (TypeError, ValueError) as e:
        print(f"  WARN: no pude actualizar ITEMS BY SECTION fila 18 automaticamente ({e}); revisar a mano.")

    wb.save(SRC)
    print(f"Guardado: {SRC}")
    print("NOTA: REVIEW CYCLES distribution NO se modifico (recuento manual; "
          "re-revisiones incrementan ciclo y 2 nuevos = ciclo 1).")


if __name__ == "__main__":
    main()
