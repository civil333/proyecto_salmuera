#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
update_register_n28.py
Actualiza el Master Deliverable Register a TM N28 (E65 + E66, submittals
25007-0065 + 25007-0066, 4 documentos de la familia de Control). Modelado en
update_register_n27.py.

Idempotencia: el backup _pre-N28.xlsx conserva el estado LIMPIO N27 (48/25/10/0).
Al re-correr, RESTAURA SRC desde ese backup antes de aplicar.

Veredicto N28: 2 - Approved as Noted. Tally 4 Code 2.
  UPDATES (3 re-revisiones de filas existentes):
    #67 (row 66) Control Philosophy    P22-BT-09-009-001 (Rev D->E, 3 -> 2-AN)
    #93 (row 79) Alarm & Interlock List P22-LI-09-008-015 (Rev B->C, 3 -> 2-AN)
    #31 (row 33) IO List               P22-LI-09-008-001 (Rev 4->5, 2 -> 2-AN)
  NEW_ITEMS (1, Section 1 Engineering):
    Control and Sequence Chart P22-LI-09-008-017 (Rev A, 2-AN) - primer issue,
    el hijo que llevaba 6-7 ciclos sin emitirse.

Tally esperado tras N28: 48 Code 1 / 28 Code 2 / 8 Code 3 / 0 Code 4;
108 items / 84 delivered; 28 TMs / 66 entregas.
"""
import os
import shutil
import collections
from copy import copy
import openpyxl

EV = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register.xlsx")
BAK = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register_pre-N28.xlsx")

TM = "N28"
TM_DATE = "20-Jul-2026"

MR_UPDATES = {
    # Plant Control Philosophy Rev E -> Code 2 (familia entregada; alinear tags/setpoints)
    "P22-BT-09-009-001": dict(
        rev="E", delivery="E65", tm=TM, verdict="2-AN",
        action=(
            "Rev E delivered (E65). Code 2 - Approved as noted (TM N28). Parent "
            "delivered with the long-outstanding children (Control and Sequence Chart "
            "Rev A, Alarm and Interlock List Rev C). Own control logic correct and "
            "verified: turbocharger bypass VE-09-002 governed by the pressure control "
            "loop (closes the O&M contradiction), clean HP pump permissive, SEC 4.71 and "
            "6x4 / 70-membrane array. Cross-document reconciliation to incorporate at Rev "
            "0: the instrument tables still tag the RO HP and CIP pump temperature "
            "sensors reversed vs the Alarm List, IO List and Instrument List "
            "(winding/bearing swap, a safety-sensor mapping); the motor-temperature, "
            "vibration and discharge low-pressure setpoints must align to the Alarm List; "
            "the winding trip stated at 155C against a Class B citation is inconsistent. "
            "OBS-01 to OBS-04, NOTE-01. Issue at IFC Rev 0 as a coordinated Control-family "
            "reconciliation; no new revision."),
    ),
    # Alarm and Interlock List Rev C -> Code 2 (cierra N17/N20; vibracion AHH propia)
    "P22-LI-09-008-015": dict(
        rev="C", delivery="E65", tm=TM, verdict="2-AN",
        action=(
            "Rev C delivered (E65). Code 2 - Approved as noted (TM N28). Closes the "
            "comments carried since TM N17/N20: permeate-conductivity setpoints corrected "
            "to microsiemens (the two-order-of-magnitude unit error is gone), RO HP pump "
            "winding/bearing correctly tagged and set (winding TE-09-001 140/120, bearing "
            "TE-09-002 95/90, resolving the Rev B swap), turbo vibration symmetric, six "
            "comment-sheet replies implemented. OBS-01 MAJOR: RO HP pump vibration "
            "high-high trip set at 10.0 mm/s on a transmitter ranged 0-8.9, so the trip "
            "the Control Philosophy requires can never fire - bring within range or "
            "re-range the transmitter. OBS-02: incoming-breaker alarm tag XT001 vs IO "
            "List XA001. NOTE-01/02: tag prefixes, missing valve/pump fault alarms, boost "
            "differential SP, housekeeping. Issue at IFC Rev 0; governs the reconciled "
            "setpoints for the Control-family Rev 0; no new revision."),
    ),
    # IO List Rev 5 (IFC) -> Code 2 (cierra N25; falta salida heater)
    "P22-LI-09-008-001": dict(
        rev="5", delivery="E66", tm=TM, verdict="2-AN",
        action=(
            "Rev 5 delivered (E66), IFC. Code 2 - Approved as noted (TM N28). Closes the "
            "TM N25 note: dosing-pump remote/running counted as BOOL, running feedback "
            "sourced from the LCP (not the HMI), four relay-contact coordination signals "
            "to the plant control system present, both motors Pt-100 winding+bearing "
            "(ET Section 5.3 - Electrical Motors), no placeholders. OBS-01: the Alarm and "
            "Interlock List commands a Stop-heater interlock on the CIP tank heater "
            "REL-09-001 but this list carries no output channel for it - add the heater "
            "start/stop output, or confirm in writing that the heater is controlled "
            "outside the module PLC (then Code 1, item transfers to the Alarm List). Tag "
            "alignments (VIT/ORPIT) and Valve List/P&ID reconciliation to Section 3. Issue "
            "at IFC Rev 0; no new revision."),
    ),
}

# Filas nuevas: (Document, Code, Rev, Delivery, Verdict, Action, Status)
NEW_ITEMS = [
    ("Control and Sequence Chart", "P22-LI-09-008-017", "A", "E65",
     "2-AN",
     ("Rev A delivered (E65), first issue. Code 2 - Approved as Noted (TM N28). The "
      "operating sequence chart outstanding for six to seven review cycles is now "
      "issued; structure complete and correct (group control, service start-up, normal "
      "and emergency shutdown, CIP, flushing), permissives match the Control Philosophy "
      "one to one. Cross-document reconciliation to incorporate at Rev 0: OBS-01 Note 5 "
      "states the turbocharger bypass by a TDS value where the Control Philosophy "
      "governs it by pressure; OBS-02 to OBS-06 sequence setpoints differ from the Alarm "
      "List (Stage-2 over-pressure abort 90 vs 93 bar, flushing flow 40 vs 48/36, "
      "suction permissive 1.5 vs 1.75) and the Control Philosophy (VFD ramp 1 vs "
      "0.1-0.3 Hz/s), plus the CIP-return valve stage description and the brine-flow "
      "formula; NOTE-01 housekeeping. Issue at IFC Rev 0 as a coordinated Control-family "
      "reconciliation; no new revision."),
     "Delivered"),
]

# Section 1 (ENGINEERING) recibe el item nuevo (Sequence Chart es I&C/control).
SECTION_ROW = 18  # '1. ENGINEERING'

RH_NEW = [
    (None, "Plant Control Philosophy", "P22-BT-09-009-001", "E", "E65",
     "25007-0065", "2-AN",
     "Rev E. Code 2 - Approved as noted. Parent delivered with its children (Sequence "
     "Chart Rev A, Alarm List Rev C). Logic correct (bypass by pressure loop, "
     "permissive, SEC, 6x4/70). Reconcile at Rev 0: instrument tables tag winding/"
     "bearing reversed vs Instrument List/Alarm List/IO List (OBS-01); motor-temp/"
     "vibration/discharge-low-P setpoints to the Alarm List (OBS-02/03/04); winding "
     "trip 155C vs Class B. Coordinated Control-family reconciliation."),
    (None, "Alarm and Interlock List", "P22-LI-09-008-015", "C", "E65",
     "25007-0065", "2-AN",
     "Rev C. Code 2 - Approved as noted. Closes N17/N20: permeate units to uS/cm, "
     "winding TE-09-001 140/120 + bearing TE-09-002 95/90 (Rev B swap resolved), turbo "
     "vibration symmetric, 6 CCS replies real. OBS-01 MAJOR: RO HP pump vibration AHH "
     "10.0 above the 0-8.9 transmitter range -> trip unreachable. OBS-02 MCCB tag vs IO "
     "List. Issue at IFC Rev 0; governs the reconciled setpoints."),
    (None, "Control and Sequence Chart", "P22-LI-09-008-017", "A", "E65",
     "25007-0065", "2-AN",
     "Rev A, first issue. Code 2 - Approved as noted. The 6-7 cycle overdue sequence "
     "chart, structure and permissives correct (match the Control Philosophy). "
     "Reconcile at Rev 0: Note 5 bypass by TDS vs pressure (OBS-01); setpoints vs Alarm "
     "List + VFD ramp vs Control Philosophy + CIP-return valve stage + brine formula "
     "(OBS-02 to OBS-06). Coordinated Control-family reconciliation."),
    (None, "IO List", "P22-LI-09-008-001", "5", "E66",
     "25007-0066", "2-AN",
     "Rev 5, IFC. Code 2 - Approved as noted. Closes N25: dosing BOOL count, RUNNING "
     "from LCP, 4 relay-contact DCS signals, Pt-100 winding+bearing both motors, no "
     "placeholders. OBS-01: no output channel for the Alarm List Stop-heater interlock "
     "(REL-09-001) - add it or confirm the heater is outside the module PLC (then Code "
     "1). Tag/Valve List alignments to Section 3. Issue at IFC Rev 0."),
]


def style_row(ws, dst_row, ref_row, ncols):
    for c in range(1, ncols + 1):
        s = ws.cell(ref_row, c)
        d = ws.cell(dst_row, c)
        d._style = copy(s._style)


def main():
    if os.path.exists(BAK):
        shutil.copyfile(BAK, SRC)
        print(f"Restaurado SRC desde baseline limpio: {BAK}")
    else:
        shutil.copyfile(SRC, BAK)
        print(f"Baseline creado: {BAK}")

    wb = openpyxl.load_workbook(SRC)
    mr = wb["Master Register"]
    rh = wb["Revision History"]
    sm = wb["Summary"]

    ncols_mr = mr.max_column
    hdr = {mr.cell(1, c).value: c for c in range(1, ncols_mr + 1)}
    C_NUM = hdr["#"]; C_DOC = hdr["Document"]; C_CODE = hdr["Code / ET Reference"]
    C_REV = hdr["Rev"]; C_DEL = hdr["Delivery"]; C_TM = hdr["TM"]
    C_VER = hdr["Verdict"]; C_STA = hdr["Status"]; C_ACT = hdr["Action Required"]

    # 1) updates (3 re-revisiones)
    seen = set()
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code in MR_UPDATES:
            u = MR_UPDATES[code]
            mr.cell(r, C_REV).value = u["rev"]
            mr.cell(r, C_DEL).value = u["delivery"]
            mr.cell(r, C_TM).value = u["tm"]
            mr.cell(r, C_VER).value = u["verdict"]
            mr.cell(r, C_STA).value = "Delivered"
            mr.cell(r, C_ACT).value = u["action"]
            seen.add(code)
            print(f"  MR update row {r}: {code} -> {u['verdict']} / {u['tm']} (Rev {u['rev']})")
    missing = set(MR_UPDATES) - seen
    if missing:
        raise SystemExit(f"ERROR: codigos N28 no encontrados en Master Register: {missing}")

    # 2) filas NUEVAS
    existing = set(); max_num = 0; ref_mr_row = 2
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code:
            existing.add(code); ref_mr_row = r
            try:
                max_num = max(max_num, int(mr.cell(r, C_NUM).value))
            except (TypeError, ValueError):
                pass
    for i, (doc, code, rev, dele, ver, act, sta) in enumerate(NEW_ITEMS, start=1):
        if code in existing:
            raise SystemExit(f"ERROR: NEW_ITEM {code} ya existe en el register.")
        dst = ref_mr_row + i
        style_row(mr, dst, ref_mr_row, ncols_mr)
        mr.cell(dst, C_NUM).value = max_num + i
        mr.cell(dst, C_DOC).value = doc
        mr.cell(dst, C_CODE).value = code
        mr.cell(dst, C_REV).value = rev
        mr.cell(dst, C_DEL).value = dele
        mr.cell(dst, C_TM).value = TM
        mr.cell(dst, C_VER).value = ver
        mr.cell(dst, C_STA).value = sta
        mr.cell(dst, C_ACT).value = act
        print(f"  MR new row {dst}: #{max_num + i} {code} -> {ver}")

    reg_num = {}
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code:
            try:
                reg_num[code] = int(mr.cell(r, C_NUM).value)
            except (TypeError, ValueError):
                pass

    # 3) Revision History
    rcols = rh.max_column
    rhh = {rh.cell(1, c).value: c for c in range(1, rcols + 1)}
    prior = collections.Counter(); rh_last = 1
    for r in range(2, rh.max_row + 1):
        cd = rh.cell(r, rhh["Code / ET Reference"]).value
        if cd:
            prior[str(cd).strip()] += 1; rh_last = r
    for i, row in enumerate(RH_NEW, start=1):
        dst = rh_last + i
        _num, doc, code, rev, dele, sub, ver, keyobs = row
        num = reg_num.get(code, "")
        cycle = prior[code] + 1; prior[code] += 1
        style_row(rh, dst, rh_last, rcols)
        rh.cell(dst, rhh["#"]).value = num
        rh.cell(dst, rhh["Document"]).value = doc
        rh.cell(dst, rhh["Code / ET Reference"]).value = code
        rh.cell(dst, rhh["Rev"]).value = rev
        rh.cell(dst, rhh["Delivery"]).value = dele
        rh.cell(dst, rhh["Submittal"]).value = sub
        rh.cell(dst, rhh["TM"]).value = TM
        rh.cell(dst, rhh["TM Date"]).value = TM_DATE
        rh.cell(dst, rhh["Verdict"]).value = ver
        rh.cell(dst, rhh["Cycle"]).value = cycle
        rh.cell(dst, rhh["Key Observations"]).value = keyobs
        print(f"  RH new row {dst}: #{num} {code} cycle {cycle} -> {ver}")

    # 4) Summary: recomputar total/delivered/partial/not-delivered + verdicts desde MR
    stat = collections.Counter(); vd = collections.Counter(); total = 0
    for r in range(2, mr.max_row + 1):
        code = mr.cell(r, C_CODE).value
        if code and str(code).strip():
            total += 1
            su = str(mr.cell(r, C_STA).value or "").strip().upper()
            stat[su] += 1
            if su == "DELIVERED":
                vd[str(mr.cell(r, C_VER).value or "").strip()] += 1
    delivered = stat.get("DELIVERED", 0)
    not_delivered = stat.get("NOT DELIVERED", 0)
    partial = stat.get("PARTIAL", 0)
    print(f"  Summary recompute: total={total} delivered={delivered} "
          f"not_delivered={not_delivered} partial={partial} verdicts={dict(vd)}")

    sm["B4"].value = total
    sm["B5"].value = delivered
    sm["B7"].value = partial
    sm["B8"].value = not_delivered
    sm["B9"].value = "28 (N1 through N28 - TM N5 issued in Rev 0 and Rev 1)"
    sm["B10"].value = "66 (E1 through E66)"
    sm["B12"].value = "17-Jul-2026 (E66)"
    sm["B13"].value = "20-Jul-2026"

    def pct(n):
        return f"{round(100 * n / delivered)}%" if delivered else "0%"
    sm["B27"].value = vd.get("1-Approved", 0); sm["C27"].value = pct(vd.get("1-Approved", 0))
    sm["B28"].value = vd.get("2-AN", 0); sm["C28"].value = pct(vd.get("2-AN", 0))
    sm["B29"].value = vd.get("3-To be revised", 0); sm["C29"].value = pct(vd.get("3-To be revised", 0))
    sm["B30"].value = vd.get("4-Rejected", 0); sm["C30"].value = pct(vd.get("4-Rejected", 0))

    # 5) ITEMS BY SECTION: +1 total y +1 delivered a Section 1 (el Sequence Chart)
    sm.cell(SECTION_ROW, 2).value = (sm.cell(SECTION_ROW, 2).value or 0) + 1
    sm.cell(SECTION_ROW, 3).value = (sm.cell(SECTION_ROW, 3).value or 0) + 1
    print(f"  ITEMS BY SECTION: Section 1 -> "
          f"{sm.cell(SECTION_ROW,2).value}/{sm.cell(SECTION_ROW,3).value}")

    _guardar_robusto(wb, SRC)


def _guardar_robusto(wb, dst):
    import zipfile
    scratch_dir = ("/private/tmp/claude-501/-Volumes-home-Documentos-NAS-DESAROLLO-"
                   "PROYECTOS-CLAUDE-MODULO-DE-SALMUERA-TALTAL/"
                   "163017ca-4212-4407-918e-c9861077b701/scratchpad")
    os.makedirs(scratch_dir, exist_ok=True)
    tmp = os.path.join(scratch_dir, "register_n28_build.xlsx")
    wb.save(tmp)
    if zipfile.ZipFile(tmp).testzip() is not None:
        raise SystemExit("ERROR: el build en scratchpad salio corrupto.")
    openpyxl.load_workbook(tmp).close()
    print(f"Build verificado en scratchpad: {tmp}")
    for intento in range(1, 6):
        shutil.copyfile(tmp, dst)
        try:
            if zipfile.ZipFile(dst).testzip() is None:
                openpyxl.load_workbook(dst).close()
                print(f"Guardado y verificado en destino (intento {intento}): {dst}")
                return
        except Exception as e:
            print(f"  intento {intento}: destino aun corrupto ({e}); recopiando...")
    raise SystemExit(f"ERROR: no se pudo dejar copia integra. Build en {tmp}.")


if __name__ == "__main__":
    main()
