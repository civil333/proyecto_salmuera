#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
update_register_n40.py — aplica el Transmittal N40 al Master Deliverable Register.

TM N40 (P22-TM-09-000-040-0), emitido el 18 de septiembre de 2026, sobre los submittals
25007-0093 (ENTREGA 93, recibida el lunes 14-Sep), 25007-0094 (ENTREGA 94, jueves 17-Sep)
y 25007-0095 (ENTREGA 95, viernes 18-Sep). CATORCE documentos: trece re-emisiones y una
primera emision, el Factory Acceptance Test Procedure del modulo.

IDEMPOTENCIA. Si existe el backup _pre-N40.xlsx, restaura SRC desde ese backup antes de
aplicar. Re-correr el script es seguro.

VEREDICTO GLOBAL: 3 — To be revised. 12 Codigo 1, 1 Codigo 2, 1 Codigo 3.

EFECTO SOBRE EL TALLY, declarado ANTES de correr. Se parte del cierre del N39: 68 / 19 / 2 / 0
con 113 items y 89 entregados.

  Entra a Delivered (uno): fila #85, que era el placeholder NOT DELIVERED del procedimiento
    FAT (codigo 'ET Seccion 8.1'), pasa a P22-BA-09-000-017 Rev A, Codigo 3. +1 Codigo 3.
  Suben de Codigo 2 a Codigo 1 (tres): Instrument Location Layout (Rev 0, cierra N37), DS CIP
    Pump (Rev 0, cierra la compatibilidad de motor del N2), DS Antiscalant Dosing Tank (Rev 0,
    cierra las notas del N4 via P&ID Rev 0 y CT-001). +3 Codigo 1, -3 Codigo 2.
  Sin cambio de codigo: Schematic Rev C (2 -> 2, tercer ciclo) y los nueve Codigo 1 restantes.

  Tally esperado tras el N40: 71 / 16 / 3 / 0. Items 113, entregados 90.

CORRECCION DE REGISTRO, independiente del veredicto: tres datasheets de estanques llevaban el
correlativo corrido en la columna de codigo (fila 13 decia -009-010 y es -009-009; fila 14
decia -009-011 y es -009-010; fila 15 decia -009-014 y es -009-011). Manda el documento
(formulario de submittal, portada del PDF, TM N3/N4/N34/N39). Se corrige la celda de codigo.
La fila 39 arrastraba "A/C thermal calc Rev C still pending", ya cerrado: se reescribe.

🔴 CORRER DESPUES DE ENVIAR el transmittal, no antes.
"""
import os
import shutil
import collections
from copy import copy
import openpyxl

EV = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register.xlsx")
BAK = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register_pre-N40.xlsx")
TM = "N40"
TM_DATE = "18-Sep-2026"
TALLY_ESPERADO = {"1-Approved": 71, "2-AN": 16, "3-To be revised": 3, "4-Rejected": 0}

# Correcciones de codigo, por numero de fila (#): codigo_viejo -> codigo_real
CODE_FIXES = {13: "P22-ET-09-009-009", 14: "P22-ET-09-009-010", 15: "P22-ET-09-009-011"}

# Fila #85: el placeholder del FAT pasa a ser el documento recibido.
ROW85 = dict(
    document="Factory Acceptance Test Procedure (module)",
    code="P22-BA-09-000-017", rev="A", delivery="E95", tm=TM, verdict="3-To be revised",
    status="Delivered",
    action=("Rev A submitted for approval (E95, 18-Sep-2026, submittal 25007-0095), first "
            "issue of the module FAT procedure required by ET Section 8.1 and by row 7.1 of "
            "the ITP P22-BA-09-000-004 Rev 0, an ADASA Hold Point. Code 3 - To be revised "
            "(TM N40). Ten pages that restate rows 7.2 to 7.9 of the ITP as headings 5.1 to "
            "5.8: no test step, no acceptance value, none of the nine record forms Section 8 "
            "names; no loop check list, no simulation method and no fault scenarios for the "
            "control system test; rows 7.3, 7.4 and 7.5 of the ITP refer the key points, the "
            "sequence and the simulation detail to this procedure and the procedure refers "
            "them back to approved drawings. Hydrostatic prerequisite of ET 8.1 not carried. "
            "Rev B required by Friday 25-Sep-2026 (OBS-01 to OBS-10, NOTE-01, NOTE-02). The FAT "
            "cannot formally open until the procedure is approved."),
)

MR_UPDATES = {
    "P22-CD-09-008-002": dict(
        rev="Rev C", delivery="E93", tm=TM, verdict="2-AN",
        action=("Rev C submitted for approval (E93, 14-Sep-2026, submittal 25007-0093). Code 2 "
                "- Approved as noted (TM N40), third cycle (N20, N37, N40). Closed: CIP heater "
                "running, item 109 of the I/O List Rev 6, now on input 7 of module -A3, sheet "
                "37, tag XB002. Open: row 11 of the bill of material (sheet 28) still reads "
                "2711P-T10C21D8S; the comment sheet proposes to revise the approved PLC and HMI "
                "Datasheet P22-ET-09-008-001 Rev 0 to that model. ADASA reaffirms the "
                "2711P-T10C22D9P declared binding at TM N30 (two Ethernet RJ45 ports, 1 GB) and "
                "does not accept a revision of the approved datasheet. To be replaced at Rev 0 "
                "(OBS-01); comment sheet to cite documents by their own code (NOTE-01)."),
    ),
    "P22-DWG-09-008-001": dict(
        rev="Rev 0", delivery="E93", tm=TM, verdict="1-Approved",
        action=("Rev 0 issued for construction (E93, 14-Sep-2026, submittal 25007-0093). Code 1 "
                "- Approved (TM N40). Both observations of TM N37 close: revision block "
                "describes each issue, Rev C row reads APR.17.26 again, comment sheet bound as "
                "sheet 4; notes box declares P&ID Rev 0, Instrument List Rev F and Equipment "
                "Layout Rev D by code. Housekeeping recorded, does not degrade: drawing status "
                "cell reads ISSUED FOR APPROVAL against a Rev 0 row that reads ISSUED FOR "
                "CONTRUCTION (sic). To be corrected when the file is next touched."),
    ),
    "P22-ET-09-009-009": dict(
        rev="Rev 0", delivery="E93", tm=TM, verdict="1-Approved",
        action=("Rev 0 issued for construction (E93, 14-Sep-2026, submittal 25007-0093). Code 1 "
                "- Approved (TM N40). Re-issued aligned to the GA of CIP Flushing Tank "
                "P22-DWG-09-005-014: nozzle schedule carries the 533 mm manhole, N42 spare DN25 "
                "and N97 sensor DN40; the comment sheet declares the GA governs. Closes the "
                "cross-document condition of TM N34 on the datasheet side."),
    ),
    "P22-LI-09-009-001": dict(
        rev="Rev 0", delivery="E94", tm=TM, verdict="1-Approved",
        action=("Rev 0 issued for construction (E94, 17-Sep-2026, submittal 25007-0094). Code 1 "
                "- Approved (TM N40). Content identical to Rev C approved at TM N22."),
    ),
    "P22-LI-09-009-002": dict(
        rev="Rev 0", delivery="E94", tm=TM, verdict="1-Approved",
        action=("Rev 0 issued for construction (E94, 17-Sep-2026, submittal 25007-0094). Code 1 "
                "- Approved (TM N40). Content identical to Rev A approved at TM N3."),
    ),
    "P22-LI-09-009-003": dict(
        rev="Rev 2", delivery="E94", tm=TM, verdict="1-Approved",
        action=("Rev 2 issued for construction (E94, 17-Sep-2026, submittal 25007-0094). Code 1 "
                "- Approved (TM N40). Condition of TM N38 met: CP-SSD-DN80-09-049 (first-stage "
                "CIP reject, 54 m3/h) is DN80 and CP-SSD-DN65-09-048 (second-stage, 36 m3/h) is "
                "DN65; revision history states the correction."),
    ),
    "P22-ET-09-009-001": dict(
        rev="Rev 0", delivery="E94", tm=TM, verdict="1-Approved",
        action=("Rev 0 issued for construction (E94, 17-Sep-2026, submittal 25007-0094). Code 1 "
                "- Approved (TM N40). Stage 1 LG SW 400 SR in six BPV-8-1200-SP-7, stage 2 LG SW "
                "400R G2 UHP in four BPV-8-1800-SP-7, seven elements per vessel, as at Rev B."),
    ),
    "P22-ET-09-009-003": dict(
        rev="Rev 0", delivery="E94", tm=TM, verdict="1-Approved",
        action=("Rev 0 issued for construction (E94, 17-Sep-2026, submittal 25007-0094). Code 1 "
                "- Approved (TM N40). Closes the motor compatibility point of TM N2: one motor, "
                "IEC 160MB, 11 kW, two poles, 2940-2950 rpm, external VFD, 10.11 kW absorbed at "
                "duty, consistent with the pump curve."),
    ),
    "P22-ET-09-009-004": dict(
        rev="Rev 0", delivery="E94", tm=TM, verdict="1-Approved",
        action=("Rev 0 issued for construction (E94, 17-Sep-2026, submittal 25007-0094). Code 1 "
                "- Approved (TM N40). ProMinent GMXa 1602, 0.02 L/h duty, 2.3 L/h max, 16 barg, "
                "as at Rev B."),
    ),
    "P22-ET-09-009-005": dict(
        rev="Rev 0", delivery="E94", tm=TM, verdict="1-Approved",
        action=("Rev 0 issued for construction (E94, 17-Sep-2026, submittal 25007-0094). Code 1 "
                "- Approved (TM N40). Adds the EPDM gasket row and restates the filtration rate "
                "per square metre; same cartridge and housing as Rev E."),
    ),
    "P22-ET-09-009-006": dict(
        rev="Rev 0", delivery="E94", tm=TM, verdict="1-Approved",
        action=("Rev 0 issued for construction (E94, 17-Sep-2026, submittal 25007-0094). Code 1 "
                "- Approved (TM N40). Closes NOTE-01 of TM N24: housing model 31SBFX4-040A-E, E "
                "suffix = EPDM cover gasket per the ordering guide; gasket row reads EPDM."),
    ),
    "P22-ET-09-009-010": dict(
        rev="Rev 0", delivery="E94", tm=TM, verdict="1-Approved",
        action=("Rev 0 issued for construction (E94, 17-Sep-2026, submittal 25007-0094). Code 1 "
                "- Approved (TM N40). Closes the notes of TM N4: the P&ID Rev 0 reads 0.34 m3 "
                "total and 0.27 effective for TK-09-002, as the datasheet does; the 0.5 ppm "
                "dosing rate was accepted in principle under CT-001 Rev 1, whose residual "
                "points remain in that query."),
    ),
    "P22-ET-09-009-011": dict(
        rev="Rev 0", delivery="E94", tm=TM, verdict="1-Approved",
        action=("Rev 0 issued for construction (E94, 17-Sep-2026, submittal 25007-0094). Code 1 "
                "- Approved (TM N40). Quantic Logic flange immersion heater, 20 kW, 380 V, SS316 "
                "wetted parts, as at Rev B."),
    ),
}

RH_NEW = [
    ("Factory Acceptance Test Procedure (module)", "P22-BA-09-000-017", "A", "E95", "25007-0095",
     "3-To be revised",
     "First issue. Not the detailed procedure of ET 8.1, BAE 31 a) and ITP 7.1: headings without "
     "steps, values or record forms; no loop checks, simulation method or fault scenarios; "
     "circular with ITP rows 7.3 to 7.5. Rev B by 25-Sep-2026."),
    ("PLC/LCP Schematic Diagram", "P22-CD-09-008-002", "Rev C", "E93", "25007-0093", "2-AN",
     "Third cycle. CIP heater running input closed (sheet 37). Terminal 2711P-T10C21D8S still "
     "in the BOM; comment sheet proposes to revise the approved datasheet instead. ADASA "
     "reaffirms the -D9P."),
    ("Instrument Location Layout", "P22-DWG-09-008-001", "Rev 0", "E93", "25007-0093", "1-Approved",
     "Both N37 observations closed. Drawing status cell still reads ISSUED FOR APPROVAL; "
     "housekeeping."),
    ("DS CIP Tank", "P22-ET-09-009-009", "Rev 0", "E93", "25007-0093", "1-Approved",
     "Aligned to the GA of CIP Flushing Tank; the GA governs."),
    ("Utility Consumption List", "P22-LI-09-009-001", "Rev 0", "E94", "25007-0094", "1-Approved",
     "Identical to Rev C."),
    ("Chemical Consumption List", "P22-LI-09-009-002", "Rev 0", "E94", "25007-0094", "1-Approved",
     "Identical to Rev A."),
    ("Line List", "P22-LI-09-009-003", "Rev 2", "E94", "25007-0094", "1-Approved",
     "Sizes of 09-048 and 09-049 corrected as required at N38."),
    ("DS UHPRO System", "P22-ET-09-009-001", "Rev 0", "E94", "25007-0094", "1-Approved",
     "Vessel and membrane assignment per stage unchanged from Rev B."),
    ("DS CIP Pump", "P22-ET-09-009-003", "Rev 0", "E94", "25007-0094", "1-Approved",
     "Motor compatibility of N2 closed: 11 kW IEC 160MB with external VFD, 10.11 kW at duty."),
    ("DS Antiscalant Pump", "P22-ET-09-009-004", "Rev 0", "E94", "25007-0094", "1-Approved",
     "Unchanged from Rev B."),
    ("DS RO Cartridge Filter", "P22-ET-09-009-005", "Rev 0", "E94", "25007-0094", "1-Approved",
     "EPDM gasket row added; filtration rate restated per square metre."),
    ("DS CIP Cartridge Filter", "P22-ET-09-009-006", "Rev 0", "E94", "25007-0094", "1-Approved",
     "NOTE-01 of N24 closed: model suffix E = EPDM cover gasket."),
    ("DS Antiscalant Dosing Tank", "P22-ET-09-009-010", "Rev 0", "E94", "25007-0094", "1-Approved",
     "N4 notes closed via P&ID Rev 0 (0.34 / 0.27 m3) and CT-001 Rev 1."),
    ("DS CIP Tank Heater", "P22-ET-09-009-011", "Rev 0", "E94", "25007-0094", "1-Approved",
     "Unchanged from Rev B."),
]


def style_row(ws, dst_row, ref_row, ncols):
    for c in range(1, ncols + 1):
        ws.cell(dst_row, c)._style = copy(ws.cell(ref_row, c)._style)


def main():
    if os.path.exists(BAK):
        shutil.copyfile(BAK, SRC)
        print(f"Restaurado SRC desde baseline limpio: {BAK}")
    else:
        shutil.copyfile(SRC, BAK)
        print(f"Baseline creado: {BAK}")
    wb = openpyxl.load_workbook(SRC)
    mr = wb["Master Register"]; rh = wb["Revision History"]; sm = wb["Summary"]
    ncols_mr = mr.max_column
    hdr = {mr.cell(1, c).value: c for c in range(1, ncols_mr + 1)}
    C_NUM = hdr["#"]; C_DOC = hdr["Document"]; C_CODE = hdr["Code / ET Reference"]
    C_REV = hdr["Rev"]; C_DEL = hdr["Delivery"]; C_TM = hdr["TM"]
    C_VER = hdr["Verdict"]; C_STA = hdr["Status"]; C_ACT = hdr["Action Required"]

    # 1) correcciones de codigo y fila 85
    for r in range(2, mr.max_row + 1):
        try:
            num = int(mr.cell(r, C_NUM).value)
        except (TypeError, ValueError):
            continue
        if num in CODE_FIXES:
            antes = mr.cell(r, C_CODE).value
            mr.cell(r, C_CODE).value = CODE_FIXES[num]
            print(f"  MR code fix row {r} (#{num}): {antes} -> {CODE_FIXES[num]}")
        if num == 85:
            assert str(mr.cell(r, C_STA).value).strip().upper() == "NOT DELIVERED", mr.cell(r, C_STA).value
            mr.cell(r, C_DOC).value = ROW85["document"]
            mr.cell(r, C_CODE).value = ROW85["code"]
            mr.cell(r, C_REV).value = ROW85["rev"]
            mr.cell(r, C_DEL).value = ROW85["delivery"]
            mr.cell(r, C_TM).value = ROW85["tm"]
            mr.cell(r, C_VER).value = ROW85["verdict"]
            mr.cell(r, C_STA).value = ROW85["status"]
            mr.cell(r, C_ACT).value = ROW85["action"]
            print(f"  MR row {r} (#85): placeholder FAT -> {ROW85['code']} Rev A, Code 3, Delivered")

    # 2) updates por codigo
    seen = set()
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code in MR_UPDATES:
            u = MR_UPDATES[code]
            mr.cell(r, C_REV).value = u["rev"]; mr.cell(r, C_DEL).value = u["delivery"]
            mr.cell(r, C_TM).value = u["tm"]; mr.cell(r, C_VER).value = u["verdict"]
            mr.cell(r, C_STA).value = "Delivered"; mr.cell(r, C_ACT).value = u["action"]
            seen.add(code)
            print(f"  MR update row {r}: {code} -> {u['verdict']} / {u['tm']} ({u['rev']})")
    missing = set(MR_UPDATES) - seen
    if missing:
        raise SystemExit(f"ERROR: codigos N40 no encontrados en Master Register: {missing}")

    # 3) Revision History
    reg_num = {}
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code:
            try:
                reg_num[code] = int(mr.cell(r, C_NUM).value)
            except (TypeError, ValueError):
                pass
    rcols = rh.max_column
    rhh = {rh.cell(1, c).value: c for c in range(1, rcols + 1)}
    prior = collections.Counter(); rh_last = 1
    for r in range(2, rh.max_row + 1):
        cd = rh.cell(r, rhh["Code / ET Reference"]).value
        if cd:
            prior[str(cd).strip()] += 1; rh_last = r
    for i, row in enumerate(RH_NEW, start=1):
        dst = rh_last + i
        doc, code, rev, dele, sub, ver, keyobs = row
        num = reg_num.get(code, "")
        cycle = prior[code] + 1; prior[code] += 1
        style_row(rh, dst, rh_last, rcols)
        for k, v in [("#", num), ("Document", doc), ("Code / ET Reference", code), ("Rev", rev),
                     ("Delivery", dele), ("Submittal", sub), ("TM", TM), ("TM Date", TM_DATE),
                     ("Verdict", ver), ("Cycle", cycle), ("Key Observations", keyobs)]:
            rh.cell(dst, rhh[k]).value = v
        print(f"  RH new row {dst}: #{num} {code} cycle {cycle} -> {ver}")

    # 4) Summary recomputado
    stat = collections.Counter(); vd = collections.Counter(); total = 0
    for r in range(2, mr.max_row + 1):
        if mr.cell(r, C_CODE).value:
            total += 1
            su = str(mr.cell(r, C_STA).value or "").strip().upper()
            stat[su] += 1
            if su == "DELIVERED":
                vd[str(mr.cell(r, C_VER).value or "").strip()] += 1
    delivered = stat.get("DELIVERED", 0); not_delivered = stat.get("NOT DELIVERED", 0)
    partial = stat.get("PARTIAL", 0)
    print(f"  Summary recompute: total={total} delivered={delivered} not_delivered={not_delivered} "
          f"partial={partial} verdicts={dict(vd)}")
    real = {k: vd.get(k, 0) for k in TALLY_ESPERADO}
    if real != TALLY_ESPERADO:
        print(f"  🔴 TALLY DISTINTO DEL ESPERADO: esperado {TALLY_ESPERADO}, real {real}")
    else:
        print(f"  Tally coincide con lo declarado: {real}")
    sm["B4"].value = total; sm["B5"].value = delivered
    sm["B7"].value = partial; sm["B8"].value = not_delivered
    sm["B9"].value = "40 (N1 through N40 - TM N5 issued in Rev 0 and Rev 1)"
    sm["B10"].value = "95 (E1 through E95; the submittal series has no missing numbers)"
    sm["B12"].value = "18-Sep-2026 (E95)"
    sm["B13"].value = "18-Sep-2026"
    def pct(n):
        return f"{round(100 * n / delivered)}%" if delivered else "0%"
    for cell, key in [("27", "1-Approved"), ("28", "2-AN"), ("29", "3-To be revised"),
                      ("30", "4-Rejected"), ("31", "No code issued")]:
        sm["B" + cell].value = vd.get(key, 0); sm["C" + cell].value = pct(vd.get(key, 0))
    # ITEMS BY SECTION: la fila 85 pertenece a la seccion 3 (FABRICATION & FAT): +1 entregado
    for r in range(1, 40):
        v = sm.cell(r, 1).value
        if isinstance(v, str) and v.startswith("3. FABRICATION"):
            antes = sm.cell(r, 3).value
            sm.cell(r, 3).value = int(antes) + 1
            print(f"  ITEMS BY SECTION seccion 3 delivered: {antes} -> {antes + 1}")
    _guardar_robusto(wb, SRC)


def _guardar_robusto(wb, dst):
    import zipfile, tempfile
    scratch_dir = os.environ.get("CLAUDE_SCRATCH") or tempfile.mkdtemp(prefix="register_n40_")
    os.makedirs(scratch_dir, exist_ok=True)
    tmp = os.path.join(scratch_dir, "register_n40_build.xlsx")
    wb.save(tmp)
    if zipfile.ZipFile(tmp).testzip() is not None:
        raise SystemExit("ERROR: el build en scratchpad salio corrupto.")
    openpyxl.load_workbook(tmp).close()
    print(f"Build verificado en scratchpad: {tmp}")
    for intento in range(1, 6):
        shutil.copyfile(tmp, dst)
        try:
            if zipfile.ZipFile(dst).testzip() is None:
                openpyxl.load_workbook(dst).close()
                print(f"Guardado y verificado en destino (intento {intento}): {dst}")
                return
        except Exception as e:
            print(f"  intento {intento}: destino aun corrupto ({e}); recopiando...")
    raise SystemExit(f"ERROR: no se pudo dejar copia integra. Build en {tmp}.")


if __name__ == "__main__":
    main()
