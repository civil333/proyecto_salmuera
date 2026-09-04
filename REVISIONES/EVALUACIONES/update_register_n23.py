#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
update_register_n23.py
Actualiza el Master Deliverable Register a TM N23 (submittals 25007-0051 E51 +
25007-0052 E52). Update EXPLICITO (no el fragil update_register.py, cuyo parser
de veredicto es case-sensitive con guion y no agrega filas).

Hace:
  1. Respalda el original (-002-0) a *_pre-N23.xlsx.
  2. Master Register: actualiza 3 filas existentes (#92 ITP, #84 NDE, #35
     Instrument Layout) y agrega 4 nuevas (#104 RO Vessel Hydro, #105 HP/LP,
     #106 Painting, #107 Structural Criteria), copiando el estilo de fila.
  3. Revision History: agrega 7 filas (1 por doc del TM N23), cycle derivado del
     conteo previo del codigo.
  4. Summary: recomputa Total/Delivered y la distribucion de veredictos DESDE la
     hoja Master Register; actualiza Transmittals (23), Deliveries (52), Latest
     Delivery (E52, 18-Jun), Status Date, y ITEMS BY SECTION (Sec.1 +1 eng,
     Sec.3 +3 fabricacion).
  NO toca REVIEW CYCLES distribution (basada en ejemplos; recuento manual).
"""
import os
import shutil
import collections
from copy import copy
import openpyxl

EV = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register.xlsx")
BAK = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register_pre-N23.xlsx")

TM = "N23"
TM_DATE = "18-Jun-2026"

# ---- Master Register: updates a filas existentes (por codigo) ----
MR_UPDATES = {
    "P22-BA-09-000-004": dict(
        rev="C", delivery="E51", tm=TM, verdict="2-AN",
        action=(
            "Rev C delivered (E51). Code 2 - Approved as Noted (TM N23): "
            "materially closes the oldest fabrication carry-forward - row 2.2 "
            "now carries the 1,800 psi x 1.1 vessel test, the ADASA witness and "
            "the dossier Hold Points (rows 7.6/8.3); ASME Section X is correct "
            "for the FRP vessels. Two edits at Rev 0: OBS-01 declare the "
            "no-code-stamp basis per the 02-Jun-2026 waiver; OBS-02 raise the "
            "vessel test from Witness (W) to Hold Point (H), consistent with "
            "row 5.2. Full closure also depends on the RO Vessel Hydrostatic "
            "Test Procedure (#104) being corrected."),
    ),
    "P22-BA-09-000-005": dict(
        rev="B", delivery="E51", tm=TM, verdict="3-To be revised",
        action=(
            "Rev B delivered (E51). Code 3 - To be revised (TM N23). Coverage "
            "is sound and matches the ET (100% VT, 100% PT root+final, 10% RT, "
            "10% PMI). OBS-01 MAJOR: governing code editions still left as "
            "placeholders - the Rev A comment was not closed in Rev B. OBS-02: "
            "AWS D1.1 (structural) and DVS 2202-1 (thermoplastic) mixed with "
            "the B31.3 Super Duplex HP circuit without mapping each code to its "
            "joints; the UT column is a baseline thickness measurement. Rev C "
            "required. The RO pressure-vessel scope gap of the prior NDE "
            "comment is now carried by the ITP (#92) and the RO Vessel "
            "Hydrostatic Test Procedure (#104)."),
    ),
    "P22-DWG-09-008-001": dict(
        rev="C", delivery="E52", tm=TM, verdict="3-To be revised",
        action=(
            "Re-issued Rev C (E52, 16-Jun-2026). Code 3 - To be revised (TM "
            "N23), reverting the TM N15 2-AN. OBS-01 MAJOR: content changed "
            "(items 8/9/31/32 descriptions + CIP relocation) but the revision "
            "letter kept at C with no new revision-history row/ECN - two "
            "drawings share the identifier Rev C. OBS-02 MAJOR: geometry "
            "aligned to the superseded Equipment Layout Rev B, which is open at "
            "Code 3 (RO Cartridge Filter horizontal); cannot be issued for "
            "construction until the upstream layouts are approved. OBS-03/04 "
            "MINOR: front title-block code -01 vs -001; Rev C date 16-Jun vs "
            "April. Rev D required."),
    ),
}

# ---- Master Register: filas nuevas (#104-107) ----
# (num, document, code, rev, delivery, tm, verdict, status, action, et_deadline)
MR_NEW = [
    (104, "RO Vessel Hydrostatic Test Procedure", "P22-BA-09-000-009", "A",
     "E51", TM, "3-To be revised", "Delivered",
     "First issue (E51). Code 3 - To be revised (TM N23). OBS-01 CRITICAL: no "
     "binding test pressure in the body (only the generic 1.1x design for "
     "ASME / 1.43x for CE rule) and 45.5 bar (~660 psi) on the embedded Protec "
     "report form, vs the 1,800 psi x 1.1 = 1,980 psi (~136.5 bar) the ITP row "
     "2.2 and the waiver require. OBS-02 MAJOR: no ADASA witness/notification "
     "section and non-traceable instrumentation (the HP/LP procedure requires "
     "certificate review + two gauges). OBS-03/04 MINOR: RT-5 hold time; "
     "ADASA-wrapper vs Protec-procedure (Rev 0, 09-10-2025) revision "
     "traceability. Rev B required - this is the test the stamp waiver was "
     "traded for.",
     "ET Sec.8: Prerequisite for HP fabrication start"),
    (105, "HP and LP Pressure Test Procedure", "P22-BA-09-000-010", "A",
     "E51", TM, "3-To be revised", "Delivered",
     "First issue (E51). Code 3 - To be revised (TM N23). OBS-01 MAJOR: the "
     "procedure writes no numeric test pressure or ASME B31.3 factor (only the "
     "generic 'required test pressure'); the ITP fixes 135 bar HP (1.5 x 90 "
     "bar) and 7.5 bar LP, and the HP design pressure (90 bar ITP vs up to 120 "
     "bar ET) is unreconciled. OBS-02 MINOR: step numbering breaks in the "
     "pneumatic section (5.6 reverts to 5.5.17/18; 5.6.5 missing). Rev B "
     "required.",
     "ET Sec.8: Prerequisite for HP fabrication start"),
    (106, "Painting Procedure", "P22-BA-09-000-011", "A",
     "E51", TM, "3-To be revised", "Delivered",
     "First issue (E51). Code 3 - To be revised (TM N23). OBS-01 MAJOR: a "
     "Jotun system is substituted for the Sherwin-Williams system fixed in the "
     "Painting Specifications Rev B (approved Code 1 at TM N11) with no "
     "product-to-product equivalence justification. OBS-02 MAJOR: marine C5-M "
     "durability not demonstrated (the Barrier 80 primer is certified only for "
     "C5-I). OBS-03 MAJOR: anchor profile contradictory (body 50-80 vs form "
     "40-75 micron) and below the 50 micron ET minimum. OBS-04..06 MINOR: RAL "
     "5012 colour, scope bounding (ASTM A-36 only; exclude SS / FRP / HDPE), "
     "QC adhesion/DFT/ISO 2808/SSPC-SP10 references. Architecture "
     "80/200/75=355 micron matches the ET. Rev B required.",
     "ET Sec.8: Prerequisite for fabrication QA"),
    (107, "UHPRO Structural Design Criteria", "P22-CD-09-005-003", "A",
     "E52", TM, "3-To be revised", "Delivered",
     "First issue (E52). Code 3 - To be revised (TM N23). Driver OBS-03 MAJOR: "
     "the criteria omit the lifting/handling load case and the lifting-point "
     "and yoke design criteria the module ET requires (lifting design = BW "
     "Water scope; crane/lifting equipment = ADASA, ET 5.6) - basis for "
     "deliverables #72/#76. OBS-01 MAJOR seismic citation inconsistent: the "
     "parameter table and combinations cite a non-existent NCh2369:2009 while "
     "the code list cites NCh 2369 Of.2003 (the ET edition; this engineering "
     "was contracted and due under Of.2003, before the 2025 oficializacion) - "
     "correct 2009 to Of.2003; the parameter values are already consistent. "
     "OBS-02 MAJOR confirm the NCh 2369 Of.2003 clause 4.5 ASD combinations "
     "govern (not the generic NCh 3171). OBS-04..07 MINOR: revision identity, "
     "Soil E basis, design weight P, wind units/edition. NOTE-01 concrete "
     "grade. Rev B required.",
     "ET Sec.7: Max. 90 days from NTP"),
]

# ---- Revision History: 7 filas nuevas (num, doc, code, rev, delivery, submittal, verdict, keyobs) ----
RH_NEW = [
    (92, "Inspection and Test Plan Offsite", "P22-BA-09-000-004", "C", "E51",
     "25007-0051", "2-AN",
     "Rev C. Code 2 - AN. Materially closes the oldest fabrication "
     "carry-forward (row 2.2: 1800 psi x 1.1 vessel test, ADASA witness, "
     "dossier Hold Points 7.6/8.3; ASME Section X correct for FRP). Rev 0 "
     "edits: declare no-code-stamp basis (OBS-01); vessel test "
     "Witness->Hold Point (OBS-02). Full closure also gated on the RO Vessel "
     "Hydrostatic Test Procedure (#104)."),
    (104, "RO Vessel Hydrostatic Test Procedure", "P22-BA-09-000-009", "A",
     "E51", "25007-0051", "3-To be revised",
     "First issue. Code 3. OBS-01 CRITICAL: no binding test pressure in the "
     "body (generic 1.1x ASME / 1.43x CE) and 45.5 bar on the Protec form vs "
     "the required 1,980 psi (1800 x 1.1). OBS-02 MAJOR: no ADASA witness "
     "section / non-traceable instrumentation. OBS-03/04 MINOR. Rev B "
     "required."),
    (105, "HP and LP Pressure Test Procedure", "P22-BA-09-000-010", "A",
     "E51", "25007-0051", "3-To be revised",
     "First issue. Code 3. OBS-01 MAJOR: no numeric test pressure / B31.3 "
     "factor (ITP fixes 135 bar HP, 7.5 bar LP); HP design 90 vs 120 bar "
     "unreconciled. OBS-02 MINOR: pneumatic-section numbering. Rev B "
     "required."),
    (84, "Plan de Ensayos No Destructivos (NDE) - Super Duplex",
     "P22-BA-09-000-005", "B", "E51", "25007-0051", "3-To be revised",
     "Rev B. Code 3. OBS-01 MAJOR: code editions still placeholders (the Rev "
     "A comment not closed); OBS-02: AWS/DVS mixed with the B31.3 Super Duplex "
     "circuit without joint mapping, UT column is a thickness measurement. "
     "Coverage matches the ET. Rev C required."),
    (106, "Painting Procedure", "P22-BA-09-000-011", "A", "E51", "25007-0051",
     "3-To be revised",
     "First issue. Code 3. OBS-01 MAJOR: Jotun substituted for the approved "
     "Sherwin-Williams system (Painting Specs Rev B, Code 1 TM N11) without "
     "equivalence. OBS-02 MAJOR: C5-M not demonstrated (primer C5-I). OBS-03 "
     "MAJOR: anchor profile 40-75 vs 50-80 micron, below the ET 50. OBS-04..06 "
     "MINOR. Architecture 355 micron matches the ET. Rev B required."),
    (35, "Instrument Location Layout", "P22-DWG-09-008-001", "C", "E52",
     "25007-0052", "3-To be revised",
     "Re-issued Rev C (E52). Code 3 (reverts the N15 2-AN). OBS-01 MAJOR: "
     "content changed (items 8/9/31/32 + CIP) without bumping the revision "
     "letter; two drawings share Rev C. OBS-02 MAJOR: aligned to the "
     "superseded Equipment Layout Rev B (open at Code 3). OBS-03/04 MINOR. Rev "
     "D required."),
    (107, "UHPRO Structural Design Criteria", "P22-CD-09-005-003", "A", "E52",
     "25007-0052", "3-To be revised",
     "First issue (E52). Code 3. Driver OBS-03 MAJOR: omits the lifting load "
     "case + lifting-point/yoke design criteria the ET requires (design=BWW; "
     "crane/equip=ADASA) - basis for #72/#76. OBS-01 MAJOR seismic citation: "
     "table/combos cite non-existent NCh2369:2009 vs code list NCh 2369 Of.2003 "
     "(ET edition; engineering due under Of.2003 pre-2025 oficializacion) - "
     "correct to Of.2003, values OK. OBS-02 MAJOR confirm NCh 2369 Of.2003 "
     "cl.4.5 ASD combos govern (not NCh 3171). OBS-04..07 MINOR + NOTE-01. Rev "
     "B required."),
]


def style_row(ws, dst_row, ref_row, ncols):
    for c in range(1, ncols + 1):
        s = ws.cell(ref_row, c)
        d = ws.cell(dst_row, c)
        d._style = copy(s._style)


def main():
    shutil.copyfile(SRC, BAK)
    print(f"Backup: {BAK}")

    wb = openpyxl.load_workbook(SRC)  # data_only=False: preserva formulas/estilos
    mr = wb["Master Register"]
    rh = wb["Revision History"]
    sm = wb["Summary"]

    ncols_mr = mr.max_column
    # localizar columnas por header
    hdr = {mr.cell(1, c).value: c for c in range(1, ncols_mr + 1)}
    C_NUM = hdr["#"]; C_DOC = hdr["Document"]; C_CODE = hdr["Code / ET Reference"]
    C_REV = hdr["Rev"]; C_DEL = hdr["Delivery"]; C_TM = hdr["TM"]
    C_VER = hdr["Verdict"]; C_STA = hdr["Status"]; C_ACT = hdr["Action Required"]
    C_ETD = hdr["ET Deadline"]

    # 1) updates a filas existentes
    last_data = 1
    for r in range(2, mr.max_row + 1):
        if mr.cell(r, C_CODE).value:
            last_data = r
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code in MR_UPDATES:
            u = MR_UPDATES[code]
            mr.cell(r, C_REV).value = u["rev"]
            mr.cell(r, C_DEL).value = u["delivery"]
            mr.cell(r, C_TM).value = u["tm"]
            mr.cell(r, C_VER).value = u["verdict"]
            mr.cell(r, C_ACT).value = u["action"]
            print(f"  MR update row {r}: {code} -> {u['verdict']} / {u['tm']}")

    # 2) filas nuevas
    ref = last_data
    for i, row in enumerate(MR_NEW, start=1):
        dst = last_data + i
        num, doc, code, rev, dele, tm, ver, sta, act, etd = row
        style_row(mr, dst, ref, ncols_mr)
        mr.cell(dst, C_NUM).value = num
        mr.cell(dst, C_DOC).value = doc
        mr.cell(dst, C_CODE).value = code
        mr.cell(dst, C_REV).value = rev
        mr.cell(dst, C_DEL).value = dele
        mr.cell(dst, C_TM).value = tm
        mr.cell(dst, C_VER).value = ver
        mr.cell(dst, C_STA).value = sta
        mr.cell(dst, C_ACT).value = act
        mr.cell(dst, C_ETD).value = etd
        print(f"  MR new row {dst}: #{num} {code} -> {ver}")

    # 3) Revision History
    rcols = rh.max_column
    rhh = {rh.cell(1, c).value: c for c in range(1, rcols + 1)}
    # conteo previo por codigo para cycle
    prior = collections.Counter()
    rh_last = 1
    for r in range(2, rh.max_row + 1):
        cd = rh.cell(r, rhh["Code / ET Reference"]).value
        if cd:
            prior[str(cd).strip()] += 1
            rh_last = r
    for i, row in enumerate(RH_NEW, start=1):
        dst = rh_last + i
        num, doc, code, rev, dele, sub, ver, keyobs = row
        cycle = prior[code] + 1
        style_row(rh, dst, rh_last, rcols)
        rh.cell(dst, rhh["#"]).value = num
        rh.cell(dst, rhh["Document"]).value = doc
        rh.cell(dst, rhh["Code / ET Reference"]).value = code
        rh.cell(dst, rhh["Rev"]).value = rev
        rh.cell(dst, rhh["Delivery"]).value = dele
        rh.cell(dst, rhh["Submittal"]).value = sub
        rh.cell(dst, rhh["TM"]).value = TM
        rh.cell(dst, rhh["TM Date"]).value = TM_DATE
        rh.cell(dst, rhh["Verdict"]).value = ver
        rh.cell(dst, rhh["Cycle"]).value = cycle
        rh.cell(dst, rhh["Key Observations"]).value = keyobs
        print(f"  RH new row {dst}: #{num} {code} cycle {cycle}")

    # 4) Summary: recomputar desde Master Register
    rows = []
    for r in range(2, mr.max_row + 1):
        code = mr.cell(r, C_CODE).value
        if code and str(code).strip():
            rows.append((str(mr.cell(r, C_VER).value or "").strip(),
                         str(mr.cell(r, C_STA).value or "")))
    total = len(rows)
    delivered = sum(1 for v, s in rows if "Delivered" in s)
    vd = collections.Counter(v for v, s in rows if "Delivered" in s)
    print(f"  Summary recompute: total={total} delivered={delivered} verdicts={dict(vd)}")

    sm["B4"].value = total
    sm["B5"].value = delivered
    sm["B9"].value = "23 (N1 through N23 - TM N5 issued in Rev 0 and Rev 1)"
    sm["B10"].value = "52 (E1 through E52)"
    sm["B12"].value = "18-Jun-2026 (E52)"
    sm["B13"].value = "18-Jun-2026"
    # verdict distribution
    def pct(n):
        return f"{round(100 * n / delivered)}%" if delivered else "0%"
    sm["B27"].value = vd.get("1-Approved", 0); sm["C27"].value = pct(vd.get("1-Approved", 0))
    sm["B28"].value = vd.get("2-AN", 0); sm["C28"].value = pct(vd.get("2-AN", 0))
    sm["B29"].value = vd.get("3-To be revised", 0); sm["C29"].value = pct(vd.get("3-To be revised", 0))
    sm["B30"].value = vd.get("4-Rejected", 0); sm["C30"].value = pct(vd.get("4-Rejected", 0))
    # ITEMS BY SECTION: Sec.1 +1 (Structural, engineering); Sec.3 +3 (procedures)
    sm["B18"].value = 81; sm["C18"].value = 73; sm["D18"].value = "90%"
    sm["B20"].value = 10; sm["C20"].value = 7; sm["D20"].value = "70%"

    wb.save(SRC)
    print(f"Guardado: {SRC}")
    print("NOTA: REVIEW CYCLES distribution (basada en ejemplos) NO se modifico.")


if __name__ == "__main__":
    main()
