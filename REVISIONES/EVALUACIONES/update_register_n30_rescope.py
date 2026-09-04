#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
update_register_n30_rescope.py
Incorpora al Master Deliverable Register el TM N30 RE-ESCOPEADO, que cubre tres
submittals: 25007-0068 (E68), 25007-0069 (E69) y 25007-0070 (E70).

SUPERSEDE A update_register_n31.py, que ya NO debe correrse. Aquel se escribio
cuando la E69 iba a ir en un TM N31 futuro y por eso dejaba el Dossier Index en
'Under review' con TM '--'. Al absorberse la E69 en el N30, el veredicto existe y
se escribe. Este script rehace ese trabajo desde el mismo baseline y agrega la
E70, de modo que correr los dos en cadena duplicaria el item 113.

Idempotencia: restaura SRC desde _pre-N31.xlsx, que es el estado LIMPIO anterior
a la E69 (108 items / 84 delivered, 50/28/6/0, 30 TMs / 68 entregas). Correrlo N
veces produce el mismo libro.

QUE CAMBIA

  #40  P22-DWG-09-005-004 Piping Layout        Rev B -> Rev C, E70, ciclo 3, 2-AN
  #43  P22-LI-09-008-006  DS DP Switch         Rev A -> Rev B, E70, ciclo 2, 1-Approved
  #113 P22-BA-09-000-013  Dossier Index        NUEVO Rev A, E69, ciclo 1, 3-To be revised
  #114 P22-DWG-09-005-007 3D Model             NUEVO Rev A, E70, ciclo 1, 2-AN
  #65  Manufacturing and Testing Dossier       sigue NOT DELIVERED, solo su accion
  #29  DS PLC & HMI Panel Component            fecha del TM: 23-Jul -> 05-Aug-2026

  Los planos de taller 25007-ME-PI-0901-0006 a -0016 NO ENTRAN al registro. El
  transmittal los devuelve declarando que no constituyen un entregable recibido:
  sin codigo ADASA, ausentes del Submittal Form, y la portada del documento que
  los contiene declara "Page 1 of 4". Darles fila aqui seria registrar como
  recibido lo que el transmittal declara no recibido. Entran cuando se sometan
  formalmente; el compromiso queda en el Registro de Compromisos, no aqui.

FECHA DEL TM. El registro anota los veredictos con fecha 05-Aug-2026. Correrlo
ANTES de que el correo salga adelanta el hecho; el script lo advierte por consola
y la casilla queda en el checklist post-envio del correo.

Tally esperado: 50 Code 1 / 29 Code 2 / 7 Code 3 / 0 Code 4 = 86, que iguala a
delivered (ningun entregado queda sin veredicto). 110 items / 86 delivered;
30 TMs (el N30 ya estaba contado) / 70 entregas. ITEMS BY SECTION: Seccion 1
84->85 total y 76->77 delivered (el modelo 3D es ingenieria); Seccion 3 11->12 y
8->9 (el indice del dossier acompana al item 65).
"""
import collections
import os
import shutil
import sys
from copy import copy

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

EV = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register.xlsx")
BAK = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register_pre-N31.xlsx")

TM = "N30"
TM_DATE = "05-Aug-2026"
STATUS_DATE = "05-Aug-2026"

# --- items existentes que cambian de revision con esta entrega ---
MR_UPDATES = {
    40: dict(
        rev="C", delivery="E70", tm=TM, verdict="2-AN", status="Delivered",
        action=(
            "Rev C reviewed and approved as noted (TM N30, 05-Aug-2026, submittal "
            "25007-0070). Two items to incorporate when issuing at IFC Rev 0, with no "
            "intermediate revision: the tie-in schedule with elevations and coordinates at "
            "the module battery limits, and the flange class at each battery-limit "
            "connection. Detailed as OBS-02, OBS-03 and NOTE-01 to NOTE-03 on the annotated "
            "PDF. The drawing itself is otherwise correct against the approved Line List "
            "(P22-LI-09-009-003 Rev 0) and the P&ID Rev D. Rev C grew from 5 pages to 22: "
            "the 17 added pages are shop fabrication sheets under BW Water numbering that "
            "the transmittal returns as not received, since they carry no ADASA document "
            "code, are absent from the Submittal Form, and the cover of the file still "
            "declares Page 1 of 4."),
    ),
    43: dict(
        rev="B", delivery="E70", tm=TM, verdict="1-Approved", status="Delivered",
        action=(
            "Rev B approved (TM N30, 05-Aug-2026, submittal 25007-0070). Accepted as-is; "
            "issue directly at IFC Rev 0, no annotated PDF. The revision changes the wetted "
            "material of the differential pressure switch, which the datasheet states "
            "explicitly, and the change is consistent with the service. No observation."),
    ),
}

# --- items nuevos ---
NEW_ITEMS = [
    dict(
        number=113,
        document="Fabrication and Testing Dossier Index",
        code="P22-BA-09-000-013",
        rev="A", delivery="E69", submittal="25007-0069", cycle=1,
        tm=TM, verdict="3-To be revised", status="Delivered",
        section="3. FABRICATION",
        rh_obs=(
            "Index of 27 chapters with no document number, revision or inclusion status "
            "against any line. Five chapters missing. Re-issue as Rev B."),
        action=(
            "Rev A delivered (E69, 27-Jul-2026, submittal 25007-0069, IFA) in response to "
            "the first tranche of the 25-Jul dossier claim. Verdict 3 - To be revised "
            "(TM N30, 05-Aug-2026): re-issue as Rev B. The index lists 27 chapters with no "
            "document number, revision or inclusion status against any of them, so it "
            "cannot serve as the checklist for the preliminary documentation review that "
            "the Inspection and Testing Base Plan (P22-IT-09-000-001-0) assigns to ADASA at "
            "item 7.6. Missing chapters: preparation for dispatch (cleaning and "
            "preservation, packaging and marking records, packing list) per Inspection and "
            "Test Plan rows 8.1 and 8.2; the FAT Approval Certificate that the Technical "
            "Specification (P22-ET-09-000-001-0), Section 8, declares an integral part of "
            "the final dossier; inspection personnel qualifications and test equipment "
            "calibration per NDE Plan Section 3.0 and Inspection and Test Plan row 2.4; "
            "manufacturer certificates and factory test records of the main equipment per "
            "row 2.3; and a non-conformance and weld repair register. Chapters B3, B4 and "
            "B6 list ultrasonic, dye penetrant and radiography procedures that have never "
            "been submitted to ADASA for review, which the NDE Plan Section 2.0 requires: "
            "these are to be approved before welding starts, since records produced under "
            "unapproved procedures are not admissible into the dossier. Detailed as OBS-01 "
            "to OBS-06 and NOTE-01 to NOTE-04 on the annotated PDF."),
        et_deadline="ET Sec.7: Max. 90 days from NTP",
    ),
    dict(
        number=114,
        document="3D Model (Navisworks)",
        code="P22-DWG-09-005-007",
        rev="A", delivery="E70", submittal="25007-0070", cycle=1,
        tm=TM, verdict="2-AN", status="Delivered",
        section="1. ENGINEERING",
        rh_obs=(
            "First submission of the federated model. File not identified by its document "
            "code and revision. Tags reconciled against the four approved lists."),
        action=(
            "Rev A approved as noted (TM N30, 05-Aug-2026, submittal 25007-0070). First "
            "submission of the federated Navisworks model. To incorporate at IFC Rev 0, "
            "with no intermediate revision: identify the file by its document code and "
            "revision - the internal title is a working name and the origin path is a "
            "personal user folder - and reconcile the tags that diverge from the approved "
            "lists. The tags were read from the AutoCAD Plant 3D property category of the "
            "model itself, not from layer names. No requirement of BIM-level metadata "
            "rigour is raised: publication properties, selection sets and viewpoints are "
            "not called for by the Technical Specification, and an observation with no "
            "requirement behind it is refutable. What is required is consistency with the "
            "Valve List (P22-LI-09-005-002 Rev D), the Instrument List (P22-LI-09-008-003 "
            "Rev E) and the Line List (P22-LI-09-009-003 Rev 0), all three approved at "
            "Code 1, and with the Equipment List (P22-LI-09-005-001 Rev B). The model is "
            "not an annotated deliverable: a Navisworks file cannot carry the annotation "
            "format used for drawings, so the observations are stated in full in the "
            "transmittal text and Section 4 declares the exception."),
        et_deadline="ET Sec.7: Max. 90 days from NTP",
    ),
]

# Clave por NUMERO DE ITEM, no por 'Code / ET Reference': los items 62, 63 y 65
# comparten la referencia 'ET Sec 7, p.28' y usarla como clave toca tres filas.
MR_ACTION_FIXES = {
    65: (
        "NOT DELIVERED. The Fabrication and Testing Dossier Index (P22-BA-09-000-013 Rev A) "
        "arrived on 27-Jul-2026 as item 113 of this register and was returned at Code 3 in "
        "TM N30, but an index is not the dossier: no fabrication or testing record has been "
        "submitted to date. The second tranche of the 25-Jul claim - complete preliminary "
        "dossier structured against the index, due Friday 31-Jul-2026 - lapsed with no "
        "evidence of delivery. Contractual basis: Technical Specification "
        "(P22-ET-09-000-001-0), Section 7, maximum 90 days from award, including the "
        "manufacturing certificates of the special steels for the super duplex equipment. "
        "Prerequisite for factory acceptance and gating item 8.3 (final dossier approval, "
        "Hold) and item 8.4 (release for dispatch) of the Inspection and Testing Base Plan, "
        "which support the 40% payment milestone. Part of the Week 1 records reached ADASA "
        "through the third-party inspector's Annex A, not from the supplier."),
}

# El TM N30 dejo de emitirse el 23-Jul y salio el 05-Aug con el alcance ampliado.
# La fila de Revision History del item 29 se escribio con la fecha del borrador.
RH_DATE_FIXES = {(29, "0"): TM_DATE}


def style_row(ws, dst_row, ref_row, ncols):
    for c in range(1, ncols + 1):
        ws.cell(dst_row, c)._style = copy(ws.cell(ref_row, c)._style)


def main():
    if not os.path.exists(BAK):
        raise SystemExit("ERROR: falta el baseline limpio %s" % os.path.basename(BAK))
    shutil.copyfile(BAK, SRC)
    print("Restaurado SRC desde baseline limpio:", os.path.basename(BAK))

    wb = openpyxl.load_workbook(SRC)
    mr = wb["Master Register"]
    rh = wb["Revision History"]
    sm = wb["Summary"]

    # Descuadre HEREDADO de ITEMS BY SECTION, medido ANTES de tocar nada. El
    # bloque esta hardcodeado en el libro y se ha editado a mano durante meses.
    # No se corrige a ciegas: la asignacion de seccion no es derivable de ninguna
    # columna. Lo que si se exige es que este updater NO lo agrande.
    delta_previo = (
        sum(int(sm.cell(r, 2).value) for r in range(18, 23)) - int(sm["B4"].value),
        sum(int(sm.cell(r, 3).value) for r in range(18, 23)) - int(sm["B5"].value),
    )

    ncols = mr.max_column
    hdr = {mr.cell(1, c).value: c for c in range(1, ncols + 1)}
    C_NUM, C_DOC, C_CODE = hdr["#"], hdr["Document"], hdr["Code / ET Reference"]
    C_REV, C_DEL, C_TM = hdr["Rev"], hdr["Delivery"], hdr["TM"]
    C_VER, C_STA, C_ACT = hdr["Verdict"], hdr["Status"], hdr["Action Required"]
    C_DDL = hdr["ET Deadline"]

    # 1) items existentes que suben de revision + higiene del item 65
    pend_upd, pend_fix = set(MR_UPDATES), set(MR_ACTION_FIXES)
    for r in range(2, mr.max_row + 1):
        num = mr.cell(r, C_NUM).value
        num = int(num) if isinstance(num, (int, float)) else None
        if num in pend_upd:
            u = MR_UPDATES[num]
            antes = "%s Rev %s (%s, %s)" % (mr.cell(r, C_CODE).value,
                                            mr.cell(r, C_REV).value,
                                            mr.cell(r, C_DEL).value,
                                            mr.cell(r, C_VER).value)
            for col, key in ((C_REV, "rev"), (C_DEL, "delivery"), (C_TM, "tm"),
                             (C_VER, "verdict"), (C_STA, "status"), (C_ACT, "action")):
                mr.cell(r, col).value = u[key]
            pend_upd.discard(num)
            print("  MR fila %3d item #%-3d %s  ->  Rev %s (%s, %s)"
                  % (r, num, antes, u["rev"], u["delivery"], u["verdict"]))
        if num in pend_fix:
            if str(mr.cell(r, C_STA).value or "").strip().upper() != "NOT DELIVERED":
                raise SystemExit(
                    "ERROR: el item #%d ya no esta NOT DELIVERED; revisar antes de tocar"
                    % num)
            mr.cell(r, C_ACT).value = MR_ACTION_FIXES[num]
            pend_fix.discard(num)
            print("  MR fila %3d item #%-3d sigue NOT DELIVERED, accion refrescada"
                  % (r, num))
    if pend_upd or pend_fix:
        raise SystemExit("ERROR: items no encontrados en Master Register: %s"
                         % (pend_upd | pend_fix))

    # 2) items nuevos al final, con el estilo de la ultima fila
    ultima = max(r for r in range(2, mr.max_row + 1) if mr.cell(r, C_CODE).value)
    max_num = max(int(mr.cell(r, C_NUM).value) for r in range(2, mr.max_row + 1)
                  if isinstance(mr.cell(r, C_NUM).value, (int, float)))
    existentes = {str(mr.cell(r, C_CODE).value or "").strip()
                  for r in range(2, mr.max_row + 1)}

    for i, it in enumerate(NEW_ITEMS, start=1):
        if it["code"] in existentes:
            raise SystemExit("ERROR: %s ya existe en el registro" % it["code"])
        if it["number"] != max_num + i:
            raise SystemExit("ERROR: el item %s esperaba el numero %d y el registro va en %d"
                             % (it["code"], it["number"], max_num + i))
        dst = ultima + i
        style_row(mr, dst, ultima, ncols)
        mr.cell(dst, C_NUM).value = it["number"]
        mr.cell(dst, C_DOC).value = it["document"]
        mr.cell(dst, C_CODE).value = it["code"]
        mr.cell(dst, C_REV).value = it["rev"]
        mr.cell(dst, C_DEL).value = it["delivery"]
        mr.cell(dst, C_TM).value = it["tm"]
        mr.cell(dst, C_VER).value = it["verdict"]
        mr.cell(dst, C_STA).value = it["status"]
        mr.cell(dst, C_ACT).value = it["action"]
        mr.cell(dst, C_DDL).value = it["et_deadline"]
        print("  MR nuevo item #%-3d fila %3d: %s Rev %s (%s) -> %s"
              % (it["number"], dst, it["code"], it["rev"], it["delivery"], it["verdict"]))

    # 3) Revision History: una fila por documento dispuesto en este TM
    rh_cols = rh.max_column
    rh_last = max(r for r in range(2, rh.max_row + 1) if rh.cell(r, 3).value)

    rh_nuevas = []
    for num in sorted(MR_UPDATES):
        u = MR_UPDATES[num]
        for r in range(2, mr.max_row + 1):
            if mr.cell(r, C_NUM).value == num:
                doc, code = mr.cell(r, C_DOC).value, mr.cell(r, C_CODE).value
                break
        rh_nuevas.append((num, doc, code, u["rev"], u["delivery"], "25007-0070",
                          TM, TM_DATE, u["verdict"], u.get("cycle")))
    for it in NEW_ITEMS:
        rh_nuevas.append((it["number"], it["document"], it["code"], it["rev"],
                          it["delivery"], it["submittal"], TM, TM_DATE, it["verdict"],
                          it["cycle"]))

    # el ciclo de un item existente = filas previas suyas en RH + 1
    for i, fila in enumerate(rh_nuevas):
        if fila[9] is None:
            previas = sum(1 for r in range(2, rh_last + 1) if rh.cell(r, 1).value == fila[0])
            rh_nuevas[i] = fila[:9] + (previas + 1,)

    for i, fila in enumerate(rh_nuevas, start=1):
        dst = rh_last + i
        style_row(rh, dst, rh_last, rh_cols)
        for c, v in enumerate(fila, start=1):
            rh.cell(dst, c).value = v
        obs = next((it["rh_obs"] for it in NEW_ITEMS if it["number"] == fila[0]), None)
        if obs and rh_cols >= 11:
            rh.cell(dst, 11).value = obs
        print("  RH fila %3d: item #%-3d %s Rev %s ciclo %s -> %s"
              % (dst, fila[0], fila[2], fila[3], fila[9], fila[8]))

    # correccion de fecha: el N30 se redacto el 23-Jul y se emitio el 05-Aug
    for r in range(2, rh.max_row + 1):
        clave = (rh.cell(r, 1).value, str(rh.cell(r, 4).value or ""))
        if clave in RH_DATE_FIXES and str(rh.cell(r, 7).value or "") == TM:
            antes = rh.cell(r, 8).value
            rh.cell(r, 8).value = RH_DATE_FIXES[clave]
            print("  RH fila %3d: item #%s fecha del TM %s -> %s"
                  % (r, clave[0], antes, RH_DATE_FIXES[clave]))

    # 4) Summary recomputado desde el Master Register
    stat, vd, total = collections.Counter(), collections.Counter(), 0
    for r in range(2, mr.max_row + 1):
        if str(mr.cell(r, C_CODE).value or "").strip():
            total += 1
            su = str(mr.cell(r, C_STA).value or "").strip().upper()
            stat[su] += 1
            if su == "DELIVERED":
                vd[str(mr.cell(r, C_VER).value or "").strip()] += 1
    delivered = stat.get("DELIVERED", 0)

    sm["B4"].value = total
    sm["B5"].value = delivered
    sm["B7"].value = stat.get("PARTIAL", 0)
    sm["B8"].value = stat.get("NOT DELIVERED", 0)
    sm["B9"].value = "30 (N1 through N30 - TM N5 issued in Rev 0 and Rev 1)"
    sm["B10"].value = "70 (E1 through E70)"
    sm["B12"].value = "05-Aug-2026 (E70)"
    sm["B13"].value = STATUS_DATE

    con_veredicto = sum(vd[k] for k in
                        ("1-Approved", "2-AN", "3-To be revised", "4-Rejected"))

    def pct(n):
        return "%d%%" % round(100 * n / con_veredicto) if con_veredicto else "0%"

    for fila, clave in ((27, "1-Approved"), (28, "2-AN"),
                        (29, "3-To be revised"), (30, "4-Rejected")):
        sm.cell(fila, 2).value = vd.get(clave, 0)
        sm.cell(fila, 3).value = pct(vd.get(clave, 0))

    # 5) ITEMS BY SECTION — hardcodeada en el libro, se edita a mano
    for it in NEW_ITEMS:
        fila_sec = next((r for r in range(18, 23)
                         if str(sm.cell(r, 1).value or "").startswith(it["section"])), None)
        if fila_sec is None:
            raise SystemExit("ERROR: seccion '%s' no encontrada en ITEMS BY SECTION"
                             % it["section"])
        sm.cell(fila_sec, 2).value = int(sm.cell(fila_sec, 2).value) + 1
        if it["status"].strip().upper() == "DELIVERED":
            sm.cell(fila_sec, 3).value = int(sm.cell(fila_sec, 3).value) + 1
        print("  Summary %-16s total=%s delivered=%s"
              % (it["section"], sm.cell(fila_sec, 2).value, sm.cell(fila_sec, 3).value))

    suma_total = sum(int(sm.cell(r, 2).value) for r in range(18, 23))
    suma_deliv = sum(int(sm.cell(r, 3).value) for r in range(18, 23))
    delta_post = (suma_total - total, suma_deliv - delivered)
    if delta_post != delta_previo:
        raise SystemExit(
            "ERROR: este updater movio el descuadre de ITEMS BY SECTION de %s a %s"
            % (delta_previo, delta_post))
    if any(delta_post):
        print("  AVISO - descuadre HEREDADO en ITEMS BY SECTION, no introducido aqui: "
              "las secciones declaran %+d items y %+d delivered respecto del registro "
              "(secciones %d/%d contra totales %d/%d). Medido tambien sobre el backup "
              "limpio pre-N31."
              % (delta_post[0], delta_post[1], suma_total, suma_deliv, total, delivered))

    if con_veredicto != delivered:
        print("  AVISO - %d entregados sin veredicto" % (delivered - con_veredicto))

    wb.save(SRC)
    print("\nGuardado:", os.path.basename(SRC))
    print("  %d items / %d delivered | veredictos %s"
          % (total, delivered,
             "/".join(str(vd.get(k, 0)) for k in
                      ("1-Approved", "2-AN", "3-To be revised", "4-Rejected"))))
    print("  30 TMs / 70 entregas | ITEMS BY SECTION %d/%d (descuadre heredado %+d/%+d)"
          % (suma_total, suma_deliv, delta_post[0], delta_post[1]))
    print("\n  Los veredictos quedan con fecha %s. Si el correo del TM N30 no ha salido,\n"
          "  esta escritura adelanta el hecho: re-correr tras el envio." % TM_DATE)


if __name__ == "__main__":
    main()
