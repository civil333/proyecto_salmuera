#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
update_register_n24.py
Actualiza el Master Deliverable Register a TM N24 (submittals 25007-0053 E53 +
25007-0054 E54). Update EXPLICITO modelado en update_register_n23.py (no el
fragil update_register.py).

Particularidad de TM N24: los 5 documentos son re-revisiones de documentos ya
entregados (ya contados como Delivered en revisiones previas N20/N22), por lo que
NO se agregan filas nuevas al Master Register; las 5 filas existentes se ACTUALIZAN
(Rev / Delivery / TM / Verdict / Action). Total y Delivered no cambian; solo la
distribucion de veredictos (recomputada desde la hoja).

Hace:
  1. Respalda el original (-002-0) a *_pre-N24.xlsx.
  2. Master Register: actualiza 5 filas existentes por codigo:
     #31 IO List (Rev 2->3, N20 3 -> N24 2-AN), #81 LCP Datasheet (Rev B->0,
     N20 2-AN -> N24 2-AN), #36 I&C Cable Schedule (Rev 1->2, N20 3 -> N24 1),
     #52 Data Transfer List Modbus (Rev 1->2, N20 2-AN -> N24 1), #10 DS CIP
     Cartridge Filter (Rev D->E, N22 3 -> N24 1).
  3. Revision History: agrega 5 filas (1 por doc del TM N24), cycle derivado del
     conteo previo del codigo.
  4. Summary: recomputa Total/Delivered y la distribucion de veredictos DESDE la
     hoja Master Register; actualiza Transmittals (24), Deliveries (54), Latest
     Delivery (E54, 23-Jun), Status Date. ITEMS BY SECTION no cambia (re-revisiones
     de docs ya entregados; sin nuevas entregas por seccion).
  NO toca REVIEW CYCLES distribution (basada en ejemplos; recuento manual).
"""
import os
import shutil
import collections
from copy import copy
import openpyxl

EV = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register.xlsx")
BAK = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register_pre-N24.xlsx")

TM = "N24"
TM_DATE = "23-Jun-2026"

# ---- Master Register: updates a filas existentes (por codigo) ----
MR_UPDATES = {
    # IO List Rev 3 -> Code 2 - Approved as Noted
    "P22-LI-09-008-001": dict(
        rev="3", delivery="E53", tm=TM, verdict="2-AN",
        action=(
            "Rev 3 delivered (E53). Code 2 - Approved as Noted (TM N24). The two "
            "module-to-external-PLC interface signals closed at TM N19 (external "
            "enable XA005, module running status YA001) are present and correctly "
            "typed as relay contacts; the field I/O scheme over Ethernet/IP is "
            "consistent with the approach accepted at TM N20. Two edits at IFC "
            "(no new review cycle): OBS-01 NEW REQUEST - ADASA extends the "
            "interface with two further hardwired relay-contact signals (SYSTEM "
            "FAULT STATUS TO DCS, SYSTEM LOCAL/REMOTE STATUS TO DCS, both "
            "outputs); OBS-02 MINOR - complete the signal-type count cell in "
            "items 130/135. Issue for construction also depends on the Plant "
            "Control Philosophy children (Section 3)."),
    ),
    # LCP Datasheet Rev 0 -> Code 2 - Approved as Noted
    "P22-ET-09-007-005": dict(
        rev="0", delivery="E53", tm=TM, verdict="2-AN",
        action=(
            "Rev 0 delivered (E53). Code 2 - Approved as Noted (TM N24). The "
            "panel enclosure reconfirms the correct marine specification (SS316L, "
            "NEMA 4X / IP66); the enclosure contradiction tracked since TM N20 "
            "resides in the Outline Panel Drawing (P22-CD-09-008-001), not here. "
            "Two minor edits at IFC: OBS-01 reconcile the panel internal "
            "consumption and the UPS SAI-09-001 figure on sheet 114; OBS-02 unify "
            "the document code to P22-ET-09-007-005 across the body. NOTE-01 HART "
            "acquisition (5069 family does not decode HART to the controller) is "
            "governed by the PLC/HMI Datasheet (P22-ET-09-008-001) and tracked in "
            "Section 3."),
    ),
    # I&C Cable Schedule Rev 2 -> Code 1 - Approved
    "P22-LI-09-008-002": dict(
        rev="2", delivery="E53", tm=TM, verdict="1-Approved",
        action=(
            "Rev 2 delivered (E53). Code 1 - Approved (TM N24). Field-instrument "
            "coverage is one-to-one with the IO List and the cable types match "
            "the signal types (shielded instrument cable for 4-20 mA + HART, "
            "Ethernet for Ethernet/IP devices, control cable for hardwired I/O); "
            "the four corrections from TM N20 are incorporated. No defect of this "
            "document; issue directly at IFC Rev 0."),
    ),
    # Data Transfer List (Modbus TCP/IP) Rev 2 -> Code 1 - Approved
    "P22-LI-09-008-004": dict(
        rev="2", delivery="E53", tm=TM, verdict="1-Approved",
        action=(
            "Rev 2 delivered (E53). Code 1 - Approved (TM N24). The Modbus "
            "interface to the supervisory system is consistent; no module safety "
            "or start signal is carried by Modbus alone (the external start and "
            "run-status interface remains hardwired in the IO List); the "
            "conductivity-scale item from the previous cycle is closed. Issue "
            "directly at IFC Rev 0. One related verification tracked in Section 3: "
            "confirm against the Control Matrix that the RO HP Pump start respects "
            "the external hardwired ENABLE (XA005), not a Modbus bit."),
    ),
    # DS CIP Cartridge Filter Rev E -> Code 1 - Approved
    "P22-ET-09-009-006": dict(
        rev="E", delivery="E54", tm=TM, verdict="1-Approved",
        action=(
            "Rev E delivered (E54). Code 1 - Approved (TM N24). Materially closes "
            "the TM N22 observation: the cartridge gasket is now EPDM and a "
            "material-compatibility statement for the FRP housing and the seal "
            "against the CIP fluid (pH 2 to 12) is attached; the prior minor "
            "items (component name, per-cartridge surface area, filtration rate) "
            "are corrected. Design pressure/temperature (7 bar / 45 C, hydrotest "
            "8 bar) and the vertical configuration match the ET. NOTE-01: confirm "
            "at IFC that the housing-cover gasket is also EPDM (vendor generic "
            "guide cites nitrile); material certificates go to the fabrication "
            "dossier. Issue directly at IFC Rev 0."),
    ),
}

# ---- Revision History: 5 filas nuevas (num, doc, code, rev, delivery, submittal, verdict, keyobs) ----
RH_NEW = [
    (31, "IO List", "P22-LI-09-008-001", "3", "E53", "25007-0053", "2-AN",
     "Rev 3. Code 2 - AN. The two module-to-external-PLC signals closed at TM "
     "N19 (enable XA005, running status YA001) present and correctly typed as "
     "relay contacts; Ethernet/IP field scheme per TM N20. OBS-01 NEW REQUEST: "
     "ADASA extends the interface with two further hardwired relay-contact "
     "outputs (SYSTEM FAULT STATUS, SYSTEM LOCAL/REMOTE STATUS to DCS). OBS-02 "
     "MINOR: signal-type count cell in items 130/135. Fold into IFC; issue also "
     "gated on the Control Philosophy children (Section 3)."),
    (81, "Datasheet Local Control Panel (LCP)", "P22-ET-09-007-005", "0", "E53",
     "25007-0053", "2-AN",
     "Rev 0. Code 2 - AN. Enclosure reconfirms the correct marine spec (SS316L, "
     "NEMA 4X / IP66); the contradiction tracked since TM N20 resides in the "
     "Outline Panel Drawing, not here. OBS-01 MINOR: panel internal "
     "consumption / UPS figure on sheet 114. OBS-02 MINOR: code unify to "
     "P22-ET-09-007-005 across the body. NOTE-01 HART acquisition governed by "
     "the PLC/HMI Datasheet (P22-ET-09-008-001), tracked. Fold into IFC."),
    (36, "I&C Cable Schedule", "P22-LI-09-008-002", "2", "E53", "25007-0053",
     "1-Approved",
     "Rev 2. Code 1 - Approved. Field-instrument coverage one-to-one with the "
     "IO List; cable types match the signal types (shielded instrument cable "
     "for 4-20 mA + HART, Ethernet for Ethernet/IP, control cable for hardwired "
     "I/O); the four TM N20 corrections incorporated. No defect; issue at IFC "
     "Rev 0."),
    (52, "Data Transfer List (Modbus TCP)", "P22-LI-09-008-004", "2", "E53",
     "25007-0053", "1-Approved",
     "Rev 2. Code 1 - Approved. Modbus interface to the supervisory system "
     "consistent; no safety/start signal carried by Modbus alone (external "
     "start and run-status remain hardwired in the IO List); the "
     "conductivity-scale item closed. Issue at IFC Rev 0. Section 3 "
     "verification: RO HP Pump start respects the hardwired ENABLE (XA005) vs a "
     "Modbus bit, against the Control Matrix."),
    (10, "DS CIP Cartridge Filter", "P22-ET-09-009-006", "E", "E54",
     "25007-0054", "1-Approved",
     "Rev E. Code 1 - Approved. Materially closes the TM N22 observation: "
     "gasket now EPDM + material-compatibility statement for the FRP housing "
     "against the CIP fluid (pH 2-12); prior minor items (component name, "
     "surface area, filtration rate) corrected. Design 7 bar / 45 C, hydrotest "
     "8 bar, vertical config match the ET. NOTE-01: confirm cover gasket also "
     "EPDM; material certs to the fabrication dossier. Issue at IFC Rev 0."),
]


def style_row(ws, dst_row, ref_row, ncols):
    for c in range(1, ncols + 1):
        s = ws.cell(ref_row, c)
        d = ws.cell(dst_row, c)
        d._style = copy(s._style)


def main():
    shutil.copyfile(SRC, BAK)
    print(f"Backup: {BAK}")

    wb = openpyxl.load_workbook(SRC)  # data_only=False: preserva formulas/estilos
    mr = wb["Master Register"]
    rh = wb["Revision History"]
    sm = wb["Summary"]

    ncols_mr = mr.max_column
    hdr = {mr.cell(1, c).value: c for c in range(1, ncols_mr + 1)}
    C_NUM = hdr["#"]; C_DOC = hdr["Document"]; C_CODE = hdr["Code / ET Reference"]
    C_REV = hdr["Rev"]; C_DEL = hdr["Delivery"]; C_TM = hdr["TM"]
    C_VER = hdr["Verdict"]; C_STA = hdr["Status"]; C_ACT = hdr["Action Required"]
    C_ETD = hdr["ET Deadline"]

    # 1) updates a filas existentes (sin filas nuevas en TM N24)
    seen = set()
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code in MR_UPDATES:
            u = MR_UPDATES[code]
            mr.cell(r, C_REV).value = u["rev"]
            mr.cell(r, C_DEL).value = u["delivery"]
            mr.cell(r, C_TM).value = u["tm"]
            mr.cell(r, C_VER).value = u["verdict"]
            mr.cell(r, C_ACT).value = u["action"]
            seen.add(code)
            print(f"  MR update row {r}: {code} -> {u['verdict']} / {u['tm']} (Rev {u['rev']})")
    missing = set(MR_UPDATES) - seen
    if missing:
        raise SystemExit(f"ERROR: codigos N24 no encontrados en Master Register: {missing}")

    # 2) Revision History
    rcols = rh.max_column
    rhh = {rh.cell(1, c).value: c for c in range(1, rcols + 1)}
    prior = collections.Counter()
    rh_last = 1
    for r in range(2, rh.max_row + 1):
        cd = rh.cell(r, rhh["Code / ET Reference"]).value
        if cd:
            prior[str(cd).strip()] += 1
            rh_last = r
    for i, row in enumerate(RH_NEW, start=1):
        dst = rh_last + i
        num, doc, code, rev, dele, sub, ver, keyobs = row
        cycle = prior[code] + 1
        style_row(rh, dst, rh_last, rcols)
        rh.cell(dst, rhh["#"]).value = num
        rh.cell(dst, rhh["Document"]).value = doc
        rh.cell(dst, rhh["Code / ET Reference"]).value = code
        rh.cell(dst, rhh["Rev"]).value = rev
        rh.cell(dst, rhh["Delivery"]).value = dele
        rh.cell(dst, rhh["Submittal"]).value = sub
        rh.cell(dst, rhh["TM"]).value = TM
        rh.cell(dst, rhh["TM Date"]).value = TM_DATE
        rh.cell(dst, rhh["Verdict"]).value = ver
        rh.cell(dst, rhh["Cycle"]).value = cycle
        rh.cell(dst, rhh["Key Observations"]).value = keyobs
        print(f"  RH new row {dst}: #{num} {code} cycle {cycle} -> {ver}")

    # 3) Summary: recomputar desde Master Register
    rows = []
    for r in range(2, mr.max_row + 1):
        code = mr.cell(r, C_CODE).value
        if code and str(code).strip():
            rows.append((str(mr.cell(r, C_VER).value or "").strip(),
                         str(mr.cell(r, C_STA).value or "")))
    total = len(rows)
    delivered = sum(1 for v, s in rows if "Delivered" in s)
    vd = collections.Counter(v for v, s in rows if "Delivered" in s)
    print(f"  Summary recompute: total={total} delivered={delivered} verdicts={dict(vd)}")

    sm["B4"].value = total
    sm["B5"].value = delivered
    sm["B9"].value = "24 (N1 through N24 - TM N5 issued in Rev 0 and Rev 1)"
    sm["B10"].value = "54 (E1 through E54)"
    sm["B12"].value = "23-Jun-2026 (E54)"
    sm["B13"].value = "23-Jun-2026"

    def pct(n):
        return f"{round(100 * n / delivered)}%" if delivered else "0%"
    sm["B27"].value = vd.get("1-Approved", 0); sm["C27"].value = pct(vd.get("1-Approved", 0))
    sm["B28"].value = vd.get("2-AN", 0); sm["C28"].value = pct(vd.get("2-AN", 0))
    sm["B29"].value = vd.get("3-To be revised", 0); sm["C29"].value = pct(vd.get("3-To be revised", 0))
    sm["B30"].value = vd.get("4-Rejected", 0); sm["C30"].value = pct(vd.get("4-Rejected", 0))
    # ITEMS BY SECTION: sin cambio en TM N24 (re-revisiones de docs ya entregados)

    wb.save(SRC)
    print(f"Guardado: {SRC}")
    print("NOTA: ITEMS BY SECTION y REVIEW CYCLES distribution NO se modificaron "
          "(re-revisiones; recuento manual).")


if __name__ == "__main__":
    main()
