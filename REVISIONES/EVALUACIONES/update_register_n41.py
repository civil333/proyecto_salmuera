#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
update_register_n41.py — aplica el Transmittal N41 al Master Deliverable Register.

TM N41 (P22-TM-09-000-041-0), emitido por Victor Gutierrez el 2 de octubre de 2026, sobre el
submittal 25007-0098 (ENTREGA 98, recibida el viernes 25-Sep): un documento, el Factory
Acceptance Test Procedure del modulo, Rev B.

IDEMPOTENCIA. Si existe el backup _pre-N41.xlsx, restaura SRC desde ese backup antes de
aplicar. Re-correr el script es seguro.

VEREDICTO GLOBAL: 3 — To be revised. 1 Codigo 3.

EFECTO SOBRE EL TALLY, declarado ANTES de correr. Se parte del cierre del N40: 71 / 16 / 3 / 0
con 113 items y 90 entregados. La fila #85 (P22-BA-09-000-017) pasa de Rev A a Rev B y sigue
en Codigo 3: el tally no cambia.

  Tally esperado tras el N41: 71 / 16 / 3 / 0. Items 113, entregados 90.

ENTREGAS: E96, E97, E99 y E100 se recibieron entre el 22-Sep y el 1-Oct y quedan sin
transmittal (van al N42). El resumen cuenta 100 entregas recibidas.

CORRER DESPUES DE ENVIAR el transmittal (ya enviado el 2-Oct-2026 a las 11:29).
"""
import os
import shutil
import collections
from copy import copy
import openpyxl

EV = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register.xlsx")
BAK = os.path.join(EV, "P22-IT-06-000-002-0_Master-Deliverable-Register_pre-N41.xlsx")
TM = "N41"
TM_DATE = "02-Oct-2026"
TALLY_ESPERADO = {"1-Approved": 71, "2-AN": 16, "3-To be revised": 3, "4-Rejected": 0}

CODE_FIXES = {}
ROW85 = None

MR_UPDATES = {
    "P22-BA-09-000-017": dict(
        rev="B", delivery="E98", tm=TM, verdict="3-To be revised",
        action=("Rev B submitted for approval (E98, 25-Sep-2026, submittal 25007-0098), in the "
                "date fixed at TM N40. Code 3 - To be revised (TM N41, issued by Victor Gutierrez "
                "on 02-Oct-2026). Of the twelve points of TM N40, OBS-04 (hydrostatic tests as "
                "prerequisite) and NOTE-02 (dry FAT scope) close; OBS-03 closes with comments "
                "(reference to the ITP responsibility matrix and Clause 37 notices). Open: OBS-01 "
                "and OBS-09 partially; OBS-02 (record forms with acceptance values for 5.1 to "
                "5.8), NOTE-01 (calibrated instruments), OBS-05 (tolerances and documents by code "
                "and revision), OBS-06 (equipment by tag and test sequence), OBS-07 (loop checks "
                "of I/O List Rev 6 and scenarios per Control Philosophy Rev 1 and Alarm and "
                "Interlock List Rev 0), OBS-08 (SEC components and certificates) and OBS-10 "
                "(written approval on Hold records; notice and attendance on Witness records). "
                "No date for Rev C in the transmittal. Row 7.1 of the ITP is an ADASA Hold Point: "
                "the FAT cannot formally open until the procedure is approved."),
    ),
}

RH_NEW = [
    ("Factory Acceptance Test Procedure (module)", "P22-BA-09-000-017", "B", "E98", "25007-0098",
     "3-To be revised",
     "Second issue. Closes OBS-04 and NOTE-02, OBS-03 with comments. Record forms still without "
     "acceptance values; no loop checks or control logic scenarios; equipment by tag, "
     "tolerances, SEC certificates and Hold/Witness approvals still open."),
]


def style_row(ws, dst_row, ref_row, ncols):
    for c in range(1, ncols + 1):
        ws.cell(dst_row, c)._style = copy(ws.cell(ref_row, c)._style)


def main():
    if os.path.exists(BAK):
        shutil.copyfile(BAK, SRC)
        print(f"Restaurado SRC desde baseline limpio: {BAK}")
    else:
        shutil.copyfile(SRC, BAK)
        print(f"Baseline creado: {BAK}")
    wb = openpyxl.load_workbook(SRC)
    mr = wb["Master Register"]; rh = wb["Revision History"]; sm = wb["Summary"]
    ncols_mr = mr.max_column
    hdr = {mr.cell(1, c).value: c for c in range(1, ncols_mr + 1)}
    C_NUM = hdr["#"]; C_DOC = hdr["Document"]; C_CODE = hdr["Code / ET Reference"]
    C_REV = hdr["Rev"]; C_DEL = hdr["Delivery"]; C_TM = hdr["TM"]
    C_VER = hdr["Verdict"]; C_STA = hdr["Status"]; C_ACT = hdr["Action Required"]

    # 1) correcciones de codigo y fila 85
    for r in range(2, mr.max_row + 1):
        try:
            num = int(mr.cell(r, C_NUM).value)
        except (TypeError, ValueError):
            continue
        if num in CODE_FIXES:
            antes = mr.cell(r, C_CODE).value
            mr.cell(r, C_CODE).value = CODE_FIXES[num]
            print(f"  MR code fix row {r} (#{num}): {antes} -> {CODE_FIXES[num]}")
        if ROW85 and num == 85:
            assert str(mr.cell(r, C_STA).value).strip().upper() == "NOT DELIVERED", mr.cell(r, C_STA).value
            mr.cell(r, C_DOC).value = ROW85["document"]
            mr.cell(r, C_CODE).value = ROW85["code"]
            mr.cell(r, C_REV).value = ROW85["rev"]
            mr.cell(r, C_DEL).value = ROW85["delivery"]
            mr.cell(r, C_TM).value = ROW85["tm"]
            mr.cell(r, C_VER).value = ROW85["verdict"]
            mr.cell(r, C_STA).value = ROW85["status"]
            mr.cell(r, C_ACT).value = ROW85["action"]
            print(f"  MR row {r} (#85): placeholder FAT -> {ROW85['code']} Rev A, Code 3, Delivered")

    # 2) updates por codigo
    seen = set()
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code in MR_UPDATES:
            u = MR_UPDATES[code]
            mr.cell(r, C_REV).value = u["rev"]; mr.cell(r, C_DEL).value = u["delivery"]
            mr.cell(r, C_TM).value = u["tm"]; mr.cell(r, C_VER).value = u["verdict"]
            mr.cell(r, C_STA).value = "Delivered"; mr.cell(r, C_ACT).value = u["action"]
            seen.add(code)
            print(f"  MR update row {r}: {code} -> {u['verdict']} / {u['tm']} ({u['rev']})")
    missing = set(MR_UPDATES) - seen
    if missing:
        raise SystemExit(f"ERROR: codigos N41 no encontrados en Master Register: {missing}")

    # 3) Revision History
    reg_num = {}
    for r in range(2, mr.max_row + 1):
        code = str(mr.cell(r, C_CODE).value or "").strip()
        if code:
            try:
                reg_num[code] = int(mr.cell(r, C_NUM).value)
            except (TypeError, ValueError):
                pass
    rcols = rh.max_column
    rhh = {rh.cell(1, c).value: c for c in range(1, rcols + 1)}
    prior = collections.Counter(); rh_last = 1
    for r in range(2, rh.max_row + 1):
        cd = rh.cell(r, rhh["Code / ET Reference"]).value
        if cd:
            prior[str(cd).strip()] += 1; rh_last = r
    for i, row in enumerate(RH_NEW, start=1):
        dst = rh_last + i
        doc, code, rev, dele, sub, ver, keyobs = row
        num = reg_num.get(code, "")
        cycle = prior[code] + 1; prior[code] += 1
        style_row(rh, dst, rh_last, rcols)
        for k, v in [("#", num), ("Document", doc), ("Code / ET Reference", code), ("Rev", rev),
                     ("Delivery", dele), ("Submittal", sub), ("TM", TM), ("TM Date", TM_DATE),
                     ("Verdict", ver), ("Cycle", cycle), ("Key Observations", keyobs)]:
            rh.cell(dst, rhh[k]).value = v
        print(f"  RH new row {dst}: #{num} {code} cycle {cycle} -> {ver}")

    # 4) Summary recomputado
    stat = collections.Counter(); vd = collections.Counter(); total = 0
    for r in range(2, mr.max_row + 1):
        if mr.cell(r, C_CODE).value:
            total += 1
            su = str(mr.cell(r, C_STA).value or "").strip().upper()
            stat[su] += 1
            if su == "DELIVERED":
                vd[str(mr.cell(r, C_VER).value or "").strip()] += 1
    delivered = stat.get("DELIVERED", 0); not_delivered = stat.get("NOT DELIVERED", 0)
    partial = stat.get("PARTIAL", 0)
    print(f"  Summary recompute: total={total} delivered={delivered} not_delivered={not_delivered} "
          f"partial={partial} verdicts={dict(vd)}")
    real = {k: vd.get(k, 0) for k in TALLY_ESPERADO}
    if real != TALLY_ESPERADO:
        print(f"  🔴 TALLY DISTINTO DEL ESPERADO: esperado {TALLY_ESPERADO}, real {real}")
    else:
        print(f"  Tally coincide con lo declarado: {real}")
    sm["B4"].value = total; sm["B5"].value = delivered
    sm["B7"].value = partial; sm["B8"].value = not_delivered
    sm["B9"].value = "41 (N1 through N41 - TM N5 issued in Rev 0 and Rev 1)"
    sm["B10"].value = "100 (E1 through E100; the submittal series has no missing numbers; E96, E97, E99 and E100 pending review)"
    sm["B12"].value = "01-Oct-2026 (E100)"
    sm["B13"].value = "02-Oct-2026"
    def pct(n):
        return f"{round(100 * n / delivered)}%" if delivered else "0%"
    for cell, key in [("27", "1-Approved"), ("28", "2-AN"), ("29", "3-To be revised"),
                      ("30", "4-Rejected"), ("31", "No code issued")]:
        sm["B" + cell].value = vd.get(key, 0); sm["C" + cell].value = pct(vd.get(key, 0))
    # ITEMS BY SECTION: sin cambio (la fila #85 ya contaba como entregada desde el N40)
    _guardar_robusto(wb, SRC)


def _guardar_robusto(wb, dst):
    import zipfile, tempfile
    scratch_dir = os.environ.get("CLAUDE_SCRATCH") or tempfile.mkdtemp(prefix="register_n41_")
    os.makedirs(scratch_dir, exist_ok=True)
    tmp = os.path.join(scratch_dir, "register_n41_build.xlsx")
    wb.save(tmp)
    if zipfile.ZipFile(tmp).testzip() is not None:
        raise SystemExit("ERROR: el build en scratchpad salio corrupto.")
    openpyxl.load_workbook(tmp).close()
    print(f"Build verificado en scratchpad: {tmp}")
    for intento in range(1, 6):
        shutil.copyfile(tmp, dst)
        try:
            if zipfile.ZipFile(dst).testzip() is None:
                openpyxl.load_workbook(dst).close()
                print(f"Guardado y verificado en destino (intento {intento}): {dst}")
                return
        except Exception as e:
            print(f"  intento {intento}: destino aun corrupto ({e}); recopiando...")
    raise SystemExit(f"ERROR: no se pudo dejar copia integra. Build en {tmp}.")


if __name__ == "__main__":
    main()
