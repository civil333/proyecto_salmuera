#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para generar documento ADASA de consulta tecnica a Van Doorn.
Fecha: 28 de febrero de 2026
Codigo: P22-CT-06-000-001-0
Titulo: TECHNICAL QUERY — Control Logic Completeness Review
"""

import sys
import os

skill_path = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..", "..", ".claude", "skills", "template-adasa",
)
sys.path.insert(0, skill_path)

from ejemplo_documento import (
    crear_documento_adasa,
    aplicar_arial_12,
    add_simple_table,
)
from docx import Document
from docx.shared import Pt

OUTPUT = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "P22-CT-06-000-001-0_Logica-Control-Gaps_ADASA.docx",
)


def add_para(doc, text, size=11):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    return para


def crear_documento():
    crear_documento_adasa(
        titulo="TECHNICAL QUERY: CONTROL LOGIC COMPLETENESS REVIEW \u2014 P22-IT-06-008-101-B Rev B",
        codigo="P22-CT-06-000-001-0",
        preparado_por="Luis Rivera",
        revisado_por="Luis Rivera",
        aprobado_por="Victor Gutierrez",
        nombre_planta="TALTAL",
        cliente="ADASA",
        output_filename=OUTPUT,
        incluir_toc=True,
    )

    doc = Document(OUTPUT)

    # Eliminar contenido placeholder
    elementos_a_eliminar = []
    encontrado = False
    for para in doc.paragraphs:
        if para.style and para.style.name == "Heading 1" and not encontrado:
            encontrado = True
        if encontrado:
            elementos_a_eliminar.append(para)
    for para in elementos_a_eliminar:
        p = para._element
        p.getparent().remove(p)

    # ============================================================
    # HEADER INFO
    # ============================================================
    add_simple_table(doc, [
        ("Field", "Value"),
        ("Date", "February 28, 2026"),
        ("Project", "BAE 12803 \u2014 Second Stage RO Brine Module, Taltal"),
        ("From", "ADASA \u2014 Aguas de Antofagasta S.A. (Luis Rivera)"),
        ("To", "Van Doorn (Engineering Detail \u2014 Internal ADASA Consultant)"),
        ("Reference Document", "P22-IT-06-008-101-B, Rev B, 22-Feb-2026"),
        ("Status", "ISSUED"),
    ])

    # ============================================================
    # 1. BACKGROUND
    # ============================================================
    doc.add_heading("BACKGROUND", level=1)

    add_para(
        doc,
        "ADASA has reviewed the Control Logic Description P22-IT-06-008-101-B Rev B submitted "
        "in Engineering Delivery 1 (Nota de Env\u00edo N\u00b01). The review was conducted against the "
        "equivalent document P13-IT-03-008-001-0 (2019), which was the document used to "
        "successfully program, commission, and operate Module 3.",
    )
    add_para(
        doc,
        "Rev B correctly identifies the main permissives, the primary control loop "
        "(CTRL_BH06_001 \u2014 3 bar at PIT-06-001), and the BW Water interface signals. "
        "However, seven gaps were identified that must be resolved before the document can "
        "be approved for PLC programming.",
    )
    add_para(
        doc,
        "A formal internal evaluation (P22-IT-06-000-004-0) documents the full comparative "
        "analysis. This query summarizes the required actions for Rev C.",
    )

    # ============================================================
    # 2. REFERENCE DOCUMENTS
    # ============================================================
    doc.add_heading("REFERENCE DOCUMENTS", level=1)

    add_simple_table(doc, [
        ("#", "Code", "Document", "Rev", "Role"),
        ("1", "P22-IT-06-008-101-B",
         "Descriptivo L\u00f3gica de Control \u2014 2da Etapa Salmuera",
         "B (22-Feb-2026)", "Document under review"),
        ("2", "P13-IT-03-008-001-0",
         "Descriptivo L\u00f3gica de Control \u2014 M\u00f3dulo 3",
         "0 (2019)", "Benchmark"),
        ("3", "P22-IT-06-000-004-0",
         "Control Logic Review Evaluation (ADASA Internal)",
         "0", "Supporting analysis"),
    ])

    # ============================================================
    # 3. IDENTIFIED GAPS AND REQUIRED ACTIONS
    # ============================================================
    doc.add_heading("IDENTIFIED GAPS AND REQUIRED ACTIONS", level=1)

    # 3.1 GAP-01
    doc.add_heading("GAP-01 \u2014 Incomplete Start-up Sequence (MAJOR)", level=2)

    add_para(
        doc,
        "Rev B lists 16 start-up permissives but does not describe the step-by-step sequence "
        "once permissives are verified. Without this, BW Water\u2019s PLC programmer cannot "
        "implement the correct interlock sequence between ADASA\u2019s feeding system and the "
        "RO module without additional verbal clarification \u2014 a risk for programming rework "
        "after the first commissioning attempt.",
    )
    add_para(doc, "Action required for Rev C \u2014 add a numbered start-up sequence addressing:")
    add_simple_table(doc, [
        ("#", "Question to Resolve"),
        ("1", "When does BH-06-001 start and at what initial VFD frequency?"),
        ("2", "How long does the system wait to confirm pressure at PIT-06-001?"),
        ("3", "What happens if pressure is not reached within that time (alarm, shutdown, retry)?"),
        ("4", "When is start pulse OI-06-003 sent to BW Water \u2014 after pump confirms, or after pressure confirmed?"),
        ("5", "How does ADASA PLC confirm the OI module is running (via OI-06-004)?"),
        ("6", "What is the timeout for OI start confirmation before a fault is declared?"),
    ])

    # 3.2 GAP-02
    doc.add_heading("GAP-02 \u2014 Incomplete Shutdown Sequence + Tag Error (MAJOR)", level=2)

    add_para(
        doc,
        "Rev B contains 3 shutdown steps. Step 3 references \u201cParada Bomba de alimentaci\u00f3n "
        "BH-\u202203\u2022-001\u201d \u2014 an incorrect tag from Module 3. The correct tag is BH-06-001.",
    )
    add_para(doc, "Action required for Rev C:")
    add_para(doc, "\u2022  Correct tag error: replace BH-03-001 \u2192 BH-06-001 in shutdown step 3.")
    add_para(doc, "\u2022  Add missing shutdown steps:")
    add_simple_table(doc, [
        ("Missing Item", "Define"),
        ("BS-06-001 (drainage submersible pump)", "Stopped before or after BH-06-001? Sequence?"),
        ("VM-06-001 / VM-06-002 (manual valves)", "Final state \u2014 remain open or closed after shutdown?"),
        ("Instrumented valves", "Final state of all instrumented valves at end of shutdown"),
        ("OI-06-004 timeout", "Max wait time for BW Water OI shutdown confirmation before fault alarm"),
    ])

    # 3.3 GAP-03
    doc.add_heading("GAP-03 \u2014 No Level Band Table for TK-06-001 (MINOR)", level=2)

    add_para(
        doc,
        "Rev B defines >70% level as start-up permissive and \u201clow level\u201d as interlock, but "
        "does not provide a formal band table for the full operational level range. The P13 "
        "benchmark included 4 level bands with thresholds and associated equipment states.",
    )
    add_para(doc, "Action required for Rev C \u2014 add a level band table for TK-06-001:")
    add_simple_table(doc, [
        ("Band", "Level Range", "Equipment State", "Action"),
        ("High-High", "Define %", "Define", "e.g., OI stop pulse / alarm"),
        ("Normal High", "Define %", "Define", "Normal operation upper limit"),
        ("Normal Low", "Define %", "Define", "Normal operation lower limit"),
        ("Low-Low", "Define %", "Define", "BH-06-001 and OI emergency stop"),
    ])

    # 3.4 GAP-04
    doc.add_heading("GAP-04 \u2014 Post-Treatment Integration Not Documented (MINOR)", level=2)

    add_para(
        doc,
        "Rev B states post-treatment \u201calready operates and requires no action from the new PLC.\u201d "
        "Correct operationally, but three coordination points are undefined.",
    )
    add_para(doc, "Action required for Rev C \u2014 confirm and document:")
    add_simple_table(doc, [
        ("#", "Item"),
        ("1", "Is TK-00-001 high level the only coordination signal between the new module and existing post-treatment?"),
        ("2", "Are there inter-PLC signals between the new ADASA PLC and the existing plant PLC, or only a shared physical tank?"),
        ("3", "Is there a flow/production confirmation for the operator (e.g., FIT on discharge) to verify 25 m\u00b3/h is reaching post-treatment?"),
    ])

    # 3.5 GAP-05
    doc.add_heading("GAP-05 \u2014 Control Loop BH-06-001 Without Complete Definition (MINOR)", level=2)

    add_para(
        doc,
        "CTRL_BH06_001 is documented as maintaining 3 bar at PIT-06-001. Control type, VFD "
        "frequency limits, and failure behavior are absent.",
    )
    add_para(doc, "Action required for Rev C \u2014 complete the loop definition:")
    add_simple_table(doc, [
        ("Parameter", "Required Information"),
        ("Control type", "PI or PID?"),
        ("VFD frequency range", "Hz minimum / maximum"),
        ("Setpoint", "3 bar (confirm) or adjustable range?"),
        ("PIT-06-001 fault behavior", "What mode does BH-06-001 enter if pressure transmitter fails?"),
        ("Start ramp", "Frequency ramp profile for BH-06-001 start"),
    ])

    # 3.6 GAP-06
    doc.add_heading("GAP-06 \u2014 Fault Alarm Signal and Restart Procedure Not Documented (MINOR)", level=2)

    add_para(
        doc,
        "A direct review of the P13 benchmark signal section confirms that the total signal "
        "count between ADASA PLC and the OI module was 4 in both documents. The gap is not "
        "a missing signal count \u2014 it is two specific documentation omissions:",
    )
    add_para(
        doc,
        "\u2022  The fault alarm signal from BW Water OI to ADASA PLC (equivalent to OI-03-003 "
        "in P13) is not listed in the Rev B signal table. This signal triggers immediate plant "
        "shutdown and must appear in the signal list with its tag, direction (DI to ADASA PLC), "
        "and the interlock it activates.",
    )
    add_para(
        doc,
        "\u2022  The restart procedure after a BW Water OI fault is not documented. In the P13, "
        "Van Doorn explicitly stated that no dedicated reset signal to the OI module is required "
        "\u2014 the operator must normalize the fault condition and reset the system via ADASA HMI "
        "before the plant can restart. Rev C must include an equivalent statement to define "
        "the restart logic for the BW Water OI interface.",
    )
    add_para(doc, "Action required for Rev C:")
    add_simple_table(doc, [
        ("#", "Action"),
        ("1", "Add fault alarm signal (equivalent to OI-03-003) to the signal list with tag, type DI, and interlock description."),
        ("2", "Document restart procedure after OI fault: confirm no dedicated reset signal to BW Water is needed and that reset is via ADASA HMI."),
    ])

    # 3.7 GAP-07
    doc.add_heading("GAP-07 \u2014 Authorship Does Not Comply with ADASA Protocol (ADMINISTRATIVE)", level=2)

    add_para(
        doc,
        "Revisions A and B show Prepared/Reviewed/Approved = L. Hughes. Per ADASA document "
        "control protocol, formal ADASA-coded documents require ADASA authorship.",
    )
    add_para(doc, "Action required for Rev C \u2014 title block must show:")
    add_simple_table(doc, [
        ("Field", "Value"),
        ("Prepared By", "Luis Rivera"),
        ("Reviewed By", "Luis Rivera"),
        ("Approved By", "Victor Gutierrez"),
    ])

    # ============================================================
    # 4. SUMMARY
    # ============================================================
    doc.add_heading("SUMMARY OF REQUIRED ACTIONS", level=1)

    add_simple_table(doc, [
        ("Gap", "Severity", "Action", "Rev C Section"),
        ("GAP-01", "Major", "Add numbered start-up sequence (6 items)", "Start-up"),
        ("GAP-02", "Major", "Correct BH-03 \u2192 BH-06 tag; complete shutdown (4 items)", "Shutdown"),
        ("GAP-03", "Minor", "Add level band table TK-06-001 (4 bands)", "Level control"),
        ("GAP-04", "Minor", "Document post-treatment coordination (3 items)", "System integration"),
        ("GAP-05", "Minor", "Complete CTRL_BH06_001 definition (5 parameters)", "Control loop"),
        ("GAP-06", "Minor", "Add fault alarm signal to signal list; document restart procedure after OI fault via HMI", "Signals"),
        ("GAP-07", "Admin", "Correct authorship in title block", "Title block"),
    ])

    # ============================================================
    # 5. RESPONSE REQUESTED
    # ============================================================
    doc.add_heading("RESPONSE REQUESTED", level=1)

    add_para(
        doc,
        "Please provide Rev C of P22-IT-06-008-101-B addressing all items above.",
    )
    add_para(
        doc,
        "Suggested deadline: 2 weeks from date of this query (target: March 14, 2026).",
    )
    add_para(
        doc,
        "Rev C will be reviewed by ADASA. If all gaps are resolved, the document will be "
        "formally approved for PLC programming by BW Water.",
    )

    # ============================================================
    # 6. DOCUMENT HISTORY
    # ============================================================
    doc.add_heading("DOCUMENT HISTORY", level=1)

    add_simple_table(doc, [
        ("Rev", "Date", "Description", "Author"),
        ("0", "28-Feb-2026", "Initial issue", "L. Rivera"),
    ])

    doc.save(OUTPUT)
    print(f"Documento generado: {OUTPUT}")


if __name__ == "__main__":
    crear_documento()
