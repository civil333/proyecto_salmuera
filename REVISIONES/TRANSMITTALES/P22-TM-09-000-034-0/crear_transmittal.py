#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TRANSMITTAL N34 ADASA-BW_WATER.
Submittals 25007-0073 (E73), 25007-0077 (E77) y 25007-0080 (E80).
Fecha de emision: 18-Ago-2026.

Veredicto global: 3 - TO BE REVISED.
Tally: 2 Code 1 + 2 Code 2 + 1 Code 3. El codigo lo fija UN solo documento.

  2.1 Alarm and Interlock List Rev 0      (P22-LI-09-008-015)  Code 1
  2.2 Control and Sequence Chart Rev 0    (P22-LI-09-008-017)  Code 1
  2.3 Quality Dossier Index Rev B         (P22-BA-09-000-013)  Code 3
  2.4 GA of Antiscalant Dosing Pump Skid Rev C (P22-DWG-09-005-011)  Code 2
  2.5 GA of CIP / Flushing Tank Rev B     (P22-DWG-09-005-014)  Code 2

REGLA DE ALCANCE. La revision de un documento que responde a comentarios previos
se limita a SI ESOS COMENTARIOS SE LEVANTARON EN FORMA. No se introducen
observaciones nuevas. Universo por documento: TM N28 para los dos de la E73,
TM N30 para el dossier, TM N26 para los dos planos.

EXCEPCION APROBADA POR EL USUARIO: el Quality Dossier Index Rev B se revisa a
fondo contra el ITP, el PIE Base y la ET, porque su funcion es ser el checklist
del dossier y un capitulo que falta deja sin ubicacion un registro que gatilla el
40% del pago.

CORRECCION DE VEREDICTO (peticion del usuario). La primera version dispuso la
Alarm and Interlock List Rev 0 en Code 3 y fue over-reach. Ante un Rev 0 la
pregunta es si el defecto es INTRINSECO al documento o VIVE EN OTRO
(feedback_rev0_solo_code1_o_code3). Los dos pendientes caen del lado que no
degrada: el rango de vibracion vive en la Instrument List -> Seccion 3, y los dos
TAG sin guion son housekeeping documental -> una linea del bloque de accion. Su
CC_ADASA se retiro a _analisis_no_anotado/.

SE RETIRA ademas la NOTE-01 del plano del estanque CIP (la caratula declara dos
paginas y trae tres): observacion NUEVA, fuera del universo del TM N26.

ADASA DECLARA EL VALOR VINCULANTE del rango de VT-09-001 en 0 a 12 mm/s rms, en
vez de preguntar cual de los dos documentos gobierna. Regla del proyecto para un
dato que dos documentos declaran distinto, con fabricacion en curso.

SIN numeros manuales en add_heading() (el template ADASA auto-numera H1/H2).
Referencias por NOMBRE de seccion. Cero simbolo de seccion.
"""

import os
import sys

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))

from ejemplo_documento import (  # noqa: E402
    crear_documento_adasa,
    aplicar_arial_12,
    add_simple_table,
    add_bullet as _add_bullet_native,
)
from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "TRANSMITTAL N34 ADASA-BW_WATER.docx")

# Enlace de descarga de la carpeta del N34, publicado el 18-Ago-2026.
# El modo de falla del N30, N31 y N32 fue verificar el enlace sobre el script y
# no sobre el texto extraido del Word. Verificar SIEMPRE sobre el .docx emitido.
DOWNLOAD_LINK = (
    "https://lrg.synology.me:6501/d/s/19WVqhVU8zwy31naz76IwRhUhCks28r6/"
    "tA3kHXW9cHDSvxetipsYW9ruR-Mk9hrr-9LBA_RPabw0"
)


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")
    rpr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_para(doc, runs):
    para = doc.add_paragraph()
    for item in runs:
        text = item[0]
        attrs = item[1] if len(item) > 1 else {}
        run = para.add_run(text)
        if attrs.get("bold"):
            run.bold = True
        if attrs.get("italic"):
            run.italic = True
    aplicar_arial_12(para)
    return para


def add_bullet(doc, text):
    _add_bullet_native(doc, text, size=11, space_after_pt=12)


# --- Seccion 1 -------------------------------------------------------------
VERDICT = (
    "Five documents from submittals 25007-0073 of 11 August, 25007-0077 of 14 August "
    "and 25007-0080 of 18 August. Tally: 2 Code 1, 2 Code 2, 1 Code 3."
)

SCOPE = (
    "Each document was reviewed against the comments ADASA raised on its previous "
    "revision, and against nothing else."
)

DISPOSITION = [
    "Alarm and Interlock List Rev 0 — Code 1. Re-issue the Instrument List at the "
    "binding vibration range ADASA declares below.",
    "Control and Sequence Chart Rev 0 — Code 1. No action.",
    "Quality Dossier Index Rev B — Code 3. Add the chapter for ADASA's FAT Approval "
    "Certificate, complete the revision and inclusion status of every line, and state "
    "which of the two dossier indices this document is.",
    "GA of Antiscalant Dosing Pump Skid Rev C — Code 2. Reconcile the per-bolt Fz with "
    "the total reaction block at Rev 0.",
    "GA of CIP / Flushing Tank Rev B — Code 2. State which document governs the top "
    "opening and the two side connections, and align the losing one.",
]

WHY_CODE3 = (
    "The chapter reported as the FAT Approval Certificate is the release for dispatch, a "
    "different record; the certificate the Technical Specification calls indispensable to "
    "the final dossier has no chapter. And the observation that set the previous code is "
    "not materially closed: the new columns are largely empty and no line carries an "
    "inclusion status, so the index still cannot work as a checklist. It is the only "
    "document that fixes the transmittal code."
)

BINDING_RANGE = (
    "The Alarm and Interlock List Rev 0 ranges VT-09-001 at 0 to 12 mm/s rms with the "
    "high-high trip at 10, while the Instrument List Rev E, approved at Code 1, still "
    "ranges the same tag at 0 to 8.9 mm/s rms. ADASA declares the binding range to be 0 "
    "to 12 mm/s rms and requires the Instrument List to be re-issued to it, so that the "
    "vibration stop of a 93 kW pump can act."
)

RECEIPT = (
    "Transmittal N33 reported both numbers as missing from the series. Both exist: the "
    "0073 reached ADASA on 11 August and the 0077 on the morning of 18 August, one day "
    "after the return date printed on its own form. Under Clause 37.2 the review period "
    "runs from formal and complete receipt, which sets the deadlines at Thursday 20 "
    "August for 25007-0073 and Thursday 27 August for the other two. This transmittal is "
    "issued within all three. Please route future submittals so that the issue date and "
    "the receipt date coincide."
)

# --- Seccion 2 -------------------------------------------------------------
S21_STATUS = (
    "Reviewed against the four points Transmittal N28 set for this issue, and against "
    "nothing else. All four are answered in the body of the document. The breaker alarm "
    "resolves as P22-PLC01-XA001 on the breaker open contact. The vibration tag is "
    "unified to VT-09-001 and the four fault alarms of the note are added, for the two "
    "valves and the two dosing pumps. The four housekeeping items are corrected, "
    "including the turbocharger boost-failure differential now stated at 10 bar and the "
    "CIP tank temperature range at 0 to 100 degrees Celsius."
)

S21_COHERENT = (
    " item 6.0 ranges VT-09-001 at 0 to 12 mm/s rms and item 6.1 sets the high-high trip "
    "at 10, inside that range. What remains is on another document, and it is tracked in "
    "Pending Observations from Previous Transmittals."
)

S21_ACTION = (
    " Related deliverable tracked in Pending Observations from Previous Transmittals: "
    "re-issue of the Instrument List (P22-LI-09-008-003) at the binding range of 0 to 12 "
    "mm/s rms for VT-09-001. Two of the added fault tags read VE09-014 and VE09-016, "
    "without the hyphen the IO List Rev 5 and the thirteen other valve rows of this list "
    "both use; that is housekeeping for the next natural issue and affects no setpoint."
)

S22_STATUS = (
    "The seven points Transmittal N28 set for this issue are closed, each verified "
    "against the document that governs it. Note 5 now assigns the turbocharger bypass to "
    "the pressure control loop on PIT-09-005, with TDS selecting only the setpoint. The "
    "second-stage abort reads 93 bar and the high-pressure pump ramps read 0.1 to 0.3 Hz "
    "per second. The new Note 9 states the flushing setpoints per stage, 48 and 36 cubic "
    "metres per hour under FIT-09-005, and the CIP return valves carry the stage "
    "assignment of the IO List. The setpoint and formula corrections are in place, "
    "including the brine flow expression, which the comment sheet does not claim."
)

S22_ACTION = (
    " ADASA notes that the two 1 Hz per second ramps remaining in the CIP operation "
    "sequence belong to the CIP and flushing pumps, for which the Control Philosophy sets "
    "no ramp limit; the 0.1 to 0.3 Hz per second range applies to the high-pressure feed "
    "pump and is correctly reflected. No annotated PDF accompanies this document."
)

S23_STATUS = (
    "The index grows from 27 chapters to 49 and closes five of the ten points. Three are "
    "claimed: the dispatch chapters against rows 8.1 and 8.2, the RO pressure vessel "
    "package at D6, and the equipment-by-equipment vendor records across Sections D and "
    "E. Two are not: the inspection personnel qualifications folded into B3 to B6 with "
    "the calibration certificate at C18, and the non-conformance and weld repair chapters "
    "at C19 and C20. What remains is set out in the Executive Summary. Itemised in "
    "P22-BA-09-000-013_B_Quality_Dossier_Index_CC_ADASA.pdf."
)

S23_NOT_THE_DOSSIER = (
    " The Fabrication and Testing Dossier required by the Technical Specification "
    "(P22-ET-09-000-001-0), Section 7 - Documentation, has not been delivered. This code "
    "does not reach it, and it continues to gate items 8.3 and 8.4 of the Inspection and "
    "Testing Base Plan."
)

S23_ACTION_LEAD = "Three items govern."
S23_ACTION = (
    " Add a chapter for the FAT Approval Certificate issued by ADASA, distinct from the "
    "release for dispatch already at C21. Complete the document number, revision and "
    "inclusion status on every line, and the Inspection and Test Plan row on Sections A, "
    "B, D and E. State whether this document is the preliminary index of row 7.6 or the "
    "final index of row 8.3, which the plan names separately and holds at different "
    "levels. Two further items: add the packing list to the dispatch chapter, and "
    "identify in C1, spool by spool, the mill certificates of the super duplex material "
    "with the PREN verification of row 2.1 (OBS-01 to OBS-03 and NOTE-01 to NOTE-02 on "
    "the annotated PDF)."
)

S24_STATUS = (
    "Rev C answers at an intermediate revision Transmittal N26 had not required. The "
    "labelling half of the item closes: Fx and Fy now equal the totals over the ten bolts "
    "and each block is identified as a total or a per-bolt value. The reconciliation half "
    "does not, for the second consecutive issue and against a comment sheet that reports "
    "the forces as updated: the sheet states a total Fz of 1.0242 kN over ten bolts and a "
    "per-bolt Fz of 0.205 kN, twice the quotient, where at Rev B the same pair differed "
    "by a factor of ten. Itemised in "
    "P22-DWG-09-005-011_C_GA_Antiscalant_Pump_Skid_CC_ADASA.pdf."
)

S24_ENDORSEMENT = (
    " The note calls the forces \"as per calculation report\" and Note 5 of the sheet "
    "still reads that the bolting details are to be finalised and endorsed. ADASA's "
    "letter of 18 August addresses the professional endorsement of the structural "
    "calculation report and sets a date for it; this item closes against that report."
)

S24_ACTION = (
    " state which of the two Fz figures governs and reconcile the per-bolt value with the "
    "total over the ten bolts, keeping the axis of each figure explicit (OBS-01 on the "
    "annotated PDF). ADASA accepts the drawing on the basis that this is a figure and its "
    "label, with no change to the anchorage arrangement."
)

S25_STATUS = (
    "Of the two items carried from Transmittal N26 the equipment tag closes: Rev A "
    "carried TK-09-001 nowhere on the sheet and Rev B carries it. The nozzle schedule "
    "does not. The comment sheet reports it reconciled with the datasheet, and the three "
    "entries in dispute are unchanged from Rev A: the top opening reads MH Manhole 533 mm "
    "internal diameter where the approved Datasheet of the CIP Tank (P22-ET-09-009-009) "
    "Rev B lists HH Handhole DN300, and the side connections N42 and N97 have no "
    "counterpart there. Neither document was re-issued; the twelve remaining nozzles "
    "agree. Itemised in P22-DWG-09-005-014_B_GA_CIP_Flushing_Tank_CC_ADASA.pdf."
)

S25_GOVERNING = (
    "The request was to state which document governs, and it stands: the datasheet is not "
    "self-consistent either, since its tank construction description calls for a welded "
    "conical cover with a manhole cover while its nozzle schedule lists a handhole."
)

S25_ACTION = (
    " state in writing whether the general arrangement or the datasheet governs the top "
    "opening and the two side connections, and align the other document accordingly, "
    "re-issuing the Datasheet of the CIP Tank if that is the one that changes (OBS-01 on "
    "the annotated PDF)."
)

# --- Seccion 3 -------------------------------------------------------------
PENDING = [
    ("N25, N29, N30", "Fabrication and testing dossier",
     "Item 65 remains NOT DELIVERED. The index has now reached Rev B; no record has "
     "followed it. It sustains items 8.3 and 8.4 of the Inspection and Testing Base "
     "Plan, on which 40 per cent of payment depends", "Open, overdue"),
    ("N26, N30", "Endorsed structural calculation report",
     "The report governs the anchorage figures of two general arrangement drawings and "
     "the foundations already built at site. Addressed in ADASA's letter of 18 August",
     "Open, overdue"),
    ("N32", "Three non-destructive testing procedures Rev A",
     "None states the ASME B31.3 acceptance criteria required by the Technical "
     "Specification (P22-ET-09-000-001-0), Section 8 - Inspections During Manufacturing",
     "Open, Code 3"),
    ("N28, N34", "Instrument List (P22-LI-09-008-003) Rev E",
     "Re-issue with VT-09-001 ranged at the binding 0 to 12 mm/s rms, so that the "
     "high-pressure pump vibration trip of 10 mm/s carried by the Alarm and Interlock "
     "List Rev 0 can act. Rev E still reads 0 to 8.9 mm/s rms", "Open, raised here"),
]

ALSO_OPEN = (
    "the liquid penetrant records of 7 August, examined five days before their procedure "
    "was submitted, and the six tag reconciliation of the HMI Display Screenshot Rev B, "
    "which carries the temperature sensor mapping of the high-pressure pump. Also the "
    "design pressure at the brine feed tie-in point, and the calibration validity of "
    "certificate 26993 at the date of the pressure test."
)

# --- Seccion 4 -------------------------------------------------------------
ATTACHMENTS = [
    ("Quality Dossier Index Rev B", "3",
     "P22-BA-09-000-013_B_Quality_Dossier_Index_CC_ADASA.pdf"),
    ("GA of Antiscalant Dosing Pump Skid Rev C", "2",
     "P22-DWG-09-005-011_C_GA_Antiscalant_Pump_Skid_CC_ADASA.pdf"),
    ("GA of CIP / Flushing Tank Rev B", "2",
     "P22-DWG-09-005-014_B_GA_CIP_Flushing_Tank_CC_ADASA.pdf"),
]

ATTACH_NOTE = (
    "All documents with open observations carry annotated PDFs, one Code 3 and two Code "
    "2. The two Code 1 documents, the Alarm and Interlock List and the Control and "
    "Sequence Chart, carry none."
)

# --- Seccion 5 -------------------------------------------------------------
RESPONSE_SUMMARY = [
    ("P22-LI-09-008-015", "Alarm and Interlock List", "0", "25007-0073", "1 — Approved"),
    ("P22-LI-09-008-017", "Control and Sequence Chart", "0", "25007-0073",
     "1 — Approved"),
    ("P22-BA-09-000-013", "Quality Dossier Index", "B", "25007-0077",
     "3 — To be revised"),
    ("P22-DWG-09-005-011", "GA of Antiscalant Dosing Pump Skid", "C", "25007-0080",
     "2 — Approved as noted"),
    ("P22-DWG-09-005-014", "GA of CIP / Flushing Tank", "B", "25007-0080",
     "2 — Approved as noted"),
]

CLOSING = (
    " The code is set by one document. The Quality Dossier Index needs the chapter for "
    "ADASA's FAT Approval Certificate, the inclusion status of its lines, and a statement "
    "of which of the two dossier indices it is. The two documents issued at Rev 0 are "
    "approved as they stand; what remains on them lives in other documents and is tracked "
    "in Pending Observations from Previous Transmittals."
)


if __name__ == "__main__":
    crear_documento_adasa(
        titulo=("TECHNICAL REVIEW TRANSMITTAL N34 — SECOND STAGE RO "
                "BRINE MODULE"),
        codigo="P22-TM-09-000-034-0",
        output_filename=OUTPUT,
        incluir_toc=True,
    )

    doc = Document(OUTPUT)

    # 1. EXECUTIVE SUMMARY
    doc.add_heading("EXECUTIVE SUMMARY", level=1)
    add_para(doc, [
        ("TRANSMITTAL VERDICT: 3 — To be revised.", {"bold": True}),
        (" " + VERDICT,),
    ])
    add_para(doc, [("Scope of this review.", {"bold": True}), (" " + SCOPE,)])
    add_para(doc, [("Disposition at a glance:", {"bold": True})])
    for item in DISPOSITION:
        add_bullet(doc, item)
    add_para(doc, [
        ("Why Code 3 — Quality Dossier Index.", {"bold": True}),
        (" " + WHY_CODE3,),
    ])
    add_para(doc, [
        ("Binding vibration range.", {"bold": True}), (" " + BINDING_RANGE,),
    ])
    add_para(doc, [
        ("Receipt of submittals 25007-0073 and 25007-0077.", {"bold": True}),
        (" " + RECEIPT,),
    ])
    add_para(doc, [
        ("Pending Observations from Previous Transmittals lists the items open from "
         "earlier transmittals.",),
    ])

    # 2. OBSERVATIONS BY DOCUMENT
    doc.add_heading("OBSERVATIONS BY DOCUMENT", level=1)

    doc.add_heading("Alarm and Interlock List Rev 0 — P22-LI-09-008-015", level=2)
    add_para(doc, [("Response Code: 1 — Approved", {"bold": True})])
    add_para(doc, [("Status.", {"bold": True}), (" " + S21_STATUS,)])
    add_para(doc, [
        ("On the vibration trip, this list is internally coherent:", {"bold": True}),
        (S21_COHERENT,),
    ])
    add_para(doc, [
        ("Action: none on this document — accepted.", {"bold": True}),
        (S21_ACTION,),
    ])

    doc.add_heading("Control and Sequence Chart Rev 0 — P22-LI-09-008-017", level=2)
    add_para(doc, [("Response Code: 1 — Approved", {"bold": True})])
    add_para(doc, [("Status.", {"bold": True}), (" " + S22_STATUS,)])
    add_para(doc, [("Action: none — accepted.", {"bold": True}), (S22_ACTION,)])

    doc.add_heading("Quality Dossier Index Rev B — P22-BA-09-000-013", level=2)
    add_para(doc, [("Response Code: 3 — To be revised", {"bold": True})])
    add_para(doc, [("Status.", {"bold": True}), (" " + S23_STATUS,)])
    add_para(doc, [
        ("The index is not the dossier.", {"bold": True}), (S23_NOT_THE_DOSSIER,),
    ])
    add_para(doc, [
        ("Action — re-issue as Rev C. ", {"bold": True}),
        (S23_ACTION_LEAD + S23_ACTION,),
    ])

    doc.add_heading(
        "GA of Antiscalant Dosing Pump Skid Rev C — P22-DWG-09-005-011", level=2)
    add_para(doc, [("Response Code: 2 — Approved as noted", {"bold": True})])
    add_para(doc, [("Status.", {"bold": True}), (" " + S24_STATUS,)])
    add_para(doc, [
        ("The reference this item points to does not yet exist in endorsed form.",
         {"bold": True}),
        (S24_ENDORSEMENT,),
    ])
    add_para(doc, [
        ("Action to issue at IFC Rev 0 — no new drawing revision required:",
         {"bold": True}),
        (S24_ACTION,),
    ])

    doc.add_heading("GA of CIP / Flushing Tank Rev B — P22-DWG-09-005-014", level=2)
    add_para(doc, [("Response Code: 2 — Approved as noted", {"bold": True})])
    add_para(doc, [("Status.", {"bold": True}), (" " + S25_STATUS,)])
    add_para(doc, [(S25_GOVERNING,)])
    add_para(doc, [
        ("Action to issue at IFC Rev 0 — no new drawing revision required:",
         {"bold": True}),
        (S25_ACTION,),
    ])

    # 3. PENDING OBSERVATIONS FROM PREVIOUS TRANSMITTALS
    doc.add_heading("PENDING OBSERVATIONS FROM PREVIOUS TRANSMITTALS", level=1)
    add_simple_table(
        doc,
        [("Origin TM", "Document", "Observation", "Status")] + PENDING)
    add_para(doc, [("Also open: ", {"bold": True}), (ALSO_OPEN,)])

    # 4. ATTACHMENTS
    doc.add_heading("ATTACHMENTS", level=1)
    add_simple_table(
        doc,
        [("Document", "Code", "Annotated PDF")] + ATTACHMENTS)
    add_para(doc, [(ATTACH_NOTE,)])
    dl = doc.add_paragraph()
    dl.add_run("Download — this transmittal and the annotated PDFs: ").bold = True
    aplicar_arial_12(dl)
    add_hyperlink(dl, DOWNLOAD_LINK, DOWNLOAD_LINK)

    # 5. RESPONSE SUMMARY
    doc.add_heading("RESPONSE SUMMARY", level=1)
    add_simple_table(
        doc,
        [("Document Code", "Title", "Rev", "Submittal", "Response Code")]
        + RESPONSE_SUMMARY)
    add_para(doc, [
        ("Overall Transmittal Verdict: 3 — TO BE REVISED.", {"bold": True}),
        (CLOSING,),
    ])

    doc.save(OUTPUT)
    print(f"Documento generado: {OUTPUT}")
