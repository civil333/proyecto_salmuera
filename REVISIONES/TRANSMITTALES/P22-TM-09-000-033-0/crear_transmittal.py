#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TRANSMITTAL N33 ADASA-BW_WATER.
Submittals 25007-0076 (E76), 25007-0078 (E78) y 25007-0079 (E79).
Fecha de emision: 17-Ago-2026.

Veredicto global: 2 - APPROVED AS NOTED.
Tally: 4 Code 1 + 1 Code 2. Ningun documento vuelve a revision.

  2.1 Civil and Loading Layout Rev B     (P22-DWG-09-005-001)  Code 2
  2.2 GA of CIP Flushing Skid Pump Rev B (P22-DWG-09-005-010)  Code 1
  2.3 Equipment Layout Rev D             (P22-DWG-09-005-003)  Code 1
  2.4 Process Flow Diagram Rev 0         (P22-DWG-09-009-01)   Code 1
  2.5 Control Philosophy Rev 1           (P22-BT-09-009-001)   Code 1

REGLA DE ALCANCE (criterio del usuario, 17-Ago-2026). La revision de un
documento que responde a comentarios previos se limita a SI ESOS COMENTARIOS
ESTAN CERRADOS. NO se introducen observaciones nuevas. Dos razones: la ingenieria
de obras civiles la ejecuta L&A y lo que sale de los planos de BW Water se ve con
ellos, no con el proveedor; y no hay tiempo para mas iteraciones de revision. Un
documento en Rev 1, que existe porque ADASA observo que un comentario de la Rev 0
no se habia levantado bien, no puede llegar con comentarios nuevos.

EL ALCANCE LO FIJA LA COLUMNA DE COMENTARIO DEL CLIENTE de la Consolidated
Comment Sheet, que trae el texto literal del pedido original. Leerla cambio dos
veredictos y corrigio un error propio:
  - El pliego del plano civil nombra literal "together with SKID FRAME", de modo
    que el marco del skid SI estaba en el pedido: es cierre pendiente, no
    comentario nuevo. El ledger de la E76 afirmaba lo contrario.
  - El Grounding Layout Rev F rotula el panel "LCP" en la propia lamina, mismo
    identificador que la fila 14 del Equipment Layout, de modo que la OBS-03 del
    TM N22 SI cerro. El P22-LCP-01 salia del texto de un comentario de ADASA.
  - El plano de la bomba cierra sus TRES partes aunque su respuesta reclame dos:
    la masa de 180 kg esta en la Nota 3. El documento manda sobre la columna de
    respuesta.

SE RETIRAN del transmittal, por ser comentarios nuevos (quedan en los ledgers
internos y, los de obra civil, en _IMPACTO_OOCC_BL.md para la conversacion con
L&A): diametros de anclaje M12 y M10, empotramiento de los plintos 5 y 7,
escalon de 22 mm, reparto de las seis filas nuevas, tabla de pesos duplicada y
el peso del panel de control, TAG del mezclador estatico, peso de operacion del
estanque CIP, factor sismico declarado, titulo del bloque de reacciones, vista de
empotramiento, informe de calculo sin identificar, codigos mal formados de la
tabla de referencias, las tres referencias del cuerpo sin fijar, el calificativo
"subject to vendor confirmation", y el MAPEO DE SENSORES DE LA BOMBA DE ALTA.

Sobre el mapeo de sensores (decision del usuario): sale del transmittal porque el
punto 1 del TM N31 declaro correcta esa pagina, de modo que objetarla ahora es
comentario nuevo sobre una Rev 1. Queda con su evidencia en el ledger de la E79.
El vehiculo para corregirlo sin abrir instrumento nuevo es el pedido de
reconciliar seis tags que el N31 dejo abierto al HMI.

SIN numeros manuales en add_heading() (el template ADASA auto-numera H1/H2).
Referencias por NOMBRE de seccion. Cero simbolo de seccion.
"""

import os
import sys

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))

from ejemplo_documento import (  # noqa: E402
    crear_documento_adasa,
    aplicar_arial_12,
    add_simple_table,
    add_bullet as _add_bullet_native,
)
from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "TRANSMITTAL N33 ADASA-BW_WATER.docx")

# El unico PDF anotado va por enlace de descarga. Enlace publicado el 17-Ago-2026.
DOWNLOAD_LINK = (
    "https://lrg.synology.me:6501/d/s/19Vh5BdPy7Bawfney5aMrUV8sZlaFy99/"
    "hGVtoOjNlRe8KS3GBP2Ek-22F4fX0UUP-4rJA83s3bw0"
)


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")
    rpr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_para(doc, runs):
    para = doc.add_paragraph()
    for item in runs:
        text = item[0]
        attrs = item[1] if len(item) > 1 else {}
        run = para.add_run(text)
        if attrs.get("bold"):
            run.bold = True
        if attrs.get("italic"):
            run.italic = True
    aplicar_arial_12(para)
    return para


def add_bullet(doc, text):
    _add_bullet_native(doc, text, size=11, space_after_pt=12)


# --- Seccion 1 -------------------------------------------------------------
VERDICT = (
    "Five documents from submittals 25007-0076 of 13 August, 25007-0078 and 25007-0079 "
    "of 17 August. Tally: 4 Code 1, 1 Code 2. No document is returned to revision."
)

SCOPE = (
    "Each document was reviewed against the comments ADASA raised on its previous "
    "revision, and against nothing else."
)

DISPOSITION = [
    "Civil and Loading Layout Rev B — Code 2. State the empty container weight and the "
    "row carrying the RO Skid frame.",
    "GA of CIP Flushing Skid Pump Rev B — Code 1. No action.",
    "Equipment Layout Rev D — Code 1. No action.",
    "Process Flow Diagram Rev 0 — Code 1. No action.",
    "Control Philosophy Rev 1 — Code 1. No action.",
]

WHY_CODE2 = (
    "Neither part of the Transmittal N16 note is closed. Note 2 gives 13,300 kg as the "
    "container and its contents, and that figure excludes both the container and all "
    "fluid inventory; the operating total is 30,529.5 kg. And the RO Skid frame, named in "
    "that note, sits in no row after rows 6 and 7 were reclassified to RO Pressure "
    "Vessels."
)

MISSING = (
    "The series runs 0076, 0078, 0079. Please confirm whether that number was issued and "
    "to whom. If it was issued and did not reach ADASA, formal and complete receipt never "
    "took place, so the seven working days of Clause 37.2 have not started to run for it "
    "and its content is not covered here."
)

REVIEW_PERIOD = (
    "Clause 37.2 sets ADASA's review at seven working days from formal and complete "
    "receipt: Monday 24 August for 25007-0076 and Wednesday 26 August for the other two. "
    "This transmittal is issued within all three, and not on the Submittal Form return "
    "dates, one of which falls on a Sunday."
)

# --- Seccion 2 -------------------------------------------------------------
S21_STATUS = (
    "Rev B answers what is counted separately, with two new notes and six new rows, and "
    "does so at an intermediate revision ADASA had not required. The two items of the "
    "Transmittal N16 note remain, as stated above. Itemised in "
    "P22-DWG-09-005-001_B_Civil_Loading_Layout_CC_ADASA.pdf."
)

S21_ACTION = (
    "state the weight of the empty modified container and reword Note 2 accordingly, and "
    "state the row carrying the RO Skid structural frame (OBS-01 and OBS-02 on the "
    "annotated PDF)."
)

S22_STATUS = (
    "The three parts of the Transmittal N11 note are closed: the bolting arrangement with "
    "diameter, spacing and embedment; the seismic base reactions; and the equipment mass "
    "with its centre of gravity, in Notes 3 and 8.1, which the written reply does not "
    "claim."
)

S23_STATUS = (
    "The three observations of Transmittal N22 are closed: the RO cartridge filter is "
    "drawn vertical, verified by ADASA on a high resolution rendering of the sheet; the "
    "Operating Weight table is embedded; and the main panel is labelled LCP, consistent "
    "with the Grounding Point and Power Panel Location Layout (P22-DWG-09-007-003) Rev F."
)

S24_STATUS = (
    "Rev B was approved at Transmittal N3 with no open comments, so the Rev 0 issue was "
    "reviewed against what was approved. Every capacity, flow, tag and material is "
    "unchanged. One figure changed and it is a correction: the high-pressure pump duty now "
    "reads 46.9 bar, the differential head of the governing Datasheet of RO HP Feed Pump "
    "(P22-ET-09-009-002) Rev D, where Rev B read 49.4 bar."
)

S24_REGISTER = (
    "ADASA has aligned its own register to this document's code, P22-DWG-09-009-01, which "
    "the document has carried since Rev A."
)

S25_STATUS = (
    "The four points Transmittal N31 set for closure at this issue are closed, each "
    "verified against the governing document: the CIP pump sensors read TE-09-003 winding "
    "and TE-09-004 bearing; the vibration alarm and trip pairs match the Alarm and "
    "Interlock List Rev C on all three machines; the winding trip is supported by the "
    "Class F insulation of the governing pump datasheet; and the reference table pins both "
    "child documents by code."
)

S25_ACTION = (
    "Adding the revision of those two child documents, P22-LI-09-008-017 Rev A and "
    "P22-LI-09-008-015 Rev C, is housekeeping for the next natural issue."
)

# --- Seccion 3 -------------------------------------------------------------
PENDING = [
    ("N25, N29", "Fabrication and testing dossier",
     "Item 65 remains NOT DELIVERED. The index arrived at Transmittal N30; no record has "
     "followed. It sustains PIE items 8.3 and 8.4, on which 40 per cent of payment "
     "depends",
     "Open, overdue"),
    ("N32", "Liquid penetrant records of 7 August",
     "The examination was carried out five days before its procedure was submitted. State "
     "which records exist and under which procedure and acceptance criteria",
     "Open, no date committed"),
    ("N32", "Three non-destructive testing procedures Rev A",
     "None states the ASME B31.3 acceptance criteria required by the Technical "
     "Specification (P22-ET-09-000-001-0), Section 8 - Inspections During Manufacturing",
     "Open, Code 3"),
]

ALSO_OPEN = (
    "the endorsed structural calculation report for the two general arrangement drawings; "
    "the six tag reconciliation of the HMI Display Screenshot Rev B; the design pressure "
    "at the brine feed tie-in point; the calibration validity of certificate 26993 at the "
    "date of the pressure test; and the Alarm and Interlock List issued at Rev 0 on 11 "
    "August, not yet reviewed and not covered here."
)

# --- Seccion 4 -------------------------------------------------------------
ATTACHMENTS = [
    ("Civil and Loading Layout Rev B", "2",
     "P22-DWG-09-005-001_B_Civil_Loading_Layout_CC_ADASA.pdf"),
]

ATTACH_NOTE = (
    "The Civil and Loading Layout is the only document with open observations and the only "
    "one carrying an annotated PDF. The other four are Code 1 — Approved."
)

# --- Seccion 5 -------------------------------------------------------------
RESPONSE_SUMMARY = [
    ("P22-DWG-09-005-001", "Civil and Loading Layout", "B", "25007-0076",
     "2 — Approved as noted"),
    ("P22-DWG-09-005-010", "GA of CIP Flushing Skid Pump", "B", "25007-0076",
     "1 — Approved"),
    ("P22-DWG-09-005-003", "Equipment Layout", "D", "25007-0078", "1 — Approved"),
    ("P22-DWG-09-009-01", "Process Flow Diagram", "0", "25007-0078", "1 — Approved"),
    ("P22-BT-09-009-001", "Control Philosophy", "1", "25007-0079", "1 — Approved"),
]

CLOSING = (
    " No document requires a new revision. The code is set by the two items of the "
    "Transmittal N16 note on the Civil and Loading Layout, both to be incorporated at the "
    "Rev 0 issue."
)


def main():
    # SIN tabla de contenidos, por decision del usuario del 17-Ago-2026: el pedido
    # fue una comunicacion mas rapida, y un indice de cinco titulos gastaba una
    # pagina completa en un documento de cuatro paginas de contenido. Es una
    # desviacion puntual del estandar de la Seccion 3.1 del CLAUDE.md, que pide
    # incluir_toc=True en todo transmittal; no cambia la regla para los proximos.
    crear_documento_adasa(
        titulo=("TECHNICAL REVIEW TRANSMITTAL N33 — SECOND STAGE RO "
                "BRINE MODULE"),
        codigo="P22-TM-09-000-033-0",
        output_filename=OUTPUT,
        incluir_toc=False,
    )

    doc = Document(OUTPUT)

    # 1. EXECUTIVE SUMMARY
    doc.add_heading("EXECUTIVE SUMMARY", level=1)
    add_para(doc, [
        ("TRANSMITTAL VERDICT: 2 — Approved as noted.", {"bold": True}),
        (" " + VERDICT,),
    ])
    add_para(doc, [("Scope of this review.", {"bold": True}), (" " + SCOPE,)])
    add_para(doc, [("Disposition at a glance:", {"bold": True})])
    for item in DISPOSITION:
        add_bullet(doc, item)
    add_para(doc, [
        ("Why Code 2 — Civil and Loading Layout.", {"bold": True}),
        (" " + WHY_CODE2,),
    ])
    add_para(doc, [
        ("Submittal 25007-0077 has not reached ADASA.", {"bold": True}),
        (" " + MISSING,),
    ])
    add_para(doc, [("Review period.", {"bold": True}), (" " + REVIEW_PERIOD,)])
    add_para(doc, [
        ("Pending Observations from Previous Transmittals lists the items open from "
         "earlier transmittals.",),
    ])

    # 2. OBSERVATIONS BY DOCUMENT
    doc.add_heading("OBSERVATIONS BY DOCUMENT", level=1)

    doc.add_heading("Civil and Loading Layout Rev B — P22-DWG-09-005-001", level=2)
    add_para(doc, [("Response Code: 2 — Approved as noted", {"bold": True})])
    add_para(doc, [("Status.", {"bold": True}), (" " + S21_STATUS,)])
    add_para(doc, [
        ("Action to issue at IFC Rev 0 — no new drawing revision required: ",
         {"bold": True}),
        (S21_ACTION,),
    ])

    doc.add_heading("GA of CIP Flushing Skid Pump Rev B — P22-DWG-09-005-010", level=2)
    add_para(doc, [("Response Code: 1 — Approved", {"bold": True})])
    add_para(doc, [("Status.", {"bold": True}), (" " + S22_STATUS,)])
    add_para(doc, [
        ("Action: none — accepted; issue directly at IFC Rev 0.", {"bold": True}),
        (" No annotated PDF accompanies this document.",),
    ])

    doc.add_heading("Equipment Layout Rev D — P22-DWG-09-005-003", level=2)
    add_para(doc, [("Response Code: 1 — Approved", {"bold": True})])
    add_para(doc, [("Status.", {"bold": True}), (" " + S23_STATUS,)])
    add_para(doc, [
        ("Action: none — accepted; issue directly at IFC Rev 0.", {"bold": True}),
        (" No annotated PDF accompanies this document.",),
    ])

    doc.add_heading("Process Flow Diagram Rev 0 — P22-DWG-09-009-01", level=2)
    add_para(doc, [("Response Code: 1 — Approved", {"bold": True})])
    add_para(doc, [("Status.", {"bold": True}), (" " + S24_STATUS,)])
    add_para(doc, [
        ("Action: none — accepted; issue directly at IFC Rev 0.", {"bold": True}),
        (" " + S24_REGISTER,),
    ])

    doc.add_heading("Control Philosophy Rev 1 — P22-BT-09-009-001", level=2)
    add_para(doc, [("Response Code: 1 — Approved", {"bold": True})])
    add_para(doc, [("Status.", {"bold": True}), (" " + S25_STATUS,)])
    add_para(doc, [
        ("Action: none — accepted.", {"bold": True}), (" " + S25_ACTION,),
    ])

    # 3. PENDING OBSERVATIONS FROM PREVIOUS TRANSMITTALS
    doc.add_heading("PENDING OBSERVATIONS FROM PREVIOUS TRANSMITTALS", level=1)
    add_para(doc, [
        ("The three most severe items open from earlier transmittals:",),
    ])
    add_simple_table(
        doc,
        [("Origin TM", "Document", "Observation", "Status")] + PENDING)
    add_para(doc, [("Also open: ", {"bold": True}), (ALSO_OPEN,)])

    # 4. ATTACHMENTS
    doc.add_heading("ATTACHMENTS", level=1)
    add_simple_table(
        doc,
        [("Document", "Code", "Annotated PDF")] + ATTACHMENTS)
    add_para(doc, [(ATTACH_NOTE,)])
    dl = doc.add_paragraph()
    dl.add_run("Download — this transmittal and the annotated PDF: ").bold = True
    aplicar_arial_12(dl)
    add_hyperlink(dl, DOWNLOAD_LINK, DOWNLOAD_LINK)

    # 5. RESPONSE SUMMARY
    doc.add_heading("RESPONSE SUMMARY", level=1)
    add_simple_table(
        doc,
        [("Document Code", "Title", "Rev", "Submittal", "Response Code")]
        + RESPONSE_SUMMARY)
    add_para(doc, [
        ("Overall Transmittal Verdict: 2 — APPROVED AS NOTED.", {"bold": True}),
        (CLOSING,),
    ])

    doc.save(OUTPUT)
    print(f"Documento generado: {OUTPUT}")


if __name__ == "__main__":
    main()
