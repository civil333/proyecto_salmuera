#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para generar TRANSMITTAL N5 ADASA-BW_WATER
Entrega 12 (Submittal 25007-0012) - 1 documento
Fecha: 23-Feb-2026

ESTRUCTURA:
1. EXECUTIVE SUMMARY (1.1 Key Findings, 1.2 Critical Observations)
2. GENERAL INFORMATION
3. DETAILED OBSERVATIONS BY DOCUMENT (1 documento: Equipment Layout)
4. PENDING OBSERVATIONS FROM PREVIOUS TRANSMITTALS (TM N2: 48d, TM N3: 26d, TM N4: 18d)
5. REQUIRED ACTIONS - BW WATER
6. ATTACHMENTS
7. RESPONSE SUMMARY

NOTAS:
- Usa skill template-adasa v7.2
- NO importar set_updatefields_true (la skill lo gestiona internamente)
- NO numeros manuales en headings (numeracion automatica del template)
- Base: patron de TM N4
"""

import sys
import os

# Agregar la ruta del skill template-adasa
skill_path = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..",
    "..",
    "..",
    ".claude",
    "skills",
    "template-adasa",
)
sys.path.insert(0, skill_path)

from ejemplo_documento import (
    crear_documento_adasa,
    aplicar_arial_12,
    add_simple_table,
)
from docx import Document
from docx.shared import Pt, Inches


def crear_transmittal():
    """Genera el Transmittal N5 en formato ADASA - E12 (1 documento)"""

    output_file = "TRANSMITTAL N5 ADASA-BW_WATER.docx"

    # Crear documento base con template ADASA
    crear_documento_adasa(
        titulo="TECHNICAL REVIEW TRANSMITTAL N5 - SECOND STAGE RO BRINE MODULE",
        codigo="P22-TM-09-000-005-0",
        preparado_por="Luis Rivera",
        revisado_por="Luis Rivera",
        aprobado_por="Victor Gutierrez",
        nombre_planta="TALTAL",
        cliente="ADASA",
        output_filename=output_file,
    )

    # Abrir documento y agregar contenido
    doc = Document(output_file)

    # Limpiar contenido de ejemplo generado por template ADASA
    paragraphs_to_remove = []
    for para in doc.paragraphs:
        texto = para.text.upper()
        if any(
            x in texto
            for x in [
                "RESUMEN EJECUTIVO",
                "INTRODUCCION",
                "INTRODUCCIÓN",
                "1. INTRO",
                "ESTE DOCUMENTO HA SIDO GENERADO",
                "AGREGUE AQUI EL CONTENIDO",
                "AGREGUE AQUÍ EL CONTENIDO",
            ]
        ):
            paragraphs_to_remove.append(para)

    for para in paragraphs_to_remove:
        p = para._element
        p.getparent().remove(p)

    # ===== 1. EXECUTIVE SUMMARY =====
    doc.add_heading("EXECUTIVE SUMMARY", level=1)

    # 1.1 Key Findings
    doc.add_heading("Key Findings", level=2)

    para = doc.add_paragraph()
    para.add_run("TRANSMITTAL VERDICT: 3 - TO BE REVISED").bold = True
    aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run(
        "One submittal reviewed containing a single document: Equipment Layout "
        "of RO Container (P22-DWG-09-005-003 Rev A). This is the first General "
        "Arrangement submitted for the UHPRO module and carries implications for "
        "thermal design, structural engineering, and spatial coordination with "
        "ADASA's existing plant infrastructure."
    )
    aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run(
        "The most significant finding: BW Water's Equipment Layout places the "
        "entire CIP system (tank, pump, cartridge filter) and chemical dosing "
        "tank inside the 40ft container. This arrangement contradicts the "
        "Technical Offer Rev.1, which explicitly specifies all CIP equipment "
        "outside the container, and contradicts the coordination meeting of "
        "February 18, 2026, where BW Water confirmed the CIP system would "
        "remain external."
    )
    aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run(
        "On a positive note, the container dimensions conform to the approved "
        "40ft standard, resolving the dimensional concern raised in Transmittal "
        "N4 (OBS-10)."
    )
    aplicar_arial_12(para)

    doc.add_paragraph()

    add_simple_table(
        doc,
        [
            ("Validation Item", "Status", "Reference"),
            (
                "Container 40ft (max 13m)",
                "COMPLIANT",
                "ET 5.1.10 L704 - 12.19m within 13m limit",
            ),
            (
                "CIP system location",
                "NOT COMPLIANT",
                "Offer Rev.1 lines 552-579 (outside); Feb-18 meeting confirmed external",
            ),
            (
                "Measurement system (MKS)",
                "NOT COMPLIANT",
                "All dimensions in imperial only (feet/inches)",
            ),
            (
                "Door specifications",
                "NOT SHOWN",
                "ET 5.1.10 L726-732 requires 4 door types",
            ),
            (
                "A/C unit locations",
                "NOT SHOWN",
                "ET 5.1.11 - n+1 confirmed Feb-17",
            ),
            (
                "PRFV non-slip floor",
                "NOT INDICATED",
                "ET 5.1.10 L733",
            ),
            (
                "Seismic anchoring",
                "NOT DETAILED",
                "ET 4.4, NCh 2369 Zone 3",
            ),
            (
                "Equipment weight table",
                "PROVIDED",
                "34,739 lb (15,758 kg) for 7 items",
            ),
        ],
    )

    para = doc.add_paragraph()
    para.add_run("\nStatistics: ").bold = True
    para.add_run(
        "0 Approved (0%) | 0 Approved as noted (0%) | 1 To be revised (100%)"
    )
    aplicar_arial_12(para)

    # 1.2 Critical Observations
    doc.add_heading("Critical Observations", level=2)

    add_simple_table(
        doc,
        [
            ("#", "Observation", "Severity", "Status"),
            (
                "OBS-01",
                "CIP system (tank 9,922 lb, pump 750 lb, filter 1,800 lb) and "
                "chemical tank (467 lb) shown inside container. Offer Rev.1 "
                "specifies all CIP equipment outside. Feb-18 meeting confirmed "
                "external CIP. Total CIP/chemical weight inside: 12,939 lb (5.9 tonnes).",
                "CRITICAL",
                "NEW",
            ),
            (
                "OBS-02",
                "All dimensions in imperial system only (feet/inches). Metric "
                "equivalents (MKS) required for Chilean engineering practice.",
                "MAJOR",
                "NEW",
            ),
            (
                "OBS-03",
                "No door specifications shown. ET 5.1.10 requires pedestrian "
                "(900x2200mm), equipment (110-degree outward opening), emergency, "
                "and sliding lateral doors.",
                "MAJOR",
                "NEW",
            ),
            (
                "OBS-04",
                "A/C units not shown in layout. ET 5.1.11 requires n+1 (minimum "
                "2 units). Location needed for spatial and thermal verification. "
                "BW Water confirmed 2 A/C (1W+1S) on Feb-17.",
                "MAJOR",
                "NEW",
            ),
            (
                "OBS-05",
                "Floor material not indicated. ET 5.1.10 L733 requires PRFV "
                "non-slip grating for operator walkways.",
                "MINOR",
                "NEW",
            ),
            (
                "OBS-06",
                "No seismic anchoring details. ET 4.4 requires NCh 2369, Zone 3. "
                "Weight table provided but anchoring provisions absent.",
                "MINOR",
                "NEW",
            ),
        ],
    )

    # ===== 2. GENERAL INFORMATION =====
    doc.add_heading("GENERAL INFORMATION", level=1)

    add_simple_table(
        doc,
        [
            ("Field", "Value"),
            ("Transmittal Code", "P22-TM-09-000-005-0"),
            ("Submittal 0012", "25007-0012 (Feb-20-2026) - 1 document"),
            ("Total Documents", "1"),
        ],
    )

    para = doc.add_paragraph()
    para.add_run("\nResponse Codes: ").bold = True
    para.add_run(
        "1=Approved, 2=Approved as noted, 3=To be revised, "
        "4=Rejected, 5=For Information"
    )
    aplicar_arial_12(para)

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Documents Reviewed:").bold = True
    aplicar_arial_12(para)

    add_simple_table(
        doc,
        [
            ("#", "Code", "Title", "Rev", "Verdict"),
            (
                "1",
                "P22-DWG-09-005-003",
                "Equipment Layout of RO Container",
                "A",
                "3 - To be revised",
            ),
        ],
    )

    # ===== 3. DETAILED OBSERVATIONS BY DOCUMENT =====
    doc.add_heading("DETAILED OBSERVATIONS BY DOCUMENT", level=1)

    # --- 3.1 Equipment Layout ---
    doc.add_heading(
        "Equipment Layout (P22-DWG-09-005-003-A) - TO BE REVISED", level=2
    )

    add_simple_table(
        doc,
        [
            ("Field", "Value"),
            ("Code", "P22-DWG-09-005-003-A"),
            ("Title", "Equipment Layout of RO Container"),
            ("Date", "19-Feb-2026"),
            ("Revision", "A (First issue)"),
            ("Drawing Status", "Issued for Approval"),
            ("Discipline", "Mechanical"),
        ],
    )

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Document Content:").bold = True
    aplicar_arial_12(para)
    para = doc.add_paragraph("- Page 1: Cover page (BW Water / Aguas Antofagasta)")
    aplicar_arial_12(para)
    para = doc.add_paragraph(
        "- Page 2: Plan view and partial elevation of 40ft RO container "
        "with equipment positions and operating weight table"
    )
    aplicar_arial_12(para)

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Equipment Weight Summary:").bold = True
    aplicar_arial_12(para)

    add_simple_table(
        doc,
        [
            ("Item", "Operating Weight (lb)", "Weight (kg)"),
            ("RO Train", "17,500", "7,938"),
            ("RO Cartridge Filter", "1,800", "816"),
            ("CIP Cartridge Filter", "1,800", "816"),
            ("CIP Tank", "9,922", "4,501"),
            ("CIP Pump", "750", "340"),
            ("Chemical Tank", "467", "212"),
            ("RO HP Pump", "2,500", "1,134"),
            ("TOTAL", "34,739", "15,758"),
        ],
    )

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Observations:").bold = True
    aplicar_arial_12(para)

    # OBS-01: CIP Inside Container
    doc.add_heading("OBS-01: CIP System Inside Container (CRITICAL)", level=3)

    para = doc.add_paragraph()
    para.add_run(
        "The Equipment Layout shows CIP Tank (9,922 lb), CIP Pump (750 lb), "
        "CIP Cartridge Filter (1,800 lb), and Chemical Tank (467 lb) located "
        "inside the 40ft container. Total CIP/chemical weight inside: "
        "12,939 lb (5,869 kg)."
    )
    aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run(
        "This contradicts two independent sources:"
    )
    aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run("Technical Offer Rev.1 (contractual document):").bold = True
    aplicar_arial_12(para)

    for item in [
        'CIP / Flushing Tank: "Located outside the 40 ft container" (line 553-555)',
        'CIP / Flushing Pump: "Located outside the 40 ft container" (line 556-559)',
        'CIP Cartridge Filter: "Located outside the 40 ft container" (line 574-576)',
        'Anti-Scalant System: "Located outside the 40 ft container" (line 577-579)',
        'Offer justification: "Vent fan is not needed since CIP system and chemical '
        'dosing system are located in outside the container" (line 550-551)',
    ]:
        para = doc.add_paragraph(f"- {item}")
        aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run("Coordination Meeting Feb-18-2026:").bold = True
    aplicar_arial_12(para)
    para = doc.add_paragraph(
        'BW Water confirmed: "CIP system external to container. '
        'Vendor will send updated layout drawing."'
    )
    aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run("ADASA Review Comments:").bold = True
    aplicar_arial_12(para)

    adasa_comments = [
        (
            "CIP System — Location and Distribution: ",
            "The CIP system must be located outside the container, with the "
            "equipment distribution and arrangement as specified by ADASA. "
            "Relocating the components outside is not sufficient — the updated "
            "drawing must reflect the specific disposition of each CIP element "
            "as requested, not an arbitrary external arrangement."
        ),
        (
            "Liquid Infrastructure Alignment: ",
            "All liquid connections (feed, permeate, concentrate, and CIP) must "
            "exit from the side indicated in ADASA's review markup. Equipment "
            "layout and piping must be configured accordingly."
        ),
        (
            "Local Control Panel Proximity: ",
            "The LCP is shown outside the container, consistent with Offer Rev.1. "
            "The panel must be physically attached or immediately adjacent to the "
            "module — cable interconnection and FAT testing are within BW Water's "
            "contractual scope. Cable tray routing between the LCP and the "
            "container interior must be resolved in the updated drawing."
        ),
    ]

    for label, text in adasa_comments:
        para = doc.add_paragraph()
        para.add_run(f"- ")
        para.add_run(label).bold = True
        para.add_run(text)
        aplicar_arial_12(para)

    # OBS-02: Imperial Units
    doc.add_heading("OBS-02: Measurements in Imperial System Only (MAJOR)", level=3)

    para = doc.add_paragraph()
    para.add_run(
        "All dimensions are given in feet and inches (e.g., 36'-10\", 6'-3\", "
        "2'-7\"). The project is located in Chile where the metric system (MKS) "
        "is standard practice. All ADASA engineering documentation, ET "
        "specifications, and referenced Chilean norms (NCh 2369) use metric units."
    )
    aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run(
        "Metric (MKS) equivalents must be included for all dimensions. "
        "A dual-unit drawing is an acceptable alternative."
    )
    aplicar_arial_12(para)

    # OBS-03: No Doors
    doc.add_heading("OBS-03: No Door Specifications Shown (MAJOR)", level=3)

    para = doc.add_paragraph()
    para.add_run("ET Section 5.1.10 (Lines 726-732) requires:")
    aplicar_arial_12(para)

    for door in [
        "Pedestrian access door: 900 mm wide x 2,200 mm high",
        "Equipment access door: sized for largest internal equipment, "
        "110-degree opening angle, opening outward",
        "Emergency door",
        "Lateral sliding door for additional access",
    ]:
        para = doc.add_paragraph(f"- {door}")
        aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run(
        "The Equipment Layout does not show any door locations, dimensions, "
        "or swing directions. For an \"Issued for Approval\" drawing, door "
        "positions are essential to verify equipment removal paths and "
        "emergency egress compliance."
    )
    aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run(
        "The revised drawing must show all four door types with positions, "
        "dimensions, and swing directions: pedestrian (900x2200mm), equipment "
        "(110-degree outward opening), emergency, and lateral sliding."
    )
    aplicar_arial_12(para)

    # OBS-04: No A/C
    doc.add_heading("OBS-04: A/C Unit Locations Not Shown (MAJOR)", level=3)

    para = doc.add_paragraph()
    para.add_run(
        "ET Section 5.1.11 requires n+1 A/C units (minimum 2) for continuous "
        "24/7 operation maintaining interior temperature below 25 degrees C. "
        "BW Water confirmed the 2-unit A/C configuration (1W+1S) on February 17, 2026."
    )
    aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run("Their location affects:")
    aplicar_arial_12(para)

    for effect in [
        "Available equipment space inside the container",
        "Air distribution effectiveness",
        "Maintenance access to the A/C units",
        "Thermal calculation accuracy (still pending 48 days)",
    ]:
        para = doc.add_paragraph(f"- {effect}")
        aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run(
        "A/C unit placement must appear in the next drawing revision. "
        "Coordination with the pending thermal calculation (TM N2, 48 days) "
        "is required to confirm both spatial and thermal adequacy."
    )
    aplicar_arial_12(para)

    # OBS-05: No PRFV floor
    doc.add_heading("OBS-05: Floor Material Not Indicated (MINOR)", level=3)

    para = doc.add_paragraph()
    para.add_run(
        "ET Section 5.1.10 (Line 733) requires PRFV (fiberglass reinforced "
        "plastic) non-slip grating for the operator walkway floor. The drawing "
        "does not indicate floor material or walkway areas."
    )
    aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run(
        "Walkway areas and PRFV non-slip grating specification should be "
        "indicated in the drawing notes."
    )
    aplicar_arial_12(para)

    # OBS-06: No Seismic
    doc.add_heading("OBS-06: No Seismic Anchoring Details (MINOR)", level=3)

    para = doc.add_paragraph()
    para.add_run(
        "ET Section 4.4 requires all internal components and the container "
        "envelope to be designed for NCh 2369, Zone 3. The drawing provides an "
        "operating weight table (34,739 lb / 15,758 kg), which is useful for "
        "structural calculations. No anchoring details, bolt patterns, or "
        "structural provisions are shown."
    )
    aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run(
        "Drawing notes should reference the seismic calculation report "
        "(required by ET Section 7) and include basic anchoring provisions "
        "for the internal equipment."
    )
    aplicar_arial_12(para)

    # Positive Findings
    doc.add_paragraph()
    para = doc.add_paragraph()
    para.add_run("Positive Findings:").bold = True
    aplicar_arial_12(para)

    add_simple_table(
        doc,
        [
            ("Item", "Status", "Details"),
            (
                "Container 40ft",
                "COMPLIANT",
                "12.19m within ET 5.1.10 limit (13m). Resolves TM N4 OBS-10",
            ),
            (
                "Equipment weight table",
                "PROVIDED",
                "Operating weights for 7 items, total 15,758 kg",
            ),
            (
                "Equipment identification",
                "ADEQUATE",
                "Major equipment clearly labeled and positioned",
            ),
            (
                "HP Pump position",
                "LOGICAL",
                "Adjacent to feed turbocharger for short HP piping runs",
            ),
            (
                "LCP outside container",
                "COMPLIANT",
                "Per ET and Offer Rev.1",
            ),
        ],
    )

    para = doc.add_paragraph()
    para.add_run("\nVERDICT: 3 - TO BE REVISED").bold = True
    para.add_run(
        " (CIP system location contradicts Offer and meeting agreement; "
        "multiple ET requirements not shown in layout)"
    )
    aplicar_arial_12(para)

    # ===== 4. PENDING OBSERVATIONS FROM PREVIOUS TRANSMITTALS =====
    doc.add_heading("PENDING OBSERVATIONS FROM PREVIOUS TRANSMITTALS", level=1)

    para = doc.add_paragraph()
    para.add_run(
        "This section tracks unresolved observations from prior transmittals, "
        "updated as of February 23, 2026."
    )
    aplicar_arial_12(para)

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Observations CLOSED since TM N4:").bold = True
    aplicar_arial_12(para)

    closed_items = [
        "Pt-100 motor windings: CLOSED Feb-17 (confirmed by BW Water per ET 5.3)",
        "Vibration transmitters: CLOSED Feb-17 (confirmed per ET 5.5.7)",
        "Duplicate TAG FIT-09-001: CLOSED Feb-17 (renumbered to FIT-09-002)",
        "A/C n+1 configuration: CLOSED Feb-17 (confirmed per ET 5.1.11)",
        "Ethernet IP PLC-VFD: CLOSED Feb-18 (ratified as communication protocol)",
        "Container >40ft (TM N4 OBS-10): CLOSED by this GA (40ft confirmed)",
    ]
    for item in closed_items:
        para = doc.add_paragraph(f"- {item}")
        aplicar_arial_12(para)

    doc.add_paragraph()

    # 4.1 From TM N2
    doc.add_heading(
        "From Transmittal N2 (January 6, 2026) - 48 DAYS PENDING", level=2
    )

    add_simple_table(
        doc,
        [
            ("#", "Observation", "Original Document", "Days", "Impact"),
            (
                "1",
                "A/C thermal calculation not delivered",
                "New document required",
                "48",
                "Blocks thermal compliance verification",
            ),
            (
                "2",
                "Modbus TCP Memory Map not delivered",
                "Control Architecture",
                "48",
                "Critical for DCS/SCADA integration",
            ),
        ],
    )

    doc.add_paragraph()

    # 4.2 From TM N3
    doc.add_heading(
        "From Transmittal N3 (January 28, 2026) - 26 DAYS PENDING", level=2
    )

    add_simple_table(
        doc,
        [
            ("#", "Observation", "Document", "Days", "Severity"),
            (
                "3",
                "VM-09-015 DN100 ANSI 900# manual actuation",
                "Valve List",
                "26",
                "CRITICAL (active dispute)",
            ),
            (
                "4",
                "IO List coordination signals (VFD variables, DO/DI "
                "for external plant signals)",
                "IO List",
                "26",
                "CRITICAL",
            ),
        ],
    )

    doc.add_paragraph()

    # 4.3 From TM N4
    doc.add_heading(
        "From Transmittal N4 (February 5, 2026) - 18 DAYS PENDING", level=2
    )

    add_simple_table(
        doc,
        [
            ("#", "Observation", "Document", "Days", "Severity"),
            (
                "5",
                "PLC 60 Hz frequency (ET 5.4.7 requires 50 Hz)",
                "Utility Consumption List",
                "18",
                "CRITICAL",
            ),
            (
                "6",
                "HP Pump power inconsistency (93/86/92/83 kW)",
                "Multiple documents",
                "18",
                "MAJOR",
            ),
            (
                "7",
                "UPS not included in BOM (ET 5.4 requires 8h autonomy)",
                "Control Architecture",
                "18",
                "MAJOR",
            ),
            (
                "8",
                "LIT TAG discrepancy (LIT-09-001 vs LIT-09-002)",
                "Multiple documents",
                "18",
                "MINOR",
            ),
        ],
    )

    para = doc.add_paragraph()
    para.add_run("\nNote: ").bold = True
    para.add_run(
        "TM N4 OBS-10 (Container >40ft) is CLOSED. The Equipment Layout "
        "confirms a standard 40ft container."
    )
    aplicar_arial_12(para)

    para = doc.add_paragraph()
    para.add_run("\nTotal open observations from previous transmittals: 8")
    aplicar_arial_12(para)

    # ===== 5. REQUIRED ACTIONS - BW WATER =====
    doc.add_heading("REQUIRED ACTIONS - BW WATER", level=1)

    # 5.1 Critical Actions (from this GA)
    doc.add_heading("Critical Actions - Equipment Layout", level=2)

    add_simple_table(
        doc,
        [
            ("#", "Action", "Document", "Reference"),
            (
                "1",
                "Revise layout: place CIP system outside container per "
                "Offer Rev.1 and Feb-18 meeting agreement. Any alternative "
                "arrangement requires formal ADASA approval.",
                "P22-DWG-09-005-003",
                "OBS-01, Offer Rev.1",
            ),
            (
                "2",
                "Include metric (MKS) equivalents for all dimensions.",
                "P22-DWG-09-005-003",
                "OBS-02",
            ),
            (
                "3",
                "Show all required doors: pedestrian (900x2200mm), "
                "equipment (110-degree outward), emergency, and sliding "
                "lateral, with locations and dimensions.",
                "P22-DWG-09-005-003",
                "OBS-03, ET 5.1.10",
            ),
            (
                "4",
                "Show A/C unit locations (2 units per confirmed n+1 "
                "configuration).",
                "P22-DWG-09-005-003",
                "OBS-04, ET 5.1.11",
            ),
        ],
    )

    doc.add_paragraph()

    # 5.2 Ongoing Critical Actions
    doc.add_heading(
        "Ongoing Critical Actions (from previous transmittals)", level=2
    )

    add_simple_table(
        doc,
        [
            ("#", "Action", "Days Pending", "Reference"),
            (
                "5",
                "Deliver A/C thermal calculation",
                "48",
                "TM N2, ET 5.1.11",
            ),
            (
                "6",
                "Deliver Modbus TCP Memory Map",
                "48",
                "TM N2",
            ),
            (
                "7",
                "Resolve VM-09-015 actuation (electric vs manual DN100 "
                "ANSI 900#)",
                "26",
                "TM N3, ET 5.2.3",
            ),
            (
                "8",
                "Update IO List: VFD Ethernet IP variables + external "
                "plant digital signals",
                "26",
                "TM N3, Feb-18 meeting",
            ),
            (
                "9",
                "Confirm PLC 50 Hz compatibility",
                "18",
                "TM N4, ET 5.4.7",
            ),
            (
                "10",
                "Unify HP Pump power value across all documents",
                "18",
                "TM N4",
            ),
            (
                "11",
                "Include UPS in BOM with 8 hours autonomy",
                "18",
                "TM N4, ET 5.4",
            ),
        ],
    )

    doc.add_paragraph()

    # 5.3 Minor Actions
    doc.add_heading("Minor Actions", level=2)

    add_simple_table(
        doc,
        [
            ("#", "Action", "Document", "Reference"),
            (
                "12",
                "Indicate PRFV non-slip floor material in layout",
                "P22-DWG-09-005-003",
                "OBS-05, ET 5.1.10",
            ),
            (
                "13",
                "Include anchoring provisions or reference seismic "
                "calculation report",
                "P22-DWG-09-005-003",
                "OBS-06, ET 4.4",
            ),
            (
                "14",
                "Resolve LIT TAG discrepancy (LIT-09-001 vs LIT-09-002)",
                "Multiple documents",
                "TM N4",
            ),
        ],
    )

    # ===== 6. ATTACHMENTS =====
    doc.add_heading("ATTACHMENTS", level=1)

    add_simple_table(
        doc,
        [
            ("#", "Attachment", "Description"),
            (
                "P",
                "P22-DWG-09-005-003-A_Equipment_Layout_Comments.pdf",
                "Equipment Layout with ADASA review comments",
            ),
        ],
    )

    # ===== 7. RESPONSE SUMMARY =====
    doc.add_heading("RESPONSE SUMMARY", level=1)

    add_simple_table(
        doc,
        [
            ("Document", "Code", "Verdict"),
            (
                "Equipment Layout of RO Container",
                "P22-DWG-09-005-003-A",
                "3 - To be revised",
            ),
        ],
    )

    para = doc.add_paragraph()
    para.add_run("\nTRANSMITTAL VERDICT: 3 - TO BE REVISED").bold = True
    aplicar_arial_12(para)

    # Guardar documento
    doc.save(output_file)
    print(f"Document generated: {output_file}")
    print(f"\nEstructura Transmittal N5 (23-Feb-2026):")
    print(f"  1. EXECUTIVE SUMMARY")
    print(f"     1.1 Key Findings (E12)")
    print(f"     1.2 Critical Observations (6 items: 1 CRITICAL, 3 MAJOR, 2 MINOR)")
    print(f"  2. GENERAL INFORMATION (Submittal 0012, 1 doc)")
    print(f"  3. DETAILED OBSERVATIONS BY DOCUMENT")
    print(f"     3.1 Equipment Layout - TO BE REVISED")
    print(f"         OBS-01: CIP inside container (CRITICAL)")
    print(f"         OBS-02: Imperial units only (MAJOR)")
    print(f"         OBS-03: No doors shown (MAJOR)")
    print(f"         OBS-04: No A/C locations (MAJOR)")
    print(f"         OBS-05: No PRFV floor (MINOR)")
    print(f"         OBS-06: No seismic anchoring (MINOR)")
    print(f"  4. PENDING OBSERVATIONS FROM PREVIOUS TRANSMITTALS")
    print(f"     4.1 From TM N2 (48 days pending) - 2 items")
    print(f"     4.2 From TM N3 (26 days pending) - 2 items")
    print(f"     4.3 From TM N4 (18 days pending) - 4 items")
    print(f"     [6 items CLOSED since TM N4, incl. OBS-10 Container]")
    print(f"  5. REQUIRED ACTIONS - BW WATER (14 total)")
    print(f"     5.1 Critical - Equipment Layout (#1-4)")
    print(f"     5.2 Ongoing Critical - Previous TMs (#5-11)")
    print(f"     5.3 Minor (#12-14)")
    print(f"  6. ATTACHMENTS (1 PDF)")
    print(f"  7. RESPONSE SUMMARY")
    print(f"\n[OK] Container 40ft CUMPLE ET 5.1.10 - cierra OBS-10 TM N4")
    print(f"[!]  CIP dentro del container - contradice Oferta Rev.1 y reunion 18-Feb")
    print(f"[OK] Observaciones pendientes: TM N2=48d, TM N3=26d, TM N4=18d")
    print(f"[OK] 6 items cerrados (Pt-100, vibration, FIT-09-002, A/C n+1, EthIP, Container)")
    return output_file


if __name__ == "__main__":
    crear_transmittal()
