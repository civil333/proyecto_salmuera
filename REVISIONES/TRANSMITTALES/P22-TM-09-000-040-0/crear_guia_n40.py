#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera la guia interna _GUIA_REVISION_FAT_N40_NO_ENVIAR.docx (espanol, es-CL).

Acompana a la plantilla TRANSMITTAL N40 ADASA-BW_WATER_PLANTILLA.docx para que
Victor Gutierrez revise y emita el N40 en Word a mano.

Fuentes verificadas el 17-Sep-2026 sobre los .md de las bases:
  - ET P22-ET-09-000-001-0 seccion 8.1 (BASES TECNICAS/md/...ET-MODULO.md,
    lineas 1728-1833). Ojo: los tres escenarios de falla van "ej." en la ET, son
    ejemplos; lo obligatorio es incluir escenarios de falla simulados.
  - ITP P22-BA-09-000-004 Rev 0, fila 7 (ENTREGA 57/md, lineas 109-118).
  - PIE Base P22-IT-09-000-001-0 seccion 2 (lineas 185-197).
  - BAE clausula 37 (30 dias de aviso internacional con tercero) y 37.2 (7 dias
    habiles desde el dia habil siguiente a la recepcion).
Criterios de over-reach: memorias del proyecto (FAT en seco, soft-I/O, HART,
presion por linea, Rev 0 solo Codigo 1 o 3).
"""

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "_GUIA_REVISION_FAT_N40_NO_ENVIAR.docx")
LANG = "es-CL"


# --------------------------------------------------------------------- helpers
def _set_lang(rPr, lang=LANG):
    for el in rPr.findall(qn("w:lang")):
        rPr.remove(el)
    el = OxmlElement("w:lang")
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)
    rPr.append(el)


def fijar_idioma_documento(doc):
    """Ultima operacion antes de guardar: sin esto Word subraya tildes y enes."""
    for nombre in ("Normal", "Heading 1", "Heading 2", "Title"):
        try:
            _set_lang(doc.styles[nombre].element.get_or_add_rPr())
        except KeyError:
            pass
    parrafos = list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parrafos.extend(cell.paragraphs)
    for p in parrafos:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr())


def _arial(style_or_run, size=None, bold=None, color=None):
    font = style_or_run.font
    font.name = "Arial"
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(att), "Arial")
    # los estilos de titulo traen fuente del tema, que manda sobre w:ascii
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(att)) is not None:
            del rfonts.attrib[qn(att)]
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color


def preparar_estilos(doc):
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.7)
        s.left_margin = s.right_margin = Inches(0.8)
    _arial(doc.styles["Normal"], size=10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    negro = RGBColor(0, 0, 0)
    _arial(doc.styles["Title"], size=16, bold=True, color=negro)
    _arial(doc.styles["Heading 1"], size=13, bold=True, color=negro)
    _arial(doc.styles["Heading 2"], size=11, bold=True, color=negro)
    for nombre in ("List Bullet", "List Number"):
        _arial(doc.styles[nombre], size=10.5)


def para(doc, *segmentos, style=None):
    """Segmentos: str (normal) o ("texto", "b") para negrita."""
    p = doc.add_paragraph(style=style)
    for seg in segmentos:
        if isinstance(seg, tuple):
            p.add_run(seg[0]).bold = True
        else:
            p.add_run(seg)
    return p


def bullet(doc, *segmentos):
    return para(doc, *segmentos, style="List Bullet")


def numerado(doc, *segmentos):
    return para(doc, *segmentos, style="List Number")


def tabla(doc, filas, anchos, size=9):
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = t._tbl
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl.tblPr.append(layout)
    for col, ancho in zip(tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")), anchos):
        col.set(qn("w:w"), str(int(ancho * 1440)))
    for i, fila in enumerate(filas):
        row = t.rows[i]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            hdr = OxmlElement("w:tblHeader")
            hdr.set(qn("w:val"), "true")
            trPr.append(hdr)
        for j, ancho in enumerate(anchos):
            row.cells[j].width = Inches(ancho)
        grupo = i > 0 and len(fila) > 1 and all(c == "" for c in fila[1:])
        if grupo:
            # fila de grupo: una sola celda a todo el ancho
            celdas = [row.cells[0].merge(row.cells[-1])]
            fila = fila[:1]
        else:
            celdas = row.cells
        for cell, texto in zip(celdas, fila):
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(texto)
            run.font.size = Pt(size)
            run.font.name = "Arial"
            if i == 0 or grupo:
                run.bold = True
    doc.add_paragraph()
    return t


# ------------------------------------------------------------------- contenido
PAUTA = [
    ("#", "Qué verificar", "Cita para el transmittal y el PDF anotado", "Resultado", "ID"),
    ("A. Identidad y forma del documento", "", "", "", ""),
    ("1", "Tiene código y revisión propios y cubre el módulo completo. El PLC/LCP FAT "
          "Procedure - Hardware (P22-PP-09-000-001) no sirve para esto: excluye las "
          "pruebas de software y la lógica de proceso.",
     "Inspection and Test Plan P22-BA-09-000-004 Rev 0, row 7.1 - Approval of "
     "Detailed FAT Procedure", "", ""),
    ("2", "Es un procedimiento con formularios en blanco para registrar resultados. Un "
          "registro de una prueba ya ejecutada no es un procedimiento.",
     "Technical Specification (P22-ET-09-000-001-0), Section 8.1 - Minimum Scope of "
     "Factory Acceptance Tests (results recorded in a specific protocol, according "
     "to the approved FAT procedure)", "", ""),
    ("3", "Los marcadores genéricos que trae el ITP ([Code], [HD Code], [Doc. Code], "
          "[Cert. Code] y los códigos PROV-...) están reemplazados por el código y la "
          "revisión de los documentos reales: P&ID, layout, datasheets, Control "
          "Philosophy, checklists.",
     "Inspection and Test Plan P22-BA-09-000-004 Rev 0, rows 7.2 to 7.7; PIE Base "
     "(P22-IT-09-000-001-0), Section 2", "", ""),
    ("4", "Cada criterio de aceptación indica documento fuente, sección y valor numérico "
          "con unidad y tolerancia. Las filas 7.3 a 7.5 del ITP delegan el criterio en "
          "este procedimiento, así que aquí tiene que estar. \"Según ET\" o \"según "
          "plano\" no sirve.",
     "PIE Base (P22-IT-09-000-001-0), Section 2 (generic references such as 'as per "
     "specification' are not acceptable)", "", ""),
    ("B. Prerrequisito", "", "", "", ""),
    ("5", "Exige que las pruebas hidrostáticas de baja y alta presión estén completadas "
          "y aprobadas antes de abrir el FAT.",
     "Technical Specification, Section 8.1", "", ""),
    ("C. Alcance mínimo de la ET sección 8.1", "", "", "", ""),
    ("6", "Inspección visual y de completitud del montaje contra el Layout y el P&ID "
          "aprobados.",
     "Technical Specification, Section 8.1; ITP row 7.2 (Hold Point)", "", ""),
    ("7", "Verificación dimensional: dimensiones generales, ubicación y tipo de las "
          "conexiones de interfaz (tie-ins de la ET sección 5.6) y puntos de izaje.",
     "Technical Specification, Section 8.1; ITP row 7.3 (Hold Point)", "", ""),
    ("8", "Integridad del montaje mecánico: instalación e identificación de tuberías, "
          "soportes y equipos.",
     "Technical Specification, Section 8.1", "", ""),
    ("9", "Pruebas funcionales en seco: giro libre y sentido de rotación de las bombas; "
          "actuación de válvulas automáticas; señalización de válvulas manuales "
          "críticas con finales de carrera.",
     "Technical Specification, Section 8.1; ITP row 7.4 (Witness Point)", "", ""),
    ("10", "Instrumentación: instalación y TAG de todos los instrumentos; continuidad y "
           "loop checks hasta el PLC; energización del tablero y del PLC.",
     "Technical Specification, Section 8.1; ITP row 7.5 (Hold Point)", "", ""),
    ("11", "Revisión de las pantallas HMI según ISA 101.",
     "Technical Specification, Section 8.1", "", ""),
    ("12", "Simulación avanzada y dinámica del control, con señales forzadas o simulador "
           "de proceso, que cubra arranque, parada normal y de emergencia, CIP, alarmas "
           "críticas, interlocks y la respuesta de los lazos importantes.",
     "Technical Specification, Section 8.1; ITP row 7.5 (Hold Point)", "", ""),
    ("13", "Escenarios de falla simulados, que la ET exige. Los ejemplos que da son "
           "pérdida de señal de un transmisor de presión crítico, disparo de bomba bajo "
           "carga y actuación de válvula de seguridad. Se exige que haya escenarios de "
           "falla con su respuesta esperada; esos tres son ejemplos, no una lista "
           "cerrada.",
     "Technical Specification, Section 8.1 (simulated fault scenarios)", "", ""),
    ("14", "Comunicaciones con sistemas externos, si aplica: las señales de coordinación "
           "con el PLC de planta.",
     "Technical Specification, Section 8.1", "", ""),
    ("15", "Revisión documental preliminar: manuales O&M preliminares, certificados de "
           "material del super dúplex, registros de las hidrostáticas y dossier de "
           "calidad.",
     "Technical Specification, Section 8.1; ITP row 7.6 (Review)", "", ""),
    ("16", "Conformidad con la normativa eléctrica chilena (SEC) como Hold Point, "
           "requisito previo al embarque.",
     "Technical Specification, Section 8.1; ITP row 7.7 (Hold Point)", "", ""),
    ("17", "Medición base de espesores por ultrasonido en la cañería de alta presión, con "
           "el plano de puntos de medición.",
     "ITP row 7.8 (Witness Point)", "", ""),
    ("18", "Cierre: protocolo de resultados, cierre de las no conformidades del FAT y "
           "Acta de Aprobación FAT emitida por ADASA.",
     "Technical Specification, Section 8.1; ITP row 7.9 (Hold Point)", "", ""),
    ("D. Presencia y aviso", "", "", "", ""),
    ("19", "Declara los puntos H y W de la fila 7 del ITP y el aviso previo a ADASA y al "
           "tercero inspector: 30 días en suministro internacional con inspección de "
           "terceros.",
     "BAE Clause 37; Inspection and Test Plan legend (H and W require notification)",
     "", ""),
    ("20", "Instrumentos de prueba identificados (modelo y serie) con certificado de "
           "calibración vigente.",
     "Sin cláusula verificada: citar la coherencia con lo pedido para el PLC/LCP FAT "
     "Procedure en Transmittal N37", "", ""),
    ("E. Uso posterior (pedido de ADASA)", "", "", "", ""),
    ("21", "Sirve de base para el procedimiento de pruebas en sitio (SIT) de la "
           "reconexión eléctrica. La ET, las BAE y la PIE no nombran un SIT: si falta "
           "algo para ese uso, es un pedido de ADASA y no un incumplimiento.",
     "Redactar como \"ADASA requests ...\", sin cita contractual y sin imputar falta",
     "", ""),
]

CODIGOS = [
    ("Código", "Cuándo corresponde", "Precedente"),
    ("1 — Approved",
     "Cubre el alcance y sus criterios son verificables. Lo que queda son detalles que "
     "no cambian cómo se ejecuta la prueba.",
     "—"),
    ("2 — Approved as noted",
     "El propio procedimiento necesita correcciones menores que se incorporan al "
     "emitirlo en Rev 0, sin revisión intermedia. No cabe si el documento ya llegó en "
     "Rev 0.",
     "PLC/LCP FAT Procedure Rev A (N27) y Rev C (N38)"),
    ("3 — To be revised",
     "Falta en sustancia alcance de la sección 8.1 (por ejemplo, la simulación con "
     "fallas o el cierre con Acta); los criterios no tienen valores verificables; "
     "llega como registro ya ejecutado; o no se identifica con código y revisión.",
     "PLC/LCP FAT Procedure Rev B (N37)"),
]


def crear_guia():
    doc = Document()
    preparar_estilos(doc)

    encabezado = doc.sections[0].header.paragraphs[0]
    run = encabezado.add_run("INTERNO — NO ENVIAR · Transmittal N40 · Guía de revisión")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.bold = True

    doc.add_paragraph("Transmittal N40: revisión del procedimiento FAT del módulo",
                      style="Title")
    para(doc, ("Uso interno, no se envía a BW Water. ", "b"),
         "Preparada el 17 de septiembre de 2026 para Victor Gutierrez, que revisa y "
         "emite el N40 mientras Luis Rivera está fuera. Acompaña a la plantilla "
         "TRANSMITTAL N40 ADASA-BW_WATER_PLANTILLA.docx, en la misma carpeta.")

    # ---- 1
    doc.add_heading("1. Qué hay en la carpeta", level=1)
    tabla(doc, [
        ("Archivo", "Para qué"),
        ("TRANSMITTAL N40 ADASA-BW_WATER_PLANTILLA.docx",
         "El transmittal a completar. Antes de editar, guardarlo como "
         "TRANSMITTAL N40 ADASA-BW_WATER.docx y trabajar sobre esa copia."),
        ("_GUIA_REVISION_FAT_N40_NO_ENVIAR.docx", "Esta guía."),
        ("COMENTARIOS", "Carpeta para el PDF recibido y su versión anotada."),
        ("crear_plantilla_tm40.py, crear_guia_n40.py, "
         "P22-TM-09-000-040-0_TRANSMITTAL.md",
         "Respaldo técnico. No se usan para emitir. No volver a correr los scripts: "
         "regeneran la plantilla con la fecha del día."),
    ], anchos=[2.6, 4.3])

    # ---- 2
    doc.add_heading("2. Pasos", level=1)
    numerado(doc, "Al llegar el correo de BW Water, anotar el número de submittal "
                  "(25007-00XX), la fecha de recepción y el código y revisión del "
                  "procedimiento. Guardar el PDF en COMENTARIOS.")
    numerado(doc, "Compartir el procedimiento con Bureau Veritas, como quedó en la "
                  "minuta del 17 de septiembre. Si Bureau Veritas manda comentarios, se "
                  "evalúan y los que se acojan entran como observaciones de ADASA.")
    numerado(doc, "Revisar el procedimiento fila por fila con la pauta de la sección 4.")
    numerado(doc, "Decidir el código con la sección 6.")
    numerado(doc, "Si el código es 2 o 3, anotar el PDF en Acrobat según la sección 7. "
                  "Con código 1 no se anota.")
    numerado(doc, "Completar el Word: reemplazar cada marcador amarillo y borrar los "
                  "bloques que no aplican (sección 8).")
    numerado(doc, "Hacer el chequeo de la sección 9 y guardar el PDF desde Word.")
    numerado(doc, "Enviar el correo de cobertura de la sección 10.")

    # ---- 3
    doc.add_heading("3. Plazo", level=1)
    bullet(doc, ("Regla: ", "b"),
           "la Cláusula 37.2 de las BAE da siete días hábiles, contados desde el día "
           "hábil siguiente a la recepción formal y completa del documento.")
    bullet(doc, ("Ejemplo: ", "b"),
           "si llega el viernes 18 de septiembre (feriado en Chile, igual que el sábado "
           "19), el conteo parte el lunes 21 y el plazo vence el martes 29 de "
           "septiembre.")
    bullet(doc, ("Objetivo interno: ", "b"),
           "emitir durante la semana del 21 de septiembre. El plazo es de control "
           "interno y no se escribe en el transmittal.")

    # ---- 4
    doc.add_heading("4. Pauta de verificación", level=1)
    para(doc, "La pauta sale de la Especificación Técnica (sección 8.1), del Inspection "
              "and Test Plan Rev 0 aprobado en Código 1 (fila 7) y de la PIE Base "
              "(sección 2). La tercera columna trae la cita en inglés para el "
              "transmittal y el PDF anotado. En Resultado: Cumple, No cumple o N.A.; "
              "en ID, la observación que se levantó.")
    tabla(doc, PAUTA, anchos=[0.35, 2.85, 2.35, 0.8, 0.55], size=8.5)

    # ---- 5
    doc.add_heading("5. Qué no observar", level=1)
    para(doc, "Son puntos ya resueltos con BW Water o que la base contractual no "
              "sostiene. Levantarlos debilita el resto de la revisión.")
    bullet(doc, ("Que el FAT sea en seco. ", "b"),
           "Se acordó que la prueba con agua y bombas en marcha va al comisionamiento y "
           "a las Pruebas de Desempeño en Taltal. Lo que sí se exige es la simulación "
           "con fallas, que incluye el disparo de bomba bajo carga.")
    bullet(doc, ("Cableado directo para las señales de campo. ", "b"),
           "Las entradas y salidas de instrumentos, válvulas y dosificadoras por "
           "Ethernet/IP están aceptadas. Solo las cuatro señales de coordinación con el "
           "PLC de planta (enable, running, fault y local/remoto) van por contacto de "
           "relé; el término en inglés es \"relay contact\".")
    bullet(doc, ("Decodificación HART central en el PLC. ", "b"),
           "Basta con que cada instrumento de campo sea 4-20 mA con HART.")
    bullet(doc, ("Una presión de prueba única para todo el super dúplex. ", "b"),
           "La presión va por línea, 1,5 veces la de diseño, según la Line List "
           "aprobada.")
    bullet(doc, ("Pedir de nuevo algo aprobado en Código 1 en otro documento. ", "b"),
           "Se cita ese documento aprobado.")
    bullet(doc, ("Datos internos. ", "b"),
           "El contrato con Bureau Veritas (jornadas, oferta), Van Doorn, los códigos "
           "del registro de compromisos (PRG-50, BV-13) y cualquier análisis interno "
           "quedan fuera del transmittal y del PDF anotado.")

    # ---- 6
    doc.add_heading("6. Cómo decidir el código", level=1)
    para(doc, ("Si el procedimiento llega en Rev 0, solo caben Código 1 o Código 3. ",
               "b"),
         "El código refleja el estado del propio procedimiento: una nota cuyo "
         "entregable está en otro documento no lo baja a Código 2.")
    tabla(doc, CODIGOS, anchos=[1.3, 3.9, 1.7])

    # ---- 7
    doc.add_heading("7. Cómo anotar el PDF en Acrobat", level=1)
    bullet(doc, "Herramienta Comentario, Cuadro de texto, sobre la página y cerca del "
                "texto observado.")
    bullet(doc, ("Texto: ", "b"),
           "\"OBS-01: <defecto>. Correct: <qué debe hacer BW Water>.\" Observaciones "
           "OBS-01, OBS-02, ... y notas NOTE-01, NOTE-02, ... en orden correlativo.")
    bullet(doc, ("Sin etiqueta de severidad en el texto: ", "b"),
           "la da el color del cuadro. Rojo, crítico (incumplimiento contractual); "
           "naranja, mayor; amarillo, menor; azul, NOTE.")
    bullet(doc, "Redactar en forma directiva: qué corregir, no una opinión. Ejemplo: "
                "\"Correct: state the numeric acceptance value and tolerance for each "
                "loop check.\"")
    bullet(doc, "Guardar como <nombre original>_CC_ADASA.pdf en COMENTARIOS. Los IDs "
                "del PDF son los que cita el bloque Action del transmittal.")

    # ---- 8
    doc.add_heading("8. Cómo completar el transmittal", level=1)
    bullet(doc, "Inglés y tercera persona: \"ADASA requires\", \"BW Water shall\".")
    bullet(doc, "Citar la especificación con código y sección escrita: \"the Technical "
                "Specification (P22-ET-09-000-001-0), Section 8.1 - Minimum Scope of "
                "Factory Acceptance Tests\". Nunca el signo de sección.")
    bullet(doc, ("Resumen ejecutivo: ", "b"),
           "la línea de disposición dice qué debe hacer BW Water. No narra avances ni "
           "cierres.")
    bullet(doc, ("Status: ", "b"),
           "de dos a cuatro frases con qué es el documento, si cubre el alcance y el "
           "hallazgo que fija el código. Las observaciones, una por una, viven en el "
           "PDF anotado; el bloque Action cita su rango de IDs.")
    bullet(doc, ("Action: ", "b"),
           "dejar solo el párrafo del código que corresponde y borrar los otros dos con "
           "su nota.")
    bullet(doc, ("Attachments: ", "b"),
           "dejar una de las dos frases. Con Código 1, borrar también la tabla.")
    bullet(doc, ("Solo el FAT: ", "b"),
           "el N40 es ejecutivo y no lleva pendientes de transmittales anteriores ni "
           "otros documentos. Lo que aparezca fuera del procedimiento se anota para el "
           "próximo transmittal.")
    bullet(doc, ("Frases a evitar: ", "b"),
           "\"not X but Y\", \"here lies the tension\", \"it is that simple\", \"there "
           "is no turning back\". Tampoco nombres del personal interno de BW Water.")

    # ---- 9
    doc.add_heading("9. Chequeo antes de emitir", level=1)
    for item in [
        "Buscar \"[\" con Ctrl+F: cero resultados.",
        "Ningún texto resaltado en amarillo: seleccionar todo y aplicar Resaltado, Sin "
        "color, después de revisar.",
        "Cambiar la fecha de generación (17/09/2026) por la de emisión, en el "
        "encabezado y en el cajetín de la portada.",
        "Tabla de contenido: clic derecho, Actualizar campo, Actualizar toda la tabla.",
        "Los IDs del Word coinciden con los del PDF anotado, y el nombre del anotado "
        "con la tabla de Attachments.",
        "Guardar como PDF desde Word (Archivo, Guardar como, PDF).",
    ]:
        bullet(doc, "☐ " + item)

    # ---- 10
    doc.add_heading("10. Correo de cobertura", level=1)
    bullet(doc, ("Para: ", "b"), "Eduardo Yamauchi (BW Water).")
    bullet(doc, ("CC: ", "b"), "Jeryl F. Regulacion (BW Water); Luis Rivera (ADASA).")
    bullet(doc, ("Asunto: ", "b"),
           "TALTAL - Transmittal N40 - Module Factory Acceptance Test Procedure")
    bullet(doc, ("Adjuntos: ", "b"),
           "el PDF del transmittal y, si existe, el _CC_ADASA.pdf.")
    para(doc, "Texto sugerido para pegar en Outlook (reemplazar los corchetes):")
    for linea in [
        "Dear Eduardo,",
        "Please find attached Transmittal N40 (P22-TM-09-000-040-0) with ADASA's review "
        "of the module Factory Acceptance Test procedure received with submittal "
        "[25007-00XX].",
        "The procedure is returned at Code [N — ...]. [One sentence with the action "
        "that sets the code, for example: it is to be re-issued as Rev B, and the "
        "Factory Acceptance Test cannot formally open until it is approved.]",
        "[Only with Code 2 or 3:] The observations are itemised in the annotated PDF "
        "attached.",
        "Best regards,",
        "Victor Gutierrez",
        "[position] - Aguas Antofagasta",
    ]:
        p = para(doc, linea)
        p.paragraph_format.left_indent = Inches(0.4)
        for r in p.runs:
            r.italic = True

    # ---- 11
    doc.add_heading("11. Lo que queda para Luis al volver", level=1)
    para(doc, "Victor no necesita hacer esto; queda anotado para el cierre del ciclo.")
    bullet(doc, "Actualizar el Master Deliverable Register: el procedimiento FAT del "
                "módulo pasa a entregado, con su código y el veredicto del N40.")
    bullet(doc, "Registro de compromisos: cerrar o ajustar PRG-50 y conciliar BV-13 y "
                "el hito H-03 con la fecha real del FAT.")
    bullet(doc, "Registrar el envío en la bitácora del README y reconciliar el .md del "
                "transmittal con el Word emitido.")

    fijar_idioma_documento(doc)
    cp = doc.core_properties
    cp.title = "Transmittal N40 - Guía de revisión del procedimiento FAT (interno)"
    cp.author = "Luis Rivera Gonzalez"
    cp.last_modified_by = "Luis Rivera Gonzalez"
    cp.comments = "Aguas de Antofagasta S.A. - uso interno"
    cp.language = LANG
    doc.save(OUTPUT)
    print(f"Guia generada: {OUTPUT}")


if __name__ == "__main__":
    crear_guia()
