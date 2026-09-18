#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
crear_transmittal.py — TRANSMITTAL N40 ADASA-BW_WATER (P22-TM-09-000-040-0).

Submittals 25007-0093 (ENTREGA 93, recibida el 14-Sep-2026), 25007-0094 (ENTREGA 94,
17-Sep-2026) y 25007-0095 (ENTREGA 95, 18-Sep-2026). Catorce documentos.

VEREDICTO GLOBAL: 3 — To be revised. 12 Codigo 1, 1 Codigo 2, 1 Codigo 3.
Determinante: el Factory Acceptance Test Procedure Rev A del modulo (P22-BA-09-000-017),
primera emision, que no es el procedimiento detallado de ET Section 8.1, BAE Clause 31 e
ITP row 7.1 (Hold Point de ADASA).

DECISIONES DEL USUARIO (18-Sep-2026):
  - Flujo estandar: este script + CC_ADASA con doc-annotator; autoria fija de la skill
    (Luis Rivera / Luis Rivera / Victor Gutierrez). La plantilla manual del 17-Sep queda
    en _plantilla_manual_superada_17sep/.
  - SIN seccion de pendientes de transmittals anteriores (pasan al N41).
  - Rev 0 de E93/E94: siempre Codigo 1 mas NOTE por lo no levantado.
  - Schematic Rev C (IFA): Codigo 2 reiterado, tercer ciclo, reafirmando el terminal
    2711P-T10C22D9P del datasheet PLC/HMI aprobado.
  - Rev B del FAT al viernes 25-Sep-2026 (cinco dias habiles desde la emision).

ESTRUCTURA (cuatro secciones): 1 Executive Summary; 2 Observations by Document (2.1 a
2.6, la 2.6 agrupa los nueve Codigo 1 sin condicion de la E94 en una tabla); 3
Attachments (dos CC_ADASA, adjuntos directos al correo, sin enlace Synology); 4
Response Summary.

Reglas: sin numeros manuales en add_heading (el template numera); sin signo de seccion;
sin nombres internos (Van Doorn, PRG-xx, BV-xx). Fuente unica del texto:
P22-TM-09-000-040-0_TRANSMITTAL.md (este script es su espejo en Python).
"""
import os
import sys

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
from ejemplo_documento import (  # noqa: E402
    crear_documento_adasa,
    aplicar_arial_12,
    add_simple_table,
    add_bullet as _add_bullet_native,
    set_repeat_table_header,
)
from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "TRANSMITTAL N40 ADASA-BW_WATER.docx")


# --------------------------------------------------------------------- helpers
def add_para(doc, runs):
    p = doc.add_paragraph()
    for item in runs:
        texto = item[0]
        fmt = item[1] if len(item) > 1 else {}
        r = p.add_run(texto)
        if fmt.get("bold"):
            r.bold = True
    return aplicar_arial_12(p)


def add_bullet(doc, text):
    return _add_bullet_native(doc, text, size=11, space_after_pt=12)


def _fijar_anchos(tabla, anchos):
    """El ancho lo manda el w:tblGrid con layout fijo; cell.width solo no basta."""
    tbl = tabla._tbl
    layout = tbl.tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl.tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.find(qn("w:tblGrid"))
    for col, ancho in zip(grid.findall(qn("w:gridCol")), anchos):
        col.set(qn("w:w"), str(int(ancho * 1440)))
    for row in tabla.rows:
        for celda, ancho in zip(row.cells, anchos):
            celda.width = Inches(ancho)


def _no_partir_filas(tabla):
    for row in tabla.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def add_table(doc, headers, filas, anchos, size=9):
    t = add_simple_table(doc, [tuple(headers)] + [tuple(f) for f in filas])
    set_repeat_table_header(t.rows[0])
    _no_partir_filas(t)
    _fijar_anchos(t, anchos)
    from docx.shared import Pt
    for row in t.rows:
        for celda in row.cells:
            for para in celda.paragraphs:
                for run in para.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(size)
    return t


# --------------------------------------------------------------------- contenido
VERDICT = ("Submittals 25007-0093, 25007-0094 and 25007-0095, fourteen documents. Tally: 12 Code 1, 1 Code 2, 1 Code 3, no Code 4. One document returns to revision: the Factory Acceptance Test Procedure of the module.")

DISPOSITION_LEAD = "Disposition at a glance:"
DISPOSITION = [
    ("Factory Acceptance Test Procedure Rev A — Code 3. ",
     "Re-issue as Rev B by Friday 25 September 2026 with the step-by-step test sheets, the acceptance values and the blank record forms for activities 5.1 to 5.8."),
    ("PLC/LCP Schematic Diagram Rev C — Code 2. ",
     "Replace the operator terminal of the bill of material with the 2711P-T10C22D9P at Rev 0; the approved PLC and HMI Datasheet is not to be revised to the terminal installed."),
    ("Instrument Location Layout Rev 0 — Code 1. ",
     "Read ISSUED FOR CONSTRUCTION in the drawing status cell when the file is next touched."),
    ("Datasheet of CIP Tank Rev 0, Line List Rev 2 and the nine datasheets and lists of submittal 25007-0094 — Code 1. ", "No action."),
]

WHY_LEAD = "Why Code 3 — Factory Acceptance Test Procedure Rev A:"
WHY = [
    "It is not the detailed procedure that the Technical Specification (P22-ET-09-000-001-0), Section 8.1 - Minimum Scope of Factory Acceptance Tests, Clause 31 of the BAE and row 7.1 of the Inspection and Test Plan P22-BA-09-000-004 Rev 0 require. Its ten pages restate the rows of the Inspection and Test Plan as headings, with no test step, no acceptance value and none of the nine record forms that its own Section 8 lists.",
    "It omits the loop checks of the field instruments to the PLC and the simulated fault scenarios that Section 8.1 makes mandatory, and it does not say against which control documents the simulation is run.",
    "Rows 7.3, 7.4 and 7.5 of the Inspection and Test Plan refer the key points, the test sequence and the simulation detail to this procedure. The procedure refers them back to approved drawings, datasheets and an approved FAT sequence that is not in it, so no document fixes the criterion.",
]

HOLD = ("Row 7.1 of the Inspection and Test Plan is an ADASA Hold Point: the Factory Acceptance Test cannot formally open until the procedure is approved, and the third-party inspector will be given the approved procedure, not this draft.")

# (titulo, codigo, status, action_lead, action)
SUBSECCIONES = [
    (
        "Factory Acceptance Test Procedure Rev A — P22-BA-09-000-017",
        "Response Code: 3 — To be revised",
        "This is the first submission of the Factory Acceptance Test procedure of the module, required by the Technical Specification (P22-ET-09-000-001-0), Section 8.1 - Minimum Scope of Factory Acceptance Tests, and by row 7.1 of the Inspection and Test Plan P22-BA-09-000-004 Rev 0, where its approval is an ADASA Hold Point. The document covers the right activities, 5.1 to 5.8 in the order of rows 7.2 to 7.9 of the Inspection and Test Plan, and correctly defines the FAT as a test without process fluid. It is not yet a procedure that can be executed or witnessed. Every acceptance criterion reads \"as per approved drawings\" or \"in accordance with approved functional design\", and the dry functional test refers to an approved FAT sequence that is not in the document. The control system test has no loop check list, no simulation method and no fault scenario. The SEC verification names no certificate, the UT baseline cites no procedure, and the nine record forms that Section 8 declares are not attached. The procedure also does not define the responsibilities that its Section 1 promises, and does not carry the hydrostatic tests as the prerequisite that Section 8.1 sets. Itemised in P22-BA-09-000-017_A_FAT_Procedure_CC_ADASA.pdf.",
        "Action — re-issue as Rev B by Friday 25 September 2026: ",
        "complete the procedure with the step-by-step test sheets, the acceptance values and the blank record forms for each activity, and with the loop check list and the simulation and fault scenarios against the approved control documents. Add the hydrostatic prerequisite, the responsibility matrix and the references by code and revision (OBS-01 to OBS-10, NOTE-01 and NOTE-02 on the annotated PDF). The FAT cannot formally open until this procedure is approved.",
    ),
    (
        "PLC/LCP Schematic Diagram Rev C — P22-CD-09-008-002",
        "Response Code: 2 — Approved as noted",
        "Of the two observations of Transmittal N37, one closes: the CIP heater running feedback, item 109 of the I/O List Rev 6, now has a terminal, input 7 of module -A3 on sheet 37, tag XB002. The second does not close, for the third time. Row 11 of the bill of material on sheet 28 still reads 2711P-T10C21D8S, and the comment sheet of this revision replies that the -D8S \"is fully sufficient\" and that BW Water will revise the PLC and HMI Datasheet P22-ET-09-008-001 to that model. ADASA declared the 2711P-T10C22D9P binding at Transmittal N30 on the basis of that datasheet, approved at Rev 0 for construction, which specifies two Ethernet RJ45 ports and 1 GB of memory. An approved datasheet is not revised to match the terminal installed; the terminal is brought to the datasheet. The comment sheet also cites the Outline drawing under the datasheet code. One more point of form: the tags XA007, XA008 and XB002 of inputs 5 to 7 on sheet 37 are PDF comments laid over the drawing, not drawing content. Itemised in P22-CD-09-008-002_C_PLC_LCP_Schematic_CC_ADASA.pdf.",
        "Action to issue at IFC Rev 0 — no new schematic revision required: ",
        "replace row 11 of the bill of material with the 2711P-T10C22D9P of the approved datasheet, draw the three input tags in the schematic, and cite each document by its own code and revision on the comment sheet (OBS-01, NOTE-01 and NOTE-02 on the annotated PDF). ADASA accepts the schematic on the basis that the terminal of the approved datasheet is the one supplied and wired.",
    ),
    (
        "Instrument Location Layout Rev 0 — P22-DWG-09-008-001",
        "Response Code: 1 — Approved",
        "Both observations of Transmittal N37 close. The revision block now describes each issue, the Rev C row reads its April date again, and the consolidated comment sheet is bound as sheet 4 of the drawing, which is where the change description of each issue lives. The notes box declares the drawings the layout is built on: P&ID Rev 0, Instrument List Rev F and Equipment Layout Rev D, each by code. One point of form remains and does not change the drawing: the drawing status cell of both sheets reads ISSUED FOR APPROVAL while the Rev 0 row reads ISSUED FOR CONSTRUCTION, and that row is spelled CONTRUCTION. The code is 1 and not 3 because the point is documentary and the drawing is issued for construction with the content ADASA approved.",
        "Action: none on this document — accepted at Rev 0. ",
        "Read ISSUED FOR CONSTRUCTION in the drawing status cell, and spell it so, when the file is next touched.",
    ),
    (
        "Datasheet of CIP Tank Rev 0 — P22-ET-09-009-009",
        "Response Code: 1 — Approved",
        "The datasheet is re-issued aligned to the GA of CIP Flushing Tank P22-DWG-09-005-014. The nozzle schedule now carries the 533 mm manhole on top, N42 as a DN25 spare and N97 as the DN40 sensor connection, and the comment sheet declares that the datasheet was revised to follow the GA. That answers the question raised at Transmittal N34 on which of the two documents governs: the GA does, and the datasheet now agrees with it.",
        "Action: none — accepted at Rev 0.", "",
    ),
    (
        "Line List Rev 2 — P22-LI-09-009-003",
        "Response Code: 1 — Approved",
        "The condition of Transmittal N38 is met: CP-SSD-DN80-09-049, the first-stage CIP reject at 54 cubic metres per hour, is DN80, and CP-SSD-DN65-09-048, the second-stage CIP reject at 36 cubic metres per hour, is DN65, and the revision history states the correction. The RO brine discharge line DA-PVC-DN65-09-016 keeps its 2 bar design and 3 bar test pressure.",
        "Action: none — accepted; issued for construction.", "",
    ),
]

REV0_TITLE = ("Datasheets and lists issued at Revision 0 — submittal 25007-0094, nine documents")
REV0_CODE = "Response Code: 1 — Approved, all nine"
REV0_STATUS = ("Each document was reviewed against the comments ADASA raised at the revision it approved, and against that revision, as a difference check. None carries an open condition and none changes a value ADASA approved.")
REV0_HEADERS = ["Document code", "Title", "Previous condition", "Verification at Rev 0"]
REV0_ROWS = [
    ("P22-LI-09-009-001", "Utility Consumption List", "None",
     "Content identical to Rev C"),
    ("P22-LI-09-009-002", "Chemical Consumption List", "None",
     "Content identical to Rev A"),
    ("P22-ET-09-009-001", "Datasheet of UHPRO System", "None",
     "Stage 1: LG SW 400 SR in six BPV-8-1200-SP-7 vessels; stage 2: LG SW 400R G2 UHP in four BPV-8-1800-SP-7 vessels; seven elements per vessel, as at Rev B"),
    ("P22-ET-09-009-003", "Datasheet of CIP / Flushing Pump",
     "Motor compatibility (Transmittal N2)",
     "One motor: IEC 160MB, 11 kW, two poles, 2940 to 2950 rpm, external VFD, 10.11 kW absorbed at duty; consistent with the pump curve"),
    ("P22-ET-09-009-004", "Datasheet of Antiscalant Dosing Pump", "None",
     "ProMinent GMXa 1602, 0.02 L/h duty, 2.3 L/h maximum, 16 barg, as at Rev B"),
    ("P22-ET-09-009-005", "Datasheet of RO Cartridge Filter", "None",
     "Adds the EPDM gasket row and restates the filtration rate per square metre; same cartridge and housing as Rev E"),
    ("P22-ET-09-009-006", "Datasheet of CIP Cartridge Filter",
     "Cover gasket in EPDM (Transmittal N24)",
     "Housing model 31SBFX4-040A-E, where the E suffix is the EPDM cover gasket of the ordering guide; gasket row reads EPDM"),
    ("P22-ET-09-009-010", "Datasheet of Antiscalant Dosing Tank",
     "P&ID capacity and dosing rate (Transmittal N4)",
     "The P&ID Rev 0 reads 0.34 cubic metres total and 0.27 effective for TK-09-002, as the datasheet does; the 0.5 ppm dosing rate was accepted in principle under Technical Query P22-CT-09-000-001, whose residual points remain in that query"),
    ("P22-ET-09-009-011", "Datasheet of CIP Tank Heater", "None",
     "Quantic Logic flange immersion heater, 20 kW, 380 V, SS316 wetted parts, as at Rev B"),
]
REV0_ACTION = "Action: none — accepted; issued for construction."

ATTACHMENTS = [
    ("Factory Acceptance Test Procedure Rev A", "3",
     "P22-BA-09-000-017_A_FAT_Procedure_CC_ADASA.pdf"),
    ("PLC/LCP Schematic Diagram Rev C", "2",
     "P22-CD-09-008-002_C_PLC_LCP_Schematic_CC_ADASA.pdf"),
]
ATTACH_NOTE = ("The two documents with open observations carry annotated PDFs, attached to the cover e-mail. The twelve Code 1 documents carry no annotated PDF.")

RESPONSE_SUMMARY = [
    ("P22-BA-09-000-017", "Factory Acceptance Test Procedure", "A", "25007-0095",
     "3 — To be revised"),
    ("P22-CD-09-008-002", "PLC/LCP Schematic Diagram", "C", "25007-0093",
     "2 — Approved as noted"),
    ("P22-DWG-09-008-001", "Instrument Location Layout", "0", "25007-0093", "1 — Approved"),
    ("P22-ET-09-009-009", "Datasheet of CIP Tank", "0", "25007-0093", "1 — Approved"),
    ("P22-LI-09-009-001", "Utility Consumption List", "0", "25007-0094", "1 — Approved"),
    ("P22-LI-09-009-002", "Chemical Consumption List", "0", "25007-0094", "1 — Approved"),
    ("P22-LI-09-009-003", "Line List", "2", "25007-0094", "1 — Approved"),
    ("P22-ET-09-009-001", "Datasheet of UHPRO System", "0", "25007-0094", "1 — Approved"),
    ("P22-ET-09-009-003", "Datasheet of CIP / Flushing Pump", "0", "25007-0094",
     "1 — Approved"),
    ("P22-ET-09-009-004", "Datasheet of Antiscalant Dosing Pump", "0", "25007-0094",
     "1 — Approved"),
    ("P22-ET-09-009-005", "Datasheet of RO Cartridge Filter", "0", "25007-0094",
     "1 — Approved"),
    ("P22-ET-09-009-006", "Datasheet of CIP Cartridge Filter", "0", "25007-0094",
     "1 — Approved"),
    ("P22-ET-09-009-010", "Datasheet of Antiscalant Dosing Tank", "0", "25007-0094",
     "1 — Approved"),
    ("P22-ET-09-009-011", "Datasheet of CIP Tank Heater", "0", "25007-0094",
     "1 — Approved"),
]


if __name__ == "__main__":
    tally = {"1": 0, "2": 0, "3": 0}
    for r in RESPONSE_SUMMARY:
        tally[r[4][0]] += 1
    assert (tally["1"], tally["2"], tally["3"]) == (12, 1, 1), tally
    assert len(RESPONSE_SUMMARY) == 14

    crear_documento_adasa(
        titulo="TECHNICAL REVIEW TRANSMITTAL N40 — SECOND STAGE RO BRINE MODULE",
        codigo="P22-TM-09-000-040-0",
        output_filename=OUTPUT,
        incluir_toc=True,
    )
    doc = Document(OUTPUT)

    # ---------------------------------------------------------- 1. Executive
    doc.add_heading("EXECUTIVE SUMMARY", level=1)
    add_para(doc, [("TRANSMITTAL VERDICT: 3 — To be revised. ", {"bold": True}),
                   (VERDICT,)])
    add_para(doc, [(DISPOSITION_LEAD, {"bold": True})])
    for lead, rest in DISPOSITION:
        add_bullet(doc, lead + rest)
    add_para(doc, [(WHY_LEAD, {"bold": True})])
    for punto in WHY:
        add_bullet(doc, punto)
    add_para(doc, [(HOLD,)])

    # ------------------------------------------------------- 2. Observations
    doc.add_heading("OBSERVATIONS BY DOCUMENT", level=1)
    for titulo, codigo, status, action_lead, action in SUBSECCIONES:
        doc.add_heading(titulo, level=2)
        add_para(doc, [(codigo, {"bold": True})])
        add_para(doc, [("Status.", {"bold": True}), (" " + status,)])
        add_para(doc, [(action_lead, {"bold": True}), (action,)])
    doc.add_heading(REV0_TITLE, level=2)
    add_para(doc, [(REV0_CODE, {"bold": True})])
    add_para(doc, [("Status.", {"bold": True}), (" " + REV0_STATUS,)])
    add_table(doc, REV0_HEADERS, REV0_ROWS, [1.3, 1.5, 1.4, 2.3], size=8)
    add_para(doc, [(REV0_ACTION, {"bold": True})])

    # -------------------------------------------------------- 3. Attachments
    doc.add_heading("ATTACHMENTS", level=1)
    add_table(doc, ["Document", "Code", "Annotated PDF"], ATTACHMENTS, [2.1, 0.6, 3.8],
              size=9)
    add_para(doc, [(ATTACH_NOTE,)])

    # ----------------------------------------------------- 4. Response Summary
    doc.add_heading("RESPONSE SUMMARY", level=1)
    add_table(doc, ["Document Code", "Title", "Rev", "Submittal", "Response Code"],
              RESPONSE_SUMMARY, [1.45, 1.95, 0.4, 0.95, 1.75], size=9)

    doc.save(OUTPUT)
    print("Documento generado: " + OUTPUT)
    print("  tally 1/2/3 = %s/%s/%s" % (tally["1"], tally["2"], tally["3"]))
