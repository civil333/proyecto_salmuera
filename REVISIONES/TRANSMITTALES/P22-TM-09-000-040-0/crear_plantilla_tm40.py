#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera la PLANTILLA del Transmittal N40 (P22-TM-09-000-040-0).

Uso: revision del procedimiento de Factory Acceptance Test del modulo, que BW
Water anuncio para el 18-Sep-2026 (minuta P22-MI-10-000-002-0). La revision la
hace y la emite Victor Gutierrez, EN WORD A MANO, durante la ausencia de Luis
Rivera. Este script solo deja el Word armado:

  - texto fijo = lo que no depende del resultado de la revision;
  - marcadores [entre corchetes, resaltados en amarillo] = lo que Victor completa
    o borra. La guia interna _GUIA_REVISION_FAT_N40_NO_ENVIAR.docx explica cada uno.

Decisiones del usuario (17-Sep): N40 y no N39 (el N39 salio el 09-Sep); solo el
procedimiento FAT; cajetin con Victor Gutierrez en las tres firmas; PDF anotado
hecho a mano en Acrobat.

Version ejecutiva, solo el FAT (usuario, 17-Sep): cuatro secciones, Executive
Summary, Observations by Document, Attachments y Response Summary. NO lleva la
seccion de pendientes de transmittales anteriores ni la "Issue for Construction"
del N39, y no importa auditar_documentos_bw.

Reglas: sin numeros manuales en add_heading (el template numera); sin signo de
seccion; sin nombres internos (Van Doorn, codigos PRG/BV del registro).

Si Victor ya empezo a llenar el Word, NO regenerar: el script sobrescribe la
plantilla, no el archivo que el guarde con otro nombre.
"""

import os
import re
import sys

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
sys.path.insert(0, os.path.expanduser("~/.claude/skills/_shared"))

from ejemplo_documento import (  # noqa: E402
    crear_documento_adasa,
    add_simple_table,
    add_bullet as _add_bullet_native,
    set_repeat_table_header,
)
import docx_metadata  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "TRANSMITTAL N40 ADASA-BW_WATER_PLANTILLA.docx")
FIRMA = "Victor Gutierrez"

# Un segmento es (texto, fmt). fmt: {"bold": True} y/o {"mark": True}.
B = {"bold": True}
M = {"mark": True}
BM = {"bold": True, "mark": True}


# --------------------------------------------------------------------- helpers
def _formatear_run(run, fmt, size):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if fmt.get("bold"):
        run.bold = True
    if fmt.get("mark"):
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def add_para(doc, segmentos, size=12):
    """Parrafo de cuerpo en Arial 12, como el N39 (aplicar_arial_12)."""
    p = doc.add_paragraph()
    for texto, *fmt in segmentos:
        _formatear_run(p.add_run(texto), fmt[0] if fmt else {}, size)
    return p


def add_bullet(doc, segmentos):
    p = _add_bullet_native(doc, "", size=11, space_after_pt=12)
    for texto, *fmt in segmentos:
        _formatear_run(p.add_run(texto), fmt[0] if fmt else {}, 11)
    return p


def add_heading_con_marcador(doc, segmentos, level):
    h = doc.add_heading("", level=level)
    for texto, *fmt in segmentos:
        run = h.add_run(texto)
        if fmt and fmt[0].get("mark"):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return h


def _fijar_anchos(tabla, anchos):
    """Copiado del N39: manda el w:tblGrid con layout fijo, no cell.width."""
    tbl = tabla._tbl
    layout = tbl.tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl.tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.find(qn("w:tblGrid"))
    for col, ancho in zip(grid.findall(qn("w:gridCol")), anchos):
        col.set(qn("w:w"), str(int(ancho * 1440)))
    for row in tabla.rows:
        for celda, ancho in zip(row.cells, anchos):
            celda.width = Inches(ancho)


def _resaltar_marcadores(para, size, bold=False):
    """Reescribe los runs del parrafo resaltando solo los tramos [entre corchetes],
    para que el texto fijo de la misma celda quede sin resaltar."""
    texto = "".join(r.text for r in para.runs)
    for r in list(para.runs):
        r._element.getparent().remove(r._element)
    for tramo in re.split(r"(\[[^\]]*\])", texto):
        if tramo:
            fmt = {"bold": bold, "mark": tramo.startswith("[")}
            _formatear_run(para.add_run(tramo), fmt, size)


def add_tabla(doc, filas, anchos, size=10):
    """Tabla con encabezado repetido, filas sin partir, anchos fijos y celdas
    alineadas a la izquierda (el justificado abre huecos en columnas angostas).
    Solo los marcadores [..] quedan resaltados."""
    t = add_simple_table(doc, [tuple(f) for f in filas])
    set_repeat_table_header(t.rows[0])
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    _fijar_anchos(t, anchos)
    for i, row in enumerate(t.rows):
        for celda in row.cells:
            for para in celda.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if i == 0:
                    para.paragraph_format.keep_with_next = True
                _resaltar_marcadores(para, size, bold=(i == 0))
    return t


def crear_plantilla():
    crear_documento_adasa(
        titulo="TECHNICAL REVIEW TRANSMITTAL N40 — SECOND STAGE RO BRINE MODULE",
        codigo="P22-TM-09-000-040-0",
        preparado_por=FIRMA,
        revisado_por=FIRMA,
        aprobado_por=FIRMA,
        output_filename=OUTPUT,
        incluir_toc=True,
    )
    doc = Document(OUTPUT)

    # ------------------------------------------------------ 1. Executive Summary
    doc.add_heading("EXECUTIVE SUMMARY", level=1)
    add_para(doc, [
        ("TRANSMITTAL VERDICT: ", B),
        ("[1 — Approved / 2 — Approved as noted / 3 — To be revised]", BM),
        (". Submittal ",), ("[25007-00XX]", M),
        (", one document: the Factory Acceptance Test procedure of the module. ",),
        ("[One sentence: whether the document returns to revision.]", M),
    ])
    add_para(doc, [("Disposition at a glance:", B)])
    add_bullet(doc, [
        ("Module Factory Acceptance Test Procedure Rev ", B), ("[X]", BM),
        (" — Code ", B), ("[N]", BM), (". ", B),
        ("[One clause with the action BW Water must take, for example: re-issue as "
         "Rev B / issue directly at IFC Rev 0.]", M),
    ])
    add_para(doc, [
        ("Why Code ", B), ("[N]", BM), (":", B),
        (" [Delete this line and its bullets if the verdict is Code 1.]", M),
    ])
    add_bullet(doc, [("[Reason that sets the code, one or two sentences.]", M)])
    add_bullet(doc, [("[Second reason, only if there is one.]", M)])

    # --------------------------------------------- 2. Observations by Document
    doc.add_heading("OBSERVATIONS BY DOCUMENT", level=1)
    add_heading_con_marcador(doc, [
        ("Module Factory Acceptance Test Procedure Rev ",), ("[X]", M),
        (" — ",), ("[document code]", M),
    ], level=2)
    add_para(doc, [
        ("Response Code: ", B),
        ("[1 — Approved / 2 — Approved as noted / 3 — To be revised]", BM),
    ])
    add_para(doc, [
        ("Status.", B),
        (" This is the first submission of the detailed Factory Acceptance Test "
         "procedure of the module, required by the Technical Specification "
         "(P22-ET-09-000-001-0), Section 8.1 - Minimum Scope of Factory Acceptance "
         "Tests, and by row 7.1 of the Inspection and Test Plan P22-BA-09-000-004 "
         "Rev 0, where its approval is an ADASA Hold Point. ",),
        ("[Two to four sentences: whether the procedure covers the scope of "
         "Section 8.1 and rows 7.1 to 7.9 of the Inspection and Test Plan, and the "
         "finding that sets the code. Do not list the observations one by one.]", M),
        (" Itemised in ",), ("[file name]", M), ("_CC_ADASA.pdf.",),
        (" [Delete this last sentence if the verdict is Code 1.]", M),
    ])
    add_para(doc, [("[Keep the one Action paragraph below that matches the Response "
                    "Code and delete the other two, together with this note.]", M)])
    add_para(doc, [
        ("Action: none on this document — accepted; issue directly at IFC Rev 0.", B),
    ])
    add_para(doc, [
        ("Action to issue at IFC Rev 0 — no new revision required: ", B),
        ("[one sentence with the changes to incorporate in the procedure itself]", M),
        (" (",), ("[OBS-01 to OBS-0X and NOTE-01]", M),
        (" on the annotated PDF). ADASA accepts the procedure on the basis that these "
         "changes are incorporated at Rev 0.",),
    ])
    add_para(doc, [
        ("Action — re-issue as Rev ", B), ("[X]", BM), (": ", B),
        ("[one sentence with what must be corrected]", M),
        (" (",), ("[OBS-01 to OBS-0X]", M),
        (" on the annotated PDF). The Factory Acceptance Test cannot formally open "
         "until this procedure is approved, since row 7.1 of the Inspection and Test "
         "Plan is an ADASA Hold Point.",),
    ])

    # --------------------------------------------------------- 3. Attachments
    doc.add_heading("ATTACHMENTS", level=1)
    add_tabla(doc, [
        ("Document", "Code", "Annotated PDF"),
        ("Module Factory Acceptance Test Procedure Rev [X]", "[N]",
         "[file name]_CC_ADASA.pdf"),
    ], anchos=[2.1, 0.6, 3.8])
    add_para(doc, [("[Keep one of the two sentences below and delete the other, "
                    "together with this note. With Code 1, also delete the table "
                    "above.]", M)])
    add_para(doc, [("The procedure carries open observations, and its annotated PDF "
                    "is attached.",)])
    add_para(doc, [("The procedure is approved without observations and carries no "
                    "annotated PDF.",)])

    # --------------------------------------------------- 4. Response Summary
    doc.add_heading("RESPONSE SUMMARY", level=1)
    add_tabla(doc, [
        ("Document Code", "Title", "Rev", "Submittal", "Response Code"),
        ("[document code]", "Module Factory Acceptance Test Procedure", "[X]",
         "[25007-00XX]", "[N — Approved / Approved as noted / To be revised]"),
    ], anchos=[1.45, 1.85, 0.5, 0.95, 1.75])

    # los titulos del template no traen "mantener con el siguiente" efectivo:
    # sin esto un titulo puede quedar solo al pie de pagina
    for para in doc.paragraphs:
        if para.style is not None and para.style.name.startswith("Heading"):
            para.paragraph_format.keep_with_next = True

    # add_simple_table deja un parrafo vacio al final: sobra
    ultimo = doc.paragraphs[-1]
    if not ultimo.text.strip():
        ultimo._element.getparent().remove(ultimo._element)

    docx_metadata.apply_core_properties(
        doc,
        title="Technical Review Transmittal N40 - Second Stage RO Brine Module",
        author=FIRMA,
        last_modified_by=FIRMA,
        subject="Second Stage RO Brine Module - Taltal",
        comments="Aguas de Antofagasta S.A.",
        language="en-US",
    )
    doc.save(OUTPUT)
    docx_metadata.fix_app_xml(OUTPUT, company="Aguas de Antofagasta S.A.")
    print(f"Plantilla generada: {OUTPUT}")


if __name__ == "__main__":
    crear_plantilla()
