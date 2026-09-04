#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TRANSMITTAL N31 ADASA-BW_WATER.
Submittals 25007-0072 (E72) y 25007-0074 (E73).
Fecha de emision: 10-Ago-2026.

ALCANCE ACOTADO POR DECISION DEL USUARIO: se atacan SOLO los comentarios
historicos que no fueron levantados. Todo hallazgo nuevo queda registrado en
los _HALLAZGOS_DETERMINISTAS.md de cada entrega y NO se emite.

Veredicto global: 2 - APPROVED AS NOTED.
Tally: 1 Code 1 + 3 Code 2 + 1 documento devuelto SIN CODIGO.

  2.1 Plant Control Philosophy Rev 0 (P22-BT-09-009-001)          SIN CODIGO
      Rev 0 ya emitido IFC: no se devuelve a revision. De las cinco
      condiciones del Codigo 2 del N28 solo cerro una. Las otras cuatro se
      declaran en texto, con la razon concreta de cada una. Sin CC_ADASA:
      un documento sin codigo no se anota, y por eso el detalle va integro
      en la subseccion (misma logica de excepcion declarada del N30).
  2.2 RO Vessel Hydrostatic Test Procedure Rev 0 (P22-BA-09-000-009) Code 1
      Las dos condiciones del N27 cumplidas. La vigencia de la calibracion
      del certificado 26993 y el rotulo del manometro del esquema van a la
      Seccion 3, no degradan.
  2.3 Tie-In Point Layout Rev B (P22-DWG-09-005-005)                Code 2
      Cierra el Codigo 3 del N7. 2 OBS.
  2.4 GA of SWRO System Skid Rev B (P22-DWG-09-005-008)             Code 2
      Cierra NOTE-11 y NOTE-12 del N15. 1 OBS (tres referencias mal citadas).
  2.5 HMI Display Screenshot Rev B (P22-LI-09-008-016)              Code 2
      Cierra 3 de 5 del N22. 7 OBS.

La ENTREGA 71 entra en la Seccion 3 SIN codigos nuevos, con una linea por
documento, citando el correo del 6 de agosto. Decision del usuario: que quede
en el transmittal y no solo en el correo, sin recodificar documentos IFC.

SIN numeros manuales en add_heading() (el template ADASA auto-numera H1/H2).
Referencias por NOMBRE de seccion. Cero simbolo de seccion. ADASA no emite
cartas: el instrumento es el transmittal y su correo de cobertura.
"""

import os
import sys

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))

from ejemplo_documento import (  # noqa: E402
    crear_documento_adasa,
    aplicar_arial_12,
    add_simple_table,
    add_bullet as _add_bullet_native,
)
from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "TRANSMITTAL N31 ADASA-BW_WATER.docx")

# Los tres PDF anotados van por enlace de descarga. Al correo se adjunta solo
# el PDF del transmittal. REEMPLAZAR por el enlace Synology real antes de emitir.
DOWNLOAD_LINK = (
    "https://lrg.synology.me:6501/d/s/19PutLaANldko09k2QUOUibPNpL295Ze/"
    "KyDEJBacXwQhWHVkM0m17KgQlqXF73CB-L7NAkDq8ag0"
)


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")
    rpr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_para(doc, runs):
    para = doc.add_paragraph()
    for item in runs:
        text = item[0]
        attrs = item[1] if len(item) > 1 else {}
        run = para.add_run(text)
        if attrs.get("bold"):
            run.bold = True
        if attrs.get("italic"):
            run.italic = True
    aplicar_arial_12(para)
    return para


def add_bullet(doc, text):
    _add_bullet_native(doc, text, size=11, space_after_pt=12)


# El resumen ejecutivo agrupa por SUBMITTAL de BW Water, que es el dato que el
# lector busca primero: a que envio suyo responde este transmittal.
DISPOSITION_0072 = [
    "Plant Control Philosophy Rev 0 — no response code. Four conditions of the "
    "Transmittal N28 approval remain open.",
    "RO Vessel Hydrostatic Test Procedure Rev 0 — Code 1. Confirm the calibration "
    "valid at the date of the test.",
    "Tie-In Point Layout Rev B — Code 2. State the design pressure at the brine feed "
    "tie-in point.",
    "GA of SWRO System Skid Rev B — Code 2. Correct three document references in the "
    "new notes.",
]

DISPOSITION_0074 = [
    "HMI Display Screenshot Rev B — Code 2. Add the electrical-variable readings and "
    "reconcile six tags.",
]

PRIORITY = (
    "The winding and bearing tags are reversed on the CIP pump in the Plant Control "
    "Philosophy and duplicated on the high-pressure pump in the HMI screens, in both cases "
    "against the Instrument List Rev E and the Alarm and Interlock List Rev C. It is a "
    "safety-sensor assignment on a 93 kW machine, and it holds the Factory Acceptance Test "
    "Procedure sign-off on motor temperature protection, open since Transmittal N27."
)

WHY_NO_CODE = (
    "because it was issued at Rev 0 for construction and ADASA does not return it to "
    "revision. Its four open points are stated in Section 2 for closure at the next issue "
    "of that document. Counting the four procedures of submittal 25007-0071, "
)

REV0_TALLY = (
    "ten conditions of ADASA approvals are now outstanding in documents already issued at "
    "Rev 0 for construction"
)

REVIEW_PERIOD = (
    "Clause 37.2 of the Special Administrative Conditions (BAE 12803) gives seven working "
    "days: Monday 17 August for 25007-0072 and Wednesday 19 August for 25007-0074. This "
    "transmittal is issued within both."
)

# ---------------------------------------------------------------------------
# Seccion 2 condensada: Response Code + Status corto + Action, citando el
# RANGO de IDs del PDF anotado (CLAUDE.md Seccion 3.2). Excepcion declarada en
# 2.1, que no lleva CC_ADASA por ir sin codigo.
# ---------------------------------------------------------------------------

OPEN_POINTS_CONTROL = [
    "The temperature-sensor mapping was corrected on the high-pressure pump and not on the "
    "CIP pump. Page 36 now reads winding TE-09-001 and TE-09-003 and bearing TE-09-002 and "
    "TE-09-004, which is correct. Page 55, Section 3.4.2, still reads item 13 CIP Pump "
    "Bearing Temperature Sensor as TE-09-003 and item 14 CIP Pump Winding Temperature Sensor "
    "as TE-09-004, which is the reverse of the Instrument List Rev E and of the Alarm and "
    "Interlock List Rev C. The document therefore also contradicts its own page 36.",
    "The vibration alarm and trip pairs still differ from the Alarm and Interlock List Rev C. "
    "Page 39 states a high alarm of 7.1 mm/s for the high-pressure pump where the list sets "
    "7.0 for VIT-09-001. Pages 41 and 43 state a high-high trip of 7.1 mm/s for both "
    "turbochargers where the list sets 6.0 for VT-09-002. The 7.1 against 6.0 difference is "
    "the one raised at Transmittal N28.",
    "The winding trip is no longer supported by a stated insulation class. Page 38 now reads "
    "a winding alarm of 120 degrees Celsius and a trip of 140, which match the list. The "
    "Class B citation that made the earlier 155 degree trip inconsistent was deleted rather "
    "than confirmed, and no insulation class appears anywhere in the document, so the 140 "
    "degree trip has nothing declared to support it.",
    "The issued child documents are still not pinned by code and revision. The reference "
    "table on page 5 continues to read SEPARATE DOCUMENT for the Controls and Sequence Chart "
    "and for the Alarm and Control Setpoint List, although both have since been issued as "
    "P22-LI-09-008-017 Rev A and P22-LI-09-008-015 Rev C.",
]

SECTIONS = [
    dict(
        heading="Plant Control Philosophy Rev 0 — P22-BT-09-009-001",
        code="Response Code: none issued",
        paras=[
            [("Status.", {"bold": True}),
             (" The document was issued at Rev 0 for construction, so no response code is "
              "stated. Of the five conditions under which ADASA approved Rev E at Transmittal "
              "N28, one is closed: the sustained low-pressure protection values were removed "
              "from the body and referred to the Alarm and Interlock List, which was one of "
              "the two routes offered. Four remain open, and the comment sheet records all "
              "four as updated and aligned. They are set out below with what ADASA found, so "
              "that each can be closed at the next issue of this document.",)],
        ],
        bullets_lead="Open points:",
        bullets=OPEN_POINTS_CONTROL,
        closing=[
            ("Action — close the four points above at the next issue of this document.",
             {"bold": True}),
            (" No annotated PDF accompanies this document, because it carries no response "
             "code; the points are stated in full here.",),
        ],
    ),
    dict(
        heading="RO Vessel Hydrostatic Test Procedure Rev 0 — P22-BA-09-000-009",
        code="Response Code: 1 — Approved",
        paras=[
            [("Status.", {"bold": True}),
             (" Both conditions of the Transmittal N27 approval are met. The binding test "
              "pressures now appear on the face of the procedure, at 91.0 bar for the "
              "BPV81200SP7 and 136.5 bar for the BPV81800SP7. The two non-conforming "
              "calibration certificates were withdrawn and replaced by a single certificate, "
              "number 26993, for an analogue gauge with a measurement range of 0 to 250 bar, "
              "which for the 136.5 bar test is 1.83 times the test pressure and therefore "
              "sits inside the 1.5 to 4 times window this procedure sets for itself.",)],
            [("Action: none on this document.", {"bold": True}),
             (" Two items are tracked in Section 3: the calibration validity of certificate "
              "26993 at the date of the test, and the gauge label on the test schematic, "
              "which still reads 0 to 160 bar.",)],
        ],
    ),
    dict(
        heading="Tie-In Point Layout Rev B — P22-DWG-09-005-005",
        code="Response Code: 2 — Approved as noted",
        paras=[
            [("Status.", {"bold": True}),
             (" Rev B closes the substance of the Code 3 issued at Transmittal N7, which had "
              "made this the longest-standing open document of the project. The tie-in points "
              "are redrawn over the consolidated layout, the line identifiers are now "
              "presented as the tie-in tags carried on the Piping and Instrumentation "
              "Diagram, and the elevation information that was missing has arrived in full: "
              "three new section views and a centreline elevation column with a value for "
              "each of the five tie-in points. Two items remain, itemised in "
              "P22-DWG-09-005-005_B_Tie_In_Point_Layout_CC_ADASA.pdf.",)],
            [("Action to issue at IFC Rev 0 — no new drawing revision required:",
              {"bold": True}),
             (" state the design pressure at the brine feed tie-in point TP-DA P8-001, so "
              "that compatibility with the ANSI 150# rating shown can be verified, and "
              "complete the flange class and flange standard cells of tie-in point TP-AS "
              "P11-001, which remain blank (OBS-01 and OBS-02 on the annotated PDF). ADASA's "
              "acceptance of the ANSI 150# rating at the battery limit is conditional on that "
              "design pressure being stated.",)],
        ],
    ),
    dict(
        heading="GA of SWRO System Skid Rev B — P22-DWG-09-005-008",
        code="Response Code: 2 — Approved as noted",
        paras=[
            [("Status.", {"bold": True}),
             (" Rev B delivers the two schedules requested at Transmittal N15. The notes "
              "block now states six pressure vessels in the first stage and four in the "
              "second, seven membrane elements per vessel with the model for each stage, and "
              "a super duplex stainless steel manifold at Class 900. The function-versus-tag "
              "matrices and the pressure, temperature and wall-thickness data are referred to "
              "the governing lists rather than repeated on the sheet, which ADASA accepts. "
              "The three references are wrong. "
              "Itemised in P22-DWG-09-005-008_B_GA_SWRO_Skid_CC_ADASA.pdf.",)],
            [("Action to issue at IFC Rev 0 — no new drawing revision required:",
              {"bold": True}),
             (" correct the three document references in the notes block — note 6 cites the "
              "Instrument List at Rev D where the current revision is Rev E, note 7 cites "
              "P22-LI-09-009-00 where the approved Line List is P22-LI-09-009-003, and note 8 "
              "cites P22-ET-09-006-01, which is not a valid document code in the project "
              "numbering (OBS-01 on the annotated PDF). As written, none of the three "
              "referred documents can be located from the sheet.",)],
        ],
    ),
    dict(
        heading="HMI Display Screenshot Rev B — P22-LI-09-008-016",
        code="Response Code: 2 — Approved as noted",
        paras=[
            [("Status.", {"bold": True}),
             (" Rev B closes three of the five observations of the Code 3 issued at "
              "Transmittal N22 and delivers the screen set in full: six process screens "
              "against two, with energy recovery and the brine path shown within the stage "
              "screens, plus trending, alarm summary, navigation and user access. The "
              "trending screen, the faceplate route by which the operator sets values and "
              "limits, and the evidence of ISA-101 conformance are all accepted. Two "
              "observations remain open, itemised in "
              "P22-LI-09-008-016_B_HMI_Display_Screenshot_CC_ADASA.pdf.",)],
            [("Action to issue at IFC Rev 0 — no new revision required:", {"bold": True}),
             (" add the electrical-variable meter readings — voltage, current and power — to "
              "the overview screen alongside the specific energy consumption it already "
              "displays, as the Technical Specification (P22-ET-09-000-001-0), Section 5.4 - "
              "Control System, requires, and express power in kilowatts; and reconcile six "
              "tags against the approved lists, naming BH-09-001 on the high-pressure feed "
              "pump, TE-09-001 on its winding, VE-09-005 on the make-up CIP branch, "
              "PIT-09-003 on the first-stage feed, FIT-09-004 on the reject line to brine "
              "drain, and LS-09-002 on the antiscalant tank low level of the overview screen "
              "(OBS-01 to OBS-07 on the annotated PDF). State also on which screen PIT-09-008 "
              "and FIT-09-002 are displayed, or why they are not. Full visual verification of "
              "the screens remains reserved for the Factory Acceptance Test.",)],
        ],
    ),
]

PENDING_OPEN = [
    ("TM N19",
     "Fabrication and testing dossier",
     "The dossier required by the Technical Specification, Section 7, has not been "
     "delivered; it gates items 8.3 and 8.4 of the Inspection and Testing Base Plan",
     "OPEN — no record delivered; the index arrived at Transmittal N30 and the dossier "
     "itself did not"),
    ("TM N30",
     "Shop fabrication set 25007-ME-PI-0901-0006 to -0016",
     "Returned as not received; threaded austenitic branch connections on super duplex "
     "lines, and the super duplex to PVC boundary without a specification break",
     "OPEN — confirm in writing whether any of these spools has already been fabricated, "
     "before the pressure test of 13 and 14 August"),
    ("TM N30",
     "Fabrication and Testing Dossier Index Rev A",
     "Four chapters absent and no document number, revision or inclusion status against "
     "any of its 27 lines",
     "OPEN — re-issue as Rev B"),
]

# Cierre del arrastre del Equipment Layout. La linea heredada del N30 decia
# "Equipment Layout Rev C (Code 3, cartridge filter orientation)", que se lee
# como si la orientacion siguiera en disputa. NO lo esta: el Piping Layout Rev C
# del submittal 25007-0070 dibuja el filtro vertical, verificado en el analisis
# de la E70 (_REVISION_62_OBSERVACIONES.md, NOTE-02) y nunca emitido, porque esa
# nota quedo entre las 48 observaciones descartadas del N30. Se emite aqui.
PENDING_CLOSED = (
    "The RO cartridge filter arrangement that has kept the Equipment Layout in Code 3 since "
    "Transmittal N22 is settled: the Piping Layout Rev C draws the filter vertical, on the "
    "plan of Sheet 1 and on Section 1-1 of Sheet 2, and ADASA accepts that arrangement. The "
    "Equipment Layout Rev C remains outstanding only to be re-issued in line with it, and "
    "its Code 3 stands until then."
)

# Los tres planos abiertos se reclaman con el DDSR del PROPIO proveedor, no con
# la redaccion heredada. En el reporte del 20-Jul la columna "Curr Rev" declara
# Equipment Layout D, GA Antiscalant C e Instrument Location D -por eso los TM
# N27 a N29 citaron esas letras: las copiaban de ahi-, pero el "Sent Date" de
# los tres sigue congelado en la revision que ADASA ya tiene, de modo que
# ninguna se envio. Las fechas prometidas (29-Jul, 30-Jul, 2-Ago) estan vencidas.
PENDING_SUMMARY = (
    "the Equipment Layout, which ADASA holds at Rev C, with the re-issue planned there for "
    "29 July after moving from 26 June and from 15 July; the GA of the Antiscalant Dosing "
    "Tank, held at Rev B, planned for 30 July; and the Instrument Location Layout, held at "
    "Rev C, planned for 2 August. None of the three has been submitted, and the Sent Date of "
    "all three still reads in that report as the revision ADASA already holds. The Factory "
    "Acceptance Test Procedure sign-off on motor temperature protection also remains held "
    "since Transmittal N27, dependent on the mapping described in the first subsection of "
    "Section 2."
)

RAISED_HERE = (
    "the calibration certificate 26993 attached to the RO Vessel Hydrostatic Test Procedure "
    "is dated 12 December 2025, while Clause 4 of that procedure sets a six-month calibration "
    "interval and the test was performed on 17 June 2026; please confirm the calibration "
    "valid at the date of the test, or provide the certificate that covers it. Separately, "
    "the test schematic of that procedure still labels the gauge as 0 to 160 bar, to be "
    "tidied at the next natural issue."
)

E71_LEAD = (
    "The five documents of submittal 25007-0071 were issued at Rev 0 for construction and "
    "were addressed by email on 6 August. Six conditions across four of those five documents "
    "remain open and are recorded here so that they are carried formally:"
)

E71_ITEMS = [
    "PMI Procedure Rev 0 — the acceptance section still refers approval and rejection to "
    "third-party refinery standards instead of conformity to UNS S32750, and the "
    "subcontractor annex has not been scoped to this project.",
    "Visual Procedure Rev 0 — no form is identified for recording the inspection, while row "
    "3.2 of the Inspection and Test Plan requires a visual report at a witness point.",
    "Painting Procedure Rev 0 — the inspection form still lacks the nominal dry film "
    "thickness and product per coat, and the Colour row remains blank against the required "
    "RAL 5012.",
    "HP and LP Pressure Test Procedure Rev 0 — clauses 5.5.2 and 5.6.3 still refer to the "
    "latest edition of ASME B31.3, so the edition governing the tests is not fixed.",
]

E71_TAIL = (
    "The UHPRO Structural Calculation Report Rev 0 incorporated both of its conditions and "
    "is not listed."
)

ATTACHMENTS = [
    ("Tie-In Point Layout Rev B (Code 2)",
     "P22-DWG-09-005-005_B_Tie_In_Point_Layout_CC_ADASA.pdf"),
    ("GA of SWRO System Skid Rev B (Code 2)",
     "P22-DWG-09-005-008_B_GA_SWRO_Skid_CC_ADASA.pdf"),
    ("HMI Display Screenshot Rev B (Code 2)",
     "P22-LI-09-008-016_B_HMI_Display_Screenshot_CC_ADASA.pdf"),
]

RESPONSE_SUMMARY = [
    ("P22-BT-09-009-001", "Plant Control Philosophy", "0", "25007-0072",
     "No response code issued"),
    ("P22-BA-09-000-009", "RO Vessel Hydrostatic Test Procedure", "0", "25007-0072",
     "1 — Approved"),
    ("P22-DWG-09-005-005", "Tie-In Point Layout", "B", "25007-0072",
     "2 — Approved as noted"),
    ("P22-DWG-09-005-008", "GA of SWRO System Skid", "B", "25007-0072",
     "2 — Approved as noted"),
    ("P22-LI-09-008-016", "HMI Display Screenshot", "B", "25007-0074",
     "2 — Approved as noted"),
]


def main() -> None:
    if DOWNLOAD_LINK.startswith("PENDIENTE"):
        print("AVISO: el enlace de descarga sigue en marcador de posicion. "
              "Reemplazar DOWNLOAD_LINK antes de emitir.")

    crear_documento_adasa(
        titulo=("TECHNICAL REVIEW TRANSMITTAL N31 — SECOND STAGE RO "
                "BRINE MODULE"),
        codigo="P22-TM-09-000-031-0",
        output_filename=OUTPUT,
        incluir_toc=True,
    )

    doc = Document(OUTPUT)

    # 1. EXECUTIVE SUMMARY
    doc.add_heading("EXECUTIVE SUMMARY", level=1)
    add_para(doc, [
        ("TRANSMITTAL VERDICT: 2 — Approved as noted.", {"bold": True}),
        (" This transmittal responds to BW Water submittals ", ),
        ("25007-0072", {"bold": True}),
        (", received on 6 August with four documents, and ", ),
        ("25007-0074", {"bold": True}),
        (", received on 10 August with one. Tally: 1 Code 1, 3 Code 2, and one document "
         "returned without a response code. The review covers only whether these revisions "
         "close the observations ADASA raised in earlier transmittals.",),
    ])
    add_para(doc, [("Submittal 25007-0072 — four documents:", {"bold": True})])
    for d in DISPOSITION_0072:
        add_bullet(doc, d)
    add_para(doc, [("Submittal 25007-0074 — one document:", {"bold": True})])
    for d in DISPOSITION_0074:
        add_bullet(doc, d)
    add_para(doc, [
        ("Act on this first: the temperature-sensor mapping. ", {"bold": True}),
        (PRIORITY,)])
    add_para(doc, [
        ("The Plant Control Philosophy carries no response code ", {"bold": True}),
        (WHY_NO_CODE,),
        (REV0_TALLY, {"bold": True}),
        (".",)])
    add_para(doc, [("Review period. ", {"bold": True}), (REVIEW_PERIOD,)])
    add_para(doc, [
        ("Section 3 lists the pending observations from previous transmittals.",)])

    # 2. OBSERVATIONS BY DOCUMENT
    doc.add_heading("OBSERVATIONS BY DOCUMENT", level=1)
    for sec in SECTIONS:
        doc.add_heading(sec["heading"], level=2)
        add_para(doc, [(sec["code"], {"bold": True})])
        for p in sec["paras"]:
            add_para(doc, p)
        if sec.get("bullets"):
            add_para(doc, [(sec["bullets_lead"], {"bold": True})])
            for b in sec["bullets"]:
                add_bullet(doc, b)
        if sec.get("closing"):
            add_para(doc, sec["closing"])

    # 3. PENDING OBSERVATIONS FROM PREVIOUS TRANSMITTALS
    doc.add_heading("PENDING OBSERVATIONS FROM PREVIOUS TRANSMITTALS", level=1)
    add_para(doc, [("The three most serious open items:", {"bold": True})])
    add_simple_table(doc, [
        ("Origin TM", "Document", "Observation", "Status")] + PENDING_OPEN)
    add_para(doc, [("Closed by submittal 25007-0070. ", {"bold": True}), (PENDING_CLOSED,)])
    add_para(doc, [
        ("Also open, and overdue against BW Water's own Document and Drawing Status Report "
         "of 20 July: ", {"bold": True}),
        (PENDING_SUMMARY,)])
    add_para(doc, [("Raised in this transmittal, on documents accepted above: ",
                    {"bold": True}), (RAISED_HERE,)])
    add_para(doc, [("Submittal 25007-0071 — conditions of approval not incorporated. ",
                    {"bold": True}), (E71_LEAD,)])
    for item in E71_ITEMS:
        add_bullet(doc, item)
    add_para(doc, [(E71_TAIL,)])

    # 4. ATTACHMENTS
    doc.add_heading("ATTACHMENTS", level=1)
    add_simple_table(doc, [("Document", "Annotated PDF")] + ATTACHMENTS)
    add_para(doc, [(
        "All documents with open observations carry an annotated PDF, with two exceptions "
        "stated here rather than left unexplained. The RO Vessel Hydrostatic Test Procedure "
        "is Code 1 — Approved and therefore carries none. The Plant Control Philosophy "
        "carries none because it is returned without a response code; its open points are "
        "stated in full in the first subsection of Section 2.",)])
    dl = doc.add_paragraph()
    dl.add_run("Download — this transmittal and the annotated PDFs: ").bold = True
    aplicar_arial_12(dl)
    add_hyperlink(dl, DOWNLOAD_LINK, DOWNLOAD_LINK)

    # 5. RESPONSE SUMMARY
    doc.add_heading("RESPONSE SUMMARY", level=1)
    add_simple_table(
        doc,
        [("Document Code", "Title", "Rev", "Submittal", "Response Code")]
        + RESPONSE_SUMMARY)
    add_para(doc, [
        ("Overall Transmittal Verdict: 2 — APPROVED AS NOTED.", {"bold": True}),
        (" The three Code 2 documents issue directly at IFC Rev 0 once the items above are "
         "incorporated, with no intermediate revision. The Plant Control Philosophy carries "
         "no response code and its four open points are to be closed at the next issue of "
         "that document. Documents not appearing in this response are unaffected by this "
         "transmittal.",),
    ])

    doc.save(OUTPUT)
    print(f"Documento generado: {OUTPUT}")


if __name__ == "__main__":
    main()
