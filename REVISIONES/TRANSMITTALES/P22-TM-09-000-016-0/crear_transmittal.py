#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para generar TRANSMITTAL N16 ADASA-BW_WATER.
Submittal 25007-0034 (Entrega 34) — documento unico.
Fecha: 24-Abr-2026

Veredicto global: 2 - APPROVED AS NOTED
Tally: 1 Code 2

Scope: Civil and Loading Drawing Rev A (P22-DWG-09-005-001) — primera emision.
NOTE-01: peso del container modificado + desglose del RO Skid.
"""

import sys
import os

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))

from ejemplo_documento import (
    crear_documento_adasa,
    aplicar_arial_12,
    add_simple_table,
)
from docx import Document


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "TRANSMITTAL N16 ADASA-BW_WATER.docx")


def add_para(doc, runs):
    """runs es lista de (texto, kwargs) con kwargs en {'bold', 'italic'}."""
    para = doc.add_paragraph()
    for item in runs:
        text = item[0]
        attrs = item[1] if len(item) > 1 else {}
        run = para.add_run(text)
        if attrs.get("bold"):
            run.bold = True
        if attrs.get("italic"):
            run.italic = True
    aplicar_arial_12(para)
    return para


def add_bullet(doc, text):
    para = doc.add_paragraph()
    para.add_run("– " + text)
    aplicar_arial_12(para)


def main() -> None:
    crear_documento_adasa(
        titulo="TECHNICAL REVIEW TRANSMITTAL N16 — SECOND STAGE RO BRINE MODULE",
        codigo="P22-TM-09-000-016-0",
        output_filename=OUTPUT,
        incluir_toc=True,
    )

    doc = Document(OUTPUT)

    # =========================================================
    # 1. EXECUTIVE SUMMARY
    # =========================================================
    doc.add_heading("1. EXECUTIVE SUMMARY", level=1)

    add_para(doc, [
        ("TRANSMITTAL VERDICT: 2 — APPROVED AS NOTED", {"bold": True}),
    ])

    add_para(doc, [(
        "Single-document review of the Civil and Loading Drawing Rev A "
        "(first submission). One note on weight disclosure, to be "
        "incorporated on Rev 0 (IFC). No new revision of Rev A required.",
    )])

    # =========================================================
    # 2. OBSERVATIONS BY DOCUMENT
    # =========================================================
    doc.add_heading("2. OBSERVATIONS BY DOCUMENT", level=1)

    doc.add_heading(
        "2.1 Civil and Loading Drawing Rev A — P22-DWG-09-005-001",
        level=2,
    )

    add_para(doc, [
        ("Response Code: 2 — Approved as Noted", {"bold": True}),
    ])

    add_para(doc, [(
        "Three-sheet set: plinth layout, sectional detail, and Equipment "
        "Load table (fifteen items with dry and operating weights). "
        "Consistent with Equipment Layout Rev B. Detailed comment: "
        "P22-DWG-09-005-001_A_Civil_and_Loading_Layout_CC_ADASA.pdf.",
    )])

    add_simple_table(doc, [
        ("ID", "Severity", "Topic"),
        (
            "NOTE-01",
            "MAJOR",
            "Weight disclosure: modified container + RO Skid breakdown",
        ),
    ])

    add_para(doc, [("NOTE-01 — Weight Disclosure", {"bold": True})])

    add_para(doc, [("Include on Rev 0 (IFC):",)])

    add_para(doc, [
        ("(a) ", {"bold": True}),
        ("Total weight of the modified 40 ft container.",),
    ])

    add_para(doc, [
        ("(b) ", {"bold": True}),
        (
            "Confirm the RO Skid operating weight (items 6–7, 8,058 kg) "
            "fully accounts for all interior piping (super-duplex HP + "
            "process), including steel mass, fluid inventory and fittings, "
            "together with skid frame, pressure vessels and wet membranes. "
            "If any piping mass is excluded, declare it separately.",
        ),
    ])

    # =========================================================
    # 3. ATTACHMENTS
    # =========================================================
    doc.add_heading("3. ATTACHMENTS", level=1)

    add_simple_table(doc, [
        ("Document", "Annotated File", "Annotations"),
        (
            "Civil and Loading Drawing Rev A",
            "P22-DWG-09-005-001_A_Civil_and_Loading_Layout_CC_ADASA.pdf",
            "NOTE-01",
        ),
    ])

    # =========================================================
    # 4. RESPONSE SUMMARY
    # =========================================================
    doc.add_heading("4. RESPONSE SUMMARY", level=1)

    add_simple_table(doc, [
        ("Document Code", "Title", "Rev", "Response Code"),
        (
            "P22-DWG-09-005-001",
            "Civil and Loading Drawing",
            "A",
            "2 — Approved as Noted",
        ),
    ])

    add_para(doc, [
        ("Overall Transmittal Verdict: 2 — APPROVED AS NOTED",
         {"bold": True}),
    ])

    doc.save(OUTPUT)
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    main()
