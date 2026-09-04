#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Actualiza LISTADO ACTIVOS.xlsx con datos de TM N14 (15-Apr-2026)
y revisiones Van Doorn (14-Apr-2026).

Cambios aplicados:
  EQUIPOS:
    - TK-06-004 -> TK-06-002 (Fosa Drenajes, reasignado Van Doorn Detalle Mecanica)
    - TK-09-002: Vol 0.27 m3 -> 0.34 m3 (total datasheet, confirmado TM N9/N13)

  INSTRUMENTOS (TM N14 — Entregas E26-E29):
    - VT-09-001/002/003: IFM VTV122 -> Wilcoxon PCH420V-M12, rango 0-127 mm/s
    - CIT-09-001 -> CIT-09-001B: Rosemount 228 toroidal, 0-200 mS/cm (brine feed)
    - CIT-09-004: Rosemount 400 -> 228 toroidal, 0-200 mS/cm (Stage 1 Reject)
    - CIT-09-005: rango 0-20 mS/cm -> 0-200 mS/cm
    - CIT-09-002: rango 0-2,000 uS/cm -> 0-200 uS/cm (permeate)
    - CIT-09-003: rango 0-2,000 uS/cm -> 0-200 uS/cm (permeate)
    - PI-09-003/004/005/006: Wika 990.10 -> 990.31 (PP/EPDM para PVC)
    - TIT-09-001 -> TIT-09-006 (TAG renombrado P&ID Rev C)
    - PHIT-09-001 -> PHIT-09-006 (TAG renombrado)
    - ORPIT-09-001 -> ORPIT-09-001A (TAG renombrado)
"""

import os
from openpyxl import load_workbook

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(SCRIPT_DIR, "LISTADO ACTIVOS.xlsx")

changes_log = []


def log(sheet, row, field, old, new):
    changes_log.append(f"  [{sheet}] Row {row}: {field} '{old}' -> '{new}'")


def update_cell(ws, row, col, new_val, sheet_name, field_name):
    old = ws.cell(row=row, column=col).value
    if old != new_val:
        log(sheet_name, row, field_name, old, new_val)
        ws.cell(row=row, column=col).value = new_val
        return True
    return False


def find_row_by_tag(ws, tag, max_row=None):
    """Find row number where column 1 == tag."""
    mr = max_row or ws.max_row
    for r in range(2, mr + 1):
        if ws.cell(row=r, column=1).value == tag:
            return r
    return None


def main():
    wb = load_workbook(XLSX)

    # ============================================================
    # EQUIPOS
    # ============================================================
    ws_eq = wb["EQUIPOS"]

    # TK-06-004 -> TK-06-002 (Fosa Drenajes, Van Doorn reasignment)
    r = find_row_by_tag(ws_eq, "TK-06-004")
    if r:
        update_cell(ws_eq, r, 1, "TK-06-002", "EQUIPOS", "TAG")

    # TK-09-002: Vol 0.27 -> 0.34 m3
    r = find_row_by_tag(ws_eq, "TK-09-002")
    if r:
        old_cap = ws_eq.cell(row=r, column=6).value
        if old_cap and "0.27" in str(old_cap):
            new_cap = str(old_cap).replace("0.27", "0.34")
            update_cell(ws_eq, r, 6, new_cap, "EQUIPOS", "Capacidad")

    # ============================================================
    # INSTRUMENTOS
    # ============================================================
    ws_inst = wb["INSTRUMENTOS"]

    # --- VT-09-001/002/003: IFM VTV122 -> Wilcoxon PCH420V-M12 ---
    for vt_tag in ["VT-09-001", "VT-09-002", "VT-09-003"]:
        r = find_row_by_tag(ws_inst, vt_tag)
        if r:
            update_cell(ws_inst, r, 5, "Wilcoxon", "INSTRUMENTOS", f"{vt_tag} Fabricante")
            update_cell(ws_inst, r, 6, "PCH420V-M12", "INSTRUMENTOS", f"{vt_tag} Modelo")
            update_cell(ws_inst, r, 7, "0-127 mm/s (HART 7.0)", "INSTRUMENTOS", f"{vt_tag} Rango")

    # --- CIT-09-001 -> CIT-09-001B (toroidal for brine feed) ---
    r = find_row_by_tag(ws_inst, "CIT-09-001")
    if r:
        update_cell(ws_inst, r, 1, "CIT-09-001B", "INSTRUMENTOS", "TAG")
        update_cell(ws_inst, r, 4, "Conductivity Analyzer (Toroidal) \u2014 Brine Feed",
                    "INSTRUMENTOS", "CIT-09-001B Descripcion")
        update_cell(ws_inst, r, 5, "Rosemount", "INSTRUMENTOS", "CIT-09-001B Fabricante")
        update_cell(ws_inst, r, 6, "228 / Tx: 1056", "INSTRUMENTOS", "CIT-09-001B Modelo")
        update_cell(ws_inst, r, 7, "0-200 mS/cm", "INSTRUMENTOS", "CIT-09-001B Rango")

    # --- CIT-09-004: contacting -> toroidal (Stage 1 Reject is brine) ---
    r = find_row_by_tag(ws_inst, "CIT-09-004")
    if r:
        update_cell(ws_inst, r, 4, "Conductivity Analyzer (Toroidal) \u2014 Stage 1 Reject",
                    "INSTRUMENTOS", "CIT-09-004 Descripcion")
        update_cell(ws_inst, r, 6, "228 / Tx: 1056", "INSTRUMENTOS", "CIT-09-004 Modelo")
        update_cell(ws_inst, r, 7, "0-200 mS/cm", "INSTRUMENTOS", "CIT-09-004 Rango")

    # --- CIT-09-005: range update ---
    r = find_row_by_tag(ws_inst, "CIT-09-005")
    if r:
        update_cell(ws_inst, r, 7, "0-200 mS/cm", "INSTRUMENTOS", "CIT-09-005 Rango")

    # --- CIT-09-002/003: permeate range update ---
    for cit_tag in ["CIT-09-002", "CIT-09-003"]:
        r = find_row_by_tag(ws_inst, cit_tag)
        if r:
            update_cell(ws_inst, r, 7, "0-200 \u00b5S/cm", "INSTRUMENTOS", f"{cit_tag} Rango")

    # --- PI-09-003/004/005/006: Wika 990.10 -> 990.31 PP/EPDM ---
    for pi_tag in ["PI-09-003", "PI-09-004", "PI-09-005", "PI-09-006"]:
        r = find_row_by_tag(ws_inst, pi_tag)
        if r:
            old_model = ws_inst.cell(row=r, column=6).value
            if old_model and "990.10" in str(old_model):
                new_model = str(old_model).replace("990.10", "990.31 (PP/EPDM)")
                update_cell(ws_inst, r, 6, new_model, "INSTRUMENTOS", f"{pi_tag} Modelo")

    # --- TIT-09-001 -> TIT-09-006 (TAG renamed P&ID Rev C) ---
    r = find_row_by_tag(ws_inst, "TIT-09-001")
    if r:
        update_cell(ws_inst, r, 1, "TIT-09-006", "INSTRUMENTOS", "TAG")
        update_cell(ws_inst, r, 5, "Rosemount", "INSTRUMENTOS", "TIT-09-006 Fabricante")
        update_cell(ws_inst, r, 6, "214C / Tx: 644", "INSTRUMENTOS", "TIT-09-006 Modelo")
        update_cell(ws_inst, r, 7, "0-100 \u00b0C", "INSTRUMENTOS", "TIT-09-006 Rango")

    # --- PHIT-09-001 -> PHIT-09-006 (TAG renamed) ---
    r = find_row_by_tag(ws_inst, "PHIT-09-001")
    if r:
        update_cell(ws_inst, r, 1, "PHIT-09-006", "INSTRUMENTOS", "TAG")

    # --- ORPIT-09-001 -> ORPIT-09-001A (TAG renamed) ---
    r = find_row_by_tag(ws_inst, "ORPIT-09-001")
    if r:
        update_cell(ws_inst, r, 1, "ORPIT-09-001A", "INSTRUMENTOS", "TAG")

    # ============================================================
    # SAVE
    # ============================================================
    wb.save(XLSX)

    print(f"LISTADO ACTIVOS.xlsx actualizado: {len(changes_log)} cambios aplicados")
    for entry in changes_log:
        print(entry)
    if not changes_log:
        print("  (sin cambios)")


if __name__ == "__main__":
    main()
