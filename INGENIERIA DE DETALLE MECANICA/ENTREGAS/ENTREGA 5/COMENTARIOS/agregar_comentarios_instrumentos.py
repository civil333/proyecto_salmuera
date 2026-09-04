"""
Comentarios ADASA sobre P22-LI-06-008-101-B
Listado de Instrumentos — Rev B

Hallazgos:
  HAL-05 (MENOR): TAGs area 06 en instrumentos dentro del modulo BW Water
    - FIT-06-005 deberia ser FIT-09-XXX (en FIL-09-002 CIP, dentro del modulo)
    - LS-06-002 deberia respetar TAG BW Water (en TK-09-002, dentro del modulo)
  HAL-18 (MENOR): CLIT-09-004/005 en lista vs CIT-09-004/005 en P&ID y BW Water
"""

import os
import shutil
import openpyxl
from openpyxl.comments import Comment

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

XLSX_LOCAL = os.path.join(SCRIPT_DIR,
    "P22-LI-06-008-101-B (LI Instrumentos)_CC_ADASA.xlsx")

# Copiar fuente
src = os.path.normpath(os.path.join(SCRIPT_DIR, "..",
    "P22-LI-06-008-101-B (LI Instrumentos).xlsx"))
if os.path.exists(src):
    shutil.copy2(src, XLSX_LOCAL)
else:
    print(f"ERROR: No se encontro {src}")
    exit(1)

wb = openpyxl.load_workbook(XLSX_LOCAL)
ws = wb["Listado"]

count = 0

for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
    for cell in row:
        val = str(cell.value).strip() if cell.value else ""

        if val == "FIT-06-005":
            cell.comment = Comment(
                "Instrumento instalado en FIL-09-002 (CIP),\n"
                "dentro del modulo BW Water.\n"
                "Corregir TAG a area 09.",
                "ADASA - Luis Rivera",
                width=300,
                height=80,
            )
            count += 1
            print(f"  Comentario en {cell.coordinate} (FIT-06-005)")

        if val == "LS-06-002":
            cell.comment = Comment(
                "Instrumento instalado en TK-09-002 (dispersante),\n"
                "dentro del modulo BW Water.\n"
                "Corregir TAG a area 09.",
                "ADASA - Luis Rivera",
                width=300,
                height=80,
            )
            count += 1
            print(f"  Comentario en {cell.coordinate} (LS-06-002)")

        if val in ("CLIT-09-004", "CLIT-09-005"):
            cell.comment = Comment(
                f"Corregir a {val.replace('CLIT', 'CIT')}.\n"
                f"P&ID y BW Water usan CIT, no CLIT.",
                "ADASA - Luis Rivera",
                width=280,
                height=60,
            )
            count += 1
            print(f"  HAL-18: Comentario en {cell.coordinate} ({val})")

if count > 0:
    wb.save(XLSX_LOCAL)
    print(f"\n{count} comentarios agregados. Guardado: {XLSX_LOCAL}")
else:
    print("ADVERTENCIA: No se encontraron celdas target")
