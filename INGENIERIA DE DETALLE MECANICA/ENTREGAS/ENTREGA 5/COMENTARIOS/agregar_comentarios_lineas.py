"""
Comentarios ADASA sobre P22-LI-06-006-101-B
Listado de Lineas — Rev B

Hallazgo:
  HAL-03 (MAYOR): TAG SA-HDPE-DN110-PN10-002 duplicado (items 2 y 12, distinto origen/destino)
"""

import os
import shutil
import openpyxl
from openpyxl.comments import Comment

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

XLSX_LOCAL = os.path.join(SCRIPT_DIR,
    "P22-LI-06-006-101-B (LI Lineas)_CC_ADASA.xlsx")

# Copiar fuente
src = os.path.normpath(os.path.join(SCRIPT_DIR, "..",
    "P22-LI-06-006-101-B (LI Líneas).xlsx"))
if os.path.exists(src):
    shutil.copy2(src, XLSX_LOCAL)
else:
    print(f"ERROR: No se encontro {src}")
    exit(1)

wb = openpyxl.load_workbook(XLSX_LOCAL)
ws = wb["Listado"]

# Buscar las dos filas con SA-HDPE-DN110-PN10-002
comentario_agregado = False
for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
    for cell in row:
        if cell.value and "SA-HDPE-DN110-PN10-002" in str(cell.value):
            cell.comment = Comment(
                "TAG SA-HDPE-DN110-PN10-002 duplicado.\n"
                "Aparece en P&ID 102 (a TK-06-001) y\n"
                "en P&ID 105 (a Fosa Drenajes).\n"
                "Renumerar una de las lineas.",
                "ADASA - Luis Rivera",
                width=350,
                height=90,
            )
            comentario_agregado = True
            print(f"  Comentario agregado en celda {cell.coordinate}: {cell.value}")

if comentario_agregado:
    wb.save(XLSX_LOCAL)
    print(f"\nGuardado: {XLSX_LOCAL}")
else:
    print("ADVERTENCIA: No se encontro SA-HDPE-DN110-PN10-002 en la hoja Listado")
