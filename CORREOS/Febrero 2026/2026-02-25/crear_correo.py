#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Follow-up: Updated Project Schedule — Committed February 23, 2026
Fecha: 25 de febrero de 2026
Destinatario: Eduardo Yamauchi — BW Water Americas Inc.

Recordatorio firme por vencimiento del catch-up schedule comprometido el 18-Feb-2026.
Sin escalar (sin Andrew Zaske en CC). Menciona BAE 43.1.a sin montos.
"""

from docx import Document
from docx.shared import Pt, Inches

OUTPUT_FILE = "2026-02-25_Recordatorio-Catch-Up-Schedule.docx"
CONTACTO = "Luis Rivera"

CC_LIST = (
    "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
    "Mauricio Vallejos, Jorge Valdes, Tanya Figueroa, Allan Valentos, "
    "Jeryl F. Regulacion, Sadeep Irugalbandara, Andrew Sia, Ghazi Ozair, "
    "Nick Huta, Marjan Arsovic, Gerald Ross, Adzlan Bin Abd Rahim"
)


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_para(doc, text, size=11):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    return para


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ── HEADER ───────────────────────────────────────────────────────────────
    fields = [
        ("Date:", "February 25, 2026"),
        ("From:", f"{CONTACTO} \u2014 Contract Administrator (ADASA)"),
        ("To:", "Eduardo Yamauchi \u2014 BW Water Americas Inc."),
        ("CC:", CC_LIST),
        (
            "Subject:",
            "Follow-up: Updated Project Schedule \u2014 Committed February 23, 2026",
        ),
        ("Ref:", "Contract C-4300 / BAE 12803 / Meeting February 18, 2026"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    # ── SALUDO ────────────────────────────────────────────────────────────────
    add_para(doc, "Dear Eduardo,")
    doc.add_paragraph()

    # ── PARRAFO 1: Contexto del compromiso ───────────────────────────────────
    add_para(
        doc,
        "During the February 18 meeting you confirmed that the updated project "
        "schedule would be delivered by February 23, 2026. As of today, "
        "February 25, ADASA has not received this document \u2014 two days "
        "past the committed date.",
    )
    doc.add_paragraph()

    # ── PARRAFO 2: Impacto ────────────────────────────────────────────────────
    add_para(
        doc,
        "Without the schedule, ADASA cannot coordinate the review of the 15 "
        "engineering documents still pending, nor plan site preparation "
        "activities. The document remains a prerequisite for project recovery "
        "and for aligning the 14 documents currently under technical review.",
    )
    doc.add_paragraph()

    # ── PARRAFO 3: Multas + solicitud ─────────────────────────────────────────
    add_para(
        doc,
        "We note that penalty provisions under BAE 43.1.a for engineering "
        "phase delays have been running since January 5, 2026. ADASA requests "
        "immediate delivery of the updated schedule or, if that is not possible "
        "today, a firm committed date by return.",
    )
    doc.add_paragraph()

    # ── CIERRE ────────────────────────────────────────────────────────────────
    add_para(doc, "We remain available for coordination.")
    doc.add_paragraph()

    # ── FIRMA ─────────────────────────────────────────────────────────────────
    add_para(doc, "Best regards,")
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in [
        "Contract Administrator",
        "ADASA \u2014 Aguas de Antofagasta S.A.",
        "Project: BAE 12803 \u2014 Second Stage RO Brine Module Taltal",
    ]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
