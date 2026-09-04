#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para generar correo DOCX - Technical Review Transmittal N4
Fecha: 02 de febrero de 2026
"""

import sys
from pathlib import Path

# Agregar skill template-adasa al path
skill_path = str(
    Path(__file__).resolve().parent.parent.parent
    / ".claude"
    / "skills"
    / "template-adasa"
)
sys.path.insert(0, skill_path)

from config_defaults import DEFAULTS
from table_utils import set_table_borders, calcular_anchos_columnas, add_simple_table

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


# set_table_borders importada de table_utils (incluye centrado automatico)


def aplicar_arial(paragraph, size=11):
    """Aplica formato Arial al parrafo"""
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def crear_correo():
    """Genera el correo en formato DOCX"""

    output_file = "2026-02-02_Transmittal-N4-Revision-Tecnica-E10.docx"

    doc = Document()

    # Configurar margenes
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ============ HEADER ============
    para = doc.add_paragraph()
    para.add_run("Date: ").bold = True
    para.add_run("February 02, 2026")
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("From: ").bold = True
    para.add_run(
        f"{DEFAULTS['contacto_nombre']} - {DEFAULTS['contacto_cargo']} ({DEFAULTS['contacto_empresa']})"
    )
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("To: ").bold = True
    para.add_run("BW Water Americas Inc. (Marjan Fariborz / Logan Maroney)")
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("CC: ").bold = True
    para.add_run("ADASA Technical Management")
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("Subject: ").bold = True
    para.add_run(
        "Technical Review Transmittal N4 (P22-TM-09-000-004-0) - Submittal 0010 Review"
    )
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("Ref: ").bold = True
    para.add_run("Contract C-4300 / BAE 12803 / Transmittals N2, N3")
    aplicar_arial(para)

    doc.add_paragraph()

    # ============ INTRO ============
    para = doc.add_paragraph("Dear BW Water Team,")
    aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(
        "We submit Technical Review Transmittal N4 (P22-TM-09-000-004-0) covering the review of Submittal 0010 received January 30, 2026."
    )
    aplicar_arial(para)

    doc.add_paragraph()

    # ============ SECTION 1: TRANSMITTAL N4 SUMMARY ============
    para = doc.add_paragraph()
    run = para.add_run("1. TRANSMITTAL N4 - TECHNICAL REVIEW SUMMARY")
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = "Arial"

    doc.add_paragraph()

    # 1.1 Documents Reviewed
    para = doc.add_paragraph()
    para.add_run("1.1 Documents Reviewed").bold = True
    aplicar_arial(para)

    table = doc.add_table(rows=3, cols=5)
    set_table_borders(table)

    docs = [
        ("#", "Code", "Title", "Rev", "Verdict"),
        (
            "1",
            "P22-LI-09-009-001",
            "Utility Consumption List",
            "A",
            "3 - To be revised",
        ),
        (
            "2",
            "P22-ET-09-009-010",
            "Antiscalant Dosing Tank DS",
            "B",
            "2 - Approved as noted",
        ),
    ]

    for i, row_data in enumerate(docs):
        row = table.rows[i]
        for j, value in enumerate(row_data):
            row.cells[j].text = value
            for para in row.cells[j].paragraphs:
                for run in para.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True
                    if j == 4 and "3 - To be revised" in value:
                        run.bold = True
                        run.font.color.rgb = RGBColor(192, 0, 0)

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Transmittal Verdict: 3 - TO BE REVISED").bold = True
    for run in para.runs:
        run.font.color.rgb = RGBColor(192, 0, 0)
    aplicar_arial(para)

    doc.add_paragraph()

    # 1.2 Key Findings
    para = doc.add_paragraph()
    para.add_run("1.2 Key Findings").bold = True
    aplicar_arial(para)

    doc.add_paragraph()

    # Positive findings (narrativo)
    para = doc.add_paragraph()
    para.add_run("Positive Findings:").bold = True
    aplicar_arial(para)

    positives = [
        "SEC (Specific Energy Consumption) validated at 3.98 kWh/m3 vs 4.71 kWh/m3 guaranteed (15% favorable margin)",
        "Antiscalant Tank material change HDPE to LMDPE accepted (justified by dimensional constraints)",
        "Production rate 21 m3/h validated",
    ]
    for item in positives:
        para = doc.add_paragraph(item, style="List Bullet")
        aplicar_arial(para, 10)

    doc.add_paragraph()

    # Critical Issues (tabla)
    para = doc.add_paragraph()
    para.add_run("Critical Issues:").bold = True
    aplicar_arial(para)

    table = doc.add_table(rows=6, cols=4)
    set_table_borders(table)

    issues = [
        ("#", "Observation", "Severity", "Status"),
        (
            "OBS-01",
            "PLC specified at 60 Hz - ET 5.4.7 requires 50 Hz (Chilean grid)",
            "CRITICAL",
            "NEW",
        ),
        (
            "OBS-02",
            "Only 1 A/C unit specified - ET 5.1.11 requires n+1 (minimum 2)",
            "CRITICAL",
            "27 DAYS PENDING",
        ),
        (
            "OBS-03",
            "A/C thermal calculation not delivered - Required per ET 5.1.11",
            "CRITICAL",
            "27 DAYS PENDING",
        ),
        (
            "OBS-04",
            "HP Pump has 4 different power values (83/86/92/93 kW)",
            "MAJOR",
            "NEW",
        ),
        (
            "OBS-05",
            "CIP Pump power discrepancy (11 kW vs 15 kW Technical Offer)",
            "MAJOR",
            "NEW",
        ),
    ]

    for i, row_data in enumerate(issues):
        row = table.rows[i]
        for j, value in enumerate(row_data):
            row.cells[j].text = value
            for para in row.cells[j].paragraphs:
                for run in para.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True
                    if j == 2 and value == "CRITICAL":
                        run.bold = True
                        run.font.color.rgb = RGBColor(192, 0, 0)
                    if j == 3 and "PENDING" in value:
                        run.bold = True
                        run.font.color.rgb = RGBColor(192, 0, 0)

    doc.add_paragraph()

    # 1.3 PLC Frequency Observation (narrativo con apertura)
    para = doc.add_paragraph()
    para.add_run("1.3 PLC Frequency Observation").bold = True
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run(
        "The Utility Consumption List specifies PLC power supply as 220V/1PH/60Hz. ET Section 5.4.7 states equipment operating at frequencies other than 50 Hz will not be accepted. "
    )
    run = para.add_run(
        "If the PLC power supply is dual-frequency (50/60 Hz auto-ranging), please confirm this specification explicitly."
    )
    run.italic = True
    para.add_run(
        " If the PLC operates only at 60 Hz, replacement with 50 Hz compatible equipment is required."
    )
    aplicar_arial(para)

    doc.add_paragraph()

    # ============ SECTION 2: PENDING OBSERVATIONS STATUS ============
    para = doc.add_paragraph()
    run = para.add_run("2. PENDING OBSERVATIONS STATUS")
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = "Arial"

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(
        "Transmittal N4 includes a new Section 4 summarizing pending observations from previous transmittals:"
    )
    aplicar_arial(para)

    table = doc.add_table(rows=3, cols=3)
    set_table_borders(table)

    pending_obs = [
        ("Origin", "Days Pending", "Key Items"),
        (
            "TM N2 (Jan-06)",
            "27 days",
            "A/C n+1 configuration, A/C thermal calculation, Static Mixer material",
        ),
        (
            "TM N3 (Jan-28)",
            "5 days",
            "Vibration transmitters, Pt-100 motor windings, VM-09-015 motorized, VFD variables",
        ),
    ]

    for i, row_data in enumerate(pending_obs):
        row = table.rows[i]
        for j, value in enumerate(row_data):
            row.cells[j].text = value
            for para in row.cells[j].paragraphs:
                for run in para.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
                    if i == 0:
                        run.bold = True
                    if j == 1 and "27" in value:
                        run.bold = True
                        run.font.color.rgb = RGBColor(192, 0, 0)

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(
        "The A/C observations (OBS-02, OBS-03) were first raised in Transmittal N2 on January 6, 2026. The Utility Consumption List in Submittal 0010 still shows only 1 A/C unit (2.64 kW), confirming this non-compliance persists."
    )
    aplicar_arial(para)

    doc.add_paragraph()

    # ============ SECTION 3: CATCH-UP SCHEDULE REMINDER ============
    para = doc.add_paragraph()
    run = para.add_run("3. CATCH-UP SCHEDULE REMINDER")
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = "Arial"

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(
        "Per our request dated January 28, 2026, we are awaiting the catch-up schedule by "
    )
    run = para.add_run("February 6, 2026")
    run.bold = True
    para.add_run(" (4 days remaining).")
    aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("The current status shows 52 documents delivered with:")
    aplicar_arial(para)

    stats = [
        "34 approved/approved as noted (65%)",
        "15 requiring revision (29%)",
        "3 rejected (6%)",
    ]
    for item in stats:
        para = doc.add_paragraph(item, style="List Bullet")
        aplicar_arial(para, 10)

    doc.add_paragraph()

    # ============ SECTION 4: REQUIRED ACTIONS ============
    para = doc.add_paragraph()
    run = para.add_run("4. REQUIRED ACTIONS")
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = "Arial"

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("4.1 Critical Actions (E10)").bold = True
    aplicar_arial(para)

    critical_actions = [
        "PLC Frequency: Confirm dual-frequency compatibility OR replace with 50 Hz equipment",
        "A/C Configuration: Add 2nd A/C unit per ET 5.1.11 and Technical Offer commitment",
        "A/C Thermal Calculation: Deliver thermal load document",
    ]
    for i, action in enumerate(critical_actions, 1):
        para = doc.add_paragraph(f"{i}. {action}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("4.2 Major Actions (E10)").bold = True
    aplicar_arial(para)

    major_actions = [
        "HP Pump Power: Unify values across all documents (Utility List, Load List, Equipment List)",
        "CIP Pump Power: Clarify correct value (11 kW vs 15 kW)",
    ]
    for i, action in enumerate(major_actions, 4):
        para = doc.add_paragraph(f"{i}. {action}")
        aplicar_arial(para)

    doc.add_paragraph()

    # ============ SECTION 5: NEXT STEPS ============
    para = doc.add_paragraph()
    run = para.add_run("5. NEXT STEPS")
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = "Arial"

    doc.add_paragraph()

    next_steps = [
        "Please confirm receipt of Transmittal N4",
        "Provide the catch-up schedule by February 6, 2026 (as previously requested)",
        "Address A/C observations that are now 27 days pending",
        "Coordinate on PLC frequency confirmation",
    ]
    for i, step in enumerate(next_steps, 1):
        para = doc.add_paragraph(f"{i}. {step}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Available for coordination as needed.")
    aplicar_arial(para)

    doc.add_paragraph()

    # ============ SIGNATURE ============
    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(DEFAULTS["contacto_nombre"]).bold = True
    aplicar_arial(para)

    para = doc.add_paragraph("Contract Administrator")
    aplicar_arial(para)

    para = doc.add_paragraph("ADASA - Aguas de Antofagasta S.A.")
    aplicar_arial(para)

    para = doc.add_paragraph("Project: BAE 12803 - Second Stage RO Brine Module Taltal")
    aplicar_arial(para)

    doc.add_paragraph()

    # ============ ATTACHMENTS ============
    para = doc.add_paragraph()
    para.add_run("Attachments:").bold = True
    aplicar_arial(para)

    attachments = [
        "P22-TM-09-000-004-0 - TRANSMITTAL N4 ADASA-BW_WATER.docx",
        "Attachment L: P22-LI-09-009-001-A_Utility_Consumption_List_Comments.pdf",
        "Attachment M: P22-ET-09-009-010-B_Antiscalant_Tank_Comments.pdf",
    ]
    for att in attachments:
        para = doc.add_paragraph(att, style="List Bullet")
        aplicar_arial(para, 10)

    # Save document
    doc.save(output_file)
    print(f"Correo generado exitosamente: {output_file}")
    return output_file


if __name__ == "__main__":
    crear_correo()
