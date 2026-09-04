#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para generar correo DOCX - Seguimiento Catch-Up Schedule vencido
Fecha: 09 de febrero de 2026
Solicitud: Catch-Up Schedule (vencido 06-02-2026)
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Configuracion manual
CONTACTO = {
    "nombre": "Luis Rivera",
    "cargo": "Contract Administrator",
    "empresa": "ADASA"
}


def aplicar_arial(paragraph, size=11):
    """Aplica formato Arial al parrafo"""
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def crear_correo():
    """Genera el correo en formato DOCX"""

    output_file = "2026-02-09_Seguimiento-Catch-Up-Schedule.docx"

    doc = Document()

    # Configurar margenes
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ============ HEADER ============
    para = doc.add_paragraph()
    para.add_run("Date: ").bold = True
    para.add_run("February 09, 2026")
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("From: ").bold = True
    para.add_run(f"{CONTACTO['nombre']} - {CONTACTO['cargo']} ({CONTACTO['empresa']})")
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("To: ").bold = True
    para.add_run("BW Water Americas Inc. (Marjan Fariborz / Logan Maroney)")
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("CC: ").bold = True
    para.add_run("ADASA Technical Management")
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("Subject: ").bold = True
    para.add_run("Catch-Up Schedule Request - Overdue Response Required")
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("Ref: ").bold = True
    para.add_run("Contract C-4300 / BAE 12803 / Request dated January 28, 2026")
    aplicar_arial(para)

    doc.add_paragraph()

    # ============ CUERPO DEL CORREO - 3 PARRAFOS EJECUTIVOS ============
    
    # Parrafo 1: Contexto con Transmittals N3 y N4
    para = doc.add_paragraph()
    para.add_run("On January 28, 2026, we submitted Technical Review Transmittal N3 (P22-TM-09-000-003-0) requesting a catch-up schedule to coordinate project recovery efforts, with a delivery deadline of February 6, 2026. This deadline has now passed without response. Additionally, Transmittal N4 (P22-TM-09-000-004-0) was submitted on February 3, 2026, also awaiting your response. The engineering phase currently shows a 35-day delay from the January 5, 2026 baseline, with 14 critical documents still pending and 15 documents requiring revision. Without this schedule, we cannot coordinate internal resources, plan procurement activities, or align site preparation with delivery timelines.")
    aplicar_arial(para)

    doc.add_paragraph()

    # Parrafo 2: Lo que se requiere
    para = doc.add_paragraph()
    run1 = para.add_run("We need the catch-up schedule immediately. ")
    run1.bold = True
    run1.font.color.rgb = RGBColor(192, 0, 0)
    para.add_run("The document must include: (1) proposed delivery dates for the 14 outstanding engineering documents, including the PIE Detallado (39 days delayed), Equipment Layout (65 days), and DS MCC (81 days); (2) re-submission dates for the 15 documents marked 'To be revised', particularly the PLC frequency specification, A/C n+1 configuration, and Modbus TCP Memory Map; (3) updated milestone dates if the current 35-day delay impacts Engineering Complete, Fabrication Start, FAT, or Ready to Ship; and (4) specific mitigation actions to recover the delay.")
    aplicar_arial(para)

    doc.add_paragraph()

    # Parrafo 3: Solicitud de accion + consulta fecha de entrega del modulo
    para = doc.add_paragraph()
    para.add_run("Please confirm receipt of this request and provide the catch-up schedule immediately. If you cannot deliver the complete document now, communicate an estimated delivery date or submit a preliminary partial schedule covering the most critical items. Additionally, please confirm whether the current delays will impact the contracted module delivery date of August 3, 2026, or if you are implementing measures to ensure this milestone remains unchanged. We remain available for coordination and stand ready to support the project recovery efforts.")
    aplicar_arial(para)

    doc.add_paragraph()
    doc.add_paragraph()

    # ============ FIRMA ============
    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO["nombre"]).bold = True
    aplicar_arial(para)

    para = doc.add_paragraph("Contract Administrator")
    aplicar_arial(para)

    para = doc.add_paragraph("ADASA - Aguas de Antofagasta S.A.")
    aplicar_arial(para)

    para = doc.add_paragraph("Project: BAE 12803 - Second Stage RO Brine Module Taltal")
    aplicar_arial(para)

    # Save document
    doc.save(output_file)
    print(f"Correo generado exitosamente: {output_file}")
    return output_file


if __name__ == "__main__":
    crear_correo()
