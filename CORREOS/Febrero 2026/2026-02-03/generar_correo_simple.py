#!/usr/bin/env python3
"""
Generador simple de DOCX para correo ejecutivo TM N4
Sin usar template ADASA - formato plano
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def crear_correo_ejecutivo():
    doc = Document()

    # Configurar estilo por defecto
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Arial"
    font.size = Pt(11)

    # Header
    doc.add_heading("Technical Review Transmittal N4 - Submittals 0010, 0011", level=1)

    p = doc.add_paragraph()
    p.add_run("Date: ").bold = True
    p.add_run("February 05, 2026")

    p = doc.add_paragraph()
    p.add_run("From: ").bold = True
    p.add_run("Luis Rivera - Contract Administrator")

    p = doc.add_paragraph()
    p.add_run("To: ").bold = True
    p.add_run("BW Water Americas Inc. (Marjan Fariborz / Logan Maroney)")

    p = doc.add_paragraph()
    p.add_run("CC: ").bold = True
    p.add_run("ADASA Technical Management")

    p = doc.add_paragraph()
    p.add_run("Subject: ").bold = True
    p.add_run(
        "Technical Review Transmittal N4 (P22-TM-09-000-004-0) - Submittals 0010, 0011 Review"
    )

    p = doc.add_paragraph()
    p.add_run("Ref: ").bold = True
    p.add_run("Contract C-4300 / BAE 12803 / Transmittals N2, N3")

    doc.add_paragraph()

    # Opening
    doc.add_paragraph("Dear BW Water Team,")
    doc.add_paragraph()
    doc.add_paragraph(
        "We submit Technical Review Transmittal N4 (P22-TM-09-000-004-0) covering the review of "
        "Submittal 0010 (received January 30, 2026) and Submittal 0011 (received February 3, 2026)."
    )

    p = doc.add_paragraph()
    p.add_run("Transmittal Verdict: 3 - TO BE REVISED").bold = True

    doc.add_paragraph()

    # Documents Reviewed
    doc.add_heading("Documents Reviewed", level=2)

    table = doc.add_table(rows=5, cols=5)
    table.style = "Table Grid"

    # Header
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "#"
    hdr_cells[1].text = "Code"
    hdr_cells[2].text = "Title"
    hdr_cells[3].text = "Rev"
    hdr_cells[4].text = "Verdict"

    # Row 1
    row = table.rows[1].cells
    row[0].text = "1"
    row[1].text = "P22-LI-09-009-001"
    row[2].text = "Utility Consumption List"
    row[3].text = "A"
    row[4].text = "3 - To be revised"

    # Row 2
    row = table.rows[2].cells
    row[0].text = "2"
    row[1].text = "P22-ET-09-009-010"
    row[2].text = "Antiscalant Dosing Tank DS"
    row[3].text = "B"
    row[4].text = "2 - Approved as noted*"

    # Row 3
    row = table.rows[3].cells
    row[0].text = "3"
    row[1].text = "P22-CD-09-004-001"
    row[2].text = "Control System Architecture"
    row[3].text = "B"
    row[4].text = "2 - Approved as noted"

    # Row 4
    row = table.rows[4].cells
    row[0].text = "4"
    row[1].text = "P22-DWG-09-007-004"
    row[2].text = "Cable Tray Layout"
    row[3].text = "A"
    row[4].text = "3 - To be revised"

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run(
        "*Subject to satisfactory response to Technical Query P22-CT-09-000-001-0 (issued separately)."
    ).italic = True

    doc.add_paragraph()

    # Critical Issues
    doc.add_heading("Critical Issues", level=2)

    table = doc.add_table(rows=6, cols=4)
    table.style = "Table Grid"

    # Header
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "#"
    hdr_cells[1].text = "Observation"
    hdr_cells[2].text = "Severity"
    hdr_cells[3].text = "Status"

    # OBS-01
    row = table.rows[1].cells
    row[0].text = "OBS-01"
    row[1].text = "PLC specified at 60 Hz - ET 5.4.7 requires 50 Hz (Chilean grid)"
    row[2].text = "CRITICAL"
    row[3].text = "NEW"

    # OBS-02
    row = table.rows[2].cells
    row[0].text = "OBS-02"
    row[1].text = "Only 1 A/C unit specified - ET 5.1.11 requires n+1 (minimum 2)"
    row[2].text = "CRITICAL"
    row[3].text = "30 DAYS PENDING"

    # OBS-03
    row = table.rows[3].cells
    row[0].text = "OBS-03"
    row[1].text = "A/C thermal calculation not delivered - Required per ET 5.1.11"
    row[2].text = "CRITICAL"
    row[3].text = "30 DAYS PENDING"

    # OBS-04
    row = table.rows[4].cells
    row[0].text = "OBS-04"
    row[1].text = "Modbus TCP Memory Map not delivered - Committed in TM N2 response"
    row[2].text = "CRITICAL"
    row[3].text = "30 DAYS PENDING"

    # OBS-10
    row = table.rows[5].cells
    row[0].text = "OBS-10"
    row[1].text = "Container dimensions exceed approved 40ft - Rejected Nov-17-2025"
    row[2].text = "CRITICAL"
    row[3].text = "NEW"

    doc.add_paragraph()

    # Required Actions
    doc.add_heading("Required Actions", level=2)

    actions = [
        "PLC Frequency: Confirm dual-frequency compatibility OR replace with 50 Hz equipment",
        "A/C Configuration: Add 2nd A/C unit per ET 5.1.11 and Technical Offer commitment (30 days pending)",
        "A/C Thermal Calculation: Deliver thermal load document (30 days pending)",
        "Modbus TCP Memory Map: Deliver document as committed (30 days pending)",
        "Container Dimensions: Confirm layout uses approved 40ft standard (60ft rejected Nov-17-2025)",
    ]

    for i, action in enumerate(actions, 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(f"{i}. ").bold = True
        p.add_run(action)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Note: ").bold = True
    p.add_run(
        "Catch-up schedule requested January 28, 2026 - awaiting response by February 6, 2026."
    )

    doc.add_paragraph()

    # Attachments
    doc.add_heading("Attachments", level=2)

    attachments = [
        "P22-TM-09-000-004-0 - TRANSMITTAL N4 ADASA-BW_WATER.docx",
        "P22-LI-09-009-001-A_Utility_Consumption_List_Comments.pdf",
        "P22-ET-09-009-010-B_Antiscalant_Tank_Comments.pdf",
        "P22-CD-09-004-001-B_Control_Architecture_Comments.pdf",
        "P22-DWG-09-007-004-A_Cable_Tray_Layout_Comments.pdf",
        "P22-CT-09-000-001-0_Technical-Query-Antiscalant-Dosing.docx",
    ]

    for att in attachments:
        doc.add_paragraph(att, style="List Bullet")

    doc.add_paragraph()
    doc.add_paragraph("Best regards,")
    doc.add_paragraph()

    # Firma
    p = doc.add_paragraph()
    p.add_run("Luis Rivera").bold = True
    doc.add_paragraph("Contract Administrator")
    doc.add_paragraph("ADASA - Aguas de Antofagasta S.A.")
    doc.add_paragraph("Project: BAE 12803 - Second Stage RO Brine Module Taltal")

    # Guardar
    output_file = "2026-02-03_Transmittal-N4-Revision-Tecnica-E10-E11.docx"
    doc.save(output_file)
    print(f"Correo ejecutivo generado: {output_file}")
    print(f"Total de párrafos: {len(doc.paragraphs)}")


if __name__ == "__main__":
    crear_correo_ejecutivo()
