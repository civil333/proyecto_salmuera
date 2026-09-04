#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Technical Review Transmittal N5 — Equipment Layout RO Container
Fecha: 23 de febrero de 2026
Submittal: 0012 — P22-DWG-09-005-003-A
"""

import sys
import os
from pathlib import Path

skill_path = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..", "..",
    ".claude", "skills", "template-adasa",
)
sys.path.insert(0, skill_path)

from ejemplo_documento import set_table_borders, calcular_anchos_columnas

from docx import Document
from docx.shared import Pt, Inches, RGBColor

OUTPUT_FILE = "2026-02-23_Transmittal-N5-Equipment-Layout.docx"
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_table(doc, data, red_col=None, red_value=None):
    """Crea tabla con anchos dinamicos y bordes."""
    anchos = calcular_anchos_columnas(data, len(data[0]))
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    set_table_borders(table)
    for i, row_data in enumerate(data):
        row = table.rows[i]
        for j, value in enumerate(row_data):
            row.cells[j].text = value
            for para in row.cells[j].paragraphs:
                for run in para.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True
                    if red_col and j == red_col and red_value and red_value in value:
                        run.bold = True
                        run.font.color.rgb = RGBColor(192, 0, 0)
    if anchos:
        for row in table.rows:
            for j, cell in enumerate(row.cells):
                if j < len(anchos):
                    cell.width = anchos[j]
    return table


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ── HEADER ──────────────────────────────────────────────────────────────
    fields = [
        ("Date:", "February 23, 2026"),
        ("From:", f"{CONTACTO} — Contract Administrator (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc."),
        (
            "CC:",
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
            "Mauricio Vallejos, Jorge Valdes, Tanya Figueroa, Allan Valentos, "
            "Jeryl F. Regulacion, Sadeep Irugalbandara, Andrew Sia, Ghazi Ozair, "
            "Nick Huta, Marjan Arsovic, Gerald Ross, Andrew Zaske, Adzlan Bin Abd Rahim",
        ),
        (
            "Subject:",
            "Technical Review Transmittal N5 (P22-TM-09-000-005-0) "
            "— Submittal 0012 — Equipment Layout",
        ),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / "
            "P22-TM-09-000-002-0 / P22-TM-09-000-003-0 / P22-TM-09-000-004-0",
        ),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    # ── CUERPO ───────────────────────────────────────────────────────────────
    para = doc.add_paragraph("Dear Eduardo,")
    aplicar_arial(para)
    doc.add_paragraph()

    # Párrafo 1: apertura + veredicto
    para = doc.add_paragraph()
    para.add_run(
        "Please find attached Technical Review Transmittal N5 (P22-TM-09-000-005-0), "
        "covering Submittal 0012 — Equipment Layout of RO Container "
        "(P22-DWG-09-005-003 Rev A). "
    )
    para.add_run("Verdict: 3 — To be revised.").bold = True
    aplicar_arial(para)
    doc.add_paragraph()

    # Párrafo 2: hallazgo crítico
    para = doc.add_paragraph(
        "The critical finding is the CIP system shown inside the container. "
        "The Technical Offer Rev.1 specifies all CIP equipment outside, and BW Water "
        "confirmed this in our February 18 meeting. The revised drawing must place the "
        "CIP system outside following the equipment distribution specified by ADASA — "
        "not an arbitrary external arrangement. "
        "Five additional observations are documented in the transmittal "
        "(metric units, doors, A/C locations, floor material, seismic anchoring)."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # Párrafo 3: positivo + pendientes
    para = doc.add_paragraph(
        "On the positive side, the drawing confirms a standard 40ft container, "
        "which closes OBS-10 from Transmittal N4. "
        "8 observations from previous transmittals remain open, the oldest dating "
        "48 days (A/C thermal calculation and Modbus TCP Memory Map, first raised in TM N2). "
        "Full detail in the attached transmittal."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # Párrafo 4: schedule
    para = doc.add_paragraph(
        "We also note that the updated project schedule was committed for today. "
        "Please send it at your earliest convenience."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # ── ADJUNTOS ─────────────────────────────────────────────────────────────
    para = doc.add_paragraph()
    para.add_run("Attachments:").bold = True
    aplicar_arial(para)
    for att in [
        "P22-TM-09-000-005-0 — TRANSMITTAL N5 ADASA-BW_WATER.docx",
        "P22-DWG-09-005-003-A_Equipment_Layout_Comments.pdf",
    ]:
        para = doc.add_paragraph(f"- {att}")
        aplicar_arial(para)

    doc.add_paragraph()

    # ── FIRMA ────────────────────────────────────────────────────────────────
    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in [
        "Contract Administrator",
        "ADASA — Aguas de Antofagasta S.A.",
        "Project: BAE 12803 — Second Stage RO Brine Module Taltal",
    ]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
