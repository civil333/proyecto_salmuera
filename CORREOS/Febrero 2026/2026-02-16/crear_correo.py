#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para generar correo DOCX - CT-001 Response Evaluation
Fecha: 16 de febrero de 2026
Asunto: Technical Query CT-001 — Response Evaluation and Pending Observations (C-4300)

Estructura fluida tipo correo profesional (sin numeracion de secciones).
Sigue el estilo del correo 2026-02-12.
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

CONTACTO = {
    "nombre": "Luis Rivera Gonzalez",
    "cargo": "Leader, Infrastructure Engineering",
    "empresa": "ADASA - Aguas de Antofagasta S.A.",
    "proyecto": "BAE 12803 - Second Stage RO Brine Module Taltal",
}


def aplicar_arial(paragraph, size=11):
    """Aplica formato Arial al parrafo"""
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_bold_label(doc, label, size=11):
    """Agrega un parrafo con texto en bold como etiqueta de seccion inline"""
    para = doc.add_paragraph()
    run = para.add_run(label)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(size)
    return para


def add_para(doc, text, size=11):
    """Agrega un parrafo normal con formato Arial"""
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    return para


def add_para_mixed(doc, fragments, size=11):
    """Agrega un parrafo con fragmentos de texto (bold/normal mezclados).
    fragments: lista de tuplas (text, bold)
    """
    para = doc.add_paragraph()
    for text, bold in fragments:
        run = para.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.bold = bold
    return para


def set_table_borders(table):
    """Aplica bordes a toda la tabla via XML"""
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else tbl.makeelement(qn("w:tblPr"), {})
    borders = tblPr.makeelement(qn("w:tblBorders"), {})
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.makeelement(
            qn(f"w:{border_name}"),
            {
                qn("w:val"): "single",
                qn("w:sz"): "4",
                qn("w:space"): "0",
                qn("w:color"): "000000",
            },
        )
        borders.append(border)
    tblPr.append(borders)
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)


def set_repeat_table_header(row):
    """Configura una fila para repetirse como header en cada pagina"""
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    tblHeader = trPr.makeelement(qn('w:tblHeader'), {qn('w:val'): 'true'})
    trPr.append(tblHeader)


def add_table_simple(doc, headers, rows):
    """Agrega tabla con bordes, header en negrita y formato Arial"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = True
    set_table_borders(table)

    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9)
        # Fondo azul claro
        shading = cell._element.get_or_add_tcPr()
        shd = shading.makeelement(
            qn("w:shd"),
            {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "D9E2F3"},
        )
        shading.append(shd)

    # Data rows
    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            run.font.name = "Arial"
            run.font.size = Pt(9)

    # Repeat header row en cada pagina
    set_repeat_table_header(table.rows[0])

    doc.add_paragraph()
    return table


def crear_correo():
    """Genera el correo en formato DOCX"""

    output_file = "2026-02-16_CT001-Response-Evaluation.docx"
    doc = Document()

    # Margenes
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ============ HEADER ============
    header_fields = [
        ("Date:", "February 16, 2026"),
        (
            "From:",
            f"{CONTACTO['nombre']} - {CONTACTO['cargo']} ({CONTACTO['empresa']})",
        ),
        ("To:", "Eduardo Yamauchi - Operations Director Americas (BW Water)"),
        ("CC:", "ADASA Technical Management; BW Water Engineering Team"),
        (
            "Subject:",
            "Technical Query CT-001 \u2014 Response Evaluation and Pending "
            "Observations (C-4300)",
        ),
        ("Ref:", "Contract C-4300 / BAE 12803"),
        (
            "Attachments:",
            "P22-CT-09-000-001-1 (Response Evaluation \u2014 Antiscalant "
            "Dosing Justification)",
        ),
    ]
    for label, value in header_fields:
        para = doc.add_paragraph()
        para.add_run(label + " ").bold = True
        para.add_run(value)
        aplicar_arial(para)

    doc.add_paragraph()

    # ============ SALUDO + APERTURA ============
    add_para(doc, "Dear Eduardo,")

    doc.add_paragraph()

    add_para_mixed(
        doc,
        [
            (
                "We have completed our evaluation of BW Water\u2019s response to "
                "Technical Query CT-001 (Antiscalant Dosing Justification). The "
                "response was received on ",
                False,
            ),
            ("February 13, 2026", True),
            (
                ", three days after the agreed deadline of February 10. "
                "Three documents were submitted: a BW Water Memorandum, an AWC "
                "projection for Pureflux SW antiscalant, and a CREST Water "
                "assessment email.",
                False,
            ),
        ],
    )

    add_para(
        doc,
        "The formal evaluation document (P22-CT-09-000-001-1) is attached. "
        "This email summarizes the main findings.",
    )

    # ============ ACCEPTANCE ============
    doc.add_paragraph()
    add_bold_label(doc, "Acceptance")

    add_para_mixed(
        doc,
        [
            (
                "ADASA accepts in principle the ",
                False,
            ),
            ("0.5 ppm dosing rate", True),
            (
                " for Pureflux SW antiscalant. The AWC PROTON simulation confirms "
                "positive safety margins for CaCO\u2083, CaSO\u2084, BaSO\u2084, "
                "SrSO\u2084, and silica scaling under the modeled conditions. The "
                "dosing pump\u2019s 115x capacity margin provides adequate operational "
                "flexibility.",
                False,
            ),
        ],
    )

    add_para(
        doc,
        "This acceptance is conditional on the satisfactory resolution of four "
        "observations identified during our review.",
    )

    # ============ OBSERVATIONS ============
    doc.add_paragraph()
    add_bold_label(doc, "Pending Observations")

    add_para_mixed(
        doc,
        [
            ("Temperature basis (Major).", True),
            (
                " The AWC projection was run at 19\u00b0C feed temperature. CT-001 "
                "specifically requested validation at 24\u00b0C, which is the "
                "worst-case condition per the Technical Specification Table 4-1. "
                "Higher temperature increases both saturation indices and crystal "
                "growth rates. We need either a revised projection at 24\u00b0C or a "
                "technical justification explaining why 19\u00b0C is representative.",
                False,
            ),
        ],
    )

    add_para_mixed(
        doc,
        [
            ("Chemical Consumption List inconsistency (Minor).", True),
            (
                " The volumetric flow listed as 0.02 L/h does not reconcile with "
                "the 0.59 kg/day daily consumption. Cross-checking against the AWC "
                "dosing rate (0.396 mL/min), the correct value should be 0.024 L/h. "
                "This is a documentation correction and does not affect equipment "
                "sizing.",
                False,
            ),
        ],
    )

    add_para_mixed(
        doc,
        [
            ("AWC projection input data (Major).", True),
            (
                " ADASA has verified the AWC PROTON input data against the ANAM "
                "brine characterization. Strontium was entered as 0.00 mg/L in "
                "the AWC simulation, while the ANAM reports measured 10\u201311 "
                "mg/L \u2014 directly relevant for SrSO\u2084 scaling. Silica was "
                "entered as 2.1 mg/L (first sampling only), not the worst-case "
                "6.4 mg/L from the second sampling. A revised AWC projection "
                "incorporating these measured values is required.",
                False,
            ),
        ],
    )

    add_para_mixed(
        doc,
        [
            ("Contradictory conclusions (Major).", True),
            (
                " AWC concludes that 0.5 ppm antiscalant is needed; CREST Water "
                "concludes that no antiscalant is necessary. The BW Water "
                "memorandum presents both assessments without stating which one "
                "forms the design basis. We need a clear statement of the adopted "
                "position and the reasoning behind it.",
                False,
            ),
        ],
    )

    # ============ SUMMARY TABLE ============
    doc.add_paragraph()
    add_bold_label(doc, "Required Actions Summary")

    add_table_simple(
        doc,
        ["#", "Action", "Priority", "Ref"],
        [
            [
                "1",
                "AWC projection at 24\u00b0C, or technical\njustification for 19\u00b0C",
                "High",
                "OBS-1",
            ],
            [
                "2",
                "Correct volumetric flow in Chemical\nConsumption List (0.02 \u2192 0.024 L/h)",
                "Low",
                "OBS-2",
            ],
            [
                "3",
                "Revised AWC projection with Sr (10\u201311\nmg/L) and worst-case SiO\u2082 (6.4 mg/L)",
                "High",
                "OBS-3",
            ],
            [
                "4",
                "State adopted design basis:\nAWC or CREST Water",
                "High",
                "OBS-4",
            ],
        ],
    )

    # ============ DEADLINE ============
    add_para_mixed(
        doc,
        [
            ("Response deadline: ", False),
            ("February 20, 2026.", True),
        ],
    )

    # ============ CIERRE ============
    doc.add_paragraph()

    add_para(
        doc,
        "The full evaluation with detailed findings is in the attached document. "
        "Please do not hesitate to reach out if you need to discuss any of the "
        "observations or coordinate the response.",
    )

    # ============ FIRMA ============
    doc.add_paragraph()
    add_para(doc, "Best regards,")
    doc.add_paragraph()

    para = doc.add_paragraph()
    run = para.add_run(CONTACTO["nombre"])
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)

    add_para(doc, CONTACTO["cargo"])
    add_para(doc, CONTACTO["empresa"])
    add_para(doc, f"Project: {CONTACTO['proyecto']}")

    # Save
    doc.save(output_file)
    print(f"Correo generado exitosamente: {output_file}")
    return output_file


if __name__ == "__main__":
    crear_correo()
