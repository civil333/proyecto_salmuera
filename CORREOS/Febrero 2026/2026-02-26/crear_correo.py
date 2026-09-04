#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Technical Information Required — Module Tie-in Points, Elevations,
        and CIP External Footprint (C-4300)
Fecha: 26 de febrero de 2026
Ref:   Contract C-4300 / BAE 12803 / P22-TM-09-000-005-0 / Offer Rev.1
"""

import sys
import os

skill_path = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..", "..",
    ".claude", "skills", "template-adasa",
)
sys.path.insert(0, skill_path)

from ejemplo_documento import set_table_borders, calcular_anchos_columnas

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT_FILE = "2026-02-26_Request-Tie-in-Definition-CIP-External-Layout.docx"
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def set_cell_background(cell, hex_color):
    """Sets cell background shading color."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_table(doc, data, header_color="D9E2F3"):
    """Crea tabla con anchos dinámicos, bordes y header coloreado."""
    anchos = calcular_anchos_columnas(data, len(data[0]))
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    set_table_borders(table)
    for i, row_data in enumerate(data):
        row = table.rows[i]
        for j, value in enumerate(row_data):
            cell = row.cells[j]
            cell.text = value
            if i == 0:
                set_cell_background(cell, header_color)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True
    if anchos:
        for row in table.rows:
            for j, cell in enumerate(row.cells):
                if j < len(anchos):
                    cell.width = anchos[j]
    return table


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ── HEADER ──────────────────────────────────────────────────────────────
    fields = [
        ("Date:", "February 26, 2026"),
        ("From:", f"{CONTACTO} — Contract Administrator (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc."),
        (
            "CC:",
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
            "Mauricio Vallejos, Jorge Valdes, Tanya Figueroa, Allan Valentos, "
            "Jeryl F. Regulacion, Sadeep Irugalbandara, Andrew Sia, Ghazi Ozair, "
            "Nick Huta, Marjan Arsovic, Gerald Ross, Andrew Zaske, Adzlan Bin Abd Rahim",
        ),
        (
            "Subject:",
            "Pending Technical Deliverables \u2014 Tie-in Definitions, CIP Footprint "
            "& Overdue P&ID Rev B (C-4300)",
        ),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / P22-TM-09-000-005-0 / Offer Rev.1",
        ),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    # ── SALUDO ───────────────────────────────────────────────────────────────
    para = doc.add_paragraph("Dear Eduardo,")
    aplicar_arial(para)
    doc.add_paragraph()

    # ── PÁRRAFO 1: Contexto de contingencia ──────────────────────────────────
    para = doc.add_paragraph(
        "Following Transmittal N5 (P22-TM-09-000-005-0, issued February 23), ADASA "
        "understands that BW Water is preparing a revised Equipment Layout with the CIP "
        "system relocated outside the container, as required by the Technical Offer Rev.1 "
        "and confirmed in the February 18 meeting. I am writing to request the information "
        "detailed below no later than Monday, March 2, 2026. The design of the external "
        "equipment foundations, structural support, and interconnection piping is currently "
        "on hold pending these definitions \u2014 each additional week of delay directly impacts "
        "the peripheral engineering schedule on the critical path to FAT (July 2026). The "
        "information requested is independent of the ongoing layout revision and can be "
        "provided as a preliminary data sheet or standalone technical note."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # ── PÁRRAFO 2: Tie-ins de proceso principales ────────────────────────────
    para = doc.add_paragraph(
        "For the design of the piping interface between the ADASA plant and the RO module, "
        "the following process connection points are required at firme: feed inlet (brine), "
        "permeate outlet, and concentrate/reject outlet. For each connection, ADASA requires: "
        "nominal diameter, flange standard and rating, container face where the nozzle emerges, "
        "and centerline elevation referenced to the bottom floor of the container "
        "(elevation 0.00). The revised layout shall confirm that all process connections "
        "are concentrated on the designated face, consistent with ADASA\u2019s annotation "
        "regarding process infrastructure alignment."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # ── PÁRRAFO 3: Nuevos tie-ins CIP y antiscalante ─────────────────────────
    para = doc.add_paragraph(
        "Relocating the CIP and antiscalant dosing equipment outside the container "
        "\u2014 per Offer Rev.1 and the February 18 agreement \u2014 generates additional "
        "connection points that must be defined at firme: CIP feed (from internal permeate "
        "header), CIP return to Stage 1, CIP return to Stage 2, CIP reject/drain, and "
        "antiscalant injection point in the feed line upstream of the cartridge filter. "
        "For each connection: nominal diameter, fitting or flange type, container face, "
        "and centerline elevation referenced to elevation 0.00."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # ── PÁRRAFO 4: Footprint CIP externo ─────────────────────────────────────
    para = doc.add_paragraph(
        "To design the external foundation and platform where the CIP equipment will be "
        "installed \u2014 also requested in TM N5 OBS-01 \u2014 ADASA requires a preliminary "
        "layout or dimensioned sketch showing the footprint of: CIP Tank (reference: "
        "71\u2033 D \u00d7 88\u2033 H per Offer Rev.1, approx. 1.8 m \u00d7 2.2 m, "
        "4,498 kg when full), CIP Pump skid, CIP Cartridge Filter, and Antiscalant dosing "
        "skid. Minimum maintenance clearances must be indicated. An IFC drawing is not "
        "required at this stage; a preliminary sketch committed at firme is sufficient. "
        "I can share the Taltal site plan to facilitate coordination of the external "
        "CIP location."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # ── PÁRRAFO 5: Recordatorio P&ID Rev B (TM N2) ───────────────────────────
    para = doc.add_paragraph(
        "ADASA also notes that P&ID Rev A (P22-DWG-09-009-002-A) was reviewed in "
        "Transmittal N2 (P22-TM-09-000-002-0, January 26) with verdict 2 \u2014 Approved as Noted. "
        "As of February 26, Rev B has not been received. Thirteen observations remain open "
        "(OBS-01 through OBS-14, except OBS-09 which was resolved with Process Calc Rev B). "
        "The corrected P&ID is required to validate the tie-in geometry and container face "
        "assignments requested above."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # ── TABLA: Resumen de información requerida ───────────────────────────────
    para = doc.add_paragraph()
    para.add_run("Summary of Required Information:").bold = True
    aplicar_arial(para)
    doc.add_paragraph()

    table_data = [
        ("#", "Information Required", "Reference", "Priority"),
        (
            "1",
            "Process tie-ins: Feed, Permeate, Concentrate\n"
            "\u2014 Nominal \u00d8, flange std & rating, container face, CL elevation",
            "ET Sec 7\nOffer Rev.1",
            "CRITICAL",
        ),
        (
            "2",
            "New CIP tie-ins: CIP Feed, CIP to Stage 1, CIP to Stage 2, CIP Reject\n"
            "\u2014 Nominal \u00d8, connection type, container face, CL elevation",
            "TM N5 OBS-01\nOffer Rev.1",
            "CRITICAL",
        ),
        (
            "3",
            "Antiscalant injection point: nominal \u00d8, location in feed line, "
            "CL elevation",
            "CT-001\nOffer Rev.1",
            "CRITICAL",
        ),
        (
            "4",
            "External CIP footprint: tank, pump skid, cartridge filter, "
            "antiscalant skid\n\u2014 Dimensions + minimum maintenance clearances",
            "TM N5 OBS-01",
            "CRITICAL",
        ),
        (
            "5",
            "Container reference elevations: floor level, roof level, platform level",
            "ET Section 7\nItem 6",
            "MAJOR",
        ),
        (
            "6",
            "P&ID Rev B with all 13 open observations addressed\n"
            "(issued TM N2, Jan 26 \u2014 31 days pending)",
            "TM N2 \u00a73.1\nP22-TM-09-000-002-0",
            "CRITICAL",
        ),
    ]
    add_table(doc, table_data)

    doc.add_paragraph()

    # ── CIERRE ────────────────────────────────────────────────────────────────
    para = doc.add_paragraph(
        "Please send your response by Monday, March 2, 2026. Without this information, "
        "the design of the external equipment foundations and interconnection piping cannot "
        "proceed. If a coordination call would help clarify any of these points, I am "
        "available at short notice \u2014 please propose a time and I will confirm immediately. "
        "Please also confirm receipt of this request and your expected delivery date by return."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # ── FIRMA ─────────────────────────────────────────────────────────────────
    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in [
        "Contract Administrator",
        "ADASA \u2014 Aguas de Antofagasta S.A.",
        "Project: BAE 12803 \u2014 Second Stage RO Brine Module Taltal",
    ]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
