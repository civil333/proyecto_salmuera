#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
crear_correo.py — Correo de cobertura del Transmittal N42 (P22-TM-09-000-042-0).

A Eduardo Yamauchi, cadena regular de transmittals. Cubre los submittals 25007-0096,
25007-0097, 25007-0099 y 25007-0100, ocho documentos con el plano del yugo. Los cuatro CC_ADASA van ADJUNTOS
(unos 6 MB en total con el transmittal), sin enlace Synology.

QUE HACE EL CORREO Y QUE NO. Remite y no repite: veredicto, una vineta por documento o
bloque con la accion que falta, y los adjuntos. La vineta del izaje cita el correo de
ADASA del 15-Sep (pedido de Luis del 5-Oct: reclamar el plano de detalle del yugo
invocando ese correo). Fuera: el detalle de cada observacion (vive en el transmittal y en
los PDF anotados), la devolucion a tres dias corridos y toda vision de proyecto.
Primera persona (lo firma Luis Rivera).

Distribucion: Eduardo destinatario; copia Jeryl F. Regulacion (BW Water) y Victor
Gutierrez (ADASA). SIN Adzlan Abd Rahim: su casilla no existe desde julio y rebota cada
envio; hay que sacarlo a mano si se usa Responder a todos (INT-13). Ingles. Document()
directo, sin template ADASA. Estado: BORRADOR.
"""
import os

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-10-07_Transmittal-N42-Lifting-Package.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for r in paragraph.runs:
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return paragraph


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph()
    p.add_run(text)
    return aplicar_arial(p, size)


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for texto, bold in segs:
        r = p.add_run(texto)
        r.bold = bold
    return aplicar_arial(p, size)


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.add_run("•  ")
    p.add_run(lead).bold = True
    p.add_run(rest)
    return aplicar_arial(p, size)


def blank(doc):
    return doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "October 7, 2026"),
        ("From:", CONTACTO + " - Leader, Infrastructure Engineering (ADASA)"),
        ("To:", "Eduardo Yamauchi - BW Water"),
        ("CC:", "Jeryl F. Regulacion - BW Water; Victor Gutierrez - ADASA"),
        ("Subject:", "TALTAL - Transmittal N42 - Lifting package and submittals "
                     "25007-0096, 0097, 0099 and 0100"),
        ("Ref:", "Contract C-4300 / BAE 12803 / Technical Specification "
                 "P22-ET-09-000-001-0, Section 7"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_segments(doc, [
        ("Please find attached Transmittal N42 (P22-TM-09-000-042-0), covering submittals "
         "25007-0096, 25007-0097, 25007-0099 and 25007-0100, eight documents. ", False),
        ("Verdict: 3 - To be revised.", True),
    ])
    blank(doc)

    add_bullet_lead(doc, "Yoke drawing P22-DWG-09-005-019 Rev A - Code 3. ",
                    "The drawing I requested in my e-mail of 15 September, due on the 17th, "
                    "arrived as sheet 1 of 3 inside the calculation addendum. It has no frame "
                    "joints, material, weights or sling data, so it cannot be fabricated "
                    "from. Please submit it complete, sheets 1 to 3, as a document of its "
                    "own, in national welded IN or HN sections, with the splices and corner "
                    "joints detailed and at least one intermediate beam at mid-length.")
    add_bullet_lead(doc, "Calculation addendum Rev A - Code 3. ",
                    "It checks only the lugs. Re-issue it under its own document code with "
                    "the design of the frame, including its torsion, racking and the "
                    "buckling of the long members, under the NCh3171 combinations of the approved "
                    "Design Criteria as a minimum, and the local check of the corner "
                    "fittings.")
    add_bullet_lead(doc, "Maintenance Lifting Points Rev A - Code 3. ",
                    "Re-issue as Rev B with the runway beam or fixed lifting points for the "
                    "pumps and both turbochargers.")
    add_bullet_lead(doc, "GA of Antiscalant Dosing Tank Rev E - Code 2. ",
                    "Set the level marking at 34 inches (864 mm) at Rev 0.")
    add_bullet_lead(doc, "Piping Layout Rev E, Valve List Rev 0, PLC/LCP FAT Procedure Rev 0 "
                         "and HMI Rev 0 - Code 1. ", "No action on these documents.")
    blank(doc)

    add_para(doc, "The observations are itemised on the four annotated PDFs. Section 3 of "
                  "the transmittal lists the items still open from previous transmittals.")
    blank(doc)

    add_para(doc, "Attachments:")
    add_bullet_lead(doc, "TRANSMITTAL N42 ADASA-BW_WATER.pdf", "")
    add_bullet_lead(doc, "P22-DWG-09-005-019_A_GA_Module_Lifting_from_Top_CC_ADASA.pdf", "")
    add_bullet_lead(doc, "P22-CD-09-005-003_A_Lifting_Addendum_CC_ADASA.pdf", "")
    add_bullet_lead(doc, "P22-DWG-09-005-006_A_Maintenance_Lifting_Points_CC_ADASA.pdf", "")
    add_bullet_lead(doc, "P22-DWG-09-005-015_E_GA_Antiscalant_Dosing_Tank_CC_ADASA.pdf", "")
    blank(doc)

    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in ["Leader, Infrastructure Engineering",
                 "ADASA, Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)

    cp = doc.core_properties
    cp.title = "TALTAL - Transmittal N42 - Lifting package"
    cp.author = "Luis Rivera Gonzalez"
    cp.last_modified_by = "Luis Rivera Gonzalez"
    cp.category = "Correo de cobertura"
    cp.comments = ("Cobertura del Transmittal N42: plano del yugo, addendum de izaje y "
                   "puntos de izaje de mantencion en Codigo 3; GA del estanque Codigo 2; "
                   "cuatro Codigo 1.")
    doc.save(OUTPUT_FILE)

    cuerpo = [p.text for p in doc.paragraphs]
    i0 = cuerpo.index("Dear Eduardo,")
    i1 = cuerpo.index("Attachments:")
    palabras = sum(len(t.split()) for t in cuerpo[i0:i1])
    print("Correo generado: " + OUTPUT_FILE)
    print(f"  palabras del cuerpo (saludo a adjuntos): {palabras}")
    assert "Adzlan" not in " ".join(cuerpo)


if __name__ == "__main__":
    crear_correo()
