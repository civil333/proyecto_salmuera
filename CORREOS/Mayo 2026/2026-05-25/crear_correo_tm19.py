#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Transmittal N19 (P22-TM-09-000-019-0). Submittals 25007-0042..0045.
Fecha: 25 de Mayo de 2026 (lunes)
Veredicto TM N19: 3 - TO BE REVISED. Tally 3 Code 1 + 5 Code 2 + 5 Code 3.
13 documentos. 10 CC_ADASA adjuntos (5 Code 3 + 5 Code 2).

La Technical Note P22-NT-09-000-001-0 se envia en un correo separado
(cadena distinta, Reply-To al cover BW Water 24-May-2026 que transmitio
el Mitigation Plan).
"""

import os
from docx import Document
from docx.shared import Pt, Inches

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-05-25_Transmittal-N19.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # =========================================================================
    # HEADER
    # =========================================================================
    fields = [
        ("Date:", "May 25, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
            "Mauricio Vallejos, Jorge Valdes, Tanya Figueroa, Allan Valentos, "
            "Jeryl F. Regulacion, Sadeep Irugalbandara, Andrew Sia, "
            "Ghazi Ozair, Nick Huta, Marjan Arsovic, Gerald Ross, "
            "Andrew Zaske, Adzlan Bin Abd Rahim",
        ),
        (
            "Subject:",
            "ADASA – Taltal Brine Module: Technical Review Transmittal N19 "
            "(25007-0042 to 25007-0045)",
        ),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / P22-TM-09-000-019-0",
        ),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    # =========================================================================
    # CUERPO
    # =========================================================================
    para = doc.add_paragraph("Dear BW Water Project Team,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Attached: ")
    para.add_run("Transmittal N19").bold = True
    para.add_run(
        " (P22-TM-09-000-019-0) covering submittals 25007-0042 to "
        "25007-0045 — thirteen documents — for ADASA's technical "
        "review."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Verdict: 3 — To Be Revised.").bold = True
    para.add_run(
        " Tally: 3 Code 1, 5 Code 2, 5 Code 3. Drivers: ITP Offsite Rev B "
        "(Section 2.10), Grounding & Power Panel Layout Rev E (Section "
        "2.5), Instrumentation & Control Cable Schedule Rev 0 (Section "
        "2.11), and both Cartridge Filter datasheets (Sections 2.12 and "
        "2.13)."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Key positive closures in this cycle: ").bold = True
    para.add_run("Cable Tray Layout Rev C").bold = True
    para.add_run(
        " closes the seven items inherited from TM N4 OBS-06/07 and TM "
        "N15 OBS-04 to OBS-08 — the longest-open inheritance in the "
        "project; "
    )
    para.add_run("PQP Rev B").bold = True
    para.add_run(
        " closes the two CRITICAL TM N17 observations on inspection "
        "matrix and FAT scope, which is a necessary prerequisite for "
        "the 40 percent payment milestone under BAE Clause 31 (release "
        "remains gated on the FAT Approval Certificate validation and "
        "the signed Acta de Aprobación FAT, as detailed in Section 2.9 "
        "of the transmittal)."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Drivers requiring action by BW Water:").bold = True
    aplicar_arial(para)

    drivers = [
        (
            "ITP Offsite Rev B (Section 2.10).",
            " Re-issue as Rev C with an explicit declaration of the "
            "ASME X certification scope for the Pressure Vessels "
            "consistent with the BW Water Technical Offer Rev1 "
            "Inspection and Testing Plan (Section 12).",
        ),
        (
            "Grounding & Power Panel Layout Rev E (Section 2.5).",
            " Re-issue as Rev F within fourteen calendar days, "
            "including the complete grounding schedule per NCh Eléct. "
            "4/2003 Section 10.0 (third consecutive transmittal "
            "carrying this item open).",
        ),
        (
            "Instrumentation & Control Cable Schedule Rev 0 "
            "(Section 2.11).",
            " Re-issue as Rev 1 with the VFD communications cable "
            "specification, the dosing-pump signal alignment with the "
            "I/O List, and the nomenclature and tag completeness "
            "corrections.",
        ),
        (
            "RO Cartridge Filter Datasheet Rev D (Section 2.12) and "
            "CIP Cartridge Filter Datasheet Rev C (Section 2.13).",
            " Re-issue addressing the configuration change H to V, "
            "vendor substitution Filtrek to Sysflo, container tie-in "
            "positions and CIP gasket compatibility, in line with "
            "the contractual scope of the Cartridge Filter "
            "specification.",
        ),
    ]
    for bold_part, rest in drivers:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.25)
        para.add_run("• ")
        para.add_run(bold_part).bold = True
        para.add_run(rest)
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph(
        "The Plant Control Philosophy Rev D, expected in this "
        "submittal cycle to address the Code 3 verdict of Transmittal "
        "N18, was not delivered. This is the third consecutive "
        "transmittal carrying the CRITICAL HP Pump permissive open "
        "(TM N15 NOTE-20 → TM N18 OBS-01 → TM N19 carry-forward). "
        "ADASA reserves contractual remedies under Contract C-4300 if "
        "Rev D is not delivered within fourteen calendar days."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "Ten annotated PDFs are attached: five for the Code 3 "
        "documents (the major findings driving the verdict) and "
        "five for the Code 2 documents (the minor refinements "
        "listed against each document). The three Code 1 — "
        "Approved documents carry no annotated PDF. Residual "
        "deliverables for IFC Rev 0 across the Code 1 and Code 2 "
        "documents are listed in Section 3 of the transmittal."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "Please confirm receipt and the target dates for Plant "
        "Control Philosophy Rev D, the five Code 3 re-issues and "
        "the Section 3 deliverables."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("We look forward to your comments.")
    aplicar_arial(para)
    doc.add_paragraph()

    # =========================================================================
    # FIRMA
    # =========================================================================
    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
