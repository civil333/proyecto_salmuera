#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo de remision para la Nota Tecnica P22-NT-09-000-001-0
(Mitigation Plan Clarifications).

Fecha: 25-May-2026 (lunes).
Idioma: ingles (BW Water).

Formato: ejecutivo y directo (~100 palabras), sin tablas, bullets minimos,
bold inline en deadline y codigo NT. Patron Document() directo per CLAUDE.md
§3.4 (correos no usan template ADASA).
"""

import os

from docx import Document
from docx.shared import Inches, Pt

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-05-25_NT-001-Submittal.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_text_with_bold(doc, segments, size=11):
    """Parrafo con segmentos (texto, bold?). segments = [(text, bool), ...]."""
    para = doc.add_paragraph()
    for text, bold in segments:
        run = para.add_run(text)
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(size)
    return para


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # =========================================================================
    # HEADER
    # =========================================================================
    fields = [
        ("Date:", "May 25, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Andrew Sia, Victor Gutierrez, Jeryl F. Regulacion, "
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Mauricio Vallejos, "
            "Jorge Valdes, Tanya Figueroa, Allan Valentos, Sadeep Irugalbandara, "
            "Ghazi Ozair, Nick Huta, Marjan Arsovic, Gerald Ross, "
            "Adzlan Bin Abd Rahim",
        ),
        (
            "Subject:",
            "ADASA – Taltal: Technical Note P22-NT-09-000-001-0 — "
            "Request for Clarifications to Mitigation Plan (C-4300)",
        ),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / "
            "Mitigation Plan + Recovery Schedule, 22-May-2026 / "
            "Cover email Eduardo Yamauchi, 24-May-2026",
        ),
        ("Attachment:", "P22-NT-09-000-001-0_Mitigation-Plan-Clarifications_ADASA.pdf"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    # =========================================================================
    # CUERPO
    # =========================================================================
    para = doc.add_paragraph("Dear Eduardo,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "ADASA acknowledges receipt of the Mitigation Plan and the Recovery "
        "Schedule issued by BW Water on 22-May-2026, transmitted under your "
        "cover email of 24-May-2026."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    add_text_with_bold(doc, [
        ("Please find attached ", False),
        ("P22-NT-09-000-001-0_Mitigation-Plan-Clarifications_ADASA.pdf", True),
        (
            ", issuing 25 technical clarifications structured across the "
            "six action items of the mitigation plan.",
            False,
        ),
    ])
    doc.add_paragraph()

    para = doc.add_paragraph(
        "ADASA's formal position on each item — acceptance, conditional "
        "acceptance, rejection, or requirement of a formal Change Order — "
        "will be issued once the requested technical clarifications are "
        "received and reviewed."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    add_text_with_bold(doc, [
        ("Written technical response is requested by ", False),
        ("end of business Friday 29-May-2026", True),
        (
            ", so that the responses can be reviewed during the weekly "
            "progress meeting scheduled for ",
            False,
        ),
        ("Tuesday 2-Jun-2026", True),
        (".", False),
    ])
    doc.add_paragraph()

    para = doc.add_paragraph(
        "We look forward to receiving the clarifications by the deadline above."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # =========================================================================
    # CIERRE
    # =========================================================================
    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)

    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
