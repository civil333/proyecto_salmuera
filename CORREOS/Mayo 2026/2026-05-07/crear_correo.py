#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Engineering Closure Proposal - Interim Acceptance for Payment Milestone (31-May-2026)
Fecha: 7 de mayo de 2026
Version ejecutiva: directa, sin sub-encabezados, foco en propuesta + tabla + reunion.
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

CONTACTO = {
    "nombre": "Luis Rivera Gonzalez",
    "cargo": "Leader, Infrastructure Engineering",
    "empresa": "ADASA - Aguas de Antofagasta S.A.",
}


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_para(doc, text, size=11):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    return para


def add_para_mixed(doc, fragments, size=11):
    """fragments: lista de tuplas (text, bold)."""
    para = doc.add_paragraph()
    for text, bold in fragments:
        run = para.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.bold = bold
    return para


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else tbl.makeelement(qn("w:tblPr"), {})
    borders = tblPr.makeelement(qn("w:tblBorders"), {})
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.makeelement(
            qn(f"w:{border_name}"),
            {
                qn("w:val"): "single",
                qn("w:sz"): "4",
                qn("w:space"): "0",
                qn("w:color"): "000000",
            },
        )
        borders.append(border)
    tblPr.append(borders)
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)


def add_table_simple(doc, headers, rows, header_size=9, body_size=10):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = True
    set_table_borders(table)

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(header_size)
        shading = cell._element.get_or_add_tcPr()
        shd = shading.makeelement(
            qn("w:shd"),
            {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "D9E2F3"},
        )
        shading.append(shd)

    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            run.font.name = "Arial"
            run.font.size = Pt(body_size)

    doc.add_paragraph()
    return table


def crear_correo():
    output_file = "2026-05-07_Engineering-Closure-Interim-Proposal.docx"
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # ============ HEADER ============
    header_fields = [
        ("Date:", "May 7, 2026"),
        ("From:", f"{CONTACTO['nombre']} - {CONTACTO['cargo']} ({CONTACTO['empresa']})"),
        ("To:", "Eduardo Yamauchi (BW Water); Andrew Sia (BW Water)"),
        ("CC:", "Victor Gutierrez (ADASA); Jeryl Regulacion (BW Water)"),
        (
            "Subject:",
            "Engineering Closure Proposal - Interim Acceptance for Payment "
            "Milestone (31-May-2026)",
        ),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / Master Deliverable Register "
            "P22-IT-06-000-002-0",
        ),
    ]
    for label, value in header_fields:
        para = doc.add_paragraph()
        para.add_run(label + " ").bold = True
        para.add_run(value)
        aplicar_arial(para)

    doc.add_paragraph()

    # ============ APERTURA + PROPUESTA ============
    add_para(doc, "Dear Eduardo, Andrew,")
    doc.add_paragraph()

    add_para(
        doc,
        "We need to close the engineering payment milestone within May. With 50 of "
        "77 Section 1 deliverables already at Code 1/2 and the remaining 27 within "
        "reach, we propose the following two-stage acceptance:",
    )

    doc.add_paragraph()

    # Parrafo ADASA con bold inline en "ADASA accepts"
    add_para_mixed(
        doc,
        [
            ("ADASA accepts ", True),
            (
                "current English revisions as engineering closure for payment "
                "certification - without prejudice to the contractual Spanish IFC "
                "Rev 0 obligation. Upon receipt of the eight items in the table, "
                "ADASA issues an interim Engineering Completion Certificate enabling "
                "the payment milestone to be processed within May.",
                False,
            ),
        ],
    )

    add_para_mixed(
        doc,
        [
            ("BW Water commits ", True),
            (
                "to (i) delivering the eight items below before May 31; "
                "(ii) confirming a binding date for Spanish IFC Rev 0 of the full "
                "Section 1 set, suggested +30 days post-payment.",
                False,
            ),
        ],
    )

    doc.add_paragraph()

    # ============ TABLA DE 8 ITEMS ============
    add_para_mixed(doc, [("Items required by 31-May:", True)])

    add_table_simple(
        doc,
        ["#", "Item", "Required action"],
        [
            ["16", "A/C Thermal Calculation",
             "Rev C - close n+1 redundancy and consolidated thermal load"],
            ["28", "Cable Tray Layout",
             "Rev C - close TM N4 OBS-05/06/07 (88 days) + 5 MAJOR from TM N15"],
            ["67", "Plant Control Philosophy",
             "Rev C - close 17 findings TM N15, including NOTE-20 CRITICAL HP Pump TAGs"],
            ["73", "Detailed PIE",
             "First issue - Hold Point for major fabrication"],
            ["75", "Modbus TCP Communication & Memory Map",
             "First issue - PLC role, byte order, register map"],
            ["81", "LCP Datasheet",
             "Rev B - close OBS-01/02/03 (I/O modules, IP rating, power)"],
            ["91", "Project Quality Plan",
             "Rev B - close PIE Base inspection matrix and FAT scope"],
            ["93", "Alarm and Interlock List",
             "Rev B - fix permeate conductivity unit error and LS-09-002 logic"],
        ],
    )

    add_para(
        doc,
        "The remaining 19 Code 3 / Not Delivered items continue under standard "
        "transmittal cadence and do not block this interim closure.",
    )

    doc.add_paragraph()

    # ============ CIERRE ============
    add_para(
        doc,
        "We remain available to discuss certificate wording, Spanish Rev 0 "
        "schedule and translation responsibility, and look forward to your "
        "comments. The constraint is the May payment milestone.",
    )

    # ============ FIRMA ============
    doc.add_paragraph()
    add_para(doc, "Best regards,")
    doc.add_paragraph()

    para = doc.add_paragraph()
    run = para.add_run(CONTACTO["nombre"])
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)

    add_para(doc, CONTACTO["cargo"])
    add_para(doc, CONTACTO["empresa"])

    doc.save(output_file)
    print(f"Correo generado exitosamente: {output_file}")
    return output_file


if __name__ == "__main__":
    crear_correo()
