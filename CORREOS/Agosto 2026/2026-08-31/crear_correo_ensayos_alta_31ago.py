#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo ADASA -> BW Water. Programa de ensayos de alta presion del 2 al 4 de
septiembre.

CORREO NUEVO, NO respuesta a ninguna cadena (decision del usuario, 31-Ago: el hilo
del Request to witness inspection 003/004 ya no aparece en su bandeja). El asunto se
escribe autonomo y sin prefijo RE. El cuerpo no depende del hilo: cita la reunion de
coordinacion del 26 de agosto y los documentos por su codigo, de modo que se sostiene
solo. Las cross-references quedan explicitas para que la trazabilidad no dependa del
encadenamiento del cliente de correo.

TRES CARGAS, en este orden:

  1. LA PRESION VINCULANTE POR LINEA, antes del ensayo del miercoles. Cada linea a
     1,5 x su presion de diseno segun la columna de hidrostatica de la Line List
     P22-LI-09-009-003 Rev 0. NO hay un valor unico para el super duplex: las once
     lineas van a 75, 90, 120 o 135 barG. Los 135 corresponden SOLO a las cuatro de
     90 barG de diseno; aplicarlos parejo sobrepresiona SIETE de las ONCE, hasta 2,7
     veces el diseno en las salidas de los dos Fedco. Y el ensayo del 20-Ago de
     DA-SSD-DN100-09-003 a 90 barG ES CORRECTO y se mantiene: su diseno es 60.

  2. LO COMPROMETIDO PARA EL VIERNES 28 Y NO RECIBIDO. El P&ID, la Line List y los
     planos de fabricacion con la numeracion unificada, y la presion exacta por
     linea que debia venir con esa entrega. Los submittals 25007-0086 a 0088 del
     28-Ago no traen ninguno de los tres.

  3. EL PROGRAMA DEL 2 AL 4 DE SEPTIEMBRE, dia por dia, con TAG y presion; la
     solicitud al tercero inspector con ADASA en copia; las tres repeticiones; la
     cuarta linea CIP retenida; y la hidrostatica del RO Vessel, fuera de alcance
     desde el 5 de agosto.

CALIBRACIONES QUE EVITAN UN RECLAMO REFUTABLE:

  - Las dos jornadas de BAJA presion del 27 y el 28 NO se objetan. Estan bien
     ejecutadas y cierran la parte de baja del alcance que el Request 003 recorto el
     5 de agosto (PRG-13). Lo que no avanza es el Punto de Detencion de alta.
  - NO se afirma que la solicitud al inspector no se emitio: se pide confirmarla y
     que ADASA quede en copia. El correo entrante esta en hold y una ausencia no se
     afirma sin barrer las tres cadenas del frente.
  - NO se cita el registro de la reunion del 26-Ago: es un resumen automatico
     interno, no una minuta oficial. La regla se reafirma en nombre propio, anclada
     en la Line List aprobada y en el Transmittal N35.

PLAZO: martes 1 de septiembre, cierre de la jornada de Penang. Penang va DOCE horas
adelante de Chile, asi que el plazo se escribe en hora de Penang. Es el unico dia
habil completo antes del ensayo del miercoles.

FUERA DEL CORREO a proposito: el plazo de entrega vencido y la multa de la Clausula
43.1 letra b; el reclamo de desempeno contra el tercero inspector, que va por su
propia cadena y en espanol; el saldo de jornadas contratadas y las tarifas, porque
el inspector figura entre los destinatarios; y el espesor de pelicula seca del marco
del skid, que es otro frente.

Ingles. Document() directo, sin template ADASA. Estado: BORRADOR.
"""
import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-08-31_HP-Tests-Programme.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11, space=6):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(space)
    aplicar_arial(p, size)
    return p


def add_segments(doc, segs, size=11, space=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    for text, bold in segs:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return p


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(3)
    r0 = p.add_run("•  " + lead)
    r0.bold = True
    r0.font.name = "Arial"
    r0.font.size = Pt(size)
    r1 = p.add_run(rest)
    r1.font.name = "Arial"
    r1.font.size = Pt(size)
    return p


def espaciador(doc):
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


# Las once lineas de super duplex de la Line List P22-LI-09-009-003 Rev 0, aprobada
# en Codigo 1 en el Transmittal N29. Sin columna de servicio: no se inventa lo que no
# esta verificado linea por linea.
SDX_LINES = [
    ("DA-SSD-DN100-09-003", "60", "90", "Tested and accepted, 20 August"),
    ("DA-SSD-DN100-09-004", "80", "120", "Not tested"),
    ("DA-SSD-DN80-09-005", "80", "120", "Not tested"),
    ("DA-SSD-DN80-09-006", "90", "135", "Not tested"),
    ("DA-SSD-DN65-09-007", "90", "135", "Not tested"),
    ("DA-SSD-DN65-09-008", "50", "75", "Not tested"),
    ("DA-SSD-DN65-09-009", "50", "75", "Not tested"),
    ("CP-SSD-DN100-09-014", "80", "120", "To be repeated"),
    ("CP-SSD-DN80-09-015", "90", "135", "To be repeated"),
    ("CP-SSD-DN80-09-044", "80", "120", "To be repeated"),
    ("CP-SSD-DN65-09-045", "90", "135", "Not tested, on hold"),
]


def add_table_lines(doc, size=9):
    headers = (
        "Line",
        "Design barG",
        "Test pressure barG\n(1.5 x design)",
        "Status",
    )
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(size)
    for row in SDX_LINES:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(value)
            run.font.name = "Arial"
            run.font.size = Pt(size)
    return table


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "August 31, 2026"),
        ("From:", CONTACTO + " - Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi, Stephane Gehant, Magdier Arias, "
                "Lokman Hakim Mat, Mohd Adnin Bin Zulkaflee - BW Water"),
        ("CC:", "Muhammad Fadhil Bin Abdul Wahid - BW Water; "
                "Ahmad Hazwan, Wan Mohd Adli W Yahya, Emylia Rosli - Bureau Veritas; "
                "Victor Gutierrez - ADASA"),
        ("Subject:", "25007 TALTAL - Super duplex test pressures and test programme "
                     "for 2 to 4 September"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    espaciador(doc)

    add_para(doc, "Dear Eduardo, Stephane, Lokman,")

    # LEAD. El estado real del punto Hold, en cifras, y lo que falta para el miercoles.
    add_segments(doc, [
        ("The high pressure tests start on Wednesday and the inputs that govern them "
         "are still open.", True),
        (" Of the eleven super duplex lines, one has a qualifying test, three were "
         "tested at 7.5 barG and must be repeated, and seven have never been tested. "
         "Last week the two witnessed days went to the low pressure PVC spools on 27 and "
         "28 August. Both are correctly tested and recorded, and they close the low "
         "pressure scope that had been open since 5 August. The Friday request form "
         "covered high and low pressure, and no super duplex line has been tested since "
         "21 August.", False),
    ])

    # CARGA 1. LA REGLA, con la advertencia de sobrepresion primero porque previene dano.
    add_segments(doc, [
        ("Each line is tested at 1.5 times its own design pressure.", True),
        (" The figure is the hydrotest column of the approved Line List "
         "(P22-LI-09-009-003 Rev 0), it follows row 5.2 of the PIE Base "
         "(P22-IT-09-000-001-0) and it agrees with ASME B31.3 paragraph 345.4.2. There "
         "is no single figure for the section: the eleven lines are tested at 75, 90, "
         "120 or 135 barG, and 135 belongs to the four designed for 90 barG. Applied "
         "across the section it would over-pressure seven of the eleven, up to 2.7 times "
         "design on the two turbocharger brine outlets. And the test of "
         "DA-SSD-DN100-09-003 on 20 August at 90 barG is correct and stands, because its "
         "design pressure is 60 barG.", False),
    ])

    add_table_lines(doc)
    espaciador(doc)

    # CARGA 2. Lo vencido el viernes. Sin citar el registro de la reunion.
    add_segments(doc, [
        ("Two of the items agreed at the coordination meeting of 26 August were due on "
         "Friday 28 and have not reached us.", True),
        (" The updated P&ID, Line List and fabrication drawings with the line numbering "
         "unified against the approved list, and the exact test pressure of each flagged "
         "spool, which was to be issued with that submittal. The three submittals "
         "received on 28 August carry none of them. Until the line numbers are unified "
         "there is nothing to verify Wednesday against.", False),
    ])

    # CARGA 3. El programa.
    add_segments(doc, [
        ("What we need in writing by the close of business on Tuesday 1 September, "
         "Penang time:", True),
    ], space=3)
    add_bullet_lead(
        doc, "The programme for 2, 3 and 4 September, day by day",
        ": which spool is tested each day, with its tag as written in the approved Line "
        "List and the pressure that will be applied to it.")
    add_bullet_lead(
        doc, "The Inspection Request issued to the third party inspector for those three "
        "days", ", with ADASA copied. If it has already gone out, please forward it to us.")
    add_bullet_lead(
        doc, "The three tests of 20 and 21 August repeated at the listed pressure",
        ": CP-SSD-DN100-09-014 at 120 barG, CP-SSD-DN80-09-015 at 135 barG and "
        "CP-SSD-DN80-09-044 at 120 barG.")
    add_bullet_lead(
        doc, "CP-SSD-DN65-09-045 at 135 barG",
        ". It is the fourth CIP line, it has never been tested, and the hold on it "
        "stands until the pressures above are confirmed.")
    add_bullet_lead(
        doc, "The RO Vessel hydrostatic test",
        ", a hold point of the Third-Party Shop Inspection Notice (P22-NT-09-000-002-0) "
        "that has been out of the inspection scope since 5 August.")
    add_bullet_lead(
        doc, "The list of every super duplex spool tested to date",
        ", with tag, pressure applied, date and gauge used. Requested on 21 August and "
        "still outstanding.")
    espaciador(doc)

    # RESERVA.
    add_segments(doc, [
        ("The hold stands.", True),
        (" Row 5.2 of the Inspection and Test Plan (P22-BA-09-000-004 Rev 0) is a hold "
         "point under ADASA control. The written confirmation of the test pressure of "
         "each line has been outstanding since 20 August, and any super duplex test run "
         "before it reaches us is treated by ADASA as not executed, whatever witness "
         "record accompanies it.", False),
    ])

    add_para(doc, "Best regards,", space=0)
    espaciador(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in [
        "Infrastructure Engineering Lead - ADASA - Aguas de Antofagasta S.A.",
        "lrivera@aguasantofagasta.cl",
    ]:
        add_para(doc, line, space=0)

    fijar_idioma(doc, LANG)

    docx_metadata.apply_core_properties(
        doc,
        title="Taltal - High pressure test programme for 2 to 4 September",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="Taltal SWRO - Binding test pressure per line and witnessed test "
                "programme for 2 to 4 September 2026",
        comments="Aguas de Antofagasta S.A.",
        language="en-US",
    )
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas Antofagasta")
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
