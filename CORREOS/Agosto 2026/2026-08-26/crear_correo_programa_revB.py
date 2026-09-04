#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo contractual sobre el Project Schedule Rev B (P22-BA-09-000-001), recibido
en la submittal 25007-0084 el 24-Ago-2026.

**CADENA SEPARADA del transmittal.** El TM N36 dispone el codigo de respuesta del
documento por el cumplimiento de sus dos observaciones documentales. Este correo
plantea el fondo: el desplazamiento del programa contra el Plazo de Entrega, que
vencio el 03-Ago-2026. Los dos salen el mismo dia con los codigos ADASA
explicitos, de modo que ninguna referencia cruzada apunta a un documento no
recibido.

REGLA DURA APLICADA — NO SE CURSA CIFRA DE MULTA. La cita es doble y el correo la
respeta: la Clausula 27 fija el plazo, la 43.1 letra b la multa diaria y la 43.4
el tope. **El correo NO calcula dias ni monto**, porque eso exige leer la fecha
de la Notificacion de Adjudicacion en el documento original de ADASA, y esa
verificacion esta pendiente. Calcularla sobre el rotulo del cronograma del
proveedor deja el reclamo atacable. El correo reserva la posicion y pide el dato.

CIFRAS DEL REV B QUE SI SE CITAN, todas verificadas en el propio documento:
  - Ex-works Penang: linea base 15-Ago-2026, proyectado 21-Sep-2026 (+37 dias)
  - Entrega en sitio: linea base 30-Sep-2026, proyectado 13-Nov-2026 (+44 dias)
  - Fin de programa:  linea base 19-Nov-2026, proyectado 02-Ene-2027 (+44 dias)
  - FAT del sistema:  proyectado 7 al 18-Sep-2026, ventana declarada de 15 a 11 dias
  - La columna de linea base conserva los cuatro anclajes adoptados el 09-Jun-2026

FUERA DEL CORREO, a proposito: los codigos de respuesta de los nueve documentos y
el detalle de las dos observaciones documentales, que viven en el transmittal.

Ingles. Document() directo, sin template ADASA. Estado: BORRADOR.
"""
import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-08-26_Project-Schedule-RevB.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph(text)
    aplicar_arial(p, size)
    return p


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for text, bold in segs:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return p


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r0 = p.add_run("•  " + lead)
    r0.bold = True
    r0.font.name = "Arial"
    r0.font.size = Pt(size)
    r1 = p.add_run(rest)
    r1.font.name = "Arial"
    r1.font.size = Pt(size)
    return p


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "August 26, 2026"),
        ("From:", CONTACTO + " - Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi - BW Water"),
        ("CC:", "Stephane Gehant, Andrew Sia, Fitri Indriyani - BW Water; Victor "
                "Gutierrez, Jorge Guevara, Ronald Pellejero - ADASA"),
        ("Subject:", "TALTAL - Project Schedule Rev B - contractual position on the "
                     "delivery period"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_segments(doc, [
        ("We acknowledge receipt of the ", False),
        ("Project Schedule Rev B (P22-BA-09-000-001)", True),
        (", issued on 24 August under submittal 25007-0084. Its response code is set "
         "out in Transmittal N36, issued today. This message addresses the substance "
         "of the programme, which we keep on this thread.", False),
    ])
    blank(doc)

    add_segments(doc, [
        ("Rev B does not re-open the recovery baseline", True),
        (" adopted on 09 June 2026, and we note that. What it does is project the "
         "following against that baseline:", False),
    ])
    blank(doc)

    add_bullet_lead(doc, "Ex-works Penang",
                    " — baseline 15 August 2026, now projected 21 September 2026, "
                    "thirty-seven days later.")
    add_bullet_lead(doc, "Delivery to site",
                    " — baseline 30 September 2026, now projected 13 November 2026, "
                    "forty-four days later.")
    add_bullet_lead(doc, "Programme finish",
                    " — baseline 19 November 2026, now projected 02 January 2027, "
                    "forty-four days later.")
    blank(doc)

    add_segments(doc, [
        ("The delivery period expired on 03 August 2026.", True),
        (" Clause 27 of the Special Administrative Conditions fixes it, Clause 43.1(b) "
         "sets the daily liquidated damages for exceeding it, and Clause 43.4 caps them. "
         "ADASA reserves its position under those clauses in full. Before quantifying "
         "anything we will verify the date of the Notice of Award against our own "
         "record, and we ask you to confirm that date in writing so that both sides "
         "compute from the same origin.", False),
    ])
    blank(doc)

    add_segments(doc, [
        ("Three points need an answer with the programme.", True),
        (" First, the Factory Acceptance Test is now projected between 7 and 18 "
         "September with a declared duration that fell from fifteen days to eleven, "
         "against a scope that has not been reduced; please state how the full test "
         "programme fits that window. Second, the shipping leg between ex-works and "
         "site now runs fifty-three days against forty-six in the baseline, with no "
         "breakdown of sea transit, customs and inland transport; please break it "
         "down. Third, the programme carries no vessel certification activity and no "
         "pressure test of any kind, which is the subject of the two observations "
         "carried in Transmittal N36.", False),
    ])
    blank(doc)

    add_segments(doc, [
        ("We ask for your written response by Wednesday 02 September 2026", True),
        (", together with the recovery measures you intend to apply to the ex-works "
         "date.", False),
    ])
    blank(doc)

    add_para(doc, "Best regards,")
    blank(doc)

    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in [
        "Infrastructure Engineering Lead - ADASA - Aguas de Antofagasta S.A.",
        "lrivera@aguasantofagasta.cl",
    ]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)

    docx_metadata.apply_core_properties(
        doc,
        title="TALTAL - Project Schedule Rev B - contractual position",
        author="Luis Rivera Gonzalez",
        language="en-US",
        comments="Correo contractual sobre el Project Schedule Rev B. Cadena de "
                 "programa, separada de la de transmittals.",
    )
    doc.save(OUTPUT_FILE)
    print("Correo generado:", OUTPUT_FILE)


if __name__ == "__main__":
    crear_correo()
