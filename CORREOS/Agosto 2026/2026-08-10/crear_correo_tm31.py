#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo de cobertura del Transmittal N31 (P22-TM-09-000-031-0) a BW Water.
Cadena regular de transmittals. Ingles. Document() directo, sin template ADASA.

Cubre los submittals 25007-0072 (E72) y 25007-0074 (E73), cinco documentos.
Veredicto global 2 - Approved as noted: 1 Code 1, 3 Code 2, y la Plant Control
Philosophy Rev 0 devuelta SIN codigo.

CORREO DE UNA PAGINA, cuatro bloques. Una linea de veredicto, la disposicion por
documento en una clausula cada uno, el HINCAPIE -diez condiciones sin incorporar
en cinco documentos ya emitidos a Rev 0 para construccion- y el primer punto a
atacar. Todo lo demas -las razones documentadas de cada punto, los seis TAG del
HMI y el detalle de la Seccion 3- vive en el transmittal y no se repite aqui.
No propone reunion.

EL RECUENTO DE DIEZ, verificado sobre los libros mayores: seis condiciones de la
ENTREGA 71 repartidas en cuatro documentos (PMI 2, Visual 1, Painting 2, HP/LP 1;
el Structural incorporo las dos suyas y queda fuera) mas cuatro de la Plant
Control Philosophy. El RO Vessel Hydrostatic Test Procedure Rev 0 NO entra: su
condicion se cumplio y es Codigo 1; su punto de vigencia de calibracion es un
pedido nuevo. Si cambia alguno de estos numeros, cambiarlo tambien en la Seccion
3 del transmittal y en su resumen ejecutivo.

El PDF del transmittal va ADJUNTO; los tres PDF anotados, por ENLACE de
descarga. El enlace queda ademas impreso en la Seccion 4 del transmittal, de
modo que sobrevive al reenvio del adjunto sin el correo: en el N30 el parrafo
del enlace no salio en el correo y los anotados quedaron alcanzables solo por
el texto impreso.

Estado: BORRADOR, pendiente de revision y envio.
"""
import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-08-10_Transmittal-N31.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"
# REEMPLAZAR por el enlace Synology real antes de enviar. Debe coincidir con el
# DOWNLOAD_LINK de crear_transmittal.py.
DOWNLOAD_LINK = (
    "https://lrg.synology.me:6501/d/s/19PutLaANldko09k2QUOUibPNpL295Ze/"
    "KyDEJBacXwQhWHVkM0m17KgQlqXF73CB-L7NAkDq8ag0"
)


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph(text)
    aplicar_arial(p, size)
    return p


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for text, bold in segs:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return p


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r0 = p.add_run("•  " + lead)
    r0.bold = True
    r0.font.name = "Arial"
    r0.font.size = Pt(size)
    r1 = p.add_run(rest)
    r1.font.name = "Arial"
    r1.font.size = Pt(size)
    return p


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), "Arial")
    rf.set(qn("w:hAnsi"), "Arial")
    rpr.append(rf)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0563C1")
    rpr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    h.append(run)
    paragraph._p.append(h)


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "August 10, 2026"),
        ("From:", CONTACTO + " - Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi - BW Water"),
        ("CC:", "Andrew Sia, Victor Gutierrez, Jeryl F. Regulacion, Jorge Guevara, "
                "Ronald Pellejero, Stephane Gehant"),
        ("Subject:", "25007 Taltal - Technical Review Transmittal N31 - submittals "
                     "25007-0072 and 25007-0074"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_segments(doc, [
        ("Transmittal N31 (", False), ("P22-TM-09-000-031-0", True),
        (") is attached. It responds to submittals 25007-0072 and 25007-0074, five documents, "
         "reviewing only whether these revisions close ADASA's earlier observations. "
         "Overall verdict: ", False),
        ("2 - Approved as noted", True), (".", False),
    ])
    blank(doc)

    add_para(doc, "Disposition:")
    add_bullet_lead(doc, "Plant Control Philosophy Rev 0 - no response code. ",
                    "Four conditions of the Transmittal N28 approval remain open.")
    add_bullet_lead(doc, "RO Vessel Hydrostatic Test Procedure Rev 0 - Code 1. ",
                    "Confirm the calibration valid at the date of the test.")
    add_bullet_lead(doc, "Tie-In Point Layout Rev B - Code 2. ",
                    "State the design pressure at the brine feed tie-in point.")
    add_bullet_lead(doc, "GA of SWRO System Skid Rev B - Code 2. ",
                    "Correct three document references in the new notes.")
    add_bullet_lead(doc, "HMI Display Screenshot Rev B - Code 2. ",
                    "Add the electrical-variable readings and reconcile six tags.")
    blank(doc)

    add_segments(doc, [
        ("Ten conditions of ADASA approvals are now outstanding in documents already issued "
         "at Rev 0 for construction. ", True),
        ("Six belong to the PMI, Visual, Painting and HP and LP Pressure Test procedures of "
         "submittal 25007-0071, raised on 6 August. Four belong to the Plant Control "
         "Philosophy of this submittal. Once a document is at Rev 0 there is no response "
         "code left to return it with, so the correction can only be made at the next issue "
         "while fabrication and inspection proceed against the document as it stands.",
         False),
    ])
    blank(doc)

    add_segments(doc, [
        ("Start with the temperature-sensor mapping. ", True),
        ("The winding and bearing tags are reversed on the CIP pump in the Plant Control "
         "Philosophy and duplicated on the high-pressure pump in the HMI screens. That item "
         "holds the Factory Acceptance Test Procedure sign-off on motor temperature "
         "protection, open since Transmittal N27. Please confirm how and when each of the "
         "ten points will be closed.", False),
    ])
    blank(doc)

    add_para(doc, "Attached: TRANSMITTAL N31 ADASA-BW_WATER.pdf")
    p = add_segments(doc, [
        ("The three annotated PDFs, listed in Section 4 of the transmittal, are available "
         "here: ", False)])
    add_hyperlink(p, DOWNLOAD_LINK, DOWNLOAD_LINK)
    blank(doc)
    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in ["Infrastructure Engineering Lead - ADASA - Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)
    docx_metadata.apply_core_properties(
        doc, title="Taltal - Technical Review Transmittal N31",
        author="Luis Rivera Gonzalez", last_modified_by="Luis Rivera Gonzalez",
        subject="Transmittal N31 - submittals 25007-0072 and 25007-0074",
        comments="Correo de cobertura del Transmittal N31 a BW Water.",
        language="en-US")
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas de Antofagasta S.A.")
    print("Correo generado: " + OUTPUT_FILE)
    if DOWNLOAD_LINK.startswith("PENDIENTE"):
        print("AVISO: enlace de descarga en marcador de posicion. "
              "Reemplazar antes de enviar, y que coincida con crear_transmittal.py.")


if __name__ == "__main__":
    crear_correo()
