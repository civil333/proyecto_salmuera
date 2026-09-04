#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo a BW Water por el paquete de reporteria semanal. Cadena de COORDINACION
SEMANAL. Ingles. Document() directo, sin template ADASA.

SOLO BW WATER. Bureau Veritas NO va: el desempeno de reporteria del proveedor
frente a ADASA no se expone a un tercero contratado por ADASA. La jornada de
inspeccion del 7-Ago, que si le compete, va en el correo conjunto
crear_correo_inspection_002_records.py.

HECHOS VERIFICADOS CARPETA POR CARPETA:
  - Ultimo DDSR en poder de ADASA: 25007_Taltal_DDSR_2026.07.20.
  - Los paquetes del 27-Jul y del 3-Ago llegaron SIN DDSR (traen informe de
    avance semanal y tracker de procura, nada mas).
  - El de esta semana aun no llega.

CUIDADO CON LA FECHA: los paquetes se archivan los MARTES (mtime 21-Jul, 28-Jul,
4-Ago), asi que hoy lunes NO se afirma atraso del paquete de la semana, solo que
no ha llegado. Lo afirmable es la ausencia del DDSR en los dos anteriores.

NO SE RECLAMA el tracker de procura: si cambia de contenido cada semana pese a
conservar el nombre de archivo (md5 distintos en las tres semanas).
NO SE MENCIONA el Plazo de Entrega vencido ni la multa: cadena contractual.
NO puede decir "como pedimos el 5 de agosto": ese correo quedo en BORRADOR y
nunca se envio.

Estado: BORRADOR, pendiente de revision y envio.
"""
import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-08-10_Weekly-Reporting-Package.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph(text)
    aplicar_arial(p, size)
    return p


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for text, bold in segs:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return p


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r0 = p.add_run("•  " + lead)
    r0.bold = True
    r0.font.name = "Arial"
    r0.font.size = Pt(size)
    r1 = p.add_run(rest)
    r1.font.name = "Arial"
    r1.font.size = Pt(size)
    return p


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), "Arial")
    rf.set(qn("w:hAnsi"), "Arial")
    rpr.append(rf)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0563C1")
    rpr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    h.append(run)
    paragraph._p.append(h)


def blank(doc):
    doc.add_paragraph()




def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "August 10, 2026"),
        ("From:", CONTACTO + " - Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi, Fitri Indriyani - BW Water"),
        ("CC:", "Victor Gutierrez, Ronald Pellejero, Jorge Guevara - ADASA; Jeryl F. "
                "Regulacion, Stephane Gehant - BW Water"),
        ("Subject:", "25007 Taltal - Weekly reporting package and Document and Drawing "
                     "Status Report"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_para(doc, "The weekly reporting package has not reached us this week. The most "
                  "recent one we hold is that of 3 August, with the progress report for "
                  "week 31, covering 27 July to 2 August.")
    blank(doc)

    add_para(doc, "The Document and Drawing Status Report is a separate matter: the last "
                  "one we hold is dated 20 July. The packages of 27 July and 3 August each "
                  "carried their progress report but not the status report. That report is "
                  "the reference ADASA uses to follow document positions and planned "
                  "re-submittal dates, and Transmittal N31, issued today, cites the one of "
                  "20 July for three drawings whose re-issue dates have passed.")
    blank(doc)

    add_para(doc, "When can we expect this week's package with the progress report for "
                  "week 32, and will the status report come with it?")
    blank(doc)

    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in ["Infrastructure Engineering Lead - ADASA - Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)
    docx_metadata.apply_core_properties(
        doc, title="Taltal - Weekly reporting package",
        author="Luis Rivera Gonzalez", last_modified_by="Luis Rivera Gonzalez",
        subject="Weekly reporting package and Document and Drawing Status Report",
        comments="Consulta a BW Water por el paquete de reporteria semanal y el DDSR, "
                 "ausente desde el 20 de julio.",
        language="en-US")
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas de Antofagasta S.A.")
    print("Correo generado: " + OUTPUT_FILE)


if __name__ == "__main__":
    crear_correo()
