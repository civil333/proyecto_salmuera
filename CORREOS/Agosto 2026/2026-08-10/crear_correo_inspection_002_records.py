#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo CONJUNTO BW Water + Bureau Veritas por los registros de la jornada de
inspeccion del viernes 7-Ago. Reply-To sobre el hilo "25007 TALTAL - Request to
witness inspection 002". Ingles. Document() directo, sin template ADASA.

CADA PEDIDO A SU DUENIO, que es lo que hace util el correo conjunto:
  - A BW Water: el formulario 002 firmado con su casilla de resultado, los
    informes de liquidos penetrantes, y los WPS y calificaciones de soldador
    que se revisaron.
  - A Bureau Veritas: el informe de inspeccion de esa jornada, en el formato en
    que emitio el de la Visita 1 (BVM-IR001-28072026).

ANCLA DE PLAZO: la Visita 1 fue el martes 28-Jul y el informe con sus anexos
llego al dia siguiente, el 29-Jul, en el zip del hilo del Request 001. Es el
precedente de los propios destinatarios.

SI SE NOMBRA A BW WATER pese a que Bureau Veritas esta en el correo: aplica la
excepcion documentada de [[feedback_bv_no_nombrar_bw_water]] para los correos
CONJUNTOS de coordinacion, donde el fabricante esta copiado y ponerlos en
contacto es el proposito. La regla de no nombrarlo rige los correos de ADASA
SOLO a BV.

EL PAQUETE DE REPORTERIA SEMANAL NO VA AQUI: es desempeno de BW Water frente a
ADASA y no se expone a un tercero contratado por ADASA. Va en
crear_correo_weekly_package.py, a BW Water solamente.

Estado: BORRADOR, pendiente de revision y envio.
"""
import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-08-10_Inspection-002-Records.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph(text)
    aplicar_arial(p, size)
    return p


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for text, bold in segs:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return p


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "August 10, 2026"),
        ("From:", CONTACTO + " - Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi, Magdier Arias, Mohd Adnin Bin Zulkaflee, Muhammad "
                "Fadhil Bin Abdul Wahid - BW Water; Ahmad Hazwan, Wan Mohd Adli W Yahya, "
                "Carlo Montecinos, Emylia Rosli - Bureau Veritas"),
        ("CC:", "Victor Gutierrez, Ronald Pellejero, Jorge Guevara - ADASA; Stephane "
                "Gehant, Lokman Hakim Bin Mat - BW Water"),
        ("Subject:", "RE: 25007 TALTAL - Request to witness inspection 002 - records of "
                     "the 7 August inspection"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear all,")
    blank(doc)

    add_para(doc, "Nothing has reached us from the witness inspection held in Penang on "
                  "Friday 7 August, and I would like to know when to expect it.")
    blank(doc)

    add_para(doc, "Inspection Request Form 002 scheduled four scopes for that day: super "
                  "duplex piping spool fabrication, penetrant testing of the root and "
                  "capping runs on those welds, skid frame fabrication, and review of the "
                  "WPS and welder qualification records.")
    blank(doc)

    add_segments(doc, [
        ("From BW Water: ", True),
        ("the signed Form 002 with its inspection result, the penetrant test reports, and "
         "the WPS and welder qualification records that were reviewed.", False),
    ])
    add_segments(doc, [
        ("From Bureau Veritas: ", True),
        ("the inspection report for that day, in the form issued for the visit of 28 July "
         "(BVM-IR001-28072026).", False),
    ])
    blank(doc)

    add_para(doc, "For the visit of 28 July the report and its annexes reached us the "
                  "following day. These are also the first fabrication records the project "
                  "has, and the fabrication and testing dossier remains outstanding.")
    blank(doc)

    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in ["Infrastructure Engineering Lead - ADASA - Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)
    docx_metadata.apply_core_properties(
        doc, title="Taltal - Records of the 7 August witness inspection",
        author="Luis Rivera Gonzalez", last_modified_by="Luis Rivera Gonzalez",
        subject="Inspection Request 002 - records of the 7 August inspection",
        comments="Correo conjunto a BW Water y Bureau Veritas por los registros de la "
                 "jornada de inspeccion del 7 de agosto.",
        language="en-US")
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas de Antofagasta S.A.")
    print("Correo generado: " + OUTPUT_FILE)


if __name__ == "__main__":
    crear_correo()
