#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo de cobertura del Transmittal N30 (P22-TM-09-000-030-0) a BW Water.
Cadena regular de transmittals. Ingles. Document() directo, sin template ADASA.

RE-ESCOPEADO. El correo se redacto el 23-Jul con la sola E68 y nunca se envio.
Al llegar la E69 y la E70 antes de la emision, el transmittal absorbe las tres
entregas y este correo se reemite con fecha 05-Ago-2026. La carpeta se movio de
CORREOS/Julio 2026/2026-07-23/ a CORREOS/Agosto 2026/2026-08-05/ porque la fecha
de la carpeta debe coincidir con la del header (CLAUDE.md Seccion 3.4).

Veredicto 3 - To be revised sobre cinco documentos y tres submittals, mas un
conjunto de planos de taller devuelto como no recibido.

CORREO DE UNA PAGINA, 279 palabras. Solo lo esencial: una linea de veredicto, la
disposicion por documento en una clausula cada uno, y el UNICO punto que exige
respuesta antes de la hidrostatica del 12 y 13 de agosto. Todo lo demas -el
detalle de los dos hallazgos de presion, el fundamento de la devolucion del
conjunto de taller, la familia de Control vencida y el computo del plazo de la
Clausula 37.2- vive en el transmittal y no se repite aqui. No propone reunion.

El PDF del transmittal va ADJUNTO (0,21 MB); los dos PDF anotados, por ENLACE de
descarga (14,3 MB juntos). Estado: BORRADOR, pendiente de revision y envio.
"""
import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-08-05_Transmittal-N30.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"
# Los dos PDF anotados pesan 14,3 MB juntos, sobre el limite habitual de adjunto
# del correo corporativo: van por enlace de descarga. Adjunto va solo el PDF del
# transmittal (0,21 MB). El enlace tambien queda impreso en la Seccion 4 del
# transmittal, de modo que sobrevive al reenvio del adjunto sin el correo.
DOWNLOAD_LINK = (
    "https://lrg.synology.me:6501/d/s/19Llot7JjfVf8lcyurHIZDfL6nZ13HIC/"
    "V4Ej-jY6AOk8cbU29kS4w7Z3SD8wrWSB-4r-gvdeEZw0"
)


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph(text)
    aplicar_arial(p, size)
    return p


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for text, bold in segs:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return p


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r0 = p.add_run("•  " + lead)
    r0.bold = True
    r0.font.name = "Arial"
    r0.font.size = Pt(size)
    r1 = p.add_run(rest)
    r1.font.name = "Arial"
    r1.font.size = Pt(size)
    return p


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), "Arial")
    rf.set(qn("w:hAnsi"), "Arial")
    rpr.append(rf)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0563C1")
    rpr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    h.append(run)
    paragraph._p.append(h)


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()
    # El estilo Normal de python-docx trae interlineado 1,08 y 8 pt despues de
    # cada parrafo. En un correo de una pagina eso son ~10 lineas regaladas y la
    # firma se cae a la hoja 2. La separacion la dan los parrafos en blanco.
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "August 5, 2026"),
        ("From:", CONTACTO + " - Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi - BW Water"),
        ("CC:", "Andrew Sia, Victor Gutierrez, Jeryl F. Regulacion, Jorge Guevara, "
                "Ronald Pellejero, Stephane Gehant"),
        ("Subject:", "25007 Taltal - Technical Review Transmittal N30 - submittals "
                     "25007-0068, 25007-0069 and 25007-0070"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_segments(doc, [
        ("Transmittal N30 (", False), ("P22-TM-09-000-030-0", True),
        (") is attached. It covers five documents across submittals 25007-0068, "
         "25007-0069 and 25007-0070. Overall verdict: ", False),
        ("3 - To be revised", True), (".", False),
    ])
    blank(doc)

    add_para(doc, "Disposition:")
    add_bullet_lead(doc, "PLC and HMI Panel Component Rev 0 - Code 1. ",
                    "Align the four documents carrying the other terminal number.")
    add_bullet_lead(doc, "Fabrication and Testing Dossier Index Rev A - Code 3. ",
                    "Re-issue as Rev B.")
    add_bullet_lead(doc, "Piping Layout Rev C - Code 2. ",
                    "Add the tie-in schedule and the flange class at Rev 0.")
    add_bullet_lead(doc, "Shop set 25007-ME-PI-0901-0006 to -0016 - not received. ",
                    "Submit it as a deliverable in its own right.")
    add_bullet_lead(doc, "3D Model Rev A - Code 2. ",
                    "Identify the file by its code and revision; reconcile its tags.")
    add_bullet_lead(doc, "Differential Pressure Switch Rev B - Code 1. ",
                    "Issue directly at IFC Rev 0.")
    blank(doc)

    add_segments(doc, [
        ("One item needs an answer before the hydrostatic test of 12 and 13 August. ",
         True),
        ("The shop fabrication sheets show ANSI 150# threaded austenitic couplings at the "
         "instrument tappings, and an unbroken boundary between the super duplex and PVC "
         "systems, on lines the approved Line List rates up to 90 barG design. Please "
         "confirm in writing whether any of these branches has already been fabricated.",
         False),
    ])
    blank(doc)

    add_para(doc, "Attached: TRANSMITTAL N30 ADASA-BW_WATER.pdf")
    p = add_segments(doc, [
        ("The two annotated PDFs, for the Dossier Index and the Piping Layout, are "
         "available here: ", False)])
    add_hyperlink(p, DOWNLOAD_LINK, DOWNLOAD_LINK)
    blank(doc)
    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in ["Infrastructure Engineering Lead - ADASA - Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)
    docx_metadata.apply_core_properties(
        doc, title="Taltal - Technical Review Transmittal N30",
        author="Luis Rivera Gonzalez", last_modified_by="Luis Rivera Gonzalez",
        subject="Transmittal N30 - submittals 25007-0068, 25007-0069 and 25007-0070",
        comments="Correo de cobertura del Transmittal N30 a BW Water.",
        language="en-US")
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas de Antofagasta S.A.")
    print("Correo generado: " + OUTPUT_FILE)


if __name__ == "__main__":
    crear_correo()
