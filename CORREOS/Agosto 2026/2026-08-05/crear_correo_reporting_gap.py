#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo de reclamo por dos faltas de reporteria del paquete semanal, a BW Water.
Cadena de coordinacion semanal, NO la de transmittals (CLAUDE.md Seccion 3.4:
cadena separada por instrumento).

Dos puntos, nada mas:
  1. El DDSR falta por segunda semana consecutiva (paquetes del 27-Jul y 03-Ago).
  2. El cronograma del 04-Ago llego colapsado a 4 paginas contra las 12 del
     27-Jul, y fuera del paquete semanal.

Tono: transaccional y directo, sin acusar. Se pide reposicion en el paquete del
lunes 10-Ago. NO se menciona el atraso del Plazo de Entrega ni la exposicion a
multa: eso vive en la cadena contractual.

Formato: parrafo simple sin tablas (patron transaccional de la Seccion 3.4).
Una pagina. Document() directo, sin template ADASA.

Estado: BORRADOR, pendiente de revision y envio.
"""
import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-08-05_Weekly-Reporting-Gap.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph(text)
    aplicar_arial(p, size)
    return p


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for text, bold in segs:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return p


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "August 5, 2026"),
        ("From:", CONTACTO + " - Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi - BW Water"),
        ("CC:", "Andrew Sia, Magdier Arias, Victor Gutierrez, Jorge Guevara, "
                "Ronald Pellejero"),
        ("Subject:", "25007 Taltal - Weekly reporting package: DDSR and full "
                     "schedule print"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_para(doc, "Two items on the weekly reporting package, both of them about what "
                  "reaches us rather than about the work itself.")
    blank(doc)

    add_segments(doc, [
        ("The Document and Drawing Status Report is missing for the second week in a "
         "row. ", True),
        ("Neither the 27 July nor the 3 August package included it. The last DDSR we "
         "hold is dated 20 July, so the documentation progress figure we are working "
         "with is now sixteen days old and we cannot reconcile it against our own "
         "register.", False),
    ])
    blank(doc)

    add_segments(doc, [
        ("The 4 August schedule came as a collapsed four-page print. ", True),
        ("The 14 July and 27 July issues ran to eleven and twelve pages and carried the "
         "engineering detail rows; this one summarises to four and omits them. It also "
         "arrived separately, outside the weekly package. The critical milestones are "
         "there and we have used them, but the level of detail no longer allows us to "
         "follow the critical path task by task.", False),
    ])
    blank(doc)

    add_para(doc, "Please restore both in the package of Monday 10 August: the DDSR, and "
                  "the schedule at the same level of detail as the previous issues and "
                  "within the package itself.")
    blank(doc)

    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in ["Infrastructure Engineering Lead - ADASA - Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)
    docx_metadata.apply_core_properties(
        doc, title="Taltal - Weekly reporting package",
        author="Luis Rivera Gonzalez", last_modified_by="Luis Rivera Gonzalez",
        subject="DDSR missing for a second week and collapsed schedule print",
        comments="Correo de reclamo de reporteria semanal a BW Water.",
        language="en-US")
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas de Antofagasta S.A.")
    print("Correo generado: " + OUTPUT_FILE)


if __name__ == "__main__":
    crear_correo()
