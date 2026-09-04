#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo conjunto ADASA -> BW Water + Bureau Veritas por las inspecciones 3 y 4,
del 13 y 14 de agosto de 2026 en Penang.

Reply al hilo del Request to witness inspection 003. Cadena de las jornadas,
SEPARADA de la de transmittals (ahi va el TM N32).

DOS INSPECCIONES DISTINTAS, no una visita de dos dias:
  - Inspeccion 3, jueves 13: ensayo de presion de la caneria de alta.
  - Inspeccion 4, viernes 14: inspeccion de preparacion de pintura.
El AQ-QAM-F027 Inspection Request (003) las cubre a las dos en una sola casilla
de fecha y una sola casilla de resultado, y no existe un Request 004. Por eso el
correo pide un formulario y un resultado por inspeccion.

EL PUNTO SUSTANTIVO es el procedimiento de pintura. El paquete que ADASA publico
a Bureau Veritas el 21-Jul tiene la Rev B, cuyo formulario declara perfil de
anclaje 40-75 um en la celda de criterio y 50-80 en la fila de aceptacion del
mismo formulario. La Rev 0 vigente lee 50-80 en los dos lugares, que es lo que
exige la Painting Specification Rev C aprobada. Verificado abriendo los dos PDF.

SIETE documentos del paquete estaban en revision superada. Se repusieron dentro
de PAQUETE_INSPECCION_BV -la carpeta NO se mueve ni se renombra, tiene enlace
publicado vivo- con los superados a subcarpetas _superseded/. Al correo se
adjuntan solo los dos que gobiernan estos dos dias.

Excepcion de correo conjunto: se nombra a BW Water aunque Bureau Veritas este
entre los destinatarios.

Cierra BV-09 del registro de compromisos (reemitir al inspector la Rev 0 del PMI).

Ingles. Document() directo, sin template ADASA. Estado: BORRADOR.
"""
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-08-12_BV-Inspections-3-and-4.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"
# Enlace del PAQUETE_INSPECCION_BV publicado el 21-Jul-2026. NO cambia: los
# archivos se reemplazaron dentro de la misma carpeta.
PACKAGE_LINK = ("https://lrg.synology.me:6501/d/s/18xc6ezTnrRYrdGebvNUYAYku3fZkKrM/"
                "mnJb6gJOkpddZgZ2vb2BZd870AMiJ0oo-mLwg5C36XQ0")

SUPERSEDED = [
    ("Painting Procedure", "P22-BA-09-000-011", "Rev B", "Rev 0"),
    ("HP and LP Pressure Test Procedure", "P22-BA-09-000-010", "Rev D", "Rev 0"),
    ("PMI Procedure", "P22-BA-09-000-006", "Rev A", "Rev 0"),
    ("Visual Procedure", "P22-BA-09-000-008", "Rev A", "Rev 0"),
    ("RO Vessel Hydrostatic Test Procedure", "P22-BA-09-000-009", "Rev C", "Rev 0"),
    ("GA of SWRO System Skid", "P22-DWG-09-005-008", "Rev A", "Rev B"),
    ("Plant Control Philosophy", "P22-BT-09-009-001", "Rev E", "Rev 0"),
]


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph(text)
    aplicar_arial(p, size)
    return p


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for text, bold in segs:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return p


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r0 = p.add_run("•  " + lead)
    r0.bold = True
    r0.font.name = "Arial"
    r0.font.size = Pt(size)
    r1 = p.add_run(rest)
    r1.font.name = "Arial"
    r1.font.size = Pt(size)
    return p


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), "Arial")
    rf.set(qn("w:hAnsi"), "Arial")
    rpr.append(rf)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0563C1")
    rpr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    h.append(run)
    paragraph._p.append(h)


def add_tabla(doc, filas):
    """Tabla Table Grid: documento, codigo, revision en el paquete, vigente."""
    tabla = doc.add_table(rows=1, cols=4)
    tabla.style = "Table Grid"
    encabezados = ["Document", "Code", "As issued 21 July", "Current"]
    for i, texto in enumerate(encabezados):
        celda = tabla.rows[0].cells[i]
        celda.text = ""
        run = celda.paragraphs[0].add_run(texto)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(10)
    for doc_name, codigo, en_paquete, vigente in filas:
        celdas = tabla.add_row().cells
        for i, texto in enumerate((doc_name, codigo, en_paquete, vigente)):
            celdas[i].text = ""
            run = celdas[i].paragraphs[0].add_run(texto)
            run.font.name = "Arial"
            run.font.size = Pt(10)
            if i == 3:
                run.bold = True
    return tabla


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "August 12, 2026"),
        ("From:", CONTACTO + " - Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi, Magdier Arias - BW Water; "
                "Carlo Montecinos - Bureau Veritas"),
        ("CC:", "Mohd Adnin Bin Zulkaflee, Muhammad Fadhil Bin Abdul Wahid - BW Water; "
                "Ahmad Hazwan, Wan Mohd Adli W Yahya, Emylia Rosli - Bureau Veritas; "
                "Victor Gutierrez, Jorge Guevara, Ronald Pellejero - ADASA"),
        ("Subject:", "RE: 25007 TALTAL - Witness inspections 3 and 4, 13 and 14 August - "
                     "governing revisions"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear all,")
    blank(doc)

    add_segments(doc, [
        ("Two witness inspections take place this week at BW Water Penang, both from "
         "9.00 am to 5.00 pm, under Inspection Request Form 003: ", False),
        ("inspection 3 on Thursday 13 August", True),
        (", the high-pressure piping pressure test, and ", False),
        ("inspection 4 on Friday 14 August", True),
        (", the painting preparation inspection.", False),
    ])

    blank(doc)
    add_segments(doc, [
        ("Governing revisions. ", True),
        ("The inspection package as issued on 21 July carried a superseded revision of both "
         "governing procedures, and in one of them the difference is not formal.", False),
    ])
    blank(doc)

    add_bullet_lead(
        doc, "Inspection 3, Thursday 13 August - high-pressure piping pressure test. ",
        "The current HP and LP Pressure Test Procedure (P22-BA-09-000-010) is Rev 0; the "
        "21 July package carried Rev D. The test pressures are those of the approved Line "
        "List, and "
        "the high-pressure test is a hold point under row 5.2 of the Inspection and Test "
        "Plan.")
    add_bullet_lead(
        doc, "Inspection 4, Friday 14 August - painting preparation. ",
        "The current Painting Procedure (P22-BA-09-000-011) is Rev 0; the 21 July package "
        "carried "
        "Rev B, whose inspection form states an anchor profile criterion of 40 to 75 "
        "micrometres in one cell and 50 to 80 in the acceptance row of that same form. Rev 0 "
        "reads 50 to 80 in both, which is what the approved Painting Specification requires. "
        "The profile to be accepted on 14 August is 50 to 80 micrometres.")
    blank(doc)

    add_segments(doc, [
        ("Finish colour. ", True),
        ("For the finish coat the colour that governs is RAL 5012 Luminous Blue, set by the "
         "Painting Specification (P22-ET-09-006-002) Rev C for the support frame inside the "
         "container. The Colour row of the inspection form in Rev 0 of the procedure states "
         "a different value, and its correction is in hand.", False),
    ])
    blank(doc)

    add_segments(doc, [
        ("Certificate of the pressure test. ", True),
        ("Row 5.2 of the Inspection and Test Plan requires a Pressure Test report with a "
         "pressure against time graphic as the certificate of that hold point. Clause 5.8.1 "
         "of the procedure names a Pressure Test Report that is not part of the submittal "
         "and that no form in the procedure produces. Please identify that form by number "
         "and revision and have it available on 13 August.", False),
    ])
    blank(doc)

    add_para(doc, "Revisions replaced in the inspection package:")
    add_tabla(doc, SUPERSEDED)
    blank(doc)

    add_para(doc, "Attached: P22-BA-09-000-010_0 HP and LP Pressure Test Procedure.pdf and "
                  "P22-BA-09-000-011_0 Painting Procedure.pdf")
    p = add_segments(doc, [
        ("All seven have been replaced in the package, which keeps the address issued on "
         "21 July: ",
         False)])
    add_hyperlink(p, PACKAGE_LINK, PACKAGE_LINK)
    blank(doc)

    add_para(doc, "Three requests:")
    add_bullet_lead(doc, "Confirmation of attendance ", "for both days.")
    add_bullet_lead(
        doc, "A separate request form and result for each inspection. ",
        "Form 003 carries both days in a single date field and a single result box, and the "
        "two days have different scopes, so one signature cannot record both.")
    add_bullet_lead(
        doc, "The signed form and the Bureau Veritas report for each day, separately, ",
        "in the format issued for the visit of 28 July (BVM-IR001-28072026).")
    blank(doc)

    add_para(doc, "Best regards,")
    blank(doc)

    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in [
        "Infrastructure Engineering Lead - ADASA - Aguas de Antofagasta S.A.",
        "lrivera@aguasantofagasta.cl",
    ]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)

    docx_metadata.apply_core_properties(
        doc,
        title="Taltal - Witness inspections 3 and 4, 13 and 14 August - governing revisions",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="Taltal SWRO - Third-party shop inspection, governing document revisions",
        comments="Aguas de Antofagasta S.A.",
        language="en-US",
    )
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas Antofagasta")
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
