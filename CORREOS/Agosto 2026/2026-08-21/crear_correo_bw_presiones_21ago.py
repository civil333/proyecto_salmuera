#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo ADASA -> BW Water sobre los ensayos de presion del circuito CIP.
Reply-to al mensaje de Adnin del 20-Ago 21:13 hora de Chile, dentro del hilo del
Request to witness inspection 003/004. Bureau Veritas en copia (decision del usuario).

VERSION EJECUTIVA (3a pasada, 21-Ago). La 2a llego a 758 palabras. El usuario pidio
"mas ejecutivo": la tabla carga las cifras, cada peticion ocupa una linea, y se
eliminaron los parrafos que repetian lo que la tabla ya dice.

EL EJE: el taller esta ensayando el CIRCUITO CIP COMPLETO como sistema de baja
presion. TRES de las CUATRO lineas CIP de super duplex, en dos jornadas consecutivas,
todas a 7,5 barG y todas con 5,0 barG de presion de diseno declarada en el registro.

  `CP-SSD-DN100-09-014`  80 / 120 barG  — ensayada a 7,5 el 20-Ago
  `CP-SSD-DN80-09-015`   90 / 135 barG  — ensayada a 7,5 el 21-Ago
  `CP-SSD-DN80-09-044`   80 / 120 barG  — ensayada a 7,5 el 21-Ago
  `CP-SSD-DN65-09-045`   90 / 135 barG  — PENDIENTE, y por eso la retencion sobre ella
                                          encabeza las peticiones

EL SPOOL 1 DEL 20-AGO QUEDA IDENTIFICADO y lo identifico el tercero inspector, que
corrigio el prefijo DA-SSD del formulario del fabricante. La peticion del correo del
20-Ago se retira.

LOS MANOMETROS NO SON EL PROBLEMA. El 20-Ago habia cuatro en la mesa, dos de 0 a 160
bar, usados esa misma manana para el ensayo a 90 barG. La pregunta es por que se
aplico el criterio de baja presion.

EL ARGUMENTO DE ADNIN CONFIRMA UN HALLAZGO ABIERTO. "This spool that has two flange
class, 900# and 150#" es el quiebre de especificacion no declarado del Transmittal
N30, sobre `CP-SSD-DN80-09-044` y `CP-SSD-DN65-09-045`, abierto en PRG-22.

FUERA DEL CORREO, a proposito: el veredicto del Transmittal N35, el plazo de entrega
vencido el 03-Ago con la multa de la Clausula 43.1 letra b, y el reclamo de desempeno
contra el tercero inspector, que va en su propia cadena.

Ingles. Document() directo, sin template ADASA. Estado: BORRADOR.
"""
import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-08-21_Pressure-Tests-21-Aug.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11, space=6):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(space)
    aplicar_arial(p, size)
    return p


def add_segments(doc, segs, size=11, space=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    for text, bold in segs:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return p


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(3)
    r0 = p.add_run("•  " + lead)
    r0.bold = True
    r0.font.name = "Arial"
    r0.font.size = Pt(size)
    r1 = p.add_run(rest)
    r1.font.name = "Arial"
    r1.font.size = Pt(size)
    return p


def espaciador(doc):
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


# Las CUATRO lineas CIP de super duplex de la Line List P22-LI-09-009-003 Rev 0,
# adjunta al propio procedimiento de ensayo del proveedor. La cuarta va en la tabla
# porque es la que convierte el listado en advertencia. La prosa NO repite estas
# cifras: ese es el recorte principal de la pasada ejecutiva.
CIP_LINES = [
    ("CP-SSD-DN100-09-014", "CIP Feed to 1st Stage RO", "80", "120",
     "5.0", "7.5 on 20 Aug"),
    ("CP-SSD-DN80-09-015", "CIP Feed to 2nd Stage RO", "90", "135",
     "5.0", "7.5 on 21 Aug"),
    ("CP-SSD-DN80-09-044", "1st Stage CIP Reject Out", "80", "120",
     "5.0", "7.5 on 21 Aug"),
    ("CP-SSD-DN65-09-045", "2nd Stage CIP Reject Out", "90", "135",
     "-", "not yet tested"),
]


def add_table_tested(doc, size=9):
    headers = (
        "Line",
        "Description",
        "Design barG\n(Line List)",
        "Hydrotest barG\n(Line List)",
        "Design barG\nas recorded",
        "Applied barG\nand date",
    )
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(size)
    for row in CIP_LINES:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(value)
            run.font.name = "Arial"
            run.font.size = Pt(size)
    return table


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "August 21, 2026"),
        ("From:", CONTACTO + " - Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi, Stephane Gehant, Magdier Arias, "
                "Mohd Adnin Bin Zulkaflee, Muhammad Fadhil Bin Abdul Wahid - BW Water"),
        ("CC:", "Ahmad Hazwan, Wan Mohd Adli W Yahya, Emylia Rosli, "
                "Carlo Montecinos - Bureau Veritas; Victor Gutierrez - ADASA"),
        ("Subject:", "RE: 25007 TALTAL - Request to witness inspection 003/004 - "
                     "super duplex test pressures"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    espaciador(doc)

    add_para(doc, "Dear Eduardo, Stephane, Adnin,")

    # LEAD. El encuadre sistemico en una frase. Las cifras van en la tabla.
    add_segments(doc, [
        ("The CIP circuit is being pressure tested as a low pressure system.", True),
        (" Three of its four super duplex lines have now been tested at 7.5 barG over "
         "two consecutive days, and all three records state a design pressure of 5.0 "
         "barG. The approved Line List (P22-LI-09-009-003 Rev 0) gives that figure to "
         "none of them, and the fourth line has not been tested yet.", False),
    ])

    add_table_tested(doc)
    espaciador(doc)

    # EL SPOOL IDENTIFICADO Y LA RETENCION VIOLADA.
    add_segments(doc, [
        ("The spool we asked you to identify on 20 August is identified, and the hold "
         "was already in writing when the next two were tested.", True),
        (" Your record gave it as DA-SSD-DN100-09-014, a tag absent from the Line List; "
         "the third party report corrects it to CP-SSD-DN100-09-014, so we withdraw that "
         "request. Row 5.2 of the Inspection and Test Plan (P22-BA-09-000-004 Rev 0) is "
         "a hold point under ADASA control, so none of the three tests discharges it.",
         False),
    ])

    # EL QUIEBRE DEL N30.
    add_segments(doc, [
        ("The two flange classes point at an observation we already raised.", True),
        (" A spool that mixes Class 900 and Class 150 components inside a line the Line "
         "List classifies as super duplex is the undeclared specification break of "
         "Transmittal N30, which named CP-SSD-DN80-09-044 and CP-SSD-DN65-09-045 by tag "
         "and is still open. While a line remains one line on that list, it is tested at "
         "the pressure the list assigns it, and no part of the super duplex section is "
         "qualified by a test at 7.5 barG.", False),
    ])

    add_segments(doc, [
        ("What we require by Monday 24 August:", True),
    ], space=3)
    add_bullet_lead(
        doc, "No test of CP-SSD-DN65-09-045",
        ", nor of any other super duplex line, until the confirmation below reaches us.")
    add_bullet_lead(
        doc, "The written confirmation requested on 20 August",
        ", due today: the test pressure of each of the eleven super duplex lines against "
        "the approved Line List (P22-LI-09-009-003 Rev 0).")
    add_bullet_lead(
        doc, "The design pressure the workshop applies to the CIP circuit",
        ", in writing, and the document it comes from.")
    add_bullet_lead(
        doc, "The list of every super duplex spool tested to date",
        ", with tag, pressure applied, date and gauge used.")
    add_bullet_lead(
        doc, "The answer to the specification break of Transmittal N30",
        ": split the line number at the joint, issue the addition to the Line List, and "
        "state how the PVC section is protected with the motorised valve closed.")
    add_bullet_lead(
        doc, "Repeat the three tests",
        " at the listed pressure, with written notice and the inspector present.")
    add_bullet_lead(
        doc, "Correct the records of 20 and 21 August",
        ", including checklist item 5 on gauge range.")
    add_bullet_lead(
        doc, "Explain the choice of gauges",
        ". Four were on the bench on 20 August, two of them 0 to 160 bar and used that "
        "morning for the 90 barG test; the CIP lines were tested with the 0 to 16 bar "
        "pair.")
    espaciador(doc)

    # RESERVA. ADASA nombrada aqui, que es la declaracion formal.
    add_segments(doc, [
        ("The hold stands.", True),
        (" Any super duplex test run before that written confirmation reaches us is "
         "treated by ADASA as not executed, whatever witness record accompanies it.",
         False),
    ])

    add_para(doc, "Best regards,", space=0)
    espaciador(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in [
        "Infrastructure Engineering Lead - ADASA - Aguas de Antofagasta S.A.",
        "lrivera@aguasantofagasta.cl",
    ]:
        add_para(doc, line, space=0)

    fijar_idioma(doc, LANG)

    docx_metadata.apply_core_properties(
        doc,
        title="Taltal - CIP circuit tested as a low pressure system",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="Taltal SWRO - Hydrostatic testing of super duplex lines against the "
                "approved Line List",
        comments="Aguas de Antofagasta S.A.",
        language="en-US",
    )
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas Antofagasta")
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
