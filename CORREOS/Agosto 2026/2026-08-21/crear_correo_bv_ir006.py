#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo ADASA -> Bureau Veritas sobre los informes BVM-IR005-20082026 y
BVM-IR006-21082026. Reply al correo de Carlo Montecinos del 21-Ago 07:55 hora de
Chile. Cadena PROPIA de Bureau Veritas, separada de la de las jornadas de inspeccion.

VERSION EJECUTIVA (3a pasada, 21-Ago). La 2a llego a 1.085 palabras. El usuario pidio
"mas ejecutivo": se recorto a la mitad SIN perder ninguno de los cinco ejes, aplicando
la disciplina de feedback_correos_remision_ejecutivos — lead con el veredicto, la
tabla carga las cifras que la prosa no repite, un pendiente por linea, cero parrafos
que dupliquen, space_after de 6pt en vez de parrafos vacios.

CORRECCION QUE ARRASTRA DE LA 2a PASADA: el borrador inicial afirmaba que el informe
del 20-Ago no habia llegado. ES FALSO. El BVM-IR005 lo remitio Emylia Rosli el 20-Ago
23:35 hora de Chile con Luis Rivera en copia, por la cadena del Request 004 — tercera
cadena del frente. Llego DENTRO del plazo de la seccion 3 de la oferta, y el correo lo
acusa expresamente para acotar el reclamo a su contenido.

LOS CINCO EJES, que sobreviven al recorte:

  1. EL ROL. La NT P22-NT-09-000-002-0 designa al Tercero Inspector REPRESENTANTE DE
     ADASA para los Puntos de Detencion y la emision de registros. La oferta 600049
     Rev.3 seccion 1: velar por los intereses del cliente revisando planos y
     procedimientos.
  2. ADASA NO TIENE COMO COMPROBARLO. 17.000 km. Sobre la identidad de lo ensayado, el
     inspector presente es el UNICO control.
  3. LAS FOTOGRAFIAS NO LO SUPLEN. Ninguna imagen de la seccion I muestra marca del
     carrete; la del IR006 rotula UNA foto con DOS numeros de linea. Constatacion de lo
     que NO aparece, verificada por render.
  4. ADASA ENTREGO LAS HERRAMIENTAS. Guia del Paquete Rev 1 seccion 3.4: P&ID Rev D
     como "Trazado maestro - soldadura, hidrostatica, montaje" y Line List Rev 0 para
     las "Visitas 2-3". Ninguno en la seccion B de los dos informes.
  5. EL CASO QUE LO PRUEBA. El inspector CORRIGIO el prefijo del TAG y aun con el
     numero correcto delante consigno 5,0 barG sobre una linea de 80.

TONO: firme, sin imputacion personal y sin adjetivos sobre el desempeno. Cierra en el
trabajo conjunto. REGLA DURA: no se nombra a BW Water; se dice "el fabricante".

ENVIADO el viernes 21-Ago-2026 a las 12:15:49 hora de Chile. Este generador quedo
alineado con el texto REALMENTE ENVIADO: al emitir, el usuario quito "sobre la
identidad de lo que se ensaya" del segundo parrafo y cambio "En ADASA mantenemos el
paquete al dia y respondemos" por "Estamos disponibles para responder". Respaldo en
"Re: 25007 TALTAL -  Request to witness inspection 006.pdf".

Espanol de Chile. Document() directo, sin template ADASA.
"""
import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-08-21_BV-Informe-IR006.docx")
CONTACTO = "Luis Rivera González"
LANG = "es-CL"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    """Sin esto Word subraya en rojo tildes y enes del cuerpo en espanol."""
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11, space=6):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(space)
    aplicar_arial(p, size)
    return p


def add_segments(doc, segs, size=11, space=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    for text, bold in segs:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return p


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(3)
    r0 = p.add_run("•  " + lead)
    r0.bold = True
    r0.font.name = "Arial"
    r0.font.size = Pt(size)
    r1 = p.add_run(rest)
    r1.font.name = "Arial"
    r1.font.size = Pt(size)
    return p


def espaciador(doc):
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


# Las TRES lineas del circuito CIP ensayadas a 7,5 barG en las dos jornadas. La tabla
# carga las cifras y por eso la prosa NO las repite: es el recorte principal de esta
# pasada ejecutiva.
LINEAS = [
    ("CP-SSD-DN100-09-014", "BVM-IR005", "20-Ago", "80", "120", "5,0", "7,5"),
    ("CP-SSD-DN80-09-015", "BVM-IR006", "21-Ago", "90", "135", "5,0", "7,5"),
    ("CP-SSD-DN80-09-044", "BVM-IR006", "21-Ago", "80", "120", "5,0", "7,5"),
]


def add_tabla_lineas(doc, size=9):
    headers = (
        "Línea",
        "Informe",
        "Jornada",
        "Diseño barG\nLine List",
        "Hidrostática barG\nLine List",
        "Diseño barG\nsegún informe",
        "Ensayo barG\naplicado",
    )
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(size)
    for row in LINEAS:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(value)
            run.font.name = "Arial"
            run.font.size = Pt(size)
    return table


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Fecha:", "21 de agosto de 2026"),
        ("De:", CONTACTO + " - Líder de Ingeniería de Infraestructura (ADASA)"),
        ("Para:", "Carlo Alberto Montecinos Zúñiga - Bureau Veritas Chile"),
        ("CC:", "Ahmad Hazwan, Wan Mohd Adli W Yahya, Emylia Rosli - Bureau Veritas; "
                "Víctor Gutiérrez - ADASA"),
        ("Asunto:", "RE: 25007 TALTAL - Request to witness inspection 006 - "
                    "informes BVM-IR005 y BVM-IR006"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    espaciador(doc)

    add_para(doc, "Estimado Carlo:")

    # LEAD. Veredicto en dos oraciones; la tabla lleva las cifras.
    add_segments(doc, [
        ("Los informes BVM-IR005 y BVM-IR006 dan por conformes tres ensayos de super "
         "dúplex ejecutados a 7,5 barG; el del 20 de agosto, además, sin comentario "
         "alguno.", True),
        (" Las tres líneas son del circuito CIP, y la Line List aprobada en Código 1 "
         "(P22-LI-09-009-003 Rev 0) les asigna entre 120 y 135 barG de hidrostática.",
         False),
    ])

    add_tabla_lineas(doc)
    espaciador(doc)

    # EL ROL Y EL UNICO CONTROL.
    add_segments(doc, [
        ("El módulo se fabrica a más de 17.000 kilómetros, el inspector presente es el "
         "único control con que contamos.", True),
        (" La Nota Técnica que lo designó (P22-NT-09-000-002-0) lo hace representante de "
         "ADASA, y la oferta lo obliga en su sección 1 a velar por los intereses del "
         "cliente revisando planos y procedimientos. Un número de línea que no "
         "corresponda a la cañería ensayada no encuentra segundo filtro.", False),
    ])

    # LAS HERRAMIENTAS ENTREGADAS Y EL REGISTRO FOTOGRAFICO.
    add_segments(doc, [
        ("Los dos documentos con los que ese contraste se hace se los entregamos hace un "
         "mes, asignados de manera expresa a la hidrostática", True),
        (": el P&ID (P22-DWG-09-009-002 Rev D) y la Line List, en la sección 3.4 de la "
         "Guía del Paquete. Ninguno figura en la sección B de los informes; el registro "
         "fotográfico tampoco permite el contraste posterior, dado que ninguna imagen de "
         "la sección I muestra la marca del carrete y la del BVM-IR006 rotula una sola "
         "fotografía con dos números de línea.", False),
    ])

    # EL CASO QUE LO PRUEBA. Sin la apertura "Un caso lo muestra completo" (U-02B).
    add_segments(doc, [
        ("El formulario del fabricante registró el carrete del 20 de agosto como "
         "DA-SSD-DN100-09-014", True),
        (", número que no existe en la Line List; el BVM-IR005 lo corrige a "
         "CP-SSD-DN100-09-014. Con dicho número delante bastaba leer una fila de la "
         "tabla que ya le habíamos entregado.", False),
    ])

    # MANOMETROS Y FLASH REPORT.
    add_segments(doc, [
        ("Ese día había cuatro manómetros sobre la mesa, dos de ellos de 0 a 160 bar",
         True),
        (", usados esa misma mañana para el ensayo a 90 barG y con sus certificados "
         "revisados (punto 6 de la sección E1). El informe llegó en plazo; lo que no "
         "hubo, en ninguna de las dos jornadas, fue el Flash Report que la sección 3.2 "
         "exige ante una desviación.", False),
    ])

    add_segments(doc, [
        ("Lo que solicitamos, con plazo al lunes 24 de agosto:", True),
    ], space=3)
    add_bullet_lead(
        doc, "Contrastar cada ensayo contra el P&ID y la Line List antes de firmar",
        ", consignando la presión de diseño junto a la aplicada.")
    add_bullet_lead(
        doc, "Fotografiar la marca de identificación del carrete",
        " en cada registro.")
    add_bullet_lead(
        doc, "Reemitir el BVM-IR005 y el BVM-IR006",
        " con la presión de diseño de la Line List.")
    add_bullet_lead(
        doc, "Levantar una No Conformidad",
        " que cubra los tres ensayos.")
    add_bullet_lead(
        doc, "Incorporar el P&ID y la Line List",
        " a la documentación de referencia de todo informe de presión.")
    add_bullet_lead(
        doc, "Instruir al inspector",
        " que la fila 5.2 del Plan de Inspección y Ensayos (P22-BA-09-000-004 Rev 0) es "
        "Punto de Detención y que la retención del 20 de agosto sigue vigente.")
    espaciador(doc)

    # RESERVA.
    add_segments(doc, [
        ("ADASA reserva su posición sobre la validez del atestiguamiento", True),
        (" y tiene los tres ensayos por no ejecutados; un testigo sin contraste contra "
         "el P&ID y la Line List no acredita el Punto de Detención de la fila 5.2.",
         False),
    ])

    # CIERRE EN EL TRABAJO CONJUNTO. Sin "Este correo apunta a..." (U-02B).
    add_segments(doc, [
        ("La presencia del inspector en terreno no está en cuestión; lo que pedimos "
         "corregir es el método de contraste.", True),
        (" Estamos disponibles para responder consultas de ingeniería el mismo día, de "
         "modo que la duda pueda levantarse en terreno. Con diez líneas de super "
         "dúplex por ensayar, es lo que necesitamos.", False),
    ])

    add_para(doc, "Quedamos a la espera de su respuesta y a disposición para coordinar.")
    add_para(doc, "Saludos cordiales,", space=0)
    espaciador(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in [
        "Líder de Ingeniería de Infraestructura - Depto. Proyectos Desalación",
        "Aguas de Antofagasta S.A.",
        "lrivera@aguasantofagasta.cl",
    ]:
        add_para(doc, line, space=0)

    fijar_idioma(doc, LANG)

    docx_metadata.apply_core_properties(
        doc,
        title="Taltal - Informes BVM-IR005 y BVM-IR006 - presiones de ensayo y alcance "
              "del Tercero Inspector",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="Taltal SWRO - Atestiguamiento de ensayos hidrostáticos contra la "
                "ingeniería aprobada",
        comments="Aguas de Antofagasta S.A.",
        language="es-CL",
    )
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas Antofagasta")
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
