#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo ADASA -> BW Water (Eduardo Yamauchi + Andrea Frezzi): confirmacion del
Incoterm del paquete de repuestos de 2 anos.

Reply-To al thread propio "Formal Re-validation of Spare Parts Quotation C4300",
que es donde llego la propuesta formal 25007-PL-0002 rev.0 el 28-Jul-2026.
Cadena SEPARADA de los transmittals y del thread de packing/EXW del modulo.

Origen: en la reunion de progreso del 11-Ago-2026 BW Water planteo que la
cotizacion declara varios puntos de entrega (Espana, China, EE.UU.) en vez de un
embarque unico desde Penang, y dejo como accion reemitirla consolidada.

Eje unico del correo: la propuesta YA declara el Incoterm en su seccion
Commercial Conditions ("EXW Penang, Malaysia"), y la columna Estimated Transit
Time del Anexo A (2 semanas desde el extranjero, 0 para lo ya cotizado DDP
Penang) solo tiene sentido como transporte hacia Penang. Se pide confirmacion
POR ESCRITO, no una Rev.1: reabrir el documento habilita a cargar el flete a
Penang sobre un total que ADASA ya formalizo en la Carta de Aumento de Monto del
04-Ago-2026 (USD 56.990,50 = repuestos 51.224,50 + adicional CIP 5.766,00).

Fuera de este correo, por decision de metodo: las once lineas mal correspondidas,
las cuatro alzas sin justificar y la ausencia de repuestos DN15. Van en correo
aparte antes de que la validez venza el 28-Ago-2026. Mezclarlas aqui permitiria
refundir todo en una reemision re-cotizada.

Plazo: se refuta con las cifras de la propia propuesta (10 semanas de lead time
maximo mas 2 de transito) que consolidar en Penang comprometa el plazo
contractual. No se ofrece ni se insinua extension.

Version ejecutiva (~200 palabras): la evidencia la cargan dos recortes CRUDOS
del propio PDF de BW Water, generados por generar_figuras_incoterm.py. Sin
resaltados anadidos: el extracto tiene que ser intacto. Fuera de esta version
quedan el parrafo de alcance logistico (embarque, internacion y aduana unicas) y
la mencion de la solicitud de aumento de monto del 04-Ago, que competian con el
argumento principal.

Sustento interno: PROGRAMA y CONTRATO/_ANALISIS_COMERCIAL_04AGO.md

Ingles (BW Water). Document() directo, sin template ADASA. Metadatos limpios.
Estado: BORRADOR.
"""

import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402  (metadatos limpios, global sec. 2.3)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(
    SCRIPT_DIR, "2026-08-11_Spare-Parts-Incoterm-Confirmation.docx"
)
FIG_DIR = os.path.join(SCRIPT_DIR, "figuras")
FIG_CONDICIONES = os.path.join(FIG_DIR, "fig1_commercial_conditions.png")
FIG_ANEXO = os.path.join(FIG_DIR, "fig2_annexA_incoterms.png")
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang_in_rPr(rPr, lang):
    lang_el = rPr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rPr.append(lang_el)
    lang_el.set(qn("w:val"), lang)


def fijar_idioma_documento(doc, lang=LANG):
    try:
        rPr = doc.styles["Normal"].element.get_or_add_rPr()
        _set_lang_in_rPr(rPr, lang)
    except KeyError:
        pass
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            _set_lang_in_rPr(run._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    para = doc.add_paragraph(text)
    aplicar_arial(para, size)
    return para


def add_punto(doc, lead, rest, size=11):
    """Punto de confirmacion con letra en negrita, sangria izquierda."""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.3)
    para.add_run(lead).bold = True
    para.add_run(rest)
    aplicar_arial(para, size)
    return para


def add_figura(doc, ruta, ancho_pulgadas, pie):
    """Recorte del PDF de BW Water, centrado, con su pie en Arial 9 cursiva."""
    if not os.path.exists(ruta):
        raise SystemExit(
            f"Falta {os.path.basename(ruta)}. Correr antes generar_figuras_incoterm.py"
        )
    para_img = doc.add_paragraph()
    para_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para_img.add_run().add_picture(ruta, width=Inches(ancho_pulgadas))

    para_pie = doc.add_paragraph()
    para_pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para_pie.add_run(pie)
    run.italic = True
    aplicar_arial(para_pie, 9)
    return para_pie


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # ---- HEADER -------------------------------------------------------------
    fields = [
        ("Date:", "August 11, 2026"),
        ("From:", f"{CONTACTO} — Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi, Andrea Frezzi — BW Water"),
        ("CC:",
         "Jorge Guevara, Victor Gutierrez — ADASA; Stephane Gehant — BW Water"),
        ("Subject:",
         "RE: Formal Re-validation of Spare Parts Quotation C4300"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    blank(doc)

    # ---- CUERPO -------------------------------------------------------------
    add_para(doc, "Dear Eduardo and Andrea,")
    blank(doc)

    add_para(
        doc,
        "Eduardo raised the delivery point of the two-year spare parts package "
        "in today's weekly coordination meeting, on the understanding that the "
        "quotation shows several delivery points instead of a single "
        "consolidated shipment from Penang. Thank you also for the formal "
        "quotation issued on 28 July 2026.",
    )
    blank(doc)

    add_para(doc, "Your proposal states a single Incoterm for the package:")
    blank(doc)

    add_figura(
        doc,
        FIG_CONDICIONES,
        5.5,
        "Proposal 25007-PL-0002 rev.0, Commercial Conditions.",
    )
    blank(doc)

    add_para(
        doc,
        "In Annex A, the Incoterms column carries the supply terms of each "
        "sub-supplier. The column beside it allows two weeks of transit for "
        "the items coming from Spain, China and the United States, and zero "
        "for those already quoted DDP Penang, so every line converges on "
        "Penang:",
    )
    blank(doc)

    add_figura(
        doc,
        FIG_ANEXO,
        6.4,
        "Proposal 25007-PL-0002 rev.0, Annex A (extract).",
    )
    blank(doc)

    add_para(
        doc,
        "A reissue is not needed; please confirm in writing:",
    )
    blank(doc)

    add_punto(
        doc,
        "a) ",
        "EXW Penang, Malaysia applies to the complete package.",
    )

    add_punto(
        doc,
        "b) ",
        "The total remains USD 51,224.50, consolidation and transport to "
        "Penang included.",
    )

    add_punto(
        doc,
        "c) ",
        "A single readiness date at Penang, in calendar weeks from the "
        "Purchase Order.",
    )
    blank(doc)

    add_para(
        doc,
        "We are ready to place the purchase order and the quotation is valid "
        "until 28 August 2026. Ten weeks of manufacturing plus two of transit "
        "leave the package at Penang well inside the 510 calendar day "
        "contractual period.",
    )
    blank(doc)

    add_para(
        doc,
        "Your reply by Thursday, 13 August 2026 would allow the order to be "
        "issued within the current validity.",
    )
    blank(doc)

    # ---- CIERRE -------------------------------------------------------------
    add_para(doc, "Best regards,")
    blank(doc)

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in [
        "Infrastructure Engineering Lead",
        "Desalination Projects Department",
        "ADASA — Aguas de Antofagasta S.A.",
        "lrivera@aguasantofagasta.cl",
    ]:
        add_para(doc, line)

    fijar_idioma_documento(doc, LANG)

    # ---- metadatos limpios (global sec. 2.3) --------------------------------
    docx_metadata.apply_core_properties(
        doc,
        title="Taltal - Two-year spare parts package - Incoterm confirmation",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="Taltal SWRO - Recommended two-year spare parts quotation",
        comments="Aguas de Antofagasta S.A.",
        language="en-US",
    )
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas Antofagasta")
    print(f"Correo generado: {OUTPUT_FILE}")
    print(f"Idioma del documento: {LANG}")


if __name__ == "__main__":
    crear_correo()
