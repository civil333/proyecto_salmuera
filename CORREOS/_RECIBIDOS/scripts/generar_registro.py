#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Regenera _REGISTRO.md y _REGISTRO.xlsx leyendo el frontmatter de cada _correo.md.

Los _correo.md son la FUENTE UNICA. El .md de registro y el .xlsx son DERIVADOS y
no se editan a mano, igual que el registro de compromisos del proyecto.

Todo campo del Excel se escribe como texto: no hay number_format que traducir ni
formulas, de modo que el gate openpyxl_lint.py pasa por construccion. Las fechas
van como cadena ISO para que ordenen bien sin depender del locale.

Uso:
    python generar_registro.py
    python generar_registro.py --check     # no escribe; exit 1 si algo cambiaria
"""
import argparse
import re
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
MD = RAIZ / "_REGISTRO.md"
XLSX = RAIZ / "_REGISTRO.xlsx"

CAMPOS = ("mensaje_id", "fecha", "remitente", "para", "cc", "asunto",
          "contraparte", "tipo")


def leer_frontmatter(ruta: Path) -> dict:
    """Parser acotado: solo los escalares que necesita el registro, mas adjuntos."""
    txt = ruta.read_text(encoding="utf-8")
    if not txt.startswith("---"):
        return {}
    fin = txt.find("\n---", 3)
    bloque = txt[3:fin if fin > 0 else 4000]
    d = {}
    for campo in CAMPOS:
        m = re.search(rf"^{campo}:\s*(.*)$", bloque, re.M)
        if m:
            d[campo] = m.group(1).strip().strip('"').strip("'")
    d["n_adjuntos"] = len(re.findall(r"^\s+- archivo:", bloque, re.M))
    destinos = re.findall(r'^\s+destino:\s*"?([^"\n]+)"?\s*$', bloque, re.M)
    fuera = sorted({x.strip().rstrip("/") for x in destinos
                    if x.strip().rstrip("/") != "adjuntos"})
    d["destino"] = "; ".join(fuera) if fuera else ("adjuntos/" if d["n_adjuntos"] else "—")
    d["carpeta"] = ruta.parent.relative_to(RAIZ).as_posix()
    return d


def recolectar() -> list:
    filas = [leer_frontmatter(p) for p in sorted(RAIZ.rglob("_correo.md"))]
    filas = [f for f in filas if f.get("mensaje_id")]
    filas.sort(key=lambda f: f.get("fecha", ""), reverse=True)
    return filas


def render_md(filas: list) -> str:
    L = ["---",
         "titulo: Registro de correo entrante — BW Water y Bureau Veritas",
         "proyecto: salmuera-taltal",
         "estado: VIGENTE",
         "second_brain: skip",
         "---",
         "",
         "# Registro de correo entrante",
         "",
         "> **Índice durable, no eje cronológico.** La Bitácora del README es el único eje "
         "cronológico del proyecto; este registro es un índice tabular del mismo género que el "
         "Índice de Transmittales. Un correo capturado aquí **no** genera entrada de Bitácora "
         "automáticamente: se propone y la aprueba el usuario.",
         ">",
         "> **Derivado.** Se regenera con `scripts/generar_registro.py` leyendo el frontmatter de "
         "cada `_correo.md`. La fuente única son esos archivos; esta tabla y el `.xlsx` no se "
         "editan a mano.",
         "",
         "Convención, ruteo y procedimiento de captura: [`_LEEME.md`](_LEEME.md).",
         "",
         "## Correspondencia recibida",
         "",
         "| Fecha | Contraparte | Remitente | Asunto | Tipo | Adj. | Destino de adjuntos | Carpeta |",
         "|---|---|---|---|---|---|---|---|"]
    if not filas:
        L.append("")
        L.append("*Sin registros. La tabla se puebla al correr el primer barrido.*")
    else:
        for f in filas:
            rem = re.sub(r"\s*<.*?>", "", f.get("remitente", "")).strip()
            L.append("| {} | {} | {} | {} | {} | {} | {} | `{}` |".format(
                f.get("fecha", "")[:16].replace("T", " "),
                f.get("contraparte", ""), rem,
                f.get("asunto", "").replace("|", "\\|"),
                f.get("tipo", ""), f.get("n_adjuntos", 0),
                f.get("destino", "—"), f.get("carpeta", "")))
    L += ["", "## Recuento", "", "| Contraparte | Correos | Último |", "|---|---|---|"]
    total_ult = ""
    for cp, etiqueta in (("BW WATER", "BW Water"), ("BUREAU VERITAS", "Bureau Veritas")):
        sub = [f for f in filas if f.get("contraparte") == cp]
        ult = sub[0].get("fecha", "")[:10] if sub else "—"
        total_ult = max(total_ult, ult) if ult != "—" else total_ult
        L.append(f"| {etiqueta} | {len(sub)} | {ult} |")
    L.append(f"| **Total** | **{len(filas)}** | {total_ult or '—'} |")
    L.append("")
    return "\n".join(L)


def escribir_xlsx(filas: list) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Correo entrante"
    cols = ["#", "Fecha", "Contraparte", "Remitente", "Asunto", "Tipo",
            "N adj", "Destino adjuntos", "Carpeta", "Mensaje ID"]
    anchos = [5, 18, 16, 26, 60, 14, 7, 34, 46, 40]

    relleno = PatternFill("solid", fgColor="003366")
    negrita = Font(bold=True, color="FFFFFF")
    for i, (c, w) in enumerate(zip(cols, anchos), start=1):
        celda = ws.cell(row=1, column=i, value=c)
        celda.fill = relleno
        celda.font = negrita
        celda.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[celda.column_letter].width = w

    for n, f in enumerate(filas, start=1):
        rem = re.sub(r"\s*<.*?>", "", f.get("remitente", "")).strip()
        valores = [str(n), f.get("fecha", "")[:16].replace("T", " "),
                   f.get("contraparte", ""), rem, f.get("asunto", "")[:200],
                   f.get("tipo", ""), str(f.get("n_adjuntos", 0)),
                   f.get("destino", ""), f.get("carpeta", ""),
                   f.get("mensaje_id", "")[:120]]
        for i, v in enumerate(valores, start=1):
            # Ningun string puede empezar con '=' (openpyxl lo guardaria como formula)
            if isinstance(v, str) and v.startswith("="):
                v = "'" + v
            ws.cell(row=n + 1, column=i, value=v)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    wb.save(XLSX)


def main() -> int:
    ap = argparse.ArgumentParser(description="Regenera el registro de correo entrante")
    ap.add_argument("--check", action="store_true",
                    help="no escribe; exit 1 si el .md cambiaria")
    args = ap.parse_args()

    filas = recolectar()
    nuevo = render_md(filas)

    if args.check:
        actual = MD.read_text(encoding="utf-8") if MD.exists() else ""
        if actual.strip() != nuevo.strip():
            print("DESACTUALIZADO: _REGISTRO.md no refleja los _correo.md", file=sys.stderr)
            return 1
        print(f"Al dia. {len(filas)} correos registrados.")
        return 0

    MD.write_text(nuevo, encoding="utf-8")
    escribir_xlsx(filas)
    bw = sum(1 for f in filas if f.get("contraparte") == "BW WATER")
    bv = sum(1 for f in filas if f.get("contraparte") == "BUREAU VERITAS")
    print(f"Registro regenerado: {len(filas)} correos ({bw} BW Water, {bv} Bureau Veritas)")
    print(f"  {MD.name}")
    print(f"  {XLSX.name}")
    print("  Gate: python ~/.claude/skills/_shared/openpyxl_lint.py " + str(XLSX))
    return 0


if __name__ == "__main__":
    sys.exit(main())
