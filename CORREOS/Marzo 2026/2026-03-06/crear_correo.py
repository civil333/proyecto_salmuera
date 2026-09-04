#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para generar correo DOCX - March 9 Meeting — EXW August 3 vs. Engineering Status
Fecha: 06 de marzo de 2026
Version: 2.2 (tabla procurement corregida: HP Pump Code 2-AN formal + Interstage TC agregado)

Correo profesional sin template ADASA (correo directo, sin portada).
Cuestion central: contradiccion EXW agosto vs. Engineering +185 dias.
Tabla comparativa Schedule vs Project Records + tabla procurement en ruta critica.
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

CONTACTO = {
    "nombre": "Luis Rivera Gonzalez",
    "cargo": "Leader, Infrastructure Engineering",
    "empresa": "ADASA \u2014 Aguas de Antofagasta S.A.",
    "proyecto": "BAE 12803 \u2014 Second Stage RO Brine Module Taltal",
}

OUTPUT_FILE = "2026-03-06_Pre-Meeting-March9-Schedule-Evaluation.docx"


# \u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500 helpers \u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500

def add_para(doc, text, size=11):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    return para


def add_bold(doc, text, size=11):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(size)
    return para


def add_bullet(doc, text, size=11, bold_prefix=None):
    para = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = para.add_run(bold_prefix)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(size)
    r = para.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(size)
    return para


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else tbl.makeelement(qn("w:tblPr"), {})
    borders = tblPr.makeelement(qn("w:tblBorders"), {})
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.makeelement(
            qn(f"w:{border_name}"),
            {
                qn("w:val"): "single",
                qn("w:sz"): "4",
                qn("w:space"): "0",
                qn("w:color"): "000000",
            },
        )
        borders.append(border)
    tblPr.append(borders)
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)


def add_data_table(doc, headers, rows, size=9, fill="D9E2F3"):
    """Tabla generica con encabezado sombreado. Acepta N columnas."""
    n_cols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=n_cols)
    table.autofit = True
    set_table_borders(table)
    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(size)
        shading = cell._element.get_or_add_tcPr()
        shd = shading.makeelement(
            qn("w:shd"),
            {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): fill},
        )
        shading.append(shd)
    # Data rows
    for r_idx, row_data in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            cells[c_idx].text = ""
            run = cells[c_idx].paragraphs[0].add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(size)
    doc.add_paragraph()
    return table


# \u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500 main \u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500

def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.1)
        section.right_margin = Inches(1.1)

    # \u2500\u2500 HEADER \u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500
    header_fields = [
        ("Date:", "March 6, 2026"),
        (
            "From:",
            f"{CONTACTO['nombre']} \u2014 {CONTACTO['cargo']} ({CONTACTO['empresa']})",
        ),
        (
            "To:",
            "Eduardo Yamauchi \u2014 Operations Director Americas (BW Water Americas Inc.)",
        ),
        (
            "CC:",
            "Jeryl F. Regulacion; Adzlan Bin Abd Rahim; Andrew Zaske (BW Water) / "
            "Cesar Malhue; Jorge Guevara; Ronald Pellejero; Victor Gutierrez (ADASA)",
        ),
        (
            "Subject:",
            "March 9 Meeting \u2014 EXW August 3 vs. Engineering Status: Clarification Required",
        ),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / TM N6 (P22-TM-09-000-006-0)",
        ),
        (
            "Attachments:",
            "P22-IT-06-000-002-0_Master-Deliverable-Register.xlsx \u2014 "
            "ADASA Master Deliverable Register (62 items, updated March 6, 2026)",
        ),
    ]
    for label, value in header_fields:
        para = doc.add_paragraph()
        run = para.add_run(label + "  ")
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run = para.add_run(value)
        run.font.name = "Arial"
        run.font.size = Pt(10)

    doc.add_paragraph()

    # \u2500\u2500 OPENING \u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500
    add_para(doc, "Eduardo,")
    doc.add_paragraph()
    add_para(
        doc,
        "ADASA acknowledges receipt of the Baseline Schedule dated March 5, 2026 "
        "(\u201c12803_Taltal Water Treatment Plant_Baseline Schedule\u201d). "
        "Before the March 9 meeting, ADASA has evaluated that document and has "
        "critical observations that require a response.",
    )
    doc.add_paragraph()

    # \u2500\u2500 SECTION 1: EXW vs SCHEDULE \u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500
    add_bold(doc, "The EXW Date vs. the Schedule", size=11)
    doc.add_paragraph()
    add_para(
        doc,
        "The schedule maintains EXW August 2\u20133, 2026 without change. "
        "The engineering data in the same document does not support that date.",
        size=11,
    )
    doc.add_paragraph()

    exw_rows = [
        (
            "Engineering end date",
            "July 9, 2026 (198 days)",
            "Contractual baseline: January 5, 2026. Current delay: +185 days.",
        ),
        (
            "General Engineering",
            "Complete as of February 20, 2026",
            "15 documents never submitted; 14 with open observations.",
        ),
        (
            "EXW delivery",
            "August 2\u20133, 2026 (no change)",
            "No recovery plan submitted. Basis for this date not explained.",
        ),
    ]
    add_data_table(
        doc,
        ["Item", "Schedule \u2014 March 5", "Project records \u2014 March 6"],
        exw_rows,
        size=9,
        fill="D9E2F3",
    )

    add_para(
        doc,
        "ADASA requests that BW Water explain, before March 9, how Engineering at "
        "+185 days over the contractual baseline is consistent with an unchanged EXW "
        "date of August 3, 2026. A recovery plan with document-level milestones \u2014 "
        "covering the 15 pending deliverables and 14 open-observation items \u2014 is required.",
        size=11,
    )
    doc.add_paragraph()
    add_para(
        doc,
        "ADASA\u2019s Master Deliverable Register (Attachment 1) lists all 62 engineering "
        "deliverables with current review verdicts and outstanding actions. BW Water\u2019s "
        "recovery plan must address each item in that register.",
        size=11,
    )
    doc.add_paragraph()

    # \u2500\u2500 SECTION 2: PROCUREMENT TABLE \u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500
    add_bold(doc, "Procurement Items on the Critical Path", size=11)
    doc.add_paragraph()
    add_para(
        doc,
        "Four procurement items are scheduled within the next 30 days. Two carry an "
        "active Code 4 \u2014 Rejected verdict; one is approved as noted but has an "
        "outstanding documentation requirement that blocks the PO; one is formally "
        "approved with a clarifying note. Issuing POs without resolving open conditions "
        "does not accelerate the project; it introduces rework that delays EXW further.",
        size=11,
    )
    doc.add_paragraph()

    procurement_rows = [
        (
            "RO HP Pump",
            "Week of March 4\u201310",
            "Code 2 \u2014 Approved as Noted (TM N6, Feb 27)",
            "PO may proceed. BW Water must add clarifying note in DS: "
            "93 kW nameplate / 78.5 kW at design point / 83\u201385 kW FEDCO internal refs.",
        ),
        (
            "Interstage Turbocharger",
            "TBD",
            "Code 2 \u2014 Approved as Noted (TM N6, Feb 27)",
            "PO on hold: coupling manufacturer\u2019s datasheet not submitted. "
            "Required: MAWP \u2265 1,845 psi, HPB service certification, vibration sensor "
            "mounting documented (ET 5.5.7). No datasheet \u2192 no PO.",
        ),
        (
            "Feed Turbocharger",
            "April 1",
            "Code 4 \u2014 Rejected (TM N6, Feb 27)",
            "PO suspended. Coupling downgraded from 2,000 psi to 1,200 psi (1.19x margin \u2014 "
            "ASME minimum 1.5x). Rev D required: \u2265 2,000 psi preferred (Piedmont Style H); "
            "\u2265 1,800 psi minimum. Submit coupling manufacturer\u2019s datasheet.",
        ),
        (
            "Valve List (109 valves)",
            "March 11",
            "Code 4 \u2014 Rejected (TM N6, Feb 27)",
            "Rev C required before March 10. Duplicate TAGs (VM-09-015, VE-09-008, "
            "VE-09-009, VM-09-065) must be resolved for correct procurement traceability "
            "and SCADA assignment. VE-09-008 (items 44/57 \u2014 ANSI 900\u0023 vs "
            "ANSI 150\u0023) carries direct procurement risk: two physically different "
            "valves share the same TAG.",
        ),
    ]
    add_data_table(
        doc,
        ["Equipment", "PO Date", "ADASA Status", "Action Required"],
        procurement_rows,
        size=9,
        fill="FCE4D6",  # naranja claro para tabla de riesgo
    )

    add_para(
        doc,
        "These four items are on the critical path to EXW. BW Water committed at "
        "the February 18, 2026 meeting: no purchase orders before ADASA approval. "
        "All four items above fall under that commitment.",
        size=11,
    )
    doc.add_paragraph()

    # \u2500\u2500 CLOSING \u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500
    add_para(
        doc,
        "Before March 9, ADASA requires: (a) BW Water\u2019s explanation of how EXW "
        "August 3 remains achievable given Engineering at +185 days over the contractual "
        "baseline, including a document-level recovery plan covering all 15 pending "
        "deliverables and 14 open-observation items with individual dates; "
        "(b) delivery status of the three March 6 commitments \u2014 IO MODBUS list, "
        "Control Architecture Rev C, VFD variables \u2014 and a firm date for the "
        "Modbus TCP Memory Map, outstanding 51 days since TM N2; (c) coupling "
        "manufacturer\u2019s datasheet for the Interstage Turbocharger confirming "
        "MAWP \u2265 1,845 psi, HPB service certification, and vibration sensor mounting "
        "per ET 5.5.7.",
        size=11,
    )
    add_para(
        doc,
        "Without these responses, the March 9 meeting cannot resolve the "
        "project\u2019s critical items.",
        size=11,
    )
    doc.add_paragraph()
    add_para(doc, "Best regards,")
    doc.add_paragraph()

    para = doc.add_paragraph()
    run = para.add_run(CONTACTO["nombre"])
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)

    add_para(doc, CONTACTO["cargo"])
    add_para(doc, CONTACTO["empresa"])
    add_para(doc, f"Project: {CONTACTO['proyecto']}")

    doc.save(OUTPUT_FILE)
    print(f"Correo generado exitosamente: {OUTPUT_FILE}")
    return OUTPUT_FILE


if __name__ == "__main__":
    crear_correo()
