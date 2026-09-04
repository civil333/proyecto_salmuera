"""
Correo interno a Ronald — Revision Control Philosophy TM N7
Fecha: 2026-03-10
Output: 2026-03-10_Revision-Control-Philosophy-Ronald.docx
"""

import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
    "2026-03-10_Revision-Control-Philosophy-Ronald.docx")


def set_font(run, bold=False, size=11, color=None):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_paragraph(doc, text="", bold=False, size=11, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if text:
        run = p.add_run(text)
        set_font(run, bold=bold, size=size)
    return p


def add_header_field(doc, label, value, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    r_label = p.add_run(f"{label}: ")
    set_font(r_label, bold=True, size=11)
    r_value = p.add_run(value)
    set_font(r_value, bold=False, size=11)


def add_bullet(doc, label, text, color=None, space_after=4):
    """Bullet con etiqueta bold coloreada + texto normal."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent = Cm(0.6)
    if label:
        r_label = p.add_run(label + " ")
        set_font(r_label, bold=True, size=11, color=color)
    r_text = p.add_run(text)
    set_font(r_text, bold=False, size=11)
    return p


GRIS = (0x40, 0x40, 0x40)

# ── Generar DOCX ───────────────────────────────────────────────────────────────
doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3.0)
    section.right_margin  = Cm(2.5)

# Titulo / asunto
p_title = doc.add_paragraph()
p_title.paragraph_format.space_after = Pt(10)
r_title = p_title.add_run("Control Philosophy Rev A — TM N7 | Comentarios ADASA listos")
set_font(r_title, bold=True, size=13)

# Header campos
add_header_field(doc, "Para",   "Ronald")
add_header_field(doc, "De",     "Luis Rivera")
add_header_field(doc, "Fecha",  "10 de marzo de 2026")
add_header_field(doc, "Asunto", "Control Philosophy Rev A — TM N7 | Comentarios ADASA listos — pendiente tu validacion",
                 space_after=10)

# Apertura
p_apertura = doc.add_paragraph()
p_apertura.paragraph_format.space_after = Pt(10)
r_ap = p_apertura.add_run(
    "Ronald, termine de revisar la Control Philosophy Rev A (P22-BT-09-009-001) para el TM N7. "
    "Salieron bastantes cosas — te paso el resumen para que veas si tienes algo que agregar antes de que lo despachemos."
)
set_font(r_ap, size=11)

# ── HALLAZGOS ─────────────────────────────────────────────────────────────────
add_paragraph(doc, "Hallazgos:", bold=True, size=11, space_before=4, space_after=3)

hallazgos = [
    ("UPS:",                          "Declaran 30 minutos de autonomia. La ET exige 8 horas. No conformidad contractual directa."),
    ("Temperatura motores:",          "No hay monitoreo continuo. La ET §5.3 exige Pt-100 en devanados y rodamientos de todos los motores. TE-09-001/002 son switches DI (on/off), no transmisores AI — no alcanzan."),
    ("Consumo energetico:",           "No hay MVE integrado al PLC para calcular CEE (kWh/m3). Es una garantia contractual — sin esto no podemos verificar nada en comisionamiento."),
    ("Protocolo:",                    "La CP declara solo 4-20mA, pero la Arquitectura Rev B ya tiene red EtherNet/IP con gateway PLX32. HART tampoco esta declarado y la ET §5.5 lo exige."),
    ("Permisivo cliente:",            "BW Water definio 2 DI individuales (Sec 3.3.3). La interfaz correcta es 1 DI de habilitacion general desde ADASA + 1 DO de estado del modulo hacia nosotros."),
    ("Permisivo arranque HP:",        "No hay umbral de presion minima como condicion de arranque. El PLC arrancaria la bomba con presion insuficiente desde ADASA — riesgo de cavitacion."),
    ("Modbus TCP/IP:",                "La interfaz no esta descrita en ningun lugar de la CP."),
    ("ISA 101:",                      "No declarado como estandar de diseno HMI. La ET lo exige."),
    ("Tags antiscalant:",             "Discrepancia entre VE-07-014/016 y VE-09-014/016."),
    ('Codigo "BT":',                  "No esta definido en el sistema de codificacion P22."),
]

for label, text in hallazgos:
    add_bullet(doc, label, text)

# ── ACCION REQUERIDA ──────────────────────────────────────────────────────────
add_paragraph(doc, "Lo que necesito de ti:", bold=True, size=11, space_before=10, space_after=3)

p_pdf = doc.add_paragraph(style="List Bullet")
p_pdf.paragraph_format.space_after = Pt(4)
p_pdf.paragraph_format.left_indent = Cm(0.6)
r_pdf1 = p_pdf.add_run("Revisar el PDF con los comentarios marcados: ")
set_font(r_pdf1, size=11)
r_pdf2 = p_pdf.add_run(
    "REVISIONES/TRANSMITTALES/P22-TM-09-000-007-0/COMENTARIOS/P22-BT-09-009-001_Rev A_CC_ADASA.pdf"
)
set_font(r_pdf2, size=10, color=GRIS)

add_bullet(doc, None, "Si tienes algo que agregar desde tu lado, avisame y lo incluimos.")
add_bullet(doc, None, "Si esta OK, confirma y despachamos el TM N7 tal como esta.")

# Cierre
add_paragraph(doc, "", space_after=6)
add_paragraph(doc, "Gracias,\nLuis", bold=False, size=11, space_after=0)

doc.save(OUTPUT_FILE)
print(f"DOCX generado: {OUTPUT_FILE}")
