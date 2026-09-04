#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: RE: 25007 Taltal: Meeting Notes — 26-MAR-2026
        Follow-up spare parts + CIP layout
Fecha:  26 de marzo de 2026
"""

import sys
import os

skill_path = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..", "..", "..",
    ".claude", "skills", "template-adasa",
)
sys.path.insert(0, skill_path)

from ejemplo_documento import set_table_borders, calcular_anchos_columnas

from docx import Document
from docx.shared import Pt, Inches

OUTPUT_FILE = "2026-03-26_Follow-Up-Spare-Parts-Turbocharger.docx"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ── HEADER ──────────────────────────────────────────────────────────────
    fields = [
        ("Date:", "March 26, 2026"),
        ("From:", "Luis Rivera \u2014 ADASA"),
        ("To:", "Eduardo Yamauchi \u2014 BW Water Americas Inc."),
        ("CC:", "Jeryl F. Regulacion; Adzlan Bin Abd Rahim; Victor Gutierrez Aqueveque"),
        ("Subject:", "RE: 25007 Taltal: Meeting Notes \u2014 26-MAR-2026"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    # ── SALUDO ───────────────────────────────────────────────────────────────
    para = doc.add_paragraph("Dear Eduardo,")
    aplicar_arial(para)
    doc.add_paragraph()

    # ── INTRO ────────────────────────────────────────────────────────────────
    para = doc.add_paragraph(
        "Thank you for the meeting notes. Two items to follow up on:"
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # ── ITEM 1: SPARE PARTS ──────────────────────────────────────────────────
    para = doc.add_paragraph()
    para.add_run("1.\u00a0 Spare Parts \u2014 Pending Response").bold = True
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "Our request of February 26, 2026 for (a) formal re-validation of the Section 3.8 "
        "recommended two-year spare parts quotation (USD\u00a043,790.00, BWWA Ref. "
        "20.24.6501.F Rev.1) and (b) service/repair kits for SIP-09-001 and SIP-09-002 "
        "turbochargers has not yet been addressed. The original deadline was March 7 \u2014 "
        "this is 19 days overdue."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "With the Feed Turbocharger (SIP-09-001) PO window opening April 1\u20137, we cannot "
        "proceed without a defined spare parts scope. Please provide your response by "
        "Tuesday, March 31, 2026."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # ── ITEM 2: CIP LAYOUT ───────────────────────────────────────────────────
    para = doc.add_paragraph()
    para.add_run("2.\u00a0 CIP Area Layout").bold = True
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "We look forward to receiving the updated CIP layout tomorrow as committed. "
        "This is on our critical path."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # ── CIERRE ────────────────────────────────────────────────────────────────
    para = doc.add_paragraph("Please confirm receipt.")
    aplicar_arial(para)
    doc.add_paragraph()

    # ── FIRMA ─────────────────────────────────────────────────────────────────
    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Luis Rivera").bold = True
    aplicar_arial(para)
    for line in [
        "Contract Administrator",
        "ADASA \u2014 Aguas de Antofagasta S.A.",
        "Project: BAE 12803 \u2014 Second Stage RO Brine Module Taltal",
    ]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
