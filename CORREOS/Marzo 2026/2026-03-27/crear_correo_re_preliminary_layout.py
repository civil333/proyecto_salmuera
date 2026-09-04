"""
Correo: Respuesta a Advanced Copy — Preliminary Equipment Layout & Tie-In Points
Fecha: 27-Mar-2026
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "2026-03-27_Re-Preliminary-Layout-Review.docx")

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def set_font(run, name="Arial", size=10, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_para(doc, text="", size=10, bold=False, space_after=4, color=None,
             align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if text:
        run = p.add_run(text)
        set_font(run, size=size, bold=bold, color=color)
    return p

def add_header_line(doc, label, value, size=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    r1 = p.add_run(label)
    set_font(r1, size=size, bold=True)
    r2 = p.add_run(value)
    set_font(r2, size=size)
    return p

def add_section_title(doc, text, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_font(run, size=size, bold=True)
    run.font.underline = True
    return p

def add_numbered_item(doc, prefix, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.8)
    r1 = p.add_run(prefix + "  ")
    set_font(r1, bold=True)
    r2 = p.add_run(text)
    set_font(r2)

def add_shading(cell, fill_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)

# ---------------------------------------------------------------------------
# Documento
# ---------------------------------------------------------------------------
doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(10)

# ---------------------------------------------------------------------------
# ENCABEZADO
# ---------------------------------------------------------------------------
add_para(doc, "CORREO ELECTR\u00d3NICO", size=12, bold=True,
         color=(31, 56, 100), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

add_header_line(doc, "Fecha:   ", "27 de Marzo de 2026")
add_header_line(doc, "De:      ", "Luis Rivera \u2014 Contract Administrator (ADASA)")
add_header_line(doc, "Para:    ", "Eduardo Yamauchi \u2014 BW Water Americas Inc.")
add_header_line(doc, "CC:      ",
    "Jeryl F. Regulacion, Adzlan Bin Abd Rahim, Sadeep Irugalbandara, Nick Huta; "
    "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
    "Mauricio Vallejos, Jorge Valdes")
add_header_line(doc, "Asunto:  ",
    "Project 12803 \u2014 Taltal SWRO | Re: Preliminary Equipment Layout & "
    "Tie-In Points \u2014 ADASA Review Comments | March 27, 2026")

doc.add_paragraph()

# ---------------------------------------------------------------------------
# SALUDO
# ---------------------------------------------------------------------------
add_para(doc, "Dear Eduardo,", bold=True, space_after=6)
add_para(doc,
    "Tuesday at 8:30\u00a0AM Chile time works for us. Please send the updated invite.",
    space_after=8)

add_para(doc,
    "We have reviewed both drawings. Comments are attached. The following points "
    "require correction before these layouts can be accepted:",
    space_after=4)

# ---------------------------------------------------------------------------
# SECCIÓN 1 — Equipment Layout
# ---------------------------------------------------------------------------
add_section_title(doc, "Equipment Layout (P22-DWG-09-005-003 Rev.\u00a0B)")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
p.paragraph_format.left_indent = Cm(0.5)
r1 = p.add_run("\u2013  ")
set_font(r1)
r2 = p.add_run("Tank must be relocated to the opposite side")
set_font(r2, bold=True)
r3 = p.add_run(" \u2014 current position blocks chemical loading access.")
set_font(r3)

# ---------------------------------------------------------------------------
# SECCIÓN 2 — Tie-In Point Layout
# ---------------------------------------------------------------------------
add_section_title(doc, "Tie-In Point Layout (P22-DWG-09-005-005 Rev.\u00a0B)")

items = [
    ("No space for a pipe rack",
     " \u2014 connections must come directly from the container."),
    ("Dosing tank has no access",
     " \u2014 relocation required."),
    ("Antiscalant loading must go directly to the dosing tank.",
     " No budget for a carrier pump."),
    ("Connection flanges for module relocation cutover are not shown.", ""),
    ("Container elevation view is missing",
     " \u2014 required to confirm tie-in flange locations."),
]

for i, (bold_text, rest) in enumerate(items, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.5)
    r0 = p.add_run(f"{i}.  ")
    set_font(r0)
    r1 = p.add_run(bold_text)
    set_font(r1, bold=True)
    if rest:
        r2 = p.add_run(rest)
        set_font(r2)

# ---------------------------------------------------------------------------
# CIERRE
# ---------------------------------------------------------------------------
add_para(doc, "", space_after=4)
add_para(doc,
    "These items must be addressed in the next revision. "
    "We will review the corrected drawings at Tuesday\u2019s meeting.",
    space_after=10)
add_para(doc, "Please confirm receipt.", space_after=10)
add_para(doc, "Best regards,", space_after=4)
add_para(doc, "Luis Rivera", bold=True, space_after=2)
add_para(doc, "Project Engineer | Contract Administrator", space_after=2)
add_para(doc, "ADASA \u2014 Aguas de Antofagasta S.A.", space_after=2)
add_para(doc, "Contract C-4300 | Project 12803 \u2014 Taltal SWRO", space_after=0)

# ---------------------------------------------------------------------------
# GUARDAR
# ---------------------------------------------------------------------------
doc.save(OUTPUT)
print(f"Generado: {OUTPUT}")
