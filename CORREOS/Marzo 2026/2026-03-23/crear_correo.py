"""
Correo: Seguimiento Log de Compras y Estado Piping Layout / 3D Model
Fecha: 23-Mar-2026
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "2026-03-23_Seguimiento-Procurement-Layout.docx")

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def set_font(run, name="Arial", size=10, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_para(doc, text="", size=10, bold=False, space_after=4, color=None,
             align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if text:
        run = p.add_run(text)
        set_font(run, size=size, bold=bold, color=color)
    return p

def add_header_line(doc, label, value, size=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    r1 = p.add_run(label)
    set_font(r1, size=size, bold=True)
    r2 = p.add_run(value)
    set_font(r2, size=size)
    return p

def add_section_title(doc, text, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_font(run, size=size, bold=True)
    # Subrayado
    run.font.underline = True
    return p

def build_table(doc, headers, rows, col_widths_cm):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    # Header row
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Cm(col_widths_cm[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_font(run, size=9, bold=True, color=(31, 56, 100))
        # Fondo azul claro en header
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "D9E1F2")
        tcPr.append(shd)
    # Data rows
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, cell_text in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = Cm(col_widths_cm[c_idx])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            bold_cell = c_idx == 3 and ("cannot" in cell_text.lower()
                                         or "no datasheet" in cell_text.lower()
                                         or "rejected" in cell_text.lower())
            run = p.add_run(cell_text)
            color = (192, 0, 0) if bold_cell else None
            set_font(run, size=9, bold=bold_cell, color=color)
    doc.add_paragraph()  # espacio post-tabla

# ---------------------------------------------------------------------------
# Documento
# ---------------------------------------------------------------------------
doc = Document()

# Márgenes
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# Fuente por defecto
style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(10)

# ---------------------------------------------------------------------------
# ENCABEZADO DEL CORREO
# ---------------------------------------------------------------------------
add_para(doc, "CORREO ELECTRÓNICO", size=12, bold=True,
         color=(31, 56, 100), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

add_header_line(doc, "Fecha:   ", "23 de Marzo de 2026")
add_header_line(doc, "De:      ", "Luis Rivera — Contract Administrator (ADASA)")
add_header_line(doc, "Para:    ", "Eduardo Yamauchi — BW Water Americas Inc.")
add_header_line(doc, "CC:      ",
    "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, Mauricio Vallejos, "
    "Jorge Valdes, Tanya Figueroa, Allan Valentos, Jeryl F. Regulacion, "
    "Sadeep Irugalbandara, Andrew Sia, Ghazi Ozair, Nick Huta, Marjan Arsovic, "
    "Gerald Ross, Andrew Zaske, Adzlan Bin Abd Rahim")
add_header_line(doc, "Asunto:  ",
    "Project 12803 — Taltal SWRO | Procurement Log Update & Piping Layout / 3D Model Status "
    "| March 23, 2026")

doc.add_paragraph()  # separador

# ---------------------------------------------------------------------------
# SALUDO
# ---------------------------------------------------------------------------
add_para(doc, "Dear Eduardo,", bold=True, space_after=6)
add_para(doc,
    "Please find below ADASA\u2019s weekly procurement log review (BW Water schedule, "
    "March\u00a05, 2026) and a status update on Proposal 25007-PL-0001.",
    space_after=10)

# ---------------------------------------------------------------------------
# SECCIÓN 1 — Semana 16-22 Mar
# ---------------------------------------------------------------------------
add_section_title(doc, "1. Procurement Log \u2014 Week of March 16\u201322, 2026")

build_table(doc,
    headers=["Equipment", "PO Window", "ADASA Code", "Status"],
    col_widths_cm=[4.5, 2.8, 4.2, 7.0],
    rows=[
        ("CIP Tank Heater", "16\u201320 Mar",
         "Code 1 \u2014 Approved",
         "Confirm PO issued."),
        ("Container (40\u2019)", "18\u201324 Mar",
         "Code 1 \u2014 Approved",
         "Confirm PO issued (window closes 24-Mar)."),
        ("All Valve", "11\u201317 Mar",
         "Code 3\n(TM N11)",
         "PO on hold. Valve List Rev D required: "
         "duplicate TAGs (VE-09-007, PSV-09-002) and area-code error (VM-07-005)."),
        ("Instrument Set /\nIO List", "11\u201317 Mar",
         "Code 3\n(TM N8 / N10)",
         "PO on hold. IL Rev C + IO List Rev C required."),
    ]
)

# ---------------------------------------------------------------------------
# SECCIÓN 2 — Semana 23-29 Mar
# ---------------------------------------------------------------------------
add_section_title(doc, "2. Procurement Log \u2014 Week of March 23\u201329, 2026")

build_table(doc,
    headers=["Equipment", "PO Window", "ADASA Code", "Action"],
    col_widths_cm=[4.5, 2.8, 4.2, 7.0],
    rows=[
        ("CIP / Flushing\nCartridge Filter", "23\u201327 Mar",
         "Code 2 \u2014 Approved\nas Noted",
         "Confirm PO or advise issuance date."),
        ("CIP / Flushing Pumps", "25\u201331 Mar",
         "Code 2 \u2014 Approved\nas Noted",
         "Confirm PO or advise issuance date."),
        ("CIP / Flushing Tank", "27 Mar \u2013 2 Apr",
         "Code 1 \u2014 Approved",
         "Confirm PO or advise issuance date."),
    ]
)

# ---------------------------------------------------------------------------
# SECCIÓN 3 — Próxima semana
# ---------------------------------------------------------------------------
add_section_title(doc, "3. Upcoming Items \u2014 Week of March 30 \u2013 April 7, 2026")

build_table(doc,
    headers=["Equipment", "PO Window", "ADASA Code", "Note"],
    col_widths_cm=[4.5, 2.8, 4.2, 7.0],
    rows=[
        ("RO Pressure Vessel /\nTubes", "31 Mar \u2013 6 Apr",
         "Code 1\n(SWRO System DS)",
         "PO issued 24-Feb. Confirm vessel housing is covered by UHPRO PO "
         "or advise if separate procurement is required."),
        ("Feed Turbocharger\n(SIP-09-001)", "1\u20137 Apr",
         "Code 2 \u2014 Approved\nas Noted (TM N11)",
         "PO may proceed. NOTE-03: update coupling label "
         "\u201cStyle 77\u201d \u2192 \u201cStyle S\u201d in next GA revision."),
        ("Structural Frames /\nSupports", "6\u201310 Apr",
         "No datasheet submitted",
         "Submit datasheet immediately \u2014 PO window opens in 14 days."),
    ]
)

# ---------------------------------------------------------------------------
# SECCIÓN 4 — Propuesta Rev.2 / Layout
# ---------------------------------------------------------------------------
add_section_title(doc, "4. Proposal 25007-PL-0001 Rev.2 \u2014 ADASA Position")
add_para(doc,
    "ADASA has received Proposal Rev.2 (March 18, 2026 \u2014 6\u00a0documents, "
    "USD\u00a05,766) and sets out its position:",
    space_after=6)

add_para(doc,
    "Phase 1 \u2014 ready to close commercially:",
    bold=True, space_after=2)
add_para(doc,
    "Piping Layout Rev B, Tie-In Points Rev B, Equipment Layout Rev B, 3D Model.",
    space_after=8)

add_para(doc,
    "Phase 2 \u2014 contingent on Phase 1:",
    bold=True, space_after=2)
add_para(doc,
    "Instrument Layout Rev B and Grounding Layout Rev B. Both carry Code 3 under "
    "TM\u00a0N11 (OBS-03 and OBS-04 \u2014 Major) because their current arrangement is based "
    "on the non-conforming Piping Layout Rev A (11,150\u00a0mm). These documents cannot be "
    "closed before Piping Layout Rev B is delivered.",
    space_after=10)

add_para(doc,
    "We request, by end of business March 25, 2026:",
    space_after=4)

items = [
    ("a)", "Confirmed delivery date for Piping Layout Rev B and 3D Model (Phase 1);"),
    ("b)", "Expected delivery date for Instrument Layout Rev B and Grounding Layout Rev B "
     "(Phase 2); and"),
    ("c)", "BW Water\u2019s commercial close confirmation for both phases."),
]
for prefix, text in items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.8)
    r1 = p.add_run(prefix + "  ")
    set_font(r1, bold=True)
    r2 = p.add_run(text)
    set_font(r2)

# ---------------------------------------------------------------------------
# CIERRE
# ---------------------------------------------------------------------------
add_para(doc,
    "Finally, we remind BW Water that an updated procurement log reflecting current PO status "
    "is due this week per the March\u00a05 schedule commitment. Please submit the updated log "
    "together with your response.",
    space_after=10)
add_para(doc, "Please confirm receipt of this message.", space_after=10)
add_para(doc, "Best regards,", space_after=4)
add_para(doc, "Luis Rivera", bold=True, space_after=2)
add_para(doc, "Project Engineer | Contract Administrator", space_after=2)
add_para(doc, "ADASA — Aguas de Antofagasta S.A.", space_after=2)
add_para(doc, "Contract C-4300 | Project 12803 — Taltal SWRO", space_after=0)

# ---------------------------------------------------------------------------
# GUARDAR
# ---------------------------------------------------------------------------
doc.save(OUTPUT)
print(f"Generado: {OUTPUT}")
