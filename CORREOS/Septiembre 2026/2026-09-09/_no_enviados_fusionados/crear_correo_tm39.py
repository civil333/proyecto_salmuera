#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
crear_correo_tm39.py
Correo de cobertura del Transmittal N39 (P22-TM-09-000-039-0), sobre los
submittals 25007-0091 (ENTREGA 91) y 25007-0092 (ENTREGA 92), tres documentos.

LO QUE HACE DISTINTO A ESTE CORREO

NO HAY NADA QUE RECLAMAR DE FONDO. Los tres documentos son re-emisiones y los tres
cierran el punto que los tenia detenidos. El correo lo dice en una linea y no narra
el avance: esa vision de proyecto no se envia, y el detalle vive en el transmittal.

ENCABEZA CON LA UNICA CONSECUENCIA DE COMPRA. La valvula VM-09-065 figura en PVC
sobre una linea que la Line List Rev 1, aprobada en Codigo 1 en el N38, lleva en
316L. La lista de valvulas es contra la que se compra, de modo que ese punto tiene
fecha aunque el documento se acepte.

LOS PLAZOS DE REVISION VAN AL FINAL Y SIN ESCALAR. Son el octavo y noveno formulario
consecutivo que pide devolucion en tres dias corridos contra los siete habiles de la
Clausula 37.2. Se registra el hecho y se deja constancia de que ADASA devuelve dentro
de plazo igual. Sin cifras de multa y sin invocar ningun umbral.

DISTRIBUCION. La fijo el usuario al planificar el N39: Eduardo Yamauchi como
destinatario, con copia a Victor Gutierrez y Jeryl Regulacion. Es mas acotada que la
del N37 y el N38, que llevaban la lista completa de BW Water.
  ATENCION: confirmar la distribucion antes de enviar.

Ingles. Document() directo, sin template ADASA. Estado: BORRADOR.
"""
import os

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-09-09_Transmittal-N39.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"

# Enlace de descarga de la carpeta COMENTARIOS del N39. Publicar la carpeta,
# pegar el enlace y REGENERAR. Verificar sobre el .docx emitido, no aqui: el
# proyecto lleva cuatro enlaces perdidos al pegar el cuerpo en Outlook.
DOWNLOAD_LINK = "PENDIENTE - publicar la carpeta COMENTARIOS y pegar el enlace"


def aplicar_arial(paragraph, size=11):
    for r in paragraph.runs:
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return paragraph


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph()
    p.add_run(text)
    return aplicar_arial(p, size)


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for texto, bold in segs:
        r = p.add_run(texto)
        r.bold = bold
    return aplicar_arial(p, size)


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.add_run("•  ")
    p.add_run(lead).bold = True
    p.add_run(rest)
    return aplicar_arial(p, size)


def blank(doc):
    return doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "September 9, 2026"),
        ("From:", CONTACTO + " - Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi - BW Water"),
        ("CC:", "Jeryl F. Regulacion - BW Water; Victor Gutierrez - ADASA"),
        ("Subject:", "TALTAL - Transmittal N39 - submittals 25007-0091 and "
                     "25007-0092"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    # --- Parrafo 1: el adjunto y el veredicto, sin narrar el avance
    add_segments(doc, [
        ("Attached is ", False),
        ("Transmittal N39 (P22-TM-09-000-039-0)", True),
        (", on submittals 25007-0091 and 25007-0092, three documents. ", False),
        ("Verdict: 2 — Approved as noted. One Code 1, two Code 2 and no "
         "Code 3.", True),
        (" No document returns to revision. All three are re-issues, and ADASA has "
         "reviewed each one only against the comments it raised before.", False),
    ])
    blank(doc)

    # --- Parrafo 2: la unica consecuencia con fecha
    add_segments(doc, [
        ("One point has a date on it. ", False),
        ("VM-09-065 is still listed with a PVC body on a line that the approved Line "
         "List carries in 316L stainless steel.", True),
        (" The CIP pump suction line CP-SS316-DN150-09-022 was changed to stainless "
         "steel on the Line List Rev 1, which ADASA approved at Code 1 in Transmittal "
         "N38, and VM-09-065 is the only DN150 valve of that circuit. Valves are "
         "bought against this list, so the material of that valve should be settled "
         "before the purchase order goes out, and not only when the list is issued at "
         "Revision 0.", False),
    ])
    blank(doc)

    # --- Parrafo 3: el segundo Code 2 y los plazos, sin escalar
    add_segments(doc, [
        ("The second Code 2 is the antiscalant tank drawing, where the level marking "
         "row still carries no size and no elevation. Both points are incorporated "
         "when each document is issued at Revision 0, with no further approval "
         "revision. ", False),
        ("Both forms again ask for return within three calendar days, against the "
         "seven working days of Clause 37.2", True),
        (", which makes eight and nine consecutive submittals. ADASA returns both well "
         "inside the contractual period.", False),
    ])
    blank(doc)

    add_para(doc, "The transmittal and the two annotated PDFs are available here:")
    add_para(doc, DOWNLOAD_LINK)
    blank(doc)

    add_bullet_lead(doc, "Valve List Rev E",
                    " — P22-LI-09-005-002_E_Valve_List_CC_ADASA.pdf")
    add_bullet_lead(doc, "GA of Antiscalant Dosing Tank Rev D",
                    " — P22-DWG-09-005-015_D_GA_Antiscalant_Dosing_Tank_CC_ADASA.pdf")
    blank(doc)

    add_para(doc, "We remain at your disposal for any clarification.")
    blank(doc)
    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in ["Infrastructure Engineering Lead - ADASA - Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)

    # Metadatos explicitos: sin esto Word deja author='python-docx'.
    cp = doc.core_properties
    cp.title = "TALTAL - Transmittal N39 - submittals 25007-0091 and 25007-0092"
    cp.author = "Luis Rivera Gonzalez"
    cp.last_modified_by = "Luis Rivera Gonzalez"
    cp.category = "Correo de cobertura"
    cp.comments = ("Correo de cobertura del Transmittal N39. Encabeza la valvula "
                   "VM-09-065, que es la unica consecuencia con fecha porque la lista "
                   "gobierna la compra. Cadena de los transmittals.")
    doc.save(OUTPUT_FILE)
    print("Correo generado: " + OUTPUT_FILE)
    if DOWNLOAD_LINK.startswith("PENDIENTE"):
        print("AVISO: el enlace de descarga sigue en placeholder. "
              "Publicar la carpeta, pegar el enlace y REGENERAR.")


if __name__ == "__main__":
    crear_correo()
