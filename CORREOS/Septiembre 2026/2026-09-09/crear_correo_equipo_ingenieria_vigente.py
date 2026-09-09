#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo al equipo de proyecto de Aguas Antofagasta: comparte para revision interna la
Nota Tecnica P22-NT-06-000-001-0 y el paquete Ingenieria Vigente para Construccion,
antes de emitirlo al contratista adjudicado del montaje.

Registro de correspondencia: primera persona, cierre "Quedo atento".
Idioma del documento fijado en es-CL.
"""
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-09-09_Equipo-Ingenieria-Vigente-Revision.docx")
CONTACTO = "Luis Rivera"
LANG = "es-CL"


def aplicar_arial(p, size=11):
    for r in p.runs:
        r.font.name = "Arial"; r.font.size = Pt(size)


def _lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang"); rPr.append(el)
    el.set(qn("w:val"), lang); el.set(qn("w:eastAsia"), lang); el.set(qn("w:bidi"), lang)


def fijar_idioma_documento(doc, lang=LANG):
    try:
        _lang(doc.styles["Normal"].element.get_or_add_rPr(), lang)
    except KeyError:
        pass
    for p in doc.paragraphs:
        for r in p.runs:
            _lang(r._element.get_or_add_rPr(), lang)


def parrafo(doc, texto, size=11, bold=False):
    p = doc.add_paragraph(); r = p.add_run(texto); r.bold = bold
    aplicar_arial(p, size); return p


def crear_correo():
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(0.8)

    campos = [
        ("Fecha:", "9 de septiembre de 2026"),
        ("De:", f"{CONTACTO} — Dirección de Ingeniería y Obras, Aguas Antofagasta S.A."),
        ("Para:", "Víctor Gutiérrez Aqueveque; Luciano Méndez Huidobro; Jorge Guevara Lizana; "
                  "Manuel Aguilera Orellana; Yohana Rodríguez Flores"),
        ("Asunto:", "PD Taltal — Ingeniería vigente para construcción, revisión del equipo antes "
                    "de emitirla al contratista"),
        ("Ref:", "Nota Técnica P22-NT-06-000-001-0 / Contrato de Montaje Mecánico y Obras Civiles, "
                 "BAE 12803"),
    ]
    for etiqueta, valor in campos:
        p = doc.add_paragraph()
        r1 = p.add_run(f"{etiqueta} "); r1.bold = True
        p.add_run(valor)
        aplicar_arial(p, 11)

    doc.add_paragraph()
    parrafo(doc, "Estimados:")
    doc.add_paragraph()

    parrafo(doc,
        "Junto con saludar, les comparto la Nota Técnica P22-NT-06-000-001-0 y el paquete "
        "Ingeniería Vigente para Construcción, para su revisión antes de que lo emitamos al "
        "contratista del montaje.")
    doc.add_paragraph()

    parrafo(doc,
        "El paquete reúne 55 documentos en 88 archivos: la ingeniería de detalle mecánica, las "
        "18 láminas y 2 especificaciones de obras civiles, y las dos especificaciones de "
        "montaje. La nota es el único documento de control y lleva la tabla de vigencia "
        "completa, documento por documento, con la revisión que rige.")
    doc.add_paragraph()

    parrafo(doc,
        "Lo que cambió respecto de la ingeniería con la que se cotizó: el módulo subió 250 "
        "milímetros y con él las cuatro cotas de conexión. Dos partidas del Capítulo 4 bajan de "
        "cantidad, la fundación del sistema CIP (4.2) de 7,36 a 5,80 metros cúbicos y la "
        "excavación común (4.6) de 115,10 a 90,58. Las dos se miden por unidad de obra, de modo "
        "que se pagan según lo realmente ejecutado.")
    doc.add_paragraph()

    parrafo(doc,
        "Una advertencia práctica sobre el modelo. La carpeta 04_ MODELO lleva dos archivos de "
        "Navisworks y el que incluye la nube de puntos pesa 6,4 gigabytes, por lo que se entrega "
        "por enlace y no dentro del comprimido. Para revisar basta el de 21 megabytes.")
    doc.add_paragraph()

    parrafo(doc,
        "Les pido sus comentarios al viernes 11 de septiembre, para emitirlo al contratista la "
        "semana entrante.")
    doc.add_paragraph()

    parrafo(doc, "Quedo atento.")
    doc.add_paragraph()
    parrafo(doc, "Saludos Cordiales,")
    doc.add_paragraph()
    parrafo(doc, CONTACTO)
    parrafo(doc, "Dirección de Ingeniería y Obras")
    parrafo(doc, "Aguas Antofagasta S.A. — Grupo EPM")

    doc.core_properties.author = "Luis Rivera Gonzalez"
    doc.core_properties.company = "Aguas Antofagasta"
    doc.core_properties.comments = ("Correo al equipo de proyecto: revision interna de la NT "
                                    "P22-NT-06-000-001-0 y del paquete de ingenieria vigente")
    fijar_idioma_documento(doc)
    doc.save(OUTPUT_FILE)
    print(f"OK: {os.path.basename(OUTPUT_FILE)}")


if __name__ == "__main__":
    crear_correo()
