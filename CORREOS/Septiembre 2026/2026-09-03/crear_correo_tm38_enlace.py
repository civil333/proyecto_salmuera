#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
crear_correo_tm38_enlace.py
Reenvio del enlace de descarga del Transmittal N38.

POR QUE EXISTE ESTE CORREO. El bloque del enlace NO sobrevivio al pegado en
Outlook del correo del 3-Sep 18:19: salieron los tres parrafos sustantivos, pero
no la linea que introduce el enlace, ni el enlace, ni las tres vinetas de los
CC_ADASA. Verificado sobre el respaldo 'transmittal 38 enviado.pdf'. Es el modo de
falla que la memoria feedback_enlace_no_sobrevive_al_pegado_en_outlook ya tenia
catalogado: el enlace va dos veces.

CONSECUENCIA QUE REPARA. BW Water tiene el transmittal en PDF pero no los tres
PDF anotados, y la Seccion 4 del transmittal los remite a un enlace que no llego.

REGLA DURA: va como RESPUESTA A LA MISMA CADENA del correo del transmittal, no
como correo nuevo. El asunto lleva el Re: y el mismo texto.

Transaccional y corto: no repite sustancia del transmittal ni del correo anterior,
no reabre el veredicto y no vuelve a exigir nada. Solo entrega el enlace.

Ingles. Document() directo, sin template ADASA. Estado: BORRADOR.
"""
import os

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-09-03_Transmittal-N38_Enlace.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"

# El mismo enlace del transmittal. Verificar sobre el .docx emitido, no aqui.
DOWNLOAD_LINK = (
    "https://lrg.synology.me:6501/d/s/19k2hKXFFerUDkcqIAD9rTIMtlMiuCaV/"
    "1RLqel9vgvdO6k-x81qjr2m6te-xWOR7-O7kgscxYeg0"
)


def aplicar_arial(paragraph, size=11):
    for r in paragraph.runs:
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return paragraph


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph()
    p.add_run(text)
    return aplicar_arial(p, size)


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.add_run("•  ")
    p.add_run(lead).bold = True
    p.add_run(rest)
    return aplicar_arial(p, size)


def blank(doc):
    return doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "September 3, 2026"),
        ("From:", CONTACTO + " - Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi, Fitri Indriyani - BW Water; "
                "Victor Gutierrez, Ronald Pellejero, Jorge Guevara - ADASA"),
        ("CC:", "Jeryl F. Regulacion, Nick Huta, Sadeep Irugalbandara, "
                "David Chee Keat Swee, Andrew Sia, Adzlan Bin Abd Rahim, "
                "Tanya Figueroa, Ahmad Iqbal Bin Azam, Stephane Gehant - BW Water"),
        ("Subject:", "RE: TALTAL - Transmittal N38 - outstanding submittals to be "
                     "closed with documents this week"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_para(doc, "The download link did not carry through in my previous message. "
                  "The three annotated PDFs referred to in Section 4 of the transmittal "
                  "are available here:")
    blank(doc)
    add_para(doc, DOWNLOAD_LINK)
    blank(doc)

    add_bullet_lead(doc, "Ultrasonic Thickness Procedure Rev C",
                    " — P22-BA-09-000-016_C_Ultrasonic_Thickness_CC_ADASA.pdf")
    add_bullet_lead(doc, "HMI Display Screenshot Rev C",
                    " — P22-LI-09-008-016_C_HMI_Display_Screenshot_CC_ADASA.pdf")
    add_bullet_lead(doc, "PLC/LCP FAT Procedure - Hardware Rev C",
                    " — P22-PP-09-000-001_C_PLC_LCP_FAT_Procedure_CC_ADASA.pdf")
    blank(doc)

    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in ["Infrastructure Engineering Lead - ADASA - Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)

    cp = doc.core_properties
    cp.title = "TALTAL - Transmittal N38 - download link"
    cp.author = "Luis Rivera Gonzalez"
    cp.last_modified_by = "Luis Rivera Gonzalez"
    cp.category = "Correo de cobertura"
    cp.comments = ("Reenvio del enlace de descarga del TM N38: el bloque no sobrevivio "
                   "al pegado en Outlook del correo de las 18:19. Responde a la misma "
                   "cadena.")
    doc.save(OUTPUT_FILE)
    print("Correo generado:", OUTPUT_FILE)


if __name__ == "__main__":
    crear_correo()
