#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
crear_correo.py — Correo de cobertura del Transmittal N40 (P22-TM-09-000-040-0).

A Eduardo Yamauchi, cadena regular de transmittals. Cubre los submittals 25007-0093,
25007-0094 y 25007-0095. Los dos CC_ADASA (FAT del modulo, 181 kB; Schematic Rev C,
4,9 MB) van ADJUNTOS al correo, sin enlace Synology: el parrafo del enlace se cayo en
cuatro envios seguidos y el peso total baja de 6 MB.

QUE HACE EL CORREO Y QUE NO. Remite y no repite (version ejecutiva pedida por el usuario
el 18-Sep: "los detalles estan en el transmittal"): veredicto, una vineta por bloque con la
accion que falta y su fecha (Rev B del FAT al viernes 25-Sep-2026, cinco dias habiles desde
la emision; terminal -D9P reafirmado; doce documentos sin accion) y los adjuntos. Fuera: el
agradecimiento por la fecha, la lista de defectos del FAT, la explicacion del -D8S, la linea
sobre la devolucion a tres dias corridos (queda en la descripcion interna). Sin vision de
proyecto, sin narrar cierres. Primera persona (lo firma Luis Rivera).

Distribucion: Eduardo destinatario; copia Jeryl F. Regulacion (BW Water) y Victor
Gutierrez (ADASA). Ingles. Document() directo, sin template ADASA. Estado: BORRADOR.
"""
import os

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-09-18_Transmittal-N40-FAT-Procedure.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"
FECHA_REV_B = "Friday 25 September 2026"


def aplicar_arial(paragraph, size=11):
    for r in paragraph.runs:
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return paragraph


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph()
    p.add_run(text)
    return aplicar_arial(p, size)


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for texto, bold in segs:
        r = p.add_run(texto)
        r.bold = bold
    return aplicar_arial(p, size)


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.add_run("•  ")
    p.add_run(lead).bold = True
    p.add_run(rest)
    return aplicar_arial(p, size)


def blank(doc):
    return doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "September 18, 2026"),
        ("From:", CONTACTO + " - Leader, Infrastructure Engineering (ADASA)"),
        ("To:", "Eduardo Yamauchi - BW Water"),
        ("CC:", "Jeryl F. Regulacion - BW Water; Victor Gutierrez - ADASA"),
        ("Subject:", "TALTAL - Transmittal N40 - Module Factory Acceptance Test "
                     "Procedure and submittals 25007-0093 to 0095"),
        ("Ref:", "Contract C-4300 / BAE 12803 / Technical Specification "
                 "P22-ET-09-000-001-0, Section 8.1 / Inspection and Test Plan "
                 "P22-BA-09-000-004 Rev 0"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_segments(doc, [
        ("Please find attached Transmittal N40 (P22-TM-09-000-040-0), covering submittals "
         "25007-0093, 25007-0094 and 25007-0095, fourteen documents. ", False),
        ("Verdict: 3 - To be revised.", True),
    ])
    blank(doc)

    add_bullet_lead(doc, "Factory Acceptance Test Procedure Rev A - Code 3. ",
                    "It is not yet the detailed procedure that Section 8.1 of the Technical "
                    "Specification and row 7.1 of the Inspection and Test Plan require. "
                    "Please issue Rev B by " + FECHA_REV_B + ". Row 7.1 is an ADASA Hold "
                    "Point: the FAT cannot formally open until the procedure is approved.")
    add_bullet_lead(doc, "PLC/LCP Schematic Diagram Rev C - Code 2. ",
                    "The operator terminal stays as declared at Transmittal N30, "
                    "2711P-T10C22D9P; the approved PLC and HMI Datasheet is not to be "
                    "revised. Replace row 11 of the bill of material at Rev 0.")
    add_bullet_lead(doc, "Twelve documents - Code 1. ",
                    "Approved at Revision 0, no action.")
    blank(doc)

    add_para(doc, "The observations are itemised on the two annotated PDFs.")
    blank(doc)

    add_para(doc, "Attachments:")
    add_bullet_lead(doc, "TRANSMITTAL N40 ADASA-BW_WATER.pdf", "")
    add_bullet_lead(doc, "P22-BA-09-000-017_A_FAT_Procedure_CC_ADASA.pdf",
                    " — Factory Acceptance Test Procedure Rev A, annotated")
    add_bullet_lead(doc, "P22-CD-09-008-002_C_PLC_LCP_Schematic_CC_ADASA.pdf",
                    " — PLC/LCP Schematic Diagram Rev C, annotated")
    blank(doc)

    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in ["Leader, Infrastructure Engineering",
                 "ADASA, Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)

    cp = doc.core_properties
    cp.title = "TALTAL - Transmittal N40 - Module Factory Acceptance Test Procedure"
    cp.author = "Luis Rivera Gonzalez"
    cp.last_modified_by = "Luis Rivera Gonzalez"
    cp.company = "Aguas Antofagasta"
    cp.category = "Correo de cobertura"
    cp.comments = ("Cobertura del Transmittal N40: FAT del modulo Rev A en Codigo 3 con "
                   "Rev B al 25-Sep-2026; Schematic Rev C Codigo 2 con el terminal -D9P "
                   "reafirmado; doce documentos en Codigo 1.")
    doc.save(OUTPUT_FILE)

    palabras = sum(len(p.text.split()) for p in doc.paragraphs)
    print("Correo generado: " + OUTPUT_FILE)
    print(f"  palabras del documento: {palabras}")


if __name__ == "__main__":
    crear_correo()
