#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
crear_correo_izaje_pe.py
Correo a Eduardo Yamauchi: el paquete de izaje del modulo y el alcance de la
revision estructural que BW Water encargo el 9 de septiembre.

POR QUE AHORA Y NO DESPUES

La minuta del 09-09-26 registra que la orden al profesional se coloco ese dia con
un compromiso de cuatro dias. El izaje y el sismo viven en el MISMO informe y en
el MISMO modelo STAAD, de modo que ampliar el alcance mientras el trabajo esta
abierto no agrega carga: la ordena. Si el correo llega despues, la revision hay
que pagarla dos veces. Por eso el correo abre como oportunidad y no como reproche.

EL ENCUADRE, QUE ES DONDE ESTE RECLAMO SE GANA O SE PIERDE

La ET NO exige un profesional chileno: cero apariciones de "profesional",
"ingeniero", "endoso" o "inscrito" en la ET, la BAE y el PIE. Fundarlo ahi seria
refutable de una linea. Lo que la ET SI exige es NCh 2369 Zona 3 y la aprobacion
de ADASA, y BW Water eligio el PE local como su via de cumplimiento, declarada por
escrito tres veces (20-May, 16-Jun y 30-Jun-2026). ADASA no le impone un
profesional: le pide cerrar la via que el mismo eligio.

EL PUNTO MAS FUERTE ES UNA CONTRADICCION DE SU PROPIO EXPEDIENTE, no una opinion
de ADASA: el Criterio de Diseno P22-CD-09-005-003 Rev 0, seccion F, pagina 10 de
10, manda disenar el padeye con FS 2,0 y carga lateral fuera de plano del 5 por
ciento; el Informe de Calculo Rev 0 no trae ninguna seccion de diseno de padeye.
Verificado sobre el indice del propio informe, que salta de "Bolt Design at the
Base" a los anexos. Y el informe se emitio igual apto para construccion.

LO QUE YA ESTA SE DICE PRIMERO. El calculo Rev 0 SI trae dos condiciones de izaje
en STAAD, los factores de API RP 2A-WSD, el centro de gravedad y la tension de
eslinga. Reclamar el izaje entero seria falso y quemaria el correo.

EL PLAZO NO SE IMPONE, por decision del usuario: se pide en la misma fecha que BW
Water comprometio para la entrega del profesional, y se le pide confirmarla.

QUEDA FUERA, disponible si escala: la multa de la Clausula 43.1 letra a), el hito
de pago, y que el PRG-32 vencio el 22-Ago hace diecinueve dias. El cierre lleva la
reserva generica de derechos.

ORTOGRAFIA: el documento se emite en en-US, de modo que la grafia es AMERICANA.
Nada de centre, modelled, metre ni analyse: Word los subraya en rojo igual que
las tildes en un correo en espanol declarado en ingles.

Ingles, primera persona. Document() directo, sin template ADASA. Responde a la
cadena del Transmittal N39. Estado: BORRADOR.
"""
import os

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-09-10_Lifting-Package-and-PE-Scope.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for r in paragraph.runs:
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return paragraph


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph()
    p.add_run(text)
    return aplicar_arial(p, size)


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for texto, bold in segs:
        r = p.add_run(texto)
        r.bold = bold
    return aplicar_arial(p, size)


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.add_run("•  ")
    p.add_run(lead).bold = True
    p.add_run(rest)
    return aplicar_arial(p, size)


def blank(doc):
    return doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "September 10, 2026"),
        ("From:", CONTACTO + " - Leader, Infrastructure Engineering (ADASA)"),
        ("To:", "Eduardo Yamauchi - BW Water"),
        ("CC:", "Jeryl F. Regulacion - BW Water; Victor Gutierrez - ADASA"),
        ("Subject:", "TALTAL - Lifting package for the module and the scope of the "
                     "structural review"),
        ("Ref:", "Contract C-4300 / BAE 12803 / Technical Specification "
                 "P22-ET-09-000-001-0"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    # --- 1. la oportunidad: el trabajo del profesional esta abierto ahora
    add_segments(doc, [
        ("Yesterday you confirmed the purchase order for the structural review was "
         "being placed that day, with a four-day commitment. Before it closes, its "
         "scope needs to cover the lifting of the module: lifting and seismic sit in "
         "the same report and the same STAAD model, so covering both now costs one "
         "review instead of two. ", False),
        ("Please confirm the professional's delivery date, because I am asking for the "
         "lifting package on that same date.", True),
    ])
    blank(doc)

    # --- 2. lo que YA esta. Las cifras del CG y de las tensiones NO van aqui: se
    #     guardan para el parrafo del yugo, que es donde prueban algo.
    add_segments(doc, [
        ("A good part is already done, and I want to say so first. The Structural "
         "Calculation Report (P22-CD-09-005-001 Rev 0) models two lifting conditions in "
         "STAAD under the API RP 2A-WSD combinations of 1.35 D and 2.0 D, and gives the "
         "center of gravity and the sling tensions.", False),
    ])
    blank(doc)

    # --- 3. el padeye y la disyuntiva de rutas, en un solo parrafo. La cita va en
    #     linea y recortada a las dos frases que prueban el punto.
    add_segments(doc, [
        ("What is missing is what turns that into something a rigger can execute. Your "
         "own criteria require a padeye (a welded lifting lug) designed to “a "
         "factor of safety, FS, of "
         "2.0” and to “a lateral out of plane load equivalent to 5% of the "
         "gravity load” (P22-CD-09-005-003 Rev 0, section F, page 10). ", False),
        ("The calculation contains no padeye design at all, and it does not lift from "
         "padeyes either", True),
        (": it lifts from the container's own ISO corner fittings. Those are two "
         "different routes, and the report does not say which is being built.", False),
    ])
    blank(doc)

    # --- 4. el yugo NO es un accesorio: es una hipotesis del calculo ya emitido
    add_segments(doc, [
        ("The yoke is not an accessory to be chosen later. It is an assumption your "
         "calculation has already made.", True),
        (" The four sling tensions are not equal: 129.2, 125.3, 90.5 and 88.5 kN. That "
         "spread comes from a center of gravity that sits off center along the module "
         "(5,486 mm one way and 6,514 the other). A yoke built to a different geometry "
         "gives different tensions and the analysis stops applying, so either the yoke "
         "follows the arrangement the model assumed, or the model is rerun against the "
         "yoke that is built.", False),
    ])
    blank(doc)

    # --- 5. las cuatro piezas
    add_segments(doc, [
        ("Together with the three deliverables the Technical Specification "
         "(P22-ET-09-000-001-0) lists in Section 7 - Engineering and Documentation "
         "to be Developed During the Assignment, page 29, inside the erection manual, "
         "this is what I need. None of it has been submitted:", False),
    ])
    blank(doc)
    add_bullet_lead(doc, "A lifting drawing.",
                    " Where the lifting points are, identified and dimensioned, with the "
                    "capacity of each and which ones each case uses. Today they exist "
                    "only as node numbers in a model.")
    add_bullet_lead(doc, "The yoke design,",
                    " consistent with the arrangement the model assumed, and who "
                    "supplies it.")
    add_bullet_lead(doc, "The lifting point design, by whichever route you build.",
                    " Either the padeye to those criteria, or, if the ISO corner "
                    "fittings are the route, written confirmation that they take the "
                    "calculated reactions with no reinforcement, and the basis for it. "
                    "The container datasheet is dated January 2026, before the module "
                    "was built inside it.")
    add_bullet_lead(doc, "The lifting calculation as a document of its own,",
                    " stating the lifting weight in kilograms.")
    blank(doc)

    # --- 6. por que tiene fecha propia: la grua es de ADASA
    add_segments(doc, [
        ("This has a date of its own because ADASA lifts the module. ", False),
        ("Section 9 - Commissioning and Start-up places the interconnection of the "
         "module and the anchoring of the container to its base within ADASA's scope, "
         "and states that this includes the lifting equipment required for the purpose, "
         "cranes among them.", True),
        (" I cannot contract a crane without the lifting weight, the center of gravity "
         "and the yoke design. I hold three weights that do not agree, 122.41 kN, "
         "13,300 kg and 32,500 kg, and ", False),
        ("none is the lifting weight.", True),
        (" Section 8.1 - Minimum Scope of Factory Acceptance Testing also puts the "
         "lifting points under dimensional verification, and that test opens within "
         "days.", False),
    ])
    blank(doc)

    # --- 7. el marco, sin imponer fecha
    add_segments(doc, [
        ("Section 7 gives ADASA seven working days to comment and states that the "
         "equipment cannot be released for transport to site without ADASA's approval "
         "of the documentation. I am asking for nothing beyond what that section sets "
         "out, and the review open this week can close it.", False),
    ])
    blank(doc)

    add_para(doc,
             "This email is sent without prejudice to ADASA's rights under the Contract.")
    blank(doc)

    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in ["Leader, Infrastructure Engineering",
                 "ADASA, Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)

    # Metadatos explicitos: sin esto Word deja author='python-docx'.
    cp = doc.core_properties
    cp.title = ("TALTAL - Lifting package for the module and the scope of the "
                "structural review")
    cp.author = "Luis Rivera Gonzalez"
    cp.last_modified_by = "Luis Rivera Gonzalez"
    cp.company = "Aguas Antofagasta"
    cp.category = "Correo"
    cp.comments = ("Pide ampliar al izaje el alcance de la revision estructural "
                   "encargada el 9 de septiembre, y reclama los tres entregables de "
                   "izaje de la ET Seccion 7 pagina 29. Cadena del Transmittal N39.")
    doc.save(OUTPUT_FILE)

    palabras = sum(len(p.text.split()) for p in doc.paragraphs)
    print("Correo generado: " + OUTPUT_FILE)
    print(f"  palabras del documento: {palabras}")


if __name__ == "__main__":
    crear_correo()
