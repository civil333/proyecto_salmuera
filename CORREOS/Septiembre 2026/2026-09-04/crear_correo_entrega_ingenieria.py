#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo de remision de la Nota Tecnica P22-NT-06-000-001-0 y del paquete
INGENIERIA VIGENTE PARA CONSTRUCCION al contratista adjudicado del Montaje
Mecanico y las Obras Civiles del Modulo de Segunda Etapa de Salmuera, Planta
Desaladora Taltal.

Es carta de cobertura: remite y fija el acuse, sin repetir la sustancia de la
nota. Registro de correspondencia, en primera persona.

Idioma del documento fijado en es-CL para que Word corrija en espanol.
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-09-04_Entrega-Ingenieria-Vigente-Construccion-Taltal.docx")
CONTACTO = "Luis Rivera"
LANG = "es-CL"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def aplicar_arial_table(table, size=10):
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(size)


def _set_lang_in_rPr(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma_documento(doc, lang=LANG):
    try:
        _set_lang_in_rPr(doc.styles["Normal"].element.get_or_add_rPr(), lang)
    except KeyError:
        pass
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang_in_rPr(r._element.get_or_add_rPr(), lang)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        _set_lang_in_rPr(r._element.get_or_add_rPr(), lang)


def add_table_with_header(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        c.text = h
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
    for i, fila in enumerate(rows, start=1):
        for j, v in enumerate(fila):
            t.rows[i].cells[j].text = str(v)
    aplicar_arial_table(t, size=10)
    return t


def parrafo(doc, texto, size=11, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold = bold
    aplicar_arial(p, size)
    return p


def crear_correo():
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(0.8)

    campos = [
        ("Fecha:", "4 de septiembre de 2026"),
        ("De:", f"{CONTACTO} — Dirección de Ingeniería y Obras, Aguas Antofagasta S.A."),
        ("Para:", "[COMPLETAR: contratista adjudicado]"),
        ("CC:", "Víctor Gutiérrez (ADASA)"),
        ("Asunto:", "PD Taltal — Nota Técnica P22-NT-06-000-001-0 e ingeniería vigente para construcción"),
        ("Ref:", "Contrato de Montaje Mecánico y Obras Civiles / BAE 12803"),
    ]
    for etiqueta, valor in campos:
        p = doc.add_paragraph()
        r1 = p.add_run(f"{etiqueta} ")
        r1.bold = True
        p.add_run(valor)
        aplicar_arial(p, 11)

    doc.add_paragraph()
    parrafo(doc, "Estimados señores:")
    doc.add_paragraph()

    parrafo(doc,
        "Junto con saludar, les remito la Nota Técnica P22-NT-06-000-001-0 y el paquete "
        "Ingeniería Vigente para Construcción, que reúne la ingeniería de detalle mecánica y "
        "de obras civiles en su revisión vigente al 4 de septiembre de 2026.")
    doc.add_paragraph()

    parrafo(doc,
        "La nota declara qué cambió respecto de la ingeniería con la que ustedes cotizaron y "
        "su efecto sobre las partidas del Formato de Presupuesto. Les pido revisarla antes de "
        "emitir órdenes de compra de accesorios y bridas y antes de iniciar las fundaciones.")
    doc.add_paragraph()

    parrafo(doc,
        "La nota es el punto de entrada: indica qué contiene cada carpeta y qué documentos "
        "cambiaron de revisión. La planilla de la carpeta 0 lista la revisión que rige para "
        "cada uno de los 54 documentos.")
    doc.add_paragraph()

    parrafo(doc,
        "Quedo atento a su acuse de recibo, que les pido a más tardar el viernes 11 de "
        "septiembre, y a cualquier consulta que surja de la revisión.")
    doc.add_paragraph()
    parrafo(doc, "Saludos cordiales,")
    doc.add_paragraph()
    parrafo(doc, CONTACTO)
    parrafo(doc, "Dirección de Ingeniería y Obras")
    parrafo(doc, "Aguas Antofagasta S.A. — Grupo EPM")

    doc.core_properties.author = "Luis Rivera Gonzalez"
    doc.core_properties.company = "Aguas Antofagasta"
    doc.core_properties.comments = "Correo de remision de la NT P22-NT-06-000-001-0 y del paquete de ingenieria vigente"
    fijar_idioma_documento(doc)
    doc.save(OUTPUT_FILE)
    print(f"OK: {os.path.basename(OUTPUT_FILE)}")


if __name__ == "__main__":
    crear_correo()
