#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo a L&A Ingenieria y Proyectos: pedido de re-emision del plano de Movimiento de
Tierra P22-DWG-00-001-001 (laminas 1 y 2) tras el cambio de nivel de la ENTREGA 12, y
consulta por el cuadro de excavacion de P22-DWG-00-002-003 LAM1.

Reply-To al hilo de la carta 067-032-032-COR-TT-013.
Idioma del documento fijado en es-CL.
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-09-04_LyA-Movimiento-Tierra-Re-emision.docx")
CONTACTO = "Luis Rivera"
LANG = "es-CL"


def aplicar_arial(p, size=11):
    for r in p.runs:
        r.font.name = "Arial"; r.font.size = Pt(size)


def aplicar_arial_table(t, size=10):
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"; r.font.size = Pt(size)


def _lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang"); rPr.append(el)
    el.set(qn("w:val"), lang); el.set(qn("w:eastAsia"), lang); el.set(qn("w:bidi"), lang)


def fijar_idioma_documento(doc, lang=LANG):
    try:
        _lang(doc.styles["Normal"].element.get_or_add_rPr(), lang)
    except KeyError:
        pass
    for p in doc.paragraphs:
        for r in p.runs:
            _lang(r._element.get_or_add_rPr(), lang)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        _lang(r._element.get_or_add_rPr(), lang)


def tabla(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]; c.text = h
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
    for i, fila in enumerate(rows, start=1):
        for j, v in enumerate(fila):
            t.rows[i].cells[j].text = str(v)
    aplicar_arial_table(t, 10)
    return t


def parrafo(doc, texto, size=11, bold=False):
    p = doc.add_paragraph(); r = p.add_run(texto); r.bold = bold
    aplicar_arial(p, size); return p


def vineta(doc, texto):
    p = doc.add_paragraph(style="List Paragraph")
    p.paragraph_format.left_indent = Inches(0.3)
    r = p.add_run("• " + texto)
    aplicar_arial(p, 11); return p


def crear_correo():
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(0.8)

    campos = [
        ("Fecha:", "4 de septiembre de 2026"),
        ("De:", f"{CONTACTO} — Aguas Antofagasta S.A. (lrivera@aguasantofagasta.cl)"),
        ("Para:", "Pablo Castillo (pcastillo@lyaingenieria.cl); Yohana Rodríguez Flores "
                  "(yrodriguez@aguasantofagasta.cl)"),
        ("CC:", "Nicolás Yanes, Cristian Sánchez, Jesús Alarcón, J. Cid (L&A); "
                "Dio Documentos, Luciano Méndez Huidobro, Víctor Gutiérrez Aqueveque, "
                "Jorge Guevara Lizana, Manuel Aguilera Orellana (ADASA)"),
        ("Asunto:", "ID Módulo RO 2da Etapa para Salmuera / TT-013 — Movimiento de tierra "
                    "pendiente de actualizar tras el cambio de nivel"),
        ("Ref:", "Carta 067-032-032-COR-TT-013 del 03-09-2026"),
    ]
    for et, val in campos:
        p = doc.add_paragraph(); r = p.add_run(f"{et} "); r.bold = True; p.add_run(val)
        aplicar_arial(p, 11)

    doc.add_paragraph()
    parrafo(doc, "Estimado Pablo:")
    doc.add_paragraph()

    parrafo(doc,
        "Recibimos la entrega del 3 de septiembre con las cinco láminas en revisión 1. Al revisarlas "
        "aparece un punto que queda abierto y que les pedimos cerrar: el plano de Movimiento de Tierra "
        "no acompañó a esa entrega y quedó referido al nivel anterior.")
    doc.add_paragraph()

    parrafo(doc, "El cambio de nivel", bold=True)
    parrafo(doc,
        "En la revisión 1 del plano P22-DWG-00-002-003 lámina 1, el sello de fundación del contenedor "
        "pasa de la cota +5,150 a la +5,400 y la cara superior de la +6,050 a la +6,300. El nivel de "
        "terreno natural se mantiene en +6,000, de modo que lo que cambia es la profundidad de "
        "excavación, que disminuye 250 milímetros.")
    doc.add_paragraph()

    parrafo(doc, "Lo que queda desactualizado", bold=True)
    parrafo(doc,
        "El plano P22-DWG-00-001-001, en revisión 0 del 22 de junio, sigue dimensionado sobre el nivel "
        "anterior en sus dos láminas:")
    doc.add_paragraph()

    tabla(doc,
        ["Lámina", "Dónde", "Qué dice hoy", "Por qué queda superado"],
        [["LAM1", "Cuadro Cubicaciones Movimiento Tierra, ítem 8",
          "Excavación contenedor 38,02 m³, área 53,51 m²",
          "El sello sube 250 mm sobre esa área"],
         ["LAM1", "Cuadro Cubicaciones Movimiento Tierra, ítem 7",
          "Excavación TK CIP y equipos 5,01 m³",
          "La revisión 1 del P22-DWG-00-002-007 LAM1 la baja a 1,60 m³"],
         ["LAM2", "Secciones E y F",
          "Fondo de excavación de la fundación del contenedor en EL. 5,15",
          "El sello vigente es +5,400"]])
    doc.add_paragraph()

    parrafo(doc, "Lo que les pedimos", bold=True)
    vineta(doc,
        "Re-emitir P22-DWG-00-001-001 láminas 1 y 2 en revisión 1, con el cuadro de cubicaciones y las "
        "secciones ajustados al sello vigente.")
    vineta(doc,
        "Confirmar o corregir el cuadro de excavación de P22-DWG-00-002-003 lámina 1. La revisión 1 lo "
        "mantiene en 38,02 metros cúbicos de excavación, 45,62 de retiro y 30,03 de relleno, los mismos "
        "valores de la revisión 0, pese a que en esa misma lámina el sello subió 250 milímetros. "
        "Entendemos que la excavación debería bajar del orden de 13 metros cúbicos y el relleno subir en "
        "proporción; la cifra la fija el cálculo de ustedes.")
    doc.add_paragraph()

    parrafo(doc,
        "El plano de Implantación General P22-DWG-00-002-001 lo revisamos y no requiere cambio: acota "
        "el nivel de terreno natural por zona, que no se movió.")
    doc.add_paragraph()

    parrafo(doc,
        "Un punto de forma, para la próxima emisión. En las cinco láminas de la revisión 1 la fila de "
        "esa revisión describe el cambio, “MODIFICACIONES INDICADAS”, y no declara el estado de "
        "emisión; la fila de la revisión 0 conserva “APTO PARA CONSTRUCCIÓN”. La carta de "
        "remisión sí las declara para construcción, pero quien construye lee el plano. Les pedimos que "
        "el cajetín declare el estado de emisión en la fila de la revisión vigente.")
    doc.add_paragraph()

    parrafo(doc,
        "Quedamos atentos a la fecha en que pueden emitir estas láminas.")
    doc.add_paragraph()
    parrafo(doc, "Saludos cordiales,")
    doc.add_paragraph()
    parrafo(doc, CONTACTO)
    parrafo(doc, "Departamento de Ingeniería y Optimización (DIO)")
    parrafo(doc, "Aguas Antofagasta S.A. — Grupo EPM")
    parrafo(doc, "lrivera@aguasantofagasta.cl")

    doc.core_properties.author = "Luis Rivera Gonzalez"
    doc.core_properties.company = "Aguas Antofagasta"
    doc.core_properties.comments = "Pedido de re-emision del plano de movimiento de tierra tras el cambio de nivel"
    fijar_idioma_documento(doc)
    doc.save(OUTPUT_FILE)
    print(f"OK: {os.path.basename(OUTPUT_FILE)}")


if __name__ == "__main__":
    crear_correo()
