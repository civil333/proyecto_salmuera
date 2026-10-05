#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
crear_correo_izaje_addendum.py
Respuesta a Eduardo Yamauchi sobre el addendum de izaje, en la misma cadena
(RE: TALTAL - Lifting package for the module and the scope of the structural
review), Reply-All sobre su correo del lunes 14-Sep a las 23:04.

VERSION EJECUTIVA (pedido del usuario, 15-Sep): cinco movimientos cortos.

1. Quien confirma: "Eng. Thomas" = Thomas Engineers (Tomas Avila). Que su
   alcance cubra el paquete completo y que manden su fecha de entrega.
2. El addendum no completa el paquete: su Figura 16 repite las reacciones de
   la Condicion 1 del Rev 0 (P22-CD-09-005-001, nodos 126-129 a 2,0 D, pag. 58-59
   de 548) y agrega W10x49 y un punto de gancho, sin cotas. ADASA fabrica el
   yugo en Chile y con ese documento no se construye.
3. La ET adjunta, Section 7 pag. 29, con los tres entregables del Manual de
   montaje y su cita literal ("Memoria de Calculo" SIN tilde, como el PDF
   original). La ET NO pide una "maniobra": no se exige.
4. El plano del yugo tiene que ser de fabricacion, y la memoria sigue debiendo
   la verificacion local de los puntos pedida el 11-Sep.
5. Paquete completo el jueves 17-Sep, ligado a la Section 7 pag. 30. Sin el
   calculo de los siete dias habiles (feriado del 18: cerraria el 29).

FUERA en esta version: la aritmetica del peso vacio (suma de la Figura 16
entre 2,0 = 100,80 kN), el codigo propio del addendum (el -003 es de los
Criterios de Diseno Rev 0), el ancho de eslinga y el izaje de prueba, y la
reserva de derechos, que sigue en el hilo desde el 11-Sep. BAE 43.1, Figura 15,
Rev A/B y "Page 2 of 2" van a la Seccion 3 del proximo transmittal.

Las citas en espanol van con idioma es-CL y tras coma, sin parentesis, para no
inflar la densidad de parentesis del registro de correspondencia.

ORTOGRAFIA: en-US. Ingles, primera persona. Document() directo, sin template ADASA.
"""
import os

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-09-15_Lifting-Addendum-Reply.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"
LANG_CITA = "es-CL"

RUNS_ES = []   # runs con texto en espanol (citas literales de la ET y nombre del adjunto)


def aplicar_arial(paragraph, size=11):
    for r in paragraph.runs:
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return paragraph


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)
    for r in RUNS_ES:
        _set_lang(r._element.get_or_add_rPr(), LANG_CITA)


def add_para(doc, text, size=11):
    p = doc.add_paragraph()
    p.add_run(text)
    return aplicar_arial(p, size)


def add_runs(doc, runs, size=11):
    """Parrafo por tramos (texto, negrita, espanol)."""
    p = doc.add_paragraph()
    for texto, bold, es in runs:
        r = p.add_run(texto)
        r.bold = bold
        if es:
            RUNS_ES.append(r)
    return aplicar_arial(p, size)


def add_bullet_cita(doc, texto, cita, size=11):
    """Vineta con el entregable en ingles y la cita literal de la ET en espanol."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.add_run("•  " + texto + ', "')
    RUNS_ES.append(p.add_run(cita + "."))
    p.add_run('"')
    return aplicar_arial(p, size)


def blank(doc):
    return doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "September 15, 2026", False),
        ("From:", CONTACTO + " - Leader, Infrastructure Engineering (ADASA)", False),
        ("To:", "Eduardo Yamauchi - BW Water", False),
        ("CC:", "Stephane Gehant; Jeryl F. Regulacion; Lokman Hakim Bin Mat; "
                "Magdier Arias; Mohd Adnin Bin Zulkaflee; Sadeep Irugalbandara; "
                "Nick Huta - BW Water; Victor Gutierrez - ADASA", False),
        ("Subject:", "RE: TALTAL - Lifting package for the module and the scope of "
                     "the structural review", False),
        ("Ref:", "Contract C-4300 / BAE 12803 / Technical Specification "
                 "P22-ET-09-000-001-0", False),
        ("Attachment:", "P22-ET-09-000-001-0 (ET Módulo).pdf", True),
    ]
    for label, value, es in fields:
        add_runs(doc, [(label, True, False), (" " + value, False, es)])
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    # --- 1. quien confirma: Thomas Engineers; que cubra el paquete completo y su fecha
    add_para(doc,
             "I understand the confirmation you are waiting for is Thomas Engineers' "
             "acceptance of the lifting scope. Please make sure it covers the complete "
             "package below, and send me their delivery date.")
    blank(doc)

    # --- 2. el addendum repite el Rev 0, sin cotas; con eso no se construye
    add_para(doc,
             "The addendum does not complete the package, because its Figure 16 repeats "
             "the Lifting Condition 1 reactions already in the Rev 0 Structural "
             "Calculation Report, P22-CD-09-005-001, and adds W10x49 members and a hook "
             "point with no dimensions. ADASA fabricates the yoke in Chile, and we cannot "
             "build it from this document.")
    blank(doc)

    # --- 3. la ET adjunta: seccion, pagina y los tres entregables literales
    add_runs(doc, [
        ("Section 7 - Engineering and Documentation to be Developed During the "
         "Assignment, page 29, of the attached Technical Specification "
         "(P22-ET-09-000-001-0) requires three documents inside the Plant Erection "
         "Manual:", False, False),
    ])
    add_bullet_cita(doc, "Lifting calculation",
                    "Memoria de Calculo para el izaje del módulo")
    add_bullet_cita(doc, "Lifting drawing with the lifting points and weights",
                    "Plano de izaje del módulo indicando los puntos de izaje, "
                    "indicando claramente pesos")
    add_bullet_cita(doc, "Drawing and design of the lifting yoke",
                    "Plano y diseño del yugo de Izaje para el módulo")
    blank(doc)

    # --- 4. plano de fabricacion del yugo y verificacion local pendiente
    add_para(doc,
             "The yoke drawing must be one we can fabricate from: dimensions, connections "
             "and welds, sling attachment points and the connection to the top corner "
             "fittings. The calculation still needs the local check of the lifting points "
             "requested on 11 September.")
    blank(doc)

    # --- 5. plazo jueves 17, ligado a la Section 7 pag. 30 (sin calculo de dias)
    add_runs(doc, [
        ("Please send the complete package by Thursday 17 September,", True, False),
        (" so that ADASA's review under Section 7, page 30, closes before the module "
         "is released for transport.", False, False),
    ])
    blank(doc)

    add_para(doc, "Best regards,")
    blank(doc)
    add_runs(doc, [(CONTACTO, True, False)])
    for line in ["Leader, Infrastructure Engineering",
                 "ADASA, Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)

    cp = doc.core_properties
    cp.title = ("RE: TALTAL - Lifting package for the module and the scope of the "
                "structural review")
    cp.author = "Luis Rivera Gonzalez"
    cp.last_modified_by = "Luis Rivera Gonzalez"
    cp.company = "Aguas Antofagasta"
    cp.category = "Correo"
    cp.comments = ("Respuesta ejecutiva del 15-Sep-2026 al addendum de izaje: repite el "
                   "Rev 0 sin cotas; ET Seccion 7 pag. 29 adjunta; plano de fabricacion "
                   "del yugo; paquete completo al jueves 17-Sep.")
    doc.save(OUTPUT_FILE)

    n_campos = len(fields)
    cuerpo = doc.paragraphs[n_campos + 2:-6]   # tras campos + blanco + "Dear", hasta la firma
    palabras_doc = sum(len(p.text.split()) for p in doc.paragraphs)
    palabras_cuerpo = sum(len(p.text.split()) for p in cuerpo)
    print("Correo generado: " + OUTPUT_FILE)
    print(f"  palabras del documento: {palabras_doc}")
    print(f"  palabras del cuerpo (sin encabezado ni firma): {palabras_cuerpo}")


if __name__ == "__main__":
    crear_correo()
