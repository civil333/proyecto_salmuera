#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
crear_correo_izaje_respuesta.py
Respuesta a Eduardo Yamauchi sobre el paquete de izaje, en la misma cadena del
correo del 10 de septiembre (RE: TALTAL - Lifting package for the module and
the scope of the structural review). BW Water contesto el 11-Sep a las 9:18.

LO QUE BW WATER ENTENDIO MAL, Y ES EL EJE DEL CORREO

Trata el yugo como un accesorio de transporte que resuelve "the forwarder /
rigging team". El yugo es para el izaje que ejecuta ADASA en Taltal, con grua
propia, para posar el modulo sobre su fundacion. Por eso la ET pide a BW Water
el DISENO y el PLANO del yugo, no el yugo: ADASA lo fabrica en Chile a partir
de ese diseno. La cadena de secciones: Section 3 (montaje en sitio por ADASA,
Ex Works) -> Section 9 (gruas alcance de ADASA) -> Section 7 pag. 29 (los tres
entregables dentro del manual de montaje, para aprobacion).

LO QUE SU RESPUESTA CONFIRMA, Y SE DEVUELVE COMO CONFIRMACION

- Padeye: declaran la ruta (esquineros ISO, sin padeye) y que FS 2,0 y 5 %
  lateral gobiernan "the lift points and members at the lift point". El Rev 0
  modela el izaje como global y reporta reacciones en nodos; NO verifica el
  estado tensional del punto. Con la excentricidad de la carga (CG a 5.486 /
  6.514 mm; tensiones 88,5 a 129,2 kN) esa verificacion local es lo que se
  pide, y validada por el profesional de la revision estructural.
- Yugo: admiten que el izaje superior necesita un lifting frame. Ese es el
  yugo. La base aprobada en el TM N25 fue la de su hoja de comentarios del
  24-Jun-2026: "spreader beams will be needed so that only vertical forces are
  induced at the lifting points". La Figura 15 del Rev 0 (verificada por
  render) lleva las cuatro eslingas inferiores a DOS puntos sobre el techo,
  uno por costado: el izaje inferior tambien depende de un arreglo de gancho
  que hay que construir, y su geometria fija las tensiones.
- Forwarder: por decision del usuario NO se menciona el "previous email" que
  citan (no existe en los registros). Se separan los dos izajes, se toma el
  punto de la revision del lifting plan, y se les pide quedar en contacto con
  el forwarder y compartirle el procedimiento de izaje como base.

FECHA FIJADA POR EL USUARIO: martes 15 de septiembre de 2026, dia en que
vencen los cuatro dias comprometidos el 9-Sep para el profesional.

PALANCAS EN EL CUERPO: solo Section 7 pag. 30 (no liberar para transporte sin
aprobacion) y Section 8.1 (verificacion dimensional de puntos de izaje en el
FAT). La BAE 43.1 a) y d) quedan en el PRG-49 y en la Seccion 3 del proximo
transmittal, no aqui: el correo persuade, el transmittal registra.

ORTOGRAFIA: en-US. Nada de centre, modelled, metre ni analyse.
Ingles, primera persona. Document() directo, sin template ADASA. ENVIADO el 11-Sep-2026.
"""
import os

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-09-11_Lifting-Package-Reply.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for r in paragraph.runs:
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return paragraph


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph()
    p.add_run(text)
    return aplicar_arial(p, size)


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for texto, bold in segs:
        r = p.add_run(texto)
        r.bold = bold
    return aplicar_arial(p, size)


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.add_run("\u2022  ")
    p.add_run(lead).bold = True
    p.add_run(rest)
    return aplicar_arial(p, size)


def blank(doc):
    return doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "September 11, 2026"),
        ("From:", CONTACTO + " - Leader, Infrastructure Engineering (ADASA)"),
        ("To:", "Eduardo Yamauchi - BW Water"),
        ("CC:", "Stephane Gehant; Jeryl F. Regulacion; Lokman Hakim Bin Mat; "
                "Magdier Arias; Mohd Adnin Bin Zulkaflee; Sadeep Irugalbandara; "
                "Nick Huta - BW Water; Victor Gutierrez - ADASA"),
        ("Subject:", "RE: TALTAL - Lifting package for the module and the scope of "
                     "the structural review"),
        ("Ref:", "Contract C-4300 / BAE 12803 / Technical Specification "
                 "P22-ET-09-000-001-0"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    # --- 1. quien iza, y por que la ET pide el diseno y no el yugo
    add_segments(doc, [
        ("Thank you for the replies. ", False),
        ("The lifting package is for the lift ADASA performs at Taltal with its own "
         "crane, to set the module on its foundation.", True),
        (" That is why the Technical Specification (P22-ET-09-000-001-0) asks BW "
         "Water for the design and drawing of the yoke and leaves the yoke itself "
         "out. Section 3 - Scope of the Assignment places site erection with ADASA "
         "under Ex Works delivery, and Section 9 - Commissioning and Start-up places "
         "the cranes on ADASA's side. Section 7 - Engineering and Documentation to be "
         "Developed During the Assignment, page 29, lists the lifting calculation, "
         "the lifting drawing with weights and the yoke design inside the erection "
         "manual.", False),
    ])
    blank(doc)

    add_para(doc, "My answers to your three replies:")
    blank(doc)

    # --- 2. padeye: la ruta declarada y la verificacion local del punto
    add_bullet_lead(doc, "Padeye.",
                    " Agreed, the module lifts from the ISO corner fittings and no "
                    "padeye is built. Your reply also says the criteria apply to the "
                    "lift points and the members at the lift point. The Rev 0 report "
                    "(P22-CD-09-005-001) models the lift globally and does not check "
                    "the corner fitting, the corner post and the members welded to "
                    "them under the reactions, at FS 2.0 with the 5% lateral load. The "
                    "load is eccentric (center of gravity 5,486 mm from one end, 6,514 "
                    "from the other, sling tensions 88.5 to 129.2 kN). That local check "
                    "is required, and it belongs in the professional's review.")
    blank(doc)

    # --- 3. yugo: el lifting frame que ellos nombran, y la base aprobada
    add_bullet_lead(doc, "Yoke.",
                    " You confirm that the top lift needs a lifting frame. That frame "
                    "is the yoke. Your comment sheet of 24 June 2026 stated that "
                    "spreader beams would keep the forces at the lifting points "
                    "vertical, and the criteria were approved on that basis. Figure 15 "
                    "brings the bottom slings to two points above the roof, so the "
                    "bottom lift also needs a built arrangement. Please designate the "
                    "site lift condition and give me the yoke or arrangement designed "
                    "for it, with the bottom corner fittings checked under the "
                    "horizontal component if the 30-degree lift is the one.")
    blank(doc)

    # --- 4. forwarder: los dos izajes separados, y la oferta tomada
    add_bullet_lead(doc, "Forwarder and rigging plan.",
                    " The forwarder moves the containerized plant module from your "
                    "factory to Taltal, and its handling in transit is the forwarder's "
                    "matter. Setting the module on its foundation is the site lift, and "
                    "that lift is ADASA's. The rigging contractor's method statement "
                    "selects crane, slings and shackles for it against your design. I "
                    "take up your offer to review that plan. "
                    "Please stay in contact with the forwarder and rigging team and "
                    "share the lifting procedure with them as the basis, once issued.")
    blank(doc)

    # --- 5. fecha fijada PARA QUE PUEDAN CUMPLIR, y es fecha ultima (usuario, 11-Sep)
    add_segments(doc, [
        ("I set the date so that you can meet it. Tuesday 15 September is the day "
         "the four days committed for the professional end. ", False),
        ("It is the latest date for the package, with the lifting covered by his "
         "review, and it cannot slip past that day.", True),
        (" Please confirm that scope.", False),
    ])
    blank(doc)

    # --- 6. las tres fuentes que enmarcan la fecha, cada una nombrada primero
    add_segments(doc, [
        ("Your comment sheet of 24 June to the Design Criteria committed the final "
         "weight, center of gravity and maximum load per lifting point once the "
         "design was finished. The Rev 0 Structural Calculation Report was issued "
         "for construction on 23 July. The Technical Specification, Section "
         "7, page 30, states that the equipment cannot be released for transport to "
         "site without ADASA's approval of this documentation. The Technical "
         "Specification, Section 8.1 - Minimum Scope of Factory Acceptance Testing, "
         "includes the lifting points in the dimensional verification at the FAT, "
         "which runs the week of the 15th.", False),
    ])
    blank(doc)

    add_para(doc,
             "This email is sent without prejudice to ADASA's rights under the Contract.")
    blank(doc)

    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in ["Leader, Infrastructure Engineering",
                 "ADASA, Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)

    cp = doc.core_properties
    cp.title = ("RE: TALTAL - Lifting package for the module and the scope of the "
                "structural review")
    cp.author = "Luis Rivera Gonzalez"
    cp.last_modified_by = "Luis Rivera Gonzalez"
    cp.company = "Aguas Antofagasta"
    cp.category = "Correo"
    cp.comments = ("Respuesta del 11-Sep-2026 a BW Water: el yugo es para el izaje "
                   "de ADASA en Taltal, no para el forwarder; verificacion local del "
                   "punto de izaje bajo la excentricidad; paquete el 15-Sep. Cadena "
                   "del correo del 10-Sep (Transmittal N39).")
    doc.save(OUTPUT_FILE)

    cuerpo = doc.paragraphs[8:-6]   # tras los 6 campos + blanco + "Dear", hasta la firma
    palabras_doc = sum(len(p.text.split()) for p in doc.paragraphs)
    palabras_cuerpo = sum(len(p.text.split()) for p in cuerpo)
    print("Correo generado: " + OUTPUT_FILE)
    print(f"  palabras del documento: {palabras_doc}")
    print(f"  palabras del cuerpo (sin encabezado ni firma): {palabras_cuerpo}")


if __name__ == "__main__":
    crear_correo()
