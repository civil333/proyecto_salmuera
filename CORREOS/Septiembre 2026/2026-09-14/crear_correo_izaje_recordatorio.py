#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
crear_correo_izaje_recordatorio.py
Recordatorio a Eduardo Yamauchi sobre el paquete de izaje, en la misma cadena
(RE: TALTAL - Lifting package for the module and the scope of the structural
review), Reply-All sobre el correo que ADASA envio el viernes 11-Sep.

POR QUE SALE HOY, LUNES 14

BW Water no respondio nada al correo del 11-Sep (confirmado por el usuario) y
la fecha ultima que ese correo fijo para el paquete es el martes 15-Sep. El
recordatorio llega antes del vencimiento, de modo que refuerza la fecha en vez
de constatar el incumplimiento.

QUE PIDE, Y POR QUE NO ES UN "PLEASE REPLY" GENERICO

Las dos cosas concretas que el correo del 11-Sep pidio y no tienen respuesta:
que la revision del profesional cubre el izaje del modulo, y la condicion de
izaje en sitio que designan con el yugo o arreglo disenado para ella. Mas el
paquete en la fecha.

LA REUNION ENTRA CON PROPOSITO

No ha llegado el enlace de la reunion semanal de coordinacion. Esta semana
tiene que ser el miercoles 16-Sep a la hora habitual (decision del usuario: dia
y misma hora de siempre, sin escribir la hora). Como el paquete vence el dia
antes, se pide que vaya en la agenda.

NO REABRE SUSTANCIA: sin secciones de la ET, sin BAE 43.1, sin Figura 15, sin
cifras, sin reserva de derechos nueva (ya viaja en el hilo).

ORTOGRAFIA: en-US. Ingles, primera persona. Document() directo, sin template ADASA.
"""
import os

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-09-14_Lifting-Package-Reminder.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for r in paragraph.runs:
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return paragraph


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph()
    p.add_run(text)
    return aplicar_arial(p, size)


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for texto, bold in segs:
        r = p.add_run(texto)
        r.bold = bold
    return aplicar_arial(p, size)


def blank(doc):
    return doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "September 14, 2026"),
        ("From:", CONTACTO + " - Leader, Infrastructure Engineering (ADASA)"),
        ("To:", "Eduardo Yamauchi - BW Water"),
        ("CC:", "Stephane Gehant; Jeryl F. Regulacion; Lokman Hakim Bin Mat; "
                "Magdier Arias; Mohd Adnin Bin Zulkaflee; Sadeep Irugalbandara; "
                "Nick Huta - BW Water; Victor Gutierrez - ADASA"),
        ("Subject:", "RE: TALTAL - Lifting package for the module and the scope of "
                     "the structural review"),
        ("Ref:", "Contract C-4300 / BAE 12803 / Technical Specification "
                 "P22-ET-09-000-001-0"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    # --- 1. sin respuesta al 11-Sep, la fecha ultima, y las dos confirmaciones
    add_segments(doc, [
        ("I have not had a reply to my email of Friday 11 September on the lifting "
         "package. ", False),
        ("Tomorrow, Tuesday 15 September, is the latest date for the package.", True),
        (" Two points from that email are still open on your side: confirmation that "
         "the professional's review covers the lifting of the module, and the site "
         "lift condition you designate, with the yoke or arrangement designed for "
         "it. Please confirm both before then.", False),
    ])
    blank(doc)

    # --- 2. enlace de la reunion: miercoles 16, hora habitual, paquete en agenda
    add_para(doc,
             "I have also not received the link for this week's coordination meeting, "
             "which needs to be on Wednesday 16 September at the usual time. Please "
             "send the invitation and put the lifting package on the agenda.")
    blank(doc)

    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in ["Leader, Infrastructure Engineering",
                 "ADASA, Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)

    cp = doc.core_properties
    cp.title = ("RE: TALTAL - Lifting package for the module and the scope of the "
                "structural review")
    cp.author = "Luis Rivera Gonzalez"
    cp.last_modified_by = "Luis Rivera Gonzalez"
    cp.company = "Aguas Antofagasta"
    cp.category = "Correo"
    cp.comments = ("Recordatorio del 14-Sep-2026 a BW Water: sin respuesta al correo "
                   "del 11-Sep; paquete de izaje el 15-Sep; confirmar alcance del "
                   "profesional y condicion de izaje con su yugo; enlace de la reunion "
                   "del miercoles 16-Sep a la hora habitual.")
    doc.save(OUTPUT_FILE)

    cuerpo = doc.paragraphs[8:-6]   # tras los 6 campos + blanco + "Dear", hasta la firma
    palabras_doc = sum(len(p.text.split()) for p in doc.paragraphs)
    palabras_cuerpo = sum(len(p.text.split()) for p in cuerpo)
    print("Correo generado: " + OUTPUT_FILE)
    print(f"  palabras del documento: {palabras_doc}")
    print(f"  palabras del cuerpo (sin encabezado ni firma): {palabras_cuerpo}")


if __name__ == "__main__":
    crear_correo()
