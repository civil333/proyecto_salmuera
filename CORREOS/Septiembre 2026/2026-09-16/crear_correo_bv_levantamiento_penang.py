#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
crear_correo_bv_levantamiento_penang.py
Correo a Bureau Veritas Chile (Jaime Martinez, administrador de contrato) con
copia a los coordinadores de BW Water en Penang y a Victor Gutierrez. Hilo nuevo.

QUE HACE

1. Informa lo que ADASA supo en la reunion del 16-Sep: desfase de unos 50 mm
   entre los turbos y sus spools, spools a cortar, resoldar y reensayar, listo
   para despacho al 12-Oct, y por las fotos el turbo interetapa.
2. Mantiene la prueba de alta del 17-Sep (Request to witness inspection 008),
   sin las lineas conectadas a los turbos, igual que el correo a Eduardo del 16.
3. Pide a Bureau Veritas un levantamiento desde la visita del 18: causa raiz
   del desfase (ITP 4.2), equipos y conformidad de los turbos con sus hojas de
   datos Rev 0 (4.1), instrumentos (4.3 y 4.4) y otros temas constructivos.
4. Pide un informe aparte del avance electrico (ITP 6.1 a 6.3).

DECISIONES DEL USUARIO (16-Sep)

- Reemplaza el borrador del 17 sobre el Request 008, que se elimino.
- Destinatario: el usuario lo llamo "Javier"; su firma es Jaime Andres Martinez
  Sanchez, Administrador de contrato, CESMEC-Bureau Veritas
  (martinez.jaime@bureauveritas.com). En su correo del 11-Sep copio a Carlo
  Montecinos y Luis Rodrigo Arcila.
- Ingles. CC Magdier Arias, Eduardo Yamauchi y Lokman Hakim Bin Mat (coordinan
  en BW Water Penang) y Victor Gutierrez.
- La prueba del 17 se mantiene sin las lineas de los turbos.
- El levantamiento parte el 18 y Bureau Veritas propone jornadas adicionales
  dentro de las contratadas, con aprobacion de ADASA.
- Adjuntos: solo lo esencial para entender el caso, no todos los planos.

EXCEPCION DE CORREO CONJUNTO: se nombra a BW Water porque esta en copia (ver
feedback_bv_no_nombrar_bw_water). Sin oferta, tarifas ni saldo de jornadas.

ORTOGRAFIA: en-US. Ingles, primera persona. Document() directo, sin template ADASA.
"""
import os

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-09-16_BV-Survey-Turbocharger-Offset.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"
SUBJECT = "25007 TALTAL - Turbocharger spool offset and inspection scope in Penang"

# Version ejecutiva (usuario, 16-Sep, pasada anti-ia en voz de Luis): cada item
# nombra que se levanta y su fila del ITP; los documentos van en los adjuntos.
SURVEY = [
    [("1. Root cause of the spool offset, with offsets, photographs and spool "
      "identification recorded before any cut (Inspection and Test Plan "
      "P22-BA-09-000-004 Rev 0, row 4.2).", False)],
    [("2. Equipment against the Equipment List Rev 0, including both turbochargers "
      "against their datasheets Rev 0 (row 4.1).", False)],
    [("3. Instruments against the Instrument List Rev F and the P&ID Rev 0 (rows 4.3 "
      "and 4.4).", False)],
    [("4. Any other construction issue found now that all the equipment is reported "
      "at the workshop.", False)],
]


def aplicar_arial(paragraph, size=11):
    for r in paragraph.runs:
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return paragraph


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    el.set(qn("w:bidi"), lang)


def fijar_idioma(doc, lang=LANG):
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph()
    p.add_run(text)
    return aplicar_arial(p, size)


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for texto, bold in segs:
        r = p.add_run(texto)
        r.bold = bold
    return aplicar_arial(p, size)


def blank(doc):
    return doc.add_paragraph()


def crear_correo():
    doc = Document()
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "September 16, 2026"),
        ("From:", CONTACTO + " - Leader, Infrastructure Engineering (ADASA)"),
        ("To:", "Jaime Martínez - Bureau Veritas"),
        ("CC:", "Carlo Montecinos, Luis Rodrigo Arcila - Bureau Veritas; Magdier "
                "Arias, Eduardo Yamauchi, Lokman Hakim Bin Mat - BW Water; Victor "
                "Gutierrez - ADASA"),
        ("Subject:", SUBJECT),
        ("Attachments:", "ADASA email to BW Water (16 September); BW Water weekly "
                         "report and Progress Report Week 37; P&ID Rev 0; Line List "
                         "Rev 1; Equipment List Rev 0; Instrument List Rev F; "
                         "turbocharger datasheets Rev 0"),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(" " + value)
        aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Jaime,")
    blank(doc)

    # --- 1. lo que se supo hoy, en primera persona (registro 11.1)
    add_para(doc,
             "BW Water informed us today of an offset of about 50 mm between the "
             "turbocharger connections and their piping spools. The affected spools have to "
             "be cut, re-welded and re-tested, and ready to ship moves to 12 October. From "
             "the photographs I understand it is the interstage turbocharger, the one "
             "with the highest operating pressure in the module. I attach the background "
             "documents, including my email to BW Water and its latest weekly report.")
    blank(doc)

    # --- 2. la prueba del 17 se mantiene sin los turbos
    add_para(doc,
             "The high-pressure test of 17 September (Request 008) goes ahead, "
             "excluding the lines connected to the turbochargers.")
    blank(doc)

    # --- 3. el levantamiento
    add_segments(doc, [
        ("Starting with the visit of 18 September, I ask Bureau Veritas for a survey "
         "at the Penang workshop covering:", True),
    ])
    for segs in SURVEY:
        add_segments(doc, segs)
    blank(doc)

    # --- 4. informes y jornadas en un solo parrafo
    add_segments(doc, [
        ("Please issue the survey in one inspection report and the electrical progress "
         "(rows 6.1 to 6.3) in a ", False),
        ("separate report", True),
        (", with a Flash Report for any deviation. Please also propose the inspection "
         "days needed, for our approval before mobilizing beyond 18 September.", False),
    ])
    blank(doc)

    # --- 5. BW Water y traspaso
    add_para(doc,
             "Magdier, Eduardo and Lokman, please give the inspector access to the "
             "equipment, spools and records.")
    blank(doc)
    # Ausencia declarada de frente (usuario, 16-Sep): "From tomorrow Victor leads"
    # sonaba a retirarse del tema. Fuera de la oficina dos semanas, segun la reunion.
    add_para(doc,
             "I will be out of the office for the next two weeks. In the meantime, Victor "
             "Gutierrez will handle communications for ADASA, so please send the reports "
             "to both of us.")
    blank(doc)

    # --- cierre del registro 11.1 ("Quedo atento")
    # Agradecimiento por la gestion (usuario, 16-Sep).
    add_para(doc, "Thank you for coordinating this. I look forward to your confirmation.")
    blank(doc)
    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph()
    p.add_run(CONTACTO).bold = True
    aplicar_arial(p)
    for line in ["Leader, Infrastructure Engineering",
                 "ADASA, Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)

    cp = doc.core_properties
    cp.title = SUBJECT
    cp.author = "Luis Rivera Gonzalez"
    cp.last_modified_by = "Luis Rivera Gonzalez"
    cp.company = "Aguas Antofagasta"
    cp.category = "Correo"
    cp.comments = ("Correo del 16-Sep-2026 a Bureau Veritas: desfase de spools del turbo "
                   "interetapa, prueba del 17 sin turbos, levantamiento desde el 18 e "
                   "informe electrico aparte.")
    doc.save(OUTPUT_FILE)

    cuerpo = doc.paragraphs[8:-8]
    print("Correo generado: " + OUTPUT_FILE)
    print(f"  palabras del cuerpo (sin encabezado ni firma): "
          f"{sum(len(p.text.split()) for p in cuerpo)}")


if __name__ == "__main__":
    crear_correo()
