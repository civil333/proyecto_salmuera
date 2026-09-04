#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Monday Commitments Follow-up — 22-Jun-2026.
Reply a "RE: 25007 Taltal: Meeting Notes" (hilo minuta 16-Jun + nuestro correo 17-Jun).
Reclamo firme y acotado de los compromisos que vencian este lunes 22-Jun y no llegaron
(spare parts quotes + ex-work shipment details) + estado de los compromisos del reporte
semanal. Fedco fuera (en memoria/seguimiento). Ingles (BW Water), Document() directo.
Cuerpo = fuente unica `2026-06-22_Monday-Commitments-Follow-up.md` (ya pasado por anti-ia, VERDE).
"""

import os
from docx import Document
from docx.shared import Pt, Inches

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-06-22_Monday-Commitments-Follow-up.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_para(doc, text, size=11):
    para = doc.add_paragraph(text)
    aplicar_arial(para, size)
    return para


def add_num_item(doc, n, lead, rest):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.25)
    para.add_run(f"{n}. ")
    para.add_run(lead).bold = True
    para.add_run(rest)
    aplicar_arial(para)
    return para


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # HEADER
    fields = [
        ("Date:", "June 22, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        (
            "To:",
            "Eduardo Yamauchi, Stephane Gehant, Adzlan Bin Abd Rahim, "
            "Lokman Hakim Bin Mat, Jeryl F. Regulacion, Sadeep Irugalbandara, "
            "Nick Huta — BW Water Americas Inc.",
        ),
        ("CC:", "Victor Gutierrez — ADASA"),
        ("Subject:", "RE: 25007 Taltal: Meeting Notes"),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / Meeting Notes 16-Jun-2026 / "
            "ADASA note 17-Jun-2026",
        ),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    # BODY (fuente: 2026-06-22_Monday-Commitments-Follow-up.md, anti-ia VERDE)
    add_para(doc, "Dear Eduardo,")
    doc.add_paragraph()

    add_para(
        doc,
        "Thank you for the 22-Jun status package (procurement tracker, "
        "Progress Report Week 25 and DDSR).",
    )
    doc.add_paragraph()

    add_para(
        doc,
        "Several deliverables from your 16-Jun meeting notes are missing from "
        "today's package, two of them committed for \"next Monday\":",
    )
    add_num_item(doc, 1, "Spare parts quotes", ": not included.")
    add_num_item(
        doc, 2, "Ex-work shipment details",
        ": not included. This is time-sensitive, as the RO pressure vessels "
        "reach ex-works Spain tomorrow (23-Jun) and the Shipping Plan "
        "(P22-BA-09-000-002) is still Not submitted.",
    )
    add_num_item(
        doc, 3, "Invoice and payment status update",
        ": still pending. Your 16-Jun notes committed BWW to report on "
        "invoicing; we need it to reconcile the payment milestones.",
    )
    doc.add_paragraph()

    add_para(
        doc,
        "Please explain why the Monday items lapsed and provide all three "
        "without further delay.",
    )
    doc.add_paragraph()

    add_para(
        doc,
        "On the weekly report, the photographic records and the super duplex "
        "piping line now appear, as we requested on 17-Jun. Two report "
        "commitments remain open. First, a fabrication schedule with actual "
        "dates: the report still shows progress percentages, not a dated "
        "schedule. Second, the \"material received status at workshop\" column "
        "carried into the Project Schedule itself; today it lives only in the "
        "progress report. Please incorporate both in the schedule re-issue.",
    )
    doc.add_paragraph()

    add_para(
        doc,
        "We would appreciate a written explanation for the missed Monday "
        "deliverables and a firm date for each open item, ahead of the weekly "
        "meeting.",
    )
    doc.add_paragraph()

    add_para(doc, "Best regards,")
    doc.add_paragraph()
    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        add_para(doc, line)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
