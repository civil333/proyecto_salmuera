#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo ADASA -> BW Water: follow-up del Recovery Schedule vencido (04-Jun) +
confirmacion del waiver ASME enviado el 02-Jun.

El Recovery Schedule actualizado vencia hoy y es deliverable de BW Water (ADASA
lo revisa, no lo produce). Sigue pendiente -> se persigue, anclado a la logica
del waiver: ADASA renuncio a la estampa para recuperar ~6 semanas, asi que el
schedule construido sobre la fecha sin estampa del 22-Jun es la ejecucion
natural de esa concesion. La parte ASME se confirma y reafirma SIN re-abrir las
tres condiciones del 02-Jun (no diluir el framing condicionado).

- Reply-To al thread del Mitigation Plan (subject "RE: 25007 Taltal - Mitigation Plan").
- Idioma: ingles (BW Water). Patron Document() directo per CLAUDE.md 3.4.
- Fuente unica del cuerpo: 2026-06-04_Recovery-Schedule-Request-and-ASME-Confirmation_Descripcion.md
"""

import os

from docx import Document
from docx.shared import Inches, Pt

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(
    SCRIPT_DIR, "2026-06-04_Recovery-Schedule-Request-and-ASME-Confirmation.docx"
)
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_para(doc, text, size=11):
    para = doc.add_paragraph(text)
    aplicar_arial(para, size)
    return para


def add_segments(doc, segments, size=11):
    """Parrafo con segmentos (texto, bold?). segments = [(text, bool), ...]."""
    para = doc.add_paragraph()
    for text, bold in segments:
        run = para.add_run(text)
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(size)
    return para


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # =========================================================================
    # HEADER
    # =========================================================================
    fields = [
        ("Date:", "June 4, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Andrew Sia, Victor Gutierrez, Jeryl F. Regulacion, "
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Mauricio Vallejos, "
            "Jorge Valdes, Tanya Figueroa, Allan Valentos, Sadeep Irugalbandara, "
            "Ghazi Ozair, Nick Huta, Marjan Arsovic, Gerald Ross, "
            "Adzlan Bin Abd Rahim, Stephane Gehant, Shane Banks, Fadey Kassim",
        ),
        ("Subject:", "RE: 25007 Taltal - Mitigation Plan"),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / "
            "Recovery Schedule due 04-Jun-2026 / "
            "ASME Stamp Waiver 02-Jun-2026 / "
            "Protec Arisawa Europe letter / "
            "Technical Note P22-NT-09-000-001-0",
        ),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    blank(doc)

    # =========================================================================
    # CUERPO
    # =========================================================================
    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_segments(
        doc,
        [
            ("The updated recovery schedule was due today, 04 June, and we have not "
             "received it. Please send it at your earliest convenience. As agreed, "
             "it should be built on the non-stamped ", False),
            ("22 June", True),
            (" date for the RO pressure vessels and state whether 22 June is "
             "ex-works Spain or delivered at Penang, with the transit time, so the "
             "downstream and EXW-Penang dates are traceable.", False),
        ],
    )
    blank(doc)

    add_para(
        doc,
        "We also want to confirm you received our message of 02 June waiving the "
        "ASME Section X stamp for these vessels. Our position there stands: ADASA "
        "accepts the non-stamped version on the terms set out in that message, and "
        "we waived the stamp specifically to recover the time. Please confirm "
        "receipt, and that the recovery schedule and the updated ITP reflect it.",
    )
    blank(doc)

    add_para(doc, "We look forward to your comments.")
    blank(doc)

    # =========================================================================
    # CIERRE
    # =========================================================================
    add_para(doc, "Best regards,")
    blank(doc)

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        add_para(doc, line)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
