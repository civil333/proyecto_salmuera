#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo de consulta tecnica a Van Doorn Ingenieria y Consultoria (autor de la
ingenieria de detalle mecanica, Cuadernillo de Isometrias Compilado Rev 0) por
la inconsistencia de material en el buje de reduccion NPT de la toma de
instrumento del plano P22-DWG-06-006-005: la descripcion dice AISI 316 mientras
la columna SPEC de la misma fila dice SAF2507. Ocurre identico en las laminas
H.1 (item 3) y H.3 (item 2). ADASA declara que rige SAF2507 (servicio salmuera)
y pide corregir la descripcion en ambas laminas.

Correo en espanol, idioma fijado es-CL (CLAUDE.md Seccion 3.4). Document() directo,
sin template ADASA. Estado: BORRADOR.
"""

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-06-11_Consulta-VanDoorn-Buje-2507.docx")
CONTACTO = "Luis Rivera"
LANG = "es-CL"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang_in_rPr(rPr, lang):
    lang_el = rPr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rPr.append(lang_el)
    lang_el.set(qn("w:val"), lang)
    lang_el.set(qn("w:eastAsia"), lang)
    lang_el.set(qn("w:bidi"), lang)


def fijar_idioma_documento(doc, lang=LANG):
    try:
        rPr = doc.styles["Normal"].element.get_or_add_rPr()
        _set_lang_in_rPr(rPr, lang)
    except KeyError:
        pass
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            _set_lang_in_rPr(run._element.get_or_add_rPr(), lang)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_lang_in_rPr(run._element.get_or_add_rPr(), lang)


def add_bullet(doc, lead, cuerpo):
    para = doc.add_paragraph(style="List Bullet")
    para.add_run(lead).bold = True
    para.add_run(cuerpo)
    aplicar_arial(para)
    return para


def aplicar_arial_table(table, size=10):
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(size)


def add_table_with_header(doc, headers, rows, widths=None):
    """Tabla con header en negrita y bordes (style Table Grid), Arial 10."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = header
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True

    for i, row_data in enumerate(rows, start=1):
        for j, value in enumerate(row_data):
            table.rows[i].cells[j].text = str(value)

    if widths:
        for j, w in enumerate(widths):
            for row in table.rows:
                row.cells[j].width = w

    aplicar_arial_table(table, size=10)
    return table


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    fields = [
        ("Fecha:", "11 de junio de 2026"),
        ("De:", f"{CONTACTO} — DIO ADASA (Aguas de Antofagasta S.A.)"),
        ("Para:", "[Contacto Van Doorn] — Van Doorn Ingeniería y Consultoría"),
        ("CC:", "Víctor Gutiérrez (ADASA)"),
        ("Asunto:",
         "Consulta — Material del buje NPT en conexión de instrumento "
         "(P22-DWG-06-006-005): AISI 316 vs SAF2507"),
        ("Ref:",
         "BAE 12803 / Ingeniería de Detalle Mecánica — Cuadernillo de Isometrías "
         "(Compilado Rev 0) / Plano P22-DWG-06-006-005-0"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Estimados:")
    aplicar_arial(para)
    doc.add_paragraph()

    # Apertura + hallazgo
    para = doc.add_paragraph(
        "Al revisar el Cuadernillo de Isometrías (Compilado Rev 0) encontramos "
        "una inconsistencia de material en el buje de reducción de la toma de "
        "instrumento del plano P22-DWG-06-006-005. El buje (BUJE DE REDUCCIÓN, "
        "CABEZA HEXAGONAL, THDM NPT, 1\"×1/2\", CL3000, ASME B16.11) figura "
        "descrito como ")
    para.add_run("AISI 316").bold = True
    para.add_run(", mientras que la columna de especificación (SPEC) de esa misma "
                 "fila lo asigna a ")
    para.add_run("SAF2507").bold = True
    para.add_run(" (Super Duplex). La incoherencia se repite idéntica en dos "
                 "láminas del mismo plano; los ítems del cuadro LISTA DE "
                 "MATERIALES afectados son:")
    aplicar_arial(para)

    doc.add_paragraph()

    add_table_with_header(
        doc,
        ["Plano", "Lámina", "Ítem", "Material en descripción", "Columna SPEC"],
        [
            ["P22-DWG-06-006-005-0 (SA-HDPE-DN110-PN10-005)", "H.1", "3",
             "AISI 316", "SAF2507"],
            ["P22-DWG-06-006-005-0 (SA-HDPE-DN110-PN10-005)", "H.3", "2",
             "AISI 316", "SAF2507"],
        ],
        widths=[Inches(2.9), Inches(0.7), Inches(0.55), Inches(1.5), Inches(1.1)],
    )

    doc.add_paragraph()

    # Posicion ADASA
    para = doc.add_paragraph(
        "Para el servicio de salmuera de rechazo concentrado el material que rige "
        "es el ")
    para.add_run("Super Duplex SAF2507 (UNS S32750)").bold = True
    para.add_run(
        "; el AISI 316 es susceptible de picado por cloruro y no es apto. El propio "
        "plano lo respalda: el codo de transición inmediatamente aguas arriba del "
        "buje (PE100 × Super Duplex, ítem 6 en H.1 e ítem 4 en H.3) ya está definido "
        "en Super Duplex. Entendemos, por tanto, que el valor correcto es SAF2507 y "
        "que la mención AISI 316 en la descripción corresponde a un arrastre de una "
        "librería genérica de fittings.")
    aplicar_arial(para)
    doc.add_paragraph()

    # Solicitud
    para = doc.add_paragraph(
        "Solicitamos corregir la descripción del buje a Super Duplex SAF2507 "
        "(UNS S32750) en ambas láminas e incorporar el cambio en la próxima "
        "revisión del cuadernillo. Si la intención de diseño hubiese sido "
        "efectivamente AISI 316, agradecemos justificarla a la luz del servicio.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Quedamos atentos a sus comentarios.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Saludos cordiales,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in [
        "Departamento de Ingeniería y Optimización (DIO)",
        "ADASA — Aguas de Antofagasta S.A.",
        "luis.rivera@adasa.cl",
    ]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    fijar_idioma_documento(doc, LANG)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")
    print(f"Idioma del documento: {LANG}")


if __name__ == "__main__":
    crear_correo()
