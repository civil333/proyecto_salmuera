#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Weekly Coordination Call minutes follow-up - 25-Jun-2026.
Reply-All al thread "25007 Project Taltal - Weekly Coordination Call" (call martes 23-Jun).
Insiste por la minuta del 23-Jun (sin respuesta) y deja por escrito los dos compromisos
Fedco asumidos en esa reunion (cronograma ajustado + reporte oficial del retraso), con
fecha firme (viernes 26-Jun). Foco: solo minuta + 2 Fedco (los pendientes del 22-Jun
viven en su propia cadena). Ingles (BW Water), Document() directo.
Cuerpo = fuente unica `2026-06-25_Weekly-Call-Minutes-Followup.md`.
"""

import os
import sys
from docx import Document
from docx.shared import Pt, Inches

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
try:
    from docx_metadata import apply_core_properties, fix_app_xml
    _HAS_META = True
except Exception:
    _HAS_META = False

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-06-25_Weekly-Call-Minutes-Followup.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_para(doc, text, size=11):
    para = doc.add_paragraph(text)
    aplicar_arial(para, size)
    return para


def add_num_item(doc, n, lead, rest):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.25)
    para.add_run(f"{n}. ")
    para.add_run(lead).bold = True
    para.add_run(rest)
    aplicar_arial(para)
    return para


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # HEADER
    fields = [
        ("Date:", "June 25, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        (
            "To:",
            "Eduardo Yamauchi, Jeryl F. Regulacion, Adzlan Bin Abd Rahim, "
            "Nick Huta, Sadeep Irugalbandara, Stephane Gehant, "
            "Lokman Hakim Bin Mat — BW Water Americas Inc.",
        ),
        ("CC:", "Victor Gutierrez — ADASA"),
        ("Subject:", "RE: 25007 Project Taltal - Weekly Coordination Call"),
        ("Ref:", "Contract C-4300 / BAE 12803 / Weekly Coordination Call 23-Jun-2026"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    # BODY (fuente: 2026-06-25_Weekly-Call-Minutes-Followup.md)
    add_para(doc, "Dear Eduardo,")
    doc.add_paragraph()

    add_para(
        doc,
        "Following my note of 23-Jun, we are still awaiting the minutes of "
        "Tuesday's coordination call. Please issue them by Friday, 26-Jun.",
    )
    doc.add_paragraph()

    add_para(
        doc,
        "We need them on record because two of the commitments made in that "
        "meeting are time-critical, and I want to confirm them here so they are "
        "not lost:",
    )
    add_num_item(
        doc, 1, "Revised schedule reflecting the Fedco delay",
        ": the re-sequenced Project Schedule that absorbs the slip on the HP "
        "pump and the two turbochargers (BH-09-001, SIP-09-001, SIP-09-002).",
    )
    add_num_item(
        doc, 2, "Formal report on the Fedco delay",
        ": the report, in report format, setting out the root cause and the "
        "sequence of events behind the delivery slippage, together with the "
        "recovery actions.",
    )
    doc.add_paragraph()

    add_para(
        doc,
        "Please confirm the committed issue date for each of these two "
        "deliverables by Friday, 26-Jun, and ensure the minutes capture them "
        "together with the remaining action items from the call.",
    )
    doc.add_paragraph()

    add_para(doc, "Best regards,")
    doc.add_paragraph()
    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        add_para(doc, line)

    if _HAS_META:
        apply_core_properties(
            doc,
            title="Weekly Coordination Call - Minutes Follow-up",
            author="Luis Rivera",
            subject="25007 Project Taltal - Weekly Coordination Call",
            category="Email",
            comments="Aguas de Antofagasta S.A. - Proyecto Taltal BAE 12803",
            language="en-US",
            revision=1,
        )

    doc.save(OUTPUT_FILE)

    if _HAS_META:
        fix_app_xml(OUTPUT_FILE, application="Microsoft Office Word",
                    company="Aguas de Antofagasta S.A.")

    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
