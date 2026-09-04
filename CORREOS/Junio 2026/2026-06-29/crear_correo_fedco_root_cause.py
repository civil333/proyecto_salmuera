#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Fedco delay - formal root-cause report still outstanding. 29-Jun-2026.
Reply-All al thread "25007 Taltal - Schedule Update" (correo de Yamauchi del 26-Jun).
ADASA pidio el 25-Jun (compromiso de la call 23-Jun) DOS entregables: (1) cronograma
re-secuenciado y (2) reporte formal de causas. El 26-Jun llego solo el cronograma (que
PROPAGA el slip 15-Ago->10-Sep) + una narrativa en el cuerpo del correo; el reporte
formal sigue pendiente. La narrativa atribuye la causa al "delayed down payment 09-Jun".
Postura del usuario: firme + reservar posicion (referencia medida al regimen de atraso);
sobre la atribucion al down-payment, PEDIR SUSTANCIACION sin rechazar aun. Exige el reporte
formal por viernes 03-Jul con 4 contenidos. Senala 02-Sep > 21-Ago rechazado el 17-Jun +
conflicto FAT/running-test (NT-001). Ingles (BW Water), Document() directo, metadatos limpios.
"""

import os
import sys
from docx import Document
from docx.shared import Pt, Inches

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
try:
    from docx_metadata import apply_core_properties, fix_app_xml
    _HAS_META = True
except Exception:
    _HAS_META = False

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-06-29_Fedco-Root-Cause-Report-Request.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_para(doc, text, size=11):
    para = doc.add_paragraph(text)
    aplicar_arial(para, size)
    return para


def add_num_item(doc, n, lead, rest):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.25)
    para.add_run(f"{n}. ")
    para.add_run(lead).bold = True
    para.add_run(rest)
    aplicar_arial(para)
    return para


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    fields = [
        ("Date:", "June 29, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        (
            "To:",
            "Eduardo Yamauchi, Stephane Gehant, Adzlan Bin Abd Rahim, "
            "Jeryl F. Regulacion, Lokman Hakim Bin Mat, Sadeep Irugalbandara, "
            "Nick Huta — BW Water Americas Inc.",
        ),
        ("CC:", "Victor Gutierrez — ADASA"),
        ("Subject:", "RE: 25007 Taltal - Schedule Update"),
        ("Ref:", "Contract C-4300 / BAE 12803 / Schedule Update of 26-Jun-2026"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    add_para(doc, "Dear Eduardo,")
    doc.add_paragraph()

    add_para(
        doc,
        "Thank you for the schedule update of 26-June (the updated Project "
        "Schedule and the Progress Update).",
    )
    doc.add_paragraph()

    add_para(
        doc,
        "In our note of 25-June we confirmed two deliverables committed in the "
        "coordination call of 23-June: (1) the re-sequenced Project Schedule "
        "reflecting the Fedco delay, and (2) a formal report on the Fedco delay "
        "setting out the root cause, the sequence of events and the recovery "
        "actions. The schedule has been received. The formal report is still "
        "outstanding: the 26-June cover email gives a narrative and the Progress "
        "Update is a Gantt printout, whereas we asked for a report in report "
        "format setting out the root cause, the sequence of events and the "
        "recovery actions.",
    )
    doc.add_paragraph()

    add_para(
        doc,
        "Two points on the schedule itself. First, it extends the slip "
        "downstream rather than absorbing it: shipping completion moves from "
        "15-August to 10-September. We asked for a re-sequenced schedule that "
        "demonstrates how the slip is absorbed within the Performance Test, "
        "backed by concrete recovery actions; a continued-monitoring statement "
        "does not meet that request. "
        "Second, the Penang delivery now indicated for 02-September is later "
        "than the 21-August date we already declined in writing on 17-June, and "
        "it reopens the integrated running-test question raised in NT-001: if "
        "the High Feed Pump arrives after the shipment window, the integrated "
        "running test cannot be performed at Penang.",
    )
    doc.add_paragraph()

    add_para(
        doc,
        "We therefore ask that the formal report be issued by Friday, "
        "03-July-2026, containing:",
    )
    add_num_item(
        doc, 1, "Reconciled completion date",
        ": the single Fedco completion date, supported by the vendor's own "
        "documentation (a Fedco letter or purchase-order confirmation), not "
        "\"as advised\".",
    )
    add_num_item(
        doc, 2, "Root cause and sequence of events",
        ": with dates and supporting documents. Where the explanation refers to "
        "the down payment received on 09-June, please substantiate the complete "
        "chain: the original Fedco purchase order and its payment terms, the "
        "dependency between the advance and the order, and the corresponding "
        "dates, so the causal sequence is on the record. ADASA reserves its "
        "position on responsibility and does not accept any allocation by "
        "implication.",
    )
    add_num_item(
        doc, 3, "Recovery actions",
        ": including the integrated running-test location decision (Penang "
        "versus a site SAT), with personnel, window and cost.",
    )
    add_num_item(
        doc, 4, "LCP panel critical-path plan",
        ": given the panel must reach site before the FAT window (19-August to "
        "08-September).",
    )
    doc.add_paragraph()

    add_para(
        doc,
        "This report is needed to establish the basis on which the delay is "
        "assessed under the contract. ADASA reserves all of its rights.",
    )
    doc.add_paragraph()

    add_para(doc, "We look forward to the report by 03-July.")
    doc.add_paragraph()

    add_para(doc, "Best regards,")
    doc.add_paragraph()
    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        add_para(doc, line)

    if _HAS_META:
        apply_core_properties(
            doc,
            title="Fedco Delay - Formal Root-Cause Report Request",
            author="Luis Rivera",
            subject="25007 Taltal - Schedule Update",
            category="Email",
            comments="Aguas de Antofagasta S.A. - Proyecto Taltal BAE 12803",
            language="en-US",
            revision=1,
        )

    doc.save(OUTPUT_FILE)

    if _HAS_META:
        fix_app_xml(OUTPUT_FILE, application="Microsoft Office Word",
                    company="Aguas de Antofagasta S.A.")

    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
