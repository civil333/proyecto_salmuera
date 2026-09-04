#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Transmittal N25 (P22-TM-09-000-025-0). Submittals 25007-0055 (E55) +
25007-0056 (E56). Fecha: 29 de Junio de 2026.
Veredicto TM N25: 3 - TO BE REVISED. Tally 4 Code 1 + 1 Code 2 + 1 Code 3 (6 docs).
Driver unico Code 3: Outline Panel Drawing Rev B (gate de fabricacion del enclosure
aun abierto). IO List Rev 4 Code 2 - Approved as Noted: las 4 senales de coordinacion
con el PLC externo estan cumplidas; unica nota = celda de conteo de dosificadoras
(menor, fold a Rev 0). Cierres Code 1: LCP Rev 1 (N24), UHPRO Structural Rev B (N23),
Static Mixer Rev 0 (N22), Painting Spec Rev C. Escala los entregables vencidos con
deadline consolidado (viernes 03-Jul-2026). 2 CC_ADASA adjuntos (IO List Code 2 +
Outline Code 3).
English, BW Water. Formato ejecutivo, sin reuniones. Cadena = thread regular de
transmittals.
"""

import os
from docx import Document
from docx.shared import Pt, Inches

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-06-29_Transmittal-N25.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    fields = [
        ("Date:", "June 29, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
            "Mauricio Vallejos, Jorge Valdes, Tanya Figueroa, Allan Valentos, "
            "Jeryl F. Regulacion, Sadeep Irugalbandara, Andrew Sia, "
            "Ghazi Ozair, Nick Huta, Marjan Arsovic, Gerald Ross, "
            "Andrew Zaske, Adzlan Bin Abd Rahim",
        ),
        (
            "Subject:",
            "ADASA – Taltal Brine Module: Technical Review Transmittal N25 "
            "(25007-0055, 25007-0056)",
        ),
        ("Ref:", "Contract C-4300 / BAE 12803 / P22-TM-09-000-025-0"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Dear BW Water Project Team,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Attached: ")
    para.add_run("Transmittal N25").bold = True
    para.add_run(
        " (P22-TM-09-000-025-0) covering submittals 25007-0055 and "
        "25007-0056 — six documents — for ADASA's technical review."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Verdict: 3 — To Be Revised.").bold = True
    para.add_run(
        " Tally: 4 Code 1, 1 Code 2, 1 Code 3. The Local Control Panel Datasheet, "
        "the UHPRO Structural Design Criteria, the Static Mixer Datasheet and the "
        "Painting Specification are Approved; the IO List is Approved as Noted; the "
        "PLC-LCP Outline Panel Drawing is To Be Revised."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "The Outline Panel Drawing Rev B only half-corrects the enclosure "
        "contradiction that gates panel fabrication: it adds an exterior SUS316L "
        "line, a sealed air-conditioning unit and the panel weight, but the "
        "material and finishing rows still build a painted sheet-steel enclosure. "
        "Please re-issue it as Rev C with a coherent SS316L specification on every "
        "weather-exposed surface. The IO List Rev 4 incorporates the four "
        "module-to-external-PLC interface signals ADASA required and is Approved "
        "as Noted; the only remaining item is the dosing-pump count cell, to "
        "complete at the construction issue (no new revision)."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "The four Approved documents close the Transmittal N24 panel cycle, the "
        "Transmittal N23 structural cycle and the Transmittal N22 static-mixer "
        "cycle, each with the minor notes set out in the transmittal."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Outstanding deliverables. ").bold = True
    para.add_run(
        "Several committed items remain outstanding and are now overdue: the "
        "Grounding Point and Power Panel Location Layout Rev F (committed for "
        "17-Jun), the FAT and SAT comparison table (committed for 15-Jun), the "
        "Plant Control Philosophy children that gate the control package (a sixth "
        "cycle), and the three mechanical installation-route plans (committed for "
        "26-Jun). Please deliver these by Friday, July 3, 2026."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "Two annotated PDFs accompany the transmittal (the IO List and the "
        "Outline Panel Drawing); the four Approved documents carry no annotated "
        "PDF. Open items from previous transmittals are inventoried in its "
        "Section 3. Please confirm receipt and the target dates for the noted "
        "items."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("We look forward to your comments.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
