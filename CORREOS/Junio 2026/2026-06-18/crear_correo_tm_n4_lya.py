#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo de remision del Transmittal N4 (flujo OOCC) a L&A Ingenieria (Pablo Castillo).
Acusa la Entrega 7 (carta 067-032-032-COR-TT-008, 17-Jun-2026), remite el TM N4
(veredicto Codigo 3 - Por revisar) e INVOCA el ciclo de revision del TdR: los planos
de obras civiles van en Rev D sin converger a Rev 0 y arrastran comentarios de fondo
desde el TM N2 (mejoramiento de suelo, plano rector de excavaciones). Reincidencia de
fallas de entrega (MC-002-003 y DWG-002-005, 3a vez) y re-emisiones parciales.

Decision del usuario (18-Jun): invocar el limite de ciclo del TdR, anclado en los
planos (NO en la Clase 2 AACE, reclasificada menor). Reply-To a la cadena de la carta
TT-008. Idioma fijado en es-CL (CLAUDE.md Seccion 3.4). Estado: BORRADOR.
Dias verificados contra calendario: correo jueves 18-jun-2026; deadlines miercoles
24-jun-2026 y jueves 02-jul-2026.
"""

import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-06-18_TM-N4-OOCC-Remision.docx")
CONTACTO = "Luis Rivera"
LANG = "es-CL"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang_in_rPr(rPr, lang):
    lang_el = rPr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rPr.append(lang_el)
    lang_el.set(qn("w:val"), lang)
    lang_el.set(qn("w:eastAsia"), lang)
    lang_el.set(qn("w:bidi"), lang)


def fijar_idioma_documento(doc, lang=LANG):
    try:
        rPr = doc.styles["Normal"].element.get_or_add_rPr()
        _set_lang_in_rPr(rPr, lang)
    except KeyError:
        pass
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            _set_lang_in_rPr(run._element.get_or_add_rPr(), lang)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_lang_in_rPr(run._element.get_or_add_rPr(), lang)


def add_bullet(doc, lead, cuerpo):
    para = doc.add_paragraph(style="List Bullet")
    para.add_run(lead).bold = True
    para.add_run(cuerpo)
    aplicar_arial(para)
    return para


def add_tabla(doc, filas):
    """filas: lista de tuplas (col1, col2, col3); la primera es el encabezado."""
    tabla = doc.add_table(rows=len(filas), cols=3)
    tabla.style = "Table Grid"
    for i, fila in enumerate(filas):
        celdas = tabla.rows[i].cells
        for j, texto in enumerate(fila):
            celdas[j].text = ""
            run = celdas[j].paragraphs[0].add_run(texto)
            run.font.name = "Arial"
            run.font.size = Pt(10)
            if i == 0:
                run.bold = True
    return tabla


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    fields = [
        ("Fecha:", "18 de junio de 2026"),
        ("De:", f"{CONTACTO} — DIO ADASA (Aguas de Antofagasta S.A.)"),
        ("Para:",
         "Pablo Castillo — L&A Ingeniería y Proyectos (pcastillo@lyaingenieria.cl)"),
        ("CC:",
         "Macarena Vera, Lucas Molina, Cristhian Sánchez, Jesús Alarcón (L&A); "
         "Yohana Rodríguez, Víctor Gutiérrez (ADASA)"),
        ("Asunto:",
         "Transmittal N4 ADASA — Revisión Entrega 7 (OOCC): convergencia de planos "
         "a Rev 0 y entregables pendientes"),
        ("Ref:",
         "BAE 12803 / TdR P22-TR-00-010-01-1 / Carta 067-032-032-COR-TT-008 / "
         "Transmittal P22-TM-00-010-004-0"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Estimado Pablo:")
    aplicar_arial(para)
    doc.add_paragraph()

    # Apertura: acuse + remision + veredicto
    para = doc.add_paragraph(
        "Recibimos el miércoles 17 de junio la Entrega 7 (carta "
        "067-032-032-COR-TT-008). Adjuntamos el Transmittal N4 con su revisión. "
        "El veredicto global es ")
    para.add_run("Código 3 — Por revisar").bold = True
    para.add_run(".")
    aplicar_arial(para)
    doc.add_paragraph()

    # Narrativa: ningun comentario es nuevo (arrastre de N2/N3)
    para = doc.add_paragraph()
    para.add_run("El transmittal no incorpora observaciones nuevas. ").bold = True
    para.add_run(
        "Cada comentario se emitió en un transmittal anterior y sigue sin "
        "levantarse: la mayoría desde el Transmittal N2 (es la tercera vez que se "
        "formulan) y el resto desde el Transmittal N3. Junto a cada observación, "
        "tanto en el transmittal como en los PDF anotados, se indica su "
        "transmittal de origen.")
    aplicar_arial(para)
    doc.add_paragraph()

    # Reconocimiento del avance + tabla de documentos a emitir en Rev 0
    para = doc.add_paragraph()
    para.add_run("Reconocemos el avance del ciclo: ").bold = True
    para.add_run(
        "trece de los diecinueve documentos quedan aprobados para emisión en "
        "Rev 0. Solicitamos emitirlos en Rev 0; los marcados como aprobados con "
        "comentarios incorporan en esa misma emisión las observaciones del "
        "Transmittal N4, sin nueva revisión intermedia:")
    aplicar_arial(para)

    add_tabla(doc, [
        ("Documento", "Título", "Acción a Rev 0"),
        ("P22-MC-00-002-002", "Fundación Dinámica Bomba BH-06-001",
         "Emitir en Rev 0"),
        ("P22-MC-00-002-005", "Fundación Sistema CIP", "Emitir en Rev 0"),
        ("P22-MC-00-003-001", "Cubierta Metálica Sistema CIP Exterior",
         "Emitir en Rev 0"),
        ("P22-ET-00-010-101", "Especificación Movimiento de Tierra",
         "Emitir en Rev 0"),
        ("P22-ET-00-010-102", "Especificación Obras Civiles", "Emitir en Rev 0"),
        ("P22-ET-00-010-103", "Especificación Estructura Metálica",
         "Emitir en Rev 0"),
        ("P22-DWG-00-002-001", "Implantación General OOCC", "Emitir en Rev 0"),
        ("P22-MC-00-002-001", "Fundación Estanque TK-06-001",
         "Emitir en Rev 0 incorporando los comentarios"),
        ("P22-MC-00-002-004", "Fundación Contenedor RO",
         "Emitir en Rev 0 incorporando los comentarios"),
        ("P22-IT-00-010-101", "Itemizado Movimiento de Tierra",
         "Emitir en Rev 0 incorporando los comentarios"),
        ("P22-IT-00-010-102", "Itemizado Obras Civiles",
         "Emitir en Rev 0 incorporando los comentarios"),
        ("P22-IT-00-010-103", "Itemizado Estructura Metálica",
         "Emitir en Rev 0 incorporando los comentarios"),
        ("P22-DWG-00-002-006", "Canalizaciones y Red de Drenajes "
         "(re-emitir láminas 1 y 2)",
         "Emitir en Rev 0 incorporando los comentarios"),
    ])
    doc.add_paragraph()

    # Invocacion del ciclo de revision del TdR
    para = doc.add_paragraph()
    para.add_run(
        "Los seis planos de obras civiles que permanecen en Código 3 requieren "
        "una nueva revisión antes de Rev 0. ").bold = True
    para.add_run(
        "Van en Rev D sin converger, y los comentarios de fondo que ADASA viene "
        "observando desde el Transmittal N2 no se han incorporado: la nota de "
        "mejoramiento de suelo exigida por la Especificación de Movimiento de "
        "Tierra sigue ausente en los planos de fundación, con sus notas "
        "particulares sin cambios respecto de la revisión anterior. El ciclo de "
        "revisión de la ingeniería previsto en los Términos de Referencia "
        "(P22-TR-00-010-01-1) contempla un número acotado de revisiones antes de "
        "la emisión para construcción (Rev 0), y estos planos ya lo han excedido "
        "sin converger. Solicitamos formalmente que la próxima revisión sea la "
        "que cierre estos planos a Rev 0, incorporando la totalidad de los "
        "comentarios del Transmittal N4.")
    aplicar_arial(para)
    doc.add_paragraph()

    # Punto rector
    para = doc.add_paragraph()
    para.add_run(
        "El respaldo de las cubicaciones sigue siendo el punto que sostiene esa "
        "convergencia. ").bold = True
    para.add_run(
        "El plano de Excavaciones (P22-DWG-00-001-001 Rev C) continúa cubriendo "
        "solo las zanjas de drenaje; mientras no incorpore las excavaciones de las "
        "fundaciones, las cubicaciones de excavación que aparecen en los "
        "Itemizados no tienen un plano que las respalde.")
    aplicar_arial(para)
    doc.add_paragraph()

    # Entregables pendientes / reincidentes
    para = doc.add_paragraph(
        "Para poder cerrar la revisión, solicitamos completar y regularizar los "
        "siguientes entregables:")
    aplicar_arial(para)

    add_bullet(
        doc, "Entregar los dos documentos que no llegan desde hace tres entregas: ",
        "la Memoria de Cálculo del Sistema de Drenajes (P22-MC-00-002-003) y el "
        "plano de Detalles de Anclaje y Conexiones (P22-DWG-00-002-005). Este "
        "último es el que permite cerrar la coherencia entre la memoria de cálculo "
        "del anclaje del estanque y el plano de fundación.")
    add_bullet(
        doc, "Re-emitir completos los planos que llegaron con láminas sueltas, "
        "incorporando en ellas los comentarios pendientes: ",
        "el plano de Canalizaciones (P22-DWG-00-002-006) llegó solo con la lámina "
        "3, por lo que faltan la 1 y la 2 (nomenclatura de cámaras CD-06-00N); el "
        "de Fundación del Sistema CIP (P22-DWG-00-002-007) llegó solo con la "
        "lámina 1, falta la 3 (mejoramiento de suelo e impermeabilización); y el "
        "de Cubierta Metálica (P22-DWG-00-003-001) llegó solo con la lámina 1, "
        "falta la 2 (corrección de la sigla de la Especificación C5-M).")
    add_bullet(
        doc, "Regularizar la revisión de la Cubierta Metálica "
        "(P22-DWG-00-003-001): ",
        "la carta de transmisión la lista como Rev D, pero el cajetín de la lámina "
        "dice Rev 0. Dejar una sola revisión coherente y corregir la fecha del "
        "cajetín.")

    doc.add_paragraph()

    # Plazos
    para = doc.add_paragraph(
        "Agradecemos reponer los archivos no entregados (P22-MC-00-002-003 y "
        "P22-DWG-00-002-005) y las láminas faltantes a más tardar el ")
    para.add_run("miércoles 24 de junio de 2026").bold = True
    para.add_run(
        ", y emitir la revisión que cierra los planos de obras civiles y los "
        "Itemizados a Rev 0, con el plano de excavación de fundaciones que "
        "respalde las cubicaciones, a más tardar el ")
    para.add_run("jueves 2 de julio de 2026").bold = True
    para.add_run(".")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Quedamos atentos a sus comentarios.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Saludos cordiales,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in [
        "Departamento de Ingeniería y Optimización (DIO)",
        "ADASA — Aguas de Antofagasta S.A.",
        "luis.rivera@adasa.cl",
    ]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    fijar_idioma_documento(doc, LANG)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")
    print(f"Idioma del documento: {LANG}")


if __name__ == "__main__":
    crear_correo()
