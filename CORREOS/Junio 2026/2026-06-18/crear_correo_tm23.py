#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Transmittal N23 (P22-TM-09-000-023-0). Submittals 25007-0051 (E51) +
25007-0052 (E52). Fecha: 18 de Junio de 2026.
Veredicto TM N23: 3 - TO BE REVISED. Tally 0 Code 1 + 1 Code 2 + 6 Code 3 (7 docs).
Driver global: RO Vessel Hydrostatic Test Procedure Rev A (presion de prueba sin
valor vinculante / 45,5 bar vs 1.980 psi del ITP). Cierre positivo: el ITP Rev C
cierra materialmente el item de fabricacion mas antiguo (base del waiver ASME).
7 CC_ADASA adjuntos (6 Code 3 + 1 Code 2; sin Code 1 este ciclo).
English, BW Water. Formato ejecutivo (sin proponer reuniones). Carry-forward
neutral (sin escalar los vencidos; inventario en Section 3 del transmittal).
"""

import os
from docx import Document
from docx.shared import Pt, Inches

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-06-18_Transmittal-N23.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    fields = [
        ("Date:", "June 18, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
            "Mauricio Vallejos, Jorge Valdes, Tanya Figueroa, Allan Valentos, "
            "Jeryl F. Regulacion, Sadeep Irugalbandara, Andrew Sia, "
            "Ghazi Ozair, Nick Huta, Marjan Arsovic, Gerald Ross, "
            "Andrew Zaske, Adzlan Bin Abd Rahim",
        ),
        (
            "Subject:",
            "ADASA – Taltal Brine Module: Technical Review Transmittal N23 "
            "(25007-0051, 25007-0052)",
        ),
        ("Ref:", "Contract C-4300 / BAE 12803 / P22-TM-09-000-023-0"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Dear BW Water Project Team,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Attached: ")
    para.add_run("Transmittal N23").bold = True
    para.add_run(
        " (P22-TM-09-000-023-0) covering submittals 25007-0051 and "
        "25007-0052 — seven documents — for ADASA's technical review."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Verdict: 3 — To Be Revised.").bold = True
    para.add_run(
        " Tally: 0 Code 1, 1 Code 2, 6 Code 3, driven by the RO Vessel "
        "Hydrostatic Test Procedure."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "On the positive side, the Inspection and Test Plan Rev C materially "
        "closes the package's oldest open fabrication item: the agreed vessel "
        "test basis (1,800 psi x 1.1, ADASA witness, production and test "
        "dossier) is now captured. The verdict is fixed by the RO Vessel "
        "Hydrostatic Test Procedure, whose test pressure does not match the "
        "1,980 psi the Inspection and Test Plan requires; the Painting "
        "Procedure and the Instrument Location Layout also require a new "
        "revision. The per-document codes, observations and required actions "
        "are set out in the attached transmittal."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "Seven annotated PDFs accompany the transmittal (no Code 1 documents "
        "this cycle); open items from previous transmittals are inventoried in "
        "its Section 3. Please confirm receipt and the target dates for the "
        "revisions."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("We look forward to your comments.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
