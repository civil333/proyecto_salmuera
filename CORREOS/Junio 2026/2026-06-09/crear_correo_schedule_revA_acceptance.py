#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo ADASA -> BW Water: respuesta al Project Schedule Rev A (08-Jun).

BW Water entrego el nuevo Project Schedule 08-06-26 Rev A en respuesta a la
escalacion de ADASA del 08-Jun. Este correo lo ADOPTA como linea base de
recuperacion CON reservas:
- El ex-works Espana (23-Jun) + 40d mar a Penang (02-Ago) + EXW Penang (15-Ago)
  quedan trazables; el waiver sin estampa es lo que sostiene el 15-Ago.
- Pero el schedule NO menciona ASME/estampa ni muestra prueba de presion de los
  vessels -> la base sin estampa + la hidrostatica de fabrica + dossier + testigo
  deben ir al ITP actualizado.
- Adoptar Rev A no cierra: 25 aclaraciones NT-001, tabla FAT/SAT 15-Jun,
  evidencia firmada Fedco, filtro RO alternativo.

- Reply-To al thread del Mitigation Plan
  (subject "RE: 25007 Taltal - Mitigation Plan - Pressure Vessel Protec";
   confirmar el subject real del correo entrante antes de enviar).
- Idioma: ingles (BW Water). Patron Document() directo per CLAUDE.md 3.4.
- Fuente unica del cuerpo:
  2026-06-09_Recovery-Schedule-RevA-Acceptance-with-Reservations_Descripcion.md
"""

import os

from docx import Document
from docx.shared import Inches, Pt

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(
    SCRIPT_DIR, "2026-06-09_Recovery-Schedule-RevA-Acceptance-with-Reservations.docx"
)
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_para(doc, text, size=11):
    para = doc.add_paragraph(text)
    aplicar_arial(para, size)
    return para


def add_segments(doc, segments, size=11):
    """Parrafo con segmentos (texto, bold?). segments = [(text, bool), ...]."""
    para = doc.add_paragraph()
    for text, bold in segments:
        run = para.add_run(text)
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(size)
    return para


def add_bullet(doc, segments, size=11):
    """Vineta nativa (List Bullet) con segmentos (texto, bold?)."""
    para = doc.add_paragraph(style="List Bullet")
    for text, bold in segments:
        run = para.add_run(text)
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(size)
    return para


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # =========================================================================
    # HEADER
    # =========================================================================
    fields = [
        ("Date:", "June 9, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Andrew Sia, Adzlan Bin Abd Rahim, Jeryl F. Regulacion, "
            "Sadeep Irugalbandara, Nick Huta, Stephane Gehant, Shane Banks, "
            "Victor Gutierrez",
        ),
        ("Subject:", "RE: 25007 Taltal - Mitigation Plan - Pressure Vessel Protec"),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / "
            "Project Schedule Rev A 08-Jun-2026 / "
            "ASME Stamp Waiver 02-Jun-2026 / "
            "Protec Arisawa Europe letter / "
            "Technical Note P22-NT-09-000-001-0",
        ),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    blank(doc)

    # =========================================================================
    # CUERPO
    # =========================================================================
    add_para(doc, "Dear Eduardo,")
    blank(doc)

    # Para 1 - acuse + adopcion + aclaracion ex-works (ejecutivo)
    add_segments(
        doc,
        [
            ("Thank you for Project Schedule Rev A (08 June). It makes the "
             "pressure-vessel route traceable: ex-works Spain on ", False),
            ("23 June", True),
            (", a 40-day sea freight to Penang (arrival 02 August), and ex-works "
             "Penang on ", False),
            ("15 August", True),
            (". ADASA ", False),
            ("adopts Rev A as the recovery baseline", True),
            ("; please carry it as such in the weekly tracker. Note that 23 June "
             "is the factory date in Spain; Penang arrival is 02 August.", False),
        ],
    )
    blank(doc)

    # Lead-in a reservas
    add_para(
        doc,
        "Adopting Rev A does not close the items below, which continue to condition "
        "FAT, dispatch and the EP-2 milestone:",
    )

    # Bullet 1 - ASME basis (load-bearing, primero)
    add_bullet(
        doc,
        [
            ("ASME basis:", True),
            (" Rev A does not state it. The 23 June ex-works date is the "
             "non-stamped route under our 02 June waiver. The ", False),
            ("updated ITP must reflect this", True),
            (", with the factory hydrostatic test at rating (1800 psi x 1.1), the "
             "test-record dossier and ADASA's witness right at Protec. Rev A shows "
             "no vessel pressure test; please confirm where and when it is "
             "performed.", False),
        ],
    )
    add_bullet(
        doc,
        [
            ("The written responses to the 25 clarifications in Technical Note "
             "P22-NT-09-000-001-0.", False),
        ],
    )
    add_bullet(
        doc,
        [
            ("The FAT/SAT scope table and dry-test procedure ", False),
            ("due 15 June", True),
            (", with the HP pump and turbocharger running test located.", False),
        ],
    )
    add_bullet(
        doc,
        [
            ("A vendor-signed Fedco readiness for 09 August and the contingency if "
             "it slips; the schedule leaves about six days to ex-works Penang.",
             False),
        ],
    )
    add_bullet(
        doc,
        [
            ("The alternative RO cartridge filter (material justification and "
             "datasheet) for approval ", False),
            ("before any PO", True),
            (".", False),
        ],
    )
    blank(doc)

    add_para(
        doc,
        "We will review these at tomorrow's meeting. We look forward to it.",
    )
    blank(doc)

    # =========================================================================
    # CIERRE
    # =========================================================================
    add_para(doc, "Best regards,")
    blank(doc)

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        add_para(doc, line)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
