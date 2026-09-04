#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Transmittal N22 (P22-TM-09-000-022-0). Submittals 25007-0049 (E49) +
25007-0050 (E50). Fecha: 15 de Junio de 2026.
Veredicto TM N22: 3 - TO BE REVISED. Tally 2 Code 1 + 1 Code 2 + 4 Code 3 (7 docs).
Driver global: Plant Control Philosophy Rev D (logica operativa en documentos
hijos no entregados). Cierres positivos: permissive CRITICAL del HP Pump cerrado;
HMI screen design entregado materialmente (cierra TM N4, ~127 dias); RO Cartridge
Filter aprobado (FRP retenido). 5 CC_ADASA adjuntos (4 Code 3 + 1 Code 2).
English, BW Water. Formato ejecutivo (sin proponer reuniones).
"""

import os
from docx import Document
from docx.shared import Pt, Inches

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-06-15_Transmittal-N22.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def bullet(doc, lead, rest):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.25)
    para.add_run("• ")
    para.add_run(lead).bold = True
    para.add_run(rest)
    aplicar_arial(para)
    return para


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    fields = [
        ("Date:", "June 15, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
            "Mauricio Vallejos, Jorge Valdes, Tanya Figueroa, Allan Valentos, "
            "Jeryl F. Regulacion, Sadeep Irugalbandara, Andrew Sia, "
            "Ghazi Ozair, Nick Huta, Marjan Arsovic, Gerald Ross, "
            "Andrew Zaske, Adzlan Bin Abd Rahim",
        ),
        (
            "Subject:",
            "ADASA – Taltal Brine Module: Technical Review Transmittal N22 "
            "(25007-0049, 25007-0050)",
        ),
        ("Ref:", "Contract C-4300 / BAE 12803 / P22-TM-09-000-022-0"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Dear BW Water Project Team,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Attached: ")
    para.add_run("Transmittal N22").bold = True
    para.add_run(
        " (P22-TM-09-000-022-0) covering submittals 25007-0049 and "
        "25007-0050 — seven documents — for ADASA's technical review."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Verdict: 3 — To Be Revised.").bold = True
    para.add_run(
        " Tally: 2 Code 1, 1 Code 2, 4 Code 3, driven by the Plant Control "
        "Philosophy Rev D."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "Two highlights this cycle: the repeated CRITICAL HP Pump start "
        "permissive is closed, and the HMI screen design is materially "
        "delivered — the project's oldest open commitment. Four documents "
        "require a new revision (the Plant Control Philosophy, the Equipment "
        "Layout, the CIP Cartridge Filter and the HMI Display Screenshot); "
        "the per-document codes, observations and required actions are set "
        "out in the attached transmittal."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "Five annotated PDFs accompany the transmittal (the two Code 1 "
        "documents carry none); open items from previous transmittals are "
        "inventoried in its Section 3. Please confirm receipt and the target "
        "dates for the four revisions."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("We look forward to your comments.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
