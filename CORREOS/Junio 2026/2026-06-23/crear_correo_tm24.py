#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Transmittal N24 (P22-TM-09-000-024-0). Submittals 25007-0053 (E53) +
25007-0054 (E54). Fecha: 23 de Junio de 2026.
Veredicto TM N24: 2 - APPROVED AS NOTED. Tally 3 Code 1 + 2 Code 2 (5 docs).
Cierre positivo: el CIP Cartridge Filter Datasheet Rev E cierra el item del TM N22
(gasket EPDM + compatibilidad FRP). IO List y LCP Datasheet Code 2 (notas menores).
En el IO List, ADASA extiende la interfaz con 2 senales adicionales como PEDIDO NUEVO
(no incumplimiento; verificado: solo se habian pedido 2). 2 CC_ADASA adjuntos
(IO List + LCP, ambos Code 2). English, BW Water. Formato ejecutivo, sin reuniones.
"""

import os
from docx import Document
from docx.shared import Pt, Inches

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-06-23_Transmittal-N24.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    fields = [
        ("Date:", "June 23, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
            "Mauricio Vallejos, Jorge Valdes, Tanya Figueroa, Allan Valentos, "
            "Jeryl F. Regulacion, Sadeep Irugalbandara, Andrew Sia, "
            "Ghazi Ozair, Nick Huta, Marjan Arsovic, Gerald Ross, "
            "Andrew Zaske, Adzlan Bin Abd Rahim",
        ),
        (
            "Subject:",
            "ADASA – Taltal Brine Module: Technical Review Transmittal N24 "
            "(25007-0053, 25007-0054)",
        ),
        ("Ref:", "Contract C-4300 / BAE 12803 / P22-TM-09-000-024-0"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Dear BW Water Project Team,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Attached: ")
    para.add_run("Transmittal N24").bold = True
    para.add_run(
        " (P22-TM-09-000-024-0) covering submittals 25007-0053 and "
        "25007-0054 — five documents — for ADASA's technical review."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Verdict: 2 — Approved as Noted.").bold = True
    para.add_run(
        " Tally: 3 Code 1, 2 Code 2. The Cable Schedule, the Modbus Data "
        "Transfer List and the CIP Cartridge Filter Datasheet are Approved; "
        "the IO List and the Local Control Panel Datasheet are Approved as "
        "Noted."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "The CIP Cartridge Filter Datasheet Rev E closes the open fabrication "
        "item from Transmittal N22: the cartridge gasket is now EPDM and the "
        "material-compatibility statement for the FRP housing against the CIP "
        "fluid is attached. The Local Control Panel Datasheet enclosure "
        "reconfirms the correct marine specification; its only notes are two "
        "minor corrections."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "On the IO List, the two module-to-external-PLC interface signals "
        "previously requested are present and correctly hardwired. ADASA is "
        "extending this interface with two further relay-contact signals — the "
        "module fault status and the module local/remote status — as a new "
        "requirement set out in the transmittal, to be incorporated at the "
        "construction issue."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "Two annotated PDFs accompany the transmittal (the IO List and the "
        "Local Control Panel Datasheet); the three Approved documents carry no "
        "annotated PDF. Open items from previous transmittals are inventoried "
        "in its Section 3. Please confirm receipt and the target date for the "
        "noted items."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("We look forward to your comments.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
