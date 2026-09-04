#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo ADASA -> BW Water: reversion de posicion sobre la estampa ASME en los
RO Pressure Vessels. Version EJECUTIVA (condensada).

En la reunion del 02-Jun ADASA reitero que la estampa ASME era requerida.
Tras revisar la carta de Protec Arisawa y el estado de atraso, ADASA se desdice
y acepta la version SIN estampa (entrega 22/06) como gesto de cooperacion con la
reprogramacion. El correo abre con la disculpa por el cambio y fija las
condiciones que protegen a ADASA (escenario 22/06 en el recovery schedule,
dossier completo de pruebas, sin precedente ni renuncia a otras exigencias).

- Reply-To al thread del Mitigation Plan (subject "RE: 25007 Taltal - Mitigation Plan").
- Idioma: ingles (BW Water). Patron Document() directo per CLAUDE.md 3.4.
- Fuente unica del cuerpo: 2026-06-02_ASME-Stamp-Waiver-RO-Pressure-Vessels_Descripcion.md
"""

import os

from docx import Document
from docx.shared import Inches, Pt

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(
    SCRIPT_DIR, "2026-06-02_ASME-Stamp-Waiver-RO-Pressure-Vessels.docx"
)
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_para(doc, text, size=11):
    para = doc.add_paragraph(text)
    aplicar_arial(para, size)
    return para


def add_segments(doc, segments, size=11):
    """Parrafo con segmentos (texto, bold?). segments = [(text, bool), ...]."""
    para = doc.add_paragraph()
    for text, bold in segments:
        run = para.add_run(text)
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(size)
    return para


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # =========================================================================
    # HEADER
    # =========================================================================
    fields = [
        ("Date:", "June 2, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Andrew Sia, Victor Gutierrez, Jeryl F. Regulacion, "
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Mauricio Vallejos, "
            "Jorge Valdes, Tanya Figueroa, Allan Valentos, Sadeep Irugalbandara, "
            "Ghazi Ozair, Nick Huta, Marjan Arsovic, Gerald Ross, "
            "Adzlan Bin Abd Rahim, Stephane Gehant, Shane Banks, Fadey Kassim",
        ),
        ("Subject:", "RE: 25007 Taltal - Mitigation Plan"),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / "
            "Protec Arisawa Europe letter (delivery times and ASME) / "
            "Recovery Schedule 22-May-2026 / "
            "Technical Note P22-NT-09-000-001-0",
        ),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    blank(doc)

    # =========================================================================
    # CUERPO (version ejecutiva)
    # =========================================================================
    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_para(
        doc,
        "In the meeting we just concluded we stated that the ASME Section X stamp "
        "was required for the RO pressure vessels. Having reviewed the Protec "
        "Arisawa letter and the recovery status, we are revising that position, "
        "and we apologize for the change so soon after the meeting.",
    )
    blank(doc)

    add_segments(
        doc,
        [
            ("ADASA will accept the non-stamped version", True),
            (". Protec is a recognized FRP vessel manufacturer; its letter confirms "
             "that design, materials, quality controls and testing are identical in "
             "both versions, with test pressure ", False),
            ("1800 psi × 1.1", True),
            (" in each case, the only difference being the ASME inspector. As a "
             "gesture to support the reprogramming, we waive the stamp for this "
             "equipment, on three conditions:", False),
        ],
    )
    blank(doc)

    add_segments(
        doc,
        [
            ("1. ", False),
            ("Availability. ", True),
            ("Build the recovery schedule due 04 June on the non-stamped ", False),
            ("22 June", True),
            (" date, stating whether 22 June is ex-works Spain or delivered at "
             "Penang and the transit time, so the downstream and EXW-Penang dates "
             "are traceable.", False),
        ],
    )
    blank(doc)

    add_segments(
        doc,
        [
            ("2. ", False),
            ("Dossier. ", True),
            ("Deliver the full test records for the supplied units (hydrostatic "
             "test at rating, material and dimensional certificates, "
             "fiber-percentage documentation), plus the BPV81200SP "
             "type-qualification reports for the cyclic and bursting tests. We "
             "set aside only the ASME inspector's sign-off, not the underlying "
             "test records.", False),
        ],
    )
    blank(doc)

    add_segments(
        doc,
        [
            ("3. ", False),
            ("Scope. ", True),
            ("This waiver covers the RO pressure vessels and this recovery only. "
             "It changes no other requirement of the Specification or the Contract, "
             "the vessel performance and warranties remain BW Water's, and no "
             "change order is pursued for the ASME item provided the time saved "
             "reaches the schedule.", False),
        ],
    )
    blank(doc)

    add_para(
        doc,
        "Please reflect this in the recovery schedule and the updated ITP. We look "
        "forward to your comments.",
    )
    blank(doc)

    # =========================================================================
    # CIERRE
    # =========================================================================
    add_para(doc, "Best regards,")
    blank(doc)

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        add_para(doc, line)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
