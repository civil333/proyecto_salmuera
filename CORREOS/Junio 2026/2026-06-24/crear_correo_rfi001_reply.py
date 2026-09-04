#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo de cobertura: Reply-To al thread del RFI-001 (25007-RO-RFI-0001),
emitido por Billy Tan (BW Water / Engineering) el 22-Jun-2026.

ADASA devuelve el form RFI con la respuesta en la seccion "Replied
Information" (adjunto) y confirma en el cuerpo, en una linea, que la
configuracion descrita cumple la ET Seccion 5.5: 4-20mA+HART es requisito
a nivel de instrumento; no se exige pass-through HART al PLC/SCADA. Sin
proponer reunion (correo transaccional, cover-only).

Fecha: 24 de Junio de 2026 (miercoles). Cadena del propio RFI, separada de
los transmittals (regla de cadenas por instrumento contractual).
"""

import os
import sys
from docx import Document
from docx.shared import Pt, Inches

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
from docx_metadata import apply_core_properties, fix_app_xml

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(
    SCRIPT_DIR, "2026-06-24_RFI-001-HART-Reply.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    fields = [
        ("Date:", "June 24, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Billy Tan — BW Water / Engineering"),
        (
            "CC:",
            "Eduardo Yamauchi, Stephane Gehant, Jeryl F. Regulacion, "
            "Adzlan Bin Abd Rahim, Nick Huta, Sadeep Irugalbandara, "
            "Victor Gutierrez",
        ),
        (
            "Subject:",
            "RE: RFI 25007-RO-RFI-0001 — Analog Input Modules and the "
            "4-20 mA + HART instrumentation requirement (Section 5.5)",
        ),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / RFI 25007-RO-RFI-0001 / "
            "P22-ET-09-000-001-0",
        ),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Dear Billy,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "We acknowledge RFI 25007-RO-RFI-0001 and return it with ADASA's "
        "reply completed in the Replied Information section (attached)."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Compliance with Section 5.5 is confirmed. ").bold = True
    para.add_run(
        "The 4-20 mA + HART protocol is a field-instrument requirement. It "
        "does not require HART pass-through to the PLC or SCADA, "
        "HART-capable I/O modules, or a dedicated HART asset-management "
        "solution. The Allen-Bradley 5069-IF8 acquiring the 4-20 mA "
        "process variable, with HART available locally at the loop through "
        "a handheld communicator, therefore complies, and no additional "
        "HART-capable I/O hardware is required."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "This confirmation is conditioned on every field instrument being "
        "supplied 4-20 mA + HART-capable; an instrument supplied as 4-20 mA "
        "only, without HART, would not comply. The completed RFI form sets "
        "out the full reply."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("We look forward to your acknowledgement.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    apply_core_properties(
        doc,
        title="RFI 25007-RO-RFI-0001 - ADASA Reply (cover)",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="RFI 25007-RO-RFI-0001 - 4-20 mA + HART",
        comments="",
        language="en-US",
    )
    doc.save(OUTPUT_FILE)
    fix_app_xml(OUTPUT_FILE, application="Microsoft Office Word",
                company="Aguas Antofagasta")
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
