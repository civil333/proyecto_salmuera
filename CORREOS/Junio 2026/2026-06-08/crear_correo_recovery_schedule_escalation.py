#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo ADASA -> BW Water: escalacion del Recovery Schedule vencido (08-Jun).

Tercer requerimiento del mismo deliverable: pedido para la reunion del 02-Jun,
re-pedido el 04-Jun, aun no recibido al 08-Jun. Escala el tono y fija deadline
firme anclado a un evento fisico (reunion semanal del martes 09-Jun). Hace
explicita la palanca que el 04-Jun dejo implicita: la condicion #1 ("Availability")
del waiver ASME del 02-Jun es justamente este schedule sobre la fecha sin estampa
del 22-Jun + declaracion ex-works Espana vs Penang; sin el, el waiver no queda
consolidado. Se reafirma por referencia, sin re-abrir las tres condiciones.

- Reply-To al thread del Mitigation Plan
  (subject "RE: 25007 Taltal - Mitigation Plan - Pressure Vessel Protec").
- Idioma: ingles (BW Water). Patron Document() directo per CLAUDE.md 3.4.
- Fuente unica del cuerpo: 2026-06-08_Recovery-Schedule-Escalation_Descripcion.md
"""

import os

from docx import Document
from docx.shared import Inches, Pt

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-06-08_Recovery-Schedule-Escalation.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_para(doc, text, size=11):
    para = doc.add_paragraph(text)
    aplicar_arial(para, size)
    return para


def add_segments(doc, segments, size=11):
    """Parrafo con segmentos (texto, bold?). segments = [(text, bool), ...]."""
    para = doc.add_paragraph()
    for text, bold in segments:
        run = para.add_run(text)
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(size)
    return para


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # =========================================================================
    # HEADER
    # =========================================================================
    fields = [
        ("Date:", "June 8, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Andrew Sia, Adzlan Bin Abd Rahim, Jeryl F. Regulacion, "
            "Sadeep Irugalbandara, Nick Huta, Stephane Gehant, Shane Banks, "
            "Victor Gutierrez",
        ),
        ("Subject:", "RE: 25007 Taltal - Mitigation Plan - Pressure Vessel Protec"),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / "
            "Recovery Schedule (overdue) / "
            "ASME Stamp Waiver 02-Jun-2026 / "
            "Protec Arisawa Europe letter / "
            "Technical Note P22-NT-09-000-001-0",
        ),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    blank(doc)

    # =========================================================================
    # CUERPO
    # =========================================================================
    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_segments(
        doc,
        [
            ("We are still without the updated recovery schedule. It was first due "
             "for the weekly progress meeting on 02 June, requested again on 04 "
             "June, and as of today, 08 June, it has not reached us. The weekly "
             "progress meeting is ", False),
            ("tomorrow, Tuesday 09 June", True),
            (", and we need the schedule in hand beforehand so it can be reviewed "
             "there.", False),
        ],
    )
    blank(doc)

    add_segments(
        doc,
        [
            ("Our 02 June message waived the ASME Section X stamp on the RO "
             "pressure vessels specifically to recover the time, and that waiver "
             "rests on three conditions. The first, ", False),
            ("Availability", True),
            (", is precisely this schedule: built on the non-stamped ", False),
            ("22 June", True),
            (" date, stating whether 22 June is ex-works Spain or delivered at "
             "Penang and the transit time, so the EXW-Penang and downstream dates "
             "are traceable. Until it reaches us, ", False),
            ("the first condition is unmet and the waiver is not consolidated", True),
            (", and we cannot confirm that the time it was meant to save actually "
             "reaches the program.", False),
        ],
    )
    blank(doc)

    add_para(
        doc,
        "The delivery period remains firm and any sub-supplier delay is recovered "
        "at BW Water's cost (BAE Clauses 27, 35 and 46), with payment milestones "
        "running against it (Clause 31). We are not reopening any of that here; we "
        "are asking only for the schedule that was already due.",
    )
    blank(doc)

    add_segments(
        doc,
        [
            ("Please send the recovery schedule, on the ", False),
            ("22 June", True),
            (" basis and with the ex-works / Penang declaration, ", False),
            ("ahead of tomorrow's meeting", True),
            (". We look forward to it.", False),
        ],
    )
    blank(doc)

    # =========================================================================
    # CIERRE
    # =========================================================================
    add_para(doc, "Best regards,")
    blank(doc)

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        add_para(doc, line)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
