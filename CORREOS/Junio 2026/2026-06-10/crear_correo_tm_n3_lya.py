#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo de remision del Transmittal N3 (flujo OOCC) a L&A Ingenieria (Pablo Castillo).
Acusa la Entrega 5 (carta 067-032-032-COR-TT-006), remite el TM N3 (veredicto
Codigo 3 - Por revisar) y eleva el punto rector: sin un plano de excavacion de las
fundaciones, las cubicaciones de excavacion de fundaciones no tienen respaldo grafico.
Lista las cuatro fallas de entrega que requieren reposicion.

Reply-To a la cadena de la carta TT-006. Idioma fijado en es-CL (CLAUDE.md Seccion 3.4).
Estado: BORRADOR.
"""

import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-06-10_TM-N3-OOCC-Remision.docx")
CONTACTO = "Luis Rivera"
LANG = "es-CL"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang_in_rPr(rPr, lang):
    lang_el = rPr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rPr.append(lang_el)
    lang_el.set(qn("w:val"), lang)
    lang_el.set(qn("w:eastAsia"), lang)
    lang_el.set(qn("w:bidi"), lang)


def fijar_idioma_documento(doc, lang=LANG):
    try:
        rPr = doc.styles["Normal"].element.get_or_add_rPr()
        _set_lang_in_rPr(rPr, lang)
    except KeyError:
        pass
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            _set_lang_in_rPr(run._element.get_or_add_rPr(), lang)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_lang_in_rPr(run._element.get_or_add_rPr(), lang)


def add_bullet(doc, lead, cuerpo):
    para = doc.add_paragraph(style="List Bullet")
    para.add_run(lead).bold = True
    para.add_run(cuerpo)
    aplicar_arial(para)
    return para


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    fields = [
        ("Fecha:", "10 de junio de 2026"),
        ("De:", f"{CONTACTO} — DIO ADASA (Aguas de Antofagasta S.A.)"),
        ("Para:",
         "Pablo Castillo — L&A Ingeniería y Proyectos (pcastillo@lyaingenieria.cl)"),
        ("CC:",
         "Macarena Vera, Lucas Molina, Cristhian Sánchez, Jesús Alarcón (L&A); "
         "Yohana Rodríguez, Víctor Gutiérrez (ADASA)"),
        ("Asunto:",
         "Transmittal N3 ADASA — Revisión Entrega 5 (OOCC) y entregables pendientes"),
        ("Ref:",
         "BAE 12803 / TdR P22-TR-00-010-01-1 / Carta 067-032-032-COR-TT-006 / "
         "Transmittal P22-TM-00-010-003-0"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Estimado Pablo:")
    aplicar_arial(para)
    doc.add_paragraph()

    # Apertura: acuse + remision + veredicto
    para = doc.add_paragraph(
        "Recibimos el martes 9 de junio la Entrega 5 (carta "
        "067-032-032-COR-TT-006). Adjuntamos el Transmittal N3 con su revisión. "
        "El veredicto global es ")
    para.add_run("Código 3 — Por revisar").bold = True
    para.add_run(
        ": las seis memorias de cálculo y las tres especificaciones técnicas "
        "quedan en Código 2, aprobadas con comentarios a incorporar en Rev 0, e "
        "incluso quedó cerrado en plano el anclaje preinstalado del estanque, el "
        "punto más antiguo de la revisión OOCC. El veredicto lo fijan los planos "
        "de obras civiles y los tres Itemizados, que todavía arrastran "
        "observaciones.")
    aplicar_arial(para)
    doc.add_paragraph()

    # Punto rector
    para = doc.add_paragraph()
    para.add_run("El punto más relevante es el respaldo de las cubicaciones. ").bold = True
    para.add_run(
        "El plano de Excavaciones (P22-DWG-00-001-001 Rev B) cubre solo las "
        "zanjas de drenaje; no incluye las excavaciones de las fundaciones. "
        "Mientras ese plano no exista, las cubicaciones de excavación de "
        "fundaciones que aparecen en los Itemizados y en las memorias no tienen "
        "un plano que las respalde. Necesitamos que la próxima revisión lo "
        "incorpore.")
    aplicar_arial(para)
    doc.add_paragraph()

    # Fallas de entrega
    para = doc.add_paragraph(
        "Registramos además cuatro fallas de entrega que requieren reposición:")
    aplicar_arial(para)

    add_bullet(
        doc, "P22-MC-00-002-003 (Sistema de Drenajes): ",
        "declarada en la carta (ítem 9), pero el archivo no llegó; en su lugar "
        "vino una Rev 0 antigua de la cubierta metálica.")
    add_bullet(
        doc, "P22-DWG-00-002-003 lámina 2: ",
        "el archivo entregado es un duplicado de la lámina 1; falta el detalle "
        "del inserto INS-1 al que remite el plano.")
    add_bullet(
        doc, "P22-DWG-00-002-006 lámina 3: ",
        "el cajetín la anuncia como “3 de 3”, pero no fue incluida (detalle "
        "típico de cámara).")
    add_bullet(
        doc, "P22-DWG-00-002-005 (Detalles de Anclaje y Conexiones): ",
        "no entregado, por segunda vez consecutiva.")

    doc.add_paragraph()

    # Solicitud + plazo
    para = doc.add_paragraph(
        "Agradecemos reponer estos cuatro archivos a más tardar el ")
    para.add_run("viernes 12 de junio de 2026").bold = True
    para.add_run(
        ", y emitir la revisión de los planos y de los Itemizados incorporando "
        "los comentarios del Transmittal N3, junto con el plano de excavación de "
        "fundaciones que respalde las cubicaciones, en el siguiente ciclo de "
        "revisión.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Quedamos atentos a sus comentarios.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Saludos cordiales,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in [
        "Departamento de Ingeniería y Optimización (DIO)",
        "ADASA — Aguas de Antofagasta S.A.",
        "luis.rivera@adasa.cl",
    ]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    fijar_idioma_documento(doc, LANG)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")
    print(f"Idioma del documento: {LANG}")


if __name__ == "__main__":
    crear_correo()
