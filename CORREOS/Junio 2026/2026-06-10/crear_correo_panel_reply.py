#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Reply-To al thread BW Water 10-Jun-2026 9:52 (Eduardo Yamauchi)
"panel drawings submitted for approval on May 28 under Transmittal No.
25007-0046" — pedido de revision expedita del paquete panel PLC-LCP.

ADASA acepta la via de revision por correo. Disposicion del paquete en
TM N20 (emitido hoy, cadena regular): Outline Rev A Code 3 (driver:
contradiccion material/IP del enclosure), Schematic Rev A Code 2,
LCP Datasheet Rev B Code 2. Ruta expedita ofrecida: confirmacion escrita
del enclosure del datasheet + Outline Rev B alineado => release.

Fecha: 10 de Junio de 2026 (miercoles). Cadena SEPARADA del TM N20,
mismo dia, cross-references explicitas (regla cadenas por instrumento).
"""

import os
from docx import Document
from docx.shared import Pt, Inches

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(
    SCRIPT_DIR, "2026-06-10_Panel-Drawings-Reply.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    fields = [
        ("Date:", "June 10, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Stephane Gehant, Jeryl F. Regulacion, Adzlan Bin Abd Rahim, "
            "Nick Huta, Sadeep Irugalbandara, Billy Tan, Victor Gutierrez",
        ),
        (
            "Subject:",
            "RE: PLC-LCP Panel Drawings (25007-0046) — ADASA review "
            "position and expedited path",
        ),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / P22-TM-09-000-020-0",
        ),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Dear Eduardo,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "We proceed with the review via email, as you proposed. The "
        "full panel package review is in Transmittal N20 "
        "(P22-TM-09-000-020-0), issued today on the regular transmittal "
        "chain together with the rest of submittals 25007-0046 and "
        "25007-0047."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("One decision gates the enclosure fabrication. ").bold = True
    para.add_run(
        "The Outline Panel Drawing's own Panel Specification Sheet "
        "(sheet 2) specifies sheet steel painted RAL 7035 with "
        "protection class IP55, while your LCP Datasheet Rev B — "
        "submitted in the same 25007-0046 — specifies the nVent Hoffman "
        "FS66S enclosure in unpainted Stainless Steel 316L, NEMA "
        "4X/IP66, and the IFC Single Line Diagram labels the panel "
        "METAL CLAD, NEMA4X/IP66. The Technical Specification requires "
        "NEMA 4X or its IP equivalent for the coastal Taltal site."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Expedited path: ").bold = True
    para.add_run(
        "confirm in writing that the enclosure to be fabricated is the "
        "one specified in the LCP Datasheet Rev B (SS316L, NEMA "
        "4X/IP66) and re-issue the Outline as Rev B with the "
        "specification sheet aligned, the cooling arrangement "
        "reconciled with the protection class, and the actual panel "
        "weight declared. On that basis ADASA can release the enclosure "
        "fabrication against Rev B without waiting for a further "
        "review cycle. Dimensions and bills of material raised no "
        "objection."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "The Schematic Diagram Rev A and the LCP Datasheet Rev B are "
        "Approved as Noted (Code 2) — neither blocks the enclosure "
        "fabrication. Details and annotated PDFs are in Transmittal "
        "N20, Sections 2.5 to 2.7."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("We look forward to your written confirmation.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
