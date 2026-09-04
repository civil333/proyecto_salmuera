#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Transmittal N20 (P22-TM-09-000-020-0). Submittals 25007-0046 + 0047.
Fecha: 10 de Junio de 2026 (miercoles)
Veredicto TM N20: 3 - TO BE REVISED. Tally 8 Code 1 + 7 Code 2 + 5 Code 3.
20 documentos. 12 CC_ADASA adjuntos (5 Code 3 + 7 Code 2).

La respuesta al correo BW Water 10-Jun sobre los panel drawings va en un
correo separado (Reply-To a ese thread) emitido el mismo dia, con
cross-reference explicita a este transmittal.
"""

import os
from docx import Document
from docx.shared import Pt, Inches

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-06-10_Transmittal-N20.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    fields = [
        ("Date:", "June 10, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
            "Mauricio Vallejos, Jorge Valdes, Tanya Figueroa, Allan Valentos, "
            "Jeryl F. Regulacion, Sadeep Irugalbandara, Andrew Sia, "
            "Ghazi Ozair, Nick Huta, Marjan Arsovic, Gerald Ross, "
            "Andrew Zaske, Adzlan Bin Abd Rahim",
        ),
        (
            "Subject:",
            "ADASA – Taltal Brine Module: Technical Review Transmittal N20 "
            "(25007-0046 and 25007-0047)",
        ),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / P22-TM-09-000-020-0",
        ),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Dear BW Water Project Team,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Attached: ")
    para.add_run("Transmittal N20").bold = True
    para.add_run(
        " (P22-TM-09-000-020-0) covering submittals 25007-0046 and "
        "25007-0047 — twenty documents — for ADASA's technical review."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Verdict: 3 — To Be Revised.").bold = True
    para.add_run(
        " Tally: 8 Code 1, 7 Code 2, 5 Code 3. Drivers: PLC-LCP Outline "
        "Panel Drawing Rev A (Section 2.6), Alarm & Interlock List Rev B "
        "(Section 2.4), I/O List Rev 2 (Section 2.13), Instrumentation & "
        "Control Cable Schedule Rev 1 (Section 2.14) and NDE Plan Rev A "
        "(Section 2.17)."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Key positive closures in this cycle: ").bold = True
    para.add_run(
        "the five electrical IFC Rev 0 documents (Load List, Power Cable "
        "Schedule, Cable Datasheet, Single Line Diagram, Typical "
        "Installation Details) are accepted with every TM N19 condition "
        "verified as incorporated — a complete closure of the electrical "
        "package; "
    )
    para.add_run("Instrument List Rev E").bold = True
    para.add_run(" and ")
    para.add_run("Pressure Transmitter Datasheet Rev C").bold = True
    para.add_run(
        " are approved as-is, and the analyser power-supply item open "
        "since TM N14 (approximately 96 days) is closed by the I/O List "
        "Rev 2."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Drivers requiring action by BW Water:").bold = True
    aplicar_arial(para)

    drivers = [
        (
            "PLC-LCP Outline Panel Drawing Rev A (Section 2.6).",
            " Re-issue as Rev B with the Panel Specification Sheet "
            "aligned to the LCP Datasheet Rev B enclosure (Stainless "
            "Steel 316L, NEMA 4X/IP66), the cooling arrangement "
            "reconciled with the protection class, and the actual panel "
            "weight. This is the decision gating enclosure fabrication "
            "— our position on the expedited path is in today's reply "
            "to your panel drawings email.",
        ),
        (
            "Plant Control Philosophy Rev D (Section 3).",
            " Fourth consecutive transmittal cycle without delivery; "
            "the fourteen-day window stated in Transmittal N19 expired "
            "on 08-Jun. The I/O List conditional acceptance has "
            "reverted to Code 3 by its own terms, and the cabling and "
            "alarm documents remain gated. ADASA's reservation of "
            "remedies under Contract C-4300 stands.",
        ),
        (
            "Alarm & Interlock List Rev B (Section 2.4).",
            " Re-issue as Rev C reconciling the winding/bearing sensor "
            "assignments with the Instrument List Rev E on both pump "
            "trains, the confirmed 95 C bearing setpoint, and the pH "
            "and CIP Tank temperature tags.",
        ),
        (
            "Instrumentation & Control Cable Schedule Rev 1 "
            "(Section 2.14).",
            " Re-issue as Rev 2 correcting the duplicate item numbers "
            "on sheet 4 and the copy-pasted valve power descriptions.",
        ),
        (
            "NDE Plan Rev A (Section 2.17).",
            " Re-issue as Rev B incorporating the RO pressure vessel "
            "test scope per the 02-Jun waiver (factory hydrostatic, "
            "dossier, witness points) — also outstanding: ITP Rev C and "
            "the Hydrostatic, Preservation and FAT procedures with "
            "firm dates.",
        ),
    ]
    for bold_part, rest in drivers:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.25)
        para.add_run("• ")
        para.add_run(bold_part).bold = True
        para.add_run(rest)
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph(
        "Twelve annotated PDFs are attached: five for the Code 3 "
        "documents and seven for the Code 2 documents. The eight "
        "Code 1 — Approved documents carry no annotated PDF. Residual "
        "deliverables for IFC Rev 0 are listed in Section 3 of the "
        "transmittal."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "Please confirm receipt and the target dates for Plant Control "
        "Philosophy Rev D, the PLC-LCP Outline Rev B, the remaining "
        "Code 3 re-issues and the Section 3 deliverables."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("We look forward to your comments.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
