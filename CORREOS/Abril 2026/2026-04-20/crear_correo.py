"""
Correo: Weekly Procurement Reporting - Proposed Format (version ejecutiva)
Fecha: 20-Apr-2026
Codigo: CORREO-2026-04-20-WEEKLY-PROCUREMENT-PROTOCOL
Version: 8 (opening simplificado: solo propuesta de reporte; referencias a Kassim/Eduardo/action item BWW movidas al correo de minuta)
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_ALIGN_VERTICAL

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "2026-04-20_Weekly-Procurement-Reporting-Protocol.docx")


def set_font(run, name="Arial", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text="", bold=False, italic=False, size=11, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        set_font(run, bold=bold, italic=italic, size=size)
    return p


def add_separator(doc):
    sep = doc.add_paragraph()
    sep.paragraph_format.space_before = Pt(0)
    sep.paragraph_format.space_after = Pt(10)
    pPr = sep._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    pBdr.append(bottom)
    pPr.append(pBdr)


def set_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "999999")
        tcBorders.append(b)
    tcPr.append(tcBorders)


def set_cell_shading(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def write_cell(cell, text, bold=False, size=10, shade=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_borders(cell)
    if shade:
        set_cell_shading(cell, shade)


def add_protocol_table(doc):
    rows = [
        ("Cadence",
         "Every Monday, 10:00 AM Santiago time. First report: Monday, 27 April 2026."),
        ("Scope",
         "17 equipment lines (16 from the Baseline + RO Membranes declared today)."),
        ("Status",
         "C / E / D / N, aligned with the engineering Codes 1\u20134 already in use on this project."),
        ("ExWorks",
         "Contractual ExWorks point: BW Water Penang workshop, Malaysia. "
         "Module final destination: Taltal, Chile."),
    ]
    t = doc.add_table(rows=len(rows), cols=2)
    t.autofit = False
    t.columns[0].width = Cm(3.0)
    t.columns[1].width = Cm(13.0)
    for i, (label, value) in enumerate(rows):
        lc = t.cell(i, 0)
        vc = t.cell(i, 1)
        lc.width = Cm(3.0)
        vc.width = Cm(13.0)
        write_cell(lc, label, bold=True, size=10, shade="E8E8E8")
        write_cell(vc, value, size=10)


def main():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.0)

    # --- HEADER ---
    add_para(doc, "ADASA \u2014 Aguas de Antofagasta S.A.", bold=True, size=12, space_after=2)
    add_para(doc, "Contract C-4300 | Project 12803 \u2014 Taltal SWRO", size=10, space_after=8)

    # --- SUBJECT / TO / FROM / DATE / ATTACHMENT ---
    subj = add_para(doc, space_before=0, space_after=4)
    r = subj.add_run("Subject: ")
    set_font(r, bold=True, size=11)
    r2 = subj.add_run(
        "Project 12803 \u2014 Taltal SWRO | "
        "Weekly Procurement Reporting \u2014 Proposed Format"
    )
    set_font(r2, size=11)

    for label, value in [
        ("Date: ", "April 20, 2026"),
        (
            "To: ",
            "Andrew J. Zaske \u2014 Vice President-Americas, BW Water Americas Inc.; "
            "Eduardo Yamauchi \u2014 BW Water Americas Inc.; "
            "Fadey Kassim \u2014 Senior VP Global Operations, BW Water; "
            "Victor Gutierrez \u2014 ADASA",
        ),
        ("From: ", "Luis Rivera \u2014 Contract Administrator (ADASA)"),
        ("Attachment: ",
         "WEEKLY-PROCUREMENT-TRACKER_12803_v0.xlsx (ADASA template, pre-populated baseline)"),
    ]:
        p = add_para(doc, space_before=0, space_after=2)
        r = p.add_run(label)
        set_font(r, bold=True, size=11)
        r2 = p.add_run(value)
        set_font(r2, size=11)

    add_separator(doc)

    # --- SALUTATION ---
    add_para(doc, "Andrew, Eduardo, Fadey, Victor,", size=11, space_after=10)

    # --- OPENING: propuesta directa ---
    p = add_para(doc, space_before=0, space_after=10)
    parts = [
        ("Please find attached ", False),
        ("ADASA's proposed format for weekly procurement reporting", True),
        (", pre-populated against the Baseline Schedule issued by BW Water on ",
         False),
        ("05-Mar-2026", True),
        (".", False),
    ]
    for text, bold in parts:
        r = p.add_run(text)
        set_font(r, size=11, bold=bold)

    # --- PROTOCOL TABLE ---
    add_protocol_table(doc)

    # --- ONE-LINE: who completes what + status is ADASA's understanding ---
    p = add_para(doc, space_before=12, space_after=10)
    parts = [
        ("BW Water completes PO reference, date, supplier, lead time, ExWorks "
         "date and arrival Penang each week. The Status column reflects ", False),
        ("ADASA's understanding", True),
        (" of the position BW Water reported today; if any line does not match "
         "your actual procurement state, please flag it by return email ", False),
        ("without waiting for the weekly cadence", True),
        (".", False),
    ]
    for text, bold in parts:
        r = p.add_run(text)
        set_font(r, size=11, bold=bold)

    # --- PARALLEL SYSTEMS ---
    p = add_para(doc, space_before=0, space_after=10)
    r = p.add_run(
        "If anything in the template is redundant against your internal tracking, "
        "please propose the adjustment \u2014 one shared source of truth, not "
        "parallel systems."
    )
    set_font(r, size=11)

    # --- CONFIRMATION ---
    p = add_para(doc, space_before=0, space_after=12)
    parts = [
        ("Please confirm acceptance of this file as the weekly reporting instrument, "
         "and the BW Water owner of the delivery, by end of business ", False),
        ("Wednesday, 22\u00a0April\u00a02026", True),
        (".", False),
    ]
    for text, bold in parts:
        r = p.add_run(text)
        set_font(r, size=11, bold=bold)

    # --- CLOSING ---
    add_para(doc, "Best regards,", size=11, space_before=6, space_after=12)

    sig_lines = [
        ("Luis Rivera", True),
        ("Project Engineer | Contract Administrator", False),
        ("ADASA \u2014 Aguas de Antofagasta S.A.", False),
        ("Contract C-4300 | Project 12803 \u2014 Taltal SWRO", False),
    ]
    for text, bold in sig_lines:
        add_para(doc, text, bold=bold, size=10, space_before=0, space_after=1)

    doc.save(OUTPUT)
    print(f"Correo generado: {OUTPUT}")


if __name__ == "__main__":
    main()
