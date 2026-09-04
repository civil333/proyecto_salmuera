"""
Correo: Layout Revisions — Commitment Not Met
Fecha: 01-Apr-2026
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "2026-04-01_Layout-Missed-Deadline.docx")

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def set_font(run, name="Arial", size=10, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_para(doc, text="", size=10, bold=False, space_after=4, color=None,
             align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if text:
        run = p.add_run(text)
        set_font(run, size=size, bold=bold, color=color)
    return p

def add_header_line(doc, label, value, size=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    r1 = p.add_run(label)
    set_font(r1, size=size, bold=True)
    r2 = p.add_run(value)
    set_font(r2, size=size)
    return p

def add_section_title(doc, text, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_font(run, size=size, bold=True)
    run.font.underline = True
    return p

def add_bullet_item(doc, bold_text, rest=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.5)
    r0 = p.add_run("\u2013  ")
    set_font(r0)
    r1 = p.add_run(bold_text)
    set_font(r1, bold=True)
    if rest:
        r2 = p.add_run(rest)
        set_font(r2)

def add_numbered_item(doc, num, bold_text, rest=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.5)
    r0 = p.add_run(f"{num}.  ")
    set_font(r0)
    r1 = p.add_run(bold_text)
    set_font(r1, bold=True)
    if rest:
        r2 = p.add_run(rest)
        set_font(r2)

# ---------------------------------------------------------------------------
# Documento
# ---------------------------------------------------------------------------
doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(10)

# ---------------------------------------------------------------------------
# ENCABEZADO
# ---------------------------------------------------------------------------
add_para(doc, "CORREO ELECTR\u00d3NICO", size=12, bold=True,
         color=(31, 56, 100), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

add_header_line(doc, "Fecha:   ", "1 de Abril de 2026")
add_header_line(doc, "De:      ", "Luis Rivera \u2014 Contract Administrator (ADASA)")
add_header_line(doc, "Para:    ", "Eduardo Yamauchi \u2014 BW Water Americas Inc.")
add_header_line(doc, "CC:      ",
    "Jeryl F. Regulacion, Adzlan Bin Abd Rahim, Sadeep Irugalbandara, Nick Huta; "
    "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
    "Mauricio Vallejos, Jorge Valdes")
add_header_line(doc, "Asunto:  ",
    "Project 12803 \u2014 Taltal SWRO | Layout Revisions \u2014 Commitment Not Met | April 1, 2026")

doc.add_paragraph()

# ---------------------------------------------------------------------------
# SALUDO + APERTURA
# ---------------------------------------------------------------------------
add_para(doc, "Dear Eduardo,", bold=True, space_after=6)
add_para(doc,
    "During our meeting on March 31, BW Water confirmed delivery of three revised "
    "drawings by April 1, 2026. As of today\u2019s close of business, none of the "
    "following documents have been received:",
    space_after=6)

# Lista de documentos pendientes
for doc_entry in [
    "Equipment Layout (P22-DWG-09-005-003) \u2014 Rev. C",
    "Piping Layout (P22-DWG-09-005-004) \u2014 Rev. B",
    "Tie-In Point Layout (P22-DWG-09-005-005) \u2014 Rev. C",
    "3D Model (CIP Area) \u2014 per Change Order 25007-PL-0001 Rev.\u00a01",
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(f"\u2013  {doc_entry}")
    set_font(r)

add_para(doc, "", space_after=4)
add_para(doc,
    "The 3D model was included in the scope of Change Order 25007-PL-0001 Rev.\u00a01 "
    "(2-week leadtime from acceptance) and has equally not been received.",
    space_after=6)
add_para(doc, "For reference, the required corrections per document are as follows:", space_after=4)

# ---------------------------------------------------------------------------
# SECCIÓN 1 — Equipment Layout Rev. C
# ---------------------------------------------------------------------------
add_section_title(doc, "Equipment Layout (P22-DWG-09-005-003) \u2014 Rev. C")
add_bullet_item(doc, "Tank relocated to opposite side",
                " to provide chemical loading access.")

# ---------------------------------------------------------------------------
# SECCIÓN 2 — Piping Layout Rev. B
# ---------------------------------------------------------------------------
add_section_title(doc, "Piping Layout (P22-DWG-09-005-004) \u2014 Rev. B")
add_bullet_item(doc,
    "CIP equipment and antiscalant dosing system consolidated within a single "
    "external footprint.",
    " Maximum footprint: 3,500\u00a0mm (per TM N5 OBS-01, February 23, 2026).")

# ---------------------------------------------------------------------------
# SECCIÓN 3 — Tie-In Point Layout Rev. C
# ---------------------------------------------------------------------------
add_section_title(doc, "Tie-In Point Layout (P22-DWG-09-005-005) \u2014 Rev. C")

tiein_items = [
    ("Space for pipe rack confirmed,",
     " or routing shown directly from container."),
    ("Dosing tank access resolved.", ""),
    ("Antiscalant loading routed directly to dosing tank",
     " \u2014 no carrier pump is budgeted."),
    ("Connection flanges for module relocation cutover included.", ""),
    ("Container elevation view included",
     " to confirm tie-in flange locations."),
]

for i, (bold_text, rest) in enumerate(tiein_items, 1):
    add_numbered_item(doc, i, bold_text, rest)

# ---------------------------------------------------------------------------
# SECCIÓN 4 — 3D Model
# ---------------------------------------------------------------------------
add_section_title(doc, "3D Model (CIP Area) \u2014 Change Order 25007-PL-0001 Rev.\u00a01")
add_bullet_item(doc,
    "3D model updated to reflect the new CIP area layout",
    " consistent with Equipment Layout Rev.\u00a0C and Piping Layout Rev.\u00a0B.")

# ---------------------------------------------------------------------------
# CIERRE
# ---------------------------------------------------------------------------
add_para(doc, "", space_after=4)
add_para(doc,
    "ADASA requests delivery of the three revised drawings and the updated 3D model "
    "by EOB April 3, 2026. Delay in these revisions directly conditions the structural "
    "frames PO window (opening April 6\u201310) and the CIP piping procurement schedule.",
    space_after=8)
add_para(doc, "Please confirm receipt and provide a revised delivery date.", space_after=10)
add_para(doc, "Best regards,", space_after=4)
add_para(doc, "Luis Rivera", bold=True, space_after=2)
add_para(doc, "Project Engineer | Contract Administrator", space_after=2)
add_para(doc, "ADASA \u2014 Aguas de Antofagasta S.A.", space_after=2)
add_para(doc, "Contract C-4300 | Project 12803 \u2014 Taltal SWRO", space_after=0)

# ---------------------------------------------------------------------------
# GUARDAR
# ---------------------------------------------------------------------------
doc.save(OUTPUT)
print(f"Generado: {OUTPUT}")
