"""
Correo: FINAL NOTICE — Layout Submittals, Operating Weight Data & Procurement Log
Fecha: 15-Apr-2026
Codigo: CORREO-2026-04-15-FINAL-NOTICE-LAYOUTS
Version: 1
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "2026-04-15_Final-Notice-Layouts-Procurement.docx")


def set_font(run, name="Arial", size=11, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text="", bold=False, size=11, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        set_font(run, bold=bold, size=size)
    return p


def add_separator(doc):
    sep = doc.add_paragraph()
    sep.paragraph_format.space_before = Pt(0)
    sep.paragraph_format.space_after = Pt(10)
    pPr = sep._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_bullet(doc, text, size=11, bold=False, indent_cm=1.27):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(indent_cm)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    r = p.add_run("\u2013  " + text)
    set_font(r, size=size, bold=bold)
    return p


def add_lettered_item(doc, letter, text, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-0.6)
    r = p.add_run(f"{letter})  ")
    set_font(r, size=size, bold=True)
    r2 = p.add_run(text)
    set_font(r2, size=size)
    return p


def main():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.0)

    # --- HEADER ---
    add_para(doc, "ADASA \u2014 Aguas de Antofagasta S.A.", bold=True, size=12, space_after=2)
    add_para(doc, "Contract C-4300 | Project 12803 \u2014 Taltal SWRO", size=10, space_after=8)

    # --- SUBJECT / FROM / TO / CC / DATE ---
    subj = add_para(doc, space_before=0, space_after=4)
    r = subj.add_run("Subject: ")
    set_font(r, bold=True, size=11)
    r2 = subj.add_run(
        "FINAL NOTICE: Project 12803 \u2014 Taltal SWRO | "
        "Layout Submittals, Operating Weight Data & Procurement Log"
    )
    set_font(r2, size=11)

    for label, value in [
        ("Date: ", "April 15, 2026"),
        ("To: ", "Eduardo Yamauchi \u2014 BW Water Americas Inc."),
        (
            "CC: ",
            "Jeryl F. Regulacion, Adzlan Bin Abd Rahim, Sadeep Irugalbandara, Nick Huta; "
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
            "Mauricio Vallejos, Jorge Valdes",
        ),
        ("From: ", "Luis Rivera \u2014 Contract Administrator (ADASA)"),
    ]:
        p = add_para(doc, space_before=0, space_after=2)
        r = p.add_run(label)
        set_font(r, bold=True, size=11)
        r2 = p.add_run(value)
        set_font(r2, size=11)

    add_separator(doc)

    # --- SALUTATION ---
    add_para(doc, "Dear Eduardo,", size=11, space_after=10)

    # --- OPENING ---
    p = add_para(doc, space_before=0, space_after=10)
    r = p.add_run(
        "We are writing for the sixth time regarding the outstanding layout "
        "revisions and procurement documentation. Our communications of "
        "April\u00a01, 6, 10, 13 and 14 have received no response."
    )
    set_font(r, size=11)

    p = add_para(doc, space_before=0, space_after=10)
    r = p.add_run(
        "Earlier today, ADASA issued Transmittal\u00a0N14 covering eleven "
        "documents from deliveries E26 through E29. Seven were approved "
        "without observations, four approved as noted. Twelve prior "
        "observations were closed. This brings the technical review current "
        "through April\u00a015 \u2014 however, the layout and procurement "
        "items below continue to block project progress."
    )
    set_font(r, size=11)

    # --- SECTION 1: LAYOUTS ---
    add_para(
        doc,
        "1. Layout Submittals \u2014 14 Calendar Days Overdue",
        bold=True, size=11, space_before=4, space_after=6,
    )

    p = add_para(doc, space_before=0, space_after=6)
    r = p.add_run(
        "The six layout documents requested on April\u00a06 remain outstanding:"
    )
    set_font(r, size=11)

    layouts = [
        "Equipment Layout P22-DWG-09-005-003 Rev\u00a0C (formal submittal \u2014 "
        "advance copy received April\u00a02 does not replace the established review process)",
        "Piping Layout P22-DWG-09-005-004 Rev\u00a0B",
        "Tie-In Point Layout P22-DWG-09-005-005 Rev\u00a0C",
        "3D Model of CIP Area (per Change Order 25007-PL-0001)",
        "Instrument Location Layout P22-DWG-09-008-001 Rev\u00a0C",
        "Grounding Point Layout P22-DWG-09-007-003 Rev\u00a0C",
    ]
    for item in layouts:
        add_bullet(doc, item, size=11)

    p = add_para(doc, space_before=6, space_after=10)
    r = p.add_run(
        "These documents were committed for April\u00a01 and have been the "
        "subject of four consecutive missed deadlines."
    )
    set_font(r, size=11)

    # --- SECTION 2: OPERATING WEIGHT ---
    add_para(
        doc,
        "2. Operating Weight Data \u2014 Unanswered Since April 14",
        bold=True, size=11, space_before=4, space_after=6,
    )

    p = add_para(doc, space_before=0, space_after=6)
    parts = [
        ("As noted yesterday, the Equipment Layout Rev\u00a0C advance copy ", False),
        ("no longer includes the Operating Weight table", True),
        (" that was present in Rev\u00a0A (seven items, 34,739\u00a0lb / "
         "15,758\u00a0kg). ADASA requires updated operating weights for "
         "foundation design, seismic verification per NCh\u00a02369 Zone\u00a03, "
         "and the module lifting plan. The formal submittal must restore this "
         "table with current values.", False),
    ]
    for text, bold in parts:
        r = p.add_run(text)
        set_font(r, bold=bold, size=11)

    # --- SECTION 3: PROCUREMENT LOG ---
    add_para(
        doc,
        "3. Procurement Log \u2014 Not Received",
        bold=True, size=11, space_before=4, space_after=6,
    )

    p = add_para(doc, space_before=0, space_after=10)
    r = p.add_run(
        "The updated Procurement Log against the Baseline Schedule "
        "05-Mar-2026, covering all 16 equipment lines with current PO "
        "status, has been requested since April\u00a06. The Structural "
        "Frames\u00a0/\u00a0Supports PO window closed on April\u00a010 "
        "without a confirmed purchase order; manufacturing capacity extends "
        "only through May\u00a05."
    )
    set_font(r, size=11)

    # --- DEADLINE AND NEXT STEPS ---
    add_para(
        doc,
        "Deadline and Next Steps",
        bold=True, size=11, space_before=6, space_after=6,
    )

    p = add_para(doc, space_before=0, space_after=6)
    r = p.add_run(
        "ADASA requests a written response by end of business "
        "April\u00a016,\u00a02026, confirming:"
    )
    set_font(r, size=11)

    add_lettered_item(
        doc, "a",
        "Delivery dates for the six layout documents listed above "
        "(PDF\u00a0+\u00a0DWG, plus STEP/IGES for the 3D model)."
    )
    add_lettered_item(
        doc, "b",
        "The Operating Weight table, either as an updated Equipment Layout "
        "submittal or as a standalone data transmittal."
    )
    add_lettered_item(
        doc, "c",
        "The updated Procurement Log with PO status for all 16 baseline lines."
    )

    p = add_para(doc, space_before=8, space_after=10)
    parts2 = [
        ("In the absence of a response by the stated deadline, ADASA will "
         "issue a ", False),
        ("formal delay notification to the project sponsor under "
         "Contract\u00a0C-4300 on April\u00a017,\u00a02026", True),
        (". This notification has been pending since our communication "
         "of April\u00a013.", False),
    ]
    for text, bold in parts2:
        r = p.add_run(text)
        set_font(r, bold=bold, size=11)

    # --- CLOSING ---
    add_para(doc, "Best regards,", size=11, space_before=10, space_after=12)

    sig_lines = [
        ("Luis Rivera", True),
        ("Project Engineer | Contract Administrator", False),
        ("ADASA \u2014 Aguas de Antofagasta S.A.", False),
        ("Contract C-4300 | Project 12803 \u2014 Taltal SWRO", False),
    ]
    for text, bold in sig_lines:
        add_para(doc, text, bold=bold, size=10, space_before=0, space_after=1)

    doc.save(OUTPUT)
    print(f"Correo generado: {OUTPUT}")


if __name__ == "__main__":
    main()
