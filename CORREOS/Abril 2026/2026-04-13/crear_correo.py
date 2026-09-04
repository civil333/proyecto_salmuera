"""
Correo: URGENT — Structural Frames PO Window Closed | Response Required Today
Fecha: 13-Apr-2026
Codigo: CORREO-2026-04-13-FOLLOWUP-OUTSTANDING-COMMITMENTS
Version: 2 (ejecutiva)
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "2026-04-13_Follow-Up-Outstanding-Commitments.docx")


def set_font(run, name="Arial", size=11, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text="", bold=False, size=11, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        set_font(run, bold=bold, size=size)
    return p


def add_separator(doc):
    sep = doc.add_paragraph()
    sep.paragraph_format.space_before = Pt(0)
    sep.paragraph_format.space_after = Pt(10)
    pPr = sep._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    pBdr.append(bottom)
    pPr.append(pBdr)


def main():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.0)

    # --- HEADER ---
    add_para(doc, "ADASA \u2014 Aguas de Antofagasta S.A.", bold=True, size=12, space_after=2)
    add_para(doc, "Contract C-4300 | Project 12803 \u2014 Taltal SWRO", size=10, space_after=8)

    # --- SUBJECT / FROM / TO / CC / DATE ---
    subj = add_para(doc, space_before=0, space_after=4)
    r = subj.add_run("Subject: ")
    set_font(r, bold=True, size=11)
    r2 = subj.add_run(
        "URGENT: Project 12803 \u2014 Taltal SWRO | "
        "Structural Frames PO Window Closed | Response Required Today"
    )
    set_font(r2, size=11)

    for label, value in [
        ("Date: ", "April 13, 2026"),
        ("To: ", "Eduardo Yamauchi \u2014 BW Water Americas Inc."),
        (
            "CC: ",
            "Jeryl F. Regulacion, Adzlan Bin Abd Rahim, Sadeep Irugalbandara, Nick Huta; "
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
            "Mauricio Vallejos, Jorge Valdes",
        ),
        ("From: ", "Luis Rivera \u2014 Contract Administrator (ADASA)"),
    ]:
        p = add_para(doc, space_before=0, space_after=2)
        r = p.add_run(label)
        set_font(r, bold=True, size=11)
        r2 = p.add_run(value)
        set_font(r2, size=11)

    add_separator(doc)

    # --- SALUTATION ---
    add_para(doc, "Dear Eduardo,", size=11, space_after=10)

    # --- OPENING PARAGRAPH ---
    # Mixed formatting: key facts highlighted in bold
    p_open = doc.add_paragraph()
    p_open.paragraph_format.space_before = Pt(0)
    p_open.paragraph_format.space_after = Pt(10)

    parts = [
        ("Our email of April 10 has received no response. As of today, the ", False),
        ("Structural Frames / Supports PO window has closed without a confirmed purchase order", True),
        (
            " \u2014 manufacturing runs only through May 5, leaving "
            "22 calendar days with no procurement margin. "
            "Layout revisions remain ", False,
        ),
        ("12 calendar days overdue", True),
        (
            " from the April\u00a01 deadline; "
            "this is the fourth consecutive commitment not met.",
            False,
        ),
    ]
    for text, bold in parts:
        r = p_open.add_run(text)
        set_font(r, bold=bold, size=11)

    # --- REQUIRED RESPONSE ---
    add_para(
        doc,
        "We require your response today, April 13, confirming:",
        size=11, space_after=6,
    )

    items = [
        (
            "a) ",
            "A delivery date for the six outstanding layout documents "
            "(PDF + DWG, plus STEP/IGES for the 3D model).",
        ),
        (
            "b) ",
            "The updated Procurement Log against the Baseline Schedule 05-Mar-2026, "
            "covering all 16 lines with current PO status.",
        ),
    ]
    for label, body in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(1.0)
        r_label = p.add_run(label)
        set_font(r_label, bold=True, size=11)
        r_body = p.add_run(body)
        set_font(r_body, size=11)

    add_para(doc, space_after=8)

    # --- ESCALATION CLAUSE ---
    p_esc = add_para(doc, space_before=0, space_after=16)
    r_esc = p_esc.add_run(
        "If we do not hear from you by end of business today, ADASA will formally notify "
        "the project sponsor and initiate the contractual delay notification process "
        "under Contract C-4300."
    )
    set_font(r_esc, bold=True, size=11)

    # --- SIGNATURE ---
    add_para(doc, "Best regards,", size=11, space_after=8)
    add_para(doc, "Luis Rivera", bold=True, size=11, space_after=2)
    add_para(doc, "Project Engineer | Contract Administrator", size=11, space_after=2)
    add_para(doc, "ADASA \u2014 Aguas de Antofagasta S.A.", size=11, space_after=2)
    add_para(doc, "Contract C-4300 | Project 12803 \u2014 Taltal SWRO", size=11, space_after=0)

    doc.save(OUTPUT)
    print(f"Generado: {OUTPUT}")


if __name__ == "__main__":
    main()
