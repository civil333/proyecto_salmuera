#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo de cobertura: Reply-To al thread del RFI-002 (25007-RO-RFI-0002),
emitido por Billy Tan (BW Water / Engineering) el 01-Jul-2026.

ADASA devuelve el form RFI con la respuesta en "Replied Information" (adjunto)
y confirma en el cuerpo, cover-only: la construccion estandar del fabricante es
aceptable (exterior + gland plates SS316L per datasheet aprobado; internos
galvanizados/laminados; enclosure mantiene NEMA 4X/IP66) y el Outline Panel
Drawing debe reemitirse. Sin proponer reunion (transaccional).

Fecha: 7 de Julio de 2026 (martes). Cadena del propio RFI, separada de los
transmittals y del thread del Notice of Delay del PLC (correo D3).

Clon de CORREOS/Junio 2026/2026-06-24/crear_correo_rfi001_reply.py.
"""

import os
import sys
from docx import Document
from docx.shared import Pt, Inches

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
from docx_metadata import apply_core_properties, fix_app_xml

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-07-07_RFI-002-Enclosure-Reply.docx")
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def crear_correo():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    fields = [
        ("Date:", "July 7, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Billy Tan — BW Water / Engineering"),
        (
            "CC:",
            "Eduardo Yamauchi, Stephane Gehant, Jeryl F. Regulacion, "
            "Adzlan Bin Abd Rahim, Nick Huta, Sadeep Irugalbandara, "
            "Victor Gutierrez",
        ),
        (
            "Subject:",
            "RE: RFI 25007-RO-RFI-0002 — LCP Enclosure Material Specification "
            "(NEMA 4X / IP66)",
        ),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / RFI 25007-RO-RFI-0002 / "
            "P22-ET-09-000-001-0",
        ),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Dear Billy,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "We acknowledge RFI 25007-RO-RFI-0002 and return it with ADASA's "
        "reply completed in the Replied Information section (attached)."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(
        "ADASA confirms the manufacturer's standard construction is "
        "acceptable. "
    ).bold = True
    para.add_run(
        "The external enclosure body, door, roof, rear panel, plinth and the "
        "gasketed gland plates shall be Stainless Steel 316L, unpainted, per "
        "the approved Datasheet of Local Control Panel (LCP); the internal "
        "mounting components may be galvanized or cold-rolled steel per the "
        "manufacturer's standard design, provided the enclosure maintains "
        "NEMA 4X / IP66. The PLC-LCP Outline Panel Drawing shall be re-issued "
        "to distinguish the external SS316L components from the internal ones, "
        "restrict the Gray RAL 7035 finishing to internal surfaces, and state "
        "the protection class in full as NEMA 4X / IP66."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "This confirmation is conditional on the external enclosure and gland "
        "plates being SS316L and on the enclosure maintaining a protection "
        "class not lower than NEMA 4X; the observations on the Outline Panel "
        "Drawing remain in force until the corrected revision is submitted and "
        "approved. The completed RFI form sets out the full reply."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("We look forward to your acknowledgement.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.add_paragraph()
    para = doc.add_paragraph()
    run = para.add_run(
        "Attachment: RFI 25007-RO-RFI-0002 — completed reply form "
        "(Replied Information section)."
    )
    run.italic = True
    run.font.name = "Arial"
    run.font.size = Pt(10)

    apply_core_properties(
        doc,
        title="RFI 25007-RO-RFI-0002 - ADASA Reply (cover)",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="RFI 25007-RO-RFI-0002 - LCP Enclosure Material",
        comments="",
        language="en-US",
    )
    doc.save(OUTPUT_FILE)
    fix_app_xml(OUTPUT_FILE, application="Microsoft Office Word",
                company="Aguas Antofagasta")
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
