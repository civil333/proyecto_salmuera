#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo de cobertura del Transmittal N28 (P22-TM-09-000-028-0) a BW Water.
Cadena regular de transmittals. Ingles. Document() directo, sin template ADASA.
Veredicto 2 - Approved as Noted (4 Code 2): la familia de Control por fin
entregada completa; se pide re-emitir los 4 como conjunto coordinado a Rev 0 con
dos definiciones fijadas (mapeo winding/bearing + trip de vibracion alcanzable).
Re-escala los 3 vencidos (HP/LP Rev D, O&M Rev B, UHPRO Structural Rev B),
deadline Vie 31-Jul-2026. Estado: BORRADOR.
"""
import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-07-20_Transmittal-N28.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"
DOWNLOAD_LINK = "https://lrg.synology.me:6501/d/s/196vffdOgHuhxkUIjBvvagziqyuR4aQr/ByGtK2Zyvp6HBEouqMYWxFJ0Lc-1c7oq-67dg5dMBXA0"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang"); rPr.append(el)
    el.set(qn("w:val"), lang)


def fijar_idioma(doc, lang=LANG):
    try:
        _set_lang(doc.styles["Normal"].element.get_or_add_rPr(), lang)
    except KeyError:
        pass
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph(text); aplicar_arial(p, size); return p


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for text, bold in segs:
        r = p.add_run(text); r.bold = bold
        r.font.name = "Arial"; r.font.size = Pt(size)
    return p


def add_bullet(doc, text, size=11):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run("•  " + text); r.font.name = "Arial"; r.font.size = Pt(size)
    return p


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    run = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts"); rf.set(qn("w:ascii"), "Arial"); rf.set(qn("w:hAnsi"), "Arial")
    rpr.append(rf)
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; run.append(t)
    h.append(run); paragraph._p.append(h)


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8); s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8); s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "July 20, 2026"),
        ("From:", f"{CONTACTO} — Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water"),
        ("CC:", "Andrew Sia, Victor Gutierrez, Jeryl F. Regulacion, Jorge Guevara, "
                "Ronald Pellejero, Stephane Gehant"),
        ("Subject:", "25007 Taltal - Technical Review Transmittal N28 - Plant Control "
                     "documents"),
        ("Ref:", "P22-TM-09-000-028-0 / submittals 25007-0065 (E65) and 25007-0066 (E66)"),
    ]
    for label, value in fields:
        p = doc.add_paragraph(); p.add_run(label).bold = True
        p.add_run(f" {value}"); aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_segments(doc, [
        ("Transmittal N28 (", False), ("P22-TM-09-000-028-0", True),
        (") is attached, covering the four Plant Control documents (submittals "
         "25007-0065 and 25007-0066): Control Philosophy Rev E, Alarm and Interlock List "
         "Rev C, Control and Sequence Chart Rev A (first issue) and IO List Rev 5.",
         False),
    ])
    blank(doc)

    add_segments(doc, [
        ("Verdict: 2 - Approved as Noted (four Code 2).", True),
        (" The Control family is delivered complete; each document is correct in what it "
         "governs and issues at IFC Rev 0 with the noted corrections. No new revision "
         "cycle.", False),
    ])
    blank(doc)

    add_para(doc, "Please re-issue the four as a coordinated set at Rev 0, fixing two "
                  "definitions across all of them:")
    add_bullet(doc,
               "winding and bearing mapping - winding TE-09-001, bearing TE-09-002 (as "
               "the Alarm List, IO List and Instrument List already hold; the Control "
               "Philosophy still reverses it, a safety-sensor mapping);")
    add_bullet(doc,
               "a reachable HP pump vibration trip - the Alarm List sets 10.0 mm/s above "
               "the transmitter's 8.9 range, so it cannot fire.")
    blank(doc)

    add_segments(doc, [
        ("Still open, and now the path to the Rev 0 control set, by ", False),
        ("Friday 31 July 2026", True),
        (": the ", False), ("HP and LP Pressure Test Procedure", True),
        (" Rev D (its Line List orders 75 bar on a PVC line), the ", False),
        ("O&M Manual", True), (" Rev B, and the ", False),
        ("UHPRO Structural Calculation", True), (" Rev B.", False),
    ])
    blank(doc)

    p = doc.add_paragraph()
    p.add_run("Download - this transmittal and the four annotated PDFs: ").bold = True
    aplicar_arial(p)
    add_hyperlink(p, DOWNLOAD_LINK, DOWNLOAD_LINK)
    blank(doc)

    add_para(doc, "We look forward to your comments.")
    blank(doc)
    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph(); p.add_run(CONTACTO).bold = True; aplicar_arial(p)
    for line in ["Infrastructure Engineering Lead",
                 "Desalination Projects Department",
                 "ADASA — Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)
    docx_metadata.apply_core_properties(
        doc, title="Taltal - Technical Review Transmittal N28",
        author="Luis Rivera Gonzalez", last_modified_by="Luis Rivera Gonzalez",
        subject="Transmittal N28 - Plant Control documents",
        comments="Aguas de Antofagasta S.A.", language="en-US")
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas Antofagasta")
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
