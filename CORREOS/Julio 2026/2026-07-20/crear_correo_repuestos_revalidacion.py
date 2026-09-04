#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo ADASA -> BW Water (Eduardo Yamauchi + Andrea Frezzi): rechazo por forma de
la cotizacion de repuestos a 2 anos recibida el 15-Jul-2026 y solicitud de
reemision como documento comercial controlado.

Reply-To al thread propio "Formal Re-validation of Spare Parts Quotation C4300"
(abierto el 26-Feb-2026). Cadena SEPARADA de la de transmittals, de la minuta y
del calendario Hold/Witness.

Decision del usuario (20-Jul): el punto de la valvula solenoide ya se acepto en
la conversacion previa, asi que NO se reabre aqui. El correo se acota a UN eje:
el documento recibido no es una oferta comercial y no puede procesarse. Cuatro
requisitos minimos para la reemision. Las observaciones de precio y alcance
(cambio Service Kit -> Mechanical Seal, alzas de 43% a 372%, diferencia de USD
1.240 en la oferta Sep-2025, repuestos DN15) se reservan para la respuesta a la
reemision: se discuten contra un documento formal, no contra una planilla.

El cuarto requisito (Scope reconciliation) es el unico agregado sobre el borrador
del usuario: sin el, la reemision llega formalmente correcta y tecnicamente
equivocada, y se pierde un ciclo completo contra el ex-works Penang 09-10 Sep.

Sustento interno: PROGRAMA y CONTRATO/REPUESTOS DE 2 ANOS/_ANALISIS_REPUESTOS_15JUL.md

Ingles (BW Water). Document() directo, sin template ADASA. Metadatos limpios.
Estado: BORRADOR.
"""

import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402  (metadatos limpios, global sec. 2.3)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(
    SCRIPT_DIR, "2026-07-20_Spare-Parts-Quotation-Formal-Response.docx"
)
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang_in_rPr(rPr, lang):
    lang_el = rPr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rPr.append(lang_el)
    lang_el.set(qn("w:val"), lang)


def fijar_idioma_documento(doc, lang=LANG):
    try:
        rPr = doc.styles["Normal"].element.get_or_add_rPr()
        _set_lang_in_rPr(rPr, lang)
    except KeyError:
        pass
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            _set_lang_in_rPr(run._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    para = doc.add_paragraph(text)
    aplicar_arial(para, size)
    return para


def add_requisito(doc, lead, rest, size=11):
    """Requisito con encabezado en negrita, sangria izquierda."""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.3)
    para.add_run(lead).bold = True
    para.add_run(rest)
    aplicar_arial(para, size)
    return para


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # ---- HEADER -------------------------------------------------------------
    fields = [
        ("Date:", "July 20, 2026"),
        ("From:", f"{CONTACTO} — Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi, Andrea Frezzi — BW Water"),
        ("CC:",
         "Jorge Guevara, Victor Gutierrez — ADASA; Stephane Gehant — BW Water"),
        ("Subject:",
         "RE: Formal Re-validation of Spare Parts Quotation C4300"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    blank(doc)

    # ---- CUERPO -------------------------------------------------------------
    add_para(doc, "Dear Eduardo and Andrea,")
    blank(doc)

    add_para(doc, "Thank you for the updated file received on 15 July 2026.")
    blank(doc)

    add_para(
        doc,
        "The document as received cannot be processed by Aguas de Antofagasta "
        "S.A. It is a working spreadsheet, not a formal commercial offer, and "
        "therefore does not support our internal review and purchase order "
        "cycle. Please reissue the Recommended Two-Year Spare Parts quotation "
        "under Contract C-4300 as a controlled commercial document meeting the "
        "following minimum requirements:",
    )
    blank(doc)

    add_requisito(
        doc,
        "Formal issuance — ",
        "BW Water letterhead and corporate identity, quotation number with "
        "revision index, issue date, and the signature of the legal "
        "representative or of the authorized commercial / quotations manager, "
        "stating name, position and date.",
    )
    blank(doc)

    add_requisito(
        doc,
        "Delivery lead time — ",
        "ex-works manufacturing lead time per line item, expressed in calendar "
        "weeks from PO date, plus transit time and named delivery point. "
        "Long-lead items (membranes, HP pump and energy recovery components, "
        "instrumentation) must be identified explicitly as such.",
    )
    blank(doc)

    add_requisito(
        doc,
        "Commercial conditions — ",
        "validity period of the offer, currency, applicable Incoterm 2020 with "
        "named place, and payment terms.",
    )
    blank(doc)

    add_requisito(
        doc,
        "Scope reconciliation — ",
        "each line referenced to the tag of the equipment or valve it serves, "
        "quoting the manufacturer part number of the item actually supplied. "
        "Several lines in the file received correspond to equipment other than "
        "that being manufactured.",
    )
    blank(doc)

    add_para(
        doc,
        "Please confirm receipt and issue the document by Friday, 31 July 2026.",
    )
    blank(doc)

    # ---- CIERRE -------------------------------------------------------------
    add_para(doc, "Best regards,")
    blank(doc)

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in [
        "Infrastructure Engineering Lead",
        "Desalination Projects Department",
        "ADASA — Aguas de Antofagasta S.A.",
        "lrivera@aguasantofagasta.cl",
    ]:
        add_para(doc, line)

    fijar_idioma_documento(doc, LANG)

    # ---- metadatos limpios (global sec. 2.3) --------------------------------
    docx_metadata.apply_core_properties(
        doc,
        title="Taltal - Spare Parts Quotation C4300 - reissue as formal document",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="Taltal SWRO - Recommended two-year spare parts quotation",
        comments="Aguas de Antofagasta S.A.",
        language="en-US",
    )
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas Antofagasta")
    print(f"Correo generado: {OUTPUT_FILE}")
    print(f"Idioma del documento: {LANG}")


if __name__ == "__main__":
    crear_correo()
