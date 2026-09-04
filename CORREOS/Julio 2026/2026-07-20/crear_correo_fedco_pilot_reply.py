#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo ADASA -> BW Water (Eduardo Yamauchi): respuesta a la invitacion de FEDCO
para asistir al pilot test de las bombas HP y turbocargadores en la fabrica de
FEDCO (EE.UU.).

Thread propio "25007 Taltal: FEDCO - Pilot Testing" (correo de Yamauchi del
15-Jul). Cadena separada del calendario Hold/Witness, de los TM y del root-cause
del atraso Fedco. Reply-To a este thread.

Posicion (consistente con el punto 3 del correo del calendario H/W, borrador
20-Jul): ADASA, via Bureau Veritas, atestigua solo el FAT en el taller de Penang;
NO asiste al pilot test en la fabrica FEDCO (EE.UU.); a cambio, toda la
documentacion de las pruebas FEDCO debe ir en el dossier de entrega del equipo,
disponible para el inspector en Penang.

Ingles (BW Water). Document() directo, sin template ADASA, sin tabla, transaccional
(~110 palabras). Metadatos limpios. Estado: BORRADOR.
"""

import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402  (metadatos limpios, global sec. 2.3)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-07-20_FEDCO-Pilot-Test-Reply.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang_in_rPr(rPr, lang):
    lang_el = rPr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rPr.append(lang_el)
    lang_el.set(qn("w:val"), lang)


def fijar_idioma_documento(doc, lang=LANG):
    try:
        rPr = doc.styles["Normal"].element.get_or_add_rPr()
        _set_lang_in_rPr(rPr, lang)
    except KeyError:
        pass
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            _set_lang_in_rPr(run._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    para = doc.add_paragraph(text)
    aplicar_arial(para, size)
    return para


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # ---- HEADER -------------------------------------------------------------
    fields = [
        ("Date:", "July 20, 2026"),
        ("From:",
         f"{CONTACTO} — Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water"),
        ("CC:",
         "Victor Gutierrez, Jorge Guevara, Ronald Pellejero — ADASA; "
         "Stephane Gehant — BW Water"),
        ("Subject:", "RE: 25007 Taltal: FEDCO - Pilot Testing"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    blank(doc)

    # ---- CUERPO -------------------------------------------------------------
    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_para(
        doc,
        "Thank you for the note on FEDCO's pilot testing programme for the RO "
        "high-pressure pumps and turbochargers.",
    )
    blank(doc)

    add_para(
        doc,
        "Aguas de Antofagasta will not attend the pilot test at FEDCO's "
        "facility. Consistent with our third-party shop inspection "
        "arrangement, ADASA's witnessing, through Bureau Veritas, is limited "
        "to the Factory Acceptance Test at the Penang workshop; we will not "
        "travel to witness pump or turbocharger tests at the FEDCO facility in "
        "the United States.",
    )
    blank(doc)

    add_para(
        doc,
        "We do, however, require the complete test documentation for the "
        "supplied pumps and turbochargers to be included in the equipment "
        "delivery dossier, together with the rest of the documentation: the "
        "pilot and factory test protocols, acceptance criteria, measured "
        "results and certificates, made available to the inspector at the "
        "Penang workshop for review.",
    )
    blank(doc)

    add_para(doc, "We look forward to your comments.")
    blank(doc)

    # ---- CIERRE -------------------------------------------------------------
    add_para(doc, "Best regards,")
    blank(doc)

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in [
        "Infrastructure Engineering Lead",
        "Desalination Projects Department",
        "ADASA — Aguas de Antofagasta S.A.",
        "lrivera@aguasantofagasta.cl",
    ]:
        add_para(doc, line)

    fijar_idioma_documento(doc, LANG)

    # ---- metadatos limpios (global sec. 2.3) --------------------------------
    docx_metadata.apply_core_properties(
        doc,
        title="Taltal - FEDCO Pilot Testing - ADASA reply",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="Taltal SWRO - FEDCO pilot test attendance",
        comments="Aguas de Antofagasta S.A.",
        language="en-US",
    )
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas Antofagasta")
    print(f"Correo generado: {OUTPUT_FILE}")
    print(f"Idioma del documento: {LANG}")


if __name__ == "__main__":
    crear_correo()
