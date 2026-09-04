#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo ADASA -> BW Water: el dossier de fabricacion y pruebas para la revision
de Bureau Veritas no fue entregado, y tres frentes conexos de la inspeccion de
taller que arranca en Penang el lunes 27-Jul-2026.

Reply-all al thread "RE: Taltal - Designation of Third-Party Shop Inspector and
Inspection Schedule" (cadena de inspeccion, separada de los transmittals y de la
replica a la minuta 07-Jul). Ingles. Document() directo, sin template ADASA
(CLAUDE.md 3.4).

Fundamento (el correo NO se apoya en el compromiso de fecha, que no consta por
escrito de BW Water, sino en la obligacion contractual):
  - ET P22-ET-09-000-001-0 Seccion 7: dossier de fabricacion y pruebas de todos
    los equipos electricos y electromecanicos, con certificados de fabricacion
    de los aceros especiales super duplex, maximo a los 90 dias de la
    adjudicacion. Master Register item 65 = NOT DELIVERED.
  - PIE Base P22-IT-09-000-001-0 (Inspection and Testing Base Plan): item 7.6
    revision del dossier preliminar (ADASA: R); items 8.3 aprobacion final del
    dossier (Hold) y 8.4 Liberacion para Despacho (Hold, soporta el 40%).
  - Contenido: lista de documentacion requerida presentada por BW Water en la
    Quality Kick-off Meeting (QKOM).

Frentes conexos, verificados en HITO BUREAU VERITAS/CORREOS VBV-BW/ (23-24 Jul):
  - Inspection Request 001 adjunto el ITP en Rev C (superada; Rev 0 es la
    aprobada Codigo 1 en el TM N26) y las copias con las anotaciones de revision
    de ADASA, no los documentos limpios vigentes.
  - Cubre una sola jornada (28-Jul, PMI super duplex) de las tres semanales
    comunicadas a BV el 21-Jul; el WQT + WPS/PQR se movio al 07-Ago. Aviso de
    4 dias contra los 30 de la BAE Clausula 37.
  - Los puntos de reconciliacion (V5/V6, FAT y Dispatch Release, Kick-off)
    siguen sin respuesta escrita; por eso ADASA no pudo confirmar V2 a V6.

NO reclamar los tres procedimientos que estaban en Codigo 3 (009 Rev C, 010
Rev D, 011 Rev B): los tres estan en Codigo 2. Estado: BORRADOR.
"""

import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402  (metadatos limpios, global sec. 2.3)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(
    SCRIPT_DIR, "2026-07-25_BWWater-Inspection-Dossier-Request.docx"
)
CONTACTO = "Luis Rivera González"
LANG = "en-US"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang_in_rPr(rPr, lang):
    lang_el = rPr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rPr.append(lang_el)
    lang_el.set(qn("w:val"), lang)


def fijar_idioma_documento(doc, lang=LANG):
    try:
        rPr = doc.styles["Normal"].element.get_or_add_rPr()
        _set_lang_in_rPr(rPr, lang)
    except KeyError:
        pass
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            _set_lang_in_rPr(run._element.get_or_add_rPr(), lang)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_lang_in_rPr(run._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    para = doc.add_paragraph(text)
    aplicar_arial(para, size)
    return para


def blank(doc):
    doc.add_paragraph()


def add_numbered(doc, number, lead, rest, size=11):
    """Item numerado manual: 'N. ' + lead en negrita + resto normal."""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.25)
    para.paragraph_format.space_after = Pt(6)
    r0 = para.add_run(f"{number}. ")
    r0.font.name = "Arial"
    r0.font.size = Pt(size)
    r1 = para.add_run(lead)
    r1.bold = True
    r1.font.name = "Arial"
    r1.font.size = Pt(size)
    r2 = para.add_run(" " + rest)
    r2.font.name = "Arial"
    r2.font.size = Pt(size)
    return para


def add_bullet_lead(doc, lead, rest, size=11):
    """Vineta con guion: lead en negrita + resto normal, sangria de item."""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.5)
    para.paragraph_format.space_after = Pt(6)
    r0 = para.add_run("- ")
    r0.font.name = "Arial"
    r0.font.size = Pt(size)
    r1 = para.add_run(lead)
    r1.bold = True
    r1.font.name = "Arial"
    r1.font.size = Pt(size)
    r2 = para.add_run(" " + rest)
    r2.font.name = "Arial"
    r2.font.size = Pt(size)
    return para


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # ---- HEADER -------------------------------------------------------------
    fields = [
        ("Date:", "July 25, 2026"),
        ("From:",
         f"{CONTACTO} — Infrastructure Engineering Lead (ADASA)"),
        ("To:",
         "Eduardo Yamauchi, Magdier Arias, Stephane Gehant — BW Water"),
        ("CC:",
         "Victor Gutierrez, Jorge Guevara, Ronald Pellejero — ADASA; "
         "Mohd Adnin Zulkifli, Lokman Hakim Mat — BW Water Penang"),
        ("Subject:",
         "RE: Taltal - Designation of Third-Party Shop Inspector and "
         "Inspection Schedule"),
        ("Ref:",
         "Contract C-4300 / BAE 12803 / Technical Specification "
         "P22-ET-09-000-001-0, Section 7 / Inspection and Testing Base Plan "
         "P22-IT-09-000-001-0 / Inspection Request 001"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    blank(doc)

    # ---- CUERPO -------------------------------------------------------------
    add_para(doc, "Dear Eduardo, dear Magdier,")
    blank(doc)

    add_para(
        doc,
        "We acknowledge Inspection Request 001 for the super duplex positive "
        "material identification on 28 July. Bureau Veritas will attend. Four "
        "points need to be closed for the surveillance period to run as "
        "planned.",
    )
    blank(doc)

    add_numbered(
        doc, 1, "The fabrication and testing dossier is still outstanding.",
        "It is a contract deliverable under the Technical Specification "
        "P22-ET-09-000-001-0, Section 7, due within 90 days of award and "
        "including the mill certificates of the super duplex materials. It "
        "is also the document ADASA reviews under item 7.6 of the Inspection "
        "and Testing Base Plan P22-IT-09-000-001-0, and its content is the "
        "documentation list BW Water presented at the Quality Kick-off "
        "Meeting. The date of Friday 24 July, recorded in our email of 21 "
        "July, passed without delivery. We ask for it in two steps:",
    )
    add_bullet_lead(
        doc, "By close of business in Penang on Monday 27 July:",
        "the dossier index required by item 7.6, plus the records supporting "
        "the first inspection week. These are the material test reports and "
        "traceability of the super duplex spools to be examined, the "
        "calibration certificate of the PMI instrument with the technician "
        "certificates, and the welding procedure and welder qualification "
        "records for the joints already welded since mid-July. Welding "
        "started before the surveillance window opened, so those records are "
        "the only means of verification left to the inspector.",
    )
    add_bullet_lead(
        doc, "By Friday 31 July:",
        "the complete preliminary dossier, structured against that index.",
    )
    add_para(
        doc,
        "Until ADASA has reviewed the dossier there is no basis for items 8.3 "
        "and 8.4, the final dossier approval and the Release for Dispatch "
        "that supports the 40% milestone. Please confirm as well that the "
        "FEDCO factory test records and certificates are incorporated into "
        "it, as set out in our previous email.",
    )
    blank(doc)

    add_numbered(
        doc, 2, "The documents issued to the inspector are not the current "
        "revisions.",
        "Inspection Request 001 attached the Inspection and Test Plan at Rev "
        "C, and both attachments are the ADASA-annotated review copies. Rev C "
        "is superseded by Rev 0, returned Code 1. Please reissue the current "
        "clean revisions to Bureau Veritas before the session of 28 July; the "
        "inspection package ADASA distributed on 21 July is the reference.",
    )
    add_numbered(
        doc, 3, "Week 1 scope and notification lead time.",
        "The plan communicated on 21 July sets three surveillance days per "
        "week. Request 001 notifies one, and the welder and procedure "
        "qualification test has moved to 7 August. We need days 2 and 3 of "
        "the first week notified with their scope, and confirmation of "
        "whether that test now falls within the second week. Request 001 was "
        "also issued four days ahead of the inspection, against the 30 days "
        "required by BAE Clause 37. ADASA accepts it for this first visit "
        "without setting a precedent, and asks that the remaining "
        "notifications observe the contractual notice.",
    )
    add_numbered(
        doc, 4, "The reconciliation points remain unanswered.",
        "Visits 5 and 6 placed on the fixed schedule, the FAT and Dispatch "
        "Release dates with confirmation that the reduced duration preserves "
        "the minimum FAT scope, and the Kick-off Meeting. Without that answer "
        "ADASA cannot confirm visits 2 to 6 to the inspector, and the "
        "surveillance period is already running. Please reply in writing "
        "before the weekly meeting of Tuesday 28 July.",
    )
    blank(doc)

    add_para(
        doc,
        "Any deviation from the fixed schedule must be notified to ADASA in "
        "writing and in advance, stating its cause, so that the affected Hold "
        "and Witness points can be reassessed.",
    )
    blank(doc)

    # ---- CIERRE -------------------------------------------------------------
    add_para(doc, "Best regards,")
    blank(doc)

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in [
        "Infrastructure Engineering Lead",
        "Desalination Projects Department",
        "ADASA — Aguas de Antofagasta S.A.",
        "lrivera@aguasantofagasta.cl",
    ]:
        add_para(doc, line)

    fijar_idioma_documento(doc, LANG)

    # ---- metadatos limpios (global sec. 2.3) --------------------------------
    docx_metadata.apply_core_properties(
        doc,
        title="Fabrication and Testing Dossier - Shop Inspection at Penang",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="Taltal SWRO - inspection dossier and week 1 coordination",
        comments="Aguas de Antofagasta S.A.",
        language="en-US",
    )
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas Antofagasta")
    print(f"Correo generado: {OUTPUT_FILE}")
    print(f"Idioma del documento: {LANG}")


if __name__ == "__main__":
    crear_correo()
