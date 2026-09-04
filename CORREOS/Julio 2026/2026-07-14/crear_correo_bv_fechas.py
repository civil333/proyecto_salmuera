#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo a Bureau Veritas Chile (Luis Rodrigo Arcila) con las fechas de las
visitas de inspeccion en el taller de BW Water (Penang) para el Modulo SWRO
Taltal, en respuesta a su consulta para reservar inspectores y pasajes.

Base: NT P22-NT-09-000-002-0 (notificada a BW Water el 08-Jul-2026) validada
contra el status update de BW Water del 14-Jul-2026 (Recovery Schedule +
Week 28 + tracker de procurement). Solo la Visita 1 (27-31 Jul) es firme hoy;
el resto queda sujeto al calendario Hold/Witness que BW Water entrega el
viernes 17-Jul (revision en reunion del martes 21-Jul). FAT: base 07-12 Sep
con posible desplazamiento de una semana (14-19 Sep).

Correo en espanol, idioma fijado es-CL (CLAUDE.md Seccion 3.4). Document()
directo, sin template ADASA. Estado: BORRADOR.
"""

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-07-14_BV-Fechas-Visitas-Penang.docx")
CONTACTO = "Luis Rivera"
LANG = "es-CL"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang_in_rPr(rPr, lang):
    lang_el = rPr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rPr.append(lang_el)
    lang_el.set(qn("w:val"), lang)
    lang_el.set(qn("w:eastAsia"), lang)
    lang_el.set(qn("w:bidi"), lang)


def fijar_idioma_documento(doc, lang=LANG):
    try:
        rPr = doc.styles["Normal"].element.get_or_add_rPr()
        _set_lang_in_rPr(rPr, lang)
    except KeyError:
        pass
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            _set_lang_in_rPr(run._element.get_or_add_rPr(), lang)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_lang_in_rPr(run._element.get_or_add_rPr(), lang)


def aplicar_arial_table(table, size=10):
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(size)


def add_table_with_header(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = header
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True

    for i, row_data in enumerate(rows, start=1):
        for j, value in enumerate(row_data):
            table.rows[i].cells[j].text = str(value)

    if widths:
        for j, w in enumerate(widths):
            for row in table.rows:
                row.cells[j].width = w

    aplicar_arial_table(table, size=10)
    return table


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    fields = [
        ("Fecha:", "14 de julio de 2026"),
        ("De:", f"{CONTACTO} — DIO ADASA (Aguas de Antofagasta S.A.)"),
        ("Para:", "Luis Rodrigo Arcila — Bureau Veritas Chile"),
        ("CC:", "Víctor Gutiérrez (ADASA)"),
        ("Asunto:",
         "Fechas de visitas de inspección en taller de fabricación (Penang) — "
         "Módulo SWRO Taltal"),
        ("Ref:",
         "Cotización BV 600049 / ITP P22-BA-09-000-004 Rev 0 / "
         "Nota Técnica ADASA P22-NT-09-000-002-0"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Estimado Luis Rodrigo:")
    aplicar_arial(para)
    doc.add_paragraph()

    # Accion principal arriba, en negrita
    para = doc.add_paragraph()
    para.add_run(
        "La Visita 1 queda confirmada para la semana del lunes 27 al viernes "
        "31 de julio; recomendamos gestionar desde ya la movilización del "
        "inspector a Penang.").bold = True
    para.add_run(
        " Para el resto del calendario enviamos las ventanas planificadas, "
        "que pueden variar; la confirmación final irá al cierre de la "
        "próxima semana.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "El 8 de julio notificamos formalmente al fabricante la designación "
        "de Bureau Veritas como Tercero Inspector de ADASA (Nota Técnica "
        "P22-NT-09-000-002-0), junto con el calendario de inspección anclado "
        "a los puntos Hold (H) y Witness (W) del Plan de Inspección y Ensayos "
        "aprobado. Las ventanas comunicadas, contrastadas con el avance de "
        "fabricación reportado el 14 de julio, quedan como sigue:")
    aplicar_arial(para)
    doc.add_paragraph()

    add_table_with_header(
        doc,
        ["Visita", "Ventana 2026", "Alcance principal", "Estado"],
        [
            ["1", "27–31 jul",
             "Fabricación spools Super Duplex: PMI, calificación de "
             "soldadores, inspección visual de soldaduras (W)",
             "CONFIRMADA — movilizar"],
            ["2", "3–7 ago",
             "Ensayos no destructivos (RT/PT/UT) de soldaduras de alta "
             "presión (W); control dimensional del skid (W)",
             "Planificada — puede variar en días"],
            ["3", "10–14 ago",
             "Hidrostáticas de piping: PVC 7,5 bar (W) y Super Duplex "
             "135 bar (H); hidrostática RO Vessel (H); pintura (W)",
             "Planificada — puede variar en días"],
            ["4", "17–21 ago",
             "Término de pintura; montaje final de piping e instrumentos "
             "sobre el skid (W)",
             "Planificada — puede variar en días"],
            ["5", "24–28 ago",
             "Posicionamiento de equipos dentro del contenedor (W); "
             "alineamiento de piping de alta presión (W)",
             "Planificada — puede variar en días"],
            ["6", "31 ago–4 sep",
             "Inspección CSC del contenedor; posicionamiento bomba de alta "
             "presión y turbocargadores (W); integración del tablero (W)",
             "Planificada — puede variar en días"],
            ["FAT", "7–12 sep (base)",
             "Atestiguamiento del Factory Acceptance Test completo, hasta el "
             "Dispatch Release (H)",
             "Base 7–12 sep; posible desplazamiento de una semana "
             "(14–19 sep)"],
        ],
        widths=[Inches(0.55), Inches(1.0), Inches(2.9), Inches(2.05)],
    )

    doc.add_paragraph()

    para = doc.add_paragraph(
        "Las ventanas de las Visitas 2 a 6 pueden variar. El fabricante debe "
        "entregarnos este viernes 17 de julio el calendario consolidado de "
        "notificaciones de puntos Hold y Witness, que revisaremos en la "
        "reunión semanal del martes 21 de julio; con ese insumo les "
        "enviaremos la confirmación final el viernes 24 de julio. Los "
        "ajustes que esperamos son del orden de días para cada visita de "
        "tres jornadas, en torno a las ventanas indicadas, no "
        "reprogramaciones de semanas completas.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph(
        "Para el FAT, la semana del 7 al 12 de septiembre se mantiene como "
        "base de planificación. La bomba de alta presión y los dos "
        "turbocargadores llegan a Penang el 2 de septiembre y las pruebas "
        "funcionales solo pueden correr una vez instalados, por lo que el "
        "cierre del FAT puede desplazarse una semana, hacia el 14 al 19 de "
        "septiembre. Sugerimos considerar esa banda al planificar la "
        "disponibilidad del inspector; la fecha definitiva se fijará con la "
        "notificación formal del punto Hold correspondiente.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Quedamos atentos a sus comentarios.")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Saludos cordiales,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in [
        "Departamento de Ingeniería y Optimización (DIO)",
        "ADASA — Aguas de Antofagasta S.A.",
        "lrivera@aguasantofagasta.cl",
    ]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    fijar_idioma_documento(doc, LANG)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")
    print(f"Idioma del documento: {LANG}")


if __name__ == "__main__":
    crear_correo()
