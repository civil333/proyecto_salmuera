#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo ADASA -> BW Water: respuesta al plan de despacho / "RFQ for packing cost"
(Faizah Bardan via Eduardo Yamauchi, 20-Jul-2026).

- Reply-To al thread "20.25.6501 TALTAL: RFQ FOR PACKING COST". Cadena SEPARADA
  del correo del panel PLC/LCP (mismo dia).
- Idioma: ingles (BW Water). Patron Document() directo per CLAUDE.md 3.4.
- Postura: confirma el alcance logistico EXW de ADASA (contenedores, haulage,
  freight forwarding/aduana/flete) y coordina el loading schedule; RESERVA el
  costo de embalaje: bajo EXW (Incoterms) el packing de exportacion en el punto
  Ex-Works es alcance/costo del vendedor. Pide a BW aclarar si pretende cargar
  un packing cost a ADASA y su base contractual.
- "20.25.6501" = numero de proyecto interno de BW Water para Taltal (no una
  cotizacion) -> no se trata como referencia comercial.
- Fuente unica del cuerpo: 2026-07-21_Packing-EXW-Response_Descripcion.md
"""

import os
import sys

from docx import Document
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402  (metadatos limpios, global sec. 2.3)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(
    SCRIPT_DIR, "2026-07-21_Packing-EXW-Response.docx"
)
CONTACTO = "Luis Rivera"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_para(doc, text, size=11):
    para = doc.add_paragraph(text)
    aplicar_arial(para, size)
    return para


def add_lead(doc, lead, rest, size=11):
    """Parrafo con un tramo inicial en bold (bold inline)."""
    para = doc.add_paragraph()
    r1 = para.add_run(lead)
    r1.bold = True
    r2 = para.add_run(rest)
    for run in (r1, r2):
        run.font.name = "Arial"
        run.font.size = Pt(size)
    return para


def add_bullet(doc, text, size=11):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.25)
    run = para.add_run("•  " + text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    return para


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # ---- HEADER -------------------------------------------------------------
    fields = [
        ("Date:", "July 21, 2026"),
        ("From:", f"{CONTACTO} - Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi - BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Victor Gutierrez, Stephane Gehant, Faizah Binti Bardan, "
            "Lokman Hakim Bin Mat",
        ),
        ("Subject:", "RE: 20.25.6501 TALTAL: RFQ FOR PACKING COST"),
        (
            "Ref:",
            "Contract C-4300 / BAE 12803 / Incoterm EXW (Penang) / "
            "BW Water dispatch plan 20-Jul-2026 "
            "(UHPRO system, loose equipment, CIP tank)",
        ),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    blank(doc)

    # ---- CUERPO -------------------------------------------------------------
    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_para(
        doc,
        "Thank you for the dispatch plan for the Taltal equipment forwarded on "
        "20 July 2026, covering the three units for shipment from Penang under "
        "EXW: one 40' HC container with the containerized UHPRO system, one "
        "20' GP container with the loose equipment and accessories, and one "
        "20' flat rack with the CIP tank.",
    )
    blank(doc)

    add_para(
        doc,
        "On the EXW allocation, ADASA confirms it will arrange, from the BW "
        "Water yard onward: the 20' GP container delivered by side-loader "
        "haulage for loading the loose equipment; haulage for the 40' HC "
        "container and the 20' flat rack; and the freight forwarding, inland "
        "transport, customs clearance and ocean shipment to the final "
        "destination. To coordinate this, please send the loading schedule and "
        "the vessel closing date so we can position the container two to three "
        "days beforehand and align the haulage with the one-day crane you will "
        "arrange.",
    )
    blank(doc)

    add_lead(
        doc,
        "On the packing cost, ",
        "we note that under EXW (Incoterms) the export packing and marking of "
        "the goods at the Ex-Works point fall within the seller's scope and "
        "cost. ADASA's position is that packing the equipment for export is "
        "part of BW Water's supply. Before ADASA can consider any packing "
        "charge, please confirm whether a packing cost is intended to be "
        "charged to ADASA and, if so, state the contractual basis for it. "
        "ADASA reserves its rights in the meantime; this does not affect the "
        "logistics coordination described above.",
    )
    blank(doc)

    add_para(doc, "We look forward to the loading schedule and to your "
                  "confirmation on the above.")
    blank(doc)

    # ---- CIERRE -------------------------------------------------------------
    add_para(doc, "Best regards,")
    blank(doc)

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA - Aguas de Antofagasta S.A."]:
        add_para(doc, line)

    # ---- metadatos limpios (global sec. 2.3) --------------------------------
    docx_metadata.apply_core_properties(
        doc,
        title="Taltal Equipment Dispatch (EXW) - ADASA Response",
        author="Luis Rivera Gonzalez",
        last_modified_by="Luis Rivera Gonzalez",
        subject="Equipment dispatch from Penang under EXW - ADASA response",
        comments="Correo ADASA a BW Water - respuesta al plan de despacho EXW "
                 "y reserva del packing cost (20-Jul-2026).",
        category="Correspondencia",
        language="en-US",
    )
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas de Antofagasta S.A.")
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
