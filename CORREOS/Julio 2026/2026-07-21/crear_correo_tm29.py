#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo de cobertura del Transmittal N29 (P22-TM-09-000-029-0) a BW Water.
Cadena regular de transmittals. Ingles. Document() directo, sin template ADASA.
Veredicto 2 - Approved as Noted (2 Code 2 + 2 Code 1). TM de recuperacion:
cierra los dos Code 3 vencidos (HP/LP Pressure Test Rev D, UHPRO Structural Calc
Rev B) y recibe el Outline y el Line List a IFC Rev 0. Pendiente: la familia de
Control (N28) como conjunto coordinado a Rev 0, deadline Vie 31-Jul-2026.
Estado: BORRADOR.
"""
import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.expanduser("~/.claude/skills/template-adasa"))
import docx_metadata  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-07-21_Transmittal-N29.docx")
CONTACTO = "Luis Rivera González"
LANG = "en-US"
DOWNLOAD_LINK = "https://lrg.synology.me:6501/d/s/199TcR9nfzQGBsdiqjFx1KBITOy7Czm5/IKYQplKAUYslUaQQSd6OflbA0m4xU-9--Vb9gGZr7XQ0"


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def _set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang"); rPr.append(el)
    el.set(qn("w:val"), lang)


def fijar_idioma(doc, lang=LANG):
    try:
        _set_lang(doc.styles["Normal"].element.get_or_add_rPr(), lang)
    except KeyError:
        pass
    for p in doc.paragraphs:
        for r in p.runs:
            _set_lang(r._element.get_or_add_rPr(), lang)


def add_para(doc, text, size=11):
    p = doc.add_paragraph(text); aplicar_arial(p, size); return p


def add_segments(doc, segs, size=11):
    p = doc.add_paragraph()
    for text, bold in segs:
        r = p.add_run(text); r.bold = bold
        r.font.name = "Arial"; r.font.size = Pt(size)
    return p


def add_bullet(doc, text, size=11):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run("•  " + text); r.font.name = "Arial"; r.font.size = Pt(size)
    return p


def add_bullet_lead(doc, lead, rest, size=11):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
    r0 = p.add_run("•  " + lead); r0.bold = True
    r0.font.name = "Arial"; r0.font.size = Pt(size)
    r1 = p.add_run(rest); r1.font.name = "Arial"; r1.font.size = Pt(size)
    return p


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    run = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts"); rf.set(qn("w:ascii"), "Arial"); rf.set(qn("w:hAnsi"), "Arial")
    rpr.append(rf)
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; run.append(t)
    h.append(run); paragraph._p.append(h)


def blank(doc):
    doc.add_paragraph()


def crear_correo():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8); s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8); s.right_margin = Inches(0.8)

    fields = [
        ("Date:", "July 21, 2026"),
        ("From:", f"{CONTACTO} - Infrastructure Engineering Lead (ADASA)"),
        ("To:", "Eduardo Yamauchi - BW Water"),
        ("CC:", "Andrew Sia, Victor Gutierrez, Jeryl F. Regulacion, Jorge Guevara, "
                "Ronald Pellejero, Stephane Gehant"),
        ("Subject:", "25007 Taltal - Technical Review Transmittal N29 - pressure test, "
                     "structural calculation, panel outline and line list"),
        ("Ref:", "P22-TM-09-000-029-0 / submittal 25007-0067 (E67)"),
    ]
    for label, value in fields:
        p = doc.add_paragraph(); p.add_run(label).bold = True
        p.add_run(f" {value}"); aplicar_arial(p)
    blank(doc)

    add_para(doc, "Dear Eduardo,")
    blank(doc)

    add_segments(doc, [
        ("Transmittal N29 (", False), ("P22-TM-09-000-029-0", True),
        (") is attached, covering the four re-submitted documents of submittal "
         "25007-0067: the HP and LP Pressure Test Procedure Rev D, the UHPRO Structural "
         "Calculation Report Rev B, the PLC/LCP Outline Panel Drawing Rev 0 and the Line "
         "List Rev 0.", False),
    ])
    blank(doc)

    add_segments(doc, [
        ("Verdict: 2 - Approved as Noted (two Code 2, two Code 1).", True),
        (" A recovery transmittal: the two documents held at Code 3 in N26 and N27 are "
         "resolved and verified. Disposition:", False),
    ])
    blank(doc)

    add_bullet_lead(doc, "HP and LP Pressure Test Procedure Rev D - Code 2. ",
                    "The 75 bar / PVC test-pressure error is corrected; state the "
                    "applicable ASME Section V and B31.3 edition.")
    add_bullet_lead(doc, "UHPRO Structural Calculation Report Rev B - Code 2. ",
                    "Seismic check confirmed (NCh 2369:2003, utilization 0.454 < 1.0); "
                    "issue a single, non-duplicated report file and complete the comment "
                    "sheet.")
    add_bullet_lead(doc, "PLC/LCP Outline Panel Drawing Rev 0 - Code 1. ",
                    "Accepted at IFC; enclosure gate closed (RFI-002 incorporated, Single "
                    "Line Diagram re-issued Rev 1).")
    add_bullet_lead(doc, "Line List Rev 0 - Code 1. ",
                    "Accepted; record the DA-PVC-DN65-09-016 design-pressure correction in "
                    "the revision history.")
    blank(doc)

    add_segments(doc, [("Near-term commitments:", True)])
    add_bullet_lead(doc, "By this Friday, 24 July 2026: ",
                    "deliver the shop-inspection dossier for Bureau Veritas' review, as "
                    "committed; the first surveillance week at the Penang workshop begins "
                    "27 July.")
    add_bullet_lead(doc, "By Friday, 31 July 2026: ",
                    "re-issue the Plant Control family from Transmittal N28 (Control "
                    "Philosophy Rev E, Alarm and Interlock List Rev C, Control and Sequence "
                    "Chart Rev A) as a coordinated set at Rev 0, with the Operating and "
                    "Maintenance Manual Rev B to follow.")
    blank(doc)

    p = doc.add_paragraph()
    p.add_run("Download - this transmittal and the annotated PDFs: ").bold = True
    aplicar_arial(p)
    add_hyperlink(p, DOWNLOAD_LINK, DOWNLOAD_LINK)
    blank(doc)

    add_para(doc, "We look forward to your comments.")
    blank(doc)
    add_para(doc, "Best regards,")
    blank(doc)
    p = doc.add_paragraph(); p.add_run(CONTACTO).bold = True; aplicar_arial(p)
    for line in ["Infrastructure Engineering Lead",
                 "Desalination Projects Department",
                 "ADASA - Aguas de Antofagasta S.A.",
                 "lrivera@aguasantofagasta.cl"]:
        add_para(doc, line)

    fijar_idioma(doc, LANG)
    docx_metadata.apply_core_properties(
        doc, title="Taltal - Technical Review Transmittal N29",
        author="Luis Rivera Gonzalez", last_modified_by="Luis Rivera Gonzalez",
        subject="Transmittal N29 - pressure test, structural calc, outline, line list",
        comments="Correo de cobertura del Transmittal N29 a BW Water.",
        language="en-US")
    doc.save(OUTPUT_FILE)
    docx_metadata.fix_app_xml(OUTPUT_FILE, company="Aguas de Antofagasta S.A.")
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
