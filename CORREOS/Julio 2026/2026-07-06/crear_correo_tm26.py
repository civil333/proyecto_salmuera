#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correo: Transmittal N26 (P22-TM-09-000-026-0). Submittals 25007-0057 (E57) a
25007-0062 (E62), 9 documentos. Fecha: 6 de Julio de 2026 (lunes).
Veredicto TM N26: 3 - TO BE REVISED. Tally 2 Code 1 + 3 Code 2 + 4 Code 3.
Version EJECUTIVA/directa (correo corto). Entrega via link de descarga Synology
(hipervinculo). Re-escala los vencidos con deadline consolidado viernes 10-Jul-2026.
English, BW Water. Sin reuniones. Cadena = thread regular de transmittals.
"""

import os
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "2026-07-06_Transmittal-N26.docx")
CONTACTO = "Luis Rivera"
DOWNLOAD_LINK = (
    "https://lrg.synology.me:6501/d/s/18x6gGzUV5gmg0nCp8tMGG5RF2wcjxD0/"
    "Dikk5IT6Dqqsuh78ZiYs2JFfpEmvYlwi-Yrxg1S1jVA0"
)


def aplicar_arial(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rpr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def crear_correo():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    fields = [
        ("Date:", "July 6, 2026"),
        ("From:", f"{CONTACTO} — Project Engineer (ADASA)"),
        ("To:", "Eduardo Yamauchi — BW Water Americas Inc. (PMO Leader)"),
        (
            "CC:",
            "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, "
            "Mauricio Vallejos, Jorge Valdes, Tanya Figueroa, Allan Valentos, "
            "Jeryl F. Regulacion, Sadeep Irugalbandara, Andrew Sia, "
            "Ghazi Ozair, Nick Huta, Marjan Arsovic, Gerald Ross, "
            "Andrew Zaske, Adzlan Bin Abd Rahim",
        ),
        (
            "Subject:",
            "ADASA – Taltal Brine Module: Technical Review Transmittal N26 "
            "(25007-0057 to 25007-0062)",
        ),
        ("Ref:", "Contract C-4300 / BAE 12803 / P22-TM-09-000-026-0"),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(f" {value}")
        aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph("Dear BW Water Project Team,")
    aplicar_arial(para)
    doc.add_paragraph()

    # 1. Entrega + verdicto (una linea)
    para = doc.add_paragraph()
    para.add_run("Transmittal N26").bold = True
    para.add_run(
        " (P22-TM-09-000-026-0) covers submittals 25007-0057 through 25007-0062 — "
        "nine documents. "
    )
    para.add_run("Verdict: 3 — To Be Revised").bold = True
    para.add_run(" (2 Code 1, 3 Code 2, 4 Code 3).")
    aplicar_arial(para)
    doc.add_paragraph()

    # 2. Link de descarga (hipervinculo)
    para = doc.add_paragraph()
    para.add_run("The transmittal and the seven annotated PDFs are available for "
                 "download at: ")
    aplicar_arial(para)
    add_hyperlink(para, DOWNLOAD_LINK, DOWNLOAD_LINK)
    doc.add_paragraph()

    # 3. Drivers -> re-issue (directo)
    para = doc.add_paragraph(
        "Four documents drive the verdict and require re-issue: the RO Vessel "
        "Hydrostatic Test Procedure and the HP and LP Pressure Test Procedure "
        "(Rev B) still state no binding test pressure — the hydrostatic form still "
        "shows 45.5 bar against the required 1,980 psi; the UHPRO Structural "
        "Calculation Report (Rev A) omits the anchorage of the main process "
        "equipment; and the GA of the Antiscalant Dosing Tank (Rev B) still defers "
        "its NCh 2369 anchor loads. The Inspection and Test Plan Rev 0 and the NDE "
        "Plan Rev C are Approved; the PLC and HMI, the CIP Flushing Tank and the "
        "Antiscalant Pump Skid documents are Approved as Noted. Per-document actions "
        "are in Section 2."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # 4. Vencidos + deadline
    para = doc.add_paragraph()
    para.add_run("Outstanding deliverables. ").bold = True
    para.add_run(
        "The items escalated for July 3 remain open — the Grounding Layout Rev F, "
        "the FAT and SAT table, the Plant Control Philosophy children and the three "
        "installation-route plans — as does the Painting Procedure. Please deliver "
        "all of these by Friday, July 10, 2026."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    # 5. Cierre
    para = doc.add_paragraph(
        "Please confirm receipt and the target dates for the noted items. We look "
        "forward to your comments."
    )
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(CONTACTO).bold = True
    aplicar_arial(para)
    for line in ["Project Engineer", "ADASA — Aguas de Antofagasta S.A."]:
        para = doc.add_paragraph(line)
        aplicar_arial(para)

    doc.save(OUTPUT_FILE)
    print(f"Correo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    crear_correo()
