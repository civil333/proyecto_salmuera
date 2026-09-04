#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para generar correo DOCX - RE: Project Taltal - Project Management Point of Contact
Fecha: 29 de enero de 2026
Proposito: Respuesta al cambio de PM + Mensaje directo a Eduardo Yamauchi sobre retraso
"""

import sys
import os
from pathlib import Path

# Agregar skill template-adasa al path
skill_path = str(
    Path(__file__).resolve().parent.parent.parent
    / ".claude"
    / "skills"
    / "template-adasa"
)
sys.path.insert(0, skill_path)

from config_defaults import DEFAULTS
from table_utils import set_table_borders, calcular_anchos_columnas, add_simple_table

# set_table_borders importada de table_utils (incluye centrado automatico)

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def aplicar_arial(paragraph, size=11):
    """Aplica formato Arial al parrafo"""
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def crear_correo():
    """Genera el correo en formato DOCX"""

    output_file = "2026-01-29_Respuesta-Cambio-PM-BW-Water.docx"

    doc = Document()

    # Configurar margenes
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ============ HEADER ============
    para = doc.add_paragraph()
    para.add_run("Date: ").bold = True
    para.add_run("January 29, 2026")
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("From: ").bold = True
    para.add_run(
        f"{DEFAULTS['contacto_nombre']} - {DEFAULTS['contacto_cargo']} ({DEFAULTS['contacto_empresa']})"
    )
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("To: ").bold = True
    para.add_run("Dr. Marco Arsovic (BW Water); Eduardo Yamauchi (BW Water)")
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("CC: ").bold = True
    para.add_run(
        "Cesar Malhue, Jorge Guevara, Ronald Pellejero, Victor Gutierrez, Mauricio Vallejos, Jorge Valdes, Sadeep Irugalbandara, Ghazi Ozair, Nick Huta, Andrew Zaske, Fadey Kassim, Elizaveta Bazarova, Tanya Figueroa, Courtney Cooper, Nicole Pearson, Kathryn Peters, Andrew Fuller, Magdier Arias, Andrea Frezzi"
    )
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("Subject: ").bold = True
    para.add_run("RE: Project Taltal - Project Management Point of Contact")
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run("Ref: ").bold = True
    para.add_run("Contract C-4300 / BAE 12803")
    aplicar_arial(para)

    doc.add_paragraph()

    # ============ PART 1: TO MARCO ============
    para = doc.add_paragraph("Dear Marco,")
    aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(
        "Thank you for the notification regarding the PM transition. We acknowledge Eduardo Yamauchi as the new Project Manager for Taltal and appreciate the two-week handover period to ensure continuity."
    )
    aplicar_arial(para)

    doc.add_paragraph()

    # ============ SEPARATOR ============
    # Linea horizontal visual
    para = doc.add_paragraph("_" * 70)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    aplicar_arial(para, 10)

    doc.add_paragraph()

    # ============ PART 2: TO EDUARDO ============
    para = doc.add_paragraph("Dear Eduardo,")
    aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("Welcome to the Taltal project. We look forward to working with you.")
    aplicar_arial(para)

    doc.add_paragraph()

    # Parrafo sobre la situacion - con enfasis en "23 days behind"
    para = doc.add_paragraph()
    para.add_run(
        "We want to bring to your immediate attention that the project's engineering phase is currently "
    )
    run = para.add_run("23 days behind")
    run.bold = True
    para.add_run(
        " the baseline schedule. Yesterday (January 28) we sent Technical Review Transmittal N3 (P22-TM-09-000-003-0) along with a detailed engineering deliverables status and a formal request for a catch-up schedule. That email is attached below for your reference."
    )
    aplicar_arial(para)

    doc.add_paragraph()

    # Parrafo sobre la importancia del recovery plan
    para = doc.add_paragraph()
    para.add_run(
        "The recovery plan requested in that communication is important for both parties to coordinate effectively and protect the contracted milestones. The deadline we proposed for its delivery is "
    )
    run = para.add_run("February 6, 2026")
    run.bold = True
    para.add_run(".")
    aplicar_arial(para)

    doc.add_paragraph()

    # Parrafo de solicitud y ofrecimiento
    para = doc.add_paragraph()
    para.add_run(
        "Given the current delay, we would appreciate if this matter could be prioritized during your transition. We remain available for a coordination call if that would help accelerate alignment on the path forward."
    )
    aplicar_arial(para)

    doc.add_paragraph()

    # ============ SIGNATURE ============
    para = doc.add_paragraph("Best regards,")
    aplicar_arial(para)

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run(DEFAULTS["contacto_nombre"]).bold = True
    aplicar_arial(para)

    para = doc.add_paragraph("Contract Administrator")
    aplicar_arial(para)

    para = doc.add_paragraph("ADASA - Aguas de Antofagasta S.A.")
    aplicar_arial(para)

    para = doc.add_paragraph("Project: BAE 12803 - Second Stage RO Brine Module Taltal")
    aplicar_arial(para)

    doc.add_paragraph()

    # ============ ATTACHMENT REFERENCE ============
    para = doc.add_paragraph()
    para.add_run("Attachment:").bold = True
    aplicar_arial(para)

    para = doc.add_paragraph()
    para.add_run(
        'Email dated January 28, 2026: "Technical Review Transmittal N3 (P22-TM-09-000-003-0) + Engineering Deliverables Status + Catch-Up Schedule Request"'
    )
    aplicar_arial(para, 10)

    # Save document
    doc.save(output_file)
    print(f"Correo generado exitosamente: {output_file}")
    return output_file


if __name__ == "__main__":
    crear_correo()
