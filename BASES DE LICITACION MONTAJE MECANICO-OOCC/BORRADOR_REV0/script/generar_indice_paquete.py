#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera `00_INDICE DEL PAQUETE.xlsx` en la raiz de `INGENIERIA VIGENTE PARA CONSTRUCCION`.

El indice solo ubica documentos: que documentos tiene cada carpeta, con su revision y su
formato. No declara cambios, vigencia ni cantidades, que viven en la Nota Tecnica
P22-NT-06-000-001-0. Un indice que empieza a explicar compite con la nota: asi termino el
00_INDICE.txt que se elimino el 04-09-2026.

Fuente: la lista VIGENCIA de generar_ingenieria_vigente.py y el arbol real del paquete.
Una fila por documento: el PDF y su DWG van en la misma fila, y los anexos de cada ET se
nombran en la fila de la ET. La subcarpeta y el formato salen del arbol, no se escriben
a mano.

Gate: aborta sin escribir si un archivo del paquete queda sin documento o en dos, si un
documento declarado no tiene archivo, o si esta en otra carpeta que la declarada.

Se corre despues de construir_paquete_construccion.py. El .xlsx es derivado y no se
edita a mano.
"""
import math
import re
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from generar_ingenieria_vigente import (
    FUENTE, GRIS, INDICE, PAQUETE, VIGENCIA,
    _laminas_declaradas, encabezado, fila_datos, titulo_hoja,
)

FECHA = "14-09-2026"
NT = ("P22-NT-06-000-001", "-", "Nota técnica de ingeniería vigente para construcción", "0",
      "0. CONTROL DE CAMBIOS")
NO_DOCUMENTO = {"LEEME.txt", INDICE}
ORDEN_FORMATO = ["PDF", "DWG", "XLSX", "NWD", "TIF"]
REV_TEXTO = {"-": "-", "sin sufijo de revision": "Sin sufijo"}


def candidatos(p, docs, por_codigo):
    """Documentos a los que puede pertenecer el archivo p."""
    hits = []
    for codigo, idxs in por_codigo.items():
        if docs[idxs[0]][3] == "-":            # modelo, folletos y sitio: sin revision
            ok = codigo in p.name
        else:                                  # el codigo no puede seguir con un digito
            ok = re.match(re.escape(codigo) + r"(?!\d)", p.stem) is not None
        if not ok:
            continue
        if len(idxs) == 1:
            hits.append(idxs[0])
        else:                                  # un codigo con varias filas: decide la lamina
            marca = p.stem.replace(" ", "")
            hits += [i for i in idxs if any(m in marca for m in _laminas_declaradas(docs[i][1]))]
    return hits


def asignar():
    docs = list(VIGENCIA) + [NT]
    por_codigo = defaultdict(list)
    for i, d in enumerate(docs):
        por_codigo[d[0]].append(i)

    archivos = sorted(p for p in PAQUETE.rglob("*")
                      if p.is_file() and not p.name.startswith(".") and p.name not in NO_DOCUMENTO)
    filas, anexos, errores = defaultdict(list), defaultdict(list), []
    for p in archivos:
        rel = p.relative_to(PAQUETE)
        if p.parent.name == "anexos":
            # El anexo pertenece a la ET cuyo PDF esta en la carpeta de arriba.
            base = [q for q in p.parent.parent.iterdir() if q.is_file()]
            hits = sorted({i for q in base for i in candidatos(q, docs, por_codigo)})
            if len(hits) == 1:
                anexos[hits[0]].append(p)
            else:
                errores.append(f"{rel}: anexo sin ET unica en su carpeta")
            continue
        hits = candidatos(p, docs, por_codigo)
        if len(hits) == 1:
            filas[hits[0]].append(p)
        else:
            errores.append(f"{rel}: " + ("sin documento" if not hits else f"en {len(hits)} documentos"))

    registros = []
    for i, (codigo, lamina, titulo, rev, dossier) in enumerate(docs):
        ps = filas[i]
        if not ps:
            errores.append(f"{codigo} {lamina}: declarado y sin archivo")
            continue
        carpetas = {q.parent for q in ps}
        if len(carpetas) > 1:
            errores.append(f"{codigo} {lamina}: archivos en {len(carpetas)} carpetas")
            continue
        partes = ps[0].parent.relative_to(PAQUETE).parts
        if partes[0] != dossier:
            errores.append(f"{codigo} {lamina}: declarado en {dossier} y esta en {partes[0]}")
            continue

        cod = codigo.rsplit(".", 1)[0] if codigo.lower().endswith(".nwd") else codigo
        if _laminas_declaradas(lamina):
            cod = f"{cod} {lamina}"
        if anexos[i]:
            letras = sorted({re.match(r"Anexo-([A-Z])", q.name).group(1) for q in anexos[i]})
            if len(letras) == 1:
                titulo = f"{titulo}, con anexo {letras[0]}"
            else:
                titulo = f"{titulo}, con anexos {', '.join(letras[:-1])} y {letras[-1]}"
        formatos = sorted({q.suffix[1:].upper() for q in ps},
                          key=lambda f: ORDEN_FORMATO.index(f) if f in ORDEN_FORMATO else 99)
        registros.append((dossier, " / ".join(partes[1:]), i, cod, titulo,
                          REV_TEXTO.get(rev, rev), " + ".join(formatos)))

    registros.sort(key=lambda r: (r[0], r[1], r[2]))
    return registros, errores


def banda(ws, fila, texto, nivel):
    """Fila separadora de carpeta (nivel 1) o de subcarpeta (nivel 2)."""
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
    c = ws.cell(row=fila, column=1, value=texto)
    if nivel == 1:
        c.font = Font(name=FUENTE, size=10, bold=True)
        for col in range(1, 5):
            ws.cell(row=fila, column=col).fill = GRIS
        ws.row_dimensions[fila].height = 18
    else:
        c.font = Font(name=FUENTE, size=9, bold=True, italic=True)
        c.alignment = Alignment(indent=1, vertical="center")
        ws.row_dimensions[fila].height = 15


def escribir(registros):
    wb = Workbook()
    ws = wb.active
    ws.title = "Índice"
    titulo_hoja(ws, "Índice del paquete · Ingeniería vigente para construcción", 4)
    c = ws.cell(row=2, column=1, value=f"{len(registros)} documentos · {FECHA}")
    c.font = Font(name=FUENTE, size=9, italic=True)
    encabezado(ws, 4, ["Código", "Documento", "Rev.", "Formato"], [34, 72, 10, 14])

    fila, carpeta, sub = 5, None, None
    for dossier, s, _i, cod, titulo, rev, fmt in registros:
        if dossier != carpeta:
            banda(ws, fila, dossier, 1)
            fila, carpeta, sub = fila + 1, dossier, None
        if s and s != sub:
            banda(ws, fila, s, 2)
            fila += 1
        sub = s
        fila_datos(ws, fila, (cod, titulo, rev, fmt))
        for col in (3, 4):
            ws.cell(row=fila, column=col).alignment = Alignment(horizontal="center", vertical="top")
        lineas = max(math.ceil(len(titulo) / 85), math.ceil(len(cod) / 38), 1)
        ws.row_dimensions[fila].height = 15 * lineas
        fila += 1

    ws.sheet_view.showGridLines = False
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(PAQUETE / INDICE)


def main():
    registros, errores = asignar()
    if errores:
        print("El indice no coincide con el paquete:")
        for e in errores:
            print("  -", e)
        sys.exit(1)
    escribir(registros)
    total = sum(1 for p in PAQUETE.rglob("*") if p.is_file() and not p.name.startswith("."))
    print(f"OK: {INDICE} con {len(registros)} documentos en "
          f"{len({r[0] for r in registros})} carpetas; el paquete queda con {total} archivos.")


if __name__ == "__main__":
    main()
