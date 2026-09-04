#!/usr/bin/env python3
"""
Apila la revision historica en el cajetin del BL y corrige la revision del encabezado: la skill template-adasa escribe una
sola fila con la revision del frontmatter en `tabla[0].rows[2]`. Este post-proceso deja
la Rev 0 historica con su fecha original en la fila 2 y la Rev 1 nueva en la fila 1,
copiando el formato de la fila existente para no perder la fuente Arial del cajetin.
El encabezado de pagina de la plantilla trae "Revision N: 0" fijo y el conversor no lo
toca, asi que aqui tambien se lleva a la revision vigente.

Uso: apilar_cajetin_rev1.py <docx>
"""
import copy, sys
from docx import Document

REV_HISTORICA = ["0", "25/06/2026", "Luis Rivera", "Luis Rivera", "Victor Gutierrez", "ADASA"]
REV_NUEVA     = ["1", "04/09/2026", "Luis Rivera", "Luis Rivera", "Victor Gutierrez", "ADASA"]


def escribir_fila(fila, valores, modelo):
    """Escribe `valores` en `fila` clonando el formato de la fila `modelo`."""
    for celda, celda_modelo, texto in zip(fila.cells, modelo.cells, valores):
        p_modelo = celda_modelo.paragraphs[0]
        celda._tc.remove(celda.paragraphs[0]._p)
        nuevo = copy.deepcopy(p_modelo._p)
        celda._tc.append(nuevo)
        from docx.text.paragraph import Paragraph
        par = Paragraph(nuevo, celda)
        for r in par.runs[1:]:
            r._r.getparent().remove(r._r)
        if par.runs:
            par.runs[0].text = texto
        else:
            par.add_run(texto)


def corregir_revision_encabezado(doc, revision):
    """La plantilla ADASA trae 'Revision N: 0' fijo en el encabezado de cada pagina."""
    import re
    tocados = 0
    for sec in doc.sections:
        for header in (sec.header, sec.first_page_header, sec.even_page_header):
            for tabla in header.tables:
                for fila in tabla.rows:
                    for celda in fila.cells:
                        for par in celda.paragraphs:
                            if "Revisi" in par.text and "N" in par.text:
                                for run in par.runs:
                                    nuevo = re.sub(r"(Revisi[oó]n\s*N[°º]?:\s*)\S+",
                                                   r"\g<1>" + revision, run.text)
                                    if nuevo != run.text:
                                        run.text = nuevo
                                        tocados += 1
    return tocados


def main(ruta):
    doc = Document(ruta)
    cajetin = doc.tables[0]
    if len(cajetin.rows) < 4:
        sys.exit("El cajetin no tiene la estructura esperada de 4 filas.")
    modelo = cajetin.rows[2]
    escribir_fila(cajetin.rows[1], REV_NUEVA, modelo)
    escribir_fila(cajetin.rows[2], REV_HISTORICA, modelo)
    n = corregir_revision_encabezado(doc, REV_NUEVA[0])
    print(f"Encabezado: {n} referencia(s) de revision actualizada(s) a {REV_NUEVA[0]}")
    doc.save(ruta)
    for i, r in enumerate(Document(ruta).tables[0].rows):
        print(f"  [{i}] " + " | ".join(c.text.strip() for c in r.cells))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else sys.exit("Falta la ruta del .docx"))
